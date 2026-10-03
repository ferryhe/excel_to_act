"""What-if data tables: geometry read from XML, corner cell not corrupted."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import DataTableFormula

from excel_to_act.ingest.data_table import read_data_tables, sheet_xml_parts
from excel_to_act.ingest.openpyxl_reader import OpenpyxlWorkbookReader
from excel_to_act.inventory.extractor import OpenpyxlInventoryExtractor


def build_fixture(path: Path) -> Path:
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    inputs["A2"] = "rate"
    inputs["B2"] = 0.05
    inputs["A3"] = "premium"
    inputs["B3"] = 100
    scenarios = wb.create_sheet("Scenarios")
    scenarios["B7"] = 1
    scenarios["C7"] = 2
    scenarios["B8"] = 3
    wb.save(path)
    wb.close()
    return path


def inject_data_table(path: Path, sheet_name: str, row_xml: str) -> None:
    """Insert a ``<row>`` carrying a ``<f t="dataTable">`` cell into a sheet part."""
    part = sheet_xml_parts(path)[sheet_name]
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        payload = {name: zf.read(name) for name in names}
    xml = payload[part].decode("utf-8")
    assert "</sheetData>" in xml
    payload[part] = xml.replace("</sheetData>", row_xml + "</sheetData>").encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in names:
            zf.writestr(name, payload[name])


def _inventory(path: Path):
    manifest = OpenpyxlWorkbookReader().read_manifest(path)
    return OpenpyxlInventoryExtractor().extract(path, manifest)


def _data_tables(inventory):
    sheet = next(s for s in inventory.sheets if s.name == "Scenarios")
    return [r for r in sheet.ranges if r.kind == "data_table"]


def test_one_dimensional_data_table(tmp_path: Path) -> None:
    fixture = build_fixture(tmp_path / "one_dim.xlsx")
    inject_data_table(
        fixture,
        "Scenarios",
        '<row r="7"><c r="B7"><f t="dataTable" ref="B7:C9" dt2D="0" dtr="0" r1="B2">=B2*B3</f><v></v></c></row>',
    )

    spec = read_data_tables(fixture)["Scenarios"][0]
    assert spec.corner_cell == "B7"
    assert spec.ref == "B7:C9"
    assert spec.row_input_cell is None
    assert spec.col_input_cell == "$B$2"
    assert spec.two_dimensional is False

    inventory = _inventory(fixture)
    tables = _data_tables(inventory)
    assert len(tables) == 1
    assert tables[0].address == "B7:C9"
    assert tables[0].metadata["row_input_cell"] is None
    assert tables[0].metadata["col_input_cell"] == "$B$2"
    assert tables[0].metadata["formula"] == "=B2*B3"

    # The corner cell must survive as a formula, not as str(DataTableFormula).
    corner = next(c for c in inventory.sheets[1].cells if c.address == "B7")
    assert corner.kind == "formula"
    assert corner.formula == "=B2*B3"
    assert corner.value is None


def test_one_dimensional_row_oriented_data_table(tmp_path: Path) -> None:
    fixture = build_fixture(tmp_path / "one_dim_row.xlsx")
    inject_data_table(
        fixture,
        "Scenarios",
        '<row r="7"><c r="B7"><f t="dataTable" ref="B7:C9" dt2D="0" dtr="1" r1="B2">=B2*B3</f><v></v></c></row>',
    )

    table = read_data_tables(fixture)["Scenarios"][0]
    assert table.row_input_cell == "$B$2"
    assert table.col_input_cell is None


def test_two_dimensional_data_table(tmp_path: Path) -> None:
    fixture = build_fixture(tmp_path / "two_dim.xlsx")
    inject_data_table(
        fixture,
        "Scenarios",
        '<row r="7"><c r="B7"><f t="dataTable" ref="B7:C9" dt2D="1" dtr="0" r1="B2" r2="D3">=B2*B3</f><v></v></c></row>',
    )
    inventory = _inventory(fixture)
    table = _data_tables(inventory)[0]
    assert table.metadata["two_dimensional"] is True
    assert table.metadata["row_input_cell"] == "$B$2"
    assert table.metadata["col_input_cell"] == "$D$3"


def test_workbook_without_data_tables_is_untouched(tmp_path: Path) -> None:
    fixture = build_fixture(tmp_path / "none.xlsx")
    inventory = _inventory(fixture)
    assert _data_tables(inventory) == []
    assert read_data_tables(fixture) == {}
    assert (
        inventory.coverage.recognized_inventory_objects + inventory.coverage.unsupported_or_opaque_objects
        == inventory.coverage.discovered_workbook_objects
    )


def test_openpyxl_exposes_data_table_as_object_hence_xml_read(tmp_path: Path) -> None:
    """Documents why the geometry is read from XML instead of openpyxl."""
    fixture = build_fixture(tmp_path / "raw.xlsx")
    inject_data_table(
        fixture,
        "Scenarios",
        '<row r="7"><c r="B7"><f t="dataTable" ref="B7:C9" dt2D="0" dtr="0" r1="B2">=B2*B3</f><v></v></c></row>',
    )
    wb = load_workbook(fixture)
    try:
        assert isinstance(wb["Scenarios"]["B7"].value, DataTableFormula)
        assert "=B2*B3" not in str(wb["Scenarios"]["B7"].value)  # text is dropped by openpyxl
    finally:
        wb.close()
