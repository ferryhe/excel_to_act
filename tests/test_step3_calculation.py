from __future__ import annotations

import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula
from typer.testing import CliRunner

from excel_to_act.interfaces.cli import app
from excel_to_act.steps.step3 import calculation
from excel_to_act.steps.step3.calculation import CalculationBlocked, ExcelError, FormulaEvaluator


def _formula_workbook(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["A10"], sheet["B10"] = "key", "value"
    sheet["A11"], sheet["B11"] = 1, 10
    sheet["A12"], sheet["B12"] = 1, 20
    sheet["A13"], sheet["B13"] = 3, None
    sheet["A14"], sheet["B14"] = 4, 40
    sheet["D1"] = "=IF(FALSE,1/0,7)"
    sheet["D2"] = "=IFERROR(1/0,3)"
    sheet["D3"] = "=IFERROR(UNSUPPORTED(),3)"
    sheet["D14"] = "=IFNA(#N/A,4)"
    sheet["D15"] = "=IFERROR('[External.xlsx]Sheet1'!A1,3)"
    sheet["D4"] = "=VLOOKUP(1,A10:B14,2,FALSE)"
    sheet["D5"] = "=VLOOKUP(1,A10:B14,2,TRUE)"
    sheet["D6"] = "=VLOOKUP(2,A10:B14,2,TRUE)"
    sheet["D7"] = "=VLOOKUP(3,A10:B14,2,TRUE)"
    sheet["D8"] = "=VLOOKUP(0,A10:B14,2,TRUE)"
    sheet["D9"] = "=VLOOKUP(9,A10:B14,2,FALSE)"
    sheet["D10"] = "=MATCH(1,A11:A12,0)"
    sheet["D11"] = "=SUMPRODUCT((A11:A14>1)*(B11:B14))"
    sheet["D12"] = "=SUM(A11:OFFSET(A11,2,0))"
    sheet["D13"] = '=INDIRECT("B11")'
    sheet["E20"], sheet["E21"] = 5, 6
    sheet["G20"] = ArrayFormula(ref="G20:H20", text="=TRANSPOSE(E20:E21)")
    sheet["A30"] = "=A30+1"

    premium = workbook.create_sheet("Premium")
    premium["H1"] = "=2+3"
    premium["H2"] = "=1"
    premium["H3"] = "=6"
    premium["J1"] = "=H3/(H1-H2)"
    main = workbook["Main"]
    for cell, value in {"C3": 0, "C4": "Male", "C6": 30, "C7": 10, "C8": 1000,
                        "J48": 0.035, "O26": "Yes"}.items():
        main[cell] = value
    workbook.save(path)
    return path


def _baseline_for(path: Path) -> Path:
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    baseline = {
        "schema_version": "gp.excel_oracle.v1",
        "tool": "step5.oracle",
        "status": "pass",
        "engine": "Microsoft Excel",
        "source_sha256": source_hash,
        "source_copy_sha256": source_hash,
        "macros_executed": False,
        "input_overrides": [],
        "primary_inputs": {
            "C3": {"value": 0}, "C4": {"value": "Male"}, "C6": {"value": 30},
            "C7": {"value": 10}, "C8": {"value": 1000}, "J48": {"value": 0.035},
            "O26": {"value": "Yes"},
        },
        "after_full_rebuild": {
            "H1": {"value": 5, "formula": "=2+3"},
            "H2": {"value": 1, "formula": "=1"},
            "H3": {"value": 6, "formula": "=6"},
            "J1": {"value": 1.5, "formula": "=H3/(H1-H2)"},
        },
        "premium_values_B9_CD115": [[None] * 81 for _ in range(107)],
        "premium_formulas_B9_CD115": [[""] * 81 for _ in range(107)],
    }
    baseline_dir = path.parent / "baseline"
    baseline_dir.mkdir()
    baseline_path = baseline_dir / "excel_oracle.json"
    baseline_path.write_text(json.dumps(baseline), encoding="utf-8")
    return baseline_path


def test_lazy_branches_lookup_modes_arrays_and_source_errors(tmp_path: Path) -> None:
    source = _formula_workbook(tmp_path / "formula.xlsx")
    evaluator = FormulaEvaluator(source)
    try:
        assert evaluator.evaluate_formula("=IF(FALSE,1/0,7)", "Main", "test") == 7
        assert evaluator.evaluate_formula("=IFERROR(1/0,3)", "Main", "test") == 3
        assert evaluator.evaluate_formula("=IFNA(#N/A,4)", "Main", "test") == 4
        assert evaluator.evaluate_formula("=VLOOKUP(1,A10:B14,2,FALSE)", "Main", "test") == 10
        assert evaluator.evaluate_formula("=VLOOKUP(1,A10:B14,2,TRUE)", "Main", "test") == 20
        assert evaluator.evaluate_formula("=VLOOKUP(2,A10:B14,2,TRUE)", "Main", "test") == 20
        assert evaluator.evaluate_formula("=VLOOKUP(3,A10:B14,2,TRUE)", "Main", "test") == 0
        assert evaluator.evaluate_formula("=VLOOKUP(0,A10:B14,2,TRUE)", "Main", "test") == ExcelError("#N/A")
        assert evaluator.evaluate_formula("=VLOOKUP(9,A10:B14,2,FALSE)", "Main", "test") == ExcelError("#N/A")
        assert evaluator.evaluate_formula("=MATCH(1,A11:A12,0)", "Main", "test") == 1
        assert evaluator.evaluate_formula("=SUMPRODUCT((A11:A14>1)*(B11:B14))", "Main", "test") == 40
        assert evaluator.evaluate_formula("=SUM(A11:OFFSET(A11,2,0))", "Main", "test") == 5
        assert evaluator.evaluate_formula('=INDIRECT("B11")', "Main", "test") == 10
        assert evaluator.evaluate_cell("Main", 20, 8, dependency=False) == 6
        assert evaluator.evaluate_cell("Main", 20, 7, dependency=False) == 5
        assert any(edge["consumer"] == "Main!H20" and edge["prerequisite"] == "Main!G20"
                   and edge["relationship"] == "array_formula_anchor" for edge in evaluator.edges.values())
        assert evaluator.evaluate_cell("Main", 30, 1, dependency=False) is not None
    except CalculationBlocked as exc:
        assert "active formula cycle" in str(exc)
    else:
        raise AssertionError("a direct active self-cycle must block")
    finally:
        evaluator.close()

    blocked = FormulaEvaluator(source)
    try:
        try:
            blocked.evaluate_formula("=IFERROR(UNSUPPORTED(),3)", "Main", "test")
        except CalculationBlocked as exc:
            assert "unsupported active function" in str(exc)
        else:
            raise AssertionError("IFERROR must not catch an evaluator capability blocker")
    finally:
        blocked.close()

    external = FormulaEvaluator(source)
    try:
        assert external.evaluate_formula("=IF(FALSE,'[External.xlsx]Sheet1'!A1,7)", "Main", "test") == 7
        for formula in ("=IF(TRUE,'[External.xlsx]Sheet1'!A1,7)",
                        "=IFERROR('[External.xlsx]Sheet1'!A1,3)"):
            try:
                external.evaluate_formula(formula, "Main", "test")
            except CalculationBlocked as exc:
                assert "external reference syntax" in str(exc)
            else:
                raise AssertionError(f"an active external dependency must block: {formula}")
    finally:
        external.close()


def test_external_cut_override_stops_formula_traversal_and_keeps_edge(tmp_path: Path) -> None:
    source = _formula_workbook(tmp_path / "cut.xlsx")
    from openpyxl import load_workbook
    workbook = load_workbook(source, data_only=False)
    workbook["Main"]["E1"] = "=E2"
    workbook["Main"]["E2"] = "=E2+1"
    workbook["Main"]["F1"] = "=E1+2"
    workbook.save(source)
    workbook.close()
    formula = "=E2"
    formula_hash = hashlib.sha256(formula.encode("utf-8")).hexdigest()
    evaluator = FormulaEvaluator(source, external_overrides={
        "Main!E1": {"boundary_variable_id": "external.fixture", "value": 11,
                    "source_formula": formula, "source_formula_sha256": formula_hash,
                    "capture_artifact_sha256": "a" * 64},
    })
    try:
        assert evaluator.evaluate_cell("Main", 1, 6, dependency=False) == 13
        assert "Main!E2" not in evaluator.cells
        assert evaluator.cells["Main!E1"]["role"] == "external_boundary_input"
        assert evaluator.cells["Main!E1"]["formula_sha256"] == formula_hash
        assert any(edge["consumer"] == "Main!F1" and edge["prerequisite"] == "Main!E1"
                   for edge in evaluator.edges.values())
    finally:
        evaluator.close()


def test_external_cut_override_rejects_array_overlap_and_stale_formula(tmp_path: Path) -> None:
    source = _formula_workbook(tmp_path / "cut-array.xlsx")
    formula = "=TRANSPOSE(E20:E21)"
    formula_hash = hashlib.sha256(formula.encode("utf-8")).hexdigest()
    evaluator = FormulaEvaluator(source, external_overrides={
        "Main!H20": {"boundary_variable_id": "external.fixture", "value": 1,
                     "source_formula": formula, "source_formula_sha256": formula_hash,
                     "capture_artifact_sha256": "b" * 64},
    })
    try:
        try:
            evaluator.evaluate_cell("Main", 20, 8, dependency=False)
        except CalculationBlocked as exc:
            assert "overlaps a source array formula" in str(exc)
        else:
            raise AssertionError("an external cut must not split a source array formula")
    finally:
        evaluator.close()

    evaluator = FormulaEvaluator(source, external_overrides={
        "Main!D1": {"boundary_variable_id": "external.fixture", "value": 1,
                    "source_formula": "=1", "source_formula_sha256": hashlib.sha256(b"=1").hexdigest(),
                    "capture_artifact_sha256": "c" * 64},
    })
    try:
        try:
            evaluator.evaluate_cell("Main", 1, 4, dependency=False)
        except CalculationBlocked as exc:
            assert "source formula changed" in str(exc)
        else:
            raise AssertionError("a cut override must verify current formula text")
    finally:
        evaluator.close()


def test_gp_baseline_binding_reconciliation_and_cli_artifacts(tmp_path: Path) -> None:
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    source = _formula_workbook(source_dir / "workbook.xlsx")
    baseline = _baseline_for(source)
    out = tmp_path / "runtime"

    result = calculation.calculate_gp(source, baseline, out)
    assert result["status"] == "pass"
    assert result["computed"] == {"AnnuityDue": 5, "PVLoading": 1, "PVFB": 6, "GP": 1.5}
    assert result["reconciliation"]["compared_premium_formulas"] == 0
    for artifact in ("gp_result.json", "gp_result.md", "active_trace.json", "active_trace.md"):
        assert (out / artifact).is_file()
    saved = json.loads((out / "gp_result.json").read_text(encoding="utf-8"))
    assert saved["status"] == "pass"
    assert "formula caches are never loaded" in saved["formula_cache_policy"]
    trace = json.loads((out / "active_trace.json").read_text(encoding="utf-8"))
    completion = trace["completion_order"]
    assert max(completion.index(f"Premium!{cell}") for cell in ("H1", "H2", "H3")) < completion.index("Premium!J1")
    assert any(edge["consumer"] == "Premium!J1" and edge["prerequisite"] == "Premium!H1"
               for edge in trace["edges"])
    assert all(item["consumer"].startswith("Premium!") or item["consumer"].startswith("Main!")
               for item in trace["edges"])

    removed_gp_command = CliRunner().invoke(app, ["step3", "gp", "--workbook", str(source), "--baseline", str(baseline),
                                                  "--out", str(tmp_path / "cli-runtime")])
    removed_oracle_command = CliRunner().invoke(app, ["step3", "oracle", "--workbook", str(source), "--out", str(tmp_path / "oracle")])
    assert removed_gp_command.exit_code != 0
    assert removed_oracle_command.exit_code != 0
    assert CliRunner().invoke(app, ["step4", "tools"]).exit_code == 0
    assert CliRunner().invoke(app, ["step5", "tools"]).exit_code == 0


def test_stale_baseline_does_not_overwrite_existing_result(tmp_path: Path) -> None:
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    source = _formula_workbook(source_dir / "workbook.xlsx")
    baseline = _baseline_for(source)
    baseline_data = json.loads(baseline.read_text(encoding="utf-8"))
    baseline_data["source_sha256"] = "0" * 64
    baseline.write_text(json.dumps(baseline_data), encoding="utf-8")
    out = tmp_path / "runtime"
    out.mkdir()
    prior = out / "gp_result.json"
    prior.write_text('{"status":"previous"}', encoding="utf-8")

    result = calculation.calculate_gp(source, baseline, out)
    assert result["status"] == "blocked"
    assert "source hash" in result["reason"]
    assert prior.read_text(encoding="utf-8") == '{"status":"previous"}'


def test_unsuccessful_or_incomplete_oracle_baselines_do_not_overwrite_result(tmp_path: Path) -> None:
    cases = [
        ("blocked", lambda data: data.update(status="blocked"), "successful"),
        ("missing-copy-hash", lambda data: data.pop("source_copy_sha256"), "source-copy hash"),
        ("incomplete-inputs", lambda data: data["primary_inputs"].pop("O26"), "all seven saved primary inputs"),
        ("input-without-value", lambda data: data["primary_inputs"].update(O26={}),
         "all seven saved primary inputs"),
    ]
    for index, (case, mutate, expected_reason) in enumerate(cases):
        source_dir = tmp_path / f"source-{index}"
        source_dir.mkdir()
        source = _formula_workbook(source_dir / "workbook.xlsx")
        baseline = _baseline_for(source)
        baseline_data = json.loads(baseline.read_text(encoding="utf-8"))
        mutate(baseline_data)
        baseline.write_text(json.dumps(baseline_data), encoding="utf-8")
        out = tmp_path / f"runtime-{case}"
        out.mkdir()
        prior = out / "gp_result.json"
        prior.write_text('{"status":"previous"}', encoding="utf-8")

        result = calculation.calculate_gp(source, baseline, out)

        assert result["status"] == "blocked", case
        assert expected_reason in result["reason"], case
        assert prior.read_text(encoding="utf-8") == '{"status":"previous"}', case


def test_native_range_error_codes_are_compared_as_excel_errors() -> None:
    values = [[None] * 81 for _ in range(107)]
    formulas = [[""] * 81 for _ in range(107)]
    errors = [[None] * 81 for _ in range(107)]
    errors[0][0] = "#N/A"
    baseline = {
        "premium_values_B9_CD115": values,
        "premium_formulas_B9_CD115": formulas,
        "ranges": {
            "Premium!B9:CD115": {"values": values, "formulas": formulas, "excel_errors": errors}
        },
    }
    actual_values, actual_formulas = calculation._baseline_matrices(baseline)
    assert actual_values[0][0] == {"excel_error": "#N/A"}
    assert actual_formulas == formulas


def test_oracle_requires_a_bounded_request_and_packaged_script(tmp_path: Path) -> None:
    from importlib.resources import files

    script = files("excel_to_act.steps.step5").joinpath("native_excel_oracle.ps1")
    assert script.is_file()
    result = calculation.capture_excel_oracle(Path(__file__), tmp_path / "oracle", targets=[], ranges=[])
    assert result["status"] == "blocked"
    assert "at least one" in result["reason"]


def test_native_oracle_resolves_requested_sheet_names_case_insensitively(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "AMR_Table"
    workbook.create_sheet("Premium")
    workbook.create_sheet("Qtable")
    for name, reference in {
        "RowVector": "'AMR_Table'!$A$1:$B$1",
        "ColumnVector": "'AMR_Table'!$D$1:$D$2",
        "SingletonVector": "'AMR_Table'!$F$1",
        "Matrix": "'AMR_Table'!$A$4:$B$5",
        "RowTable": "'AMR_Table'!$D$4:$E$4",
        "ColumnTable": "'AMR_Table'!$G$4:$G$5",
    }.items():
        workbook.defined_names.add(DefinedName(name, attr_text=reference))
    workbook.save(source)
    workbook.close()

    calls = []
    requests = []
    target_shapes = {
        "RowVector": [2], "ColumnVector": [2], "SingletonVector": [1],
        "Matrix": [2, 2], "RowTable": [1, 2], "ColumnTable": [2, 1],
    }

    def fake_run(args, **_kwargs):
        calls.append(args)
        request_path = Path(args[args.index("-RequestPath") + 1])
        output_path = Path(args[args.index("-OutputPath") + 1])
        request = json.loads(request_path.read_text(encoding="utf-8"))
        requests.append(request)
        ranges = {}
        for record in request["ranges"]:
            start, *end = record["address"].split(":")
            from openpyxl.utils.cell import range_boundaries

            min_col, min_row, max_col, max_row = range_boundaries(
                start if not end else f"{start}:{end[0]}"
            )
            shape = [[None for _ in range(max_col - min_col + 1)]
                     for _ in range(max_row - min_row + 1)]
            ranges[record["qualified_address"]] = {
                "values": shape,
                "formulas": [["" for _ in row] for row in shape],
                "excel_errors": [[None for _ in row] for row in shape],
            }
        output_path.write_text(json.dumps({
            "engine": "Microsoft Excel",
            "engine_settings": {"version": "16.0"},
            "full_rebuild_seconds": 0.1,
            "named_targets": {},
            "after_full_rebuild": {},
            "ranges": ranges,
        }), encoding="utf-8")
        return SimpleNamespace(returncode=0, stderr="", stdout="")

    monkeypatch.setattr(calculation.os, "name", "nt")
    monkeypatch.setattr(calculation.shutil, "which", lambda _name: "powershell.exe")
    monkeypatch.setattr(calculation.subprocess, "run", fake_run)
    oracle_dir = tmp_path.parent / f"{tmp_path.name}-oracle"
    result = calculation.capture_excel_oracle(
        source, oracle_dir,
        ranges=["amr_table!A1:B2", "premium!C3", "qtable!D5:D6"],
        targets=list(target_shapes), target_shapes=target_shapes,
    )

    assert result["status"] == "pass", result
    saved = json.loads(Path(result["artifacts"][0]).read_text(encoding="utf-8"))
    assert saved["requested_ranges"] == ["AMR_Table!A1:B2", "Premium!C3", "Qtable!D5:D6"]
    assert set(saved["ranges"]) == set(saved["requested_ranges"])
    assert requests[0]["target_shapes"] == target_shapes
    assert len(calls) == 1

    missing_sheet = calculation.capture_excel_oracle(
        source, tmp_path.parent / f"{tmp_path.name}-missing-sheet", targets=[], ranges=["not_a_sheet!A1"],
    )
    assert missing_sheet["status"] == "blocked"
    assert missing_sheet["reason"] == "worksheet does not exist: not_a_sheet"
    assert len(calls) == 1
