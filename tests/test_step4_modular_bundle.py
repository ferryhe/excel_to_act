from __future__ import annotations

import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import types
from concurrent.futures import ThreadPoolExecutor

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.formula import ArrayFormula

from excel_to_act.steps.step3.calculation import CalculationBlocked
from excel_to_act.steps.step4.generator import _bind_current_support_domain, _bundle_files, _smoke_bundle
from excel_to_act.steps.step4.modular_bundle import (
    _augment_literal_sources,
    _canonical_hash,
    _named_bindings,
    _raw_bindings,
    _variable_contract_layout,
    _model_source,
    add_adapter_metadata,
    build_modular_bundle_files,
)
from excel_to_act.steps.step4.modular_compiler import build_source_cell_index, compile_formula_function
from excel_to_act.steps.step4.modular_runtime import binary
from excel_to_act.steps.step4 import modular_runtime as runtime


def _workbook(path: Path) -> str:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    for row, key in enumerate((1, 2, 3), 1):
        sheet.cell(row, 1).value = key
        sheet.cell(row, 2).value = key + 10
        sheet.cell(row, 3).value = f"=A{row}*10"
    sheet["D1"] = "=VLOOKUP(2,RawTable,3,FALSE)"
    sheet["F1"] = "No"
    sheet["E1"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    sheet["G1"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    workbook.save(path)
    workbook.close()
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _input_records():
    return [
        {"variable_id": "raw.key", "role": "raw", "kind": "projection_series", "shape": [3],
         "axes": [{"name": "year"}], "dependencies": [], "initial_condition": None,
         "source_extents": [{"ref": "Main!A1:A3", "role": "raw_input",
             "coordinate_mapping": {"kind": "row_ordinal", "axes": [{"axis": "year", "dimension": 0,
                 "source_coordinate": "row", "origin": 1, "index_origin": 0}]}}]},
        {"variable_id": "raw.value", "role": "raw", "kind": "projection_series", "shape": [3],
         "axes": [{"name": "year"}], "dependencies": [], "initial_condition": None,
         "source_extents": [{"ref": "Main!B1:B3", "role": "raw_input",
             "coordinate_mapping": {"kind": "row_ordinal", "axes": [{"axis": "year", "dimension": 0,
                 "source_coordinate": "row", "origin": 1, "index_origin": 0}]}}]},
        {"variable_id": "raw.config", "role": "raw", "kind": "scalar", "shape": [],
         "axes": [], "dependencies": [], "initial_condition": None,
         "source_extents": [{"ref": "Main!F1", "role": "raw_input",
             "coordinate_mapping": {"kind": "scalar", "fixed_axes": []}}]},
        {"variable_id": "table.value", "role": "derived", "kind": "projection_series", "shape": [3],
         "axes": [{"name": "year"}], "dependencies": [{"variable_id": "raw.key"}],
         "initial_condition": None,
         "source_extents": [{"ref": "Main!C1:C3", "role": "formula_output",
             "coordinate_mapping": {"kind": "row_ordinal", "axes": [{"axis": "year", "dimension": 0,
                 "source_coordinate": "row", "origin": 1, "index_origin": 0}]}}],
         "equation_segments": [{"segment_id": "value_pass", "source_family_ids": ["family-value"],
             "equation_family_ids": ["eq-value"], "index_axis": "year", "index_start": 0,
             "index_stop_exclusive": 3, "execution": "vectorized", "recurrence_group": None,
             "snapshot_before_update": False}]},
        {"variable_id": "model.answer", "role": "derived", "kind": "scalar", "shape": [],
         "axes": [], "dependencies": [{"variable_id": "raw.key"}, {"variable_id": "table.value"}],
         "initial_condition": None,
         "source_extents": [{"ref": "Main!D1", "role": "formula_output",
             "coordinate_mapping": {"kind": "scalar", "fixed_axes": []}}],
         "equation_segments": [{"segment_id": "answer_pass", "source_family_ids": ["family-answer"],
             "equation_family_ids": ["eq-answer"], "index_axis": None, "index_start": 0,
             "index_stop_exclusive": 1, "execution": "scalar", "recurrence_group": None,
             "snapshot_before_update": False}]},
    ]


def _variable_interface(source_hash: str) -> dict:
    return {
        "schema_version": "step3.tool_interface.v1",
        "contract_readiness": {"status": "defined", "blockers": []},
        "bindings": {"source": {"workbook_sha256": source_hash}, "design": {}, "input_boundary": {}},
        "support_domain": {"status": "fixed_trace", "active_trace_sha256": "t" * 64,
                           "external_policy": {}, "tolerances": {"absolute": 0, "relative": 0}},
        "objects": [{"variable_id": variable_id, "logical_name": name,
                     "mutability": "mutable" if variable_id in {"raw.value", "raw.config"} else "immutable",
                     "shape": shape, "axes": axes, "value_type": value_type,
                     "required": variable_id != "raw.config",
                     "default_policy": {"status": "saved" if variable_id == "raw.config" else "none"},
                     "allowed_range": allowed_range}
                    for variable_id, name, shape, axes, value_type, allowed_range in (
                        ("raw.key", "Key", [3], [{"name": "year"}], "integer", {"status": "unbounded"}),
                        ("raw.value", "Value", [3], [{"name": "year", "keys": [1, 2, 3]}], "number",
                         {"status": "declared", "minimum": 0, "maximum": 25}),
                        ("raw.config", "Config", [], [], "integer", {"status": "enum", "values": [2, 5]}),
                    )],
        "source_provenance": {"inputs": {variable_id: {"role": "source_raw"}
                                         for variable_id in ("raw.key", "raw.value", "raw.config")},
                              "outputs": {"Answer": {"selector": "Answer"}}},
        "outputs": {"properties": {"Answer": {"x-axes": []}}},
    }


def test_post_discovery_support_binds_current_artifacts_and_rejects_stale_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from excel_to_act.steps.conversion_workflow import json_bytes

    source_hash, design_hash = "a" * 64, "d" * 64
    definition = _variable_interface(source_hash)
    definition["bindings"]["design"] = {
        "artifact": {"json_sha256": design_hash},
        "review_receipts": {name: {"decision": "approve", "artifact_sha256": design_hash}
                            for name in ("agent", "human")}}
    definition["support_domain"] = {
        "status": "exact_cases", "trace_binding": "post_discovery",
        "cases": [{name: {"path": "case.json", "sha256": "c" * 64} for name in
                   ("trace", "raw_request", "public_request", "producer")}],
        "external_policy": {}, "tolerances": {"absolute": 0, "relative": 0}}
    interface_path = tmp_path / "interface.json"
    interface_path.write_bytes(json_bytes(definition))
    capture_path = tmp_path / "external_inputs.json"
    capture_path.write_bytes(json_bytes({"schema_version": "step4.external_inputs.v1",
        "status": "pass", "source_sha256": source_hash,
        "stage3": {"artifact_json_sha256": design_hash}}))
    capture_hash = hashlib.sha256(capture_path.read_bytes()).hexdigest()
    trace_path = tmp_path / "active_trace.json"
    trace_path.write_bytes(json_bytes({"source_sha256": source_hash,
        "design_sha256": design_hash, "external_capture_sha256": capture_hash, "cells": []}))
    trace_hash = hashlib.sha256(trace_path.read_bytes()).hexdigest()
    discovery_path = tmp_path / "discovery.json"
    discovery_path.write_bytes(json_bytes({"schema_version": "step4.discovery.v1",
        "status": "pass", "source_sha256": source_hash,
        "stage3_artifact_sha256": design_hash, "active_trace_sha256": trace_hash,
        "active_trace_path": str(trace_path)}))
    kwargs = {"interface_path": interface_path, "design_sha256": design_hash,
              "source_sha256": source_hash, "capture_path": capture_path,
              "capture_sha256": capture_hash, "discovery_path": discovery_path,
              "trace_path": trace_path}
    bound = _bind_current_support_domain(definition, **kwargs)
    assert bound["support_domain"]["active_trace_sha256"] == trace_hash
    assert bound["support_certificate"]["capture_artifact_sha256"] == capture_hash
    assert bound["support_certificate"]["approved_interface_sha256"] == hashlib.sha256(
        interface_path.read_bytes()).hexdigest()
    assert "active_trace_sha256" not in definition["support_domain"]
    capture_path.write_bytes(json_bytes({"schema_version": "step4.external_inputs.v1",
        "status": "pass", "source_sha256": source_hash,
        "stage3": {"artifact_json_sha256": design_hash}, "dependency_version": "new"}))
    refreshed_hash = hashlib.sha256(capture_path.read_bytes()).hexdigest()
    with pytest.raises(CalculationBlocked, match="capture.*trace"):
        _bind_current_support_domain(definition, **{**kwargs, "capture_sha256": refreshed_hash})
    capture_path.write_bytes(json_bytes({"schema_version": "step4.external_inputs.v1",
        "status": "pass", "source_sha256": source_hash,
        "stage3": {"artifact_json_sha256": design_hash}}))
    blocked = json.loads(json.dumps(definition))
    blocked["contract_readiness"]["status"] = "blocked"
    with pytest.raises(CalculationBlocked, match="contract is blocked"):
        _bind_current_support_domain(blocked, **kwargs)
    source_path = tmp_path / "source.xlsx"
    source_path.write_bytes(b"synthetic source")
    bindings = [{"address": f"main!B{row}", "variable_id": "raw.value", "indices": [row - 1]}
                for row in range(1, 4)]
    bindings.append({"address": "main!F1", "variable_id": "raw.config", "indices": []})
    saved = {f"main!B{row}": row + 10 for row in range(1, 4)} | {"main!F1": 2}
    proofs = [{"input_sha256": "i" * 64, "dependency_sha256": "e" * 64}]
    monkeypatch.setattr("excel_to_act.steps.step4.modular_bundle._verified_case_layout",
                        lambda *_args: proofs)
    contract = _variable_contract_layout(bound, bindings, [], trace_hash, {}, saved,
        source_path=source_path, semantic_plan={}, source_profile={}, target_names=["Answer"],
        names={}, capture_artifact_sha256=capture_hash)
    assert contract["support_certificate"]["verified_cases_sha256"] == _canonical_hash(proofs)
    with pytest.raises(CalculationBlocked, match="certificate"):
        _variable_contract_layout(bound, bindings, [], trace_hash, {}, saved,
            source_path=source_path, semantic_plan={}, source_profile={}, target_names=["Answer"],
            names={}, capture_artifact_sha256="0" * 64)
    stale_case = json.loads(json.dumps(bound))
    stale_case["support_domain"]["cases"][0]["producer"]["sha256"] = "0" * 64
    with pytest.raises(CalculationBlocked, match="certificate"):
        _variable_contract_layout(stale_case, bindings, [], trace_hash, {}, saved,
            source_path=source_path, semantic_plan={}, source_profile={}, target_names=["Answer"],
            names={}, capture_artifact_sha256=capture_hash)
    def missing_closure(*_args: object) -> list[dict[str, str]]:
        raise CalculationBlocked("exact-case source cells closure differs from trace")
    monkeypatch.setattr("excel_to_act.steps.step4.modular_bundle._verified_case_layout", missing_closure)
    with pytest.raises(CalculationBlocked, match="closure differs"):
        _variable_contract_layout(bound, bindings, [], trace_hash, {}, saved,
            source_path=source_path, semantic_plan={}, source_profile={}, target_names=["Answer"],
            names={}, capture_artifact_sha256=capture_hash)
    for changed in ("design_sha256", "source_sha256", "capture_sha256"):
        with pytest.raises(CalculationBlocked, match="stale|current"):
            _bind_current_support_domain(definition, **{**kwargs, changed: "0" * 64})
    discovery_path.write_text("{}", encoding="utf-8")
    with pytest.raises(CalculationBlocked, match="discovery|current"):
        _bind_current_support_domain(definition, **kwargs)
    discovery_path.write_bytes(json_bytes({"schema_version": "step4.discovery.v1",
        "status": "pass", "source_sha256": source_hash,
        "stage3_artifact_sha256": design_hash, "active_trace_sha256": trace_hash,
        "active_trace_path": str(trace_path)}))
    trace_path.write_text("changed", encoding="utf-8")
    with pytest.raises(CalculationBlocked, match="discovery|trace"):
        _bind_current_support_domain(definition, **kwargs)


@pytest.mark.parametrize("exact_cases", [False, True])
def test_empty_interface_registry_never_silently_drops_exact_cases(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, exact_cases: bool,
) -> None:
    from excel_to_act.steps.conversion_workflow import json_bytes

    def evidence(name: str, payload: dict) -> tuple[Path, str]:
        path = tmp_path / name
        path.write_bytes(json_bytes(payload))
        return path, hashlib.sha256(path.read_bytes()).hexdigest()

    source_path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Main"
    workbook.active["A1"] = 3
    workbook.save(source_path)
    workbook.close()
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    design_hash = "d" * 64
    trace_path, trace_hash = evidence("active_trace.json", {
        "source_sha256": source_hash, "design_sha256": design_hash,
        "cells": [{"address": "Main!A1", "sheet": "Main", "role": "source_value"}],
        "names": [{"scope": "workbook", "name": "Result"}],
    })
    evidence("discovery.json", {"schema_version": "step4.discovery.v1", "status": "pass",
        "source_sha256": source_hash, "stage3_artifact_sha256": design_hash,
        "active_trace_path": str(trace_path), "active_trace_sha256": trace_hash})
    map_path, map_hash = evidence("map.json", {})
    profile_path, profile_hash = evidence("profile.json", {})
    plan_path, plan_hash = evidence("plan.json", {"status": "pass",
        "source": {"source_sha256": source_hash},
        "semantic_map_input": {"path": str(map_path), "sha256": map_hash},
        "evidence": {"source_family_profile": {"path": str(profile_path), "sha256": profile_hash}}})
    check_path, check_hash = evidence("check.json", {"status": "pass",
        "semantic_plan_sha256": plan_hash, "semantic_map_input_sha256": map_hash,
        "source_family_profile_sha256": profile_hash})
    capture_path, capture_hash = evidence("capture.json", {})
    implementation_path, implementation_hash = evidence("implementation.json", {})
    implementation_input_path, _ = evidence("implementation-input.json", {})
    stage3 = {"source": {"workbook_path": str(source_path), "workbook_sha256": source_hash},
        "_artifact_sha256": design_hash,
        "analysis": {"artifacts": {"semantic_plan": {"path": str(plan_path), "sha256": plan_hash},
                                   "semantic_check": {"path": str(check_path), "sha256": check_hash}}},
        "design": {"targets": [{"selector": "Result"}], "validation_plan": [],
                   "historical_evidence": [], "tool_interface": {"support_domain": {
                       "status": "exact_cases" if exact_cases else "fixed_trace"}}}}
    monkeypatch.setattr("excel_to_act.steps.step4.generator._require_modular_design", lambda *_: None)
    monkeypatch.setattr("excel_to_act.steps.step4.generator.load_current_implementation_plan",
                        lambda *_: ({"implementation_input": {"path": str(implementation_input_path)}},
                                   implementation_path, implementation_hash,
                                   {"trace_path": trace_path, "projected_plan": {},
                                    "projected_profile": {}, "capture": {}, "capture_path": capture_path,
                                    "capture_file_sha256": capture_hash}))
    monkeypatch.setattr("excel_to_act.steps.step4.generator._saved_scenario_guard",
                        lambda *_: (["saved"], {}))
    monkeypatch.setattr("excel_to_act.steps.step4.generator._path_implementation_evidence",
                        lambda *_: {})
    monkeypatch.setattr("excel_to_act.steps.step4.generator.load_workflow",
                        lambda *_: (tmp_path, {"tool_interfaces": []}))
    calls: list[str] = []
    def export(_root: Path) -> dict:
        calls.append("export")
        return {"status": "blocked"}
    def build(**_kwargs: object) -> dict:
        calls.append("build")
        return {"manifest": {}}
    monkeypatch.setattr("excel_to_act.steps.step4.generator.export_tool_interface", export)
    monkeypatch.setattr("excel_to_act.steps.step4.generator.build_modular_bundle_files", build)
    monkeypatch.setattr("excel_to_act.steps.step4.generator._modular_oracle_ranges", lambda *_: [])

    if exact_cases:
        with pytest.raises(CalculationBlocked, match="exact-case interface cannot be exported"):
            _bundle_files(stage3, trace_path, tmp_path / "bundle", tmp_path)
        assert calls == ["export"]
    else:
        assert _bundle_files(stage3, trace_path, tmp_path / "bundle", tmp_path)["manifest"] == {}
        assert calls == ["build"]


def test_exact_case_replay_rejects_stale_case_reference(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from excel_to_act.steps.step4.modular_bundle import _verified_case_layout

    monkeypatch.setattr("excel_to_act.steps.step4.modular_bundle.compile_mapped_families",
                        lambda *_args: {})
    monkeypatch.setattr("excel_to_act.steps.step4.modular_bundle.render_grouped_driver",
                        lambda *_args: "driver")
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"synthetic source")
    cases = [{name: {"path": str(tmp_path / "absent.json"), "sha256": "a" * 64}
              for name in ("trace", "raw_request", "public_request", "producer")}]
    with pytest.raises(CalculationBlocked, match="exact-case trace artifact is missing or stale"):
        _verified_case_layout(cases, {}, [], source, {}, {}, {}, [], {}, {}, [])


def test_fixed_trace_guard_includes_lookup_return_used_by_a_branch() -> None:
    bindings = [{"address": f"main!B{row}", "variable_id": "raw.value", "indices": [row - 1]}
                for row in range(1, 4)]
    bindings.append({"address": "main!F1", "variable_id": "raw.config", "indices": []})
    trace = {"lookups": [{"formula_cell": "Main!D1", "function": "VLOOKUP",
                          "table_range": "Main!A1:E3", "lookup_column": 1,
                          "return_column": 2, "matched_row": 2}],
             "branches": [{"formula_cell": "Main!D1", "function": "IF", "selected_argument": 1}],
             "edges": [{"consumer": "Main!D1", "prerequisite": "Main!F1",
                        "relationship": "cell_reference"}]}
    contract = _variable_contract_layout(_variable_interface("a" * 64), bindings, [], "t" * 64,
                                         trace, {"main!B1": 11, "main!B2": 12, "main!B3": 13,
                                                 "main!F1": 2})
    assert contract["fixed_trace_source_values"] == {"main!B2": 12, "main!F1": 2}


def test_derived_seed_alias_must_match_an_approved_fixed_source_cell() -> None:
    definition = _variable_interface("a" * 64)
    bindings = ([{"address": f"main!A{row}", "variable_id": "raw.key", "indices": [row - 1]}
                 for row in range(1, 4)]
                + [{"address": f"main!B{row}", "variable_id": "raw.value", "indices": [row - 1]}
                   for row in range(1, 4)]
                + [{"address": "main!F1", "variable_id": "raw.config", "indices": []},
                   {"address": "main!A1", "variable_id": "state", "indices": [0]}])
    plan = {"variables": [{"variable_id": "state", "role": "derived", "initial_condition": {
        "input_variable_id": "raw.key", "input_index": [0], "logical_index": 0,
        "seed_ref": "Main!A1"}}]}
    saved = {**{f"main!A{row}": row for row in range(1, 4)},
             **{f"main!B{row}": row + 10 for row in range(1, 4)}, "main!F1": 2}
    contract = _variable_contract_layout(definition, bindings, [], "t" * 64, {}, saved,
                                         semantic_plan=plan)
    assert "state" not in contract["mutable"]
    plan["variables"][0]["initial_condition"]["seed_ref"] = "Main!A2"
    with pytest.raises(CalculationBlocked, match="approved fixed seed"):
        _variable_contract_layout(definition, bindings, [], "t" * 64, {}, saved,
                                  semantic_plan=plan)
    plan["variables"][0]["initial_condition"].update({
        "seed_ref": "Main!A1", "input_variable_id": "raw.value"})
    with pytest.raises(CalculationBlocked, match="approved fixed seed"):
        _variable_contract_layout(definition, bindings, [], "t" * 64, {}, saved,
                                  semantic_plan=plan)


def test_fixed_trace_guard_verifies_finite_match_selected_vlookup_column() -> None:
    definition = _variable_interface("a" * 64)
    bindings = [{"address": f"main!B{row}", "variable_id": "raw.value", "indices": [row - 1]}
                for row in range(1, 4)]
    bindings.append({"address": "main!F1", "variable_id": "raw.config", "indices": []})
    trace = {"source_sha256": "a" * 64, "lookups": [
        {"formula_cell": "Main!D1", "function": "MATCH", "lookup_range": "Main!H1:J1",
         "lookup_value": "Blue", "match_type": 0},
        {"formula_cell": "Main!D1", "function": "VLOOKUP", "table_range": "Main!A1:E3",
         "lookup_column": 1, "return_column": 2, "matched_row": 2}],
        "cells": [{"address": "Main!D1", "role": "calculated_formula",
                   "formula": "=VLOOKUP(A1,A1:E3,MATCH(F1,H1:J1,0),FALSE)"},
                  *[{"address": f"Main!{column}1", "role": "source_value", "value": value}
                    for column, value in (("H", "Red"), ("I", "Blue"), ("J", "Green"))]],
        "branches": [{"formula_cell": "Main!D1", "function": "IF", "selected_argument": 1}],
        "edges": [{"consumer": "Main!D1", "prerequisite": "Main!F1",
                   "relationship": "cell_reference"},
                  {"consumer": "Main!D1", "prerequisite": "Main!H1:J1",
                   "relationship": "range_read"}]}
    saved = {"main!B1": 11, "main!B2": 12, "main!B3": 13, "main!F1": 2}
    contract = _variable_contract_layout(definition, bindings, [], "t" * 64, trace, saved)
    assert contract["fixed_trace_source_values"] == {"main!B2": 12, "main!F1": 2}
    trace["cells"][2]["value"] = "Other"
    with pytest.raises(CalculationBlocked, match="MATCH key|MATCH selection"):
        _variable_contract_layout(definition, bindings, [], "t" * 64, trace, saved)
    trace["cells"][2]["value"] = "Blue"
    trace["lookups"][1]["return_column"] = 3
    with pytest.raises(CalculationBlocked, match="MATCH selection"):
        _variable_contract_layout(definition, bindings, [], "t" * 64, trace, saved)


def test_partial_raw_table_contract_selects_only_declared_cells() -> None:
    definition = _variable_interface("a" * 64)
    value = next(item for item in definition["objects"] if item["variable_id"] == "raw.value")
    value["mutable_cells"] = [{"indices": [1], "value_type": "number", "allow_blank": False,
                               "allowed_range": {"status": "enum", "values": [12, 18]}}]
    bindings = [{"address": f"main!B{row}", "variable_id": "raw.value", "indices": [row - 1]}
                for row in range(1, 4)]
    bindings.append({"address": "main!F1", "variable_id": "raw.config", "indices": []})
    contract = _variable_contract_layout(definition, bindings, [], "t" * 64, {},
                                         {f"main!B{row}": row + 10 for row in range(1, 4)})
    assert contract["mutable"]["Value"]["bindings"] == [bindings[1]]
    value["mutable_cells"] = [{**value["mutable_cells"][0], "indices": [9]}]
    with pytest.raises(CalculationBlocked, match="missing or duplicate per-cell"):
        _variable_contract_layout(definition, bindings, [], "t" * 64, {}, {})


def test_generated_value_checks_keep_partial_indices_and_dependent_enum_typed(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    adapter = types.ModuleType("input_adapter")
    adapter.build_values = lambda **_kwargs: None
    adapter._allocate = lambda shape, missing: missing if not shape else [missing] * shape[0]
    adapter._canonical_hash = lambda value: hashlib.sha256(json.dumps(
        value, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    pricing = types.ModuleType("pricing")
    pricing.run_pricing = lambda values, _functions: values
    families = types.ModuleType("formula_families")
    families.FAMILY_FUNCTIONS = {}
    for name, module in (("modular_runtime", runtime), ("input_adapter", adapter),
                         ("pricing", pricing), ("formula_families", families)):
        monkeypatch.setitem(sys.modules, name, module)
    namespace = {"__file__": "model.py"}
    exec(_model_source([], "s" * 64, "m" * 64), namespace)
    check = namespace["_checked_values"]
    table = {"shape": [2, 2], "mutable_cells": [
        {"indices": [0, 1], "value_type": "string", "allow_blank": False,
         "allowed_range": {"status": "enum", "values": ["A", "B"]}}]}
    assert check([{"indices": [0, 1], "value": "B"}], table, "Table") == {(0, 1): "B"}
    with pytest.raises(runtime.CalculationBlocked, match="indices differ"):
        check([{"indices": [1, 0], "value": "B"}], table, "Table")
    with pytest.raises(runtime.CalculationBlocked, match="approved values"):
        check([{"indices": [0, 1], "value": "C"}], table, "Table")
    period = {"shape": [], "value_type": "integer", "allow_blank": False,
              "allowed_range": {"status": "dependent_enum", "values": [1, 2],
                                "variable_id": "term"}}
    assert check(7, period, "Period", dependent_value=7) == {(): 7}
    with pytest.raises(runtime.CalculationBlocked, match="approved values"):
        check(8, period, "Period", dependent_value=7)
    with pytest.raises(runtime.CalculationBlocked, match="invalid integer"):
        check(True, period, "Period", dependent_value=True)
    axes = [{"name": "row", "keys": ["first", "second"]},
            {"name": "column", "keys": ["fixed", "choice"]}]
    contract = {"source_sha256": "a" * 64, "fixed_trace_source_values": {},
                "support_domain": {"external_policy": {}}, "external_specs": {},
                "mutable": {
                    "Table": {**table, "variable_id": "table", "axes": axes, "required": True,
                              "allowed_range": {"status": "unbounded"},
                              "default_policy": {"status": "none"},
                              "bindings": [{"address": "main!B1", "indices": [0, 1]}]},
                    "Term": {"shape": [], "variable_id": "term", "axes": [], "value_type": "integer",
                             "allowed_range": {"status": "declared", "minimum": 2, "maximum": 7},
                             "required": True, "default_policy": {"status": "none"},
                             "bindings": [{"address": "main!C1", "indices": []}]},
                    "Period": {**period, "variable_id": "period", "axes": [], "required": True,
                               "default_policy": {"status": "none"},
                               "bindings": [{"address": "main!D1", "indices": []}]}}}
    namespace["VARIABLE_CONTRACT_SHA256"] = adapter._canonical_hash(contract)
    namespace["__file__"] = str(tmp_path / "model.py")
    files = {}
    for filename in ("model.py", "input_layout.json", "source_values.json",
                     "pricing.py", "formula_families.py"):
        path = tmp_path / filename
        path.write_text("{}", encoding="utf-8")
        files[filename] = hashlib.sha256(path.read_bytes()).hexdigest()
    (tmp_path / "model_manifest.json").write_text(json.dumps({
        "source_sha256": "a" * 64, "files": files,
        "variable_input_contract_sha256": namespace["VARIABLE_CONTRACT_SHA256"]}), encoding="utf-8")
    layout = {"source_sha256": "a" * 64, "variable_contract": contract}
    inputs = {"Table": {"axes": axes, "values": [{"indices": [0, 1], "value": "B"}]},
              "Term": {"axes": [], "values": 7}, "Period": {"axes": [], "values": 7}}
    overrides, _external, _digest, _dependency, _contract = namespace["_variable_inputs"](
        layout, inputs, {})
    assert overrides == {"main!B1": "B", "main!C1": 7, "main!D1": 7}
    with pytest.raises(runtime.CalculationBlocked, match="approved values"):
        namespace["_variable_inputs"](layout, {**inputs, "Period": {"axes": [], "values": 8}}, {})
    contract["verified_cases"] = [{"input_sha256": adapter._canonical_hash(inputs),
                                    "dependency_sha256": adapter._canonical_hash({})}]
    namespace["VARIABLE_CONTRACT_SHA256"] = adapter._canonical_hash(contract)
    (tmp_path / "model_manifest.json").write_text(json.dumps({
        "source_sha256": "a" * 64, "files": files,
        "variable_input_contract_sha256": namespace["VARIABLE_CONTRACT_SHA256"]}), encoding="utf-8")
    assert namespace["_variable_inputs"](layout, inputs, {})[2] == adapter._canonical_hash(inputs)
    changed = {**inputs, "Table": {"axes": axes,
                                   "values": [{"indices": [0, 1], "value": "A"}]}}
    with pytest.raises(runtime.CalculationBlocked, match="no verified formula"):
        namespace["_variable_inputs"](layout, changed, {})
    assert namespace["_variable_inputs"](layout, inputs, {})[0] == overrides
    contract["verified_cases"][0]["dependency_sha256"] = "b" * 64
    namespace["VARIABLE_CONTRACT_SHA256"] = adapter._canonical_hash(contract)
    (tmp_path / "model_manifest.json").write_text(json.dumps({
        "source_sha256": "a" * 64, "files": files,
        "variable_input_contract_sha256": namespace["VARIABLE_CONTRACT_SHA256"]}), encoding="utf-8")
    with pytest.raises(runtime.CalculationBlocked, match="no verified formula"):
        namespace["_variable_inputs"](layout, inputs, {})


def test_generated_model_selects_only_exact_distinct_formula_variants(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path,
) -> None:
    source_hash = "a" * 64
    input_records = {value: {"Choice": {"axes": [], "values": value}} for value in (1, 2, 3)}
    dependency_hash = _canonical_hash({})
    monkeypatch.setitem(sys.modules, "modular_runtime", runtime)
    metadata = {}
    for variant_id, address, increment in (("v0001", "main!B1", 10),
                                           ("v0002", "main!C1", 100)):
        family = types.ModuleType(f"formula_families_{variant_id}")
        source = compile_formula_function(
            "calculate", "=A1+10" if increment == 10 else "=A1+D1", "Main", address,
            lambda reference, _sheet, _owner: ("values['raw.choice']" if reference.min_col == 1
                                                else "runtime.source_literal(values, 'main!D1')"),
            lambda _name, _sheet, _owner: "",
        )
        exec("import modular_runtime as runtime\n" + source, family.__dict__)
        family.FAMILY_FUNCTIONS = {"calculate": family.calculate}
        driver = types.ModuleType(f"pricing_{variant_id}")
        driver.run_pricing = lambda values, functions: {
            **values, "model.answer": functions["calculate"](values),
        }
        monkeypatch.setitem(sys.modules, family.__name__, family)
        monkeypatch.setitem(sys.modules, driver.__name__, driver)
        metadata[variant_id] = {
            "source_sha256": source_hash, "driver_sha256": hashlib.sha256(source.encode()).hexdigest(),
            "trace_sha256": increment.__str__().zfill(64), "plan_sha256": "b" * 64,
            "profile_sha256": "c" * 64, "named_bindings": {},
            "source_literals": {} if increment == 10 else {"main!D1": 100},
            "extra_variable_shapes": {} if increment == 10 else {"extra.helper": []},
            "formula_cells": [{"address": address, "variable_id": "model.answer", "indices": []}],
        }
    contract = {
        "source_sha256": source_hash, "interface_sha256": "i" * 64,
        "fixed_trace_source_values": {}, "support_domain": {"external_policy": {}},
        "external_specs": {}, "outputs": {"properties": {}}, "output_provenance": {},
        "mutable": {"Choice": {"variable_id": "raw.choice", "shape": [], "axes": [],
                              "value_type": "integer", "allow_blank": False, "required": True,
                              "default_policy": {"status": "none"},
                              "allowed_range": {"status": "enum", "values": [1, 2, 3]},
                              "bindings": [{"address": "main!A1", "indices": []}]}},
        "verified_cases": [
            {"input_sha256": _canonical_hash(input_records[value]),
             "dependency_sha256": dependency_hash, "variant_id": variant_id,
             "variant_metadata_sha256": _canonical_hash(metadata[variant_id]),
             "driver_sha256": metadata[variant_id]["driver_sha256"],
             "trace_sha256": metadata[variant_id]["trace_sha256"]}
            for value, variant_id in ((1, "v0001"), (2, "v0002"))
        ],
    }
    contract["support_certificate"] = {
        "verified_cases_sha256": _canonical_hash(contract["verified_cases"]),
        "active_trace_sha256": "t" * 64,
        "capture_artifact_sha256": "c" * 64,
    }
    adapter = types.ModuleType("input_adapter")
    adapter._canonical_hash = _canonical_hash
    adapter._allocate = lambda shape, missing: missing if not shape else [missing] * shape[0]
    adapter.build_values = lambda **kwargs: {
        "raw.choice": kwargs["raw_overrides"]["main!A1"], "model.answer": 0,
        "__source_values__": {"main!A1": kwargs["raw_overrides"]["main!A1"]},
    }
    pricing = types.ModuleType("pricing")
    pricing.run_pricing = lambda _values, _functions: pytest.fail("stale base driver ran")
    families = types.ModuleType("formula_families")
    families.FAMILY_FUNCTIONS = {}
    for name, module in (("modular_runtime", runtime), ("input_adapter", adapter),
                         ("pricing", pricing), ("formula_families", families)):
        monkeypatch.setitem(sys.modules, name, module)
    layout = {"source_sha256": source_hash, "variable_contract": contract,
              "target_bindings": {"Answer": "model.answer"},
              "variable_shapes": {"model.answer": []}, "formula_cells": []}
    files = {}
    for filename, payload in {
        "model.py": _model_source([], "s" * 64, "m" * 64, _canonical_hash(contract)),
        "input_layout.json": json.dumps(layout), "source_values.json": "{}",
        "pricing.py": "", "formula_families.py": "",
        **{f"variant_{key}.json": json.dumps(value) for key, value in metadata.items()},
        **{f"pricing_{key}.py": "" for key in metadata},
        **{f"formula_families_{key}.py": "" for key in metadata},
    }.items():
        path = tmp_path / filename
        path.write_text(payload, encoding="utf-8")
        files[filename] = hashlib.sha256(path.read_bytes()).hexdigest()
    (tmp_path / "model_manifest.json").write_text(json.dumps({
        "source_sha256": source_hash, "files": files,
        "variable_input_contract_sha256": _canonical_hash(contract),
        "support_certificate_sha256": _canonical_hash(contract["support_certificate"]),
        "trace_sha256": "t" * 64, "external_capture_artifact_sha256": "c" * 64,
    }), encoding="utf-8")
    namespace = {"__file__": str(tmp_path / "model.py")}
    exec((tmp_path / "model.py").read_text(encoding="utf-8"), namespace)
    assert namespace["run_model"](input_records[1], {})["targets"]["Answer"] == 11
    assert namespace["run_model"](input_records[2], {})["targets"]["Answer"] == 102
    assert namespace["run_model"](input_records[1], {})["cells"] == {"main!B1": 11}
    (tmp_path / "model_manifest.json").write_text(json.dumps({
        "source_sha256": source_hash, "files": files,
        "variable_input_contract_sha256": _canonical_hash(contract),
        "support_certificate_sha256": "0" * 64,
        "trace_sha256": "t" * 64, "external_capture_artifact_sha256": "c" * 64,
    }), encoding="utf-8")
    with pytest.raises(runtime.CalculationBlocked, match="support certificate is stale"):
        namespace["run_model"](input_records[1], {})
    (tmp_path / "model_manifest.json").write_text(json.dumps({
        "source_sha256": source_hash, "files": files,
        "variable_input_contract_sha256": _canonical_hash(contract),
        "support_certificate_sha256": _canonical_hash(contract["support_certificate"]),
        "trace_sha256": "t" * 64, "external_capture_artifact_sha256": "c" * 64,
    }), encoding="utf-8")
    with pytest.raises(runtime.CalculationBlocked, match="no verified formula"):
        namespace["run_model"](input_records[3], {})
    with pytest.raises(runtime.CalculationBlocked, match="external producer inputs"):
        namespace["run_model"](input_records[1], {"stale": {}})
    (tmp_path / "variant_v0001.json").write_text("{}", encoding="utf-8")
    with pytest.raises(runtime.CalculationBlocked, match="bundle asset changed"):
        namespace["run_model"](input_records[1], {})


def test_modular_bundle_runs_named_table_and_grouped_family_driver_in_isolation(tmp_path: Path) -> None:
    source_path = tmp_path / "source.xlsx"
    _workbook(source_path)
    workbook = load_workbook(source_path)
    workbook["Main"]["D1"] = "=VLOOKUP(2,RawTable,F1,FALSE)"
    workbook["Main"]["E2"] = "=A2*100"
    workbook["Main"]["F1"] = 2
    workbook.save(source_path)
    workbook.close()
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    variables = _input_records()
    formulas = [
        ("Main!C1", "=A1*10"), ("Main!C2", "=A2*10"), ("Main!C3", "=A3*10"),
        ("Main!D1", "=VLOOKUP(2,RawTable,F1,FALSE)"),
    ]
    trace = {"source_sha256": source_hash, "names": [
        {"scope": "workbook", "name": "RawTable"}, {"scope": "workbook", "name": "Answer"}],
        "cells": [
            *[{"address": f"Main!A{row}", "sheet": "Main", "role": "source_value"} for row in range(1, 4)],
            *[{"address": f"Main!B{row}", "sheet": "Main", "role": "source_value"} for row in range(1, 4)],
            {"address": "Main!F1", "sheet": "Main", "role": "source_value"},
            *[{"address": address, "sheet": "Main", "role": "calculated_formula", "formula": formula}
              for address, formula in formulas],
        ],
        "edges": [{"consumer": "Main!D1", "prerequisite": "Main!F1", "relationship": "cell_reference"},
                  {"consumer": "Main!D1", "prerequisite": "Main!A1:E3", "relationship": "lookup_table"}],
        "lookups": [{"formula_cell": "Main!D1", "function": "VLOOKUP",
                     "table_range": "Main!A1:E3", "lookup_column": 1, "return_column": 2,
                     "matched_row": 2, "match_mode": "exact"}]}
    plan = {
        "schema_version": "step3.semantic_plan.v1", "source": {"source_sha256": source_hash},
        "variables": variables,
        "equation_families": [
            {"family_id": "eq-value", "function_name": "table_value"},
            {"family_id": "eq-answer", "function_name": "answer_value"},
        ],
        "source_mappings": [
            {"source_family_id": "family-value", "equation_family_id": "eq-value",
             "variable_id": "table.value", "source_members": ["Main!C1", "Main!C2", "Main!C3"],
             "member_count": 3},
            {"source_family_id": "family-answer", "equation_family_id": "eq-answer",
             "variable_id": "model.answer", "source_members": ["Main!D1"], "member_count": 1},
        ],
        "array_mappings": [],
        "execution_plan": {"passes": [
            {"pass_id": "table-values", "source_family_ids": ["family-value"]},
            {"pass_id": "answer", "source_family_ids": ["family-answer"]},
        ]},
        "coverage": {
            "ordinary_formula_members_mapped": 4,
            "array_formula_followers_mapped": 0,
            "active_formula_members_mapped": 4,
            "source_family_count": 2,
            "semantic_family_count": 2,
            "array_instance_count": 0,
        },
        "reference_resolution": {"named_references": [
            {"name": "RawTable", "kind": "range", "source_ref": "Main!A1:E3",
             "table_binding": {"kind": "column_overlay_table", "key_variable_id": "raw.key",
                               "value_variable_id": "raw.value", "value_column_index": 2,
                               "formula_columns": {"C": "table.value"}}},
            {"name": "Answer", "kind": "scalar", "source_ref": "Main!D1",
             "variable_id": "model.answer", "resolution": "derived_formula_output"},
        ]},
    }
    profile = {"source_sha256": source_hash, "families": [
        {"family_id": "family-value", "source_members": ["Main!C1", "Main!C2", "Main!C3"],
         "member_count": 3},
        {"family_id": "family-answer", "source_members": ["Main!D1"], "member_count": 1},
    ]}
    raw_values = {f"main!A{row}": row for row in range(1, 4)}
    raw_values.update({f"main!B{row}": row + 10 for row in range(1, 4)})
    # The compiler must source every declared raw binding from the workbook, even
    # when the active formula trace omitted a literal cell from its frontier.
    trace["cells"] = [cell for cell in trace["cells"] if cell["address"] != "Main!A2"]
    raw_values.pop("main!A2")
    output = tmp_path / "revision"
    built = build_modular_bundle_files(
        semantic_plan=plan, source_profile=profile, trace=trace, raw_source_values=raw_values,
        source_path=source_path, semantic_plan_path=tmp_path / "semantic_plan.json",
        source_profile_path=tmp_path / "source_profile.json", semantic_map_path=tmp_path / "semantic_map.json",
        revision_dir=output, target_names=["Answer"], path_coverage={"implemented_path_ids": []},
        source_sha256=source_hash, design_sha256="d" * 64,
        semantic_plan_sha256="p" * 64, source_profile_sha256="f" * 64,
        trace_sha256="t" * 64, trace_path=tmp_path / "trace.json", trace_source="synthetic_fixture",
        scenario_ids=["synthetic"],
        variable_interface=_variable_interface(source_hash),
    )

    bundle = Path(built["bundle"])
    result = _smoke_bundle(bundle, output / "smoke", ["Answer"])
    payload = json.loads(Path(result["result_path"]).read_text(encoding="utf-8"))
    assert payload["targets"]["Answer"] == 12
    isolated = tmp_path / "isolated"
    shutil.copytree(bundle, isolated)
    assert not list(isolated.glob("*.xlsx"))

    def varied(values: list[int], label: str, axes: list[dict] | None = None) -> subprocess.CompletedProcess[str]:
        request = isolated / f"{label}.json"
        request.write_text(json.dumps({"inputs": {"Value": {"values": values,
                                                               "axes": axes if axes is not None else [{"name": "year", "keys": [1, 2, 3]}]}},
                                       "external_data": {}}), encoding="utf-8")
        return subprocess.run([sys.executable, "model.py", "--inputs", str(request),
                               "--out", str(isolated / f"{label}-result.json")], cwd=isolated,
                              capture_output=True, text=True, check=False)

    with ThreadPoolExecutor(max_workers=2) as pool:
        first, second = list(pool.map(lambda pair: varied(*pair),
                                      [([11, 25, 13], "first"), ([11, 20, 13], "second")]))
    assert first.returncode == second.returncode == 0, (first.stderr, second.stderr)
    varied_result = json.loads((isolated / "first-result.json").read_text(encoding="utf-8"))
    baseline_result = json.loads((isolated / "second-result.json").read_text(encoding="utf-8"))
    assert varied_result["targets"]["Answer"] == 25
    assert baseline_result["targets"]["Answer"] == 20
    partial_interface = _variable_interface(source_hash)
    next(item for item in partial_interface["objects"] if item["variable_id"] == "raw.value")[
        "mutable_cells"] = [{"indices": [1], "value_type": "number", "allow_blank": False,
                            "allowed_range": {"status": "enum", "values": [12, 25]}}]
    partial = build_modular_bundle_files(
        semantic_plan=plan, source_profile=profile, trace=trace, raw_source_values=raw_values,
        source_path=source_path, semantic_plan_path=tmp_path / "semantic_plan.json",
        source_profile_path=tmp_path / "source_profile.json", semantic_map_path=tmp_path / "semantic_map.json",
        revision_dir=tmp_path / "partial-revision", target_names=["Answer"],
        path_coverage={"implemented_path_ids": []}, source_sha256=source_hash,
        design_sha256="d" * 64, semantic_plan_sha256="p" * 64, source_profile_sha256="f" * 64,
        trace_sha256="t" * 64, trace_path=tmp_path / "trace.json", trace_source="synthetic_fixture",
        scenario_ids=["synthetic"], variable_interface=partial_interface)
    partial_bundle = Path(partial["bundle"])
    partial_request = tmp_path / "partial-request.json"
    partial_request.write_text(json.dumps({"inputs": {"Value": {"axes": [{"name": "year", "keys": [1, 2, 3]}],
        "values": [{"indices": [1], "value": 25}]}}, "external_data": {}}), encoding="utf-8")
    partial_run = subprocess.run([sys.executable, "model.py", "--inputs", str(partial_request),
                                  "--out", str(tmp_path / "partial-result.json")],
                                 cwd=partial_bundle, capture_output=True, text=True, check=False)
    assert partial_run.returncode == 0, partial_run.stderr
    partial_result = json.loads((tmp_path / "partial-result.json").read_text(encoding="utf-8"))
    assert partial_result["targets"]["Answer"] == 25
    assert partial_result["raw_inputs"]["main!B2"] == 25
    assert json.loads((partial_bundle / "source_values.json").read_text(encoding="utf-8"))[
        "main!B1"] == 11
    assert json.loads((partial_bundle / "source_values.json").read_text(encoding="utf-8"))[
        "main!B3"] == 13
    partial_request.write_text(json.dumps({"inputs": {"Value": {"axes": [{"name": "year", "keys": [1, 2, 3]}],
        "values": [{"indices": [0], "value": 25}]}}, "external_data": {}}), encoding="utf-8")
    rejected_fixed = subprocess.run([sys.executable, "model.py", "--inputs", str(partial_request)],
                                    cwd=partial_bundle, capture_output=True, text=True, check=False)
    assert rejected_fixed.returncode != 0 and "indices differ" in rejected_fixed.stderr
    assert "main!e2" not in {address.casefold() for address in built["manifest"]["formula_addresses"]}
    assert varied_result["raw_inputs"]["main!F1"] == 2
    assert json.loads((bundle / "input_layout.json").read_text(encoding="utf-8"))[
        "variable_contract"]["fixed_trace_source_values"] == {"main!F1": 2}
    assert varied_result["input_sha256"] != baseline_result["input_sha256"]
    selector_request = isolated / "selector.json"
    selector_request.write_text(json.dumps({"inputs": {
        "Value": {"values": [11, 25, 13], "axes": [{"name": "year", "keys": [1, 2, 3]}]},
        "Config": {"values": 5, "axes": []}}, "external_data": {}}), encoding="utf-8")
    unsupported = subprocess.run([sys.executable, "model.py", "--inputs", str(selector_request)],
                                 cwd=isolated, capture_output=True, text=True, check=False)
    assert unsupported.returncode != 0 and "unsupported fixed-trace formula or lookup selection: main!F1" in unsupported.stderr
    selector_request.write_text(json.dumps({"inputs": {
        "Value": {"values": [11, 25, 13], "axes": [{"name": "year", "keys": [1, 2, 3]}]},
        "Config": None}, "external_data": {}}), encoding="utf-8")
    null_optional = subprocess.run([sys.executable, "model.py", "--inputs", str(selector_request)],
                                   cwd=isolated, capture_output=True, text=True, check=False)
    assert null_optional.returncode != 0 and "Config must have values and axes" in null_optional.stderr
    assert varied([11, 25, 13], "repeat").returncode == 0
    assert json.loads((isolated / "repeat-result.json").read_text(encoding="utf-8"))["input_sha256"] == varied_result["input_sha256"]
    assert "outside its approved range" in varied([11, 26, 13], "range").stderr
    assert "invalid number" in varied([11, True, 13], "type").stderr
    assert "wrong shape" in varied([11, 25], "shape").stderr
    assert "axes or key order" in varied([11, 25, 13], "axes", [{"name": "other"}]).stderr
    assert "axes or key order" in varied([11, 25, 13], "reordered", [{"name": "year", "keys": [2, 1, 3]}]).stderr
    typed_axis = varied([11, 25, 13], "typed-axis", [{"name": "year", "keys": [True, 2, 3]}])
    assert typed_axis.returncode != 0 and "axes or key order" in typed_axis.stderr
    assert "duplicate or missing keys" in varied([11, 25, 13], "duplicate", [{"name": "year", "keys": [1, 1, 3]}]).stderr
    null_request = isolated / "null-request.json"
    null_request.write_text("null", encoding="utf-8")
    null_run = subprocess.run([sys.executable, "model.py", "--inputs", str(null_request),
                               "--out", str(isolated / "null-result.json")], cwd=isolated,
                              capture_output=True, text=True, check=False)
    assert null_run.returncode != 0 and "variable request must contain inputs and external_data" in null_run.stderr
    assert not (isolated / "null-result.json").exists()
    request = isolated / "missing.json"
    request.write_text(json.dumps({"inputs": {}, "external_data": {}}), encoding="utf-8")
    missing = subprocess.run([sys.executable, "model.py", "--inputs", str(request)],
                             cwd=isolated, capture_output=True, text=True, check=False)
    assert "required variable input is missing" in missing.stderr
    request.write_text(json.dumps({"inputs": {"Key": {"values": [1, 2, 3], "axes": [{"name": "year"}]}},
                                   "external_data": {}}), encoding="utf-8")
    immutable = subprocess.run([sys.executable, "model.py", "--inputs", str(request)],
                               cwd=isolated, capture_output=True, text=True, check=False)
    assert "unknown or immutable variable input" in immutable.stderr
    source_values_path = isolated / "source_values.json"
    source_values_path.write_bytes(source_values_path.read_bytes() + b" ")
    assert "bundle asset changed" in varied([11, 25, 13], "stale").stderr
    assert payload["calculated_variables"]["table.value"] == {
        "shape": [3],
        "demanded_values": [
            {"indices": [0], "value": 10},
            {"indices": [1], "value": 20},
            {"indices": [2], "value": 30},
        ],
    }
    assert payload["calculated_variables"]["model.answer"]["demanded_values"] == [
        {"indices": [], "value": 12}]
    assert payload["calculated_value_count"] == 4
    assert payload["cells"] == {"main!C1": 10, "main!C2": 20, "main!C3": 30, "main!D1": 12}
    assert built["manifest"]["compiled_variant_count"] == 2
    assert built["manifest"]["array_instance_count"] == built["coverage"]["array_instance_count"] == 0
    assert built["coverage"]["active_formula_member_count"] == 4
    assert "Main!" not in (bundle / "pricing.py").read_text(encoding="utf-8")
    assert "excel_to_act" not in "\n".join(path.read_text(encoding="utf-8")
                                            for path in bundle.glob("*.py"))
    model_source = (bundle / "model.py").read_text(encoding="utf-8")
    assert "main!A1" not in model_source and "main!F1" not in model_source
    raw_bundle = json.loads((bundle / "source_values.json").read_text(encoding="utf-8"))
    bundle_manifest = json.loads((bundle / "model_manifest.json").read_text(encoding="utf-8"))
    assert bundle_manifest["array_instance_count"] == built["coverage"]["array_instance_count"]
    assert set(raw_bundle) == {"main!A1", "main!A2", "main!A3", "main!B1", "main!B2", "main!B3", "main!F1"}
    assert raw_bundle["main!A2"] == 2
    layout = json.loads((bundle / "input_layout.json").read_text(encoding="utf-8"))
    assert layout["scenario_guard"]["source_values"]["main!A1"] == 1
    assert built["manifest"]["scenario_guard_coordinate_count"] == 7
    layout_path = bundle / "input_layout.json"
    layout["scenario_guard"]["source_values"]["main!A1"] = -1
    layout_path.write_text(json.dumps(layout), encoding="utf-8")
    changed_layout = subprocess.run(
        [sys.executable, "model.py", "--out", "changed_layout.json"], cwd=bundle,
        capture_output=True, text=True, check=False,
    )
    assert changed_layout.returncode != 0
    assert "bundle scenario guard differs from its generated contract" in changed_layout.stderr
    layout["scenario_guard"]["source_values"]["main!A1"] = 1
    layout_path.write_text(json.dumps(layout), encoding="utf-8")
    raw_bundle["main!A1"] = 99
    (bundle / "source_values.json").write_text(json.dumps(raw_bundle), encoding="utf-8")
    changed_scenario = subprocess.run(
        [sys.executable, "model.py", "--out", "changed.json"], cwd=bundle,
        capture_output=True, text=True, check=False,
    )
    assert changed_scenario.returncode != 0
    assert "saved-scenario source selector changed" in changed_scenario.stderr
    raw_bundle["main!A1"] = 1
    raw_bundle["main!F1"] = "Yes"
    (bundle / "source_values.json").write_text(json.dumps(raw_bundle), encoding="utf-8")
    changed_configuration = subprocess.run(
        [sys.executable, "model.py", "--out", "changed_config.json"], cwd=bundle,
        capture_output=True, text=True, check=False,
    )
    assert changed_configuration.returncode != 0
    assert "saved-scenario source selector changed: main!F1" in changed_configuration.stderr


def test_approved_stored_empty_adapter_metadata_reaches_bundle_and_unapproved_blank_blocks(tmp_path: Path) -> None:
    source_path = tmp_path / "source.xlsx"
    source_hash = _workbook(source_path)
    variables = _input_records()
    formulas = [
        ("Main!C1", "=A1*10"), ("Main!C2", "=A2*10"), ("Main!C3", "=A3*10"),
        ("Main!D1", "=VLOOKUP(2,RawTable,3,FALSE)"),
    ]
    base_cells = [
        *[{"address": f"Main!A{row}", "sheet": "Main", "role": "source_value", "value": row}
          for row in range(1, 4)],
        *[{"address": f"Main!B{row}", "sheet": "Main", "role": "source_value", "value": row + 10}
          for row in range(1, 4)],
        *[{"address": address, "sheet": "Main", "role": "calculated_formula", "formula": formula}
          for address, formula in formulas],
    ]
    trace = {"source_sha256": source_hash, "names": [
        {"scope": "workbook", "name": "RawTable"}, {"scope": "workbook", "name": "Answer"}],
        "cells": [*base_cells, {"address": "Main!E1", "sheet": "Main", "role": "source_value", "value": None}]}
    plan = {
        "schema_version": "step3.semantic_plan.v1", "source": {"source_sha256": source_hash},
        "variables": variables,
        "equation_families": [
            {"family_id": "eq-value", "function_name": "table_value"},
            {"family_id": "eq-answer", "function_name": "answer_value"},
        ],
        "source_mappings": [
            {"source_family_id": "family-value", "equation_family_id": "eq-value",
             "variable_id": "table.value", "source_members": ["Main!C1", "Main!C2", "Main!C3"],
             "member_count": 3},
            {"source_family_id": "family-answer", "equation_family_id": "eq-answer",
             "variable_id": "model.answer", "source_members": ["Main!D1"], "member_count": 1},
        ],
        "array_mappings": [],
        "execution_plan": {"passes": [
            {"pass_id": "table-values", "source_family_ids": ["family-value"]},
            {"pass_id": "answer", "source_family_ids": ["family-answer"]},
        ]},
        "coverage": {"ordinary_formula_members_mapped": 4, "array_formula_followers_mapped": 0,
                     "active_formula_members_mapped": 4, "source_family_count": 2,
                     "semantic_family_count": 2, "array_instance_count": 0},
        "reference_resolution": {"named_references": [
            {"name": "RawTable", "kind": "range", "source_ref": "Main!A1:C3",
             "table_binding": {"kind": "column_overlay_table", "key_variable_id": "raw.key",
                               "value_variable_id": "raw.value", "value_column_index": 2,
                               "formula_columns": {"C": "table.value"}}},
            {"name": "Answer", "kind": "scalar", "source_ref": "Main!D1",
             "variable_id": "model.answer", "resolution": "derived_formula_output"},
        ]},
    }
    profile = {"source_sha256": source_hash, "families": [
        {"family_id": "family-value", "source_members": ["Main!C1", "Main!C2", "Main!C3"],
         "member_count": 3},
        {"family_id": "family-answer", "source_members": ["Main!D1"], "member_count": 1},
    ]}
    scope = {"allowed_regions": [{"source_ref": "Main!E1",
                                  "rationale": "source-audited stored-empty adapter boundary"}]}
    candidate_trace = {"unknowns": [
        {"category": "unrepresented_source_coordinate", "evidence": "Main!E1"}]}
    blank_plan = json.loads(json.dumps(plan))
    metadata = add_adapter_metadata(blank_plan, {"axis_metadata": [], "source_records": []},
                                    scope, trace, source_path, source_hash, candidate_trace)
    assert metadata["source_coordinate_count"] == 1
    assert blank_plan["adapter_metadata_groups"]["source.verified_empty"]["addresses"] == ["main!E1"]

    output = tmp_path / "approved-empty-revision"
    built = build_modular_bundle_files(
        semantic_plan=blank_plan, source_profile=profile, trace=trace,
        raw_source_values={f"main!A{row}": row for row in range(1, 4)}
                         | {f"main!B{row}": row + 10 for row in range(1, 4)},
        source_path=source_path, semantic_plan_path=tmp_path / "semantic_plan.json",
        source_profile_path=tmp_path / "source_profile.json", semantic_map_path=tmp_path / "semantic_map.json",
        revision_dir=output, target_names=["Answer"], path_coverage={"implemented_path_ids": []},
        source_sha256=source_hash, design_sha256="d" * 64,
        semantic_plan_sha256="p" * 64, source_profile_sha256="f" * 64,
        trace_sha256="t" * 64, trace_path=tmp_path / "trace.json", trace_source="synthetic_fixture",
        scenario_ids=["synthetic"],
    )
    metadata_file = json.loads((Path(built["bundle"]) / "adapter_metadata.json").read_text(encoding="utf-8"))
    assert metadata_file["groups"]["source.verified_empty"]["values"] == [None]

    # A confirmed catalog can carry an audited adapter blank without a separate
    # single-cell model-scope region.
    catalog_blank = {"axis_metadata": [], "source_records": [{
        "record_id": "terminal", "classification": "adapter_metadata",
        "source_extents": [{"sheet": "Main", "range": "E1", "source_role": "raw_source"}],
        "provenance": {"source_workbook_sha256": source_hash},
    }], "source_scope_notes": {"stored_coordinate_reconciliation": {
        "source_copy_sha256": source_hash,
        "historical_empty_coordinates": [{"address": "Main!E1",
                                          "source_value_status": "stored_blank_no_literal"}],
    }}}
    catalog_plan = json.loads(json.dumps(plan))
    add_adapter_metadata(catalog_plan, catalog_blank, {"allowed_regions": []},
                         trace, source_path, source_hash, candidate_trace)
    assert catalog_plan["adapter_metadata_groups"]["source.terminal"]["classification"] == "source_audited_blank"
    assert _augment_literal_sources(source_path, source_hash, catalog_plan, trace, {})["main!E1"] is None

    stale_catalog = json.loads(json.dumps(catalog_blank))
    stale_catalog["source_records"][0]["provenance"]["source_workbook_sha256"] = "0" * 64
    with pytest.raises(CalculationBlocked, match="catalog blank.*source"):
        add_adapter_metadata(json.loads(json.dumps(plan)), stale_catalog, {"allowed_regions": []},
                             trace, source_path, source_hash, candidate_trace)

    missing_evidence = {**catalog_blank, "source_scope_notes": {}}
    missing_plan = json.loads(json.dumps(plan))
    add_adapter_metadata(missing_plan, missing_evidence, {"allowed_regions": []},
                         trace, source_path, source_hash, candidate_trace)
    with pytest.raises(CalculationBlocked, match="catalog metadata source coordinate is unexpectedly blank"):
        _augment_literal_sources(source_path, source_hash, missing_plan, trace, {})

    unapproved_trace = {**trace, "cells": [*trace["cells"], {
        "address": "Main!G1", "sheet": "Main", "role": "source_value", "value": None}]}
    unapproved_plan = json.loads(json.dumps(plan))
    add_adapter_metadata(unapproved_plan, {"axis_metadata": [], "source_records": []},
                         scope, unapproved_trace, source_path, source_hash, candidate_trace)
    with pytest.raises(CalculationBlocked, match="source values are outside the approved logical inputs and metadata: main!G1"):
        build_modular_bundle_files(
            semantic_plan=unapproved_plan, source_profile=profile, trace=unapproved_trace,
            raw_source_values={f"main!A{row}": row for row in range(1, 4)}
                             | {f"main!B{row}": row + 10 for row in range(1, 4)},
            source_path=source_path, semantic_plan_path=tmp_path / "semantic_plan.json",
            source_profile_path=tmp_path / "source_profile.json", semantic_map_path=tmp_path / "semantic_map.json",
            revision_dir=tmp_path / "unapproved-empty-revision", target_names=["Answer"],
            path_coverage={"implemented_path_ids": []}, source_sha256=source_hash,
            design_sha256="d" * 64, semantic_plan_sha256="p" * 64,
            source_profile_sha256="f" * 64, trace_sha256="u" * 64,
            trace_path=tmp_path / "unapproved-trace.json", trace_source="synthetic_fixture",
            scenario_ids=["synthetic"],
        )


def test_axis_metadata_initial_condition_is_checked_metadata_not_raw_input(tmp_path: Path) -> None:
    source_path = tmp_path / "axis-seed.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["A1"] = 7
    sheet["A2"] = "=A1*2"
    sheet["B1"] = "=A1+E1"
    sheet["E1"] = 11
    workbook.save(source_path)
    workbook.close()
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()

    variables = [
        {"variable_id": "raw.operand", "role": "raw", "kind": "scalar", "shape": [],
         "axes": [], "dependencies": [], "initial_condition": None,
         "source_extents": [{"ref": "Main!E1", "role": "raw_input",
                              "coordinate_mapping": {"kind": "scalar", "fixed_axes": []}}]},
        {"variable_id": "raw.seed_alias", "role": "derived", "kind": "scalar", "shape": [],
         "axes": [], "dependencies": [],
         "initial_condition": {"seed_ref": "Main!E1", "source_role": "raw_input_seed"},
         "source_extents": []},
        {"variable_id": "amr.age_key", "role": "derived", "kind": "projection_series", "shape": [2],
         "axes": [{"name": "age_index"}], "dependencies": [{"variable_id": "amr.age_key"}],
         "initial_condition": {"seed_ref": "Main!A1", "logical_index": 0,
                               "source_role": "accepted_axis_metadata"},
         "source_extents": [{"ref": "Main!A2", "role": "formula_output",
                              "coordinate_mapping": {"kind": "row_ordinal", "axes": [
                                  {"axis": "age_index", "dimension": 0,
                                   "source_coordinate": "row", "origin": 1, "index_origin": 0}]}}],
         "equation_segments": [{"segment_id": "age_update", "source_family_ids": ["family-age"],
                                "equation_family_ids": ["eq-age"], "index_axis": "age_index",
                                "index_start": 1, "index_stop_exclusive": 2,
                                "execution": "vectorized", "recurrence_group": None,
                                "snapshot_before_update": False}]},
        {"variable_id": "model.answer", "role": "derived", "kind": "scalar", "shape": [],
         "axes": [], "dependencies": [{"variable_id": "amr.age_key"},
                                       {"variable_id": "raw.operand"}],
         "initial_condition": None,
         "source_extents": [{"ref": "Main!B1", "role": "formula_output",
                              "coordinate_mapping": {"kind": "scalar", "fixed_axes": []}}],
         "equation_segments": [{"segment_id": "answer", "source_family_ids": ["family-answer"],
                                "equation_family_ids": ["eq-answer"], "index_axis": None,
                                "index_start": 0, "index_stop_exclusive": 1,
                                "execution": "scalar", "recurrence_group": None,
                                "snapshot_before_update": False}]},
    ]
    plan = {
        "schema_version": "step3.semantic_plan.v1", "source": {"source_sha256": source_hash},
        "variables": variables,
        "equation_families": [{"family_id": "eq-age", "function_name": "age_update"},
                              {"family_id": "eq-answer", "function_name": "answer"}],
        "source_mappings": [
            {"source_family_id": "family-age", "equation_family_id": "eq-age",
             "variable_id": "amr.age_key", "source_members": ["Main!A2"], "member_count": 1},
            {"source_family_id": "family-answer", "equation_family_id": "eq-answer",
             "variable_id": "model.answer", "source_members": ["Main!B1"], "member_count": 1},
        ],
        "array_mappings": [],
        "execution_plan": {"passes": [
            {"pass_id": "age", "source_family_ids": ["family-age"]},
            {"pass_id": "answer", "source_family_ids": ["family-answer"]},
        ]},
        "coverage": {"ordinary_formula_members_mapped": 2, "array_formula_followers_mapped": 0,
                     "active_formula_members_mapped": 2, "source_family_count": 2,
                     "semantic_family_count": 2, "array_instance_count": 0},
        "reference_resolution": {"named_references": [
            {"name": "Answer", "kind": "scalar", "source_ref": "Main!B1",
             "variable_id": "model.answer", "resolution": "derived_formula_output"}]},
    }
    formulas = {"Main!A2": "=A1*2", "Main!B1": "=A1+E1"}
    trace = {"source_sha256": source_hash,
             "names": [{"scope": "workbook", "name": "Answer", "destination": "Main!B1"}],
             "cells": [
                 {"address": "Main!A1", "sheet": "Main", "role": "source_value", "value": 7},
                 {"address": "Main!E1", "sheet": "Main", "role": "source_value", "value": 11},
                 *[{"address": address, "sheet": "Main", "role": "calculated_formula", "formula": formula}
                   for address, formula in formulas.items()],
             ]}
    profile = {"source_sha256": source_hash, "families": [
        {"family_id": "family-age", "source_members": ["Main!A2"], "member_count": 1},
        {"family_id": "family-answer", "source_members": ["Main!B1"], "member_count": 1},
    ]}
    projection = add_adapter_metadata(plan, {"axis_metadata": [{"axis_id": "age_origin",
                                                                  "source_extents": [{"sheet": "Main", "range": "A1"}]}],
                                             "source_records": []},
                                      {"allowed_regions": []}, trace, source_path, source_hash, {"unknowns": []})
    assert projection["raw_source_coordinate_count"] == 1
    assert projection["source_metadata_coordinate_count"] == 1
    assert projection["metadata_seed_binding_count"] == 1
    assert plan["metadata_seed_bindings"] == [{
        "variable_id": "amr.age_key", "indices": [0], "metadata_group": "axis.age_origin",
        "metadata_indices": [0], "source_address": "main!A1"}]
    source_index = build_source_cell_index(plan)
    raw_bindings = _raw_bindings(plan, source_index, {item["variable_id"]: item for item in variables})
    assert {item["address"].casefold() for item in raw_bindings} == {"main!e1"}
    assert {item["variable_id"] for item in raw_bindings} == {"raw.operand", "raw.seed_alias"}

    revision = tmp_path / "bundle-revision"
    built = build_modular_bundle_files(
        semantic_plan=plan, source_profile=profile, trace=trace,
        raw_source_values={"main!E1": 11}, source_path=source_path,
        semantic_plan_path=tmp_path / "semantic_plan.json", source_profile_path=tmp_path / "profile.json",
        semantic_map_path=tmp_path / "map.json", revision_dir=revision, target_names=["Answer"],
        path_coverage={"implemented_path_ids": []}, source_sha256=source_hash,
        design_sha256="d" * 64, semantic_plan_sha256="p" * 64,
        source_profile_sha256="f" * 64, trace_sha256="t" * 64,
        trace_path=tmp_path / "trace.json", trace_source="synthetic_fixture", scenario_ids=["saved"],
    )
    bundle = Path(built["bundle"])
    result = _smoke_bundle(bundle, revision / "smoke", ["Answer"])
    output = json.loads(Path(result["result_path"]).read_text(encoding="utf-8"))
    assert output["targets"]["Answer"] == 18
    assert built["manifest"]["raw_source_coordinate_count"] == 1
    assert built["manifest"]["source_metadata_coordinate_count"] == 1
    assert built["manifest"]["metadata_seed_binding_count"] == 1
    raw_values = json.loads((bundle / "source_values.json").read_text(encoding="utf-8"))
    assert raw_values == {"main!E1": 11}
    metadata = json.loads((bundle / "adapter_metadata.json").read_text(encoding="utf-8"))
    original_metadata_bytes = (bundle / "adapter_metadata.json").read_bytes()
    assert metadata["groups"]["axis.age_origin"]["values"] == [7]

    def run_model(label: str) -> subprocess.CompletedProcess[str]:
        return subprocess.run([sys.executable, "model.py", "--out", f"{label}.json"], cwd=bundle,
                              capture_output=True, text=True, check=False)

    metadata["groups"].clear()
    (bundle / "adapter_metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
    missing_group = run_model("missing_group")
    assert missing_group.returncode != 0
    assert "adapter metadata file is stale or incomplete" in missing_group.stderr
    metadata["groups"]["axis.age_origin"] = {"addresses": ["main!A1"], "values": [7],
                                               "sha256": "bad"}
    (bundle / "adapter_metadata.json").write_text(json.dumps(metadata), encoding="utf-8")
    changed_group = run_model("changed_group")
    assert changed_group.returncode != 0
    assert "adapter metadata group changed: axis.age_origin" in changed_group.stderr
    (bundle / "adapter_metadata.json").write_bytes(original_metadata_bytes)
    layout_path = bundle / "input_layout.json"
    layout = json.loads(layout_path.read_text(encoding="utf-8"))
    layout["metadata_seed_bindings"] = []
    layout_path.write_text(json.dumps(layout), encoding="utf-8")
    changed_binding = run_model("changed_seed_binding")
    assert changed_binding.returncode != 0
    assert "metadata seed bindings are missing or differ" in changed_binding.stderr


def test_named_scalar_can_bind_external_coordinate_or_declared_source_literal() -> None:
    plan = {"variables": [
        {"variable_id": "external.rate", "role": "external", "shape": [2],
         "axes": [{"name": "age"}], "source_extents": [{
             "ref": "Rates!W5:W6", "role": "formula_output",
             "coordinate_mapping": {"kind": "row_ordinal", "axes": [{
                 "axis": "age", "dimension": 0, "source_coordinate": "row",
                 "origin": 5, "index_origin": 0}]}}]},
        {"variable_id": "raw.rate_header", "role": "raw", "shape": [], "axes": [],
         "source_extents": [{"ref": "Rates!W3", "role": "adapter_metadata",
             "coordinate_mapping": {"kind": "scalar", "axes": []}}]},
    ], "source_mappings": [], "array_mappings": [], "reference_resolution": {
        "named_references": [
            {"name": "SelectedRate", "kind": "scalar", "source_ref": "Rates!W5",
             "binding": {"kind": "variable_coordinate", "variable_id": "external.rate", "indices": [0]}},
            {"name": "SelectedRateDirect", "kind": "variable_coordinate", "source_ref": "Rates!W5",
             "variable_id": "external.rate", "indices": [0]},
            {"name": "RateLabel", "kind": "scalar", "source_ref": "Rates!W3",
             "binding": {"kind": "source_literal", "address": "Rates!W3"}},
        ]}}

    index = build_source_cell_index(plan)
    names = _named_bindings(plan, index, {"rates!W3"})
    values = {"external.rate": [0.25, 0.5], "__source_values__": {"rates!W3": "CI_2017_M"},
              "__names__": names}

    assert runtime.lookup_name(values, "SelectedRate") == 0.25
    assert runtime.lookup_name(values, "SelectedRateDirect") == 0.25
    assert runtime.lookup_name(values, "RateLabel") == "CI_2017_M"
    values["__source_values__"].clear()
    with pytest.raises(runtime.CalculationBlocked, match="missing"):
        runtime.lookup_name(values, "RateLabel")


def test_raw_binding_that_is_now_a_formula_is_rejected(tmp_path: Path) -> None:
    source_path = tmp_path / "source.xlsx"
    source_hash = _workbook(source_path)
    plan = {"source": {"source_sha256": source_hash}, "variables": [
        {"variable_id": "raw.formula", "role": "raw", "kind": "scalar", "shape": [],
         "axes": [], "dependencies": [], "initial_condition": None,
         "source_extents": [{"ref": "Main!C1", "role": "raw_input",
                              "coordinate_mapping": {"kind": "scalar", "fixed_axes": []}}]},
    ]}
    with pytest.raises(CalculationBlocked, match="declared raw source boundary is a formula cell"):
        _augment_literal_sources(source_path, source_hash, plan, {"cells": []}, {})


def test_declared_source_literal_is_read_and_formula_is_rejected(tmp_path: Path) -> None:
    source_path = tmp_path / "source.xlsx"
    source_hash = _workbook(source_path)
    plan = {"variables": [], "adapter_source_literals": ["Main!B2"]}
    assert _augment_literal_sources(source_path, source_hash, plan, {"cells": []}, {}) == {
        "main!B2": 12,
    }
    plan["adapter_source_literals"] = ["Main!C2"]
    with pytest.raises(CalculationBlocked, match="declared raw source boundary is a formula cell"):
        _augment_literal_sources(source_path, source_hash, plan, {"cells": []}, {})


def test_raw_binding_that_overlaps_an_array_follower_is_rejected(tmp_path: Path) -> None:
    source_path = tmp_path / "array.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["A1"] = 1
    sheet["A2"] = 2
    sheet["C1"] = ArrayFormula(ref="C1:C2", text="=A1:A2")
    workbook.save(source_path)
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    plan = {"source": {"source_sha256": source_hash}, "variables": [
        {"variable_id": "raw.array_follower", "role": "raw", "kind": "scalar", "shape": [],
         "axes": [], "dependencies": [], "initial_condition": None,
         "source_extents": [{"ref": "Main!C2", "role": "raw_input",
                              "coordinate_mapping": {"kind": "scalar", "fixed_axes": []}}]},
    ]}
    trace = {"cells": [{"address": "Main!C2", "role": "calculated_array_formula",
                        "array_ref": "C1:C2"}]}
    with pytest.raises(CalculationBlocked, match="formula or array member"):
        _augment_literal_sources(source_path, source_hash, plan, trace, {})


def test_blank_cells_compare_as_zero_in_numeric_conditions() -> None:
    assert binary("<=", lambda: None, lambda: 1) is True
    assert binary(">", lambda: 1, lambda: None) is True
    assert binary("=", lambda: None, lambda: 0) is True
    assert binary("<", lambda: None, lambda: None) is False
