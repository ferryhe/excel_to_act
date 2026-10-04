"""Focused acceptance coverage for the human Step1 conversion boundary."""

from __future__ import annotations

import json
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils.datetime import CALENDAR_MAC_1904
import pytest
from typer.testing import CliRunner

from excel_to_act.interfaces.cli import app
from excel_to_act.ingest.openpyxl_reader import OpenpyxlWorkbookReader
from excel_to_act.inventory.extractor import OpenpyxlInventoryExtractor
from excel_to_act.schemas import CellInventory
from excel_to_act.steps.step1.source_scan import SourceScanError, scan_step1_source
from excel_to_act.steps.step1.workflow import _hash_file, auto_recover, convert_directory, execute_tool, finalize_run


MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _make_book(path: Path, *, custom_part: bool = False, marker: str = "first") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.epoch = CALENDAR_MAC_1904
    sheet = workbook.active
    sheet.title = "Facts"
    sheet.append(["label", "number", "date", "cached", "empty", "data table"])
    sheet.append(["Premium", 1.0, datetime(2024, 1, 1), "=B2*2", "", "=1+1"])
    sheet["D3"] = "=FALSE()"
    sheet["D4"] = '=""'
    sheet["D5"] = "=1+1"
    sheet["A6"] = "Field"
    sheet["B6"] = "Value"
    sheet["A7"] = marker
    sheet["B7"] = 0.05
    sheet["A9"] = "Merged title"
    sheet.merge_cells("A9:B9")
    sheet["A2"].comment = Comment("source note", "Step1 test")
    sheet["A3"].hyperlink = "https://example.test/facts"
    sheet["A3"].style = "Hyperlink"
    sheet["A8"].value = "valid"
    sheet.row_dimensions[8].hidden = True
    sheet.column_dimensions["G"].width = 14
    sheet.freeze_panes = "B2"
    sheet.add_table(Table(displayName="InputTable", ref="A6:B7"))
    validation = DataValidation(type="decimal", operator="greaterThan", formula1="0")
    sheet.add_data_validation(validation)
    validation.add("B7")
    sheet.conditional_formatting.add("B7", FormulaRule(formula=["B7>0"]))
    workbook.defined_names.add(DefinedName("Rate", attr_text="Facts!$B$7"))
    workbook.defined_names.add(DefinedName("Rate", attr_text="Facts!$B$7", localSheetId=0))
    workbook.save(path)
    workbook.close()

    temp_path = path.with_name(f"{path.stem}.rewrite{path.suffix}")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(payload)
                cells = {cell.attrib.get("r"): cell for cell in root.iter() if cell.tag == MAIN + "c"}
                if "E2" not in cells:
                    row = next(row for row in root.iter() if row.tag == MAIN + "row" and row.attrib.get("r") == "2")
                    empty_cell = ET.SubElement(row, MAIN + "c", {"r": "E2", "t": "inlineStr"})
                    cells["E2"] = empty_cell
                empty_cell = cells["E2"]
                empty_cell.attrib["t"] = "inlineStr"
                for child in list(empty_cell):
                    empty_cell.remove(child)
                inline = ET.SubElement(empty_cell, MAIN + "is")
                ET.SubElement(inline, MAIN + "t")
                cells["B2"].find(MAIN + "v").text = "12345678901234567890.1234"
                _set_formula(cells["D2"], "B2*2", {"t": "shared", "ref": "D2:D3", "si": "0"}, "0")
                cells["D3"].attrib["t"] = "b"
                _set_formula(cells["D3"], None, {"t": "shared", "si": "0"}, "0")
                cells["D4"].attrib["t"] = "str"
                _set_formula(cells["D4"], "1+1", {"t": "array", "ref": "D4"}, "")
                _set_formula(cells["D5"], None, {"t": "dataTable", "ref": "D5:E6", "dt2D": "1"}, None)
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif info.filename == "xl/workbook.xml":
                root = ET.fromstring(payload)
                defined_names = next((element for element in root if element.tag == MAIN + "definedNames"), None)
                if defined_names is None:
                    defined_names = ET.SubElement(root, MAIN + "definedNames")
                if not any(element.tag == MAIN + "definedName" and element.attrib.get("localSheetId") is None for element in defined_names):
                    name = ET.SubElement(defined_names, MAIN + "definedName", {"name": "Rate"})
                    name.text = "Facts!$B$7"
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, payload)
        if custom_part:
            target.writestr("xl/custom-opaque.bin", b"opaque-bytes")
    temp_path.replace(path)
    return path


def _set_formula(cell: ET.Element, text: str | None, attributes: dict[str, str], cache: str | None) -> None:
    formula = cell.find(MAIN + "f")
    if formula is None:
        formula = ET.SubElement(cell, MAIN + "f")
    formula.attrib.clear()
    formula.attrib.update(attributes)
    formula.text = text
    value = cell.find(MAIN + "v")
    if cache is None:
        if value is not None:
            cell.remove(value)
    else:
        if value is None:
            value = ET.SubElement(cell, MAIN + "v")
        value.text = cache


def _convert_single(tmp_path: Path, *, custom_part: bool = False) -> tuple[Path, Path, dict]:
    raw = tmp_path / "raw"
    source = _make_book(raw / "book.xlsx", custom_part=custom_part)
    out = tmp_path / "converted"
    result = convert_directory(raw, out)
    run = out / result["entries"][0]["run_path"]
    return source, run, result


def _read(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def _write(path: Path, value) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _rewrite_package(path: Path, *, replacements: dict[str, bytes] | None = None, additions: dict[str, bytes] | None = None) -> None:
    replacements = replacements or {}
    additions = additions or {}
    temporary = path.with_name(f"{path.stem}.package-rewrite{path.suffix}")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        names = set()
        for info in source.infolist():
            names.add(info.filename)
            target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
        for name, payload in additions.items():
            assert name not in names
            target.writestr(name, payload)
    temporary.replace(path)


def _append_relationship(path: Path, rels_part: str, attributes: dict[str, str]) -> None:
    with zipfile.ZipFile(path, "r") as package:
        root = ET.fromstring(package.read(rels_part))
    ET.SubElement(root, PKG_REL + "Relationship", attributes)
    _rewrite_package(path, replacements={rels_part: ET.tostring(root, encoding="utf-8", xml_declaration=True)})


def _make_typed_date_book(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "N/A"
    sheet["A1"].number_format = "yyyy-mm-dd"
    workbook.save(path)
    workbook.close()

    sheet_root = ET.Element(MAIN + "worksheet")
    ET.SubElement(sheet_root, MAIN + "dimension", {"ref": "A1:I1"})
    sheet_data = ET.SubElement(sheet_root, MAIN + "sheetData")
    row = ET.SubElement(sheet_data, MAIN + "row", {"r": "1"})
    cell_xml = (
        ("A1", {"s": "1", "t": "s"}, None, "0"),
        ("B1", {"s": "1", "t": "str"}, None, "42"),
        ("C1", {"s": "1", "t": "e"}, None, "#N/A"),
        ("D1", {"s": "1", "t": "d"}, None, "2024-01-02T00:00:00"),
        ("E1", {"s": "1", "t": "inlineStr"}, None, None),
        ("F1", {"s": "1", "t": "b"}, None, "1"),
        ("G1", {"s": "1"}, None, "45293"),
        ("H1", {"s": "1"}, "45293", "45293"),
        ("I1", {"s": "1", "t": "str"}, 'IF(1,"N/A","")', "N/A"),
    )
    for address, attributes, formula, value in cell_xml:
        cell = ET.SubElement(row, MAIN + "c", {"r": address, **attributes})
        if formula is not None:
            ET.SubElement(cell, MAIN + "f").text = formula
        if attributes.get("t") == "inlineStr":
            inline = ET.SubElement(cell, MAIN + "is")
            ET.SubElement(inline, MAIN + "t").text = "文本"
        elif value is not None:
            ET.SubElement(cell, MAIN + "v").text = value

    shared_strings = ET.Element(MAIN + "sst", {"count": "1", "uniqueCount": "1"})
    item = ET.SubElement(shared_strings, MAIN + "si")
    ET.SubElement(item, MAIN + "t").text = "N/A"

    with zipfile.ZipFile(path, "r") as package:
        relationships = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        content_types = ET.fromstring(package.read("[Content_Types].xml"))
    ET.SubElement(
        relationships,
        PKG_REL + "Relationship",
        {
            "Id": "rIdStep1SharedStrings",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
            "Target": "sharedStrings.xml",
        },
    )
    ET.SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        {
            "PartName": "/xl/sharedStrings.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
        },
    )
    _rewrite_package(
        path,
        replacements={
            "xl/worksheets/sheet1.xml": ET.tostring(sheet_root, encoding="utf-8", xml_declaration=True),
            "xl/_rels/workbook.xml.rels": ET.tostring(relationships, encoding="utf-8", xml_declaration=True),
            "[Content_Types].xml": ET.tostring(content_types, encoding="utf-8", xml_declaration=True),
        },
        additions={"xl/sharedStrings.xml": ET.tostring(shared_strings, encoding="utf-8", xml_declaration=True)},
    )
    return path


def _add_phonetic_string_cases(path: Path) -> tuple[bytes, bytes]:
    xml_namespace = "http://www.w3.org/XML/1998/namespace"
    strings_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="2" uniqueCount="2">\n'
        "  <si>\n    <t>漢字</t>\n    <rPh sb=\"0\" eb=\"2\"><t>かんじ</t></rPh>\n    <phoneticPr fontId=\"1\"/>\n  </si>\n"
        "  <si>\n    <r><rPr><b/></rPr><t>漢</t></r>\n    <r><t>字</t></r>\n    <rPh sb=\"0\" eb=\"2\"><t>かんじ</t></rPh>\n    <phoneticPr fontId=\"1\"/>\n  </si>\n"
        "</sst>"
    ).encode("utf-8")
    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(path, "r") as source:
        sheet = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
        cells = {cell.attrib["r"]: cell for cell in sheet.iter() if cell.tag == MAIN + "c"}
        for row_number in range(1, 7):
            address = f"A{row_number}"
            if address not in cells:
                row = next(row for row in sheet.iter() if row.tag == MAIN + "row" and row.attrib.get("r") == str(row_number))
                cells[address] = ET.Element(MAIN + "c", {"r": address})
                row.insert(0, cells[address])
        for index, address in enumerate(("A1", "A2")):
            cell = cells[address]
            cell.attrib["t"] = "s"
            for child in list(cell):
                cell.remove(child)
            ET.SubElement(cell, MAIN + "v").text = str(index)

        def inline(address: str, *, rich: bool = False, text: str | None = "漢字", spaces: bool = False) -> None:
            cell = cells[address]
            cell.attrib["t"] = "inlineStr"
            for child in list(cell):
                cell.remove(child)
            container = ET.SubElement(cell, MAIN + "is")
            container.text = "\n  "
            if rich:
                first = ET.SubElement(container, MAIN + "r")
                ET.SubElement(first, MAIN + "rPr")
                ET.SubElement(first, MAIN + "t").text = "漢"
                first.tail = "\n  "
                second = ET.SubElement(container, MAIN + "r")
                ET.SubElement(second, MAIN + "t").text = "字"
                second.tail = "\n  "
            else:
                body = ET.SubElement(container, MAIN + "t")
                if spaces:
                    body.attrib[f"{{{xml_namespace}}}space"] = "preserve"
                body.text = text
                body.tail = "\n  "
            phonetic = ET.SubElement(container, MAIN + "rPh", {"sb": "0", "eb": "2"})
            ET.SubElement(phonetic, MAIN + "t").text = "かんじ"
            phonetic.tail = "\n  "
            ET.SubElement(container, MAIN + "phoneticPr", {"fontId": "1"})
            container.tail = "\n  "

        inline("A3")
        inline("A4", rich=True)
        inline("A5", text="  spacing  ", spaces=True)
        inline("A6", text="")
        replacements["xl/worksheets/sheet1.xml"] = ET.tostring(sheet, encoding="utf-8", xml_declaration=True)

        workbook_rels = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        relation_ids = [int(rel.attrib["Id"][3:]) for rel in workbook_rels if rel.attrib.get("Id", "").startswith("rId") and rel.attrib["Id"][3:].isdigit()]
        ET.SubElement(
            workbook_rels,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            {"Id": f"rId{max(relation_ids, default=0) + 1}", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings", "Target": "sharedStrings.xml"},
        )
        replacements["xl/_rels/workbook.xml.rels"] = ET.tostring(workbook_rels, encoding="utf-8", xml_declaration=True)

        content_types = ET.fromstring(source.read("[Content_Types].xml"))
        ET.SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            {"PartName": "/xl/sharedStrings.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"},
        )
        replacements["[Content_Types].xml"] = ET.tostring(content_types, encoding="utf-8", xml_declaration=True)
    _rewrite_package(path, replacements=replacements, additions={"xl/sharedStrings.xml": strings_xml})
    with zipfile.ZipFile(path, "r") as source:
        return source.read("xl/sharedStrings.xml"), source.read("xl/worksheets/sheet1.xml")


def _make_shared_string_index_case(path: Path, index: str, *, formula: bool) -> Path:
    _make_book(path)
    _add_phonetic_string_cases(path)
    address = "A2" if formula else "A1"
    with zipfile.ZipFile(path, "r") as package:
        sheet = ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
    cell = next(item for item in sheet.iter() if item.tag == MAIN + "c" and item.attrib.get("r") == address)
    cell.attrib["t"] = "s"
    for child in list(cell):
        cell.remove(child)
    if formula:
        ET.SubElement(cell, MAIN + "f").text = "1+1"
    ET.SubElement(cell, MAIN + "v").text = index
    _rewrite_package(path, replacements={"xl/worksheets/sheet1.xml": ET.tostring(sheet, encoding="utf-8", xml_declaration=True)})
    return path


def _make_multi_sheet_book(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    first = workbook.active
    first.title = "One"
    first["A1"] = 0
    second = workbook.create_sheet("Two")
    second["B2"] = 2
    workbook.defined_names.add(DefinedName("Multi", attr_text="'One'!$A$1,'Two'!$B$2"))
    workbook.save(path)
    workbook.close()
    return path


def test_mixed_recursive_batch_duplicate_identity_and_output_exclusion(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    first = _make_book(raw / "west" / "book.xlsx", marker="west")
    duplicate = raw / "copy" / "book-copy.xlsx"
    duplicate.parent.mkdir(parents=True)
    shutil.copyfile(first, duplicate)
    _make_book(raw / "east" / "book.xlsx", marker="east")
    (raw / "broken.xlsx").write_bytes(b"not an OOXML workbook")
    (raw / "old.xls").write_bytes(b"legacy input")
    (raw / "table.csv").write_text("A,B\n1,2\n", encoding="utf-8")

    output = raw / "step1-output"
    runner = CliRunner()
    first_result = runner.invoke(app, ["step1", "convert", str(raw), "--out", str(output)])
    assert first_result.exit_code == 1
    first_envelope = json.loads(first_result.output)
    assert first_envelope["status"] == "fail"
    assert first_envelope["metrics"]["input_count"] == 6
    entries = first_envelope["entries"]
    assert {entry["relative_path"] for entry in entries} == {
        "west/book.xlsx", "east/book.xlsx", "copy/book-copy.xlsx", "broken.xlsx", "old.xls", "table.csv"
    }
    supported = [entry for entry in entries if entry["relative_path"].endswith(".xlsx") and entry["status"] != "fail"]
    assert len({entry["run_path"] for entry in supported}) == 3
    by_path = {entry["relative_path"]: entry for entry in entries}
    assert by_path["west/book.xlsx"]["sha256"] == by_path["copy/book-copy.xlsx"]["sha256"]
    assert by_path["west/book.xlsx"]["run_path"] != by_path["copy/book-copy.xlsx"]["run_path"]
    assert by_path["west/book.xlsx"]["sha256"] != by_path["east/book.xlsx"]["sha256"]

    second_result = runner.invoke(app, ["step1", "convert", str(raw), "--out", str(output)])
    second_envelope = json.loads(second_result.output)
    assert second_envelope["metrics"]["input_count"] == 6
    assert second_envelope["source"]["output_subtree_excluded"] is True
    batch_path = output / second_envelope["artifacts"][0]["path"]
    batch = _read(batch_path)
    assert batch["discovery"]["output_subtree_excluded"] is True
    assert batch["discovery"]["excluded_output_path"] == "step1-output"

    empty_input = tmp_path / "empty-input"
    empty_input.mkdir()
    empty_batch = convert_directory(empty_input, tmp_path / "empty-output")
    assert empty_batch["status"] == "fail"
    assert empty_batch["entries"][0]["diagnostics"][0]["code"] == "empty_input_directory"


def test_unreadable_source_hash_is_one_failed_batch_entry_and_check_is_structured(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    unreadable = _make_book(raw / "a-unreadable.xlsx")
    _make_book(raw / "b-readable.xlsx", marker="readable")
    original_hash = _hash_file

    def hash_with_one_unreadable(path: Path) -> str:
        if path == unreadable:
            raise PermissionError("fixture denies read")
        return original_hash(path)

    with patch("excel_to_act.steps.step1.workflow._hash_file", side_effect=hash_with_one_unreadable):
        batch = convert_directory(raw, tmp_path / "converted")
    assert batch["status"] == "fail"
    assert batch["metrics"]["input_count"] == 2
    failed = next(entry for entry in batch["entries"] if entry["relative_path"] == "a-unreadable.xlsx")
    readable = next(entry for entry in batch["entries"] if entry["relative_path"] == "b-readable.xlsx")
    assert failed["sha256"] is None and failed["status"] == "fail"
    assert failed["diagnostics"][0]["code"] == "source_hash_failed"
    assert readable["status"] == "pass"
    assert readable["metrics_state"] == "measured"
    assert failed["metrics_state"] == "unavailable"
    assert failed["metrics"]["logical_objects_total"] is None
    assert (tmp_path / "converted" / batch["artifacts"][0]["path"]).is_file()
    assert (tmp_path / "converted" / batch["entries"][0]["run_path"] / "handoff.json").is_file()

    readable_run = tmp_path / "converted" / readable["run_path"]

    def deny_fresh_hash(path: Path) -> str:
        if path == unreadable:
            raise PermissionError("source became unreadable")
        return original_hash(path)

    with patch("excel_to_act.steps.step1.workflow._hash_file", side_effect=deny_fresh_hash):
        unreadable_check = execute_tool("step1.check", tmp_path / "converted" / failed["run_path"])
        readable_check = execute_tool("step1.check", readable_run)
        readable_source = _read(readable_run / "source.json")["source_path"]
        with patch("excel_to_act.steps.step1.workflow._hash_file", side_effect=lambda path: (_ for _ in ()).throw(PermissionError("source became unreadable")) if str(path) == readable_source else original_hash(path)):
            unreadable_after_convert = execute_tool("step1.check", readable_run)
    assert unreadable_check["status"] == "fail"
    assert any(item["code"] == "source_hash_unavailable" for item in unreadable_check["diagnostics"])
    assert readable_check["status"] == "pass"
    assert unreadable_after_convert["status"] == "fail"
    assert any(item["code"] == "source_hash_unreadable" for item in unreadable_after_convert["diagnostics"])


def test_raw_facts_normalized_fidelity_and_persisted_recovery(tmp_path: Path) -> None:
    source, run, converted = _convert_single(tmp_path)
    assert converted["status"] == "pass"
    facts = _read(run / "source_facts.json")
    assert facts["date_system"] == "1904"
    cells = {(cell["sheet"], cell["address"]): cell for cell in facts["cells"]}
    number = cells[("Facts", "B2")]
    assert number["raw_value_text"] == "12345678901234567890.1234"
    assert number["normalized_value"] != number["raw_value_text"]
    assert number["workbook_date_system"] == "1904"
    assert cells[("Facts", "C2")]["date_serial_text"] == "43830"
    shared_master = cells[("Facts", "D2")]
    assert shared_master["raw_formula_attributes"] == {"t": "shared", "ref": "D2:D3", "si": "0"}
    assert shared_master["cached_text"] == "0" and shared_master["normalized_cached_value"] == 0
    assert shared_master["normalized_cached_value_available"] is True
    assert cells[("Facts", "D3")]["normalized_cached_value"] is False
    assert cells[("Facts", "D4")]["normalized_cached_value"] == ""
    assert cells[("Facts", "D5")]["formula_present"] is True
    assert cells[("Facts", "D2")]["cached_text_present"] is True
    assert cells[("Facts", "D5")]["cached_text_present"] is False
    assert all(cell["workbook_date_system"] == "1904" for cell in facts["cells"])
    assert cells[("Facts", "D5")]["normalized_cached_value_available"] is False
    assert cells[("Facts", "E2")]["normalized_value"] == ""
    inventory = _read(run / "inventory.json")
    data_table_formula = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "D5")
    assert data_table_formula["kind"] == "formula"
    assert data_table_formula["data_type"] == "f"
    assert data_table_formula["formula"] is None
    assert data_table_formula["raw_formula_attributes"]["t"] == "dataTable"
    assert data_table_formula["raw_formula_attributes"]["ref"] == "D5:E6"

    names = [item for item in inventory["workbook_ranges"] if item["kind"] == "defined_name"]
    assert {(item["name"], item["metadata"]["scope"]) for item in names} == {
        ("Rate", "workbook"), ("Rate", "Facts")
    }
    assert execute_tool("step1.check", run)["status"] == "pass"
    cell = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "B2")
    cell["value"] = 7
    _write(run / "inventory.json", inventory)
    broken = execute_tool("step1.check", run)
    assert broken["status"] == "fail"
    assert broken["metrics"]["supported_fidelity_ratio"] < 1
    assert broken["metrics"]["supported_facts_exact"] == broken["metrics"]["supported_facts_total"] - 1
    assert any(item["code"] == "normalized_fidelity_mismatch" for item in broken["diagnostics"])

    repaired = execute_tool("inventory.extract", run)
    assert repaired["status"] == "pass"
    assert repaired["attempt"] == 1
    assert _read(run / "attempt_history.json")["attempts"][0]["improved"] is True
    assert execute_tool("step1.check", run)["status"] == "pass"


def test_legacy_inventory_source_facts_remain_unknown_until_step1_measures(tmp_path: Path) -> None:
    source = _make_book(tmp_path / "raw" / "legacy.xlsx")
    manifest = OpenpyxlWorkbookReader().read_manifest(source)
    legacy = OpenpyxlInventoryExtractor().extract(source, manifest)
    formula = next(cell for cell in legacy.sheets[0].cells if cell.address == "D2")
    assert formula.formula == "=B2*2"
    assert formula.cached_value == 0 and formula.cached_value_available is True
    assert formula.formula_present is None
    assert formula.cached_text_present is None
    assert formula.workbook_date_system is None

    previous_json = formula.model_dump(mode="json")
    for field in ("formula_present", "cached_text_present", "workbook_date_system"):
        previous_json.pop(field)
    loaded = CellInventory.model_validate(previous_json)
    assert loaded.formula_present is loaded.cached_text_present is loaded.workbook_date_system is None

    output = tmp_path / "legacy-artifacts"
    inspected = CliRunner().invoke(app, ["inspect", str(source), "--out", str(output)])
    assert inspected.exit_code == 0, inspected.output
    inventory = _read(output / "inventory.json")
    serialized = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "D2")
    assert serialized["formula"] == "=B2*2"
    assert serialized["cached_value"] == 0 and serialized["cached_value_available"] is True
    assert serialized["formula_present"] is serialized["cached_text_present"] is serialized["workbook_date_system"] is None


@pytest.mark.parametrize("formula", [False, True])
@pytest.mark.parametrize("index", ["-1", "9", "not-an-index", "", "0_0"])
def test_invalid_shared_string_indexes_are_source_errors(tmp_path: Path, index: str, formula: bool) -> None:
    source = _make_shared_string_index_case(tmp_path / f"invalid-{formula}-{index or 'empty'}.xlsx", index, formula=formula)
    with pytest.raises(SourceScanError, match="shared-string index"):
        scan_step1_source(source)


@pytest.mark.parametrize("index", ["+0", " +0 "])
def test_shared_string_indexes_keep_decimal_whitespace_and_sign_support(tmp_path: Path, index: str) -> None:
    source = _make_shared_string_index_case(tmp_path / "valid.xlsx", index, formula=False)
    scan = scan_step1_source(source)
    assert next(cell for cell in scan["cells"] if cell["address"] == "A1")["normalized_value"] == "漢字"


def test_shared_string_index_zero_and_mixed_batch_failure_are_handled_truthfully(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    valid = _make_shared_string_index_case(raw / "valid.xlsx", "0", formula=False)
    with zipfile.ZipFile(valid, "r") as package:
        sheet = ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
    formula_cell = next(item for item in sheet.iter() if item.tag == MAIN + "c" and item.attrib.get("r") == "A2")
    formula_cell.attrib["t"] = "s"
    for child in list(formula_cell):
        formula_cell.remove(child)
    ET.SubElement(formula_cell, MAIN + "f").text = "1+1"
    ET.SubElement(formula_cell, MAIN + "v").text = "0"
    _rewrite_package(valid, replacements={"xl/worksheets/sheet1.xml": ET.tostring(sheet, encoding="utf-8", xml_declaration=True)})
    _make_shared_string_index_case(raw / "invalid.xlsx", "0_0", formula=True)

    output = tmp_path / "converted"
    result = convert_directory(raw, output)
    entries = {entry["relative_path"]: entry for entry in result["entries"]}
    assert entries["valid.xlsx"]["status"] == "pass"
    assert entries["invalid.xlsx"]["status"] == "fail"
    assert entries["invalid.xlsx"]["metrics_state"] == "unavailable"
    bad_run = output / entries["invalid.xlsx"]["run_path"]
    assert execute_tool("step1.check", bad_run)["status"] == "fail"
    assert finalize_run(bad_run)["final_output"] is None
    assert not list((output / "final").rglob("promotion.json"))

    good_run = output / entries["valid.xlsx"]["run_path"]
    valid_facts = _read(good_run / "source_facts.json")
    fact_cells = {cell["address"]: cell for cell in valid_facts["cells"]}
    assert fact_cells["A1"]["normalized_value"] == "漢字"
    assert fact_cells["A2"]["normalized_cached_value"] == "漢字"
    assert fact_cells["A2"]["cached_text"] == "0"


def test_date_serial_requires_numeric_source_type_and_raw_number(tmp_path: Path) -> None:
    source = _make_typed_date_book(tmp_path / "raw" / "typed-dates.xlsx")
    scan = scan_step1_source(source)
    source_cells = {cell["address"]: cell for cell in scan["cells"]}

    text = source_cells["A1"]
    assert text["cell_type"] == "s"
    assert text["raw_value_text"] == "0" and text["normalized_value"] == "N/A"
    assert text["number_format"] == "yyyy-mm-dd"
    assert text["date_serial_text"] is None
    for address in ("B1", "C1", "D1", "E1", "F1", "I1"):
        assert source_cells[address]["date_serial_text"] is None, (address, source_cells[address])
    assert source_cells["B1"]["normalized_value"] == "42"
    assert source_cells["C1"]["normalized_value"] == "#N/A"
    assert source_cells["D1"]["normalized_value"] == "2024-01-02T00:00:00"
    assert source_cells["E1"]["normalized_value"] == "文本"
    assert source_cells["F1"]["normalized_value"] is True
    assert source_cells["G1"]["raw_value_text"] == "45293"
    assert source_cells["G1"]["date_serial_text"] == "45293"
    assert source_cells["G1"]["normalized_value"] == "2024-01-02T00:00:00"
    assert source_cells["H1"]["cached_text"] == "45293"
    assert source_cells["H1"]["date_serial_text"] == "45293"
    assert source_cells["H1"]["normalized_cached_value"] == "2024-01-02T00:00:00"
    assert source_cells["I1"]["cached_text"] == "N/A"
    assert source_cells["I1"]["normalized_cached_value"] == "N/A"

    converted = convert_directory(source.parent, tmp_path / "converted")
    assert converted["entries"][0]["status"] == "pass"
    run = tmp_path / "converted" / converted["entries"][0]["run_path"]
    inventory = _read(run / "inventory.json")
    cells = {cell["address"]: cell for cell in inventory["sheets"][0]["cells"]}
    assert cells["A1"]["value"] == "N/A" and cells["A1"]["raw_value_text"] == "0"
    assert cells["A1"]["ooxml_cell_type"] == "s" and cells["A1"]["date_serial_text"] is None
    assert cells["G1"]["date_serial_text"] == cells["H1"]["date_serial_text"] == "45293"
    assert execute_tool("step1.check", run)["status"] == "pass"

    cells["A1"]["date_serial_text"] = "0"
    _write(run / "inventory.json", inventory)
    rejected = execute_tool("step1.check", run)
    assert rejected["status"] == "fail"
    assert any(item["code"] == "normalized_fidelity_mismatch" for item in rejected["diagnostics"])
    assert execute_tool("inventory.extract", run)["status"] == "pass"
    assert finalize_run(run)["ready_for_next_step"] is True


def test_directory_order_ties_have_stable_entries_handoff_and_next_tool(tmp_path: Path) -> None:
    raw = tmp_path / "tie-raw"
    raw.mkdir()
    paths = [raw / "SS.xlsx", raw / "ß.xlsx"]
    for index, path in enumerate(paths):
        workbook = Workbook()
        workbook.active["A1"] = f"source {index}"
        workbook.save(path)
        workbook.close()

    original_rglob = Path.rglob

    def convert_in_order(enumeration: list[Path], output: Path) -> dict:
        def controlled_rglob(directory: Path, pattern: str):
            if directory.resolve() == raw.resolve() and pattern == "*":
                return iter(enumeration)
            return original_rglob(directory, pattern)

        with patch.object(Path, "rglob", controlled_rglob):
            return convert_directory(raw, output)

    forward = convert_in_order(paths, tmp_path / "out-forward")
    reverse = convert_in_order(list(reversed(paths)), tmp_path / "out-reverse")
    expected_order = ["SS.xlsx", "ß.xlsx"]
    assert [entry["relative_path"] for entry in forward["entries"]] == expected_order
    assert [entry["relative_path"] for entry in reverse["entries"]] == expected_order
    assert forward["next_tool"]["name"] == reverse["next_tool"]["name"] == "step1.check"
    selected = []
    for result in (forward, reverse):
        run = Path(result["next_tool"]["arguments"]["run"])
        selected.append(_read(run / "source.json")["relative_path"])
        batch = run.parents[3] / "batch_handoff.json"
        handoff = _read(batch)
        assert [entry["relative_path"] for entry in handoff["entries"]] == expected_order
        assert handoff["discovery"]["ordering"] == "relative_path_casefold_then_original"
    assert selected == ["SS.xlsx", "SS.xlsx"]


def test_visible_shared_and_inline_text_ignores_phonetics_and_preserves_raw_xml(tmp_path: Path) -> None:
    source = _make_book(tmp_path / "raw" / "phonetic.xlsx")
    original_shared, original_sheet = _add_phonetic_string_cases(source)
    converted = convert_directory(source.parent, tmp_path / "converted")
    run = tmp_path / "converted" / converted["entries"][0]["run_path"]
    assert converted["status"] == "pass"

    facts = _read(run / "source_facts.json")
    values = {cell["address"]: cell for cell in facts["cells"]}
    assert [values[f"A{row}"]["normalized_value"] for row in range(1, 7)] == [
        "漢字", "漢字", "漢字", "漢字", "  spacing  ", ""
    ]
    assert values["A1"]["raw_value_text"] == "0"
    assert values["A2"]["raw_value_text"] == "1"
    assert values["A3"]["inline_string_text"].count("かんじ") == 1
    assert values["A4"]["inline_string_text"].count("かんじ") == 1
    assert values["A3"]["raw_value_text"] is None

    assert values["D2"]["raw_formula_text"] == "B2*2"
    assert values["D2"]["cached_text"] == "0"
    assert values["D2"]["normalized_cached_value"] == 0
    validation = next(obj for obj in scan_step1_source(source)["objects"] if obj["kind"] == "data_validation")
    assert validation["details"]["formulas"] == ["0"]

    parts = {part["name"]: part for part in _read(run / "package_parts.json")}
    for name, original in (("xl/sharedStrings.xml", original_shared), ("xl/worksheets/sheet1.xml", original_sheet)):
        preserved = run / parts[name]["preserved_path"]
        assert preserved.read_bytes() == original

    inventory = _read(run / "inventory.json")
    inventory_cell = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "A1")
    inventory_cell["value"] = "漢字かんじ"
    _write(run / "inventory.json", inventory)
    broken = execute_tool("step1.check", run)
    assert broken["status"] == "fail"
    assert broken["metrics_state"] == "measured"
    assert any(item["code"] == "normalized_fidelity_mismatch" for item in broken["diagnostics"])
    assert execute_tool("inventory.extract", run)["status"] == "pass"
    assert finalize_run(run)["final_output"]


def test_missing_object_opaque_part_and_fresh_finalization_gate(tmp_path: Path) -> None:
    source, run, converted = _convert_single(tmp_path, custom_part=True)
    assert converted["status"] == "partial"
    assert converted["entries"][0]["metrics"]["opaque_parts"] == 1
    quality = execute_tool("step1.check", run)
    assert "data_table" in quality["coverage_scope"]["logical_object_kinds"]
    assert "style-only blank cells do not" in quality["coverage_scope"]["cell_unit"]
    assert quality["metrics"]["logical_objects_total"] == quality["metrics"]["logical_objects_accounted"]
    assert quality["metrics"]["package_parts_preserved"] == quality["metrics"]["package_parts_total"]

    inventory = _read(run / "inventory.json")
    inventory["sheets"][0]["cells"] = [cell for cell in inventory["sheets"][0]["cells"] if cell["address"] != "B2"]
    _write(run / "inventory.json", inventory)
    missing = execute_tool("step1.check", run)
    assert missing["status"] == "fail"
    assert missing["metrics"]["supported_fidelity_ratio"] < 1
    assert missing["metrics"]["traceability_ratio"] < 1

    repaired = execute_tool("inventory.extract", run)
    assert repaired["status"] == "partial"
    package_parts = _read(run / "package_parts.json")
    opaque = next(part for part in package_parts if part["name"] == "xl/custom-opaque.bin")
    missing_bytes = run / opaque["preserved_path"]
    missing_bytes.unlink()
    assert execute_tool("step1.check", run)["status"] == "fail"
    assert execute_tool("package.preserve", run)["status"] == "partial"
    assert (run / opaque["preserved_path"]).is_file()

    package_parts = _read(run / "package_parts.json")
    opaque = next(part for part in package_parts if part["name"] == "xl/custom-opaque.bin")
    opaque["opaque"] = False
    _write(run / "package_parts.json", package_parts)
    assert execute_tool("step1.check", run)["status"] == "fail"
    assert execute_tool("package.preserve", run)["status"] == "partial"

    finalized = finalize_run(run)
    assert finalized["ready_for_next_step"] is True
    assert finalized["final_output"]
    output = tmp_path / "converted" / finalized["final_output"]
    promoted_handoff = _read(output / "handoff.json")
    candidate_handoff = _read(run / "handoff.json")
    assert promoted_handoff["final_output"] == finalized["final_output"]
    assert promoted_handoff["coverage_scope"] == quality["coverage_scope"]
    assert "original_path" not in promoted_handoff["source"]
    assert promoted_handoff["run_path"].startswith("final/")
    assert candidate_handoff["final_output"] == finalized["final_output"]
    assert promoted_handoff["source"]["sha256"] == candidate_handoff["source"]["sha256"]
    assert execute_tool("step1.check", run)["status"] == "partial"
    assert _read(run / "handoff.json")["final_output"] == finalized["final_output"]

    current = _read(run / "inventory.json")
    next(cell for cell in current["sheets"][0]["cells"] if cell["address"] == "B2")["value"] = 9
    _write(run / "inventory.json", current)
    refused = finalize_run(run)
    assert refused["status"] == "fail"
    assert refused["final_output"] is None
    assert _read(run / "handoff.json")["final_output"] is None
    assert len(list((tmp_path / "converted" / "final").rglob("promotion.json"))) == 1


@pytest.mark.parametrize("damage", ["modified", "deleted", "missing_hash", "invalid_hash"])
def test_final_link_is_retained_only_for_hash_valid_promoted_payloads(tmp_path: Path, damage: str) -> None:
    _, run, _ = _convert_single(tmp_path)
    final = finalize_run(run)
    output = tmp_path / "converted" / final["final_output"]
    promotion_path = output / "promotion.json"
    promotion = _read(promotion_path)
    if damage == "modified":
        (output / "inventory.json").write_text("tampered\n", encoding="utf-8")
    elif damage == "deleted":
        (output / "inventory.json").unlink()
    elif damage == "missing_hash":
        promotion["files"].pop("inventory.json")
        _write(promotion_path, promotion)
    else:
        promotion["files"]["inventory.json"] = "bad-hash"
        _write(promotion_path, promotion)

    checked = execute_tool("step1.check", run)
    assert checked["status"] == "pass"
    assert _read(run / "handoff.json")["final_output"] is None
    assert output.exists()  # An invalid link is cleared; the previous final folder is preserved.


def test_final_link_survives_an_unrecorded_sidecar_in_final_directory(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    final = finalize_run(run)
    assert final["final_output"]
    output = tmp_path / "converted" / final["final_output"]
    note = output / "review-note.txt"
    note.write_text("Review annotation\n", encoding="utf-8")

    checked = execute_tool("step1.check", run)

    assert checked["status"] == "pass"
    assert _read(run / "handoff.json")["final_output"] == final["final_output"]
    assert note.read_text(encoding="utf-8") == "Review annotation\n"


def test_final_link_survives_output_relocation_when_payload_hashes_match(tmp_path: Path) -> None:
    _, run, converted = _convert_single(tmp_path)
    final = finalize_run(run)
    moved_output = tmp_path / "relocated-output"
    (tmp_path / "converted").replace(moved_output)
    moved_run = moved_output / converted["entries"][0]["run_path"]

    checked = execute_tool("step1.check", moved_run)
    assert checked["status"] == "pass"
    assert _read(moved_run / "handoff.json")["final_output"] == final["final_output"]


def test_finalize_refreshes_only_target_batch_diagnostics_and_human_report(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    _make_book(raw / "opaque.xlsx", custom_part=True, marker="opaque")
    _make_book(raw / "healthy.xlsx", marker="healthy")
    output = tmp_path / "converted"
    converted = convert_directory(raw, output)
    entries = {entry["relative_path"]: entry for entry in converted["entries"]}
    run = output / entries["opaque.xlsx"]["run_path"]
    other_before = entries["healthy.xlsx"]
    batch_path = run.parents[3] / "batch.json"
    handoff_path = run.parents[3] / "batch_handoff.json"
    initial_other = next(entry for entry in _read(batch_path)["entries"] if entry["run_id"] == other_before["run_id"])

    inventory = _read(run / "inventory.json")
    next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "B2")["value"] = 999
    _write(run / "inventory.json", inventory)
    failed = finalize_run(run)
    failed_codes = {item["code"] for item in failed["diagnostics"]}
    assert failed["status"] == "fail"
    assert "normalized_fidelity_mismatch" in failed_codes
    mismatch = next(item for item in failed["diagnostics"] if item["code"] == "normalized_fidelity_mismatch")

    failed_entry = next(entry for entry in _read(handoff_path)["entries"] if entry["run_id"] == failed["run_id"])
    assert failed_entry["status"] == "fail"
    assert failed_entry["diagnostics"] == failed["diagnostics"]
    assert {item["code"] for item in failed_entry["diagnostics"]} == failed_codes
    report = (run.parents[3] / "batch_handoff.md").read_text(encoding="utf-8")
    assert "normalized_fidelity_mismatch [error]" in report
    assert mismatch["message"] in report and "B2" in report
    unchanged_other = next(entry for entry in _read(batch_path)["entries"] if entry["run_id"] == other_before["run_id"])
    assert unchanged_other == initial_other

    repaired = execute_tool("inventory.extract", run)
    assert repaired["status"] == "partial"
    accepted = finalize_run(run)
    assert accepted["status"] == "partial" and accepted["ready_for_next_step"] is True
    current_entry = next(entry for entry in _read(handoff_path)["entries"] if entry["run_id"] == accepted["run_id"])
    current_codes = {item["code"] for item in current_entry["diagnostics"]}
    assert current_entry["diagnostics"] == accepted["diagnostics"]
    assert "opaque_parts_preserved" in current_codes
    assert "normalized_fidelity_mismatch" not in current_codes
    report = (run.parents[3] / "batch_handoff.md").read_text(encoding="utf-8")
    warning = next(item for item in accepted["diagnostics"] if item["code"] == "opaque_parts_preserved")
    assert "opaque_parts_preserved [warning]" in report
    assert warning["message"] in report and "normalized_fidelity_mismatch" not in report
    unchanged_other = next(entry for entry in _read(batch_path)["entries"] if entry["run_id"] == other_before["run_id"])
    assert unchanged_other == initial_other


def test_legacy_manifest_counts_workbook_and_sheet_scoped_name_declarations(tmp_path: Path) -> None:
    workbook = Workbook()
    facts = workbook.active
    facts.title = "Facts"
    facts["A1"] = 1
    other = workbook.create_sheet("Other")
    other["C2"] = 2
    workbook.defined_names.add(DefinedName("GlobalRate", attr_text="'Facts'!$A$1"))
    facts.defined_names.add(DefinedName("LocalRate", attr_text="'Facts'!$A$1"))
    facts.defined_names.add(DefinedName("LocalMulti", attr_text="'Facts'!$A$1,'Other'!$C$2"))
    facts.defined_names.add(DefinedName("LocalCross", attr_text="'Other'!$C$2"))
    source = tmp_path / "four-names.xlsx"
    workbook.save(source)
    workbook.close()

    manifest = OpenpyxlWorkbookReader().read_manifest(source)
    assert manifest.named_ranges_count == 4


def test_identityless_projection_multiset_and_parent_ownership_are_checked(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    source = _make_multi_sheet_book(raw / "multi.xlsx")
    converted = convert_directory(raw, tmp_path / "converted")
    run = tmp_path / "converted" / converted["entries"][0]["run_path"]
    assert execute_tool("step1.check", run)["status"] == "pass"

    inventory = _read(run / "inventory.json")
    names = [item for item in inventory["workbook_ranges"] if item["name"] == "Multi"]
    assert len(names) == 2
    assert sum(item["source_identity"] is None for item in names) == 1
    # Multiple destinations are a legitimate legacy projection: one source
    # name can emit a second address row without its own logical identity.
    assert execute_tool("step1.check", run)["status"] == "pass"

    extra = json.loads(json.dumps(inventory["sheets"][0]["cells"][0]))
    extra["source_identity"] = None
    extra["source_location"]["source_identity"] = None
    extra["address"] = "Z99"
    extra["row"] = 99
    extra["column"] = 26
    extra["source_location"]["address"] = "Z99"
    extra["source_location"]["object_id"] = "One!Z99"
    inventory["sheets"][0]["cells"].append(extra)
    _write(run / "inventory.json", inventory)
    checked = execute_tool("step1.check", run)
    assert checked["status"] == "fail"
    assert checked["next_tool"]["name"] == "inventory.extract"
    assert any(item["code"] == "inventory_projection_mismatch" for item in checked["diagnostics"])
    assert finalize_run(run)["final_output"] is None
    assert execute_tool("inventory.extract", run)["status"] == "pass"

    inventory = _read(run / "inventory.json")
    moved = inventory["sheets"][0]["cells"].pop(0)
    inventory["sheets"][1]["cells"].append(moved)
    _write(run / "inventory.json", inventory)
    checked = execute_tool("step1.check", run)
    assert checked["status"] == "fail"
    assert checked["next_tool"]["name"] == "inventory.extract"
    assert any(item["code"] == "inventory_parent_mismatch" for item in checked["diagnostics"])
    assert finalize_run(run)["final_output"] is None
    assert execute_tool("inventory.extract", run)["status"] == "pass"
    assert execute_tool("step1.check", run)["status"] == "pass"
    assert source.is_file()


def test_canonical_source_and_inventory_comparisons_distinguish_false_from_zero(tmp_path: Path) -> None:
    source, run, _ = _convert_single(tmp_path)
    sheet_xml = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(source) as package:
        root = ET.fromstring(package.read(sheet_xml))
    zero_cell = next(cell for cell in root.iter() if cell.tag == MAIN + "c" and cell.attrib.get("r") == "B7")
    zero_cell.find(MAIN + "v").text = "0"
    _rewrite_package(source, replacements={sheet_xml: ET.tostring(root, encoding="utf-8", xml_declaration=True)})

    # Reconvert against the now-zero source so each mutation is checked against
    # a fresh source-derived candidate.
    output = run.parents[4]
    raw = source.parent
    batch = convert_directory(raw, output / "zero-converted")
    run = output / "zero-converted" / batch["entries"][0]["run_path"]
    inventory = _read(run / "inventory.json")
    zero = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "B7")
    assert zero["value"] == 0 and type(zero["value"]) is int
    zero["value"] = False
    _write(run / "inventory.json", inventory)
    checked = execute_tool("step1.check", run)
    assert checked["status"] == "fail"
    mismatch = next(item for item in checked["metrics"]["fidelity_deviations"] if item["identity"] == zero["source_identity"])
    assert mismatch["fields"]["value"] == {"expected": 0, "actual": False}
    assert finalize_run(run)["final_output"] is None
    assert execute_tool("inventory.extract", run)["status"] == "pass"

    inventory = _read(run / "inventory.json")
    cached = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "D2")
    assert cached["cached_value"] == 0
    cached["cached_value"] = False
    _write(run / "inventory.json", inventory)
    checked = execute_tool("step1.check", run)
    assert checked["status"] == "fail"
    mismatch = next(item for item in checked["metrics"]["fidelity_deviations"] if item["identity"] == cached["source_identity"])
    assert mismatch["fields"]["cached_value"] == {"expected": 0, "actual": False}
    assert execute_tool("inventory.extract", run)["status"] == "pass"

    batch = convert_directory(source.parent, tmp_path / "facts-converted")
    run = tmp_path / "facts-converted" / batch["entries"][0]["run_path"]
    facts = _read(run / "source_facts.json")
    fact = next(cell for cell in facts["cells"] if cell["address"] == "B7")
    fact["normalized_value"] = False
    _write(run / "source_facts.json", facts)
    checked = execute_tool("step1.check", run)
    assert checked["status"] == "fail"
    assert any(item["code"] == "raw_source_facts_mismatch" for item in checked["diagnostics"])
    assert checked["next_tool"]["name"] == "source_facts.refresh"
    assert execute_tool("source_facts.refresh", run)["status"] == "pass"

    ledger = _read(run / "logical_objects.json")
    sheet_entry = next(item for item in ledger if item["kind"] == "sheet")
    sheet_entry["details"]["index"] = False
    _write(run / "logical_objects.json", ledger)
    checked = execute_tool("step1.check", run)
    assert checked["status"] == "fail"
    assert any(item["code"] == "logical_ledger_mismatch" for item in checked["diagnostics"])
    assert checked["next_tool"]["name"] == "source_facts.refresh"
    assert execute_tool("source_facts.refresh", run)["status"] == "pass"
    assert execute_tool("step1.check", run)["status"] == "pass"


def test_unvisited_relationship_parts_are_opaque_and_preserved(tmp_path: Path) -> None:
    source, _, _ = _convert_single(tmp_path / "default")
    unknown_rels = "xl/custom/_rels/item1.xml.rels"
    _rewrite_package(source, additions={unknown_rels: b"<broken"})

    converted = convert_directory(source.parent, tmp_path / "default" / "converted")
    entry = converted["entries"][0]
    run = tmp_path / "default" / "converted" / entry["run_path"]
    assert entry["status"] == "partial"
    check = execute_tool("step1.check", run)
    assert check["status"] == "partial" and _read(run / "handoff.json")["ready_for_next_step"] is True
    parts = _read(run / "package_parts.json")
    unknown = next(part for part in parts if part["name"] == unknown_rels)
    root_relationships = next(part for part in parts if part["name"] == "_rels/.rels")
    assert unknown["opaque"] is True
    assert root_relationships["opaque"] is False
    preserved = run / unknown["preserved_path"]
    assert preserved.read_bytes() == b"<broken"
    assert check["metrics"]["package_parts_preserved"] == check["metrics"]["package_parts_total"]
    assert finalize_run(run)["ready_for_next_step"] is True

    strict_batch = convert_directory(source.parent, tmp_path / "strict" / "converted", allow_opaque=False)
    strict_entry = strict_batch["entries"][0]
    strict_run = tmp_path / "strict" / "converted" / strict_entry["run_path"]
    assert strict_entry["status"] == "fail"
    strict_check = execute_tool("step1.check", strict_run)
    assert _read(strict_run / "handoff.json")["ready_for_next_step"] is False
    assert any(item["code"] == "opaque_parts_blocked" for item in strict_check["diagnostics"])

    known_rels_source = _make_book(tmp_path / "known-rels" / "raw" / "bad.xlsx")
    _rewrite_package(known_rels_source, replacements={"xl/_rels/workbook.xml.rels": b"<broken"})
    failed_batch = convert_directory(known_rels_source.parent, tmp_path / "known-rels" / "converted")
    failed = failed_batch["entries"][0]
    assert failed["status"] == "fail"
    assert any(item["code"] == "source_read_failed" for item in failed["diagnostics"])


def test_missing_internal_relationship_targets_fail_mixed_batch_and_never_finalize(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    cases = (
        ("broken-root.xlsx", "_rels/.rels", "customXml/missing.xml"),
        ("broken-workbook.xlsx", "xl/_rels/workbook.xml.rels", "../customXml/missing.xml"),
        ("broken-sheet.xlsx", "xl/worksheets/_rels/sheet1.xml.rels", "../customXml/missing.xml"),
        ("empty-target.xlsx", "xl/worksheets/_rels/sheet1.xml.rels", ""),
    )
    for filename, rels_part, target in cases:
        source = _make_book(raw / filename)
        _append_relationship(
            source,
            rels_part,
            {
                "Id": "rId99",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml",
                "Target": target,
            },
        )
        with pytest.raises(SourceScanError, match="missing internal target"):
            scan_step1_source(source)
    valid_source = _make_book(raw / "valid.xlsx", marker="healthy")

    output = tmp_path / "converted"
    batch = convert_directory(
        raw,
        output,
        traceability_min=0.0,
        fidelity_min=0.0,
        parsed_package_min=0.0,
        opaque_max=1.0,
    )
    entries = {entry["relative_path"]: entry for entry in batch["entries"]}
    assert batch["status"] == "fail"
    assert batch["metrics"]["input_count"] == 5
    assert entries["valid.xlsx"]["status"] == "pass"
    for filename, rels_part, _ in cases:
        entry = entries[filename]
        assert entry["status"] == "fail"
        assert any(item["code"] == "source_read_failed" for item in entry["diagnostics"])
        run = output / entry["run_path"]
        handoff = _read(run / "handoff.json")
        assert handoff["ready_for_next_step"] is False
        assert any(rels_part in item["message"] and "missing internal target" in item["message"] for item in handoff["diagnostics"])
        checked = execute_tool("step1.check", run)
        assert checked["status"] == "fail"
        assert finalize_run(run)["final_output"] is None
        assert _read(run / "handoff.json")["ready_for_next_step"] is False
    assert valid_source.is_file()
    assert not list((output / "final").rglob("promotion.json"))


def test_external_and_existing_opaque_relationship_targets_remain_supported(tmp_path: Path) -> None:
    source = _make_book(tmp_path / "raw" / "book.xlsx", custom_part=True)
    _append_relationship(
        source,
        "_rels/.rels",
        {
            "Id": "rId99",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml",
            "Target": "https://example.test/external-root.xml",
            "TargetMode": "External",
        },
    )
    _append_relationship(
        source,
        "xl/worksheets/_rels/sheet1.xml.rels",
        {
            "Id": "rId99",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml",
            "Target": "../custom-opaque.bin",
        },
    )
    _append_relationship(
        source,
        "xl/worksheets/_rels/sheet1.xml.rels",
        {
            "Id": "rId98",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink",
            "Target": "https://example.test/external-sheet.xml",
            "TargetMode": "External",
        },
    )
    unknown_rels = "xl/custom/_rels/item1.xml.rels"
    _rewrite_package(source, additions={unknown_rels: b"<unparsed-relationship-bytes"})

    fresh = scan_step1_source(source)
    existing_target = "xl/custom-opaque.bin"
    assert existing_target not in fresh["parsed_parts"]
    assert fresh["parts"][existing_target]["opaque"] is True
    assert unknown_rels not in fresh["parsed_parts"]
    assert fresh["parts"][unknown_rels]["opaque"] is True

    batch = convert_directory(source.parent, tmp_path / "converted")
    entry = batch["entries"][0]
    run = tmp_path / "converted" / entry["run_path"]
    assert entry["status"] == "partial"
    check = execute_tool("step1.check", run)
    assert check["status"] == "partial"
    parts = {part["name"]: part for part in _read(run / "package_parts.json")}
    assert (run / parts[existing_target]["preserved_path"]).read_bytes() == b"opaque-bytes"
    assert (run / parts[unknown_rels]["preserved_path"]).read_bytes() == b"<unparsed-relationship-bytes"
    assert finalize_run(run)["ready_for_next_step"] is True


def test_out_of_range_local_defined_name_fails_but_valid_scopes_remain(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    invalid_source = _make_book(raw / "invalid-local.xlsx")
    _make_book(raw / "valid-scopes.xlsx")
    workbook_part = "xl/workbook.xml"
    with zipfile.ZipFile(invalid_source) as package:
        root = ET.fromstring(package.read(workbook_part))
    local_name = next(
        item
        for item in root.iter()
        if item.tag == MAIN + "definedName" and item.attrib.get("localSheetId") is not None
    )
    local_name.attrib["localSheetId"] = "42"
    _rewrite_package(invalid_source, replacements={workbook_part: ET.tostring(root, encoding="utf-8", xml_declaration=True)})

    converted = convert_directory(raw, tmp_path / "converted")
    entries = {entry["relative_path"]: entry for entry in converted["entries"]}
    assert entries["invalid-local.xlsx"]["status"] == "fail"
    assert any(item["code"] == "source_read_failed" for item in entries["invalid-local.xlsx"]["diagnostics"])
    invalid_run = tmp_path / "converted" / entries["invalid-local.xlsx"]["run_path"]
    assert _read(invalid_run / "handoff.json")["ready_for_next_step"] is False
    assert finalize_run(invalid_run)["final_output"] is None

    assert entries["valid-scopes.xlsx"]["status"] == "pass"
    valid_run = tmp_path / "converted" / entries["valid-scopes.xlsx"]["run_path"]
    valid_names = [item for item in _read(valid_run / "inventory.json")["workbook_ranges"] if item["kind"] == "defined_name"]
    assert {(item["name"], item["metadata"]["scope"]) for item in valid_names} == {
        ("Rate", "workbook"), ("Rate", "Facts")
    }


def test_package_part_location_mismatch_is_repaired_and_gated(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    package_path = run / "package_parts.json"
    parts = _read(package_path)
    target = next(part for part in parts if part["name"] == "docProps/app.xml")
    target["source_location"]["ooxml_part"] = "xl/workbook.xml"
    target["source_location"]["object_id"] = "xl/workbook.xml"
    _write(package_path, parts)

    checked = execute_tool("step1.check", run)
    assert checked["status"] == "fail"
    assert checked["next_tool"]["name"] == "package.preserve"
    assert any(item["code"] == "package_parts_mismatch" for item in checked["diagnostics"])
    refused = finalize_run(run)
    assert refused["status"] == "fail" and refused["final_output"] is None
    assert execute_tool("package.preserve", run)["status"] == "pass"
    assert execute_tool("step1.check", run)["status"] == "pass"
    repaired = next(part for part in _read(package_path) if part["name"] == "docProps/app.xml")
    assert repaired["source_location"] == {
        "ooxml_part": "docProps/app.xml",
        "object_type": "package_part",
        "object_id": "docProps/app.xml",
    }


def test_package_ratio_thresholds_are_independent_and_use_package_denominators(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    _make_book(raw / "book.xlsx", custom_part=True)

    def convert_case(name: str, *, parsed_min: float, opaque_max: float, allow_opaque: bool = True):
        output = tmp_path / name
        result = convert_directory(raw, output, parsed_package_min=parsed_min, opaque_max=opaque_max, allow_opaque=allow_opaque)
        run = output / result["entries"][0]["run_path"]
        return result, run, _read(run / "quality.json")

    baseline, baseline_run, baseline_quality = convert_case("baseline", parsed_min=0.0, opaque_max=1.0)
    assert baseline["status"] == "partial"
    baseline_check = execute_tool("step1.check", baseline_run)
    assert baseline_check["thresholds"]["parsed_package_min"] == 0.0
    assert baseline_check["thresholds"]["opaque_max"] == 1.0
    metrics = baseline_quality["metrics"]
    parsed_ratio = metrics["parsed_package_parts_ratio"]
    opaque_rate = metrics["opaque_rate"]
    total = metrics["package_parts_total"]
    assert total > 0 and 0 < parsed_ratio < 1 and 0 < opaque_rate < 1
    assert metrics["parsed_package_parts"] < total and metrics["opaque_parts"] > 0
    assert baseline_quality["coverage_scope"]["parsed_package_parts_ratio"]
    assert "all non-directory OOXML ZIP members" in baseline_quality["coverage_scope"]["parsed_package_parts_ratio"]

    # A scanner-read member marked opaque remains opaque and cannot inflate
    # the parsed numerator.
    from excel_to_act.ingest.ooxml_package import OPAQUE_MARKERS

    scanner_markers = {**OPAQUE_MARKERS, "xl/worksheets/sheet1.xml": "test opaque marker"}
    with patch("excel_to_act.steps.step1.source_scan.OPAQUE_MARKERS", scanner_markers):
        fresh_scan = scan_step1_source(raw / "book.xlsx")
        assert "xl/worksheets/sheet1.xml" in fresh_scan["parsed_parts"]
        assert fresh_scan["parts"]["xl/worksheets/sheet1.xml"]["opaque"] is True
        marked_batch = convert_directory(raw, tmp_path / "marked-part", parsed_package_min=0.0, opaque_max=1.0)
    marked_quality = _read(tmp_path / "marked-part" / marked_batch["entries"][0]["run_path"] / "quality.json")
    assert marked_quality["metrics"]["parsed_package_parts"] == sum(
        name in fresh_scan["parsed_parts"] and not info["opaque"]
        for name, info in fresh_scan["parts"].items()
    )
    assert marked_quality["metrics"]["opaque_parts"] >= 1

    boundary, boundary_run, boundary_quality = convert_case("boundary", parsed_min=parsed_ratio, opaque_max=opaque_rate)
    assert boundary["status"] == "partial"
    assert boundary_quality["thresholds"]["parsed_package_min"] == parsed_ratio
    assert boundary_quality["thresholds"]["opaque_max"] == opaque_rate
    boundary_final = finalize_run(boundary_run)
    assert boundary_final["ready_for_next_step"] is True
    assert boundary_final["thresholds"] == boundary_quality["thresholds"]

    loose, _, loose_quality = convert_case("loose", parsed_min=0.0, opaque_max=1.0)
    assert loose["status"] == "partial"
    assert loose_quality["thresholds"]["parsed_package_min"] == 0.0
    assert loose_quality["thresholds"]["opaque_max"] == 1.0

    parsed_too_strict, parsed_run, parsed_quality = convert_case("parsed-strict", parsed_min=parsed_ratio + 0.01, opaque_max=1.0)
    assert parsed_too_strict["status"] == "fail"
    assert any(item["code"] == "parsed_package_threshold_not_met" for item in parsed_quality["diagnostics"])
    assert not any(item["code"] == "opaque_rate_threshold_not_met" for item in parsed_quality["diagnostics"])
    assert finalize_run(parsed_run)["final_output"] is None

    opaque_too_strict, _, opaque_quality = convert_case("opaque-strict", parsed_min=0.0, opaque_max=opaque_rate - 0.01)
    assert opaque_too_strict["status"] == "fail"
    assert any(item["code"] == "opaque_rate_threshold_not_met" for item in opaque_quality["diagnostics"])
    assert not any(item["code"] == "parsed_package_threshold_not_met" for item in opaque_quality["diagnostics"])

    independent, _, independent_quality = convert_case("independent", parsed_min=0.0, opaque_max=opaque_rate - 0.01)
    assert independent["status"] == "fail"
    assert any(item["code"] == "opaque_rate_threshold_not_met" for item in independent_quality["diagnostics"])
    assert not any(item["code"] == "parsed_package_threshold_not_met" for item in independent_quality["diagnostics"])

    both_strict, _, both_quality = convert_case("both-strict", parsed_min=parsed_ratio + 0.01, opaque_max=opaque_rate - 0.01)
    assert both_strict["status"] == "fail"
    both_codes = {item["code"] for item in both_quality["diagnostics"]}
    assert {"parsed_package_threshold_not_met", "opaque_rate_threshold_not_met"} <= both_codes

    strict_boolean, _, strict_quality = convert_case("strict-boolean", parsed_min=0.0, opaque_max=1.0, allow_opaque=False)
    assert strict_boolean["status"] == "fail"
    assert any(item["code"] == "opaque_parts_blocked" for item in strict_quality["diagnostics"])

    endpoints, _, endpoint_quality = convert_case("endpoints", parsed_min=1.0, opaque_max=0.0)
    assert endpoints["status"] == "fail"
    assert endpoint_quality["thresholds"]["parsed_package_min"] == 1.0
    assert endpoint_quality["thresholds"]["opaque_max"] == 0.0

    cli = CliRunner()
    invalid_cli = cli.invoke(app, ["step1", "convert", str(raw), "--out", str(tmp_path / "invalid-cli"), "--parsed-package-min", "1.01"])
    assert invalid_cli.exit_code == 1
    invalid_envelope = json.loads(invalid_cli.output)
    assert invalid_envelope["status"] == "error"
    assert invalid_envelope["diagnostics"][0]["code"] == "invalid_step1_options"
    invalid_opaque_cli = cli.invoke(app, ["step1", "convert", str(raw), "--out", str(tmp_path / "invalid-opaque-cli"), "--opaque-max", "1.01"])
    assert invalid_opaque_cli.exit_code == 1
    invalid_opaque_envelope = json.loads(invalid_opaque_cli.output)
    assert invalid_opaque_envelope["status"] == "error"
    assert invalid_opaque_envelope["diagnostics"][0]["code"] == "invalid_step1_options"
    invalid_bool = convert_directory(raw, tmp_path / "invalid-bool", parsed_package_min=True)
    assert invalid_bool["status"] == "error"
    assert invalid_bool["diagnostics"][0]["code"] == "invalid_step1_options"


def test_package_threshold_defaults_and_invalid_source_values_are_stable(tmp_path: Path) -> None:
    source, run, _ = _convert_single(tmp_path, custom_part=True)
    source_json_path = run / "source.json"
    original = _read(source_json_path)
    original["thresholds"].pop("parsed_package_min")
    original["thresholds"].pop("opaque_max")
    _write(source_json_path, original)
    legacy_quality = execute_tool("step1.check", run)
    assert legacy_quality["status"] == "partial"
    stored_quality = _read(run / "quality.json")
    assert stored_quality["thresholds"]["parsed_package_min"] == 0.0
    assert stored_quality["thresholds"]["opaque_max"] == 1.0

    stable_source = _read(source_json_path)
    invalid_cases = (
        ("parsed_package_min", True),
        ("parsed_package_min", 1.01),
        ("opaque_max", False),
        ("opaque_max", -0.01),
    )
    cli = CliRunner()
    for threshold, invalid_value in invalid_cases:
        corrupted = json.loads(json.dumps(stable_source))
        corrupted["thresholds"][threshold] = invalid_value
        _write(source_json_path, corrupted)
        checked = cli.invoke(app, ["step1", "check", "--run", str(run)])
        assert checked.exit_code == 1
        check_envelope = json.loads(checked.output)
        assert check_envelope["status"] == "fail"
        assert any(item["code"] == "source_record_unreadable" for item in check_envelope["diagnostics"])
        finalized = cli.invoke(app, ["step1", "finalize", "--run", str(run)])
        assert finalized.exit_code == 1
        final_envelope = json.loads(finalized.output)
        assert final_envelope["status"] == "fail" and final_envelope["final_output"] is None
        assert not list((run.parents[5] / "final").rglob("promotion.json"))
        assert _read(run / "handoff.json")["ready_for_next_step"] is False
    source_json_path.write_text(json.dumps(stable_source), encoding="utf-8")
    assert finalize_run(run)["ready_for_next_step"] is True

    bad_source = _make_book(tmp_path / "bad-source" / "raw" / "bad.xlsx")
    bad_source.write_bytes(b"not an OOXML workbook")
    failed_batch = convert_directory(bad_source.parent, tmp_path / "bad-source" / "converted")
    failed = failed_batch["entries"][0]
    assert failed["status"] == "fail"
    bad_run = tmp_path / "bad-source" / "converted" / failed["run_path"]
    bad_metrics = _read(bad_run / "quality.json")["metrics"]
    assert failed["metrics_state"] == "unavailable"
    assert _read(bad_run / "quality.json")["metrics_state"] == "unavailable"
    assert bad_metrics["package_parts_total"] is None
    assert bad_metrics["parsed_package_parts_ratio"] is None
    assert finalize_run(bad_run)["final_output"] is None


def test_unavailable_source_metrics_stay_null_across_cli_and_handoffs(tmp_path: Path) -> None:
    source, run, _ = _convert_single(tmp_path)
    catalogue_result = CliRunner().invoke(app, ["step1", "tools"])
    assert catalogue_result.exit_code == 0
    catalogue = json.loads(catalogue_result.output)
    assert catalogue["metrics_contract"]["state_field"] == "metrics_state"
    assert "null" in catalogue["metrics_contract"]["states"]["unavailable"]

    source_record = _read(run / "source.json")
    source_record["thresholds"] = {"traceability": 0.0, "fidelity": 0.0, "parsed_package_min": 0.0, "opaque_max": 1.0}
    _write(run / "source.json", source_record)
    source.write_bytes(b"not an OOXML workbook")

    runner = CliRunner()
    commands = (
        ["step1", "check", "--run", str(run)],
        ["step1", "coverage", "--run", str(run)],
        ["step1", "fidelity", "--run", str(run)],
        ["step1", "tool", "report.handoff", "--run", str(run)],
        ["step1", "tool", "inventory.extract", "--run", str(run)],
        ["step1", "auto-recover", "--run", str(run), "--max-attempts", "1"],
        ["step1", "finalize", "--run", str(run)],
    )
    for command in commands:
        result = runner.invoke(app, command)
        assert result.exit_code == 1, (command, result.output)
        envelope = json.loads(result.output)
        assert envelope["metrics_state"] == "unavailable", (command, envelope)
        assert envelope["status"] in {"fail", "error"}
        for key, value in envelope["metrics"].items():
            if key != "input_count" and key != "tool_count":
                assert value is None, (command, key, value)
        assert envelope.get("ready_for_next_step") is not True
        if command[1] == "finalize":
            assert envelope["final_output"] is None

    assert any(item["code"] == "source_changed" for item in _read(run / "quality.json")["diagnostics"])
    handoff = _read(run / "handoff.json")
    assert handoff["metrics_state"] == "unavailable"
    assert handoff["metrics"]["logical_objects_total"] is None
    assert "unavailable" in (run / "handoff.md").read_text(encoding="utf-8")
    assert not list((run.parents[5] / "final").rglob("promotion.json"))


def test_initial_malformed_and_unsupported_inputs_have_unavailable_metrics(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    _make_book(raw / "broken.xlsx").write_bytes(b"not an OOXML workbook")
    (raw / "legacy.xls").write_bytes(b"legacy workbook")
    out = tmp_path / "converted"
    batch = convert_directory(raw, out)
    assert batch["status"] == "fail"
    assert {entry["relative_path"] for entry in batch["entries"]} == {"broken.xlsx", "legacy.xls"}
    for entry in batch["entries"]:
        assert entry["status"] == "fail"
        assert entry["metrics_state"] == "unavailable"
        assert entry["metrics"]["logical_objects_total"] is None
        assert entry["metrics"]["package_parts_total"] is None
        run = out / entry["run_path"]
        handoff = _read(run / "handoff.json")
        assert handoff["metrics_state"] == "unavailable"
        assert handoff["metrics"]["supported_fidelity_ratio"] is None
        assert "unavailable" in (run / "handoff.md").read_text(encoding="utf-8")
    batch_handoff = _read(out / batch["artifacts"][0]["path"])
    assert {entry["metrics_state"] for entry in batch_handoff["entries"]} == {"unavailable"}
    assert "unavailable" in (out / batch["artifacts"][1]["path"]).read_text(encoding="utf-8")


def test_duplicate_candidate_identities_block_finalize_and_offer_real_repairs(tmp_path: Path) -> None:
    runner = CliRunner()
    _, run, _ = _convert_single(tmp_path)
    repair_cases = (
        ("inventory.json", "duplicate_inventory_identity", "inventory.extract", "sheets.0.cells"),
        ("source_facts.json", "duplicate_source_fact_identity", "source_facts.refresh", "cells"),
        ("logical_objects.json", "duplicate_logical_object_identity", "source_facts.refresh", None),
        ("package_parts.json", "duplicate_package_part_name", "package.preserve", None),
    )
    for index, (filename, diagnostic_code, repair_tool, nested_key) in enumerate(repair_cases):
        _, run, _ = _convert_single(tmp_path / f"case-{index}")
        artifact = run / filename
        original = _read(artifact)
        if nested_key == "sheets.0.cells":
            target = original["sheets"][0]["cells"]
        elif nested_key:
            target = original[nested_key]
        else:
            target = original
        target.append(json.loads(json.dumps(target[0])))
        _write(artifact, original)

        checked = runner.invoke(app, ["step1", "check", "--run", str(run)])
        assert checked.exit_code == 1
        check_envelope = json.loads(checked.output)
        assert check_envelope["status"] == "fail"
        assert check_envelope["next_tool"]["name"] == repair_tool
        assert any(item["code"] == diagnostic_code for item in check_envelope["diagnostics"])
        refused = runner.invoke(app, ["step1", "finalize", "--run", str(run)])
        assert refused.exit_code == 1
        final_envelope = json.loads(refused.output)
        assert final_envelope["status"] == "fail" and final_envelope["final_output"] is None
        assert not list((run.parents[5] / "final").rglob("promotion.json"))

        repaired = runner.invoke(app, ["step1", "tool", repair_tool, "--run", str(run)])
        assert repaired.exit_code == 0 and json.loads(repaired.output)["status"] == "pass"
        checked_again = runner.invoke(app, ["step1", "check", "--run", str(run)])
        assert checked_again.exit_code == 0 and json.loads(checked_again.output)["status"] == "pass"


def test_invalid_candidate_collection_shapes_offer_real_repairs(tmp_path: Path) -> None:
    repair_cases = (
        ("inventory.json", [], "inventory_unreadable", "inventory.extract"),
        ("source_facts.json", [], "source_facts_invalid", "source_facts.refresh"),
        ("logical_objects.json", ["invalid row"], "logical_ledger_invalid", "source_facts.refresh"),
        ("package_parts.json", [None], "package_ledger_invalid", "package.preserve"),
    )
    for index, (filename, invalid_value, diagnostic_code, repair_tool) in enumerate(repair_cases):
        _, run, _ = _convert_single(tmp_path / f"shape-{index}")
        _write(run / filename, invalid_value)
        checked = execute_tool("step1.check", run)
        assert checked["status"] == "fail"
        assert checked["next_tool"]["name"] == repair_tool
        assert any(item["code"] == diagnostic_code for item in checked["diagnostics"])
        refused = finalize_run(run)
        assert refused["status"] == "fail" and refused["final_output"] is None
        assert not list((run.parents[5] / "final").rglob("promotion.json"))

        repaired = execute_tool(repair_tool, run)
        assert repaired["status"] == "pass"
        assert execute_tool("step1.check", run)["status"] == "pass"


def test_malformed_source_facts_candidate_returns_json_and_is_not_promoted(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    facts_path = run / "source_facts.json"
    original = facts_path.read_bytes()
    runner = CliRunner()
    for malformed in ([], {"date_system": "1900", "cells": [1]}):
        _write(facts_path, malformed)
        checked = runner.invoke(app, ["step1", "check", "--run", str(run)])
        assert checked.exit_code == 1
        check_envelope = json.loads(checked.output)
        assert check_envelope["status"] == "fail"
        assert check_envelope["next_tool"]["name"] == "source_facts.refresh"
        assert any(item["code"] == "source_facts_invalid" for item in check_envelope["diagnostics"])

        finalized = runner.invoke(app, ["step1", "finalize", "--run", str(run)])
        assert finalized.exit_code == 1
        final_envelope = json.loads(finalized.output)
        assert final_envelope["status"] == "fail" and final_envelope["final_output"] is None
        assert final_envelope["next_tool"]["name"] == "source_facts.refresh"
        assert (run / "quality.json").is_file() and (run / "handoff.json").is_file()
        assert not list((tmp_path / "converted" / "final").rglob("promotion.json"))
    facts_path.write_bytes(original)


def test_source_change_malformed_xml_and_invalid_run_return_structured_errors(tmp_path: Path) -> None:
    source, run, _ = _convert_single(tmp_path / "changed")
    source.write_bytes(source.read_bytes() + b"tamper")
    result = execute_tool("step1.check", run)
    assert result["status"] == "fail"
    assert any(item["code"] == "source_changed" for item in result["diagnostics"])
    assert finalize_run(run)["status"] == "fail"

    broken_source = _make_book(tmp_path / "malformed" / "raw" / "bad.xlsx")
    corrupted = tmp_path / "malformed" / "raw" / "broken.xlsx"
    with zipfile.ZipFile(broken_source, "r") as original, zipfile.ZipFile(corrupted, "w", zipfile.ZIP_DEFLATED) as target:
        for info in original.infolist():
            payload = b"<broken" if info.filename == "xl/workbook.xml" else original.read(info.filename)
            target.writestr(info, payload)
    malformed_batch = convert_directory(corrupted.parent, tmp_path / "malformed" / "output")
    malformed = next(entry for entry in malformed_batch["entries"] if entry["relative_path"] == "broken.xlsx")
    malformed_run = tmp_path / "malformed" / "output" / malformed["run_path"]
    assert malformed["status"] == "fail"
    assert execute_tool("step1.check", malformed_run)["status"] == "fail"

    runner = CliRunner()
    missing_run = runner.invoke(app, ["step1", "check", "--run", str(tmp_path / "missing-run")])
    assert missing_run.exit_code == 1
    envelope = json.loads(missing_run.output)
    assert envelope["status"] == "error"
    assert envelope["diagnostics"][0]["code"] == "run_unavailable"


def test_invalid_source_metadata_is_json_error_and_never_promotes(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    runner = CliRunner()
    source_path = run / "source.json"
    original = source_path.read_bytes()
    final_root = run.parents[5] / "final"
    commands = [
        ["step1", "check", "--run", str(run)],
        ["step1", "coverage", "--run", str(run)],
        ["step1", "fidelity", "--run", str(run)],
        ["step1", "tool", "report.handoff", "--run", str(run)],
        ["step1", "tool", "inventory.extract", "--run", str(run)],
        ["step1", "auto-recover", "--run", str(run), "--max-attempts", "1"],
        ["step1", "finalize", "--run", str(run)],
    ]
    for bad_metadata in (b"{", b"[]", b"{}"):
        source_path.write_bytes(bad_metadata)
        for command in commands:
            result = runner.invoke(app, command)
            assert result.exit_code == 1, (command, result.output)
            envelope = json.loads(result.output)
            assert envelope["status"] in {"fail", "error"}, (command, envelope)
            assert envelope.get("ready_for_next_step") is not True
            assert any(item["code"] == "source_record_unreadable" for item in envelope["diagnostics"])
        assert not final_root.exists() or not list(final_root.rglob("promotion.json"))
        assert (run / "quality.json").is_file()
        assert (run / "handoff.json").is_file()
    source_path.write_bytes(original)


def test_no_progress_attempts_stop_and_do_not_repeat_forever(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    inventory = _read(run / "inventory.json")
    next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "B2")["value"] = 9
    _write(run / "inventory.json", inventory)
    with patch("excel_to_act.steps.step1.workflow._tool_action", return_value={"artifacts": [], "metrics": {}}):
        first = execute_tool("inventory.extract", run)
        second = execute_tool("inventory.extract", run)
    assert first["status"] == "fail"
    assert second["status"] == "blocked"
    assert any(item["code"] == "no_progress_stop" for item in second["diagnostics"])
    history = _read(run / "attempt_history.json")
    assert len(history["attempts"]) == 2
    assert history["no_progress_stopped"] is True
    checked = execute_tool("step1.check", run)
    auto = auto_recover(run)
    handoff = _read(run / "handoff.json")
    assert checked["retryable"] is False and checked["next_tool"] is None
    assert checked["stop_reason"] == auto["stop_reason"] == handoff["stop_reason"] == "no_progress_stop"
    assert checked["remaining_tool"] == auto["remaining_tool"] == handoff["remaining_tool"]
    assert auto["attempts_used"] == 0 and auto["recovery_stopped"] is True


def test_auto_recover_exhaustion_is_consistent_and_blocks_suggested_tool(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    facts = _read(run / "source_facts.json")
    facts["cells"][0]["raw_value_text"] = "tampered"
    _write(run / "source_facts.json", facts)
    inventory = _read(run / "inventory.json")
    cell = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "B2")
    cell["value"] = 88
    _write(run / "inventory.json", inventory)

    initial = execute_tool("step1.check", run)
    assert initial["next_tool"]["name"] == "source_facts.refresh"
    recovered = auto_recover(run, max_attempts=1)
    checked = execute_tool("step1.check", run)
    handoff = _read(run / "handoff.json")
    assert recovered["status"] == checked["status"] == "fail"
    assert recovered["recovery_stopped"] is checked["recovery_stopped"] is handoff["recovery_stopped"] is True
    assert recovered["stop_reason"] == checked["stop_reason"] == handoff["stop_reason"] == "attempt_limit_reached"
    assert recovered["retryable"] is checked["retryable"] is handoff["ready_for_next_step"] is False
    assert recovered["next_tool"] is checked["next_tool"] is None
    assert recovered["remaining_tool"]["name"] == checked["remaining_tool"]["name"] == "inventory.extract"
    assert len(recovered["actions"]) == 1

    attempt_count = len(_read(run / "attempt_history.json")["attempts"])
    blocked = execute_tool("inventory.extract", run)
    repeated = auto_recover(run, max_attempts=1)
    assert blocked["status"] == "blocked" and blocked["retryable"] is False
    assert blocked["stop_reason"] == repeated["stop_reason"] == "attempt_limit_reached"
    assert len(_read(run / "attempt_history.json")["attempts"]) == attempt_count
    assert repeated["attempts_used"] == 0 and repeated["next_tool"] is None

    raised_limit_request = auto_recover(run, max_attempts=3)
    assert raised_limit_request["stop_reason"] == "attempt_limit_reached"
    assert raised_limit_request["attempts_used"] == 0
    assert _read(run / "attempt_history.json")["attempt_limit"] == 1


def test_manifest_repair_success_is_recorded_as_a_cumulative_attempt(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    (run / "workbook_manifest.json").write_text("{", encoding="utf-8")
    before = execute_tool("step1.check", run)
    assert before["next_tool"]["name"] == "manifest.read"

    repaired = execute_tool("manifest.read", run)
    history = _read(run / "attempt_history.json")
    assert repaired["status"] == "pass"
    assert repaired["attempt"] == 1 and repaired["attempt_limit"] == 3
    assert history["attempts"][0]["tool"] == "manifest.read"
    assert history["attempts"][0]["improved"] is True


def test_manifest_repair_failures_share_persisted_attempt_and_no_progress_bounds(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    (run / "workbook_manifest.json").write_text("{", encoding="utf-8")
    before = execute_tool("step1.check", run)
    assert before["next_tool"]["name"] == "manifest.read"

    with patch("excel_to_act.steps.step1.workflow._tool_action", side_effect=OSError("read-only manifest")):
        first = execute_tool("manifest.read", run)
    assert first["status"] == "error"
    assert first["attempt"] == 1
    history = _read(run / "attempt_history.json")
    assert len(history["attempts"]) == 1 and history["attempts"][0]["error"]

    with patch("excel_to_act.steps.step1.workflow._tool_action", return_value={"artifacts": []}):
        second = execute_tool("manifest.read", run)
    assert second["recovery_stopped"] is True
    assert second["stop_reason"] == "no_progress_stop"
    assert second["attempt"] == 2
    assert len(_read(run / "attempt_history.json")["attempts"]) == 2

    checked = execute_tool("step1.check", run)
    handoff = _read(run / "handoff.json")
    auto = auto_recover(run)
    blocked = execute_tool("manifest.read", run)
    assert checked["stop_reason"] == handoff["stop_reason"] == auto["stop_reason"] == blocked["stop_reason"] == "no_progress_stop"
    assert auto["actions"] == [] and auto["attempts_total"] == 2
    assert blocked["status"] == "blocked" and blocked["attempt"] == 2


@pytest.mark.parametrize("summary_flag", [None, False], ids=["missing-flag", "false-flag"])
def test_last_no_progress_count_keeps_recovery_stopped_without_summary_flag(
    tmp_path: Path, summary_flag: bool | None
) -> None:
    _, run, _ = _convert_single(tmp_path)
    manifest_path = run / "workbook_manifest.json"
    manifest_path.write_text("{", encoding="utf-8")
    assert execute_tool("step1.check", run)["next_tool"]["name"] == "manifest.read"

    with patch("excel_to_act.steps.step1.workflow._tool_action", side_effect=OSError("read-only manifest")):
        first = execute_tool("manifest.read", run)
    with patch("excel_to_act.steps.step1.workflow._tool_action", return_value={"artifacts": []}):
        second = execute_tool("manifest.read", run)
    assert first["status"] == "error" and first["attempt"] == 1
    assert second["recovery_stopped"] is True and second["attempt"] == 2

    history_path = run / "attempt_history.json"
    history = _read(history_path)
    assert history["attempts"][-1]["no_progress_count"] == 2
    if summary_flag is None:
        history.pop("no_progress_stopped")
    else:
        history["no_progress_stopped"] = summary_flag
    _write(history_path, history)
    persisted_history = history_path.read_bytes()
    persisted_manifest = manifest_path.read_bytes()

    checked = execute_tool("step1.check", run)
    handoff = _read(run / "handoff.json")
    auto = auto_recover(run)
    blocked = execute_tool("manifest.read", run)
    finalized = finalize_run(run)

    assert checked["recovery_stopped"] is True and checked["retryable"] is False
    assert checked["next_tool"] is None and checked["remaining_tool"]["name"] == "manifest.read"
    assert checked["stop_reason"] == handoff["stop_reason"] == auto["stop_reason"] == blocked["stop_reason"] == finalized["stop_reason"] == "no_progress_stop"
    assert handoff["recovery_stopped"] is True and handoff["ready_for_next_step"] is False
    assert auto["actions"] == [] and auto["attempts_total"] == 2
    assert blocked["status"] == "blocked" and blocked["attempt"] == 2
    assert finalized["status"] == "fail" and finalized["final_output"] is None
    assert history_path.read_bytes() == persisted_history
    assert manifest_path.read_bytes() == persisted_manifest
    assert not list((run.parents[5] / "final").rglob("promotion.json"))


def test_one_no_progress_attempt_still_allows_manifest_repair(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    manifest_path = run / "workbook_manifest.json"
    manifest_path.write_text("{", encoding="utf-8")
    assert execute_tool("step1.check", run)["next_tool"]["name"] == "manifest.read"

    with patch("excel_to_act.steps.step1.workflow._tool_action", side_effect=OSError("read-only manifest")):
        failed = execute_tool("manifest.read", run)
    checked = execute_tool("step1.check", run)
    repaired = execute_tool("manifest.read", run)
    history = _read(run / "attempt_history.json")

    assert failed["status"] == "error" and failed["attempt"] == 1
    assert checked["recovery_stopped"] is False
    assert checked["retryable"] is True and checked["next_tool"]["name"] == "manifest.read"
    assert repaired["status"] == "pass" and repaired["attempt"] == 2
    assert history["attempts"][0]["no_progress_count"] == 1
    assert history["attempts"][1]["improved"] is True
    assert history["attempts"][1]["no_progress_count"] == 0


@pytest.mark.parametrize(
    ("history_bytes", "reason"),
    [
        (b"{", "attempt_history_unreadable"),
        (b"[]", "attempt_history_invalid"),
        (b"{}", "attempt_history_invalid"),
        (b'{"attempt_limit":3}', "attempt_history_invalid"),
        (b'{"attempts":[]}', "attempt_history_invalid"),
        (b'{"attempts":{},"attempt_limit":3}', "attempt_history_invalid"),
        (b'{"attempts":[],"attempt_limit":true}', "attempt_history_invalid"),
        (b'{"attempts":[],"attempt_limit":null}', "attempt_history_invalid"),
        (b'{"attempts":[[]],"attempt_limit":3}', "attempt_history_invalid"),
        (b'{"attempts":[{"no_progress_count":false}],"attempt_limit":3}', "attempt_history_invalid"),
        (b'{"attempts":[],"attempt_limit":3,"no_progress_stopped":"false"}', "attempt_history_invalid"),
    ],
)
def test_invalid_attempt_history_stops_all_feedback_without_reset_or_promotion(
    tmp_path: Path, history_bytes: bytes, reason: str
) -> None:
    _, run, _ = _convert_single(tmp_path)
    inventory_path = run / "inventory.json"
    inventory = _read(inventory_path)
    next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "B2")["value"] = 88
    _write(inventory_path, inventory)
    candidate_bytes = inventory_path.read_bytes()

    history_path = run / "attempt_history.json"
    history_path.write_bytes(history_bytes)

    checked = execute_tool("step1.check", run)
    assert checked["status"] == "fail"
    assert checked["retryable"] is False and checked["next_tool"] is None
    assert checked["recovery_stopped"] is True and checked["stop_reason"] == reason
    assert checked["remaining_tool"]["name"] == "inventory.extract"

    blocked = execute_tool("inventory.extract", run)
    assert blocked["status"] == "blocked" and blocked["retryable"] is False
    assert blocked["stop_reason"] == reason and blocked["next_tool"] is None
    assert blocked["attempt"] is None and blocked["attempt_limit"] is None
    assert any(item["code"] == reason for item in blocked["diagnostics"])

    recovered = auto_recover(run, max_attempts=1)
    assert recovered["status"] == "fail"
    assert recovered["actions"] == [] and recovered["attempts_used"] == 0
    assert recovered["recovery_stopped"] is True and recovered["stop_reason"] == reason
    assert recovered["attempt_limit"] is None and recovered["attempts_total"] is None

    finalized = finalize_run(run)
    assert finalized["status"] == "fail" and finalized["ready_for_next_step"] is False
    assert finalized["final_output"] is None and finalized["stop_reason"] == reason
    assert not list((run.parents[5] / "final").rglob("promotion.json"))

    handoff = _read(run / "handoff.json")
    quality = _read(run / "quality.json")
    for result in (handoff, quality):
        assert result["recovery_stopped"] is True
        assert result["stop_reason"] == reason
        assert result["remaining_tool"]["name"] == "inventory.extract"
        assert result["ready_for_next_step"] is False
        assert any(item["code"] == reason for item in result["diagnostics"])
    assert history_path.read_bytes() == history_bytes
    assert inventory_path.read_bytes() == candidate_bytes


def test_invalid_attempt_history_cli_emits_json_without_traceback(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    inventory_path = run / "inventory.json"
    inventory = _read(inventory_path)
    next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "B2")["value"] = 88
    _write(inventory_path, inventory)
    (run / "attempt_history.json").write_text("[]", encoding="utf-8")

    result = CliRunner().invoke(app, ["step1", "tool", "inventory.extract", "--run", str(run)])
    assert result.exit_code == 1
    envelope = json.loads(result.output)
    assert envelope["status"] == "blocked"
    assert envelope["stop_reason"] == "attempt_history_invalid"
    assert "Traceback" not in result.output


def test_attempt_history_keeps_legacy_optional_fields_compatible(tmp_path: Path) -> None:
    _, run, _ = _convert_single(tmp_path)
    inventory_path = run / "inventory.json"
    inventory = _read(inventory_path)
    next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "B2")["value"] = 88
    _write(inventory_path, inventory)
    history = {"attempt_limit": 3, "attempts": [{"tool": "cached_values.read"}]}
    _write(run / "attempt_history.json", history)

    result = execute_tool("step1.check", run)
    assert result["status"] == "fail"
    assert result["retryable"] is True
    assert result["next_tool"]["name"] == "inventory.extract"
    assert result["stop_reason"] is None
