from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from typer.testing import CliRunner

from excel_to_act.interfaces.cli import app
from excel_to_act.ingest.openpyxl_reader import OpenpyxlWorkbookReader
from excel_to_act.inventory.extractor import OpenpyxlInventoryExtractor
from excel_to_act.steps import conversion_workflow
from excel_to_act.steps.step3 import workflow
from excel_to_act.steps.step3 import exploration
from excel_to_act.steps.step3 import semantic
from excel_to_act.steps.step3 import design
from excel_to_act.steps.step3 import profiling


def _location(sheet: str | None, address: str | None, object_type: str, object_id: str | None = None) -> dict:
    return {"workbook_path": "fixture.xlsx", "sheet_name": sheet, "sheet_index": 0 if sheet else None,
            "address": address, "object_type": object_type, "object_id": object_id}


def _cell(sheet: str, address: str, row: int, column: int, kind: str, value=None, formula=None, **extra) -> dict:
    return {"source_location": _location(sheet, address, "cell", f"{sheet}!{address}"),
            "source_identity": f"cell:{sheet}!{address}", "address": address, "row": row, "column": column,
            "kind": kind, "value": value, "formula": formula, "data_type": "f" if formula else "n", **extra}


def _inventory() -> dict:
    main_cells = [
        _cell("Main", "C3", 3, 3, "value", 10),
        _cell("Main", "C4", 4, 3, "value", 20),
        _cell("Main", "F1", 1, 6, "formula", formula="=C3+1", cached_value=11, cached_value_available=True),
        _cell("Main", "F2", 2, 6, "formula", formula="=C4+1", cached_value=21, cached_value_available=True),
        _cell("Main", "C2", 2, 3, "formula", formula="=SUM(B2:B5)"),
        _cell("Main", "D2", 2, 4, "formula", formula="=SUM(C3:C4)"),
        _cell("Main", "D3", 3, 4, "formula", formula='=INDIRECT("Main!C3")'),
        _cell("Main", "D4", 4, 4, "formula", formula="=#REF!"),
        _cell("Main", "D5", 5, 4, "formula", formula="=Main!Z99"),
        _cell("Main", "E1", 1, 5, "formula", formula="=Ignored!A1"),
        _cell("Main", "G1", 1, 7, "formula", formula="=GP"),
        _cell("Main", "G2", 2, 7, "formula", formula="=F1+$C$3"),
        _cell("Main", "G3", 3, 7, "formula", formula="=F2+$C$3"),
        _cell("Main", "H1", 1, 8, "value", "Age", data_type="s"),
        _cell("Main", "I1", 1, 9, "value", "Rate", data_type="s"),
        _cell("Main", "H2", 2, 8, "value", 30),
        _cell("Main", "I2", 2, 9, "value", 0.5),
        _cell("Main", "H3", 3, 8, "value", 40),
        _cell("Main", "I3", 3, 9, "value", 0.6),
        _cell("Main", "H4", 4, 8, "value", "Total", data_type="s"),
        _cell("Main", "I4", 4, 9, "value", 1.1),
        _cell("Main", "J1", 1, 10, "formula", formula="=SUM(Rates[Rate])"),
        _cell("Main", "J2", 2, 10, "formula", formula="=SUM(Rates[#All])"),
        _cell("Main", "J3", 3, 10, "formula", formula="=SUM(Rates[#Data])"),
        _cell("Main", "K2", 2, 11, "formula", formula="=SUM(Rates[@Rate])"),
    ]
    ignored_cells = [_cell("Ignored", "A1", 1, 1, "formula", formula="=1+1", cached_value=2, cached_value_available=True)]
    names = []
    for name, address in (("GP", "C3"), ("NLP", "C4")):
        names.append({"source_location": _location("Main", address, "defined_name", name), "source_identity": None,
                      "name": name, "address": address, "kind": "defined_name", "metadata": {"scope": "workbook"}})
    names.append({"source_location": _location("Main", None, "defined_name", "_xlnm._FilterDatabase"), "source_identity": None,
                  "name": "_xlnm._FilterDatabase", "address": "", "kind": "defined_name", "metadata": {"scope": "workbook"}})
    table = {"source_location": _location("Main", "H1:I4", "table", "Rates"), "name": "Rates", "address": "H1:I4",
             "kind": "table", "metadata": {"columns": ["Age", "Rate"], "header_row_count": 1, "totals_row_count": 1}}
    return {"schema_version": "phase1.v1", "artifact_type": "workbook_inventory", "workbook_sha256": "source-sha",
            "sheets": [
                {"source_location": _location("Main", None, "sheet"), "name": "Main", "index": 0, "max_row": 5, "max_column": 9,
                 "state": "visible", "cells": main_cells, "ranges": [table], "layout_objects": []},
                {"source_location": _location("Ignored", None, "sheet"), "name": "Ignored", "index": 1, "max_row": 1, "max_column": 1,
                 "state": "visible", "cells": ignored_cells, "ranges": [], "layout_objects": []},
            ], "workbook_ranges": names, "vba_modules": [], "unsupported_features": [],
            "coverage": {"recognized_inventory_objects": len(main_cells) + len(ignored_cells) + len(names) + 1,
                         "unsupported_or_opaque_objects": 0, "discovered_workbook_objects": len(main_cells) + len(ignored_cells) + len(names) + 1}}


def _write_json(path: Path, value: dict) -> str:
    data = json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8") + b"\n"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return hashlib.sha256(data).hexdigest()


def _prepared_fixture(tmp_path: Path, monkeypatch, *, scope: dict | None = None, entries: int = 1, select_source: bool = True,
                      inventory: dict | None = None, workflow_path: Path | None = None) -> tuple[Path, Path, dict]:
    root = tmp_path / "step1"
    root.mkdir(parents=True)
    inventory_sha = _write_json(root / "inventory.json", inventory if inventory is not None else _inventory())
    manifest = {"schema_version": "phase1.v1", "artifact_type": "workbook_manifest", "sha256": "source-sha",
                "sheets": [{"name": "Main"}, {"name": "Ignored"}]}
    manifest_sha = _write_json(root / "workbook_manifest.json", manifest)
    entry_list = []
    for ordinal in range(entries):
        entry_list.append({"source_id": f"source-{ordinal}", "source_path": f"book-{ordinal}.xlsx", "source_sha256": "source-sha",
                           "run_id": f"run-{ordinal}", "run_path": "run", "status": "pass", "ready_for_next_step": True,
                           "artifacts": [{"name": "inventory.json", "path": "inventory.json", "sha256": inventory_sha},
                                         {"name": "workbook_manifest.json", "path": "workbook_manifest.json", "sha256": manifest_sha}]})
    index = {"schema_version": "step2.index.v1", "step1_root": str(root), "input_handoff_path": "unused.json",
             "input_handoff_sha256": "0" * 64, "status": "pass", "validation_status": "pass", "entries": entry_list}
    index_path = tmp_path / "index.json"
    _write_json(index_path, index)
    monkeypatch.setattr(workflow, "_validate_saved_index", lambda *_args: {"status": "pass", "diagnostics": []})
    scope_path = None
    if scope is not None:
        scope_path = tmp_path / "scope.json"
        _write_json(scope_path, scope)
    out = tmp_path / "analysis"
    result = workflow.prepare_analysis(index_path, root, out, source_id="source-0" if entries > 1 and select_source else None,
                                       scope_path=scope_path, workflow_path=workflow_path)
    return out, root, result


def test_fields_keep_named_scalars_and_group_only_copied_formulas() -> None:
    inventory = _inventory()
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    by_address = {(item["sheet"], member["address"]): item for item in fields["fields"] for member in item["members"]}

    assert by_address[("Main", "C3")]["member_count"] == 1
    assert by_address[("Main", "C3")]["names"] == ["GP"]
    assert by_address[("Main", "C4")]["names"] == ["NLP"]
    assert by_address[("Main", "H1")]["member_count"] == 1
    assert by_address[("Main", "I1")]["member_count"] == 1
    assert by_address[("Main", "H1")]["physical_shape"] == "0D"
    assert by_address[("Main", "I1")]["physical_shape"] == "0D"
    assert by_address[("Main", "H1")]["field_id"] != by_address[("Main", "I1")]["field_id"]
    name_descriptors = {item["name"]: item for item in fields["descriptors"]["defined_names"]}
    assert name_descriptors["GP"]["physical_dimensions"] == 0
    assert name_descriptors["_xlnm._FilterDatabase"]["physical_dimensions"] is None
    assert by_address[("Main", "F1")] is by_address[("Main", "F2")]
    assert by_address[("Main", "G2")] is by_address[("Main", "G3")]
    assert by_address[("Main", "F1")]["role"] == "calculated"
    assert by_address[("Main", "F1")]["cached_formula_cells"] == 2
    assert fields["summary"]["excluded_stored_cells"] == 1
    assert next(item for item in fields["descriptors"]["tables"] if item["name"] == "Rates")["physical_dimensions"] == 2


def test_array_formula_followers_are_calculated_members_without_becoming_raw_inputs() -> None:
    inventory = _inventory()
    main = next(item for item in inventory["sheets"] if item["name"] == "Main")
    main["cells"].extend([
        _cell("Main", "L1", 1, 12, "formula", formula="=TRANSPOSE(P1:P2)",
              raw_formula_attributes={"t": "array", "ref": "L1:M2"},
              cached_value_present=False, cached_value_available=False),
        _cell("Main", "M1", 1, 13, "value", 0.1, formula=None,
              formula_present=False, cached_value_present=False, cached_value_available=False),
        _cell("Main", "L2", 2, 12, "value", 0.2, formula=None,
              formula_present=False, cached_value_present=False, cached_value_available=False),
        _cell("Main", "M2", 2, 13, "value", 0.3, formula=None,
              formula_present=False, cached_value_present=False, cached_value_available=False),
    ])

    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    by_address = {(item["sheet"], member["address"]): (item, member)
                  for item in fields["fields"] for member in item["members"]}
    array_field = by_address[("Main", "L1")][0]

    assert array_field["role"] == "calculated"
    assert array_field["physical_shape"] == "2D"
    assert array_field["address_extent"] == "L1:M2"
    assert array_field["member_count"] == 4
    assert array_field["formula_member_count"] == 1
    assert array_field["array_formula_follower_count"] == 3
    for address in ("L1", "M1", "L2", "M2"):
        assert by_address[("Main", address)][0] is array_field
    follower = by_address[("Main", "M1")][1]
    assert follower["source_kind"] == "value"
    assert follower["formula_present"] is False
    assert follower["formula"] is None
    assert follower["cached_value_available"] is False
    assert follower["array_formula_member"]["role"] == "follower"
    assert fields["summary"]["ordinary_formula_cells"] == fields["summary"]["formula_cells"] - 3
    assert fields["summary"]["array_formula_followers"] == 3
    assert fields["descriptors"]["array_formulas"] == [{
        "sheet": "Main", "anchor": "L1", "address": "L1:M2", "formula": "=TRANSPOSE(P1:P2)",
        "physical_dimensions": 2, "physical_shape": [2, 2], "stored_member_count": 4,
        "stored_follower_count": 3,
    }]


def test_dependencies_keep_ranges_names_empty_dynamic_and_excluded_boundaries() -> None:
    inventory = _inventory()
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    field_map = {item["field_id"]: item for item in fields["fields"]}
    by_addr = {(item["sheet"], member["address"]): item["field_id"] for item in fields["fields"] for member in item["members"]}
    consumer = by_addr[("Main", "D2")]

    prerequisites = {item["prerequisite"] for item in dependencies["edges"] if item["consumer"] == consumer}
    assert by_addr[("Main", "C3")] in prerequisites
    assert by_addr[("Main", "C4")] in prerequisites
    assert field_map[by_addr[("Main", "C3")]]["names"] == ["GP"]
    assert any(item["boundary_type"] == "excluded_sheet" and item["target_details"]["cached_formula_count"] == 1 for item in dependencies["boundaries"])
    assert any(item["category"] == "dynamic_reference" for item in dependencies["unresolved"])
    assert any(item["boundary_type"] == "empty_cell" for item in dependencies["boundaries"])
    assert dependencies["direction"] == "prerequisite_before_consumer"


def test_partial_empty_range_keeps_boundary_evidence_and_prerequisite_edges() -> None:
    inventory = _inventory()
    inventory["sheets"][0]["cells"] = [cell for cell in inventory["sheets"][0]["cells"] if cell["address"] != "C4"]
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    by_addr = {(item["sheet"], member["address"]): item["field_id"] for item in fields["fields"] for member in item["members"]}
    consumer = by_addr[("Main", "D2")]

    prerequisites = {item["prerequisite"] for item in dependencies["edges"] if item["consumer"] == consumer}
    assert by_addr[("Main", "C3")] in prerequisites
    partial = [item for item in dependencies["boundaries"]
               if item["consumer"] == consumer and item["boundary_type"] == "partial_empty_range"]
    assert len(partial) == 1
    assert partial[0]["status"] == "boundary"
    assert partial[0]["target_details"]["range_area"] == 2
    assert partial[0]["target_details"]["stored_cell_count"] == 1
    assert partial[0]["target_details"]["empty_cells"] == 1
    assert "C4" in partial[0]["target_details"]["empty_examples"]
    assert not any(item["category"] == "partial_empty_range" for item in dependencies["unresolved"])
    assert dependencies["summary"]["empty_reference_occurrences"] == sum(
        item["occurrences"] for item in dependencies["boundaries"]
        if item["status"] == "empty" or item["boundary_type"] == "partial_empty_range"
    )
    plan = workflow._build_plan(fields, dependencies, "binding", "fields", "dependencies")
    assert plan["readiness"]["generation_ready"] is False


def test_named_scalar_self_reference_is_cycle_not_recurrence() -> None:
    inventory = _inventory()
    d2 = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "D2")
    d2["formula"] = "=SelfRef+1"
    inventory["workbook_ranges"].append({
        "source_location": _location("Main", "D2", "defined_name", "SelfRef"),
        "source_identity": None,
        "name": "SelfRef",
        "address": "D2",
        "kind": "defined_name",
        "metadata": {"scope": "workbook"},
    })
    inventory["sheets"][0]["cells"].extend([
        _cell("Main", "L1", 1, 12, "value", 0),
        _cell("Main", "L2", 2, 12, "formula", formula="=L1+1"),
        _cell("Main", "L3", 3, 12, "formula", formula="=L2+1"),
    ])
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    by_address = {(field["sheet"], member["address"]): field["field_id"]
                  for field in fields["fields"] for member in field["members"]}
    scalar_consumer = by_address[("Main", "D2")]
    scalar_internal = next(item for item in dependencies["internal_dependencies"]
                           if item["consumer"] == scalar_consumer)

    assert scalar_internal["same_cell"] is True
    assert scalar_internal["recurrence_candidate"] is False
    plan = workflow._build_plan(fields, dependencies, "binding", "fields", "dependencies")
    scalar_block = next(item for item in plan["blocks"] if scalar_consumer in item["fields"])
    assert scalar_block["cycle_candidate"] is True
    assert scalar_consumer not in plan["recurrence_candidates"]

    lag_field = by_address[("Main", "L2")]
    lag_internal = next(item for item in dependencies["internal_dependencies"]
                        if item["consumer"] == lag_field)
    assert lag_internal["same_cell"] is False
    assert lag_internal["recurrence_candidate"] is True
    assert lag_field in plan["recurrence_candidates"]


def test_named_scalar_self_reference_in_multicell_range_is_cycle() -> None:
    inventory = _inventory()
    c3 = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "C3")
    c3["kind"] = "formula"
    c3["formula"] = "=SUM(C3:C4)"
    c3["value"] = None
    c3["data_type"] = "f"
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    gp_field = next(item["field_id"] for item in fields["fields"] if item["names"] == ["GP"])
    nlp_field = next(item["field_id"] for item in fields["fields"] if item["names"] == ["NLP"])
    self_dependency = next(item for item in dependencies["internal_dependencies"]
                           if item["consumer"] == gp_field and item["prerequisite"] == gp_field)

    assert self_dependency["same_cell"] is True
    assert self_dependency["recurrence_candidate"] is False
    assert any(item["consumer"] == gp_field and item["prerequisite"] == nlp_field
               for item in dependencies["edges"])
    plan = workflow._build_plan(fields, dependencies, "binding", "fields", "dependencies")
    gp_block = next(item for item in plan["blocks"] if gp_field in item["fields"])
    assert gp_block["cycle_candidate"] is True
    assert gp_field not in plan["recurrence_candidates"]


def test_structured_table_references_keep_header_data_totals_and_current_row_boundaries() -> None:
    inventory = _inventory()
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    by_addr = {(item["sheet"], member["address"]): item["field_id"] for item in fields["fields"] for member in item["members"]}
    header = by_addr[("Main", "H1")]
    age_data = by_addr[("Main", "H2")]
    rate_data = by_addr[("Main", "I2")]
    rate_last_data = by_addr[("Main", "I3")]
    rate_total = by_addr[("Main", "I4")]
    assert rate_total != rate_data
    formulas = {address: by_addr[("Main", address)] for address in ("J1", "J2", "J3", "K2")}

    def prerequisites(address: str) -> set[str]:
        return {item["prerequisite"] for item in dependencies["edges"] if item["consumer"] == formulas[address]}

    assert header not in prerequisites("J1")  # Rates[Rate] means data cells by default.
    assert prerequisites("J1") == {rate_data, rate_last_data}
    assert header in prerequisites("J2")  # #All explicitly includes headers.
    assert age_data in prerequisites("J2") and rate_data in prerequisites("J2")
    assert rate_total in prerequisites("J2")
    assert header not in prerequisites("J3")
    assert age_data in prerequisites("J3") and rate_data in prerequisites("J3")
    assert rate_total not in prerequisites("J3")
    assert prerequisites("K2") == {rate_data}  # [@Rate] is detected from graph metadata.


def test_bare_current_row_reference_is_explicitly_unresolved() -> None:
    inventory = _inventory()
    rate_formula = next(cell for cell in inventory["sheets"][0]["cells"] if cell["address"] == "H2")
    rate_formula["kind"] = "formula"
    rate_formula["formula"] = "=[@Rate]*2"
    rate_formula["value"] = None
    rate_formula["data_type"] = "f"
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    consumer = next(field["field_id"] for field in fields["fields"]
                    if field["sheet"] == "Main" and any(member["address"] == "H2" for member in field["members"]))

    assert not any(item["consumer"] == consumer for item in dependencies["edges"])
    unresolved = next(item for item in dependencies["unresolved"] if item["consumer"] == consumer)
    assert unresolved["category"] == "unresolved_reference"
    assert any("[@Rate]" in sample for sample in unresolved["evidence_samples"])


def test_compound_structured_reference_is_an_explicit_unresolved_boundary() -> None:
    inventory = _inventory()
    inventory["sheets"][0]["cells"].append(
        _cell("Main", "J4", 4, 10, "formula", formula="=SUM(Rates[[#All],[Rate]])")
    )
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    consumer = next(
        field["field_id"] for field in fields["fields"]
        if field["sheet"] == "Main" and any(member["address"] == "J4" for member in field["members"])
    )

    assert not any(item["consumer"] == consumer for item in dependencies["edges"])
    boundary = next(item for item in dependencies["boundaries"]
                    if item["consumer"] == consumer and item["boundary_type"] == "structured_reference_unsupported")
    assert boundary["status"] == "unresolved"
    assert boundary["evidence"] == "Rates[[#All],[Rate]]"
    assert boundary["target_details"]["selector"] == "[#All],[Rate]"
    unresolved = next(item for item in dependencies["unresolved"] if item["consumer"] == consumer)
    assert unresolved["category"] == "unresolved_reference"
    assert "Rates[[#All],[Rate]]" in unresolved["evidence_samples"]


def test_plain_structured_column_without_verified_header_mapping_is_unresolved() -> None:
    inventory = _inventory()
    inventory["sheets"][0]["ranges"][0]["metadata"].pop("columns")
    inventory["sheets"][0]["ranges"][0]["metadata"]["header_row_count"] = 0
    inventory["sheets"][0]["cells"] = [
        cell for cell in inventory["sheets"][0]["cells"] if cell["address"] not in {"H1", "I1"}
    ]
    inventory["sheets"][0]["cells"].append(
        _cell("Main", "J4", 4, 10, "formula", formula="=SUM(Rates[Rate])")
    )
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    consumer = next(
        field["field_id"] for field in fields["fields"]
        if field["sheet"] == "Main" and any(member["address"] == "J4" for member in field["members"])
    )

    assert not any(item["consumer"] == consumer for item in dependencies["edges"])
    assert any(item["consumer"] == consumer and item["boundary_type"] == "structured_reference_unsupported"
               for item in dependencies["boundaries"])


def test_legacy_table_without_native_columns_does_not_infer_from_blank_header_row() -> None:
    inventory = _inventory()
    inventory["sheets"][0]["ranges"][0]["metadata"].pop("columns")
    # A blank first header shifts the remaining visible header away from its true table column.
    inventory["sheets"][0]["cells"] = [cell for cell in inventory["sheets"][0]["cells"] if cell["address"] != "H1"]
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    by_address = {(field["sheet"], member["address"]): field["field_id"]
                  for field in fields["fields"] for member in field["members"]}
    consumer = by_address[("Main", "J1")]

    assert not any(item["consumer"] == consumer for item in dependencies["edges"])
    boundary = next(item for item in dependencies["boundaries"]
                    if item["consumer"] == consumer and item["boundary_type"] == "structured_reference_unsupported")
    assert boundary["evidence"] == "Rates[Rate]"
    assert any(item["consumer"] == consumer and item["category"] == "unresolved_reference"
               for item in dependencies["unresolved"])


def test_legacy_table_without_row_counts_exposes_boundary_uncertainty() -> None:
    inventory = _inventory()
    inventory["sheets"][0]["ranges"][0]["metadata"].pop("header_row_count")
    inventory["sheets"][0]["ranges"][0]["metadata"].pop("totals_row_count")
    fields = workflow._build_fields(inventory, ["Ignored"], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, ["Ignored"], "binding", "fields-sha")
    by_addr = {(item["sheet"], member["address"]): item["field_id"] for item in fields["fields"] for member in item["members"]}
    consumers = {address: by_addr[("Main", address)] for address in ("J1", "J2", "J3", "K2")}

    # #All names its full physical range, while data-only selectors cannot safely infer the body.
    all_prereqs = {item["prerequisite"] for item in dependencies["edges"] if item["consumer"] == consumers["J2"]}
    assert by_addr[("Main", "H1")] in all_prereqs
    assert by_addr[("Main", "I4")] in all_prereqs
    unknown = [item for item in dependencies["boundaries"] if item["boundary_type"] == "structured_reference_boundary_unknown"]
    assert len(unknown) == 3
    assert any(item["category"] == "table_boundary_unknown" for item in dependencies["unresolved"])


def test_inventory_preserves_table_header_and_totals_row_counts(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Age", "Rate"])
    sheet.append([10, 0.5])
    sheet.append([20, 0.6])
    sheet.append(["Total", 1.1])
    table = Table(displayName="Rates", ref="A1:B4", headerRowCount=1, totalsRowCount=1)
    sheet.add_table(table)
    no_header = workbook.create_sheet("NoHeaders")
    no_header.append([10, 0.5])
    no_header.append([20, 0.6])
    no_header_table = Table(displayName="NoHeaderRates", ref="A1:B2", headerRowCount=0, totalsRowCount=0)
    no_header.add_table(no_header_table)
    source = tmp_path / "table-boundaries.xlsx"
    workbook.save(source)
    workbook.close()

    manifest = OpenpyxlWorkbookReader().read_manifest(source)
    inventory = OpenpyxlInventoryExtractor().extract(source, manifest)
    descriptor = next(item for item in inventory.sheets[0].ranges if item.kind == "table")
    assert descriptor.metadata["columns"] == ["Age", "Rate"]
    assert descriptor.metadata["header_row_count"] == 1
    assert descriptor.metadata["totals_row_count"] == 1
    no_header_descriptor = next(item for item in inventory.sheets[1].ranges if item.kind == "table")
    assert no_header_descriptor.metadata["columns"] == ["Column1", "Column2"]
    assert no_header_descriptor.metadata["header_row_count"] == 0
    assert no_header_descriptor.metadata["totals_row_count"] == 0


def test_headerless_table_uses_native_columns_not_first_data_row_text(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inputs"
    sheet["A1"] = "Rate"
    sheet["B1"] = "North"
    sheet["A2"] = 10
    sheet["B2"] = 20
    sheet["C1"] = "=SUM(NoHeaderRates[Column1])"
    sheet["C2"] = "=SUM(NoHeaderRates[Rate])"
    sheet.add_table(Table(displayName="NoHeaderRates", ref="A1:B2", headerRowCount=0, totalsRowCount=0))
    source = tmp_path / "headerless-table.xlsx"
    workbook.save(source)
    workbook.close()

    manifest = OpenpyxlWorkbookReader().read_manifest(source)
    inventory = OpenpyxlInventoryExtractor().extract(source, manifest).model_dump(mode="json")
    table = next(item for item in inventory["sheets"][0]["ranges"] if item["kind"] == "table")
    assert table["metadata"]["columns"] == ["Column1", "Column2"]
    fields = workflow._build_fields(inventory, [], "binding")
    dependencies = workflow._build_dependencies(inventory, fields, [], "binding", "fields-sha")
    by_address = {(field["sheet"], member["address"]): field["field_id"]
                  for field in fields["fields"] for member in field["members"]}
    valid_consumer = by_address[("Inputs", "C1")]
    invalid_consumer = by_address[("Inputs", "C2")]
    valid_prerequisites = {item["prerequisite"] for item in dependencies["edges"] if item["consumer"] == valid_consumer}

    assert valid_prerequisites == {by_address[("Inputs", "A1")], by_address[("Inputs", "A2")]}
    assert not any(item["consumer"] == invalid_consumer for item in dependencies["edges"])
    assert any(item["consumer"] == invalid_consumer and item["category"] == "unresolved_reference"
               for item in dependencies["unresolved"])


def test_scc_orders_components_and_keeps_cycles_as_blocks() -> None:
    components = workflow._scc(["a", "b", "c", "d"], {"a": {"b", "c"}, "b": set(), "c": {"b"}, "d": {"d"}})
    assert components == [["d"], ["a"], ["c"], ["b"]] or sorted(components) == [["a"], ["b"], ["c"], ["d"]]
    fields = {field_id: {"field_id": field_id, "role": "calculated", "physical_shape": "0D"} for field_id in "abcd"}
    dependencies = {"edges": [{"consumer": "a", "prerequisite": "b"}, {"consumer": "b", "prerequisite": "a"},
                               {"consumer": "c", "prerequisite": "a"}], "internal_dependencies": [], "unresolved": []}
    plan = workflow._build_plan({"fields": list(fields.values())}, dependencies, "binding", "fields", "deps")
    by_field = {field_id: block for block in plan["blocks"] for field_id in block["fields"]}
    assert by_field["a"]["block_id"] == by_field["b"]["block_id"]
    assert by_field["a"]["cycle_candidate"] is True
    assert by_field["a"]["order"] < by_field["c"]["order"]
    assert plan["readiness"]["gpu_executable"] is False


def test_prepare_stages_validate_and_detect_tampering(tmp_path: Path, monkeypatch) -> None:
    scope = {"schema_version": "analysis.scope.v1", "workbook_sha256": "source-sha", "source_run_id": "older-run", "ignored_sheets": ["Ignored"]}
    out, _root, prepared = _prepared_fixture(tmp_path, monkeypatch, scope=scope)
    assert prepared["status"] == "pass"
    assert workflow.build_fields(out)["status"] == "pass"
    assert workflow.build_dependencies(out)["status"] == "pass"
    assert workflow.build_plan(out)["status"] == "pass"
    checked = workflow.validate_analysis(out)
    assert checked["status"] == "pass"
    assert checked["handoff"]["readiness"] == {"draft_only": True, "runtime_verified": False, "generation_ready": False, "gpu_executable": False}
    assert (out / "fields.md").is_file() and (out / "execution_plan.md").is_file()

    fields_path = out / "fields.json"
    fields_path.write_text("{}", encoding="utf-8")
    stale = workflow.build_dependencies(out)
    assert stale["status"] == "blocked"
    assert "checksum mismatch" in stale["diagnostics"][0]["message"]


def test_input_catalog_cli_uses_analysis_top_level_binding_hash(tmp_path: Path, monkeypatch) -> None:
    workflow_root = tmp_path / "reviewed-workflow"
    conversion_workflow.create_workflow(
        workflow_root, {"source_id": "source-0", "run_id": "run-0", "workbook_sha256": "source-sha"})
    stage1 = conversion_workflow.append_stage_artifact(
        workflow_root, 1, "stage1.json", {"status": "pass"}, "Step 1")
    for reviewer in ("agent", "human"):
        conversion_workflow.record_decision(workflow_root, 1, reviewer, "approve", "Fixture approval.")
    conversion_workflow.append_stage_artifact(
        workflow_root, 2, "stage2.json", {"status": "pass"}, "Step 2",
        input_stages={1: stage1["artifact"]})
    for reviewer in ("agent", "human"):
        conversion_workflow.record_decision(workflow_root, 2, reviewer, "approve", "Fixture approval.")

    analysis_dir, _step1_root, prepared = _prepared_fixture(
        tmp_path / "prepared", monkeypatch, workflow_path=workflow_root)
    assert prepared["status"] == "pass"
    assert workflow.build_fields(analysis_dir)["status"] == "pass"
    assert workflow.build_dependencies(analysis_dir)["status"] == "pass"
    assert workflow.build_plan(analysis_dir)["status"] == "pass"

    analysis_bytes = (analysis_dir / "analysis.json").read_bytes()
    analysis = json.loads(analysis_bytes)
    binding = analysis["binding"]
    source = {"source_id": binding["source"]["source_id"],
              "run_id": binding["source"]["run_id"],
              "workbook_sha256": binding["source"]["source_sha256"],
              "analysis_binding_sha256": analysis["binding_sha256"]}
    scenario_basis = {"scenario_id": "fixture-saved-inputs", "selectors": {"source": "fixture"},
                      "overrides": {}}
    scenario = {**scenario_basis,
                "scenario_sha256": conversion_workflow.hash_bytes(conversion_workflow.json_bytes(scenario_basis))}
    target = {"target_id": "fixture_output", "selector": "Main!G1", "result_order": 0,
              "shape": [], "start_cells": ["Main!G1"], "result_kind": "scalar",
              "units": "USD", "axes": []}
    target_selection = {"source": source, "scenario": scenario, "targets": [target]}
    fields = json.loads((analysis_dir / "fields.json").read_text(encoding="utf-8"))["fields"]
    variables = []
    for field in fields:
        members = field["members"]
        is_formula = field["role"] == "calculated"
        if is_formula:
            continue
        selected = members
        role = "source_raw"
        kind = "scalar" if len(selected) == 1 else "vector"
        variables.append({
            "variable_id": f"fixture.{field['field_id']}", "group_id": "fixture_inputs",
            "logical_name": f"Fixture field {field['field_id']}", "role": role, "kind": kind,
            "logical_role": "business_input" if kind == "scalar" else "series",
            "shape": [] if kind == "scalar" else [len(selected)],
            "axes": [] if kind == "scalar" else [{"name": "stored_member_order", "role": "observation"}],
            "source_extents": [{"sheet": field["sheet"], "range": member["address"],
                                "source_role": "raw_source"}
                               for member in selected],
            "consumer_targets": ["fixture_output"],
            "provenance": {"description": "Exact fixture field membership from current Step 3 fields."},
            "value_source_policy": {"origin": "workbook_formula_mode_source", "formula_cache_allowed": False},
        })
    variables[0]["source_column_labels"] = [{"source_cell": "Main!H1", "label": "Age"}]
    catalog = {
        "schema_version": "step3.input_boundary.input.v1", "source": source, "scenario": scenario,
        "target_selection": target_selection,
        "groups": [{"group_id": "fixture_inputs", "name": "Fixture inputs", "classification": "business_input"}],
        "variables": variables, "axis_metadata": [], "source_records": [], "upstream_exclusions": [],
        "value_source_policy": {"formula_cache_allowed": False}, "open_questions": [],
    }
    targets_path, catalog_path = tmp_path / "targets.json", tmp_path / "catalog.json"
    targets_input = {"schema_version": "step3.input_targets.v1", **target_selection}
    _write_json(targets_path, targets_input)
    _write_json(catalog_path, catalog)

    result = CliRunner().invoke(app, ["step3", "input-catalog", "--analysis", str(analysis_dir),
                                     "--workflow", str(workflow_root), "--targets", str(targets_path),
                                     "--catalog", str(catalog_path)])

    assert result.exit_code == 0, result.stdout
    response = json.loads(result.stdout)
    assert response["status"] == "pass"
    pruned = response["pruned_source_analysis"]
    assert not {"_reachable_raw_keys", "_declared_boundary_keys", "_formula_by_address"} & pruned.keys()
    assert pruned["unknown_count"] is None
    def strings(value):
        if isinstance(value, dict):
            for child in value.values():
                yield from strings(child)
        elif isinstance(value, list):
            for child in value:
                yield from strings(child)
        elif isinstance(value, str):
            yield value

    assert not any(value.startswith("{'") for value in strings(pruned))
    assert "_reachable_raw_keys" not in result.stdout
    assert (analysis_dir / "analysis.json").read_bytes() == analysis_bytes
    _, manifest = conversion_workflow.load_workflow(workflow_root)
    boundary = conversion_workflow.current_revision(manifest, 3)
    assert boundary is not None
    assert boundary["artifact"]["json_sha256"] == response["artifact"]["json_sha256"]
    assert manifest["input_boundary"]["source_sha256"] == source["workbook_sha256"]
    saved_boundary = json.loads((workflow_root / boundary["artifact"]["json"]).read_text(encoding="utf-8"))
    assert saved_boundary["target_selection"]["targets"] == [target]
    saved_report = (workflow_root / boundary["artifact"]["md"]).read_text(encoding="utf-8")
    saved_front = saved_report.split("## Detailed source and boundary evidence", 1)[0]
    assert "kind scalar; shape []; units USD; declared axes []" in saved_front
    saved_labels = [label for variable in saved_boundary["variables"]
                    for label in variable.get("source_column_labels", [])]
    assert {"source_cell": "Main!H1", "label": "Age"} in saved_labels

    revision_count = len(manifest["stages"]["3"]["revisions"])

    def submit_invalid(target_doc: dict, catalog_doc: dict, suffix: str, expected_reason: str) -> None:
        invalid_targets_path = tmp_path / f"targets-with-{suffix}.json"
        invalid_catalog_path = tmp_path / f"catalog-with-{suffix}.json"
        _write_json(invalid_targets_path, target_doc)
        _write_json(invalid_catalog_path, catalog_doc)
        rejected = CliRunner().invoke(app, ["step3", "input-catalog", "--analysis", str(analysis_dir),
                                            "--workflow", str(workflow_root), "--targets", str(invalid_targets_path),
                                            "--catalog", str(invalid_catalog_path)])
        assert rejected.exit_code == 1
        blocked = json.loads(rejected.stdout)
        assert blocked["status"] == "blocked"
        assert expected_reason in blocked["diagnostics"][0]["message"]
        _, after_manifest = conversion_workflow.load_workflow(workflow_root)
        assert len(after_manifest["stages"]["3"]["revisions"]) == revision_count

    for injected in ("sample_rates", "rates"):
        invalid_catalog = json.loads(json.dumps(catalog))
        if injected == "sample_rates":
            invalid_catalog["variables"][0]["provenance"][injected] = [0.123, 0.456]
        else:
            invalid_catalog["variables"][0][injected] = [0.123, 0.456]
        submit_invalid(targets_input, invalid_catalog, injected, "unsupported fields")

    mismatched_target_catalog = json.loads(json.dumps(catalog))
    mismatched_target_catalog["target_selection"]["targets"][0].pop("units")
    submit_invalid(targets_input, mismatched_target_catalog, "target-metadata", "target order and selectors must exactly match")

    invalid_policy = json.loads(json.dumps(catalog))
    invalid_policy["value_source_policy"]["formula_derived_external_values_available_now"] = {
        "CI": [0.03, 0.04]}
    submit_invalid(targets_input, invalid_policy, "policy-values", "must remain unavailable")

    invalid_question = json.loads(json.dumps(catalog))
    invalid_question["open_questions"] = [{"question_id": "Q1", "question": {"rates": [0.03, 0.04]}}]
    submit_invalid(targets_input, invalid_question, "question-values", "must be non-empty text")

    invalid_policy_list = json.loads(json.dumps(catalog))
    invalid_policy_list["value_source_policy"]["allowed_future_external_sources"] = {"CI": [0.03, 0.04]}
    submit_invalid(targets_input, invalid_policy_list, "policy-list-values", "non-empty list of text metadata")

    invalid_binding_policy = json.loads(json.dumps(catalog))
    invalid_binding_policy["value_source_policy"]["external_values_must_bind"] = {"rates": [0.03, 0.04]}
    submit_invalid(targets_input, invalid_binding_policy, "policy-binding-values", "non-empty list of text metadata")

    invalid_policy_choice = json.loads(json.dumps(catalog))
    invalid_policy_choice["value_source_policy"]["allowed_future_external_sources"] = ["CI rates 0.03, 0.04"]
    submit_invalid(targets_input, invalid_policy_choice, "policy-list-choice", "declared input-boundary policy values")

    invalid_raw_policy = json.loads(json.dumps(catalog))
    invalid_raw_policy["variables"][0]["value_source_policy"]["formula_cache_allowed"] = {
        "CI": [0.03, 0.04]}
    submit_invalid(targets_input, invalid_raw_policy, "raw-cache-policy", "prohibit formula-cache inputs")

    invalid_source_origin = json.loads(json.dumps(catalog))
    invalid_source_origin["value_source_policy"]["raw_source_origin"] = {"rates": [0.03, 0.04]}
    submit_invalid(targets_input, invalid_source_origin, "source-origin-policy", "raw_source_origin")

    invalid_question_id = json.loads(json.dumps(catalog))
    invalid_question_id["open_questions"] = [{"question_id": {"rates": [0.03, 0.04]}, "question": "Confirm."}]
    submit_invalid(targets_input, invalid_question_id, "question-id-values", "must be non-empty text")

    invalid_record_question = json.loads(json.dumps(catalog))
    invalid_record_question["source_records"] = [{
        "record_id": "fixture.source_constant", "group_id": "fixture_inputs",
        "classification": "source_constant", "open_question_id": {"rates": [0.03, 0.04]},
        "source_extents": [{"sheet": "Main", "range": "H1", "source_role": "raw_source"}],
    }]
    submit_invalid(targets_input, invalid_record_question, "record-question-values", "open_question_id must be non-empty text")

    invalid_axis_provenance = json.loads(json.dumps(catalog))
    invalid_axis_provenance["axis_metadata"] = [{
        "axis_id": "fixture_axis", "group_id": "fixture_inputs", "name": "Fixture axis",
        "shape": [1], "source_role": "raw_source",
        "source_extents": [{"sheet": "Main", "range": "H1", "source_role": "raw_source"}],
        "provenance": {"description": {"rates": [0.03, 0.04]}},
    }]
    submit_invalid(targets_input, invalid_axis_provenance, "axis-provenance-values", "needs source provenance")

    def scenario_variant(edit, suffix: str, expected_reason: str) -> None:
        invalid_targets = json.loads(json.dumps(targets_input))
        invalid_catalog = json.loads(json.dumps(catalog))
        scenario = invalid_targets["scenario"]
        edit(scenario)
        scenario_basis = {key: value for key, value in scenario.items() if key != "scenario_sha256"}
        scenario["scenario_sha256"] = conversion_workflow.hash_bytes(
            conversion_workflow.json_bytes(scenario_basis))
        invalid_catalog["scenario"] = scenario
        invalid_catalog["target_selection"]["scenario"] = scenario
        submit_invalid(invalid_targets, invalid_catalog, suffix, expected_reason)

    scenario_variant(lambda scenario: scenario["overrides"].update({"CI_rate": [0.123, 0.456]}),
                     "scenario-overrides", "scenario overrides must be an empty object")
    scenario_variant(lambda scenario: scenario["selectors"].update({"CI_rate": [0.123, 0.456]}),
                     "selector-array", "scenario selectors must be a non-empty map")
    scenario_variant(lambda scenario: scenario["selectors"].update({"CI_rate": 0.123}),
                     "selector-number", "scenario selectors must be a non-empty map")
    scenario_variant(lambda scenario: scenario.update({"rates": [0.123, 0.456]}),
                     "scenario-extra", "scenario metadata contains unsupported fields")

    invalid_source = json.loads(json.dumps(targets_input))
    invalid_source["source"]["rates"] = [0.123, 0.456]
    submit_invalid(invalid_source, catalog, "source-extra", "targets source contains unsupported fields")
    invalid_target_root = json.loads(json.dumps(targets_input))
    invalid_target_root["rates"] = [0.123, 0.456]
    submit_invalid(invalid_target_root, catalog, "targets-extra", "targets input contains unsupported fields")
    invalid_target_item = json.loads(json.dumps(targets_input))
    invalid_target_item["targets"][0]["rates"] = [0.123, 0.456]
    submit_invalid(invalid_target_item, catalog, "target-extra", "each target contains unsupported fields")


def _candidate_source_inventory() -> dict:
    inventory = _inventory()
    main = next(sheet for sheet in inventory["sheets"] if sheet["name"] == "Main")
    target = next(cell for cell in main["cells"] if cell["address"] == "G1")
    target["formula"] = '=IF(C3>0,B1,INDIRECT("Main!C4"))+SUM(F1:F2)+J1+J2+L1+D3'
    target["cached_value"] = 987654321
    target["cached_value_available"] = True
    lookup = next(cell for cell in main["cells"] if cell["address"] == "J1")
    lookup["formula"] = "=VLOOKUP(C3,H2:I3,2,TRUE)"
    extra_cells = [
        _cell("Main", "B1", 1, 2, "formula", formula="=C4+1", cached_value=987654322,
              cached_value_available=True),
        _cell("Main", "L1", 1, 12, "formula", formula="=TRANSPOSE(P1:P2)",
              raw_formula_attributes={"t": "array", "ref": "L1:M2"},
              cached_value=987654323, cached_value_available=True),
        _cell("Main", "M1", 1, 13, "value", 123456.1, formula=None,
              formula_present=False, cached_value=123456.1, cached_value_available=True),
        _cell("Main", "L2", 2, 12, "value", 123456.2, formula=None,
              formula_present=False, cached_value=123456.2, cached_value_available=True),
        _cell("Main", "M2", 2, 13, "value", 123456.3, formula=None,
              formula_present=False, cached_value=123456.3, cached_value_available=True),
        _cell("Main", "P1", 1, 16, "value", 0.2),
        _cell("Main", "P2", 2, 16, "value", 0.4),
    ]
    main["cells"].extend(extra_cells)
    main["max_column"] = 16
    inventory["workbook_ranges"].append({
        "source_location": _location("Main", "D3", "defined_name", "ExcludedDynamic"),
        "source_identity": None, "name": "ExcludedDynamic", "address": "D3", "kind": "defined_name",
        "metadata": {"scope": "workbook"},
    })
    for key in ("recognized_inventory_objects", "discovered_workbook_objects"):
        inventory["coverage"][key] += len(extra_cells) + 1
    return inventory


def _confirmed_candidate_fixture(tmp_path: Path, monkeypatch, inventory: dict) -> tuple[Path, Path, Path, Path, bytes]:
    workflow_root = tmp_path / "reviewed-workflow"
    conversion_workflow.create_workflow(
        workflow_root, {"source_id": "source-0", "run_id": "run-0", "workbook_sha256": "source-sha"})
    stage1 = conversion_workflow.append_stage_artifact(
        workflow_root, 1, "stage1.json", {"status": "pass"}, "Step 1")
    for reviewer in ("agent", "human"):
        conversion_workflow.record_decision(workflow_root, 1, reviewer, "approve", "Synthetic test fixture.")
    conversion_workflow.append_stage_artifact(
        workflow_root, 2, "stage2.json", {"status": "pass"}, "Step 2", input_stages={1: stage1["artifact"]})
    for reviewer in ("agent", "human"):
        conversion_workflow.record_decision(workflow_root, 2, reviewer, "approve", "Synthetic test fixture.")

    original_analysis, step1_root, prepared = _prepared_fixture(
        tmp_path / "prepared", monkeypatch, inventory=inventory, workflow_path=workflow_root)
    assert prepared["status"] == "pass"
    for stage in (workflow.build_fields, workflow.build_dependencies, workflow.build_plan):
        assert stage(original_analysis)["status"] == "pass"
    original_analysis_bytes = (original_analysis / "analysis.json").read_bytes()
    original_analysis_json = json.loads(original_analysis_bytes)
    binding = original_analysis_json["binding"]
    source = {"source_id": binding["source"]["source_id"], "run_id": binding["source"]["run_id"],
              "workbook_sha256": binding["source"]["source_sha256"],
              "analysis_binding_sha256": original_analysis_json["binding_sha256"]}
    scenario_basis = {"scenario_id": "fixture-saved-inputs", "selectors": {"source": "fixture"}, "overrides": {}}
    scenario = {**scenario_basis,
                "scenario_sha256": conversion_workflow.hash_bytes(conversion_workflow.json_bytes(scenario_basis))}
    target = {"target_id": "fixture_output", "selector": "Main!G1", "result_order": 0,
              "shape": [], "start_cells": ["Main!G1"]}
    target_selection = {"source": source, "scenario": scenario, "targets": [target]}
    current_fields = json.loads((original_analysis / "fields.json").read_text(encoding="utf-8"))["fields"]
    variables = []
    for field in current_fields:
        if field["role"] != "source":
            continue
        members = field["members"]
        kind = "scalar" if len(members) == 1 else "vector"
        variables.append({
            "variable_id": f"fixture.raw.{field['field_id']}", "group_id": "fixture_inputs",
            "logical_name": f"Fixture raw field {field['field_id']}", "role": "source_raw", "kind": kind,
            "logical_role": "business_input", "shape": [] if kind == "scalar" else [len(members)],
            "axes": [] if kind == "scalar" else [{"axis_id": "member_order", "name": "member_order", "role": "observation"}],
            "source_extents": [{"sheet": field["sheet"], "range": member["address"], "source_role": "raw_source"}
                               for member in members],
            "consumer_targets": ["fixture_output"],
            "provenance": {"description": "Exact synthetic source field membership."},
            "value_source_policy": {"origin": "workbook_formula_mode_source", "formula_cache_allowed": False},
        })
    variables.append({
        "variable_id": "external.fixture_rate", "group_id": "fixture_inputs",
        "logical_name": "Fixture formula-derived input", "role": "formula_derived_external", "kind": "scalar",
        "logical_role": "other", "shape": [], "axes": [],
        "source_extents": [{"sheet": "Main", "range": "B1", "source_role": "formula_output"}],
        "consumer_targets": ["fixture_output"],
        "provenance": {"description": "Synthetic formula-derived cut; value unavailable."},
        "value_source_policy": {"allowed_future_sources": ["approved_excel_capture", "user_rate_file"],
                                 "formula_cache_allowed": False, "available_now": False},
    })
    catalog = {
        "schema_version": "step3.input_boundary.input.v1", "source": source, "scenario": scenario,
        "target_selection": target_selection,
        "groups": [{"group_id": "fixture_inputs", "name": "Fixture inputs", "classification": "business_input"}],
        "variables": variables, "axis_metadata": [], "source_records": [],
        "upstream_exclusions": [{"boundary_variable_id": "external.fixture_rate", "named_source": "ExcludedDynamic",
                                 "source_range": "Main!D3", "reason": "Keep this source branch excluded in the fixture."}],
        "value_source_policy": {"formula_cache_allowed": False}, "open_questions": [],
    }
    targets_path, catalog_path = tmp_path / "targets.json", tmp_path / "catalog.json"
    _write_json(targets_path, {"schema_version": "step3.input_targets.v1", **target_selection})
    _write_json(catalog_path, catalog)
    result = CliRunner().invoke(app, ["step3", "input-catalog", "--analysis", str(original_analysis),
                                     "--workflow", str(workflow_root), "--targets", str(targets_path),
                                     "--catalog", str(catalog_path)])
    assert result.exit_code == 0, result.stdout
    for reviewer in ("agent", "human"):
        conversion_workflow.record_decision(workflow_root, 3, reviewer, "approve", "Synthetic test fixture.")

    # Use the same checked Step 1 and index inputs but a fresh analysis directory so the
    # accepted boundary's original analysis.json input remains immutable during Step 3 work.
    continuation = tmp_path / "continuation-analysis"
    continuation_result = workflow.prepare_analysis(tmp_path / "prepared" / "index.json", step1_root,
                                                    continuation, workflow_path=workflow_root)
    assert continuation_result["status"] == "pass"
    continuation_manifest = json.loads((continuation / "analysis.json").read_text(encoding="utf-8"))
    assert continuation_manifest["binding_sha256"] == original_analysis_json["binding_sha256"]
    for stage in (workflow.build_fields, workflow.build_dependencies, workflow.build_plan):
        assert stage(continuation)["status"] == "pass"
    return workflow_root, original_analysis, continuation, step1_root, original_analysis_bytes


def test_source_candidate_trace_and_profile_need_no_step4_history(tmp_path: Path, monkeypatch) -> None:
    workflow_root, original_analysis, continuation, _step1_root, original_analysis_bytes = _confirmed_candidate_fixture(
        tmp_path, monkeypatch, _candidate_source_inventory())
    trace_dir = tmp_path / "artifacts" / "candidate-trace"
    runner = CliRunner()
    source_trace = runner.invoke(app, ["step3", "source-trace", "--analysis", str(continuation),
                                       "--workflow", str(workflow_root), "--out", str(trace_dir)])
    assert source_trace.exit_code == 0, source_trace.output
    result = json.loads(source_trace.output)
    trace_path = Path(result["trace"])
    payload = json.loads(trace_path.read_text(encoding="utf-8"))
    assert payload["schema_version"] == "gp.source_candidate_trace.v1"
    assert payload["selection_basis"] == "static_candidate"
    assert payload["formula_calculation_performed"] is False
    assert payload["formula_cache_used"] is False
    assert payload["native_excel_called"] is False
    assert payload["summary"]["candidate_closure"] == "unknown"
    assert "candidate_scope" not in payload and "model_scope" not in payload
    assert "scope_frontiers" not in payload and "candidate_starts" not in payload
    assert payload["summary"]["reached_cut_formula_coordinate_count"] == 1
    assert payload["external_formula_cuts"][0]["formula_members"][0]["address"] == "Main!B1"
    assert payload["arrays"][0]["shape"] == [2, 2]
    assert payload["arrays"][0]["follower_count"] == 3
    categories = {item["category"] for item in payload["unknowns"]}
    assert {"conditional_branch_selection", "dynamic_reference", "lookup_selection_unknown",
            "unapproved_upstream_exclusion"} <= categories
    reachability = payload["raw_candidate_reachability"]
    assert reachability["status_counts"]["reached"] > 0
    assert reachability["status_counts"]["unresolved_possibly_reached"] > 0
    assert reachability["not_reached_is_known_only"] is False
    serialized = trace_path.read_text(encoding="utf-8")
    assert "987654321" not in serialized and "123456.1" not in serialized
    assert "cached_value" not in serialized and "cached_value_available" not in serialized
    assert not (workflow_root / "workflow.json").read_text(encoding="utf-8").find('"stage4"') >= 0

    profile_dir = tmp_path / "artifacts" / "candidate-profile"
    profiled = runner.invoke(app, ["step3", "profile", "--analysis", str(continuation),
                                   "--source-trace", str(trace_path), "--out", str(profile_dir)])
    assert profiled.exit_code == 0, profiled.output
    profile = json.loads((profile_dir / "source_candidate_family_profile.json").read_text(encoding="utf-8"))
    assert profile["schema_version"] == "gp.source_candidate_family_profile.v1"
    assert profile["selection_basis"] == "static_candidate"
    assert profile["summary"]["candidate_closure"] == "unknown"
    assert (profile_dir / "source_candidate_family_profile.md").is_file()

    before_profile = (profile_dir / "source_candidate_family_profile.json").read_bytes()
    tampered = json.loads(trace_path.read_text(encoding="utf-8"))
    tampered["targets"][0]["selector"] = "Main!F1"
    tampered_path = tmp_path / "tampered_candidate_trace.json"
    _write_json(tampered_path, tampered)
    blocked = runner.invoke(app, ["step3", "profile", "--analysis", str(continuation),
                                  "--source-trace", str(tampered_path), "--out", str(profile_dir)])
    assert blocked.exit_code == 1
    assert "target or scenario metadata changed" in json.loads(blocked.stdout)["diagnostics"][0]["message"]
    assert (profile_dir / "source_candidate_family_profile.json").read_bytes() == before_profile

    tampered_cut = json.loads(trace_path.read_text(encoding="utf-8"))
    tampered_cut["external_formula_cuts"][0]["formula_members"][0]["formula"] = "=C3+999"
    tampered_cut_path = tmp_path / "tampered_candidate_cut.json"
    _write_json(tampered_cut_path, tampered_cut)
    blocked_cut = runner.invoke(app, ["step3", "profile", "--analysis", str(continuation),
                                      "--source-trace", str(tampered_cut_path), "--out", str(profile_dir)])
    assert blocked_cut.exit_code == 1
    assert "external formula cuts differ" in json.loads(blocked_cut.stdout)["diagnostics"][0]["message"]
    assert (profile_dir / "source_candidate_family_profile.json").read_bytes() == before_profile
    assert (original_analysis / "analysis.json").read_bytes() == original_analysis_bytes


def test_source_candidate_model_scope_clips_ranges_and_binds_profile_evidence(tmp_path: Path, monkeypatch) -> None:
    workflow_root, _original_analysis, continuation, _step1_root, _original_analysis_bytes = _confirmed_candidate_fixture(
        tmp_path, monkeypatch, _candidate_source_inventory())
    runner = CliRunner()
    full_dir = tmp_path / "artifacts" / "full-candidate-trace"
    full_result = runner.invoke(app, ["step3", "source-trace", "--analysis", str(continuation),
                                     "--workflow", str(workflow_root), "--out", str(full_dir)])
    assert full_result.exit_code == 0, full_result.output
    full_payload = json.loads(full_result.output)
    full_trace_path = Path(full_payload["trace"])
    full_bytes = full_trace_path.read_bytes()
    full_trace = json.loads(full_bytes)
    scope_path = tmp_path / "agent-model-scope.json"
    scope = {
        "schema_version": "gp.source_model_scope.input.v1",
        "status": "agent_proposed",
        "source": full_trace["source"],
        "scenario": full_trace["scenario"],
        "targets": full_trace["targets"],
        "analysis": full_trace["analysis"],
        "input_boundary": full_trace["input_boundary"],
        "full_candidate_trace": {"path": str(full_trace_path), "sha256": workflow._hash_file(full_trace_path)},
        "allowed_regions": [
            {"source_ref": "Main!B1", "rationale": "Confirmed formula-derived input cut."},
            {"source_ref": "Main!C3", "rationale": "Required target input."},
            {"source_ref": "Main!F1", "rationale": "First in-scope member of a source range."},
            {"source_ref": "Main!G1", "rationale": "Confirmed target start."},
            {"source_ref": "Main!J1", "rationale": "Additional candidate formula start."},
        ],
        "proposed_formula_starts": [
            {"source_ref": "Main!J1", "rationale": "Trace the saved lookup formula as a proposed candidate root."},
        ],
    }
    _write_json(scope_path, scope)
    scoped_dir = tmp_path / "artifacts" / "scoped-candidate-trace"
    scoped_result = runner.invoke(app, ["step3", "source-trace", "--analysis", str(continuation),
                                        "--workflow", str(workflow_root), "--model-scope", str(scope_path),
                                        "--out", str(scoped_dir)])
    assert scoped_result.exit_code == 0, scoped_result.output
    scoped_payload = json.loads(scoped_result.output)
    trace_path = Path(scoped_payload["trace"])
    trace = json.loads(trace_path.read_text(encoding="utf-8"))
    assert trace["candidate_scope"] == "scope_bounded_static_candidate_design"
    assert trace["summary"]["candidate_closure"] == "unknown"
    assert trace["summary"]["candidate_formula_selection_complete_for_known_references"] is False
    assert trace["summary"]["fresh_step4_active_formula_members"] is None
    assert trace["formula_calculation_performed"] is False
    assert trace["formula_cache_used"] is False
    assert trace["candidate_starts"]["target_starts"] == [{"target_id": "fixture_output", "source_ref": "Main!G1"}]
    proposed = trace["candidate_starts"]["proposed_formula_starts"]
    assert [item["source_ref"] for item in proposed] == ["Main!J1"]
    assert proposed[0]["formula_members"][0]["address"] == "Main!J1"
    candidate_addresses = {item["address"] for item in trace["cells"]}
    assert "Main!G1" in candidate_addresses and "Main!J1" in candidate_addresses
    assert all(item["address"] != "Main!J2" for item in trace["cells"])
    assert any(item["reference"] == "F1:F2" and item["excluded_source_ref"] == "Main!F2"
               for item in trace["scope_frontiers"])
    assert trace["summary"]["scope_frontier_count"] == len(trace["scope_frontiers"])
    assert full_trace_path.read_bytes() == full_bytes

    profile_dir = tmp_path / "artifacts" / "scoped-profile"
    profile_result = runner.invoke(app, ["step3", "profile", "--analysis", str(continuation),
                                        "--source-trace", str(trace_path), "--out", str(profile_dir)])
    assert profile_result.exit_code == 0, profile_result.output
    profile_path = profile_dir / "source_candidate_family_profile.json"
    profile_before = profile_path.read_bytes()
    profile = json.loads(profile_before)
    assert profile["candidate_scope"] == trace["candidate_scope"]
    assert profile["model_scope"] == trace["model_scope"]
    assert "Candidate scope:" in (profile_dir / "source_candidate_family_profile.md").read_text(encoding="utf-8")

    analysis_json = json.loads((continuation / "analysis.json").read_text(encoding="utf-8"))
    _current_analysis, manifest, binding, _step1_root, _inventory_path = workflow._load_context(continuation)
    _fields, fields_sha = workflow._read_stage(continuation, manifest, "fields")
    _dependencies, dependencies_sha = workflow._read_stage(continuation, manifest, "dependencies")
    variables = []
    equation_families = []
    source_mappings = []
    for index, family in enumerate(profile["families"]):
        assert len(family["source_members"]) == 1
        family_id = family["family_id"]
        variable_id = f"candidate.scalar_{index}"
        equation_id = f"candidate.equation_{index}"
        segment_id = f"candidate.segment_{index}"
        variables.append({
            "variable_id": variable_id, "role": "derived", "kind": "scalar", "shape": [], "axes": [],
            "source_extents": [{"ref": family["source_members"][0], "role": "formula_output",
                                "coordinate_mapping": {"kind": "scalar", "axes": []}}],
            "demanded_indices": "all", "initial_condition": None,
            "equation_segments": [{"segment_id": segment_id}], "dependencies": [],
            "error_policy": "Preserve source formula errors.", "output_usage": "Static scope validation fixture.",
        })
        equation_families.append({"family_id": equation_id, "module_id": "scope_module",
                                  "function_name": f"candidate_formula_{index}",
                                  "description": "One formula family used to validate scope evidence.",
                                  "source_family_ids": [family_id]})
        source_mappings.append({"source_family_id": family_id, "equation_family_id": equation_id,
                                "variable_id": variable_id, "module_id": "scope_module",
                                "equation_segment": segment_id,
                                "index_mapping": {"kind": "scalar", "axes": []}})
    variables.extend([
        {"variable_id": "input.issue_age", "role": "raw", "kind": "scalar", "shape": [], "axes": [],
         "source_extents": [{"ref": "Main!C3", "role": "raw_input",
                              "coordinate_mapping": {"kind": "scalar", "axes": []}}],
         "demanded_indices": [], "initial_condition": None, "equation_segments": [], "dependencies": [],
         "error_policy": "Preserve raw source values.", "output_usage": "Static scope validation fixture."},
        {"variable_id": "external.fixture_rate", "boundary_variable_id": "external.fixture_rate",
         "role": "external", "kind": "scalar", "shape": [], "axes": [],
         "source_extents": [{"ref": "Main!B1", "role": "external_input",
                              "coordinate_mapping": {"kind": "scalar", "axes": []}}],
         "demanded_indices": [], "initial_condition": None, "equation_segments": [], "dependencies": [],
         "error_policy": "Keep the confirmed external value unavailable.",
         "output_usage": "Static scope validation fixture."},
    ])
    semantic_map_path = tmp_path / "scoped_semantic_map.json"
    _write_json(semantic_map_path, {
        "schema_version": "step3.semantic_map.input.v1", "selection_basis": "static_candidate",
        "source": {"source_id": binding["source"]["source_id"], "run_id": binding["source"]["run_id"],
                   "source_sha256": binding["source"]["source_sha256"],
                   "binding_sha256": manifest["binding_sha256"],
                   "fields_sha256": fields_sha, "dependencies_sha256": dependencies_sha},
        "source_family_profile": {"path": str(profile_path), "sha256": workflow._hash_file(profile_path)},
        "source_candidate_trace": {"path": str(trace_path), "sha256": workflow._hash_file(trace_path)},
        "modules": [{"module_id": "scope_module", "order": 1, "purpose": "Validate scope-bound candidate evidence."}],
        "variables": variables, "equation_families": equation_families,
        "source_mappings": source_mappings, "array_mappings": [],
    })
    semantic_result = semantic.build_semantic_plan(continuation, semantic_map_path)
    assert semantic_result["status"] == "pass", semantic_result
    plan_path = continuation / "semantic_plan.json"
    plan_before = plan_path.read_bytes()
    plan = json.loads(plan_before)
    assert plan["candidate_scope"] == trace["candidate_scope"]
    assert plan["coverage"]["candidate_closure"] == "unknown"
    assert plan["readiness"]["generation_ready"] is False
    assert plan["evidence"]["model_scope"] == {
        "path": trace["model_scope"]["path"], "sha256": trace["model_scope"]["sha256"]}
    checked = semantic.validate_semantic_plan(continuation)
    assert checked["status"] == "pass", checked
    check = json.loads((continuation / "semantic_check.json").read_text(encoding="utf-8"))
    assert check["candidate_scope"] == trace["candidate_scope"]
    assert check["model_scope_sha256"] == trace["model_scope"]["sha256"]
    direct_inputs = design._semantic_evidence_inputs(continuation, plan, check)
    assert direct_inputs["step3_model_scope"] == scope_path.resolve()

    scope["allowed_regions"].append({"source_ref": "Main!E99", "rationale": "Changed after tracing."})
    _write_json(scope_path, scope)
    stale_profile = runner.invoke(app, ["step3", "profile", "--analysis", str(continuation),
                                       "--source-trace", str(trace_path), "--out", str(profile_dir)])
    assert stale_profile.exit_code == 1
    assert "model-scope input is missing or changed" in json.loads(stale_profile.stdout)["diagnostics"][0]["message"]
    assert profile_path.read_bytes() == profile_before
    semantic_plan_before = plan_path.read_bytes()
    semantic_result = semantic.build_semantic_plan(continuation, semantic_map_path)
    assert semantic_result["status"] == "blocked"
    assert "model-scope input is missing or changed" in semantic_result["diagnostics"][0]["message"]
    assert plan_path.read_bytes() == semantic_plan_before == plan_before
    assert analysis_json["binding_sha256"] == manifest["binding_sha256"]


def test_source_candidate_model_scope_rejects_nonformula_starts_and_requires_cuts(tmp_path: Path, monkeypatch) -> None:
    workflow_root, _original_analysis, continuation, _step1_root, _original_analysis_bytes = _confirmed_candidate_fixture(
        tmp_path, monkeypatch, _candidate_source_inventory())
    full_dir = tmp_path / "artifacts" / "full-candidate-trace"
    full_result = profiling.build_source_candidate_trace(continuation, workflow_root, full_dir)
    assert full_result["status"] == "pass", full_result
    full_trace_path = Path(full_result["trace"])
    full_trace = json.loads(full_trace_path.read_text(encoding="utf-8"))
    common = {
        "schema_version": "gp.source_model_scope.input.v1", "status": "agent_proposed",
        "source": full_trace["source"], "scenario": full_trace["scenario"], "targets": full_trace["targets"],
        "analysis": full_trace["analysis"], "input_boundary": full_trace["input_boundary"],
        "full_candidate_trace": {"path": str(full_trace_path), "sha256": workflow._hash_file(full_trace_path)},
        "allowed_regions": [
            {"source_ref": "Main!B1", "rationale": "Confirmed cut."},
            {"source_ref": "Main!C3", "rationale": "Literal source coordinate."},
            {"source_ref": "Main!G1", "rationale": "Target."},
        ],
    }
    scope_path = tmp_path / "bad_model_scope.json"
    bad_start = {**common, "proposed_formula_starts": [
        {"source_ref": "Main!C3", "rationale": "A literal source cell is not a formula start."},
    ]}
    _write_json(scope_path, bad_start)
    out_dir = tmp_path / "artifacts" / "bad-start-output"
    bad_start_result = profiling.build_source_candidate_trace(continuation, workflow_root, out_dir, scope_path)
    assert bad_start_result["status"] == "blocked"
    assert "non-formula source coordinate" in bad_start_result["diagnostics"][0]["message"]
    assert not (out_dir / "source_candidate_trace.json").exists()

    missing_cut = {**common, "allowed_regions": [{"source_ref": "Main!G1", "rationale": "Target only."}],
                   "proposed_formula_starts": []}
    _write_json(scope_path, missing_cut)
    missing_cut_result = profiling.build_source_candidate_trace(continuation, workflow_root, out_dir, scope_path)
    assert missing_cut_result["status"] == "blocked"
    assert "contain every confirmed formula-derived input cut coordinate" in missing_cut_result["diagnostics"][0]["message"]
    assert not (out_dir / "source_candidate_trace.json").exists()


def test_source_candidate_model_scope_clips_whole_column_before_expansion(tmp_path: Path, monkeypatch) -> None:
    inventory = _candidate_source_inventory()
    main = next(sheet for sheet in inventory["sheets"] if sheet["name"] == "Main")
    target = next(cell for cell in main["cells"] if cell["address"] == "G1")
    target["formula"] = "=SUM(F:F)"
    workflow_root, _original_analysis, continuation, _step1_root, _original_analysis_bytes = _confirmed_candidate_fixture(
        tmp_path, monkeypatch, inventory)
    full_result = profiling.build_source_candidate_trace(continuation, workflow_root,
                                                         tmp_path / "artifacts" / "full")
    assert full_result["status"] == "pass", full_result
    full_trace_path = Path(full_result["trace"])
    full = json.loads(full_trace_path.read_text(encoding="utf-8"))
    scope_path = tmp_path / "whole_column_scope.json"
    _write_json(scope_path, {
        "schema_version": "gp.source_model_scope.input.v1", "status": "agent_proposed",
        "source": full["source"], "scenario": full["scenario"], "targets": full["targets"],
        "analysis": full["analysis"], "input_boundary": full["input_boundary"],
        "full_candidate_trace": {"path": str(full_trace_path), "sha256": workflow._hash_file(full_trace_path)},
        "allowed_regions": [
            {"source_ref": "Main!B1", "rationale": "Confirmed formula-derived cut."},
            {"source_ref": "Main!F1:F2", "rationale": "Small candidate portion of a whole-column read."},
            {"source_ref": "Main!G1", "rationale": "Target start."},
        ],
        "proposed_formula_starts": [],
    })
    result = profiling.build_source_candidate_trace(continuation, workflow_root,
                                                    tmp_path / "artifacts" / "scoped", scope_path)
    assert result["status"] == "pass", result
    trace = json.loads(Path(result["trace"]).read_text(encoding="utf-8"))
    assert {item["address"] for item in trace["cells"]} == {"Main!F1", "Main!F2", "Main!G1"}
    assert any(item["reference"] == "F:F" and item["excluded_source_ref"] == "Main!F3:F1048576"
               and item["coordinate_count"] == 1_048_574 for item in trace["scope_frontiers"])
    assert trace["summary"]["candidate_closure"] == "unknown"
    assert trace["summary"]["candidate_formula_selection_complete_for_known_references"] is False


def test_source_candidate_model_scope_clips_whole_row_before_expansion(tmp_path: Path, monkeypatch) -> None:
    inventory = _candidate_source_inventory()
    main = next(sheet for sheet in inventory["sheets"] if sheet["name"] == "Main")
    target = next(cell for cell in main["cells"] if cell["address"] == "G1")
    target["formula"] = "=SUM(1:1)"
    workflow_root, _original_analysis, continuation, _step1_root, _original_analysis_bytes = _confirmed_candidate_fixture(
        tmp_path, monkeypatch, inventory)
    full_result = profiling.build_source_candidate_trace(continuation, workflow_root,
                                                         tmp_path / "artifacts" / "full")
    assert full_result["status"] == "pass", full_result
    full_trace_path = Path(full_result["trace"])
    full = json.loads(full_trace_path.read_text(encoding="utf-8"))
    scope_path = tmp_path / "whole_row_scope.json"
    _write_json(scope_path, {
        "schema_version": "gp.source_model_scope.input.v1", "status": "agent_proposed",
        "source": full["source"], "scenario": full["scenario"], "targets": full["targets"],
        "analysis": full["analysis"], "input_boundary": full["input_boundary"],
        "full_candidate_trace": {"path": str(full_trace_path), "sha256": workflow._hash_file(full_trace_path)},
        "allowed_regions": [
            {"source_ref": "Main!B1", "rationale": "Confirmed formula-derived cut."},
            {"source_ref": "Main!G1", "rationale": "Target start and one in-scope row coordinate."},
        ],
        "proposed_formula_starts": [],
    })
    result = profiling.build_source_candidate_trace(continuation, workflow_root,
                                                    tmp_path / "artifacts" / "scoped", scope_path)
    assert result["status"] == "pass", result
    trace = json.loads(Path(result["trace"]).read_text(encoding="utf-8"))
    assert {item["address"] for item in trace["cells"]} == {"Main!G1"}
    row_frontiers = [item for item in trace["scope_frontiers"] if item["reference"] == "1:1"]
    assert sum(item["coordinate_count"] for item in row_frontiers) == 16_382
    assert all(item["extent_kind"] == "region" for item in row_frontiers)
    assert trace["summary"]["candidate_closure"] == "unknown"
    assert trace["summary"]["candidate_formula_selection_complete_for_known_references"] is False


def test_source_candidate_model_scope_reuses_range_expansion_but_keeps_each_frontier(tmp_path: Path, monkeypatch) -> None:
    inventory = _candidate_source_inventory()
    main = next(sheet for sheet in inventory["sheets"] if sheet["name"] == "Main")
    for address in ("G1", "J1"):
        cell = next(item for item in main["cells"] if item["address"] == address)
        cell["formula"] = "=SUM(F1:F3)"
    workflow_root, _original_analysis, continuation, _step1_root, _original_analysis_bytes = _confirmed_candidate_fixture(
        tmp_path, monkeypatch, inventory)
    full_result = profiling.build_source_candidate_trace(continuation, workflow_root,
                                                         tmp_path / "artifacts" / "full")
    assert full_result["status"] == "pass", full_result
    full_trace_path = Path(full_result["trace"])
    full = json.loads(full_trace_path.read_text(encoding="utf-8"))
    scope_path = tmp_path / "shared_range_scope.json"
    _write_json(scope_path, {
        "schema_version": "gp.source_model_scope.input.v1", "status": "agent_proposed",
        "source": full["source"], "scenario": full["scenario"], "targets": full["targets"],
        "analysis": full["analysis"], "input_boundary": full["input_boundary"],
        "full_candidate_trace": {"path": str(full_trace_path), "sha256": workflow._hash_file(full_trace_path)},
        "allowed_regions": [
            {"source_ref": "Main!B1", "rationale": "Confirmed formula-derived cut."},
            {"source_ref": "Main!C3", "rationale": "Required source coordinate."},
            {"source_ref": "Main!F1", "rationale": "One in-scope member of the shared read."},
            {"source_ref": "Main!G1", "rationale": "Target start."},
            {"source_ref": "Main!J1", "rationale": "Proposed formula start."},
        ],
        "proposed_formula_starts": [
            {"source_ref": "Main!J1", "rationale": "Inspect a second consumer of the same source range."},
        ],
    })
    original_clip = profiling._clip_to_scope
    original_contains = profiling._scope_contains_key
    matching_clips = []
    included_coordinate_checks = []

    def count_shared_range_clip(bounds, allowed):
        if bounds == (6, 1, 6, 3):
            matching_clips.append(bounds)
        return original_clip(bounds, allowed)

    def count_included_coordinate_check(regions, address_key):
        if address_key == "main!F1":
            included_coordinate_checks.append(address_key)
        return original_contains(regions, address_key)

    monkeypatch.setattr(profiling, "_clip_to_scope", count_shared_range_clip)
    monkeypatch.setattr(profiling, "_scope_contains_key", count_included_coordinate_check)
    result = profiling.build_source_candidate_trace(continuation, workflow_root,
                                                    tmp_path / "artifacts" / "scoped", scope_path)
    assert result["status"] == "pass", result
    trace = json.loads(Path(result["trace"]).read_text(encoding="utf-8"))
    assert matching_clips == [(6, 1, 6, 3)]
    assert included_coordinate_checks == ["main!F1"]
    frontiers = [item for item in trace["scope_frontiers"] if item["reference"] == "F1:F3"]
    assert {(item["referring_formula"], item["excluded_source_ref"], item["coordinate_count"])
            for item in frontiers} == {
        ("Main!G1", "Main!F2:F3", 2),
        ("Main!J1", "Main!F2:F3", 2),
    }
    assert {item["address"] for item in trace["cells"]} == {"Main!F1", "Main!G1", "Main!J1"}


def test_static_candidate_semantic_plan_preserves_catalog_gate_and_continuation_binding(tmp_path: Path, monkeypatch) -> None:
    workflow_root, original_analysis, continuation, _step1_root, original_analysis_bytes = _confirmed_candidate_fixture(
        tmp_path, monkeypatch, _candidate_source_inventory())
    trace_dir, profile_dir = tmp_path / "artifacts" / "candidate-trace", tmp_path / "artifacts" / "candidate-profile"
    trace_result = profiling.build_source_candidate_trace(continuation, workflow_root, trace_dir)
    assert trace_result["status"] == "pass", trace_result
    profile_result = profiling.profile_source_candidates(continuation, Path(trace_result["trace"]), profile_dir)
    assert profile_result["status"] == "pass", profile_result
    profile_path = Path(profile_result["profile"])
    profile = json.loads(profile_path.read_text(encoding="utf-8"))
    analysis = json.loads((continuation / "analysis.json").read_text(encoding="utf-8"))
    _fields, fields_sha = workflow._read_stage(continuation, analysis, "fields")
    _dependencies, dependencies_sha = workflow._read_stage(continuation, analysis, "dependencies")
    variables: list[dict] = []
    source_mappings: list[dict] = []
    equation_families: list[dict] = []
    array_mappings: list[dict] = []
    array_by_anchor = {item["anchor"]: item for item in json.loads(Path(trace_result["trace"]).read_text(encoding="utf-8"))["arrays"]}
    for index, family in enumerate(profile["families"]):
        family_id = family["family_id"]
        variable_id = f"candidate.formula_{index + 1}"
        members = family["source_members"]
        array = array_by_anchor.get(family["representative_address"])
        if array:
            kind, shape = "matrix", array["shape"]
            axes = [{"name": "array_row", "role": "index"}, {"name": "array_column", "role": "index"}]
            index_mapping = {"kind": "matrix", "array_ref": array["array_ref"], "axes": [
                {"axis": "array_row", "dimension": 0, "source_coordinate": "row",
                 "origin": int("".join(ch for ch in array["array_ref"].split(":")[0] if ch.isdigit())), "index_origin": 0},
                {"axis": "array_column", "dimension": 1, "source_coordinate": "column",
                 "origin": "L", "index_origin": 0},
            ]}
        elif len(members) > 1:
            kind, shape = "projection_series", [len(members)]
            axes = [{"name": "source_row", "role": "index"}]
            index_mapping = {"kind": "row_ordinal", "axes": [
                {"axis": "source_row", "dimension": 0, "source_coordinate": "row",
                 "origin": 1, "index_origin": 0}]}
        else:
            kind, shape, axes = "scalar", [], []
            index_mapping = {"kind": "scalar", "axes": []}
        extent_ref = array["anchor"] if array else members[0]
        source_extents = [{"ref": extent_ref, "role": "formula_output",
                           "coordinate_mapping": index_mapping}]
        variables.append({"variable_id": variable_id, "role": "derived", "kind": kind, "shape": shape,
                          "axes": axes, "source_extents": source_extents, "demanded_indices": "all",
                          "initial_condition": None,
                          "equation_segments": [{"segment_id": f"segment_{index + 1}"}],
                          "dependencies": [], "error_policy": "Preserve source errors.",
                          "output_usage": "Static candidate fixture."})
        equation_id = f"equation_{index + 1}"
        equation_families.append({"family_id": equation_id, "function_name": f"formula_{index + 1}",
                                  "module_id": "candidate_module", "description": "Synthetic source formula family.",
                                  "source_family_ids": [family_id]})
        source_mappings.append({"source_family_id": family_id, "equation_family_id": equation_id,
                                "variable_id": variable_id, "module_id": "candidate_module",
                                "equation_segment": f"segment_{index + 1}", "index_mapping": index_mapping})
        if array:
            array_mappings.append({"array_family_id": "candidate_arrays", "module_id": "candidate_module",
                                   "instances": [{"anchor": array["anchor"], "array_ref": array["array_ref"],
                                                  "anchor_source_family_id": family_id, "shape": array["shape"],
                                                  "orientation": "source row-column array", "variable_id": variable_id,
                                                  "module_id": "candidate_module", "axes": axes}]})
    for address in ("Main!C3", "Main!C4", "Main!P1", "Main!P2"):
        variables.append({"variable_id": "input." + address.split("!")[1].lower(), "role": "raw",
                          "kind": "scalar", "shape": [], "axes": [],
                          "source_extents": [{"ref": address, "role": "raw_input",
                                               "coordinate_mapping": {"kind": "scalar", "axes": []}}],
                          "demanded_indices": [], "initial_condition": None, "equation_segments": [],
                          "dependencies": [], "error_policy": "Preserve raw source value.",
                          "output_usage": "Synthetic source reference."})
    variables.append({"variable_id": "external.fixture_rate", "boundary_variable_id": "external.fixture_rate",
                      "role": "external", "kind": "scalar", "shape": [], "axes": [],
                      "source_extents": [{"ref": "Main!B1", "role": "external_input",
                                           "coordinate_mapping": {"kind": "scalar", "axes": []}}],
                      "demanded_indices": [], "initial_condition": None, "equation_segments": [],
                      "dependencies": [], "error_policy": "Value remains unavailable.",
                      "output_usage": "Confirmed external source boundary."})
    semantic_map = {"schema_version": "step3.semantic_map.input.v1", "selection_basis": "static_candidate",
                    "source": {"source_id": analysis["binding"]["source"]["source_id"],
                               "run_id": analysis["binding"]["source"]["run_id"],
                               "source_sha256": analysis["binding"]["source"]["source_sha256"],
                               "binding_sha256": analysis["binding_sha256"],
                               "fields_sha256": fields_sha, "dependencies_sha256": dependencies_sha},
                    "source_family_profile": {"path": str(profile_path), "sha256": workflow._hash_file(profile_path)},
                    "source_candidate_trace": {"path": trace_result["trace"], "sha256": trace_result["trace_sha256"]},
                    "modules": [{"module_id": "candidate_module", "order": 1,
                                 "purpose": "Map synthetic candidate formulas."}],
                    "variables": variables, "equation_families": equation_families,
                    "source_mappings": source_mappings, "array_mappings": array_mappings}
    semantic_map_path = tmp_path / "candidate_semantic_map.json"
    _write_json(semantic_map_path, semantic_map)
    planned = semantic.build_semantic_plan(continuation, semantic_map_path)
    assert planned["status"] == "pass", planned
    semantic_plan = json.loads((continuation / "semantic_plan.json").read_text(encoding="utf-8"))
    assert semantic_plan["selection_basis"] == "static_candidate"
    assert semantic_plan["coverage"]["candidate_closure"] == "unknown"
    assert semantic_plan["coverage"]["active_formula_member_count"] is None
    assert len(semantic_plan["external_boundary_variables"]) == 1
    assert semantic_plan["readiness"]["generation_ready"] is False
    assert "Candidate closure: **unknown**" in (continuation / "semantic_plan.md").read_text(encoding="utf-8")
    checked = semantic.validate_semantic_plan(continuation)
    assert checked["status"] == "pass", checked
    check = json.loads((continuation / "semantic_check.json").read_text(encoding="utf-8"))
    assert check["selection_basis"] == "static_candidate"
    assert check["readiness"]["generation_ready"] is False
    assert "Active formula selection or closure proven: no" in (continuation / "semantic_check.md").read_text(encoding="utf-8")
    assert (original_analysis / "analysis.json").read_bytes() == original_analysis_bytes
    _current_dir, current_analysis, current_binding, _root, _inventory_path = workflow._load_context(original_analysis)
    from excel_to_act.steps.step3.input_boundary import require_confirmed_boundary_for_analysis
    _, boundary_state = require_confirmed_boundary_for_analysis(original_analysis, current_analysis, current_binding)
    assert boundary_state["state"]["status"] == "input_boundary_confirmed"


def test_stage_commands_reject_changed_bound_manifest_before_writing(tmp_path: Path, monkeypatch) -> None:
    out, root, prepared = _prepared_fixture(tmp_path, monkeypatch)
    assert prepared["status"] == "pass"
    manifest_path = root / "workbook_manifest.json"
    manifest_path.write_bytes(manifest_path.read_bytes() + b"tampered")

    for stage in (workflow.build_fields, workflow.build_dependencies, workflow.build_plan, workflow.validate_analysis):
        result = stage(out)
        assert result["status"] == "blocked"
        assert "workbook_manifest.json changed after step3 prepare" in result["diagnostics"][0]["message"]
    assert not (out / "fields.json").exists()
    assert not (out / "dependencies.json").exists()
    assert not (out / "execution_plan.json").exists()


def test_unbounded_defined_names_remain_unresolved_without_blocking_workflow(tmp_path: Path, monkeypatch) -> None:
    inventory = _inventory()
    inventory["sheets"][0]["cells"].extend([
        _cell("Main", "M1", 1, 13, "formula", formula="=SUM(EntireColumn)"),
        _cell("Main", "M2", 2, 13, "formula", formula="=SUM(EntireRow)"),
    ])
    inventory["workbook_ranges"].extend([
        {"source_location": _location("Main", "A:A", "defined_name", "EntireColumn"), "source_identity": None,
         "name": "EntireColumn", "address": "A:A", "kind": "defined_name", "metadata": {"scope": "workbook"}},
        {"source_location": _location("Main", "1:1", "defined_name", "EntireRow"), "source_identity": None,
         "name": "EntireRow", "address": "1:1", "kind": "defined_name", "metadata": {"scope": "workbook"}},
    ])
    inventory["coverage"]["recognized_inventory_objects"] += 4
    inventory["coverage"]["discovered_workbook_objects"] += 4
    out, _root, prepared = _prepared_fixture(tmp_path, monkeypatch, inventory=inventory)
    assert prepared["status"] == "pass"
    assert workflow.build_fields(out)["status"] == "pass"
    assert workflow.build_dependencies(out)["status"] == "pass"
    dependencies = json.loads((out / "dependencies.json").read_text(encoding="utf-8"))

    unbounded = [item for item in dependencies["boundaries"] if item["boundary_type"] == "unsupported_unbounded_range"]
    assert {(item["evidence"], item["target_details"]["address"]) for item in unbounded} == {
        ("EntireColumn", "A:A"), ("EntireRow", "1:1")
    }
    assert all(item["status"] == "unresolved" for item in unbounded)
    assert workflow.build_plan(out)["status"] == "pass"
    checked = workflow.validate_analysis(out)
    assert checked["status"] == "pass"
    assert checked["handoff"]["readiness"]["generation_ready"] is False


def test_prepare_rejects_wrong_source_scope_and_ambiguous_index(tmp_path: Path, monkeypatch) -> None:
    scope = {"schema_version": "analysis.scope.v1", "workbook_sha256": "another-workbook", "ignored_sheets": ["Ignored"]}
    _out, _root, wrong_scope = _prepared_fixture(tmp_path / "scope", monkeypatch, scope=scope)
    assert wrong_scope["status"] == "blocked"
    assert "does not match" in wrong_scope["diagnostics"][0]["message"]

    _out, _root, ambiguous = _prepared_fixture(tmp_path / "many", monkeypatch, entries=2, select_source=False)
    assert ambiguous["status"] == "blocked"
    assert "multiple eligible sources" in ambiguous["diagnostics"][0]["message"]


def _ready_analysis(tmp_path: Path, monkeypatch, *, scope: dict | None = None, inventory: dict | None = None) -> Path:
    out, _root, prepared = _prepared_fixture(tmp_path, monkeypatch, scope=scope, inventory=inventory)
    assert prepared["status"] == "pass"
    assert workflow.build_fields(out)["status"] == "pass"
    assert workflow.build_dependencies(out)["status"] == "pass"
    assert workflow.build_plan(out)["status"] == "pass"
    return out


def _packet_facts(packet: dict, item: dict) -> dict:
    return packet["evidence"][item["evidence_id"]]["facts"]


def test_query_pages_exact_selectors_and_does_not_invent_cells(tmp_path: Path, monkeypatch) -> None:
    inventory = _inventory()
    inventory["vba_modules"] = [{"code": "Sub Run()\nEnd Sub", "kind": "standard", "name": "Module1", "procedures": ["Run"]}]
    out = _ready_analysis(tmp_path, monkeypatch, inventory=inventory)

    cell_result = exploration.query(out, "Main!C3")
    assert cell_result["status"] == "pass"
    packet = cell_result["packet"]
    assert packet["schema_version"] == "step3.query.v1"
    assert packet["page"]["total_matches"] == packet["page"]["shown"] == 1
    assert _packet_facts(packet, packet["records"][0])["value"] == 10
    assert packet["fields"][0]["summary"]["names"] == ["GP"]
    assert packet["stage_hashes"]["fields"] == json.loads((out / "analysis.json").read_text(encoding="utf-8"))["stage_hashes"]["fields"]

    page_one = exploration.query(out, "Main!H1:I3", offset=1, limit=2)["packet"]
    page_repeat = exploration.query(out, "Main!H1:I3", offset=1, limit=2)["packet"]
    assert [item["summary"]["address"] for item in page_one["records"]] == ["I1", "H2"]
    assert page_one["page"]["next_offset"] == 3
    assert page_one == page_repeat

    empty = exploration.query(out, "Main!A100")["packet"]
    assert empty["records"] == []
    assert empty["unknowns"] and empty["page"]["total_matches"] == 0
    assert exploration.query(out, "A1")["status"] == "blocked"
    assert exploration.query(out, "Main!A:A")["status"] == "blocked"

    module = exploration.query(out, "vba:Module1")["packet"]
    assert _packet_facts(module, module["records"][0]) == inventory["vba_modules"][0]
    assert "not executed" in module["unknowns"][0]
    assert exploration.trace(out, "vba:Module1")["status"] == "blocked"


def test_query_resolves_scoped_names_and_blocks_excluded_sheet(tmp_path: Path, monkeypatch) -> None:
    inventory = _inventory()
    for sheet, address in (("Main", "H2"), ("Ignored", "A1")):
        inventory["workbook_ranges"].append({
            "source_location": _location(sheet, address, "defined_name", "LocalRate"),
            "source_identity": None, "name": "LocalRate", "address": address, "kind": "defined_name",
            "metadata": {"scope": "worksheet"},
        })
    inventory["workbook_ranges"].extend([
        {"source_location": _location("Main", "C3", "defined_name", "ScopedOrGlobal"), "source_identity": None,
         "name": "ScopedOrGlobal", "address": "C3", "kind": "defined_name", "metadata": {"scope": "workbook"}},
        {"source_location": _location("Main", "H2", "defined_name", "ScopedOrGlobal"), "source_identity": None,
         "name": "ScopedOrGlobal", "address": "H2", "kind": "defined_name", "metadata": {"scope": "worksheet"}},
    ])
    scope = {"schema_version": "analysis.scope.v1", "workbook_sha256": "source-sha", "ignored_sheets": ["Ignored"]}
    out = _ready_analysis(tmp_path, monkeypatch, scope=scope, inventory=inventory)

    ambiguous = exploration.query(out, "LocalRate")
    assert ambiguous["status"] == "blocked" and ambiguous["needs_selection"] is True
    selected = exploration.query(out, "LocalRate", sheet="Main")
    assert selected["status"] == "pass"
    assert _packet_facts(selected["packet"], selected["packet"]["records"][0])["address"] == "H2"
    assert exploration.query(out, "ScopedOrGlobal")["needs_selection"] is True
    local_wins = exploration.query(out, "ScopedOrGlobal", sheet="Main")
    assert _packet_facts(local_wins["packet"], local_wins["packet"]["records"][0])["address"] == "H2"
    assert exploration.query(out, "Ignored!A1")["status"] == "blocked"
    assert exploration.query(out, "LocalRate", sheet="Ignored")["status"] == "blocked"


def test_spec_packet_replay_preserves_cross_sheet_local_name_context(tmp_path: Path, monkeypatch) -> None:
    inventory = _inventory()
    inventory["workbook_ranges"].append({
        "source_location": _location("Control", "Main!C3", "defined_name", "LocalInput"),
        "source_identity": None, "name": "LocalInput", "address": "Main!C3", "kind": "defined_name",
        "metadata": {"scope": "worksheet"},
    })
    inventory["sheets"].append({
        "source_location": _location("Control", None, "sheet"), "name": "Control", "index": 2,
        "max_row": 1, "max_column": 1, "state": "visible", "cells": [], "ranges": [], "layout_objects": [],
    })
    inventory["coverage"]["recognized_inventory_objects"] += 2
    inventory["coverage"]["discovered_workbook_objects"] += 2
    out = _ready_analysis(tmp_path, monkeypatch, inventory=inventory)
    packet_dir = tmp_path / "local-name-packet"
    packet = exploration.query(out, "LocalInput", sheet="Control", out_dir=packet_dir)["packet"]
    assert packet["selector"]["sheet"] == "Main"
    assert packet["selector"]["sheet_context"] == "Control"
    spec_path = tmp_path / "local-name-spec.input.json"
    _write_spec(spec_path, packet_dir / "query.json", packet)
    result = workflow.validate_analysis(out, spec_path=spec_path)
    assert result["status"] == "pass", result.get("model_spec")


def test_trace_direction_multistart_frontier_and_boundaries(tmp_path: Path, monkeypatch) -> None:
    inventory = _inventory()
    inventory["sheets"][0]["cells"].append(_cell("Main", "L1", 1, 12, "formula", formula="=C3+1"))
    scope = {"schema_version": "analysis.scope.v1", "workbook_sha256": "source-sha", "ignored_sheets": ["Ignored"]}
    out = _ready_analysis(tmp_path, monkeypatch, scope=scope, inventory=inventory)
    fields = json.loads((out / "fields.json").read_text(encoding="utf-8"))
    by_id, by_cell = workflow._field_index(fields)
    gp = by_cell[("main", "C3")]
    g1 = by_cell[("main", "G1")]
    l1 = by_cell[("main", "L1")]

    upstream = exploration.trace(out, "Main!G1", direction="upstream")["packet"]
    assert any(_packet_facts(upstream, item)["consumer"] == g1 and _packet_facts(upstream, item)["prerequisite"] == gp for item in upstream["edges"])
    downstream_result = exploration.trace(out, "GP", direction="downstream")
    assert downstream_result["status"] == "pass", downstream_result
    downstream = downstream_result["packet"]
    assert any(_packet_facts(downstream, item)["prerequisite"] == gp and _packet_facts(downstream, item)["consumer"] == l1 for item in downstream["edges"])
    assert downstream["direction"] == "downstream"

    capped_depth = exploration.trace(out, "Main!G1", max_depth=0)["packet"]
    assert capped_depth["truncated"] is True and "max_depth" in capped_depth["stopping_reasons"]
    assert any(item["reason"] == "max_depth" for item in capped_depth["frontier"])
    capped_starts = exploration.trace(out, "Main!C3:C4", max_fields=1)["packet"]
    assert capped_starts["truncated"] is True and "max_fields" in capped_starts["stopping_reasons"]
    assert any(item["reason"] == "max_fields_start_limit" for item in capped_starts["frontier"])

    monkeypatch.setattr(exploration, "_TRACE_EDGE_MAX", 2)
    capped_edges = exploration.trace(out, "GP", direction="downstream")["packet"]
    assert len(capped_edges["edges"]) == 2
    assert capped_edges["truncated"] is True and "edge_limit" in capped_edges["stopping_reasons"]
    assert any(item["reason"] == "edge_record_limit" for item in capped_edges["frontier"])

    excluded = exploration.trace(out, "Main!E1")["packet"]
    assert any(_packet_facts(excluded, item).get("boundary_type") == "excluded_sheet" for item in excluded["boundaries"])
    dynamic = exploration.trace(out, "Main!D3")["packet"]
    assert any(item["kind"] == "unresolved" for item in dynamic["boundaries"])


def test_trace_keeps_scalar_self_reference_and_cycles(tmp_path: Path, monkeypatch) -> None:
    inventory = _inventory()
    inventory["sheets"][0]["cells"] = [
        {**cell, "kind": "formula", "formula": "=SUM(C3:C4)", "data_type": "f", "value": None}
        if cell["address"] == "C3" else cell
        for cell in inventory["sheets"][0]["cells"]
    ]
    out = _ready_analysis(tmp_path, monkeypatch, inventory=inventory)
    trace_result = exploration.trace(out, "GP")
    assert trace_result["status"] == "pass", trace_result
    packet = trace_result["packet"]
    assert any(item["kind"] == "internal" and _packet_facts(packet, item)["same_cell"] is True for item in packet["edges"])
    assert packet["cycles"]
    nlp = next(item["summary"]["field_id"] for item in packet["fields"] if item["summary"].get("names") == ["NLP"])
    assert any(_packet_facts(packet, item).get("prerequisite") == nlp
               for item in packet["edges"] if item["kind"] == "edge")


def _write_spec(path: Path, packet_path: Path, packet: dict, *, field_fact: dict | None = None) -> dict:
    cell_record = packet["records"][0]
    field_record = packet["fields"][0]
    cell_facts = _packet_facts(packet, cell_record)
    field_facts = _packet_facts(packet, field_record)
    value = {
        "schema_version": "step3.model_spec.input.v1",
        "source": {key: packet["source"].get(key) for key in ("source_id", "run_id", "source_sha256")},
        "binding_sha256": packet["binding_sha256"],
        "stage_hashes": packet["stage_hashes"],
        "evidence_packets": [{"path": str(packet_path.relative_to(path.parent)), "sha256": workflow._hash_file(packet_path)}],
        "field_roles": [{"status": "fact", "claim": "C3 is a stored source record.", "evidence": [
            {"id": cell_record["evidence_id"], "facts": cell_facts}]}],
        "logical_axes": [{"status": "inferred", "statement": "The field may represent an input scalar.",
                          "evidence_refs": [field_record["evidence_id"]],
                          "details": {"field_ids": [field_facts["field_id"]], "logical_shape": [1], "axes": []}}],
        "calculation_groups": [], "outputs": [], "macro_workflow": [],
        "open_questions": [{"status": "pending", "question": "Confirm units.", "missing_evidence": ["unit definition"], "details": {"unit": None}}],
        "validation_plan": [],
    }
    if field_fact is not None:
        value["field_roles"][0]["evidence"].append(field_fact)
    _write_json(path, value)
    return value


def test_validate_semantic_spec_exact_facts_and_preserves_invalid_outputs(tmp_path: Path, monkeypatch) -> None:
    out = _ready_analysis(tmp_path, monkeypatch)
    packet_dir = tmp_path / "packets"
    query_result = exploration.query(out, "Main!C3", out_dir=packet_dir)
    packet = query_result["packet"]
    packet_path = packet_dir / "query.json"
    spec_path = tmp_path / "model_spec.input.json"
    _write_spec(spec_path, packet_path, packet)

    result = workflow.validate_analysis(out, spec_path=spec_path)
    assert result["status"] == "pass", result.get("model_spec")
    spec_output = json.loads((out / "model_spec.json").read_text(encoding="utf-8"))
    assert spec_output["readiness"] == {"draft_only": True, "runtime_verified": False, "generation_ready": False, "gpu_executable": False}
    assert spec_output["logical_axes"][0]["details"]["field_ids"] == [packet["fields"][0]["summary"]["field_id"]]
    assert "logical_shape" in (out / "model_spec.md").read_text(encoding="utf-8")
    handoff = json.loads((out / "handoff.json").read_text(encoding="utf-8"))
    assert handoff["model_spec"]["json_sha256"] == workflow._hash_file(out / "model_spec.json")
    old_json, old_markdown = (out / "model_spec.json").read_bytes(), (out / "model_spec.md").read_bytes()

    bad_path = tmp_path / "bad.model_spec.input.json"
    changed = json.loads(spec_path.read_text(encoding="utf-8"))
    changed["field_roles"][0]["evidence"][0]["facts"]["value"] = 999
    _write_json(bad_path, changed)
    rejected_fact = workflow.validate_analysis(out, spec_path=bad_path)
    assert rejected_fact["status"] == "blocked"
    assert (out / "model_spec.json").read_bytes() == old_json
    assert (out / "model_spec.md").read_bytes() == old_markdown

    unknown = json.loads(spec_path.read_text(encoding="utf-8"))
    unknown["logical_axes"][0]["evidence_refs"] = ["field:unknown"]
    _write_json(bad_path, unknown)
    rejected_reference = workflow.validate_analysis(out, spec_path=bad_path)
    assert rejected_reference["status"] == "blocked"
    assert "not delivered" in rejected_reference["model_spec"]["diagnostics"][0]["message"]
    assert (out / "model_spec.json").read_bytes() == old_json

    stale = json.loads(spec_path.read_text(encoding="utf-8"))
    stale["binding_sha256"] = "0" * 64
    _write_json(bad_path, stale)
    rejected_stale = workflow.validate_analysis(out, spec_path=bad_path)
    assert rejected_stale["status"] == "blocked"
    assert (out / "model_spec.json").read_bytes() == old_json


def test_validate_rejects_rehashed_packet_facts_and_spec_field_ids(tmp_path: Path, monkeypatch) -> None:
    out = _ready_analysis(tmp_path, monkeypatch)
    packet_dir = tmp_path / "packets"
    packet = exploration.query(out, "Main!C3", out_dir=packet_dir)["packet"]
    packet_path = packet_dir / "query.json"
    spec_path = tmp_path / "model_spec.input.json"
    spec = _write_spec(spec_path, packet_path, packet)
    first_validation = workflow.validate_analysis(out, spec_path=spec_path)
    assert first_validation["status"] == "pass", first_validation
    old_spec = (out / "model_spec.json").read_bytes()

    tampered_packet = json.loads(packet_path.read_text(encoding="utf-8"))
    source_id = next(evidence_id for evidence_id, value in tampered_packet["evidence"].items() if value["kind"] == "source_cell")
    tampered_packet["evidence"][source_id]["facts"]["value"] = 999
    packet_path.write_bytes((json.dumps(tampered_packet, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8"))
    changed_spec = json.loads(spec_path.read_text(encoding="utf-8"))
    changed_spec["evidence_packets"][0]["sha256"] = workflow._hash_file(packet_path)
    _write_json(spec_path, changed_spec)
    rejected = workflow.validate_analysis(out, spec_path=spec_path)
    assert rejected["status"] == "blocked"
    assert "canonical artifacts" in rejected["model_spec"]["diagnostics"][0]["message"]
    assert (out / "model_spec.json").read_bytes() == old_spec

    original_packet = exploration.query(out, "Main!C3")["packet"]
    packet_path.write_bytes(workflow._json_bytes(original_packet))
    tampered_summary = json.loads(packet_path.read_text(encoding="utf-8"))
    tampered_summary["records"][0]["summary"]["value"] = 999
    packet_path.write_bytes(workflow._json_bytes(tampered_summary))
    changed_spec = json.loads(spec_path.read_text(encoding="utf-8"))
    changed_spec["evidence_packets"][0]["sha256"] = workflow._hash_file(packet_path)
    _write_json(spec_path, changed_spec)
    rejected_summary = workflow.validate_analysis(out, spec_path=spec_path)
    assert rejected_summary["status"] == "blocked"
    assert "differs from its checked evidence" in rejected_summary["model_spec"]["diagnostics"][0]["message"]
    assert (out / "model_spec.json").read_bytes() == old_spec

    # Restore the packet, then ensure a current packet hash cannot bless an undelivered field ID.
    original_packet = exploration.query(out, "Main!C3")["packet"]
    packet_path.write_bytes(workflow._json_bytes(original_packet))
    spec["evidence_packets"][0]["sha256"] = workflow._hash_file(packet_path)
    spec["logical_axes"][0]["details"]["field_ids"] = ["field:not-delivered"]
    _write_json(spec_path, spec)
    rejected_fields = workflow.validate_analysis(out, spec_path=spec_path)
    assert rejected_fields["status"] == "blocked"
    assert "details.field_ids" in rejected_fields["model_spec"]["diagnostics"][0]["message"]
    assert (out / "model_spec.json").read_bytes() == old_spec


def test_validate_rejects_rehashed_query_selector_and_page_tampering(tmp_path: Path, monkeypatch) -> None:
    out = _ready_analysis(tmp_path, monkeypatch)
    packet_dir = tmp_path / "query-packet"
    packet_path = packet_dir / "query.json"
    packet = exploration.query(out, "Main!C3", out_dir=packet_dir)["packet"]
    spec_path = tmp_path / "query-spec.input.json"
    _write_spec(spec_path, packet_path, packet)
    assert workflow.validate_analysis(out, spec_path=spec_path)["status"] == "pass"
    old_spec = (out / "model_spec.json").read_bytes()
    original_packet = packet_path.read_bytes()

    changes = [
        ("invalid selector", lambda value: value["selector"].update(target="Main!NOT_THE_SELECTED_CELL")),
        ("different valid selector", lambda value: value["selector"].update(target="Main!C4")),
        ("different page", lambda value: value["page"].update(offset=1)),
    ]
    for _label, mutate in changes:
        changed_packet = json.loads(original_packet)
        mutate(changed_packet)
        _write_json(packet_path, changed_packet)
        changed_spec = json.loads(spec_path.read_text(encoding="utf-8"))
        changed_spec["evidence_packets"][0]["sha256"] = workflow._hash_file(packet_path)
        _write_json(spec_path, changed_spec)
        rejected = workflow.validate_analysis(out, spec_path=spec_path)
        assert rejected["status"] == "blocked", (_label, rejected)
        assert (out / "model_spec.json").read_bytes() == old_spec


def test_validate_rejects_rehashed_trace_selector_and_traversal_tampering(tmp_path: Path, monkeypatch) -> None:
    out = _ready_analysis(tmp_path, monkeypatch)
    query_dir, trace_dir = tmp_path / "query-packet", tmp_path / "trace-packet"
    query_packet = exploration.query(out, "Main!C3", out_dir=query_dir)["packet"]
    fields = workflow._read_json(out / "fields.json")["fields"]
    field_by_address = {item["address_extent"]: item["field_id"] for item in fields}
    exploration.trace(out, field_by_address["D2"], direction="upstream", out_dir=trace_dir)
    query_path, trace_path = query_dir / "query.json", trace_dir / "trace.json"
    spec_path = tmp_path / "trace-spec.input.json"
    spec = _write_spec(spec_path, query_path, query_packet)
    spec["evidence_packets"].append({"path": "trace-packet/trace.json", "sha256": workflow._hash_file(trace_path)})
    _write_json(spec_path, spec)
    assert workflow.validate_analysis(out, spec_path=spec_path)["status"] == "pass"
    old_spec = (out / "model_spec.json").read_bytes()
    original_packet = trace_path.read_bytes()

    changes = [
        ("different valid trace start", lambda value: value["selector"].update(target=field_by_address["C3"])),
        ("different direction", lambda value: value.update(direction="downstream")),
        ("different depth limit", lambda value: value["limits"].update(max_depth=0)),
    ]
    for _label, mutate in changes:
        changed_packet = json.loads(original_packet)
        mutate(changed_packet)
        _write_json(trace_path, changed_packet)
        changed_spec = json.loads(spec_path.read_text(encoding="utf-8"))
        changed_spec["evidence_packets"][1]["sha256"] = workflow._hash_file(trace_path)
        _write_json(spec_path, changed_spec)
        rejected = workflow.validate_analysis(out, spec_path=spec_path)
        assert rejected["status"] == "blocked", (_label, rejected)
        assert (out / "model_spec.json").read_bytes() == old_spec


def test_validate_accepts_canonical_structural_edge_evidence(tmp_path: Path, monkeypatch) -> None:
    out = _ready_analysis(tmp_path, monkeypatch)
    query_dir, trace_dir = tmp_path / "packets" / "query", tmp_path / "packets" / "trace"
    query_packet = exploration.query(out, "Main!C3", out_dir=query_dir)["packet"]
    trace_packet = exploration.trace(out, "Main!D2", direction="upstream", out_dir=trace_dir)["packet"]
    edge_id = next(evidence_id for evidence_id, entry in trace_packet["evidence"].items() if entry["kind"] == "edge")
    path = tmp_path / "structural-spec.input.json"
    spec = _write_spec(path, query_dir / "query.json", query_packet)
    spec["evidence_packets"].append({"path": "packets/trace/trace.json", "sha256": workflow._hash_file(trace_dir / "trace.json")})
    spec["calculation_groups"] = [{"status": "fact", "claim": "D2 has this checked prerequisite edge.",
                                   "evidence": [{"id": edge_id, "facts": trace_packet["evidence"][edge_id]["facts"]}]}]
    _write_json(path, spec)
    result = workflow.validate_analysis(out, spec_path=path)
    assert result["status"] == "pass", json.dumps(result.get("model_spec", result), indent=2)
    saved = json.loads((out / "model_spec.json").read_text(encoding="utf-8"))
    assert saved["calculation_groups"][0]["evidence"][0]["id"] == edge_id


def test_exploration_cli_writes_packets_and_validate_accepts_spec(tmp_path: Path, monkeypatch) -> None:
    out = _ready_analysis(tmp_path, monkeypatch)
    query_out = tmp_path / "cli-packets"
    query_result = CliRunner().invoke(app, ["step3", "query", "--analysis", str(out), "--target", "Main!C3", "--out", str(query_out)])
    assert query_result.exit_code == 0, query_result.stdout
    query_json = json.loads(query_result.stdout)
    assert query_json["status"] == "pass"
    assert (query_out / "query.json").is_file() and (query_out / "query.md").is_file()

    trace_result = CliRunner().invoke(app, ["step3", "trace", "--analysis", str(out), "--target", "GP", "--direction", "downstream", "--max-depth", "2"])
    assert trace_result.exit_code == 0, trace_result.stdout
    assert json.loads(trace_result.stdout)["packet"]["direction"] == "downstream"

    packet_path = query_out / "query.json"
    packet = json.loads(packet_path.read_text(encoding="utf-8"))
    spec_path = tmp_path / "cli-spec.json"
    _write_spec(spec_path, packet_path, packet)
    validate_result = CliRunner().invoke(app, ["step3", "check", "--analysis", str(out), "--spec", str(spec_path)])
    assert validate_result.exit_code == 0, validate_result.stdout
    assert json.loads(validate_result.stdout)["model_spec"]["status"] == "pass"

    bad_query = CliRunner().invoke(app, ["step3", "query", "--analysis", str(out), "--target", "not-a-name"])
    assert bad_query.exit_code == 1
    assert json.loads(bad_query.stdout)["status"] == "blocked"
    bad_limit = CliRunner().invoke(app, ["step3", "query", "--analysis", str(out), "--target", "Main!C3", "--limit", "501"])
    assert bad_limit.exit_code == 1
    assert json.loads(bad_limit.stdout)["status"] == "blocked"


def test_step3_agent_prints_packaged_result_driven_contract() -> None:
    agent_result = CliRunner().invoke(app, ["step3", "agent"])
    assert agent_result.exit_code == 0
    for required in ("Result-Driven Analysis", "scenario", "known-only", "frontier", "not implemented", "Documentation/handoff completeness", "Stage 4", "feedback solver"):
        assert required in agent_result.stdout

    tools_result = CliRunner().invoke(app, ["step3", "tools"])
    assert tools_result.exit_code == 0
    catalog = json.loads(tools_result.stdout)
    assert any("input-boundary subgate always requires both Agent and actual human approval" in item
               for item in catalog["review_workflow"])
    agent_tool = next(item for item in catalog["tools"] if item["name"] == "step3.agent")
    assert agent_tool["command"] == "step3 agent"
    assert agent_tool["outputs"]
    profile_tool = next(item for item in catalog["tools"] if item["name"] == "step3.profile")
    assert profile_tool["command"] == (
        "step3 profile --analysis DIR (--trace ACTIVE_TRACE.json | --source-trace SOURCE_CANDIDATE_TRACE.json) --out DIR"
    )
    source_trace_tool = next(item for item in catalog["tools"] if item["name"] == "step3.source-trace")
    assert "--model-scope INPUT.json" in source_trace_tool["command"]
    assert any("clipped ranges are reported as explicit frontiers" in item for item in catalog["limits"])
    assert {item["name"] for item in catalog["tools"]}.isdisjoint({"step3.gp", "step3.oracle"})
    from importlib.resources import files
    assert files("excel_to_act.steps.step5").joinpath("native_excel_oracle.ps1").is_file()


def test_step3_profile_is_reusable_source_bound_and_preserves_prior_outputs_on_bad_trace(tmp_path, monkeypatch) -> None:
    inventory_input = _inventory()
    for sheet in inventory_input["sheets"]:
        for cell in sheet["cells"]:
            cell["address"] = f"{sheet['name']}!{cell['address']}"
    analysis_dir, _root, prepared = _prepared_fixture(tmp_path, monkeypatch, inventory=inventory_input)
    assert prepared["status"] == "pass"
    inventory = workflow._read_json(analysis_dir / ".." / "step1" / "inventory.json")
    formula_cells = [{"sheet": sheet["name"], "address": cell["address"], "formula": cell["formula"],
                      "role": "calculated_formula"}
                     for sheet in inventory["sheets"] for cell in sheet["cells"]
                     if cell.get("kind") == "formula"]
    trace_path = tmp_path / "active_trace.json"
    _write_json(trace_path, {"schema_version": "step4.discovery_trace.v1", "source_sha256": "source-sha",
                             "native_excel_called": False, "formula_cache_inputs": False,
                             "cells": formula_cells})
    out = tmp_path / "profile"
    runner = CliRunner()
    command = ["step3", "profile", "--analysis", str(analysis_dir), "--trace", str(trace_path), "--out", str(out)]
    first = runner.invoke(app, command)
    assert first.exit_code == 0, first.output
    first_result = json.loads(first.output)
    profile_path = out / "source_family_profile.json"
    markdown_path = out / "source_family_profile.md"
    before_json, before_md = profile_path.read_bytes(), markdown_path.read_bytes()
    profile = json.loads(before_json)
    assert profile["summary"]["ordinary_formula_members"] == len(formula_cells)
    assert profile["summary"]["syntactic_families"] < len(formula_cells)
    assert profile["policy"]["formula_cache_used"] is False
    assert profile["families"]
    profiled_members = [address for family in profile["families"] for address in family["source_members"]]
    assert len(profiled_members) == len(set(profiled_members)) == len(formula_cells)
    assert set(profiled_members) == {cell["address"] for cell in formula_cells}
    second = runner.invoke(app, command)
    assert second.exit_code == 0, second.output
    assert json.loads(second.output)["profile_sha256"] == first_result["profile_sha256"]

    changed_trace = json.loads(trace_path.read_text(encoding="utf-8"))
    changed_trace["cells"][0]["formula"] += "+1"
    _write_json(trace_path, changed_trace)
    rejected = runner.invoke(app, command)
    assert rejected.exit_code == 1
    assert "differs from the current canonical inventory" in json.loads(rejected.output)["diagnostics"][0]["message"]
    assert profile_path.read_bytes() == before_json
    assert markdown_path.read_bytes() == before_md


def test_query_rejects_changed_source_and_output_over_bound_inputs(tmp_path: Path, monkeypatch) -> None:
    out, root, prepared = _prepared_fixture(tmp_path, monkeypatch)
    assert prepared["status"] == "pass"
    assert workflow.build_fields(out)["status"] == "pass"
    assert workflow.build_dependencies(out)["status"] == "pass"
    blocked_output = exploration.query(out, "Main!C3", out_dir=root)
    assert blocked_output["status"] == "blocked"
    assert "overlaps a bound input" in blocked_output["diagnostics"][0]["message"]
    inventory_path = root / "inventory.json"
    inventory_path.write_bytes(inventory_path.read_bytes() + b" ")
    stale = exploration.query(out, "Main!C3")
    assert stale["status"] == "blocked"
    assert "inventory.json changed" in stale["diagnostics"][0]["message"]


def test_cli_reports_stage_order_errors_as_json_and_nonzero(tmp_path: Path) -> None:
    result = CliRunner().invoke(app, ["step3", "fields", "--analysis", str(tmp_path / "missing")])
    assert result.exit_code == 1
    assert json.loads(result.stdout)["status"] == "blocked"


def test_design_status_message_matches_the_active_reviewer_pair() -> None:
    assert "delegated TypeSafe" in design._review_policy_sentence({"reviewers": ["agent", "typesafe"]})
    assert "Agent and human" in design._review_policy_sentence({"reviewers": ["agent", "human"]})


def test_design_scenario_source_literals_are_checked_against_formula_mode_cells(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["O26"] = "Yes"
    sheet["P26"] = "Yes"
    sheet["A1"] = "=1+1"
    workbook.save(source)
    source_sha = hashlib.sha256(source.read_bytes()).hexdigest()

    def make_design(selector: str, source_cell: str, literal: object) -> dict:
        return {
            "targets": [{"target_id": "gp", "selector": "Premium!J1", "result_order": 0, "shape": []}],
            "scenarios": [{"scenario_id": "saved", "primary_inputs": {
                selector: {"source_cell": source_cell, "source_literal": literal},
            }}],
            "paths": [], "coverage": {"known_path_ids": []},
        }

    wrong = make_design("CIDecrement", "Main!O26", "No")
    design._validate_design_content(wrong)
    with pytest.raises(ValueError, match="scenario saved selector CIDecrement source literal mismatch at Main!O26: expected 'No', actual 'Yes'"):
        design._check_scenario_source_literals(wrong, source, source_sha)

    correct = make_design("CIDecrement", "Main!O26", "Yes")
    design._validate_design_content(correct)
    checked = design._check_scenario_source_literals(correct, source, source_sha)
    assert checked["status"] == "pass"
    assert checked["assertion_count"] == 1
    assert checked["formula_cache_used"] is False
    report = design._design_markdown({
        "source": {"run_id": "run", "source_id": "source", "workbook_sha256": source_sha},
        "binding_sha256": "binding", "design": correct,
        "scenario_source_validation": checked,
    })
    assert "## Saved scenario source-literal checks" in report
    assert "`Main!O26` = `\"Yes\"`" in report

    formula_literal = make_design("Calculated", "Main!A1", "=1+1")
    design._validate_design_content(formula_literal)
    with pytest.raises(ValueError, match="source_cell Main!A1 is a formula, not a source literal"):
        design._check_scenario_source_literals(formula_literal, source, source_sha)

    partial = make_design("Partial", "Main!O26", "Yes")
    del partial["scenarios"][0]["primary_inputs"]["Partial"]["source_literal"]
    with pytest.raises(ValueError, match="must declare both source_cell and source_literal"):
        design._validate_design_content(partial)


def test_stage3_design_report_renders_bound_source_kind_and_raw_extent_summary(tmp_path: Path) -> None:
    classification = {
        "source_kind_basis": "Promoted source formulas and array ranges; no caches.",
        "variables_by_role_and_kind": {"derived": {"projection_series": 2}, "raw": {"matrix": 1}},
        "demanded_variables_by_role": {"raw": 1, "derived": 2},
        "raw_extent_count": 3,
        "distinct_raw_extent_coordinate_count": 24,
        "raw_seed_variable_count": 1,
        "raw_seed_extent_count": 1,
        "demanded_raw_formula_overlaps": 0,
        "demanded_raw_array_follower_overlaps": 0,
        "formula_coordinates_outside_raw_demand": 1,
        "array_follower_coordinates_outside_raw_demand": 0,
    }
    plan = {"coverage": {"active_formula_members_mapped": 3, "active_formula_members": 3,
                          "ordinary_formula_members_mapped": 2, "ordinary_formula_members": 2,
                          "array_formula_followers_mapped": 1, "array_formula_followers": 1,
                          "semantic_family_count": 1, "array_family_count": 1},
            "source_classification": classification}
    check = {"source_classification": classification}
    summary = design._static_semantic_summary(plan, check)
    plan_path = tmp_path / "semantic_plan.json"
    plan_path.write_text(json.dumps(plan), encoding="utf-8")
    design_content = {
        "targets": [{"target_id": "T", "selector": "GP", "result_order": 0, "shape": []}],
        "scenarios": [{"scenario_id": "saved", "primary_inputs": {}}],
        "paths": [], "coverage": {"known_path_ids": []},
        "tool_execution_ledger": [{
            "tool": "step3 source-trace", "command": "excel-to-act step3 source-trace --model-scope INPUT.json",
            "purpose": "Build a source-only candidate trace.",
            "outputs": ["source_candidate_trace.json", "source_candidate_trace.md"],
            "status": "completed",
        }],
        "external_value_capture_plan": {
            "phase": "Step 4 after separate Step 3 approval",
            "ranges": ["Qtable!W5:W110", "Qtable!Y5:Y110"],
            "method": "Capture from a private recalculated copy.",
            "binding": "Bind source, scenario, catalog and range hashes.",
            "validation": "Require complete finite vectors and matching age keys.",
            "current_values_available": False,
            "formula_cache_allowed": False,
            "permitted_origins": ["approved_excel_capture"],
            "upstream_python_translation": False,
        },
    }
    design._validate_design_content(design_content)
    bad_ledger = {**design_content, "tool_execution_ledger": [
        {**design_content["tool_execution_ledger"][0], "outputs": [123]},
    ]}
    with pytest.raises(ValueError, match="outputs must be a non-empty list of text"):
        design._validate_design_content(bad_ledger)
    bad_capture = {**design_content, "external_value_capture_plan": {
        **design_content["external_value_capture_plan"], "current_values_available": True,
    }}
    with pytest.raises(ValueError, match="must remain unavailable and formula caches prohibited"):
        design._validate_design_content(bad_capture)
    payload = {
        "source": {"run_id": "run", "source_id": "source", "workbook_sha256": "sha"},
        "binding_sha256": "binding",
        "design": design_content,
        "analysis": {"artifacts": {
            "semantic_plan": {"path": str(plan_path), "sha256": "plan-sha"},
            "semantic_check": {"path": "semantic_check.json", "sha256": "check-sha"},
        }},
        "static_semantic": summary,
    }
    report = design._design_markdown(payload)
    assert "| raw | matrix | 1 |" in report
    assert "Raw source extent checks:" in report
    assert "Demanded formula overlaps: 0" in report
    assert "Demanded array-follower overlaps: 0" in report
    assert "outside raw demand: 1 / 0" in report
    assert "## Tool execution ledger" in report
    assert "excel-to-act step3 source-trace --model-scope INPUT.json" in report
    assert "source_candidate_trace.json<br>source_candidate_trace.md" in report
    assert "rendering this report does not execute" in report
    assert "## External-value capture prerequisite" in report
    assert "Step 4 after separate Step 3 approval" in report
    assert "Qtable!W5:W110" in report and "Qtable!Y5:Y110" in report
    assert "Current values available: `false`" in report
    assert "Formula cache allowed: `false`" in report
    assert "Stage 3 does not capture or include these values" in report

    changed_check = {"source_classification": {**classification, "raw_extent_count": 4}}
    with pytest.raises(ValueError, match="inconsistent source-kind summaries"):
        design._static_semantic_summary(plan, changed_check)


def test_step3_design_reader_front_shows_field_groups_order_and_question_dispositions() -> None:
    source_hash = "a" * 64
    payload = {
        "source": {"run_id": "run-1", "source_id": "pricing", "workbook_sha256": source_hash},
        "binding_sha256": "binding-hash",
        "scenario_source_validation": {},
        "design": {
            "targets": [{"target_id": "T_GP", "selector": "GP", "result_order": 0,
                         "shape": [], "units": "currency basis not extended.",
                         "axes": [{"name": "Policy duration", "axis_id": "policy-duration",
                                   "role": "observation", "units": None,
                                   "keys": [0, False, "None", 3]}]},
                        {"target_id": "T_MATRIX", "selector": "Matrix", "result_order": 1,
                         "shape": [4, 2], "result_kind": "table", "units": "USD",
                         "axes": [{"name": "Policy duration", "axis_id": "policy-duration",
                                   "role": "time", "keys": [0, False, "None", 3]},
                                  {"name": "Benefit category", "axis_id": "benefit-category",
                                   "role": "category", "keys": ["None", "D", "C"]}]}],
            "scenarios": [],
            "coverage": {"known_path_ids": ["P1"], "principal_numerical_path_ids": ["P1"],
                         "principal_design_coverage": {"numerator": 1, "denominator": 1}},
            "scope": {"boundary": "Saved GP scenario."},
            "paths": [{"path_id": "P1"}],
            "field_groups": [{"id": "projection_state", "shape": "year vectors",
                              "meaning": "Ordered recurrence and snapshots from Premium H:Q.",
                              "semantic_grouping_status": "Proposed; runtime unverified."}],
            "shared_modules": [
                {"module_id": "state_projection", "order": 3, "purpose": "Run ordered annual updates.",
                 "implementation_status": "Planned.", "verification_status": "Not runtime verified."},
                {"module_id": "input_adapter", "order": 0, "purpose": "Load accepted inputs.",
                 "implementation_status": "Planned.", "verification_status": "Not runtime verified."},
            ],
            "execution_design": {"ordering": "Inputs -> state recurrence -> reductions.",
                                 "preserve_formulas": ["Prior H leads the I recurrence; T9 is year zero."]},
            "open_questions": [{"question_id": "Q_AXIS", "question": "Which axis is used?",
                                "disposition": "Step 4 discovery must resolve or block."}],
        },
    }

    design_payload = payload["design"]
    design_payload["targets"][0].update({"selector": "None", "result_kind": None, "units": None, "shape": []})
    design_payload["scope"]["boundary"] = None
    design_payload["execution_design"]["ordering"] = False
    design_payload["open_questions"].append({"question_id": "Q_UNKNOWN", "question": "What is the unit?",
                                              "answer": None, "status": None, "disposition": None})
    before = json.dumps(payload, sort_keys=True)
    front = design._design_markdown(payload).split("## Detailed design evidence", 1)[0]

    assert "**None** — Not recorded; shape []; units Not recorded; axes " in front
    assert 'axes name: Policy duration; axis_id: policy-duration; role: observation; units: Not recorded; keys: [0, False, "None", 3]' in front
    matrix_target_line = next(line for line in front.splitlines() if line.startswith("- **Matrix**"))
    assert matrix_target_line.index("name: Policy duration") < matrix_target_line.index("name: Benefit category")
    assert "shape [4, 2]; units USD" in matrix_target_line
    assert 'keys: [0, False, "None", 3] / name: Benefit category' in matrix_target_line
    assert "Boundary: Not recorded" in front
    assert "Recorded calculation order: False." in front
    assert "**Q_UNKNOWN** — What is the unit? Recorded answer: Not recorded; recorded disposition/status: Not recorded." in front
    assert json.dumps(payload, sort_keys=True) == before
    design_payload["open_questions"] = None
    null_questions = design._design_markdown(payload).split("## Detailed design evidence", 1)[0]
    assert "Open questions recorded: Not recorded." in null_questions
    assert source_hash not in front and "Workbook SHA-256" not in front
    assert "Ordered recurrence and snapshots from Premium H:Q." in front
    assert front.index("| 0 | input_adapter") < front.index("| 3 | state_projection")
    assert "Prior H leads the I recurrence; T9 is year zero." in front
    assert "**Q_AXIS** — Which axis is used? Recorded answer: Not recorded; recorded disposition/status: Step 4 discovery must resolve or block." in front
    assert "{'question_id'" not in front

    evidence_payload = {**payload, "analysis": {"artifacts": {}}, "review_policy": {}}
    evidence = design._design_evidence_markdown(evidence_payload)
    target_evidence = next(line for line in evidence.splitlines() if line.startswith("- **T_GP**"))
    assert 'axes name: Policy duration; axis_id: policy-duration; role: observation; units: Not recorded; keys: [0, False, "None", 3]' in target_evidence
    matrix_evidence = next(line for line in evidence.splitlines() if line.startswith("- **T_MATRIX**"))
    assert matrix_evidence.index("name: Policy duration") < matrix_evidence.index("name: Benefit category")


def test_stage3_binds_nested_semantic_evidence_and_stales_when_trace_changes(tmp_path: Path) -> None:
    evidence_dir = tmp_path / "semantic"
    paths = {name: evidence_dir / f"{name}.json"
             for name in ("map", "profile", "active_trace")}
    hashes = {name: _write_json(path, {"kind": name}) for name, path in paths.items()}
    semantic_plan = {
        "semantic_map_input": {"path": str(paths["map"]), "sha256": hashes["map"]},
        "evidence": {
            "source_family_profile": {"path": str(paths["profile"]), "sha256": hashes["profile"]},
            "active_trace": {"path": str(paths["active_trace"]), "sha256": hashes["active_trace"]},
        },
    }
    semantic_check = {
        "semantic_map_input_sha256": hashes["map"],
        "source_family_profile_sha256": hashes["profile"],
        "active_trace_sha256": hashes["active_trace"],
    }
    bound_inputs = design._semantic_evidence_inputs(evidence_dir, semantic_plan, semantic_check)
    assert set(bound_inputs) == {
        "step3_semantic_map_input", "step3_source_family_profile", "step3_active_trace",
    }

    workflow_root, _ = conversion_workflow.create_workflow(
        tmp_path / "workflow", {"workbook_sha256": "source-sha", "run_id": "run-1"})
    revision = conversion_workflow.append_stage_artifact(
        workflow_root, 3, "analysis_design.json", {"status": "ready_for_review"}, "Stage 3",
        input_files=bound_inputs)
    assert conversion_workflow._revision_staleness(
        workflow_root, conversion_workflow.load_workflow(workflow_root)[1], 3, revision) == []

    paths["active_trace"].write_text('{"kind":"changed"}', encoding="utf-8")
    reasons = conversion_workflow._revision_staleness(
        workflow_root, conversion_workflow.load_workflow(workflow_root)[1], 3, revision)
    assert any("step3_active_trace" in reason for reason in reasons)

    candidate_path = evidence_dir / "source_candidate_trace.json"
    candidate_sha = _write_json(candidate_path, {"kind": "source_candidate_trace"})
    scope_path = evidence_dir / "model_scope.json"
    scope_sha = _write_json(scope_path, {"kind": "model_scope"})
    candidate_plan = {
        "selection_basis": "static_candidate",
        "candidate_scope": "scope_bounded_static_candidate_design",
        "semantic_map_input": {"path": str(paths["map"]), "sha256": hashes["map"]},
        "evidence": {
            "source_family_profile": {"path": str(paths["profile"]), "sha256": hashes["profile"]},
            "source_candidate_trace": {"path": str(candidate_path), "sha256": candidate_sha},
            "model_scope": {"path": str(scope_path), "sha256": scope_sha},
        },
    }
    candidate_check = {
        "candidate_scope": "scope_bounded_static_candidate_design",
        "semantic_map_input_sha256": hashes["map"],
        "source_family_profile_sha256": hashes["profile"],
        "source_candidate_trace_sha256": candidate_sha,
        "model_scope_sha256": scope_sha,
    }
    candidate_inputs = design._semantic_evidence_inputs(evidence_dir, candidate_plan, candidate_check)
    assert "step3_source_candidate_trace" in candidate_inputs
    assert candidate_inputs["step3_source_candidate_trace"] == candidate_path.resolve()
    assert candidate_inputs["step3_model_scope"] == scope_path.resolve()
    candidate_revision = conversion_workflow.append_stage_artifact(
        workflow_root, 3, "analysis_design.json", {"status": "candidate-ready"}, "Stage 3 candidate",
        input_files=candidate_inputs)
    candidate_path.write_text('{"kind":"changed_candidate_trace"}', encoding="utf-8")
    candidate_reasons = conversion_workflow._revision_staleness(
        workflow_root, conversion_workflow.load_workflow(workflow_root)[1], 3, candidate_revision)
    assert any("step3_source_candidate_trace" in reason for reason in candidate_reasons)
    scope_path.write_text('{"kind":"changed_model_scope"}', encoding="utf-8")
    scope_reasons = conversion_workflow._revision_staleness(
        workflow_root, conversion_workflow.load_workflow(workflow_root)[1], 3, candidate_revision)
    assert any("step3_model_scope" in reason for reason in scope_reasons)


def test_semantic_plan_maps_formula_families_and_array_followers_once(tmp_path: Path, monkeypatch) -> None:
    analysis_dir, _root, prepared = _prepared_fixture(tmp_path, monkeypatch)
    assert prepared["status"] == "pass"
    assert workflow.build_fields(analysis_dir)["status"] == "pass"
    assert workflow.build_dependencies(analysis_dir)["status"] == "pass"
    assert workflow.build_plan(analysis_dir)["status"] == "pass"
    analysis = json.loads((analysis_dir / "analysis.json").read_text(encoding="utf-8"))
    fields, fields_sha = workflow._read_stage(analysis_dir, analysis, "fields")
    _dependencies, dependencies_sha = workflow._read_stage(analysis_dir, analysis, "dependencies")
    binding = analysis["binding"]
    profile = {"schema_version": "gp.source_family_profile.v1", "status": "pass",
               "selection_basis": "active_trace",
               "source_sha256": binding["source"]["source_sha256"],
               "binding_sha256": analysis["binding_sha256"],
               "inventory": {"sha256": binding["artifacts"]["inventory_sha256"]},
               "families": [{"family_id": "source-family-1", "sheet": "Main", "column": "F",
                             "normalized_formula": "=C3+1", "representative_formula": "=C3+1",
                             "source_members": ["Main!F1"], "member_count": 1}],
               "summary": {"ordinary_formula_members": 1}}
    trace = {"schema_version": "step4.active_trace.v1",
             "source_sha256": binding["source"]["source_sha256"],
             "native_excel_called": False, "formula_cache_inputs": False,
             "cells": [{"address": "Main!F1", "formula": "=C3+1", "role": "calculated_formula"},
                       {"address": "Main!F2", "formula": "=C3+1", "role": "calculated_array_formula",
                        "array_ref": "F1:F2"}]}
    evidence_dir = tmp_path / "semantic-evidence"
    evidence_dir.mkdir()
    profile_path, trace_path = evidence_dir / "profile.json", evidence_dir / "trace.json"
    _write_json(profile_path, profile)
    _write_json(trace_path, trace)
    semantic_input = {"schema_version": "step3.semantic_map.input.v1",
                      "source": {"source_id": binding["source"]["source_id"],
                                 "run_id": binding["source"]["run_id"],
                                 "source_sha256": binding["source"]["source_sha256"],
                                 "binding_sha256": analysis["binding_sha256"],
                                 "fields_sha256": fields_sha,
                                 "dependencies_sha256": dependencies_sha},
                      "source_family_profile": {"path": str(profile_path), "sha256": workflow._hash_file(profile_path)},
                      "active_trace": {"path": str(trace_path), "sha256": workflow._hash_file(trace_path)},
                      "modules": [{"module_id": "rate_tables", "order": 1, "purpose": "Selected source formula family."}],
                      "variables": [{"variable_id": "rates.matrix", "role": "derived", "kind": "matrix",
                                     "shape": [2, 1], "axes": [{"name": "age"}, {"name": "category"}],
                                     "source_extents": [{"ref": "Main!F1:F2", "role": "formula_output",
                                                         "coordinate_mapping": {"kind": "row_column_ordinal", "axes": [
                                                             {"axis": "age", "dimension": 0, "source_coordinate": "row", "origin": 1, "index_origin": 0},
                                                             {"axis": "category", "dimension": 1, "source_coordinate": "column", "origin": "F", "index_origin": 0}]}}],
                                     "demanded_indices": {"age": "0..1", "category": "0"},
                                     "initial_condition": None,
                                     "equation_segments": [{"segment_id": "array_anchor", "source_family_ids": ["source-family-1"],
                                                            "equation_family_ids": ["eq-rate-anchor"],
                                                            "index_axis": "age", "execution": "vectorized", "index_start": 0,
                                                            "index_stop_exclusive": 2, "axis_intervals": [
                                                                {"axis": "age", "dimension": 0, "start": 0, "stop_exclusive": 2},
                                                                {"axis": "category", "dimension": 1, "start": 0, "stop_exclusive": 1}],
                                                            "source_row_start": 1,
                                                                "source_row_stop_exclusive": 2,
                                                                "snapshot_before_update": False}], "dependencies": [],
                                     "error_policy": "Preserve source errors.", "output_usage": "Rate table values."},
                                    {"variable_id": "input.gp", "role": "raw", "kind": "scalar", "shape": [],
                                     "axes": [], "source_extents": [{"ref": "Main!C3", "role": "raw_input",
                                                                        "coordinate_mapping": {"kind": "scalar", "axes": []}}], "demanded_indices": [],
                                     "initial_condition": None, "equation_segments": [], "dependencies": [],
                                     "error_policy": "Preserve raw source value.", "output_usage": "Scenario input."}],
                      "reference_resolution": {
                          "named_references": [{"name": "GP", "kind": "scalar", "source_ref": "Main!C3",
                                                "variable_id": "input.gp", "resolution": "raw_scalar"}],
                          "cell_ranges": [{"source_ref": "Main!F1:F2", "variable_id": "rates.matrix", "role": "formula_output",
                                           "coordinate_mapping": {"kind": "row_column_ordinal", "axes": [
                                               {"axis": "age", "dimension": 0, "source_coordinate": "row", "origin": 1, "index_origin": 0},
                                               {"axis": "category", "dimension": 1, "source_coordinate": "column", "origin": "F", "index_origin": 0}]}},
                          {"source_ref": "Main!C3", "variable_id": "input.gp", "role": "raw_input",
                           "coordinate_mapping": {"kind": "scalar", "axes": []}}]},
                      "execution_plan": {"passes": [
                          {"pass_id": "load_inputs", "order": 1, "module_id": "rate_tables",
                           "source_family_ids": [], "variable_ids": ["input.gp"],
                           "purpose": "Bind the scenario input."},
                          {"pass_id": "map_rate_table", "order": 2, "module_id": "rate_tables",
                           "source_family_ids": ["source-family-1"], "variable_ids": ["rates.matrix"],
                           "purpose": "Map the source table anchor."}],
                                         "constraints": ["Use explicit source coordinates."]},
                      "equation_families": [{"family_id": "eq-rate-anchor", "function_name": "rate_anchor",
                                             "module_id": "rate_tables", "description": "Source array anchor.",
                                             "source_family_ids": ["source-family-1"]}],
                      "source_mappings": [{"source_family_id": "source-family-1",
                                           "equation_family_id": "eq-rate-anchor", "variable_id": "rates.matrix",
                                           "module_id": "rate_tables", "equation_segment": "array_anchor",
                                           "index_mapping": {"kind": "matrix", "array_ref": "F1:F2",
                                                             "axes": [
                                                                 {"axis": "age", "dimension": 0,
                                                                  "source_coordinate": "row", "origin": 1,
                                                                  "index_origin": 0},
                                                                 {"axis": "category", "dimension": 1,
                                                                  "source_coordinate": "column", "origin": "F",
                                                                  "index_origin": 0}]}}],
                      "array_mappings": [{"array_family_id": "ci-row", "instances": [{
                          "anchor": "Main!F1", "array_ref": "F1:F2", "shape": [2, 1],
                          "anchor_source_family_id": "source-family-1",
                          "orientation": "source-order", "variable_id": "rates.matrix",
                          "module_id": "rate_tables", "axes": [{"name": "age"}, {"name": "category"}]}]}]}
    input_path = tmp_path / "semantic_map.json"
    _write_json(input_path, semantic_input)

    planned = semantic.build_semantic_plan(analysis_dir, input_path)
    assert planned["status"] == "pass", planned
    plan_path = analysis_dir / "semantic_plan.json"
    original_plan = plan_path.read_bytes()
    plan = json.loads(original_plan)
    assert plan["coverage"]["active_formula_members"] == 2
    assert plan["coverage"]["active_formula_members_mapped"] == 2
    assert plan["source_mappings"][0]["source_members"] == ["Main!F1"]
    assert plan["array_mappings"][0]["instances"][0]["source_members"] == ["Main!F2"]
    assert plan["equation_families"]["eq-rate-anchor"]["family_fingerprint"]
    assert plan["reference_resolution"]["named_references"][0]["variable_id"] == "input.gp"
    assert plan["variables"]["rates.matrix"]["equation_segments"][0]["source_family_ids"] == ["source-family-1"]
    assert plan["source_classification"]["demanded_raw_formula_overlaps"] == 0
    assert plan["source_classification"]["demanded_raw_array_follower_overlaps"] == 0
    assert plan["source_classification"]["raw_seed_variable_count"] == 0
    assert plan["source_classification"]["demanded_variables_by_role"]["raw"] == 1

    checked = semantic.validate_semantic_plan(analysis_dir, plan_path)
    assert checked["status"] == "pass", checked
    assert checked["source_classification"] == plan["source_classification"]
    check_path = analysis_dir / "semantic_check.json"
    original_check = check_path.read_bytes()

    runner = CliRunner()
    cli_plan = runner.invoke(app, ["step3", "plan", "--analysis", str(analysis_dir),
                                  "--semantic-map", str(input_path)])
    assert cli_plan.exit_code == 0, cli_plan.stdout
    assert json.loads(cli_plan.stdout)["semantic_mapping"]["status"] == "pass"
    cli_check = runner.invoke(app, ["step3", "check", "--analysis", str(analysis_dir),
                                   "--semantic-plan", str(plan_path)])
    assert cli_check.exit_code == 0, cli_check.stdout
    assert json.loads(cli_check.stdout)["semantic_check"]["status"] == "pass"
    original_plan = plan_path.read_bytes()
    original_check = check_path.read_bytes()


    invalid_reference_map = json.loads(json.dumps(semantic_input))
    invalid_reference_map["reference_resolution"]["cell_ranges"][0]["variable_id"] = "unknown.variable"
    invalid_reference_path = tmp_path / "invalid_reference_semantic_map.json"
    _write_json(invalid_reference_path, invalid_reference_map)
    blocked_reference = semantic.build_semantic_plan(analysis_dir, invalid_reference_path)
    assert blocked_reference["status"] == "blocked"
    assert "unknown variable" in blocked_reference["diagnostics"][0]["message"]
    assert plan_path.read_bytes() == original_plan
    assert check_path.read_bytes() == original_check


    invalid_execution_plan = json.loads(json.dumps(semantic_input))
    invalid_execution_plan["execution_plan"]["passes"][1]["source_family_ids"] = []
    invalid_execution_path = tmp_path / "invalid_execution_plan.json"
    _write_json(invalid_execution_path, invalid_execution_plan)
    blocked_execution_plan = semantic.build_semantic_plan(analysis_dir, invalid_execution_path)
    assert blocked_execution_plan["status"] == "blocked"
    assert "assign every source family exactly once" in blocked_execution_plan["diagnostics"][0]["message"]
    assert plan_path.read_bytes() == original_plan
    assert check_path.read_bytes() == original_check

    invalid_anchor = json.loads(json.dumps(semantic_input))
    invalid_anchor["array_mappings"][0]["instances"][0]["anchor_source_family_id"] = "wrong-source-family"
    invalid_anchor_path = tmp_path / "invalid_anchor_semantic_map.json"
    _write_json(invalid_anchor_path, invalid_anchor)
    blocked_anchor = semantic.build_semantic_plan(analysis_dir, invalid_anchor_path)
    assert blocked_anchor["status"] == "blocked"
    assert "mapped anchor source_family_id" in blocked_anchor["diagnostics"][0]["message"]
    assert plan_path.read_bytes() == original_plan
    assert check_path.read_bytes() == original_check

    omitted_ordinary_demand = json.loads(json.dumps(semantic_input))
    omitted_ordinary_demand["variables"][0]["demanded_indices"] = {"age": "1..1", "category": "0"}
    omitted_ordinary_path = tmp_path / "omitted_ordinary_demand.json"
    _write_json(omitted_ordinary_path, omitted_ordinary_demand)
    blocked_ordinary_demand = semantic.build_semantic_plan(analysis_dir, omitted_ordinary_path)
    assert blocked_ordinary_demand["status"] == "blocked"
    assert "mapped source member Main!F1 requires rates.matrix[0,0] outside demanded_indices" in \
        blocked_ordinary_demand["diagnostics"][0]["message"]
    assert plan_path.read_bytes() == original_plan
    assert check_path.read_bytes() == original_check

    omitted_array_demand = json.loads(json.dumps(semantic_input))
    omitted_array_demand["variables"][0]["demanded_indices"] = {"age": "0..0", "category": "0"}
    omitted_array_path = tmp_path / "omitted_array_demand.json"
    _write_json(omitted_array_path, omitted_array_demand)
    blocked_array_demand = semantic.build_semantic_plan(analysis_dir, omitted_array_path)
    assert blocked_array_demand["status"] == "blocked"
    assert "mapped array member Main!F2 requires rates.matrix[1,0] outside demanded_indices" in \
        blocked_array_demand["diagnostics"][0]["message"]
    assert plan_path.read_bytes() == original_plan
    assert check_path.read_bytes() == original_check

    invalid_map = json.loads(json.dumps(semantic_input))
    invalid_map["array_mappings"][0]["instances"][0]["shape"] = [1, 2]
    invalid_path = tmp_path / "invalid_semantic_map.json"
    _write_json(invalid_path, invalid_map)
    blocked = semantic.build_semantic_plan(analysis_dir, invalid_path)
    assert blocked["status"] == "blocked"
    assert plan_path.read_bytes() == original_plan
    assert check_path.read_bytes() == original_check

    trace_path.write_text(trace_path.read_text(encoding="utf-8").replace("=C3+1", "=C4+1"), encoding="utf-8")
    stale = semantic.validate_semantic_plan(analysis_dir, plan_path)
    assert stale["status"] == "blocked"
    assert check_path.read_bytes() == original_check


def test_execution_segments_cover_ordinary_and_array_member_indices() -> None:
    variable_id = "rates.matrix"
    family_id = "source-family-1"
    equation_id = "eq-rate"
    segment = {"segment_id": "array_anchor", "source_family_ids": [family_id],
               "equation_family_ids": [equation_id], "index_axis": "age", "execution": "vectorized",
               "index_start": 0, "index_stop_exclusive": 2,
               "axis_intervals": [{"axis": "age", "dimension": 0, "start": 0, "stop_exclusive": 2},
                                  {"axis": "category", "dimension": 1, "start": 0, "stop_exclusive": 1}],
               "source_row_start": 1, "source_row_stop_exclusive": 2,
               "fixed_axis_indices": [], "snapshot_before_update": False,
               "update_order": "none", "recurrence_group": None, "lagged_dependencies": []}
    variables = {variable_id: {"shape": [2, 1], "axes": [{"name": "age"}, {"name": "category"}],
                               "equation_segments": [segment]}}
    mapping = {"source_family_id": family_id, "equation_family_id": equation_id,
               "variable_id": variable_id, "equation_segment": "array_anchor",
               "index_mapping": {"kind": "matrix", "array_ref": "F1:F2", "axes": [
                   {"axis": "age", "dimension": 0, "source_coordinate": "row", "origin": 1, "index_origin": 0},
                   {"axis": "category", "dimension": 1, "source_coordinate": "column", "origin": "F", "index_origin": 0}]},
               "source_members": ["Main!F1"]}
    array_mappings = [{"array_family_id": "transpose", "instances": [{
        "anchor_source_family_id": family_id, "variable_id": variable_id,
        "anchor": "Main!F1", "array_ref": "F1:F2", "source_members": ["Main!F2"]}]}]

    semantic._validate_execution_segments(variables, [mapping], array_mappings)

    narrow_segment = json.loads(json.dumps(segment))
    narrow_segment["index_stop_exclusive"] = 1
    narrow_segment["axis_intervals"][0]["stop_exclusive"] = 1
    variables[variable_id]["equation_segments"] = [narrow_segment]
    with pytest.raises(ValueError, match="outside segment"):
        semantic._validate_execution_segments(variables, [mapping], array_mappings)

    variables[variable_id]["equation_segments"] = [segment]
    missing_segment = {**mapping, "equation_segment": "missing"}
    with pytest.raises(ValueError, match="source families do not match"):
        semantic._validate_execution_segments(variables, [missing_segment], array_mappings)

    fixed_segment = json.loads(json.dumps(segment))
    fixed_segment["axis_intervals"][1]["stop_exclusive"] = 2
    fixed_segment["fixed_axis_indices"] = [{"axis": "category", "dimension": 1, "index": 0}]
    variables[variable_id] = {"shape": [2, 2], "axes": [{"name": "age"}, {"name": "category"}],
                              "equation_segments": [fixed_segment]}
    fixed_mapping = json.loads(json.dumps(mapping))
    fixed_mapping["source_members"] = ["Main!G1"]
    with pytest.raises(ValueError, match="fixed-axis constraint"):
        semantic._validate_execution_segments(variables, [fixed_mapping], [])


def test_declared_table_blank_inputs_are_verified_against_source_bytes(tmp_path: Path) -> None:
    source_path = tmp_path / "source.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Qtable"
    worksheet["W110"] = "=1+1"
    workbook.save(source_path)
    workbook.close()
    source_sha = hashlib.sha256(source_path.read_bytes()).hexdigest()

    step1_root = tmp_path / "step1"
    run_dir = step1_root / "run"
    run_dir.mkdir(parents=True)
    _write_json(run_dir / "source.json", {"run_id": "run-1", "sha256": source_sha,
                                           "source_path": str(source_path)})
    source_binding = {"run_id": "run-1", "run_path": "run", "source_sha256": source_sha}
    variables = {
        "input.qtable_formula_tail_blanks": {
            "role": "raw", "shape": [1], "axes": [{"name": "rate_category"}],
            "demanded_indices": {"start": 0, "stop_exclusive": 1},
            "source_extents": [{"ref": "Qtable!W111", "coordinate_mapping": {
                "kind": "column_ordinal", "axes": [{"axis": "rate_category", "dimension": 0,
                    "source_coordinate": "column", "origin": "W", "index_origin": 0}]}}]},
        "source_bridge.qtable.c": {
            "role": "raw", "shape": [1], "axes": [{"name": "source_row_index"}],
            "demanded_indices": {"start": 0, "stop_exclusive": 1},
            "source_extents": [{"ref": "Qtable!C111", "coordinate_mapping": {
                "kind": "row_ordinal", "axes": [{"axis": "source_row_index", "dimension": 0,
                    "source_coordinate": "row", "origin": 111, "index_origin": 0}]}}]},
    }
    reference_resolution = {"named_references": [{"name": "Qtable", "table_binding": {
        "source_blank_inputs": [{"source_ref": "Qtable!C111", "role": "key_trailing_blank"},
                                {"source_ref": "Qtable!W111", "role": "formula_tail_blank",
                                 "variable_id": "input.qtable_formula_tail_blanks", "index": 0}]}}]}
    inventory = {"sheets": [{"name": "Qtable", "cells": [
        {"address": "W110", "formula_present": True, "kind": "formula", "data_type": "f", "formula": "=1+1"}]}]}

    verified = semantic._validate_source_blank_inputs(reference_resolution, variables, inventory,
                                                       source_binding, step1_root)
    assert len(verified) == 2
    assert {item["source_ref"] for item in verified} == {"Qtable!C111", "Qtable!W111"}

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Qtable"
    worksheet["W110"] = "=1+1"
    worksheet["W111"] = 7
    workbook.save(source_path)
    workbook.close()
    changed_sha = hashlib.sha256(source_path.read_bytes()).hexdigest()
    _write_json(run_dir / "source.json", {"run_id": "run-1", "sha256": changed_sha,
                                           "source_path": str(source_path)})
    source_binding["source_sha256"] = changed_sha
    with pytest.raises(ValueError, match="not a blank raw cell"):
        semantic._validate_source_blank_inputs(reference_resolution, variables, inventory,
                                               source_binding, step1_root)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Qtable"
    worksheet["W110"] = "=1+1"
    workbook.save(source_path)
    workbook.close()
    updated_sha = hashlib.sha256(source_path.read_bytes()).hexdigest()
    _write_json(run_dir / "source.json", {"run_id": "run-1", "sha256": updated_sha,
                                           "source_path": str(source_path)})
    source_binding["source_sha256"] = updated_sha
    inventory["sheets"][0]["cells"][0]["raw_formula_attributes"] = {"t": "array", "ref": "W110:W111"}
    with pytest.raises(ValueError, match="overlaps a source array formula"):
        semantic._validate_source_blank_inputs(reference_resolution, variables, inventory,
                                               source_binding, step1_root)



def test_raw_source_extents_reject_formula_cells_and_array_followers_but_allow_undemanded_formula() -> None:
    inventory = {"workbook_sha256": "source-sha", "sheets": [{"name": "Main", "cells": [
        {"address": "A1", "kind": "formula", "formula_present": True, "formula": "=1+1",
         "data_type": "f", "raw_formula_attributes": {}},
        {"address": "A2", "kind": "value", "formula_present": False, "data_type": "n"},
        {"address": "C1", "kind": "formula", "formula_present": True, "formula": "=TRANSPOSE(D1:D2)",
         "data_type": "f", "raw_formula_attributes": {"t": "array", "ref": "C1:D2"}},
        {"address": "D2", "kind": "value", "formula_present": False, "data_type": "n"},
    ]}]}

    def raw_variable(variable_id: str, ref: str, *, shape: list[int], axes: list[dict],
                     mapping: dict, demand: object) -> dict:
        return {"variable_id": variable_id, "role": "raw", "kind": "scalar" if not shape else "lookup_vector",
                "shape": shape, "axes": axes, "demanded_indices": demand,
                "source_extents": [{"ref": ref, "coordinate_mapping": mapping, "role": "raw_input"}]}

    formula_scalar = raw_variable("input.formula", "Main!A1", shape=[], axes=[],
                                  mapping={"kind": "scalar", "axes": []}, demand=[])
    with pytest.raises(ValueError, match="claims demanded formula cell Main!A1"):
        semantic.validate_raw_source_extents({"input.formula": formula_scalar}, inventory, "source-sha")

    array_follower = raw_variable("input.follower", "Main!D2", shape=[], axes=[],
                                  mapping={"kind": "scalar", "axes": []}, demand=[])
    with pytest.raises(ValueError, match="claims demanded array formula follower Main!D2"):
        semantic.validate_raw_source_extents({"input.follower": array_follower}, inventory, "source-sha")

    partial = raw_variable("input.partial", "Main!A1:A2", shape=[2], axes=[{"name": "row"}],
                           mapping={"kind": "row_ordinal", "axes": [{"axis": "row", "dimension": 0,
                               "source_coordinate": "row", "origin": 1, "index_origin": 0}]},
                           demand=[1])
    summary = semantic.validate_raw_source_extents({"input.partial": partial}, inventory, "source-sha")
    assert summary["demanded_raw_formula_overlaps"] == 0
    assert summary["formula_coordinates_outside_raw_demand"] == 1


def test_semantic_demand_includes_last_discount_source_index() -> None:
    variable = {"shape": [107], "axes": [{"name": "discount_period"}],
                "demanded_indices": {"start": 0, "stop_exclusive": 106, "step": 1}}
    index_mapping = {"kind": "projection_series", "axes": [
        {"axis": "discount_period", "dimension": 0, "source_coordinate": "row",
         "origin": 9, "index_origin": 0}]}
    last_discount_index = semantic._mapped_logical_indices(index_mapping, "Premium!T115")
    assert last_discount_index == (106,)
    assert not semantic._is_demanded(variable, last_discount_index)
    variable["demanded_indices"]["stop_exclusive"] = 107
    assert semantic._is_demanded(variable, last_discount_index)
