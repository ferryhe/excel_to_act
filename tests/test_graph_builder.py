from __future__ import annotations

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from excel_to_act.graph.builder import RegexFormulaGraphBuilder
from excel_to_act.ingest.openpyxl_reader import OpenpyxlWorkbookReader
from excel_to_act.inventory.extractor import OpenpyxlInventoryExtractor
from excel_to_act.schemas import (
    CellInventory,
    CellKind,
    CoverageSummary,
    FormulaGraph,
    GraphNodeKind,
    ModuleClassification,
    SheetInventory,
    SheetManifest,
    SourceLocation,
    WorkbookInventory,
    WorkbookManifest,
)
from excel_to_act.verify.completeness import verify_completeness


def _inventory(
    formulas: dict[tuple[str, str], str],
    *,
    sheet_names: tuple[str, ...] = ("Inputs",),
) -> WorkbookInventory:
    sheets = []
    for index, name in enumerate(sheet_names):
        cells = [
            CellInventory(
                source_location=SourceLocation(
                    workbook_path="book.xlsx",
                    sheet_name=name,
                    sheet_index=index,
                    address=address,
                    object_type="cell",
                    object_id=f"{name}!{address}",
                ),
                address=address,
                row=1,
                column=1,
                kind=CellKind.formula,
                formula=formula,
            )
            for (sheet_name, address), formula in formulas.items()
            if sheet_name == name
        ]
        sheets.append(
            SheetInventory(
                source_location=SourceLocation(
                    workbook_path="book.xlsx", sheet_name=name, sheet_index=index, object_type="sheet"
                ),
                name=name,
                index=index,
                max_row=1,
                max_column=max(1, len(cells)),
                cells=cells,
            )
        )
    count = sum(len(sheet.cells) for sheet in sheets)
    return WorkbookInventory(
        workbook_sha256="a" * 64,
        sheets=sheets,
        coverage=CoverageSummary(
            recognized_inventory_objects=count,
            unsupported_or_opaque_objects=0,
            discovered_workbook_objects=count,
        ),
    )


def _target(graph: FormulaGraph, source: str) -> str:
    return next(edge.target for edge in graph.edges if edge.source == f"cell:{source}")


def test_text_and_arithmetic_formulas_are_valid_zero_dependency(tmp_path) -> None:
    inventory = _inventory(
        {
            ("Inputs", "A1"): '="A1"',
            ("Inputs", "A2"): "=1+1",
            ("Inputs", "A3"): "=#DIV/0!",
            ("Inputs", "A4"): "=-1",
        }
    )
    graph = RegexFormulaGraphBuilder().build(inventory)

    assert graph.edges == []
    assert graph.unsupported_features == []
    assert all(node.metadata.get("references_resolved") for node in graph.nodes)

    manifest = WorkbookManifest(
        workbook_path=str(tmp_path / "not-present.xlsx"),
        file_name="not-present.xlsx",
        file_size=1,
        sha256="a" * 64,
        sheets=[SheetManifest(name="Inputs", index=0, max_row=4, max_column=1)],
    )
    check = next(
        check
        for check in verify_completeness(
            manifest, inventory, graph, ModuleClassification()
        ).checks
        if check.name == "formulas_linked"
    )
    assert check.passed


def test_empty_formula_is_unresolved_with_source_location() -> None:
    inventory = _inventory({("Inputs", "A1"): "="})
    graph = RegexFormulaGraphBuilder().build(inventory)

    assert graph.edges == []
    assert graph.nodes[0].metadata["references_resolved"] is False
    assert len(graph.unsupported_features) == 1
    diagnostic = graph.unsupported_features[0]
    assert diagnostic.feature_type == "formula_reference_unresolved"
    assert diagnostic.source_location == inventory.sheets[0].cells[0].source_location
    assert "empty" in diagnostic.description.lower()


def test_referenced_formula_cell_keeps_zero_reference_resolution(tmp_path) -> None:
    inventory = _inventory({("Inputs", "A1"): "=B1", ("Inputs", "B1"): "=1+1"})
    graph = RegexFormulaGraphBuilder().build(inventory)

    assert [(edge.source, edge.target) for edge in graph.edges] == [
        ("cell:Inputs!A1", "cell:Inputs!B1")
    ]
    assert next(node for node in graph.nodes if node.id == "cell:Inputs!B1").metadata[
        "references_resolved"
    ] is True

    manifest = WorkbookManifest(
        workbook_path=str(tmp_path / "not-present.xlsx"),
        file_name="not-present.xlsx",
        file_size=1,
        sha256="a" * 64,
        sheets=[SheetManifest(name="Inputs", index=0, max_row=1, max_column=2)],
    )
    check = next(
        check
        for check in verify_completeness(
            manifest, inventory, graph, ModuleClassification()
        ).checks
        if check.name == "formulas_linked"
    )
    assert check.passed


def test_a1_references_stay_within_excel_grid() -> None:
    inventory = _inventory(
        {
            ("Inputs", "A1"): "=XFE1",
            ("Inputs", "A2"): "=A1048577",
            ("Inputs", "A3"): "=XFD1",
            ("Inputs", "A4"): "=A1048576",
            ("Inputs", "A5"): "=XFD1048576",
        }
    )
    graph = RegexFormulaGraphBuilder().build(inventory)

    assert {edge.target for edge in graph.edges} == {
        "cell:Inputs!XFD1",
        "cell:Inputs!A1048576",
        "cell:Inputs!XFD1048576",
    }
    assert {feature.source_location.object_id for feature in graph.unsupported_features} == {
        "Inputs!A1",
        "Inputs!A2",
    }
    assert all(feature.source_location.sheet_name == "Inputs" for feature in graph.unsupported_features)


def test_trailing_infix_operator_is_unresolved_without_edges() -> None:
    inventory = _inventory({("Inputs", "A1"): "=A1+"})
    graph = RegexFormulaGraphBuilder().build(inventory)

    assert graph.edges == []
    assert graph.nodes[0].metadata["references_resolved"] is False
    assert len(graph.unsupported_features) == 1
    diagnostic = graph.unsupported_features[0]
    assert diagnostic.feature_type == "formula_reference_unresolved"
    assert diagnostic.source_location == inventory.sheets[0].cells[0].source_location
    assert "infix operator" in diagnostic.description.lower()


def test_structured_reference_and_defined_name_scopes(tmp_path) -> None:
    path = tmp_path / "names_and_table.xlsx"
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    inputs["A1"] = "Col"
    inputs["B1"] = "Value"
    inputs["A2"] = 1
    inputs["B2"] = 2
    inputs["C2"] = "=SUM(Table1[Col])"
    inputs["D2"] = "=SUM(MortRate)"
    inputs.add_table(Table(displayName="Table1", ref="A1:B2"))
    inputs.defined_names.add(DefinedName("MortRate", attr_text="'Inputs'!$A$2"))
    workbook.defined_names.add(DefinedName("MortRate", attr_text="'Inputs'!$B$2"))
    workbook.create_sheet("Other")["A1"] = "=SUM(MortRate)"
    workbook.save(path)
    workbook.close()

    manifest = OpenpyxlWorkbookReader().read_manifest(path)
    inventory = OpenpyxlInventoryExtractor().extract(path, manifest)
    graph = RegexFormulaGraphBuilder().build(inventory)

    table_node = next(node for node in graph.nodes if node.id == _target(graph, "Inputs!C2"))
    assert table_node.kind == GraphNodeKind.range
    assert table_node.metadata["table"] == "Table1"
    assert table_node.metadata["column"] == "Col"
    assert "Col" in table_node.metadata["table_columns"]
    assert all(node.kind != GraphNodeKind.external for node in graph.nodes)
    assert _target(graph, "Inputs!D2") == "name:Inputs!MortRate"
    assert _target(graph, "Other!A1") == "name:MortRate"


def test_sheet_references_keep_their_full_identity() -> None:
    inventory = _inventory(
        {
            ("Inputs", "A1"): "=Assumptions!B7",
            ("Inputs", "A2"): "='O''Brien'!$B$2",
        },
        sheet_names=("Inputs", "Assumptions", "O'Brien"),
    )
    graph = RegexFormulaGraphBuilder().build(inventory)

    assert _target(graph, "Inputs!A1") == "cell:Assumptions!B7"
    assert _target(graph, "Inputs!A2") == "cell:O'Brien!B2"


def test_external_reference_keeps_identity_without_local_edge() -> None:
    graph = RegexFormulaGraphBuilder().build(
        _inventory({("Inputs", "A3"): "=[Book.xlsx]Inputs!A1"})
    )

    external = _target(graph, "Inputs!A3")
    assert external.startswith("external:")
    assert "Book.xlsx" in external
    assert all(edge.target != "cell:Inputs!A1" for edge in graph.edges)


def test_unresolved_and_malformed_references_keep_source_location() -> None:
    inventory = _inventory(
        {
            ("Inputs", "A1"): '=INDIRECT("A1")',
            ("Inputs", "A2"): "=SUM(",
            ("Inputs", "A3"): "=UnknownName",
            ("Inputs", "A4"): "=#REF!",
        }
    )
    graph = RegexFormulaGraphBuilder().build(inventory)

    assert len(graph.unsupported_features) == 4
    assert {feature.source_location.object_id for feature in graph.unsupported_features} == {
        "Inputs!A1",
        "Inputs!A2",
        "Inputs!A3",
        "Inputs!A4",
    }
    assert all(feature.description for feature in graph.unsupported_features)
    assert any("UnknownName" in feature.description for feature in graph.unsupported_features)
    assert any("#REF!" in feature.description for feature in graph.unsupported_features)
