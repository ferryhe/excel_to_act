from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

from excel_to_act.steps.conversion_workflow import (
    append_stage_artifact,
    create_workflow,
    load_workflow,
    record_decision,
)
from excel_to_act.steps.step4 import discover_active_trace, generate_model
from excel_to_act.steps.step4.generator import _ordered_active_addresses
from excel_to_act.steps.step4.discovery import (
    _approved_stored_empty_addresses,
    _cell_key,
    _record_external_override,
)
from excel_to_act.steps.step5 import validate_generated
from excel_to_act.steps.step5.workflow import _copy_bundle_files, _verify_python


def _approve(workflow_dir: Path, stage: int) -> None:
    for reviewer in ("agent", "human"):
        record_decision(workflow_dir, stage, reviewer, "approve", f"Synthetic fixture review for stage {stage}.")


def _approved_gp_workflow(tmp_path: Path, *, array_formula: bool = False,
                          formula_text: str | None = None) -> tuple[Path, Path]:
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    source_path = source_dir / "model.xlsx"
    workbook = Workbook()
    main = workbook.active
    main.title = "Main"
    main["A1"] = 3
    main["A2"] = 4
    premium = workbook.create_sheet("Premium")
    formula = formula_text or ("=TRANSPOSE(Main!A1:A2)" if array_formula else "=Main!A1*2")
    if array_formula:
        premium["J1"] = ArrayFormula(ref="J1:K1", text=formula)
    else:
        premium["J1"] = formula
    workbook.defined_names.add(DefinedName("GP", attr_text="'Premium'!$J$1"))
    workbook.save(source_path)
    workbook.close()

    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    workflow_dir = tmp_path / "workflow"
    create_workflow(workflow_dir, {"workbook_sha256": source_hash, "run_id": "fixture-run",
                                   "source_id": "fixture-source", "workbook_path": str(source_path)})
    _, manifest = load_workflow(workflow_dir)
    manifest["mode"] = "synthetic_fixture"
    from excel_to_act.steps.conversion_workflow import _write_manifest
    _write_manifest(workflow_dir, manifest)

    stage1 = append_stage_artifact(workflow_dir, 1, "checkpoint.json",
                                   {"status": "ready_for_review", "source_sha256": source_hash},
                                   "Synthetic Stage 1 source checkpoint.", input_files={"source": source_path})
    _approve(workflow_dir, 1)
    stage2 = append_stage_artifact(workflow_dir, 2, "checkpoint.json",
                                   {"status": "ready_for_review", "source_sha256": source_hash},
                                   "Synthetic Stage 2 index checkpoint.", input_stages={1: stage1["artifact"]})
    _approve(workflow_dir, 2)
    stage3_payload = {
        "schema_version": "step3.analysis_design.v1",
        "status": "ready_for_review",
        "source": {"workbook_path": str(source_path), "workbook_sha256": source_hash,
                   "source_id": "fixture-source", "run_id": "fixture-run"},
        "design": {"targets": [{"target_id": "T_GP", "selector": "GP", "result_order": 0, "shape": []}],
                    "scenarios": [{"scenario_id": "saved-fixture"}], "validation_plan": [],
                    "coverage": {"known_path_ids": ["R_ONE"], "planned_path_ids": ["R_ONE"]},
                    "options": [{"selected_for_draft": True, "numerical_paths": ["R_ONE"]}],
                    "paths": [{"path_id": "R_ONE", "source_formula": formula}],
                    "historical_evidence": []},
    }
    append_stage_artifact(workflow_dir, 3, "analysis_design.json", stage3_payload,
                          "Synthetic approved-design fixture.", input_stages={2: stage2["artifact"]})
    _approve(workflow_dir, 3)
    return workflow_dir, source_path


def test_stored_empty_adapter_reads_need_candidate_and_exact_scope_evidence(tmp_path: Path) -> None:
    source_path = tmp_path / "stored-empty.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["C3"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    workbook.save(source_path)
    workbook.close()
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    scope = {"allowed_regions": [{"source_ref": "Main!C3", "rationale": "fixture"},
                                 {"source_ref": "Main!C3:C5", "rationale": "range is not a blank waiver"}]}
    trace = {"unknowns": [
        {"category": "unrepresented_source_coordinate", "evidence": "main!C3"},
        {"category": "unsupported_formula", "evidence": "Main!C3"},
        {"category": "unrepresented_source_coordinate", "evidence": "Main!C4"},
    ]}

    assert _approved_stored_empty_addresses(source_path, source_hash, scope, trace) == {"main!C3"}

    changed_scope = {"allowed_regions": [{"source_ref": "Main!C4", "rationale": "fixture"}]}
    with pytest.raises(ValueError, match="stored, formula-free blank"):
        _approved_stored_empty_addresses(source_path, source_hash, changed_scope, trace)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["C3"] = 12
    workbook.save(source_path)
    workbook.close()
    changed_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    with pytest.raises(ValueError, match="stored, formula-free blank"):
        _approved_stored_empty_addresses(source_path, changed_hash, scope, trace)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["C3"] = "=1"
    workbook.save(source_path)
    workbook.close()
    formula_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    with pytest.raises(ValueError, match="stored, formula-free blank"):
        _approved_stored_empty_addresses(source_path, formula_hash, scope, trace)


def test_external_override_keys_match_casefolded_trace_membership() -> None:
    overrides: dict[str, dict[str, object]] = {}
    _record_external_override(overrides, "Qtable!W5", {"value": 0.25})

    assert set(overrides) == {"qtable!W5"}
    assert _cell_key("qTaBlE!w5") in overrides

    with pytest.raises(ValueError, match="duplicated"):
        _record_external_override(overrides, "qTaBlE!w5", {"value": 0.5})


def test_generated_active_addresses_follow_formula_completion_order() -> None:
    trace = {"completion_order": ["Main!A1", "Premium!J1", "Main!B1"]}
    assert _ordered_active_addresses(trace, ["Main!B1", "Premium!J1"]) == ["Premium!J1", "Main!B1"]


def test_discovery_generation_smoke_and_stage5_isolated_validation(tmp_path: Path) -> None:
    workflow_dir, source_path = _approved_gp_workflow(tmp_path)

    discovery = discover_active_trace(workflow_dir)
    assert discovery["status"] == "pass", discovery
    assert discovery["target_values"]["T_GP"]["value"] == 6
    assert discovery["native_excel_called"] is False
    assert discovery["formula_cache_inputs"] is False

    generated = generate_model(workflow_dir)
    assert generated["status"] == "pass", generated
    assert generated["counts"]["formula_count"] == 1
    bundle = Path(generated["bundle"]["path"])
    assert bundle.is_relative_to(workflow_dir / "stage4" / "bundles")
    assert not (bundle / source_path.name).exists()
    assert Path(generated["artifact"]["json"]).parts[:2] == ("stage4", "revision-0001")

    from excel_to_act.steps.conversion_workflow import read_json
    generation_report = read_json(workflow_dir / generated["artifact"]["json"])
    assert generation_report["smoke_execution"] == "pass"
    assert generation_report["smoke"]["targets"]["GP"] == 6
    assert generation_report["active_trace"]["source"] == "stage4_discovery"
    assert generation_report["path_coverage"]["implemented_path_ids"] == ["R_ONE"]
    generation_markdown = (workflow_dir / generated["artifact"]["md"]).read_text(encoding="utf-8")
    assert "Covered source formula positions: 1" in generation_markdown
    assert "Generated formula expressions:" not in generation_markdown
    assert "Review decisions and current receipts are recorded in workflow.json." in generation_markdown
    reader_front = generation_markdown.split("## Detailed generation evidence", 1)[0]
    assert "## Requested targets" in reader_front
    assert "**GP**" in reader_front and "Generated smoke result: `6.0`" in reader_front
    assert "python model.py --out RESULT.json" in reader_front
    assert "accepted design hash" not in reader_front
    entry_line = next(line for line in reader_front.splitlines() if "Generated entry point:" in line)
    relative_entry = entry_line.split("<", 1)[1].split(">", 1)[0]
    report_dir = (workflow_dir / generated["artifact"]["md"]).parent
    assert (report_dir / relative_entry).resolve().is_file()

    _approve(workflow_dir, 4)
    validation = validate_generated(workflow_dir)
    assert validation["status"] == "pass", validation
    assert validation["result_summary"]["targets"]["GP"] == 6



def test_generation_reader_shows_ordered_non_scalar_values_and_unknowns_without_mutation() -> None:
    import json
    from excel_to_act.steps.step4.generator import _generation_markdown

    design = {
        "targets": [
            {"selector": "Scalar", "result_order": 2, "shape": [], "result_kind": None, "units": None},
            {"selector": "Vector", "result_order": 0, "shape": [3], "result_kind": "vector", "units": "USD",
             "axes": [{"name": "duration", "role": "time", "keys": [0, False, "None"]}]},
            {"selector": "Matrix", "result_order": 1, "shape": [2, 2], "result_kind": "table", "units": "count",
             "axes": []},
        ],
        "open_questions": [
            {"question_id": "Q_CAPTURE", "question": "Capture is required before execution.",
             "status": "resolved by recorded capture", "disposition": "execution prerequisite"},
            {"question_id": "Q_MEANING", "question": "What does the selected axis mean?",
             "status": None, "disposition": None},
        ],
    }
    payload = {
        "source_sha256": "source-fixture", "design_sha256": "design-fixture",
        "active_trace": {"source": "synthetic fixture", "sha256": "trace-fixture"},
        "formula_count": 0, "raw_input_count": 0, "array_member_count": 0,
        "array_instance_count": 0, "compiled_function_count": 0, "source_family_count": 0,
        "semantic_equation_group_count": None,
        "business_input_count": None, "raw_business_input_count": 0,
        "external_business_input_count": False,
        "smoke": {"status": None, "formula_count": 0,
                  "targets": {"Matrix": [[1, 2], [3, 4]], "Scalar": False,
                              "Vector": [0, False, None], "Diagnostic": [7, 8]}},
        "path_coverage": {"implemented_path_ids": [], "numerical_path_ids": [],
                          "implementation_evidence": [], "condition_only_path_ids": []},
        "bundle": {"path": "bundle"},
    }
    before = json.dumps(payload, sort_keys=True)
    report = _generation_markdown(payload, design=design, machine_link="machine.json")
    front = report.split("## Detailed generation evidence", 1)[0]

    assert front.index("**Vector**") < front.index("**Matrix**") < front.index("**Scalar**")
    assert "position [0] = 0" in front and "position [1] = False" in front
    assert "position [2] = Not recorded" in front
    assert "position [0, 0] = 1" in front and "position [1, 1] = 4" in front
    assert "shape []; units Not recorded" in front
    assert "Generated smoke result: `False`" in front
    assert "keys: position [0] = 0<br>position [1] = False<br>position [2] = None" in front
    assert "business input objects: Not recorded (0 raw, False formula-derived external)" in front
    assert "**Q_CAPTURE**" in front and "execution prerequisite" in front
    assert "resolved by recorded capture" in front and "**Q_MEANING**" in front
    assert "Recorded answer: Not recorded" in front and "status: Not recorded" in front
    assert "[generation_report.json](<machine.json>)" in front
    assert "[[1, 2], [3, 4]]" not in front
    assert json.dumps(payload, sort_keys=True) == before

def test_native_array_anchor_is_emitted_as_a_formula_not_a_raw_leaf(tmp_path: Path) -> None:
    workflow_dir, _source_path = _approved_gp_workflow(tmp_path, array_formula=True)
    discovery = discover_active_trace(workflow_dir)
    assert discovery["status"] == "pass", discovery
    generated = generate_model(workflow_dir)
    assert generated["status"] == "pass", generated
    assert generated["counts"]["formula_count"] == 1
    assert generated["counts"]["array_member_count"] == 1
    assert generated["counts"]["raw_input_count"] == 2
    from excel_to_act.steps.conversion_workflow import read_json
    report = read_json(workflow_dir / generated["artifact"]["json"])
    assert report["smoke"]["targets"]["GP"] == 3
    assert report["path_coverage"]["implemented_path_ids"] == ["R_ONE"]


def test_single_cell_reference_returned_from_if_is_scalarized(tmp_path: Path) -> None:
    workflow_dir, _source_path = _approved_gp_workflow(
        tmp_path, formula_text="=IF(TRUE,Main!A1,0)")
    assert discover_active_trace(workflow_dir)["status"] == "pass"
    generated = generate_model(workflow_dir)
    assert generated["status"] == "pass", generated
    from excel_to_act.steps.conversion_workflow import read_json
    report = read_json(workflow_dir / generated["artifact"]["json"])
    assert report["smoke"]["targets"]["GP"] == 3


def test_step5_validates_and_copies_every_local_bundle_module(tmp_path: Path) -> None:
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    contents = {
        "model.py": "from pricing import run\nfrom runtime import Runtime\n",
        "pricing.py": "from formula_families import compute\nfrom input_adapter import inputs\n",
        "formula_families.py": "def compute(value):\n    return value\n",
        "input_adapter.py": "def inputs():\n    return {}\n",
        "runtime.py": "class Runtime:\n    pass\n",
        "source_values.json": "{}\n",
        "source_map.json": "{}\n",
        "family_manifest.json": "{}\n",
    }
    file_hashes = {}
    for name, source in contents.items():
        path = bundle / name
        path.write_text(source, encoding="utf-8")
        file_hashes[name] = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = {"files": file_hashes, "raw_input_addresses": [], "formula_addresses": ["Premium!J1"],
                "formula_cache_inputs": False}

    assert _verify_python(bundle, manifest) == []
    destination = tmp_path / "isolated"
    copied = _copy_bundle_files(bundle, file_hashes, destination)
    assert set(copied) == set(contents)
    assert (destination / "formula_families.py").is_file()
    assert (destination / "source_map.json").read_text(encoding="utf-8") == "{}\n"

    bad_source = "import excel_to_act\n"
    bad_path = bundle / "formula_families.py"
    bad_path.write_text(bad_source, encoding="utf-8")
    file_hashes["formula_families.py"] = hashlib.sha256(bad_path.read_bytes()).hexdigest()
    assert any("imports unsupported module excel_to_act" in failure
               for failure in _verify_python(bundle, {**manifest, "files": file_hashes}))

