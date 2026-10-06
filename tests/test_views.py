from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from excel_to_act.schemas import SourceLocation, ViewRecord, WorkbookView
from excel_to_act.steps.step1.workflow import convert_directory
from excel_to_act.views import compile_views, serialize_views, validate_agent_output


def _source(tmp_path: Path) -> tuple[Path, Path, Path]:
    raw = tmp_path / "raw"
    raw.mkdir()
    book_path = raw / "book.xlsx"
    book = Workbook()
    book.active.title = "Facts"
    book.active["A1"] = "fact"
    book.active["A2"] = 7
    book.active["A3"] = 9
    book.active.add_table(Table(displayName="FactsTable", ref="A1:A3"))
    other = book.create_sheet("Calc")
    other["B2"] = "=Facts!A2"
    other["C3"] = "=MissingName"
    book.save(book_path)
    book.close()
    out = tmp_path / "step1"
    result = convert_directory(raw, out)
    batch = out / result["artifacts"][0]["path"]
    return out / result["entries"][0]["run_path"] / "handoff.json", out, batch


def _view() -> WorkbookView:
    loc = SourceLocation(workbook_path="book.xlsx", sheet_name="S", address="A1", object_type="cell", object_id="S!A1")
    record = ViewRecord(record_id="r1", record_type="cell", source_location=loc, facts={"value": 7})
    return WorkbookView(view_id="v1", source_id="s1", source_sha256="sha", source_run_id="run",
                        source_schema_version="phase1.v1", scope="sheet", sheet_name="S",
                        records=[record], chunks=[], opaque_report=[])


def test_compile_is_deterministic_budgeted_and_traceable_across_sheets(tmp_path: Path) -> None:
    handoff, root, _ = _source(tmp_path)
    first = compile_views(handoff, root, budget=1)
    second = compile_views(handoff, root, budget=1)
    assert serialize_views(first) == serialize_views(second)
    assert {view.sheet_name for view in first} >= {"Facts", "Calc"}
    facts_view = next(view for view in first if view.sheet_name == "Facts")
    cell = next(record for record in facts_view.records if record.record_type == "cell" and record.facts["address"] == "A2")
    assert cell.source_location.sheet_name == "Facts"
    assert cell.source_location.address == "A2"
    assert cell.facts["value"] == 7
    assert any(cell.record_id in chunk.over_budget_record_ids for chunk in facts_view.chunks)
    workbook_view = next(view for view in first if view.scope == "workbook")
    part = next(record for record in workbook_view.records if record.record_type == "package_part")
    assert part.source_location.ooxml_part and part.source_location.object_id
    assert all(chunk.over_budget_record_ids for view in first for chunk in view.chunks if chunk.estimated_tokens > 1)
    for view in first:
        chunk_ids = [record_id for chunk in view.chunks for record_id in chunk.record_ids]
        assert chunk_ids == [record.record_id for record in view.records]
        assert len(chunk_ids) == len(set(chunk_ids))
        records = {record.record_id: record for record in view.records}
        for chunk in view.chunks:
            for record_id in chunk.over_budget_record_ids:
                assert records[record_id] is not None
                assert chunk.estimated_tokens >= 1
                assert chunk.over_budget_record_ids == chunk.record_ids
    calc_view = next(view for view in first if view.sheet_name == "Calc")
    edge = next(record for record in calc_view.records if record.record_type == "dependency")
    assert edge.facts["target"].endswith("Facts!A2")
    assert edge.facts["cross_sheet"] is True
    unresolved = next(record for record in calc_view.records if record.record_type == "unresolved_reference")
    assert unresolved.source_location.sheet_name == "Calc"
    assert unresolved.source_location.address == "C3"
    assert unresolved.facts["feature_type"] == "formula_reference_unresolved"
    assert "MissingName" in unresolved.facts["description"]
    assert not any(
        record.record_type == "dependency" and record.facts["source"] == "cell:Calc!C3"
        for record in calc_view.records
    )
    region = next(view for view in first if view.scope == "region" and view.sheet_name == "Facts")
    assert region.region_address == "A1:A3"
    range_record = next(record for record in region.records if record.record_type == "range")
    assert range_record.source_location.sheet_name == "Facts"
    assert range_record.source_location.address == "A1:A3"
    ranged_cell = next(record for record in region.records if record.record_type == "cell" and record.facts["address"] == "A2")
    assert ranged_cell.source_location.sheet_name == "Facts"
    assert ranged_cell.source_location.address == "A2"
    sheet_cell = next(record for record in facts_view.records if record.record_type == "cell" and record.facts["address"] == "A2")
    assert ranged_cell.facts == sheet_cell.facts


def test_compile_accepts_batch_handoff_without_guessing_source(tmp_path: Path) -> None:
    _, root, batch_handoff = _source(tmp_path)
    views = compile_views(batch_handoff, root)
    assert views
    assert len({(view.source_id, view.source_run_id, view.source_sha256) for view in views}) == 1


def test_agent_validator_accepts_three_explicit_claim_types_and_opaque_report(tmp_path: Path) -> None:
    view = _view()
    loc = view.records[0].source_location.model_dump(mode="json")
    common = {"view_id": view.view_id, "source_run_id": view.source_run_id,
              "source_schema_version": view.source_schema_version, "record_id": "r1", "source_location": loc}
    result = validate_agent_output({"claims": [
        {**common, "kind": "fact", "value": {"value": 7}},
        {**common, "kind": "derivation", "derivation": "A stated calculation"},
        {**common, "kind": "unverified", "verified": False},
    ]}, [view])
    assert result == {"valid": True, "opaque_reports": []}


def test_agent_validator_rejects_missing_location_unknown_view_and_wrong_run() -> None:
    view = _view()
    good = {"view_id": view.view_id, "source_run_id": view.source_run_id,
            "source_schema_version": view.source_schema_version, "record_id": "r1",
            "source_location": view.records[0].source_location.model_dump(mode="json"), "kind": "fact", "value": {"value": 7}}
    invalid = [
        {k: v for k, v in good.items() if k != "source_location"},
        {**good, "view_id": "missing"},
        {**good, "source_run_id": "wrong"},
        {**good, "source_schema_version": "wrong"},
        {**good, "record_id": "missing"},
        {**good, "value": {"value": 8}},
    ]
    for claim in invalid:
        with pytest.raises(ValueError):
            validate_agent_output({"claims": [claim]}, [view])


def test_agent_validator_rejects_json_type_changes_in_facts_and_locations() -> None:
    view = _view()
    record = view.records[0].model_copy(update={"facts": {"value": 1}})
    location = record.source_location.model_copy(update={"sheet_index": 0})
    record = record.model_copy(update={"source_location": location})
    view = view.model_copy(update={"records": [record]})
    claim = {
        "view_id": view.view_id,
        "source_run_id": view.source_run_id,
        "source_schema_version": view.source_schema_version,
        "record_id": record.record_id,
        "source_location": location.model_dump(mode="json"),
        "kind": "fact",
        "value": {"value": 1},
    }
    assert claim["source_location"]["sheet_index"] == 0
    with pytest.raises(ValueError, match="fact differs"):
        validate_agent_output({"claims": [{**claim, "value": {"value": True}}]}, [view])
    with pytest.raises(ValueError, match="source_location mismatch"):
        validate_agent_output({"claims": [{**claim, "source_location": {**claim["source_location"], "sheet_index": False}}]}, [view])


def test_agent_validator_reports_opaque_without_guessing() -> None:
    view = _view()
    record = view.records[0].model_copy(update={"facts": {"opaque": True}})
    view = view.model_copy(update={"records": [record]})
    claim = {"view_id": view.view_id, "source_run_id": view.source_run_id,
             "source_schema_version": view.source_schema_version, "record_id": record.record_id,
             "source_location": record.source_location.model_dump(mode="json"),
             "kind": "opaque", "reported_opaque": True}
    assert len(validate_agent_output({"claims": [claim]}, [view])["opaque_reports"]) == 1


def test_exported_schema_is_the_runtime_model_schema() -> None:
    schema = json.loads(Path("schemas/workbook_view.schema.json").read_text(encoding="utf-8"))
    assert schema == WorkbookView.model_json_schema()
