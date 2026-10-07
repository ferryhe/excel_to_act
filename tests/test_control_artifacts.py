from __future__ import annotations

import zipfile
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
from typer.testing import CliRunner

from excel_to_act.ingest import control_artifacts
from excel_to_act.ingest.control_artifacts import (
    build_vba_handoff,
    enrich_form_controls,
    identify_activex,
    identify_checkboxes,
)
from excel_to_act.ingest.vba import VbaProject
from excel_to_act.interfaces.cli import app
from excel_to_act.schemas import SourceLocation, VbaHandoff, VbaModule
from excel_to_act.ingest.vba import _PROC_RE
from excel_to_act.steps.step1.workflow import convert_directory, execute_tool, finalize_run
from excel_to_act.steps.step2.workflow import build_index, validate_saved_index


def _fixture(path: Path) -> Path:
    workbook = Workbook()
    workbook.active.title = "Main"
    workbook.create_sheet("Other Sheet")
    workbook.save(path)
    workbook.close()

    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    controls = [
        (400, "Check Box", "rIdCtrl", "'Other Sheet'!$B$2"),
        (401, "Named", "rIdNamed", "MyRange"),
        (402, "Dynamic", "rIdDynamic", "OFFSET(A1,1,0)"),
        (403, "Broken", "rIdBroken", "#REF!"),
        (404, "MissingSheet", "rIdMissing", "Missing!B2"),
    ]
    declarations = "".join(f'<control shapeId="{shape}" r:id="{rel}" name="{name}"/>' for shape, name, rel, _raw in controls)
    sheet = sheet.replace("</worksheet>", f'<sheetPr codeName="RenamedCode"/><controls xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">{declarations}<control shapeId="900" r:id="rIdAx" name="Cmd"/><control shapeId="900" r:id="rIdAx" name="Cmd"/></controls></worksheet>')
    parts["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    rels = ['<Relationship Id="rIdVml" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" Target="../drawings/vmlDrawing1.vml"/>', '<Relationship Id="rIdAx" Type="http://schemas.microsoft.com/office/2006/relationships/control" Target="../activeX/activeX1.xml"/>']
    rels.extend(f'<Relationship Id="{rel}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/ctrlProp" Target="../ctrlProps/{rel}.xml"/>' for _shape, _name, rel, _raw in controls)
    parts["xl/worksheets/_rels/sheet1.xml.rels"] = ("<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">" + "".join(rels) + "</Relationships>").encode()
    for _shape, _name, rel, raw in controls:
        parts[f"xl/ctrlProps/{rel}.xml"] = f'<formControlPr xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" objectType="CheckBox" fmlaLink="{raw}"/>'.encode()
    parts["xl/drawings/vmlDrawing1.vml"] = (
        '<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:x="urn:schemas-microsoft-com:office:excel">'
        '<v:shape id="Check_x0020_Box" o:spid="_x0000_s400"><x:ClientData ObjectType="Checkbox"><x:FmlaLink>&apos;Other Sheet&apos;!$B$2</x:FmlaLink></x:ClientData></v:shape></xml>'
    ).encode()
    parts["xl/activeX/activeX1.xml"] = b'<ax:ocx xmlns:ax="http://schemas.microsoft.com/office/2006/activeX" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ax:classid="{test}" r:id="rId1"/>'
    parts["xl/activeX/_rels/activeX1.xml.rels"] = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2006/relationships/activeXControlBinary" Target="activeX1.bin"/></Relationships>'
    parts["xl/activeX/activeX1.bin"] = b"opaque"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return path


def test_checkbox_sources_merge_and_preserve_unresolved_links(tmp_path: Path) -> None:
    source = _fixture(tmp_path / "controls.xlsx")
    artifact = identify_checkboxes(source)
    assert len(artifact.bindings) == 5
    assert artifact.status == "partial"
    by_id = {item.shape_id: item for item in artifact.bindings}
    assert by_id["400"].linked_cell == "'Other Sheet'!B2"
    assert by_id["400"].binding_status == "resolved"
    assert len(by_id["400"].sources) == 2
    assert by_id["401"].binding_status == "named"
    assert by_id["402"].binding_status == "dynamic"
    assert by_id["403"].binding_status == "invalid"
    assert by_id["403"].linked_cell_raw == "#REF!"
    assert by_id["404"].binding_status == "unresolved"

    record = SimpleNamespace(
        kind="form_control",
        source_location=SourceLocation(sheet_name="Main", ooxml_part="xl/drawings/vmlDrawing1.vml", object_type="form_control"),
        metadata={"control_id": "Check_x0020_Box"},
    )
    inventory = SimpleNamespace(sheets=[SimpleNamespace(name="Main", ranges=[record])])
    enrich_form_controls(inventory, artifact)
    assert record.metadata["linked_cell"] == "'Other Sheet'!B2"
    assert record.metadata["linked_sheet"] == "Other Sheet"
    assert record.metadata["linked_address"] == "B2"


def test_activex_maps_only_to_sheet_codename_module_and_deduplicates(tmp_path: Path) -> None:
    source = _fixture(tmp_path / "controls.xlsm")
    project = VbaProject(
        available=True,
        modules=[
            VbaModule(name="RenamedCode", kind="ClassModule", procedures=["Cmd_Click"]),
            VbaModule(name="OtherSheet", kind="ClassModule", procedures=["Cmd_Click"]),
        ],
    )
    artifact = identify_activex(source, project)
    assert len(artifact.controls) == 1
    control = artifact.controls[0]
    assert control.sheet == "Main"
    assert control.sheet_code_name == "RenamedCode"
    assert control.binary_part == "xl/activeX/activeX1.bin"
    assert control.event_procedures == ["Cmd_Click"]
    assert control.binding_status == "resolved"


def test_vba_proc_declarations_and_standalone_conversion_are_checked(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "book.xlsx"
    _fixture(source)
    modules = [
        VbaModule(name="ModuleA", kind="StdModule", code="Public Sub First()\nEnd Sub\nPrivate Sub Second()\nEnd Sub\n", procedures=["First", "Second"]),
        VbaModule(name="ModuleB", kind="ClassModule", code="Private Sub Before()\nEnd Sub\n", procedures=["Before"]),
    ]
    project = VbaProject(available=True, modules=modules)
    assert _PROC_RE.findall(modules[0].code) == ["First", "Second"]
    monkeypatch.setattr(control_artifacts, "extract_vba_project", lambda _path: project)
    vba, source_files = build_vba_handoff(source, project)
    assert [item.procedures for item in vba.modules] == [["First", "Second"], ["Before"]]
    assert all(source_files[item.source_file].decode("utf-8") == module.code for item, module in zip(vba.modules, modules))

    runner = CliRunner()
    out = tmp_path / "vba-output"
    unmanaged = out / "vba_sources" / "keep.txt"
    unmanaged.parent.mkdir(parents=True)
    unmanaged.write_bytes(b"user file")
    result = runner.invoke(app, ["step1", "vba", "convert", str(source), "--out", str(out)])
    assert result.exit_code == 0, result.output
    assert unmanaged.read_bytes() == b"user file"

    monkeypatch.setattr(control_artifacts, "extract_vba_project", lambda _path: VbaProject(available=True, modules=modules[:1]))
    result = runner.invoke(app, ["step1", "vba", "convert", str(source), "--out", str(out)])
    assert result.exit_code == 0, result.output
    assert not (out / vba.modules[1].source_file).exists()
    assert unmanaged.read_bytes() == b"user file"
    result = runner.invoke(app, ["step1", "vba", "evaluate", str(source), "--out", str(out)])
    assert result.exit_code == 0, result.output
    manifest = VbaHandoff.model_validate_json((out / "vba_handoff.json").read_bytes())
    exported = out / manifest.modules[0].source_file
    exported.write_bytes(b"tampered")
    result = runner.invoke(app, ["step1", "vba", "evaluate", str(source), "--out", str(out)])
    assert result.exit_code == 1


def test_checkbox_dry_run_does_not_touch_existing_or_absent_output(tmp_path: Path) -> None:
    source = _fixture(tmp_path / "book.xlsx")
    runner = CliRunner()
    absent = tmp_path / "absent"
    result = runner.invoke(app, ["step1", "checkbox", "convert", str(source), "--out", str(absent), "--dry-run"])
    assert result.exit_code == 0, result.output
    assert not absent.exists()

    existing = tmp_path / "existing"
    existing.mkdir()
    sentinel = existing / "sentinel.txt"
    sentinel.write_bytes(b"unchanged")
    result = runner.invoke(app, ["step1", "checkbox", "convert", str(source), "--out", str(existing), "--dry-run"])
    assert result.exit_code == 0, result.output
    assert sentinel.read_bytes() == b"unchanged"
    assert list(existing.iterdir()) == [sentinel]


def test_step1_modules_survive_finalize_and_step2_validation(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    workbook = Workbook()
    workbook.active["A1"] = "source"
    workbook.save(input_dir / "book.xlsx")
    workbook.close()
    output = tmp_path / "step1"
    converted = convert_directory(input_dir, output)
    assert converted["status"] == "pass"
    run_dir = output / converted["entries"][0]["run_path"]

    checked = execute_tool("step1.check", run_dir)
    assert checked["status"] == "pass"
    ref_names = {item["name"] for item in checked["artifacts"]}
    assert {"checkbox_bindings.json", "activex_events.json", "vba_handoff.json", "controls_handoff.md", "vba_handoff.md"} <= ref_names
    finalized = finalize_run(run_dir)
    assert finalized["ready_for_next_step"] is True
    final_dir = output / finalized["final_output"]
    assert (final_dir / "vba_handoff.md").is_file()

    step2_out = tmp_path / "step2"
    indexed = build_index(final_dir / "handoff.json", output, step2_out)
    assert indexed["status"] != "blocked"
    validated = validate_saved_index(step2_out / "index.json", output)
    assert validated["status"] == "pass"
    markdown = (step2_out / "INDEX.md").read_text(encoding="utf-8")
    assert "Human control/VBA handoff" in markdown
