from __future__ import annotations

import json
import hashlib
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Barrier

import pytest

from openpyxl import Workbook

from excel_to_act.steps.conversion_workflow import (
    _write_manifest,
    append_stage_artifact,
    create_workflow,
    hash_file,
    load_workflow,
    record_decision,
)
from excel_to_act.steps.step5.workflow import (
    _active_formula_addresses,
    _active_formula_details,
    _bound_stage3_design,
    _compare_named_targets,
    _compare_active_formula_members,
    _compare_external_boundary,
    _oracle_range_cells,
    _preflight_oracle,
    _matches,
    _target_shapes,
    _value_has_shape,
    _reconciliation_markdown,
    _verify_python,
    reconcile,
    validate_generated,
)
from excel_to_act.steps.step5 import workflow as step5_workflow


def test_exact_case_formula_trace_replaces_saved_scenario_members(tmp_path: Path) -> None:
    trace = tmp_path / "case-trace.json"
    trace.write_text(json.dumps({"source_sha256": "source", "cells": [
        {"address": "Premium!F20", "role": "calculated_formula", "formula": "=A1+1"},
        {"address": "Premium!G20", "role": "calculated_array_formula", "formula": "=SUM(A1:A2)"},
    ]}), encoding="utf-8")
    report = {"source_sha256": "source", "active_trace": {"path": str(tmp_path / "missing.json")}}
    manifest = {"active_formula_addresses": ["Premium!A1"]}
    ref = {"path": str(trace), "sha256": hash_file(trace)}
    addresses, formulas = _active_formula_details(report, manifest, case_trace=ref)
    assert addresses == ["Premium!F20", "Premium!G20"]
    assert set(formulas) == {"premium!F20", "premium!G20"}
    assert step5_workflow._case_formula_ranges(json.loads(trace.read_text())) == ["Premium!F20:G20"]
    with pytest.raises(ValueError, match="changed"):
        _active_formula_details(report, manifest, case_trace={**ref, "sha256": "stale"})


def _synthetic_receipt_set(tmp_path: Path, monkeypatch) -> tuple[dict, list[Path]]:
    """Small synthetic receipt set; no workbook or real acceptance evidence."""
    root = tmp_path / "workflow"
    cases, variants, reports = [], [], []
    by_input = {}

    def put(path: Path, value: dict) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(value), encoding="utf-8")
        return path

    for index in range(32):
        case_id = f"case-{index:02d}"
        input_sha = f"{index + 1:064x}"
        dependency_sha = f"{index + 101:064x}"
        case_root = tmp_path / "source-cases" / case_id
        refs = {}
        for label in ("trace", "raw_request", "public_request", "producer", "profile", "plan"):
            value = {"case_identity": {"case_id": case_id}} if label == "plan" else {"case_id": case_id}
            path = put(case_root / f"{label}.json", value)
            refs[label] = {"path": str(path), "sha256": hash_file(path)}
        nested_refs = {}
        for label in ("semantic_map", "case_manifest", "inventory"):
            path = put(case_root / f"{label}.json", {"case_id": case_id})
            nested_refs[label] = {"path": str(path), "sha256": hash_file(path)}
        plan_path = Path(refs["plan"]["path"])
        put(plan_path, {"case_identity": {"case_id": case_id,
                                          "case_manifest": nested_refs["case_manifest"]},
                        "semantic_map_input": nested_refs["semantic_map"]})
        refs["plan"]["sha256"] = hash_file(plan_path)
        profile_path = Path(refs["profile"]["path"])
        put(profile_path, {"case_id": case_id, "inventory": nested_refs["inventory"]})
        refs["profile"]["sha256"] = hash_file(profile_path)
        cases.append(refs)
        variant = {"input_sha256": input_sha, "dependency_sha256": dependency_sha}
        variants.append(variant)
        formula_address = f"Premium!A{index + 1}"
        trace = {"cells": [{"address": formula_address, "role": "calculated_formula", "formula": "=1+1"}]}
        by_input[input_sha] = (case_id, refs, variant, trace)
        execution = root / "stage5" / "execution" / case_id
        request = put(execution / "input_request.json", {"case_id": case_id})
        refs["public_request"]["sha256"] = hash_file(request)
        model = put(execution / "model_result.json", {"status": "pass", "input_sha256": input_sha,
                                                      "dependency_sha256": dependency_sha,
                                                      "cells": {formula_address: 2}, "targets": {"GP": 7}})
        validation = put(execution / "execution_validation.json", {
            "schema_version": "step5.execution_validation.v1", "status": "pass",
            "compile_pass": True, "standalone_run_pass": True, "source_sha256": "source",
            "stage4_artifact_sha256": "stage4", "input_sha256": input_sha,
            "dependency_sha256": dependency_sha, "request_sha256": hash_file(request),
            "bundle_files": {"model.py": "synthetic-hash"},
            "result_path": str(model), "result_sha256": hash_file(model),
        })
        oracle = put(root / "stage5" / "oracle" / case_id / "excel_oracle.json", {
            "engine": "Microsoft Excel", "source_copy_sha256": "source",
            "macros_executed": False, "validation_sha256": hash_file(validation),
            "named_targets": {"GP": {"value": 7}},
            "ranges": {formula_address: {"qualified_address": formula_address,
                                        "values": [[2]], "formulas": [["=1+1"]],
                                        "excel_errors": [[None]]}},
        })
        report = put(root / "stage5" / "cases" / case_id / "case_report.json", {
            "schema_version": "step5.validation_report.v1", "execution_mode": "variable_input",
            "status": "ready_for_review", "native_verified_exact_input": True,
            "mismatch_count": 0, "source_sha256": "source",
            "stage4_artifact_sha256": "stage4", "stage4_bundle_manifest_sha256": "manifest",
            "stage3_design_sha256": "design", "input_sha256": input_sha,
            "dependency_sha256": dependency_sha, "tolerances": {"absolute": 1e-12, "relative": 1e-12},
            "formula_cache_inputs": False, "standalone_project_imports": False,
            "exact_case": {"case_id": case_id, "references": refs, "variant": variant,
                           "formula_member_count": 1},
            "compared": {"comparison_addresses": [formula_address], "active_formula_members": 1,
                         "active_formula_identity_matches": 1, "named_targets": 1, "required_targets": 1},
            "target_results": {"GP": {"model": 7, "excel": 7, "matched": True}},
            "standalone_validation": {"path": str(validation), "sha256": hash_file(validation),
                                      "result_sha256": hash_file(model)},
            "oracle": {"path": str(oracle), "sha256": hash_file(oracle)},
        })
        reports.append(report)
    monkeypatch.setattr(step5_workflow, "_current_exact_case",
                        lambda _contract, validation, _source: by_input[validation["input_sha256"]])
    monkeypatch.setattr(step5_workflow, "_preflight_oracle",
                        lambda oracle, _source, _stage4, _input: oracle)
    contract = {"design": {"artifact": {"json_sha256": "design"}},
                "support_domain": {"cases": cases, "status": "exact_cases",
                                   "tolerances": {"absolute": 1e-12, "relative": 1e-12}},
                "verified_cases": variants}
    return contract, reports


def test_exact_case_finalization_requires_all_current_distinct_receipts(tmp_path: Path, monkeypatch) -> None:
    contract, reports = _synthetic_receipt_set(tmp_path, monkeypatch)
    root = tmp_path / "workflow"
    records, inputs = step5_workflow._verify_case_receipts(
        root, reports, contract, "source", "stage4", "manifest", ["GP"],
        {"model.py": "synthetic-hash"})
    assert len(records) == 32 and len({record["trace_sha256"] for record in records}) == 32
    assert len([name for name in inputs if name.endswith("_native_oracle")]) == 32
    for bad in (reports[:-1], [*reports[:-1], reports[0]]):
        with pytest.raises(ValueError, match="32|duplicate"):
            step5_workflow._verify_case_receipts(root, bad, contract, "source", "stage4", "manifest", ["GP"],
                                                 {"model.py": "synthetic-hash"})
    first = json.loads(reports[0].read_text())
    first["compared"]["comparison_addresses"] = ["Premium!B1"]
    reports[0].write_text(json.dumps(first), encoding="utf-8")
    with pytest.raises(ValueError, match="complete formula set"):
        step5_workflow._verify_case_receipts(root, reports, contract, "source", "stage4", "manifest", ["GP"],
                                             {"model.py": "synthetic-hash"})
    first["compared"]["comparison_addresses"] = ["Premium!A1"]
    first["tolerances"]["absolute"] = 0
    reports[0].write_text(json.dumps(first), encoding="utf-8")
    with pytest.raises(ValueError, match="complete formula set"):
        step5_workflow._verify_case_receipts(root, reports, contract, "source", "stage4", "manifest", ["GP"],
                                             {"model.py": "synthetic-hash"})
    first["tolerances"]["absolute"] = 1e-12
    first["exact_case"]["references"] = contract["support_domain"]["cases"][1]
    reports[0].write_text(json.dumps(first), encoding="utf-8")
    with pytest.raises(ValueError, match="complete formula set"):
        step5_workflow._verify_case_receipts(root, reports, contract, "source", "stage4", "manifest", ["GP"],
                                             {"model.py": "synthetic-hash"})
    first["exact_case"]["references"] = contract["support_domain"]["cases"][0]
    oracle_path = Path(first["oracle"]["path"])
    oracle = json.loads(oracle_path.read_text())
    oracle["ranges"]["Premium!A1"]["values"] = [[3]]
    oracle_path.write_text(json.dumps(oracle), encoding="utf-8")
    first["oracle"]["sha256"] = hash_file(oracle_path)
    reports[0].write_text(json.dumps(first), encoding="utf-8")
    with pytest.raises(ValueError, match="do not reconcile"):
        step5_workflow._verify_case_receipts(root, reports, contract, "source", "stage4", "manifest", ["GP"],
                                             {"model.py": "synthetic-hash"})


def test_exact_case_finalization_rejects_changed_package_file_ledger(tmp_path: Path, monkeypatch) -> None:
    contract, reports = _synthetic_receipt_set(tmp_path, monkeypatch)
    root = tmp_path / "workflow"
    with pytest.raises(ValueError, match="package result or exact request"):
        step5_workflow._verify_case_receipts(root, reports, contract, "source", "stage4", "manifest", ["GP"],
                                             {"model.py": "different-package-hash"})


def test_step5_evidence_directories_are_reserved_exclusively(tmp_path: Path) -> None:
    for parent, prefix in ((tmp_path / "stage5" / "execution", "attempt-"),
                           (tmp_path / "stage5", "oracle-attempt-"),
                           (tmp_path / "stage5" / "cases", f"{'a' * 64}-")):
        parent.mkdir(parents=True, exist_ok=True)
        existing = parent / f"{prefix}0001"
        existing.mkdir()
        (existing / "receipt.json").write_text("preserve", encoding="utf-8")
        barrier = Barrier(8)

        def reserve(_index: int) -> Path:
            barrier.wait(timeout=10)
            return step5_workflow._reserve_evidence_dir(parent, prefix)

        with ThreadPoolExecutor(max_workers=8) as pool:
            reserved = list(pool.map(reserve, range(8)))
        assert {path.name for path in reserved} == {f"{prefix}{index:04d}" for index in range(2, 10)}
        assert all(path.is_dir() for path in reserved)
        assert (existing / "receipt.json").read_text(encoding="utf-8") == "preserve"


def test_stage5_reconciles_declared_vector_and_table_values_recursively() -> None:
    specs = {
        "Scalar": {"shape": []},
        "Vector": {"shape": [2], "result_kind": "vector"},
        "Matrix": {"shape": [2, 2], "result_kind": "table"},
        "Singleton": {"shape": [1], "result_kind": "vector"},
        "RowTable": {"shape": [1, 2], "result_kind": "table"},
        "ColumnTable": {"shape": [2, 1], "result_kind": "table"},
        "BadShape": {"shape": [2, 2], "result_kind": "table"},
    }
    model = {
        "Scalar": 5,
        "Vector": [1.0000000001, {"excel_error": "#N/A"}],
        "Matrix": [[1.0, 2.0], [3.0, 4.0]],
        "Singleton": [False],
        "RowTable": [[None, 2]],
        "ColumnTable": [[3], [4]],
        "BadShape": [[1, 2]],
    }
    native = {
        "Scalar": {"value": 5},
        "Vector": {"shape": [2], "value": [1.0, None], "excel_errors": [None, "#N/A"]},
        "Matrix": {"shape": [2, 2], "value": [[1.0, 2.0000000001], [3.0, 4.0]],
                   "excel_errors": [[None, None], [None, None]]},
        "Singleton": {"shape": [1], "value": [False], "excel_errors": [None]},
        "RowTable": {"shape": [1, 2], "value": [[None, 2]], "excel_errors": [[None, None]]},
        "ColumnTable": {"shape": [2, 1], "value": [[3], [4]], "excel_errors": [[None], [None]]},
        "BadShape": {"shape": [2, 2], "value": [[1, 2], [3, 4]],
                     "excel_errors": [[None, None], [None, None]]},
    }

    results, mismatches, compared = _compare_named_targets(
        set(specs), model, native, specs, abs_tol=1e-9, rel_tol=0.0)

    assert compared == len(specs)
    assert {name for name, record in results.items() if record["matched"]} == {
        "Scalar", "Vector", "Matrix", "Singleton", "RowTable", "ColumnTable",
    }
    assert results["Vector"]["excel"] == [1.0, {"excel_error": "#N/A"}]
    assert [item["kind"] for item in mismatches] == ["target_shape"]
    assert _matches([[1.0, 2.0000000001], [3.0, 4.0]], [[1, 2], [3, 4]], 1e-9, 0)
    assert not _matches([[1.0, 2.1]], [[1.0], [2.1]], 0.5, 0)
    assert not _matches([True], [1], 1.0, 0.0)
    assert _matches([None], [None], 0.0, 0.0)
    assert not _matches({"excel_error": "#N/A"}, "#N/A", 0, 0)
    assert _matches({"excel_error": "#N/A"}, {"excel_error": "#N/A"}, 0, 0)
    assert not _matches("", None, 0, 0)
    assert not _matches(0, False, 0, 0)
    error_target, error_mismatches, _ = _compare_named_targets(
        {"Error"}, {"Error": {"excel_error": "#N/A"}},
        {"Error": {"value": None, "excel_error": "#N/A"}}, {"Error": {"shape": []}}, 0, 0)
    assert error_target["Error"]["matched"] and not error_mismatches
    string_target, string_mismatches, _ = _compare_named_targets(
        {"Error"}, {"Error": "#N/A"},
        {"Error": {"value": None, "excel_error": "#N/A"}}, {"Error": {"shape": []}}, 0, 0)
    assert not string_target["Error"]["matched"] and string_mismatches
    assert _value_has_shape([[3], [4]], [2, 1])
    assert not _value_has_shape([[3, 4]], [2, 1])


def test_step5_reads_target_shapes_only_from_the_stage3_pair_bound_by_stage4(tmp_path: Path) -> None:
    root = tmp_path / "workflow"
    source = root / "source.xlsx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"synthetic source")
    create_workflow(root, {"workbook_sha256": hash_file(source), "workbook_path": str(source),
                           "run_id": "synthetic-run"})
    stage3 = append_stage_artifact(
        root, 3, "analysis_design.json",
        {"schema_version": "step3.analysis_design.v1", "design": {
            "targets": [{"selector": "Vector", "result_order": 0, "shape": [2], "result_kind": "vector"},
                        {"selector": "Scalar", "result_order": 1, "shape": [], "result_kind": "scalar"}],
        }},
        "Synthetic design.",
    )
    stage4 = append_stage_artifact(
        root, 4, "generation_report.json", {"schema_version": "step4.generation_report.v1"},
        "Synthetic generation.", input_stages={3: stage3["artifact"]["json_sha256"]},
    )
    _root, manifest = load_workflow(root)

    design = _bound_stage3_design(root, manifest, stage4)

    assert _target_shapes(design) == {"Vector": [2], "Scalar": []}


def test_step5_capture_passes_bound_design_shapes_to_native_oracle(tmp_path: Path, monkeypatch) -> None:
    root = tmp_path / "workflow"
    source = root / "source.xlsx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"synthetic source")
    create_workflow(root, {"workbook_sha256": hash_file(source), "workbook_path": str(source),
                           "run_id": "synthetic-run"})
    stage3 = append_stage_artifact(
        root, 3, "analysis_design.json",
        {"schema_version": "step3.analysis_design.v1", "design": {
            "targets": [{"selector": "Vector", "result_order": 0, "shape": [2]}],
        }},
        "Synthetic design.",
    )
    stage4 = append_stage_artifact(
        root, 4, "generation_report.json", {"schema_version": "step4.generation_report.v1"},
        "Synthetic generation.", input_stages={3: stage3["artifact"]["json_sha256"]},
    )
    _root, manifest = load_workflow(root)
    captured = {}

    def fake_capture(_source, output_dir, *, targets, ranges, target_shapes):
        captured.update(targets=targets, ranges=ranges, target_shapes=target_shapes)
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "excel_oracle.json").write_text(json.dumps({"status": "pass"}), encoding="utf-8")
        return {"status": "pass", "tool": "step5.oracle"}

    monkeypatch.setattr(step5_workflow, "_stage4", lambda _root: (
        manifest, stage4, {"target_names": ["Vector"], "oracle_ranges": []}, root / "stage4" / "bundle"))
    monkeypatch.setattr(step5_workflow, "capture_excel_oracle", fake_capture)

    result = step5_workflow.capture_oracle(root)

    assert result["status"] == "pass", result
    assert captured == {"targets": ["Vector"], "ranges": [], "target_shapes": {"Vector": [2]}}


def _canonical_hash(value: object) -> str:
    return hashlib.sha256((json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")) + "\n")
                           .encode("utf-8")).hexdigest()


def test_stage5_compares_every_active_formula_across_ranges_and_preserves_errors() -> None:
    oracle = {"ranges": {
        "Premium!C1:D2": {
            "qualified_address": "Premium!C1:D2",
            "values": [[2.5, None], [None, None]],
            "formulas": [["=1+1.5", ""], ["", "=#N/A"]],
            "excel_errors": [[None, None], [None, "#N/A"]],
        },
        "Main!F34:F34": {
            "qualified_address": "Main!F34",
            "values": [[7]], "formulas": [["=3+4"]], "excel_errors": [[None]],
        },
    }}
    native, ranges = _oracle_range_cells(oracle)
    assert ranges == ["Main!F34", "Premium!C1:D2"]
    addresses = ["Premium!C1", "Premium!D2", "Main!F34"]
    model = {"premium!C1": 2.5000000000001, "premium!D2": {"excel_error": "#N/A"}, "main!F34": 7}
    expected_formulas = {"premium!C1": "=1+1.5", "premium!D2": "=#N/A",
                         "main!F34": "=SUM(3,\n 4)"}
    native["main!F34"]["formula"] = "=SUM(3,4)"

    compared, identity_matches, max_abs, matched_addresses, mismatches = _compare_active_formula_members(
        addresses, native, model, expected_formulas, 1e-12, 1e-12)
    assert compared == 3
    assert identity_matches == 3
    assert max_abs is not None and 0 < max_abs < 1e-12
    assert matched_addresses == {"premium!C1", "premium!D2", "main!F34"}
    assert mismatches == []

    changed_formula = dict(native)
    changed_formula["main!F34"] = {**native["main!F34"], "formula": "=SUM(2,5)"}
    compared, identity_matches, _max_abs, _matched, mismatches = _compare_active_formula_members(
        addresses, changed_formula, model, expected_formulas, 1e-12, 1e-12)
    assert compared == 3
    assert identity_matches == 2
    assert any(item["kind"] == "active_formula_identity" and item["address"] == "Main!F34"
               for item in mismatches)
    assert not any(item["kind"] == "formula_value" and item["address"] == "Main!F34"
                   for item in mismatches)

    missing_formula = dict(native)
    missing_formula["main!F34"] = {**native["main!F34"], "formula": ""}
    compared, identity_matches, _max_abs, _matched, mismatches = _compare_active_formula_members(
        addresses, missing_formula, model, expected_formulas, 1e-12, 1e-12)
    assert compared == 3
    assert identity_matches == 2
    assert any(item["kind"] == "active_formula_identity" and item["address"] == "Main!F34"
               for item in mismatches)

    compared, _identity_matches, _max_abs, _matched_addresses, mismatches = _compare_active_formula_members(
        addresses, native, {"premium!C1": 2.5}, expected_formulas, 1e-12, 1e-12)
    assert compared == 1
    assert [item["kind"] for item in mismatches] == [
        "active_formula_missing_from_model", "active_formula_missing_from_model"]


def test_step5_oracle_requires_completed_native_calculation_state() -> None:
    oracle = {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "pass",
              "engine": "Microsoft Excel", "source_sha256": "source", "source_copy_sha256": "source",
              "stage4_artifact_sha256": "stage4", "macros_executed": False, "input_overrides": [],
              "primary_inputs": {address: {"value": 1} for address in
                                 ("C3", "C4", "C6", "C7", "C8", "J48", "O26")}}
    _preflight_oracle({**oracle, "calculation_state": 0}, "source", "stage4")
    varied = {**oracle, "calculation_state": 0, "primary_inputs": {},
              "input_overrides": [{"qualified_address": "Main!C3", "value": 7}],
              "applied_inputs": {"Main!C3": {"value": 7}},
              "input_sha256": "a" * 64, "scenario_id": "a" * 64}
    _preflight_oracle(varied, "source", "stage4", "a" * 64)
    try:
        _preflight_oracle(varied, "source", "stage4", "b" * 64)
    except ValueError as exc:
        assert "exact input" in str(exc)
    else:
        raise AssertionError("oracle for input A must not serve input B")
    for state in (None, 1, 2):
        try:
            _preflight_oracle({**oracle, "calculation_state": state}, "source", "stage4")
        except ValueError as exc:
            assert "calculation must be complete" in str(exc)
        else:
            raise AssertionError(f"native calculation state {state!r} must block reconciliation")


def test_stage5_active_member_set_comes_from_the_bound_trace(tmp_path: Path) -> None:
    trace_path = tmp_path / "active_trace.json"
    trace = {"source_sha256": "source-hash", "cells": [
        {"address": "Premium!J1", "role": "calculated_formula", "formula": "=1+1"},
        {"address": "CI_MaleChoosen!F9", "role": "calculated_array_formula", "formula": "=A1+1"},
        {"address": "Main!C3", "role": "source_value"},
    ]}
    trace_path.write_text(json.dumps(trace), encoding="utf-8")
    report = {"source_sha256": "source-hash",
              "active_trace": {"path": str(trace_path), "sha256": hash_file(trace_path)}}
    manifest = {"active_formula_addresses": ["Premium!J1", "CI_MaleChoosen!F9"]}

    assert _active_formula_addresses(report, manifest) == ["Premium!J1", "CI_MaleChoosen!F9"]
    stale = {**manifest, "active_formula_addresses": ["Premium!J1"]}
    try:
        _active_formula_addresses(report, stale)
    except ValueError as exc:
        assert "differ from the bound trace" in str(exc)
    else:
        raise AssertionError("a manifest that omits an active array follower must be rejected")


def test_legacy_formula_addresses_are_bound_to_formula_mode_source(tmp_path: Path) -> None:
    source_path = tmp_path / "legacy_source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Main"
    workbook.active["A1"] = "=1+1"
    workbook.save(source_path)
    workbook.close()
    source_hash = hash_file(source_path)
    addresses, formulas = _active_formula_details(
        {}, {"active_formula_addresses": ["Main!A1"]}, source_path, source_hash)
    assert addresses == ["Main!A1"]
    assert formulas == {"main!A1": "=1+1"}

    try:
        _active_formula_details({}, {"active_formula_addresses": ["Main!A1"]}, source_path, "stale")
    except ValueError as exc:
        assert "differs from the bound workbook" in str(exc)
    else:
        raise AssertionError("a legacy formula address set must not detach from its source workbook")


def test_overlapping_native_ranges_must_agree_on_values_and_errors() -> None:
    base = {"qualified_address": "Premium!C1", "values": [[None]],
            "formulas": [["=#N/A"]], "excel_errors": [["#N/A"]]}
    oracle = {"ranges": {"Premium!C1": base, "Premium!C1:C1": dict(base)}}
    cells, _ranges = _oracle_range_cells(oracle)
    assert cells["premium!C1"]["value"] == {"excel_error": "#N/A"}

    conflicting = {"ranges": {"Premium!C1": base,
                               "Premium!C1:C1": {**base, "excel_errors": [[None]]}}}
    try:
        _oracle_range_cells(conflicting)
    except ValueError as exc:
        assert "overlapping native oracle ranges disagree" in str(exc)
    else:
        raise AssertionError("overlapping native ranges with different error semantics must be rejected")


def test_stage5_compares_producer_shaped_external_cuts_and_literal_native_axis(tmp_path: Path) -> None:
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    values = {"external.rate": [0.1, 0.2]}
    age_values = [0, 1]
    formulas = ["=A1*0+0.1", "=A2*0+0.2"]
    source_path = tmp_path / "source.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Rates"
    worksheet["A1"], worksheet["A2"] = age_values
    worksheet["B1"], worksheet["B2"] = formulas
    workbook.save(source_path)
    workbook.close()
    source_hash = hash_file(source_path)
    # This matches the Step 4 producer: source identity is bound at payload level;
    # each cut member carries its own exact formula text and formula hash.
    members = [{"address": f"Rates!B{row}", "boundary_variable_id": "external.rate",
                "formula": formula,
                "formula_sha256": hashlib.sha256(formula.encode("utf-8")).hexdigest()}
               for row, formula in zip((1, 2), formulas)]
    values_hash, age_hash = _canonical_hash(values), _canonical_hash(age_values)
    external_payload = {
        "schema_version": "step4.external_values.v1", "source_sha256": source_hash,
        "capture_sha256": "capture-hash", "values_sha256": values_hash, "values": values,
        "boundaries": [{"boundary_variable_id": "external.rate", "source_range": "Rates!B1:B2",
                        "shape": [2]}],
        "formula_cut_members": members, "formula_cut_member_count": 2,
        "formula_cut_formula_set_sha256": "cut-set-hash",
        # The native producer does not duplicate the generated metadata-group binding.
        "age_axis": {"source_range": "Rates!A1:A2", "values": age_values, "value_sha256": age_hash},
    }
    (bundle / "external_values.json").write_text(json.dumps(external_payload), encoding="utf-8")
    (bundle / "input_layout.json").write_text(json.dumps({
        "external_bindings": [{"variable_id": "external.rate"}],
        "external_capture": {"age_axis": {"source_range": "Rates!A1:A2", "metadata_group": "axis.age"}},
    }), encoding="utf-8")
    (bundle / "adapter_metadata.json").write_text(json.dumps({"groups": {
        "axis.age": {"addresses": ["rates!A1", "rates!A2"], "values": age_values},
    }}), encoding="utf-8")
    manifest = {
        "source_sha256": source_hash, "external_bindings": [{"variable_id": "external.rate"}],
        "external_input_addresses": ["rates!B1", "rates!B2"], "external_input_count": 2,
        "external_capture_sha256": "capture-hash", "external_values_sha256": values_hash,
        "external_formula_cut_formula_set_sha256": "cut-set-hash",
        "external_formula_cut_member_count": 2, "external_age_axis_sha256": age_hash,
        "external_age_axis": {"source_range": "Rates!A1:A2", "values_sha256": age_hash,
                              "metadata_group": "axis.age"},
    }
    native = {
        # Excel COM returns literal constants through Range.Formula, e.g. "0"/"1".
        "rates!A1": {"address": "Rates!A1", "value": 0, "formula": "0"},
        "rates!A2": {"address": "Rates!A2", "value": 1, "formula": "1"},
        "rates!B1": {"address": "Rates!B1", "value": 0.1, "formula": formulas[0]},
        "rates!B2": {"address": "Rates!B2", "value": 0.2, "formula": formulas[1]},
    }
    model = {"external_inputs": values, "external_age_axis": age_values}
    metrics, mismatches = _compare_external_boundary(
        manifest, bundle, model, native, 1e-12, 1e-12, source_path)
    assert metrics == {"external_input_coordinates": 2, "external_age_axis_keys": 2,
                       "external_formula_cut_members": 2}
    assert mismatches == []

    wrong_source_manifest = {**manifest, "source_sha256": "b" * 64}
    _metrics, mismatches = _compare_external_boundary(
        wrong_source_manifest, bundle, model, native, 1e-12, 1e-12, source_path)
    assert any(item["kind"] == "external_boundary_binding" for item in mismatches)

    layout_path = bundle / "input_layout.json"
    layout_record = json.loads(layout_path.read_text(encoding="utf-8"))
    wrong_layout = json.loads(json.dumps(layout_record))
    wrong_layout["external_capture"]["age_axis"]["metadata_group"] = "axis.wrong"
    layout_path.write_text(json.dumps(wrong_layout), encoding="utf-8")
    _metrics, mismatches = _compare_external_boundary(
        manifest, bundle, model, native, 1e-12, 1e-12, source_path)
    assert any(item["kind"] == "external_age_axis_binding" for item in mismatches)
    layout_path.write_text(json.dumps(layout_record), encoding="utf-8")

    metadata_path = bundle / "adapter_metadata.json"
    metadata_record = json.loads(metadata_path.read_text(encoding="utf-8"))
    metadata_record["groups"]["axis.age"]["addresses"].reverse()
    metadata_path.write_text(json.dumps(metadata_record), encoding="utf-8")
    _metrics, mismatches = _compare_external_boundary(
        manifest, bundle, model, native, 1e-12, 1e-12, source_path)
    assert any(item["kind"] == "external_age_axis_bundle_value" for item in mismatches)
    metadata_record["groups"]["axis.age"]["addresses"].reverse()
    metadata_path.write_text(json.dumps(metadata_record), encoding="utf-8")

    changed_native = dict(native)
    changed_native["rates!B2"] = {**native["rates!B2"], "formula": "=A2/200"}
    metrics, mismatches = _compare_external_boundary(
        manifest, bundle, model, changed_native, 1e-12, 1e-12, source_path)
    assert any(item["kind"] == "external_formula_cut_changed" for item in mismatches)

    # A value in the source range that is actually a formula remains invalid, even
    # when Excel's calculated value matches the captured age key.
    formula_source_path = tmp_path / "formula_source.xlsx"
    formula_workbook = Workbook()
    formula_sheet = formula_workbook.active
    formula_sheet.title = "Rates"
    formula_sheet["A1"], formula_sheet["A2"] = 0, "=1"
    formula_sheet["B1"], formula_sheet["B2"] = formulas
    formula_workbook.save(formula_source_path)
    formula_workbook.close()
    changed_source_hash = hash_file(formula_source_path)
    external_payload["source_sha256"] = changed_source_hash
    (bundle / "external_values.json").write_text(
        json.dumps(external_payload), encoding="utf-8")
    manifest["source_sha256"] = changed_source_hash
    changed_native["rates!A2"] = {**native["rates!A2"], "formula": "=1"}
    _metrics, mismatches = _compare_external_boundary(
        manifest, bundle, model, changed_native, 1e-12, 1e-12, formula_source_path)
    assert any(item["kind"] == "external_age_axis_source_formula" for item in mismatches)


def test_step5_report_states_all_member_coverage_and_delegated_reviewer_pair() -> None:
    payload = {"status": "ready_for_review", "source_sha256": "source", "stage4_artifact_sha256": "stage4",
               "oracle": {"engine": "Microsoft Excel", "source_copy_sha256": "source",
                          "macros_executed": False, "input_overrides": []},
               "compared": {"primary_inputs": 7, "named_targets": 4, "required_targets": 4,
                            "active_formula_members": 23098, "required_active_formula_members": 23098,
                            "active_formula_identity_matches": 23098,
                            "required_active_formula_identities": 23098,
                            "active_formula_max_abs_difference": 2.22e-16,
                            "active_premium_formula_cells": 7767},
               "tolerances": {"absolute": 1e-12, "relative": 1e-12},
               "mismatch_count": 0, "mismatch_examples": [], "target_results": {},
               "path_coverage": {"implemented_path_ids": [], "numerical_path_ids": [],
                                 "runtime_verified_path_ids": [], "condition_only_path_ids": [],
                                 "runtime_evidence": [], "condition_evidence": []},
               "review_policy": {"required_reviewer_pair": ["agent", "typesafe"],
                                 "delegation": {"stages": [3, 4, 5, 6],
                                                "requested_model": "jev-latest",
                                                "authorization_sha256": "auth-sha"}}}

    markdown = _reconciliation_markdown(payload)
    assert "Active source formula members compared: 23098 / 23098" in markdown
    assert "Active source formula identities matched: 23098 / 23098" in markdown
    assert "Maximum absolute difference across numeric active formula values: 2.22e-16" in markdown
    assert "Agent and delegated TypeSafe confirmation are pending." in markdown
    assert "Agent and human confirmation are pending." not in markdown
    assert "TypeSafe delegation: stages [3, 4, 5, 6]" in markdown


def test_step5_reader_front_orders_requested_result_before_diagnostics_and_shows_tolerances() -> None:
    from excel_to_act.steps.step5.workflow import _reconciliation_markdown

    payload = {
        "status": "ready_for_review",
        "source_sha256": "source-hash",
        "stage4_artifact_sha256": "stage4-hash",
        "oracle": {"engine": "Fixture", "source_copy_sha256": "source-hash",
                   "macros_executed": False, "input_overrides": []},
        "compared": {"primary_inputs": 7, "named_targets": 2, "required_targets": 2,
                      "active_formula_members": 4, "required_active_formula_members": 4,
                      "active_formula_identity_matches": 4, "required_active_formula_identities": 4,
                      "active_formula_max_abs_difference": 0.0, "active_premium_formula_cells": 4,
                      "external_input_coordinates": 0, "required_external_input_coordinates": 0,
                      "external_formula_cut_members": 0, "required_external_formula_cut_members": 0},
        "mismatch_count": 0,
        "mismatch_examples": [],
        "tolerances": {"absolute": 1e-12, "relative": 2e-12},
        "target_results": {
            "PVFB": {"model": 14.7, "excel": 14.7, "matched": True},
            "GP": {"model": 2.7, "excel": 2.8, "matched": False},
        },
        "path_coverage": {"numerical_path_ids": [], "runtime_verified_path_ids": [],
                          "condition_only_path_ids": [], "condition_evidence": [],
                          "runtime_evidence": []},
        "review_policy": {"required_reviewer_pair": ["agent", "human"]},
    }
    design = {"targets": [{"selector": "GP", "result_order": 0,
                            "diagnostic_intermediates": ["PVFB"]}]}

    front = _reconciliation_markdown(payload, design=design).split("## Detailed comparison evidence", 1)[0]

    assert "| GP | Requested target | Not recorded | Not recorded | Not recorded | Not recorded | 2.7 | 2.8 | False |" in front
    assert "| PVFB | Diagnostic | Not recorded | Not recorded | Not recorded | Not recorded | 14.7 | 14.7 | True |" in front
    assert front.index("| GP | Requested target") < front.index("| PVFB | Diagnostic")
    assert "absolute=1e-12, relative=2e-12" in front



def test_step5_reader_shows_declared_shapes_ordered_array_values_and_unknowns_without_mutation() -> None:
    import copy

    payload = {
        "status": "ready_for_review", "source_sha256": "source-fixture",
        "stage4_artifact_sha256": "stage4-fixture", "mismatch_count": None,
        "oracle": {"engine": "fixture", "source_copy_sha256": "copy-fixture",
                   "macros_executed": False, "input_overrides": []},
        "compared": {"primary_inputs": 0, "named_targets": 3, "required_targets": 3,
                     "active_formula_members": 0, "required_active_formula_members": 0,
                     "active_formula_identity_matches": False, "required_active_formula_identities": None,
                     "active_formula_max_abs_difference": None, "active_premium_formula_cells": 0,
                     "external_input_coordinates": 0, "required_external_input_coordinates": 0,
                     "external_formula_cut_members": 0, "required_external_formula_cut_members": 0,
                     "external_age_axis_keys": 0, "required_external_age_axis_keys": 0},
        "tolerances": {"absolute": None, "relative": 0}, "mismatch_examples": [],
        "target_results": {
            "Diagnostic": {"model": [5, 6], "excel": [5, 6], "matched": True},
            "Matrix": {"model": [[1, 2], [3, 4]], "excel": [[1, 2], [3, 0]], "matched": False},
            "Scalar": {"model": False, "excel": False, "matched": None},
            "Vector": {"model": [0, False, None], "excel": [0, False, 8], "matched": True},
        },
        "path_coverage": {"numerical_path_ids": [], "runtime_verified_path_ids": [],
                          "condition_only_path_ids": [], "condition_evidence": []},
        "review_policy": {"required_reviewer_pair": ["agent", "human"]},
    }
    original = copy.deepcopy(payload)
    design = {
        "targets": [
            {"selector": "Scalar", "result_order": 2, "shape": [], "result_kind": None, "units": None},
            {"selector": "Vector", "result_order": 0, "shape": [3], "result_kind": "vector", "units": "USD",
             "axes": [{"name": "duration", "role": "projection", "keys": ["None", 0, False]}]},
            {"selector": "Matrix", "result_order": 1, "shape": [2, 2], "result_kind": "table", "units": "count",
             "axes": []},
            {"selector": "Other", "result_order": 3, "shape": None, "diagnostic_intermediates": ["Diagnostic"]},
        ],
        "open_questions": [
            {"question_id": "Q_CAPTURE", "question": "A capture was required before execution.",
             "answer": "Capture is recorded in the current run.", "status": "captured", "disposition": "execution prerequisite"},
            {"question_id": "Q_MEANING", "question": "What does the selected axis mean?",
             "answer": None, "status": None, "disposition": None},
        ],
    }
    front = _reconciliation_markdown(payload, design=design, machine_link="validation.json").split(
        "## Detailed comparison evidence", 1)[0]

    assert front.index("| Vector | Requested target") < front.index("| Matrix | Requested target") < front.index("| Scalar | Requested target")
    assert front.index("| Scalar | Requested target") < front.index("| Diagnostic | Diagnostic")
    assert "| Vector | Requested target | vector | [3] | USD |" in front
    assert "keys: position [0] = None" in front and "position [2] = False" in front
    assert "position [0] = 0" in front and "position [2] = Not recorded" in front
    assert "position [2] = 8" in front and "| True |" in front
    assert "| Scalar | Requested target | Not recorded | [] | Not recorded | Not recorded | False | False | Not recorded |" in front
    assert "position [0, 0] = 1" in front and "position [1, 1] = 4" in front
    assert "position [1, 0] = 3" in front and "position [1, 1] = 0" in front
    assert "Mismatch count: Not recorded" in front
    assert "Primary inputs matched: 0 / 7" in front and "absolute=Not recorded, relative=0" in front
    assert "Recorded answer: Capture is recorded in the current run." in front
    assert "execution prerequisite" in front and "**Q_MEANING**" in front
    assert "status: Not recorded; disposition: Not recorded" in front
    assert "Position labels do not imply a quarter, year, or other business-axis origin" in front
    assert "[validation_report.json](<validation.json>)" in front
    assert "[[1, 2], [3, 4]]" not in front and "[0, False, None]" not in front
    assert payload == original

def test_step5_accepts_modular_runtime_but_keeps_legacy_runtime_contract(tmp_path: Path) -> None:
    bundle = tmp_path / "bundle"
    bundle.mkdir()
    (bundle / "model.py").write_text("def main(): pass\n", encoding="utf-8")
    (bundle / "source_values.json").write_text("{}\n", encoding="utf-8")
    (bundle / "modular_runtime.py").write_text("VALUE = 1\n", encoding="utf-8")
    (bundle / "input_layout.json").write_text('{"external_bindings": []}\n', encoding="utf-8")
    (bundle / "input_adapter.py").write_text(
        "import hashlib\n"
        "def content_hash(payload):\n"
        "    return hashlib.sha256(payload).hexdigest()\n",
        encoding="utf-8",
    )
    modular_files = {path.name: hash_file(path) for path in bundle.iterdir()}
    modular_manifest = {
        "schema_version": "step4.modular_model.v1", "execution_kind": "modular",
        "files": modular_files, "raw_input_addresses": [], "formula_addresses": [],
        "formula_cache_inputs": False,
    }
    assert _verify_python(bundle, modular_manifest) == []

    legacy_manifest = {**modular_manifest, "schema_version": "step4.generated_model.v1",
                       "execution_kind": "legacy"}
    assert "generated bundle is missing runtime.py" in _verify_python(bundle, legacy_manifest)


def test_step5_validates_separate_external_ledger_and_disallows_overlap(tmp_path: Path) -> None:
    bundle = tmp_path / "external-bundle"
    bundle.mkdir()
    values = {"rates.external": [0.1, 0.2]}
    values_sha = _canonical_hash(values)
    layout = {
        "external_bindings": [{"variable_id": "rates.external", "source_range": "Rates!B1:B2",
                               "shape": [2], "addresses": ["rates!B1", "rates!B2"],
                               "indices": [[0], [1]]}],
        "external_capture": {"artifact_sha256": "artifact-sha", "capture_sha256": "capture-sha",
                             "values_sha256": values_sha},
    }
    external = {"schema_version": "step4.external_values.v1", "source_sha256": "source-sha",
                "capture_artifact_sha256": "artifact-sha", "capture_sha256": "capture-sha",
                "values_sha256": values_sha, "values": values}
    contents = {"model.py": "def main(): pass\n", "modular_runtime.py": "VALUE = 1\n",
                "source_values.json": "{}\n", "input_layout.json": json.dumps(layout),
                "external_values.json": json.dumps(external)}
    for name, content in contents.items():
        (bundle / name).write_text(content, encoding="utf-8")
    manifest = {
        "schema_version": "step4.modular_model.v1", "execution_kind": "modular",
        "source_sha256": "source-sha", "files": {name: hash_file(bundle / name) for name in contents},
        "raw_input_addresses": [], "formula_addresses": [],
        "external_input_addresses": ["Rates!B1", "Rates!B2"], "external_input_count": 2,
        "external_capture_artifact_sha256": "artifact-sha", "external_capture_sha256": "capture-sha",
        "external_values_sha256": values_sha, "formula_cache_inputs": False,
    }
    assert _verify_python(bundle, manifest) == []

    overlap = {**manifest, "formula_addresses": ["rates!B2"]}
    assert any("overlaps a raw or calculated formula address" in failure
               for failure in _verify_python(bundle, overlap))
    stale = {**manifest, "external_capture_artifact_sha256": "stale-artifact"}
    assert any("external_values.json is stale" in failure for failure in _verify_python(bundle, stale))


def test_step5_reconciles_synthetic_generation_and_separates_extra_capture(tmp_path: Path, monkeypatch) -> None:
    root = tmp_path / "workflow"
    source = root / "source" / "fixture.xlsm"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"synthetic source identity; not loaded by standalone validation")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    create_workflow(root, {"workbook_sha256": source_hash, "workbook_path": str(source),
                           "run_id": "synthetic-run", "source_id": "synthetic-source"})
    _, workflow = load_workflow(root)
    workflow["mode"] = "synthetic_fixture"
    _write_manifest(root, workflow)

    previous = None
    for stage in (1, 2, 3):
        stage_payload = {"status": "ready_for_review"}
        if stage == 3:
            stage_payload["design"] = {"targets": [
                {"selector": "GP", "result_order": 0, "shape": [], "result_kind": "scalar"},
            ]}
        revision = append_stage_artifact(
            root, stage, f"stage{stage}.json", stage_payload,
            f"Synthetic stage {stage}.",
            input_stages={stage - 1: previous["artifact"]["json_sha256"]} if previous else None)
        for reviewer in ("agent", "human"):
            record_decision(root, stage, reviewer, "approve", f"Synthetic stage {stage} review.")
        previous = revision

    bundle = root / "stage4" / "bundles" / "revision-0001" / "bundle"
    bundle.mkdir(parents=True)
    contract = {"mutable": {"X": {"bindings": [{"address": "main!C3", "indices": []}]}},
                "support_domain": {"external_policy": {}, "tolerances": {"absolute": 1e-12, "relative": 1e-12}}}
    files = {
        "model.py": (
            "import argparse, hashlib, json\n"
            "from pathlib import Path\n"
            "def main():\n"
            "    parser = argparse.ArgumentParser()\n"
            "    parser.add_argument('--out', type=Path, default=Path('model_result.json'))\n"
            "    parser.add_argument('--inputs', type=Path)\n"
            "    args = parser.parse_args()\n"
            "    request = json.loads(args.inputs.read_text()) if args.inputs else None\n"
            "    value = request['inputs']['X']['values'] if request else 1\n"
            "    effective = request['inputs'] if request else None\n"
            "    digest = hashlib.sha256((json.dumps(effective,sort_keys=True,separators=(',',':'))+'\\n').encode()).hexdigest() if request else None\n"
            "    args.out.write_text(json.dumps({'status':'pass','targets':{'GP':value},"
            "'cells':{'Main!A1':2},'cell_count':1,'formula_count':1,"
            "'execution_mode':'variable_input' if request else 'saved_scenario',"
            "'input_sha256':digest,'raw_inputs':{'main!C3':value} if request else {},"
            "'external_inputs':{},'target_axes':{'GP':[]},"
            "'dependency_sha256':'dependency'}), encoding='utf-8')\n"
            "if __name__ == '__main__': main()\n"
        ),
        "modular_runtime.py": "VALUE = 1\n",
        "input_layout.json": json.dumps({"external_bindings": [], "external_capture": None,
                                           "variable_contract": contract}) + "\n",
        "source_values.json": json.dumps({f"main!{address}": 1 for address in
                                           ("C3", "C4", "C6", "C7", "C8", "J48", "O26")}) + "\n",
    }
    for name, content in files.items():
        (bundle / name).write_text(content, encoding="utf-8")
    bundle_files = {name: hash_file(bundle / name) for name in files}
    manifest = {
        "schema_version": "step4.modular_model.v1", "execution_kind": "modular",
        "source_sha256": source_hash, "files": bundle_files,
        "formula_count": 1, "array_member_count": 0, "array_instance_count": 0,
        "formula_addresses": ["Main!A1"],
        "raw_input_addresses": ["main!C3", "main!C4", "main!C6", "main!C7", "main!C8", "main!J48", "main!O26"],
        "active_formula_addresses": ["Main!A1"], "formula_cache_inputs": False,
    }
    (bundle / "model_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    report_files = {**bundle_files, "model_manifest.json": hash_file(bundle / "model_manifest.json")}
    trace_path = root / "stage4" / "active_trace.json"
    trace_path.parent.mkdir(parents=True, exist_ok=True)
    trace_path.write_text(json.dumps({"source_sha256": source_hash, "cells": [
        {"address": "Main!A1", "role": "calculated_formula", "formula": "=1+1"},
    ]}), encoding="utf-8")
    append_stage_artifact(
        root, 4, "generation_report.json",
        {"schema_version": "step4.generation_report.v1", "status": "ready_for_review",
         "source_sha256": source_hash, "target_names": ["GP"], "oracle_ranges": ["Main!A1"],
         "active_trace": {"path": str(trace_path), "sha256": hash_file(trace_path)},
         "bundle": {"path": str(bundle), "files": report_files}},
        "Synthetic modular generation.", input_stages={3: previous["artifact"]["json_sha256"]})
    for reviewer in ("agent", "human"):
        record_decision(root, 4, reviewer, "approve", "Synthetic Stage 4 review.")

    result = validate_generated(root)
    assert result["status"] == "pass", result
    evidence = json.loads(Path(result["validation"]).read_text(encoding="utf-8"))
    assert evidence["status"] == "pass"
    assert evidence["result_summary"]["targets"] == {"GP": 1}

    current_root, current_workflow = load_workflow(root)
    stage4_hash = current_workflow["stages"]["4"]["revisions"][-1]["artifact"]["json_sha256"]
    stage5_root = current_root / "stage5"
    incomplete_oracle_path = stage5_root / "oracle-incomplete.json"
    incomplete_oracle_path.write_text(json.dumps({
        "schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "pass",
        "engine": "Microsoft Excel", "source_sha256": source_hash, "source_copy_sha256": source_hash,
        "stage4_artifact_sha256": stage4_hash, "macros_executed": False, "input_overrides": [],
        "primary_inputs": {address: {"value": 1} for address in
                           ("C3", "C4", "C6", "C7", "C8", "J48", "O26")},
        # calculation_state intentionally omitted
    }), encoding="utf-8")
    revisions_before = len(current_workflow["stages"]["5"]["revisions"])
    blocked = reconcile(root, Path(result["validation"]), incomplete_oracle_path)
    assert blocked["status"] == "blocked"
    assert "calculation must be complete" in blocked["reason"]
    _current_root, after_workflow = load_workflow(root)
    assert len(after_workflow["stages"]["5"]["revisions"]) == revisions_before

    def fake_capture(_source, output_dir, *, targets, ranges, target_shapes):
        assert targets == ["GP", "ExtraTarget"]
        assert ranges == ["Main!A1", "Main!B1"]
        assert target_shapes == {"GP": []}
        range_records = {}
        for selector, value, formula in (("Main!A1", 2, "=1+1"), ("Main!B1", 42, "=40+2")):
            range_records[selector] = {
                "qualified_address": selector, "values": [[value]], "formulas": [[formula]],
                "excel_errors": [[None]],
            }
        native = {
            "schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "pass",
            "engine": "Microsoft Excel", "calculation_state": 0,
            "source_sha256": source_hash, "source_copy_sha256": source_hash,
            "macros_executed": False, "input_overrides": [],
            "primary_inputs": {address: {"value": 1} for address in
                               ("C3", "C4", "C6", "C7", "C8", "J48", "O26")},
            "named_targets": {"GP": {"value": 1}, "ExtraTarget": {"value": 999}},
            "ranges": range_records,
        }
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "excel_oracle.json").write_text(json.dumps(native), encoding="utf-8")
        return {"status": "pass", "tool": "step5.oracle"}

    monkeypatch.setattr(step5_workflow, "capture_excel_oracle", fake_capture)
    captured = step5_workflow.capture_oracle(
        root, targets=["GP", "ExtraTarget"], ranges=["Main!A1", "Main!B1"])
    assert captured["status"] == "pass", captured
    assert captured["stage_advanced"] is False

    reconciled = reconcile(root, Path(result["validation"]), Path(captured["artifact"]))
    assert reconciled["status"] == "pass", reconciled
    report_path = root / reconciled["artifact"]["json"]
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["status"] == "ready_for_review"
    assert report["compared"]["named_targets"] == 1
    assert report["compared"]["active_formula_members"] == 1
    assert report["compared"]["comparison_addresses"] == ["Main!A1"]
    assert "comparison_ranges" not in report["compared"]
    assert report["captured"]["range_selectors"] == ["Main!A1", "Main!B1"]
    coverage = {item["selector"]: item for item in report["captured"]["active_formula_range_coverage"]}
    assert coverage["Main!A1"]["status"] == "contains_compared_active_formula_addresses"
    assert coverage["Main!A1"]["compared_active_formula_address_count"] == 1
    assert coverage["Main!B1"]["status"] == "captured_only_for_active_formula_comparison"
    assert coverage["Main!B1"]["compared_active_formula_address_count"] == 0
    markdown = (root / reconciled["artifact"]["md"]).read_text(encoding="utf-8")
    assert "Exact active formula addresses compared: 1 / 1" in markdown
    assert "Captured selector `Main!B1`: captured-only for active-formula comparison; no active formula address from this selector was compared." in markdown
    assert "The report lists named targets, primary inputs, external values, formula cuts, and age keys separately." in markdown
    assert "ExtraTarget" not in report["target_results"]

    request_a = root / "request-a.json"
    request_a.write_text(json.dumps({"inputs": {"X": {"values": 7, "axes": []}},
                                     "external_data": {}}), encoding="utf-8")
    varied_a = validate_generated(root, request_a)
    assert varied_a["status"] == "pass" and varied_a["input_sha256"]

    def varied_capture(_source, output_dir, *, targets, ranges, target_shapes,
                       input_overrides, input_sha256, scenario_id):
        assert input_overrides == [{"address": "main!C3", "value": 7}]
        assert scenario_id == input_sha256
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "excel_oracle.json").write_text(json.dumps({
            "schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "pass",
            "engine": "Microsoft Excel", "calculation_state": 0,
            "source_sha256": source_hash, "source_copy_sha256": source_hash,
            "macros_executed": False, "input_overrides": [{"qualified_address": "Main!C3", "value": 7}],
            "input_sha256": input_sha256, "scenario_id": scenario_id,
            "applied_inputs": {"Main!C3": {"value": 7, "excel_error": None}},
            "named_targets": {"GP": {"value": 7}},
            "ranges": {"Main!A1": {"qualified_address": "Main!A1", "values": [[2]],
                                    "formulas": [["=1+1"]], "excel_errors": [[None]]}},
        }), encoding="utf-8")
        return {"status": "pass", "tool": "step5.oracle"}

    monkeypatch.setattr(step5_workflow, "capture_excel_oracle", varied_capture)
    varied_oracle = step5_workflow.capture_oracle(root, validation_path=Path(varied_a["validation"]))
    assert varied_oracle["status"] == "pass", varied_oracle
    revisions_before = len(load_workflow(root)[1]["stages"]["5"]["revisions"])
    varied_case = reconcile(root, Path(varied_a["validation"]), Path(varied_oracle["artifact"]))
    assert varied_case["status"] == "pass" and varied_case["native_verified_exact_input"] is True, varied_case
    assert varied_case["input_sha256"] == varied_a["input_sha256"]
    case_markdown = Path(varied_case["artifact"]).with_suffix(".md").read_text(encoding="utf-8")
    assert "Primary inputs matched: 1 / 1" in case_markdown
    assert "[case_report.json](<case_report.json>)" in case_markdown
    assert len(load_workflow(root)[1]["stages"]["5"]["revisions"]) == revisions_before
    request_b = root / "request-b.json"
    request_b.write_text(json.dumps({"inputs": {"X": {"values": 8, "axes": []}},
                                     "external_data": {}}), encoding="utf-8")
    varied_b = validate_generated(root, request_b)
    assert varied_b["status"] == "pass" and varied_b["input_sha256"] != varied_a["input_sha256"]
    cross_input = reconcile(root, Path(varied_b["validation"]), Path(varied_oracle["artifact"]))
    assert cross_input["status"] == "blocked" and "exact input" in cross_input["reason"]
