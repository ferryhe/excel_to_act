from __future__ import annotations

import hashlib
import json
import struct
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils.exceptions import InvalidFileException
from typer.testing import CliRunner

from excel_to_act.ingest.openpyxl_reader import OpenpyxlWorkbookReader
from excel_to_act.interfaces.cli import app
from excel_to_act.schemas import (
    CompletenessReport,
    Handoff,
    RunMetadata,
    UnsupportedSeverity,
    WorkbookManifest,
)
from excel_to_act.steps.step1.source_scan import scan_step1_source
from excel_to_act.store.local_store import LocalArtifactStore
from excel_to_act.ingest.cached_values import read_cached_values

FIXTURES = Path(__file__).parent / "fixtures"
FAILURE_CASES = [
    "encrypted",
    "not-zip",
    "truncated",
    "missing-content-types",
    "malformed-content-types",
    "missing-workbook-xml",
    "malformed-workbook-xml",
    "missing-sheet-name",
    "malformed-sheet-id",
    "malformed-calc-id",
    "malformed-style-index",
    "malformed-chart-xml",
    "xlsb",
    "xls",
    "csv",
]
MISSING_OUTPUTS = {
    "inventory.json",
    "dependency_graph.json",
    "module_classification.json",
    "confirmation_template.json",
}


def _assert_handoff_artifact_hashes(handoff: Handoff) -> None:
    assert not {artifact.name for artifact in handoff.artifacts} & {
        "run_metadata.json",
        "artifact_index.json",
    }
    for artifact in handoff.artifacts:
        assert hashlib.sha256(Path(artifact.path).read_bytes()).hexdigest() == artifact.sha256


def _valid_workbook(path: Path) -> Path:
    workbook = Workbook()
    workbook.active["A1"] = "input"
    workbook.active["B1"] = "=1+1"
    workbook.save(path)
    workbook.close()
    return path


def _malformed_chart_workbook(path: Path) -> Path:
    source = path.with_name(f"{path.stem}-source.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["category", "value"])
    sheet.append(["one", 1])
    sheet.append(["two", 2])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, "D2")
    workbook.save(source)
    workbook.close()

    with zipfile.ZipFile(source) as original, zipfile.ZipFile(path, "w") as output:
        for item in original.infolist():
            content = (
                b"<broken-xml"
                if item.filename == "xl/charts/chart1.xml"
                else original.read(item.filename)
            )
            output.writestr(item, content)
    return path


def _corrupt_first_deflate_byte(path: Path, member_name: str) -> None:
    with zipfile.ZipFile(path) as archive:
        info = archive.getinfo(member_name)
        assert info.compress_type == zipfile.ZIP_DEFLATED
        assert member_name in archive.namelist()
    raw = bytearray(path.read_bytes())
    name_length, extra_length = struct.unpack_from("<HH", raw, info.header_offset + 26)
    data_offset = info.header_offset + 30 + name_length + extra_length
    raw[data_offset] = 0x06
    path.write_bytes(raw)


def _failure_input(case: str, tmp_path: Path) -> Path:
    if case == "encrypted":
        return FIXTURES / "encrypted.xlsx"
    if case == "not-zip":
        path = tmp_path / "not-zip.xlsx"
        path.write_bytes(b"this is not an OOXML package")
        return path
    if case == "truncated":
        path = tmp_path / "truncated.xlsx"
        path.write_bytes(b"PK\x03\x04" + b"truncated" * 8)
        return path
    if case == "malformed-chart-xml":
        return _malformed_chart_workbook(tmp_path / "malformed-chart.xml.xlsx")
    if case in {
        "missing-content-types",
        "malformed-content-types",
        "missing-workbook-xml",
        "malformed-workbook-xml",
        "missing-sheet-name",
        "malformed-sheet-id",
        "malformed-calc-id",
        "malformed-style-index",
    }:
        source = _valid_workbook(tmp_path / "source.xlsx")
        path = tmp_path / f"{case}.xlsx"
        part = (
            "[Content_Types].xml"
            if "content-types" in case
            else "xl/worksheets/sheet1.xml"
            if case == "malformed-style-index"
            else "xl/workbook.xml"
        )
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(path, "w") as output:
            for item in original.infolist():
                if item.filename == part:
                    if case == "malformed-style-index":
                        root = ET.fromstring(original.read(item.filename))
                        cell = next(
                            node for node in root.iter()
                            if node.tag.endswith("}c") and node.attrib.get("r") == "A1"
                        )
                        cell.set("s", "1")
                        output.writestr(
                            item,
                            ET.tostring(root, encoding="utf-8", xml_declaration=True),
                        )
                        continue
                    if case in {"missing-sheet-name", "malformed-sheet-id", "malformed-calc-id"}:
                        root = ET.fromstring(original.read(item.filename))
                        if case == "malformed-calc-id":
                            calc = next(node for node in root.iter() if node.tag.endswith("}calcPr"))
                            calc.set("calcId", "oops")
                        else:
                            sheet = next(node for node in root.iter() if node.tag.endswith("}sheet"))
                        if case == "missing-sheet-name":
                            del sheet.attrib["name"]
                        elif case == "malformed-sheet-id":
                            sheet.set("sheetId", "bad")
                        output.writestr(
                            item,
                            ET.tostring(root, encoding="utf-8", xml_declaration=True),
                        )
                        continue
                    if case.startswith("missing-"):
                        continue
                    output.writestr(item, b"<broken-xml")
                else:
                    output.writestr(item, original.read(item.filename))
        return path
    path = tmp_path / f"unsupported.{case}"
    path.write_bytes(b"unsupported legacy workbook input")
    return path


@pytest.mark.parametrize("case", FAILURE_CASES)
def test_failed_inspect_writes_diagnostic_only_and_stops_downstream(
    case: str, tmp_path: Path
) -> None:
    workbook = _failure_input(case, tmp_path)
    if case == "malformed-style-index":
        with zipfile.ZipFile(workbook) as archive:
            assert archive.testzip() is None
            sheet_root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            styles_root = ET.fromstring(archive.read("xl/styles.xml"))
        cell = next(
            node for node in sheet_root.iter()
            if node.tag.endswith("}c") and node.attrib.get("r") == "A1"
        )
        cell_xfs = next(node for node in styles_root if node.tag.endswith("}cellXfs"))
        assert cell.attrib["s"] == "1"
        assert len(cell_xfs) == 1
        assert scan_step1_source(workbook)["cells"]
    if case == "malformed-chart-xml":
        with zipfile.ZipFile(workbook) as archive:
            assert archive.testzip() is None
            assert archive.read("xl/charts/chart1.xml") == b"<broken-xml"
        assert scan_step1_source(workbook)["cells"]
    if case in {"missing-sheet-name", "malformed-sheet-id", "malformed-calc-id"}:
        with zipfile.ZipFile(workbook) as archive:
            assert archive.testzip() is None
            root = ET.fromstring(archive.read("xl/workbook.xml"))
        if case == "malformed-calc-id":
            calc = next(node for node in root.iter() if node.tag.endswith("}calcPr"))
            assert calc.attrib["calcId"] == "oops"
        else:
            sheet = next(node for node in root.iter() if node.tag.endswith("}sheet"))
        if case == "missing-sheet-name":
            assert "name" not in sheet.attrib
        elif case == "malformed-sheet-id":
            assert sheet.attrib["sheetId"] == "bad"
    out = tmp_path / "artifacts"
    with (
        patch("excel_to_act.orchestrator.phase1.OpenpyxlInventoryExtractor.extract") as extract,
        patch("excel_to_act.orchestrator.phase1.RegexFormulaGraphBuilder.build") as graph,
        patch("excel_to_act.orchestrator.phase1.RuleBasedClassifier.classify") as classify,
        patch("excel_to_act.orchestrator.phase1.extract_vba_project") as vba,
        patch("excel_to_act.orchestrator.phase1.ConfirmationTemplateBuilder.build") as confirmation,
    ):
        result = CliRunner().invoke(app, ["inspect", str(workbook), "--out", str(out)])

    for downstream in (extract, graph, classify, vba, confirmation):
        downstream.assert_not_called()
    assert result.exit_code == 1, result.output
    assert "Traceback" not in result.output
    assert "Reason:" in result.output
    assert "completeness.json" in result.output

    metadata_data = json.loads((out / "run_metadata.json").read_text(encoding="utf-8"))
    metadata = RunMetadata.model_validate(metadata_data)
    store = LocalArtifactStore(out)
    reloaded = store.read_run(metadata.workbook_sha256, metadata.run_id)
    manifest = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "workbook_manifest.json"),
        WorkbookManifest,
    )
    handoff_path = next(a.path for a in reloaded.artifacts if a.name == "handoff.json")
    handoff = store.read_json(handoff_path, Handoff)
    assert handoff.status == "fail"
    assert handoff.next_step == "rerun_inspect"
    assert not handoff.opaque_summary
    assert any("Not produced" in action for action in handoff.next_actions)
    next_actions = " ".join(handoff.next_actions)
    assert all(name in next_actions for name in MISSING_OUTPUTS)
    assert "Resolve the input error and rerun inspect" in next_actions
    assert "Step 2" not in next_actions
    _assert_handoff_artifact_hashes(handoff)
    assert any(
        feature.severity == UnsupportedSeverity.error
        for feature in manifest.unsupported_features
    )
    if case in {"missing-sheet-name", "malformed-sheet-id", "malformed-calc-id", "malformed-style-index", "malformed-chart-xml"}:
        feature = manifest.unsupported_features[0]
        completeness = LocalArtifactStore(out).read_json(
            next(a.path for a in reloaded.artifacts if a.name == "completeness.json"),
            CompletenessReport,
        )
        if case in {"malformed-calc-id", "malformed-style-index", "malformed-chart-xml"}:
            assert "InvalidFileException" in feature.description
            assert (
                "TypeError" in feature.description
                if case == "malformed-calc-id"
                else "IndexError" in feature.description
                if case == "malformed-style-index"
                else "XMLSyntaxError" in feature.description
            )
        else:
            assert "SourceScanError" in feature.description
            assert "xl/workbook.xml" in feature.description
            assert (
                "required name" in feature.description
                if case == "missing-sheet-name"
                else "invalid sheetId" in feature.description
            )
        completeness_path = next(
            a.path for a in reloaded.artifacts if a.name == "completeness.json"
        )
        assert f"Diagnostic: {completeness_path}" in result.output
        assert Path(completeness_path).exists()
        assert completeness.status == "fail"
    if case == "encrypted":
        feature = manifest.unsupported_features[0]
        assert "password-protected" in feature.description
        assert "decrypt" in feature.description.lower()
        assert "resave" in feature.description.lower()
        completeness_path = next(
            a.path for a in reloaded.artifacts if a.name == "completeness.json"
        )
        assert f"Diagnostic: {completeness_path}" in result.output
    if case == "not-zip":
        assert "BadZipFile" in manifest.unsupported_features[0].description
    assert {a.name for a in reloaded.artifacts}.isdisjoint(MISSING_OUTPUTS)
    assert not any((out / name).exists() for name in MISSING_OUTPUTS)
    assert (out / "completeness.json").exists()
    assert (out / "handoff.md").exists()
    markdown = (out / "handoff.md").read_text(encoding="utf-8")
    assert "Full detail lives in `inventory.json` / `dependency_graph.json`." not in markdown


def _assert_controlled_cli_failure(workbook: Path, out: Path, reason: str) -> None:
    result = CliRunner().invoke(app, ["inspect", str(workbook), "--out", str(out)])
    assert result.exit_code == 1, result.output
    assert "Traceback" not in result.output
    assert reason in result.output
    assert "Diagnostic:" in result.output
    assert (out / "completeness.json").exists()


def test_package_read_oserror_becomes_manifest_and_cli_diagnostic(tmp_path: Path) -> None:
    workbook = _valid_workbook(tmp_path / "valid.xlsx")
    original_read = zipfile.ZipFile.read

    def fail_package_read(zf: zipfile.ZipFile, name: str, *args, **kwargs):
        if name == "xl/workbook.xml":
            raise OSError("simulated package read failure")
        return original_read(zf, name, *args, **kwargs)

    with patch.object(zipfile.ZipFile, "read", autospec=True, side_effect=fail_package_read):
        manifest = OpenpyxlWorkbookReader().read_manifest(workbook)
    feature = manifest.unsupported_features[0]
    assert feature.feature_type == "input_read_error"
    assert feature.severity == UnsupportedSeverity.error
    assert not feature.opaque
    assert "SourceScanError: xl/workbook.xml: simulated package read failure" in feature.description

    out = tmp_path / "artifacts"
    with patch.object(zipfile.ZipFile, "read", autospec=True, side_effect=fail_package_read):
        _assert_controlled_cli_failure(
            workbook, out, "Reason: Could not read workbook: SourceScanError"
        )


@pytest.mark.parametrize(
    "loader_module",
    ["excel_to_act.inventory.extractor", "excel_to_act.ingest.cached_values"],
)
def test_later_openpyxl_read_oserror_writes_diagnostic_only_and_stops_downstream(
    loader_module: str, tmp_path: Path
) -> None:
    workbook = _valid_workbook(tmp_path / "valid.xlsx")
    out = tmp_path / "artifacts"
    with (
        patch(
            f"{loader_module}.load_workbook",
            side_effect=OSError("simulated later workbook read failure"),
        ) as later_read,
        patch("excel_to_act.orchestrator.phase1.RegexFormulaGraphBuilder.build") as graph,
        patch("excel_to_act.orchestrator.phase1.extract_vba_project") as vba,
        patch("excel_to_act.orchestrator.phase1.RuleBasedClassifier.classify") as classify,
        patch("excel_to_act.orchestrator.phase1.ConfirmationTemplateBuilder.build") as confirmation,
    ):
        result = CliRunner().invoke(app, ["inspect", str(workbook), "--out", str(out)])

    later_read.assert_called_once()
    for downstream in (graph, vba, classify, confirmation):
        downstream.assert_not_called()
    assert result.exit_code == 1, result.output
    assert "Traceback" not in result.output
    assert (
        "Reason: Could not read workbook: InvalidFileException: openpyxl could not read workbook: "
        "OSError: simulated later workbook read failure"
    ) in result.output
    assert "Diagnostic:" in result.output

    metadata = RunMetadata.model_validate_json((out / "run_metadata.json").read_bytes())
    store = LocalArtifactStore(out)
    reloaded = store.read_run(metadata.workbook_sha256, metadata.run_id)
    manifest = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "workbook_manifest.json"),
        WorkbookManifest,
    )
    completeness = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "completeness.json"),
        CompletenessReport,
    )
    handoff = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "handoff.json"), Handoff
    )
    completeness_path = next(a.path for a in reloaded.artifacts if a.name == "completeness.json")
    assert f"Diagnostic: {completeness_path}" in result.output
    assert manifest.unsupported_features[0].severity == UnsupportedSeverity.error
    assert "simulated later workbook read failure" in manifest.unsupported_features[0].description
    assert completeness.status == "fail"
    assert handoff.status == "fail"
    _assert_handoff_artifact_hashes(handoff)
    assert {a.name for a in reloaded.artifacts}.isdisjoint(MISSING_OUTPUTS)
    assert not any((out / name).exists() for name in MISSING_OUTPUTS)


@pytest.mark.parametrize(
    "reader_path",
    [
        "excel_to_act.inventory.extractor.read_data_tables",
        "excel_to_act.inventory.extractor.read_form_controls",
        "excel_to_act.ingest.cached_values._empty_string_formula_caches",
    ],
)
def test_package_read_oserror_is_normalized_at_helper_boundary(
    reader_path: str, tmp_path: Path
) -> None:
    workbook = _valid_workbook(tmp_path / "valid.xlsx")
    out = tmp_path / "artifacts"
    with (
        patch(reader_path, side_effect=OSError("simulated package read failure")) as reader,
        patch("excel_to_act.orchestrator.phase1.RegexFormulaGraphBuilder.build") as graph,
        patch("excel_to_act.orchestrator.phase1.extract_vba_project") as vba,
        patch("excel_to_act.orchestrator.phase1.RuleBasedClassifier.classify") as classify,
        patch("excel_to_act.orchestrator.phase1.ConfirmationTemplateBuilder.build") as confirmation,
    ):
        result = CliRunner().invoke(app, ["inspect", str(workbook), "--out", str(out)])

    reader.assert_called_once()
    for downstream in (graph, vba, classify, confirmation):
        downstream.assert_not_called()
    assert result.exit_code == 1, result.output
    assert "Traceback" not in result.output
    assert "InvalidFileException" in result.output
    assert "OSError: simulated package read failure" in result.output
    assert "Diagnostic:" in result.output

    metadata = RunMetadata.model_validate_json((out / "run_metadata.json").read_bytes())
    store = LocalArtifactStore(out)
    reloaded = store.read_run(metadata.workbook_sha256, metadata.run_id)
    manifest = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "workbook_manifest.json"),
        WorkbookManifest,
    )
    completeness = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "completeness.json"),
        CompletenessReport,
    )
    handoff = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "handoff.json"), Handoff
    )
    completeness_path = next(a.path for a in reloaded.artifacts if a.name == "completeness.json")
    assert f"Diagnostic: {completeness_path}" in result.output
    assert manifest.unsupported_features[0].severity == UnsupportedSeverity.error
    assert "simulated package read failure" in manifest.unsupported_features[0].description
    assert completeness.status == "fail"
    assert handoff.status == "fail"
    _assert_handoff_artifact_hashes(handoff)
    assert {a.name for a in reloaded.artifacts}.isdisjoint(MISSING_OUTPUTS)
    assert not any((out / name).exists() for name in MISSING_OUTPUTS)


def test_malformed_chart_after_success_clears_stale_aliases_and_stops_downstream(
    tmp_path: Path,
) -> None:
    out = tmp_path / "artifacts"
    good = _valid_workbook(tmp_path / "good.xlsx")
    success = CliRunner().invoke(app, ["inspect", str(good), "--out", str(out)])
    assert success.exit_code == 0, success.output
    success_metadata = RunMetadata.model_validate_json((out / "run_metadata.json").read_bytes())
    store = LocalArtifactStore(out)
    successful_run = store.read_run(success_metadata.workbook_sha256, success_metadata.run_id)
    successful_inventory = Path(
        next(a.path for a in successful_run.artifacts if a.name == "inventory.json")
    )
    previous_metadata_bytes = (out / "run_metadata.json").read_bytes()
    assert (out / "inventory.json").exists()

    broken = _malformed_chart_workbook(tmp_path / "broken-chart.xlsx")
    with (
        patch("excel_to_act.orchestrator.phase1.OpenpyxlInventoryExtractor.extract") as extract,
        patch("excel_to_act.orchestrator.phase1.RegexFormulaGraphBuilder.build") as graph,
        patch("excel_to_act.orchestrator.phase1.extract_vba_project") as vba,
        patch("excel_to_act.orchestrator.phase1.RuleBasedClassifier.classify") as classify,
        patch("excel_to_act.orchestrator.phase1.ConfirmationTemplateBuilder.build") as confirmation,
    ):
        result = CliRunner().invoke(app, ["inspect", str(broken), "--out", str(out)])

    for downstream in (extract, graph, vba, classify, confirmation):
        downstream.assert_not_called()
    assert result.exit_code == 1, result.output
    assert "Traceback" not in result.output
    assert "Reason: Could not read workbook: InvalidFileException" in result.output
    assert "XMLSyntaxError" in result.output
    assert "Diagnostic:" in result.output
    assert not (out / "inventory.json").exists()
    assert successful_inventory.exists()
    assert (out / "run_metadata.json").read_bytes() != previous_metadata_bytes

    failed_metadata = RunMetadata.model_validate_json((out / "run_metadata.json").read_bytes())
    assert failed_metadata.run_id != success_metadata.run_id
    failed_run = store.read_run(failed_metadata.workbook_sha256, failed_metadata.run_id)
    assert {a.name for a in failed_run.artifacts}.isdisjoint(MISSING_OUTPUTS)
    manifest = store.read_json(
        next(a.path for a in failed_run.artifacts if a.name == "workbook_manifest.json"),
        WorkbookManifest,
    )
    completeness = store.read_json(
        next(a.path for a in failed_run.artifacts if a.name == "completeness.json"),
        CompletenessReport,
    )
    handoff = store.read_json(
        next(a.path for a in failed_run.artifacts if a.name == "handoff.json"), Handoff
    )
    completeness_path = next(a.path for a in failed_run.artifacts if a.name == "completeness.json")
    assert f"Diagnostic: {completeness_path}" in result.output
    assert "XMLSyntaxError" in manifest.unsupported_features[0].description
    assert completeness.status == "fail"
    assert handoff.status == "fail"
    _assert_handoff_artifact_hashes(handoff)


def test_cached_value_lazy_iterator_errors_are_normalized_at_advance(
    tmp_path: Path,
) -> None:
    workbook = _valid_workbook(tmp_path / "valid.xlsx")

    class LazyWorksheet:
        title = "Sheet"

        def iter_rows(self):
            yield ()
            raise ValueError("simulated lazy XML read failure")

    class LoadedWorkbook:
        worksheets = [LazyWorksheet()]

        def close(self) -> None:
            pass

    with patch(
        "excel_to_act.ingest.cached_values.load_workbook", return_value=LoadedWorkbook()
    ):
        with pytest.raises(InvalidFileException, match="ValueError: simulated lazy XML read failure"):
            read_cached_values(workbook)


def test_corrupt_deflate_member_writes_controlled_failure_artifacts(tmp_path: Path) -> None:
    workbook = _valid_workbook(tmp_path / "corrupt-member.xlsx")
    _corrupt_first_deflate_byte(workbook, "[Content_Types].xml")
    with zipfile.ZipFile(workbook) as archive:
        assert "[Content_Types].xml" in archive.namelist()

    out = tmp_path / "artifacts"
    result = CliRunner().invoke(app, ["inspect", str(workbook), "--out", str(out)])
    assert result.exit_code == 1, result.output
    assert "Traceback" not in result.output
    assert "Reason: Could not read workbook: error:" in result.output

    metadata = RunMetadata.model_validate_json((out / "run_metadata.json").read_bytes())
    store = LocalArtifactStore(out)
    reloaded = store.read_run(metadata.workbook_sha256, metadata.run_id)
    manifest = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "workbook_manifest.json"),
        WorkbookManifest,
    )
    completeness = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "completeness.json"),
        CompletenessReport,
    )
    handoff = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "handoff.json"), Handoff
    )
    completeness_path = next(a.path for a in reloaded.artifacts if a.name == "completeness.json")
    assert f"Diagnostic: {completeness_path}" in result.output
    assert Path(completeness_path).exists()
    assert manifest.unsupported_features[0].severity == UnsupportedSeverity.error
    assert "error" in manifest.unsupported_features[0].description.lower()
    assert completeness.status == "fail"
    assert handoff.status == "fail"
    _assert_handoff_artifact_hashes(handoff)


def test_openpyxl_invalid_file_after_scan_becomes_manifest_and_cli_diagnostic(
    tmp_path: Path,
) -> None:
    workbook = _valid_workbook(tmp_path / "valid.xlsx")
    with (
        patch(
            "excel_to_act.ingest.openpyxl_reader.scan_step1_source", wraps=scan_step1_source
        ) as scan,
        patch(
            "excel_to_act.ingest.openpyxl_reader.load_workbook",
            side_effect=InvalidFileException("simulated openpyxl read failure"),
        ),
    ):
        manifest = OpenpyxlWorkbookReader().read_manifest(workbook)
    scan.assert_called_once_with(workbook.resolve())
    feature = manifest.unsupported_features[0]
    assert feature.feature_type == "input_read_error"
    assert feature.severity == UnsupportedSeverity.error
    assert not feature.opaque
    assert "InvalidFileException: simulated openpyxl read failure" in feature.description

    out = tmp_path / "artifacts"
    with patch(
        "excel_to_act.ingest.openpyxl_reader.load_workbook",
        side_effect=InvalidFileException("simulated openpyxl read failure"),
    ):
        _assert_controlled_cli_failure(
            workbook, out, "Reason: Could not read workbook: InvalidFileException"
        )


def test_missing_input_writes_error_manifest_and_cli_diagnostic(tmp_path: Path) -> None:
    workbook = tmp_path / "missing.xlsx"
    out = tmp_path / "artifacts"
    result = CliRunner().invoke(app, ["inspect", str(workbook), "--out", str(out)])

    assert result.exit_code == 1, result.output
    assert "Reason: Could not read workbook: FileNotFoundError" in result.output

    metadata = RunMetadata.model_validate_json((out / "run_metadata.json").read_bytes())
    store = LocalArtifactStore(out)
    reloaded = store.read_run(metadata.workbook_sha256, metadata.run_id)
    completeness_path = next(a.path for a in reloaded.artifacts if a.name == "completeness.json")
    assert f"Diagnostic: {completeness_path}" in result.output
    assert Path(completeness_path).exists()
    manifest = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "workbook_manifest.json"),
        WorkbookManifest,
    )
    completeness = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "completeness.json"),
        CompletenessReport,
    )
    handoff = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "handoff.json"), Handoff
    )
    assert manifest.unsupported_features[0].severity == UnsupportedSeverity.error
    assert "FileNotFoundError" in manifest.unsupported_features[0].description
    assert completeness.status == "fail"
    assert handoff.status == "fail"
    assert any(a.name == "handoff.md" for a in reloaded.artifacts)


def test_directory_input_writes_read_failure_diagnostic(tmp_path: Path) -> None:
    workbook = tmp_path / "directory-input"
    workbook.mkdir()
    out = tmp_path / "artifacts"
    result = CliRunner().invoke(app, ["inspect", str(workbook), "--out", str(out)])

    assert result.exit_code == 1, result.output
    assert "Reason: Could not read workbook:" in result.output

    metadata = RunMetadata.model_validate_json((out / "run_metadata.json").read_bytes())
    store = LocalArtifactStore(out)
    reloaded = store.read_run(metadata.workbook_sha256, metadata.run_id)
    completeness_path = next(a.path for a in reloaded.artifacts if a.name == "completeness.json")
    assert f"Diagnostic: {completeness_path}" in result.output
    assert Path(completeness_path).exists()
    manifest = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "workbook_manifest.json"),
        WorkbookManifest,
    )
    completeness = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "completeness.json"),
        CompletenessReport,
    )
    handoff = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "handoff.json"), Handoff
    )
    assert manifest.unsupported_features[0].severity == UnsupportedSeverity.error
    assert manifest.unsupported_features[0].description.startswith("Could not read workbook:")
    assert completeness.status == "fail"
    assert handoff.status == "fail"


def test_failure_removes_stale_root_aliases_and_preserves_prior_run(tmp_path: Path) -> None:
    out = tmp_path / "artifacts"
    good = _valid_workbook(tmp_path / "good.xlsx")
    success = CliRunner().invoke(app, ["inspect", str(good), "--out", str(out)])
    assert success.exit_code == 0, success.output

    store = LocalArtifactStore(out)
    success_metadata = RunMetadata.model_validate_json((out / "run_metadata.json").read_bytes())
    canonical = store.read_run(success_metadata.workbook_sha256, success_metadata.run_id)
    aliases = {
        "inventory.json",
        "dependency_graph.json",
        "module_classification.json",
        "confirmation_template.json",
    }
    prior_paths = {artifact.name: Path(artifact.path) for artifact in canonical.artifacts}
    assert aliases <= prior_paths.keys()
    assert all((out / name).exists() for name in aliases)
    success_handoff = Path(next(a.path for a in canonical.artifacts if a.name == "handoff.md"))
    assert "Full detail lives in `inventory.json` / `dependency_graph.json`." in success_handoff.read_text(encoding="utf-8")

    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not an OOXML workbook")
    failure = CliRunner().invoke(app, ["inspect", str(bad), "--out", str(out)])
    assert failure.exit_code == 1, failure.output
    assert not any((out / name).exists() for name in aliases)
    assert all(path.exists() for path in prior_paths.values())
    reloaded_prior = store.read_run(success_metadata.workbook_sha256, success_metadata.run_id)
    assert aliases <= {artifact.name for artifact in reloaded_prior.artifacts}


@pytest.mark.parametrize("suffix", ["xlsx", "xlsm"])
def test_successful_inspect_artifacts_read_back(suffix: str, tmp_path: Path) -> None:
    workbook = _valid_workbook(tmp_path / f"valid.{suffix}")
    out = tmp_path / "artifacts"
    result = CliRunner().invoke(app, ["inspect", str(workbook), "--out", str(out)])
    assert result.exit_code == 0, result.output

    metadata_data = json.loads((out / "run_metadata.json").read_text(encoding="utf-8"))
    metadata = RunMetadata.model_validate(metadata_data)
    store = LocalArtifactStore(out)
    reloaded = store.read_run(metadata.workbook_sha256, metadata.run_id)
    assert reloaded.completeness_status in {"pass", "warn"}
    handoff_data = store.read_json(
        next(a.path for a in reloaded.artifacts if a.name == "handoff.json"), Handoff
    )
    _assert_handoff_artifact_hashes(handoff_data)
    handoff_md = next(a.path for a in reloaded.artifacts if a.name == "handoff.md")
    assert "Full detail lives in `inventory.json` / `dependency_graph.json`." in Path(handoff_md).read_text(encoding="utf-8")
    assert {a.name for a in reloaded.artifacts} >= {
        "workbook_manifest.json",
        "inventory.json",
        "dependency_graph.json",
        "module_classification.json",
        "confirmation_template.json",
        "completeness.json",
        "handoff.json",
    }
