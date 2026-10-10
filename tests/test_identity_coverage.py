"""Identity accounting across the legacy inspect artifact boundary."""

from __future__ import annotations

import json
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName
from typer.testing import CliRunner

from excel_to_act.ingest.openpyxl_reader import OpenpyxlWorkbookReader
from excel_to_act.interfaces.cli import app
from excel_to_act.inventory.extractor import OpenpyxlInventoryExtractor
from excel_to_act.schemas import CompletenessReport, FormulaGraph, Handoff, ModuleClassification, SourceLocation, UnsupportedFeature, WorkbookManifest
from excel_to_act.steps.step2.workflow import build_index
from excel_to_act.steps.step1.source_scan import object_identity, scan_step1_source
from excel_to_act.store.local_store import LocalArtifactStore
from excel_to_act.verify.completeness import verify_completeness


def _fixture(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "Item"
    ws["B1"] = "Amount"
    ws["A2"] = "Rate"
    ws["B2"] = 0.05
    ws.add_table(Table(displayName="Rates", ref="A1:B2"))
    ws.conditional_formatting.add("B2", CellIsRule(operator="greaterThan", formula=["0"]))
    wb.defined_names.add(DefinedName("Rate", attr_text="Inputs!$B$2"))
    wb.save(path)
    wb.close()
    return path


def _verified(path: Path):
    manifest = OpenpyxlWorkbookReader().read_manifest(path)
    inventory = OpenpyxlInventoryExtractor().extract(path, manifest)
    report = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    return manifest, inventory, report


def test_normal_fixture_closes_separate_identity_ledgers(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "normal.xlsx")
    manifest, inventory, report = _verified(path)
    scan = scan_step1_source(path)
    assert report.status == "pass"
    assert report.discovered_workbook_objects == len(scan["objects"])
    assert inventory.coverage.recognized_inventory_objects == len(scan["objects"])
    assert report.unsupported_or_opaque_objects == 0
    assert all(row.expected == row.actual and not row.missing_identities for row in report.object_coverage)
    assert sum(row.expected for row in report.object_coverage if row.kind != "package_part") == len(scan["objects"])
    package = next(row for row in report.object_coverage if row.kind == "package_part")
    assert package.sheet_name is None
    assert package.expected == package.actual == len(manifest.package_parts)
    assert all(part.source_location.source_identity == object_identity("package_part", part.name) for part in manifest.package_parts)
    assert all(item.source_identity for sheet in inventory.sheets for item in [sheet, *sheet.cells, *sheet.ranges, *sheet.layout_objects])
    assert all(item.source_identity for item in inventory.workbook_ranges)


def test_zip_directory_entry_is_not_a_package_part(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "directory.xlsx")
    with zipfile.ZipFile(path, "a") as package:
        package.writestr("xl/media/", b"")
    manifest, _, report = _verified(path)
    assert report.status == "pass"
    assert not any(part.name == "xl/media/" for part in manifest.package_parts)


def test_print_area_declared_name_is_accounted_and_removal_fails(tmp_path: Path) -> None:
    path = tmp_path / "print-area.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = 1
    ws.print_area = "A1:B2"
    wb.save(path)
    wb.close()
    manifest, inventory, report = _verified(path)
    source = next(obj for obj in scan_step1_source(path)["objects"] if obj["kind"] == "defined_name" and obj["details"]["name"] == "_xlnm.Print_Area")
    names = [item for item in inventory.workbook_ranges if item.name == "_xlnm.Print_Area"]
    assert source["details"]["scope"] == "Inputs"
    assert report.status == "pass"
    assert len(names) == 1 and names[0].source_identity == source["identity"]
    cli = CliRunner().invoke(app, ["inspect", str(path), "--out", str(tmp_path / "print-area-out")])
    assert cli.exit_code == 0, cli.output
    inventory.workbook_ranges.remove(names[0])
    missing = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    assert missing.status == "fail"
    assert any(gap.source_identity == source["identity"] for gap in missing.gaps)


@pytest.mark.parametrize("kind", ["conditional_formatting", "data_validation"])
def test_reordered_multi_area_sqref_matches_and_removal_fails(tmp_path: Path, kind: str) -> None:
    path = tmp_path / f"{kind}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = 1
    ws["B2"] = 2
    if kind == "conditional_formatting":
        ws.conditional_formatting.add("A1 B2", CellIsRule(operator="greaterThan", formula=["0"]))
    else:
        validation = DataValidation(type="whole", operator="greaterThan", formula1="0")
        ws.add_data_validation(validation)
        validation.add("A1")
        validation.add("B2")
    wb.save(path)
    wb.close()
    rewritten = BytesIO()
    with zipfile.ZipFile(path) as original, zipfile.ZipFile(rewritten, "w") as output:
        for member in original.infolist():
            data = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                assert data.count(b'sqref="A1 B2"') == 1
                data = data.replace(b'sqref="A1 B2"', b'sqref="B2 A1"')
            output.writestr(member, data)
    path.write_bytes(rewritten.getvalue())
    manifest, inventory, report = _verified(path)
    source = next(obj for obj in scan_step1_source(path)["objects"] if obj["kind"] == kind)
    record = next(item for item in inventory.sheets[0].layout_objects if item.kind == kind)
    assert source["details"]["address"] == "B2 A1"
    assert record.address == "A1 B2"
    assert report.status == "pass"
    assert record.source_identity == source["identity"]
    cli = CliRunner().invoke(app, ["inspect", str(path), "--out", str(tmp_path / f"{kind}-out")])
    assert cli.exit_code == 0, cli.output
    inventory.sheets[0].layout_objects.remove(record)
    missing = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    assert missing.status == "fail"
    assert any(gap.source_identity == source["identity"] for gap in missing.gaps)


@pytest.mark.parametrize("kind", ["cell", "defined_name", "conditional_formatting", "table"])
def test_missing_declared_identity_fails_with_source_gap(tmp_path: Path, kind: str) -> None:
    path = _fixture(tmp_path / "missing.xlsx")
    manifest = OpenpyxlWorkbookReader().read_manifest(path)
    inventory = OpenpyxlInventoryExtractor().extract(path, manifest)
    if kind == "cell":
        inventory.sheets[0].cells = [c for c in inventory.sheets[0].cells if c.address != "B2"]
    elif kind == "defined_name":
        inventory.workbook_ranges = [r for r in inventory.workbook_ranges if r.kind != kind]
    else:
        for sheet in inventory.sheets:
            sheet.ranges = [r for r in sheet.ranges if r.kind != kind]
            sheet.layout_objects = [r for r in sheet.layout_objects if r.kind != kind]
    report = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    row = next(row for row in report.object_coverage if row.kind == kind)
    assert report.status == "fail"
    assert row.expected > row.actual
    assert len(row.missing_identities) == 1
    assert any(gap.source_identity == row.missing_identities[0] and gap.object_type == kind for gap in report.gaps)
    assert any(kind in reason for reason in report.blocking_reasons)


def test_missing_package_part_and_scan_error_are_failures(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "parts.xlsx")
    manifest, inventory, _ = _verified(path)
    missing = manifest.package_parts.pop()
    report = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    assert report.status == "fail"
    assert object_identity("package_part", missing.name) in next(
        row for row in report.object_coverage if row.kind == "package_part"
    ).missing_identities
    with patch("excel_to_act.verify.completeness.scan_step1_source", side_effect=ValueError("bad XML")):
        report = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    assert report.status == "fail"
    assert any("bad XML" in reason for reason in report.blocking_reasons)


def test_duplicate_output_identity_does_not_close_coverage(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "duplicate.xlsx")
    manifest, inventory, _ = _verified(path)
    inventory.sheets[0].cells.append(inventory.sheets[0].cells[0].model_copy(deep=True))
    report = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    row = next(row for row in report.object_coverage if row.kind == "cell")
    assert report.status == "fail"
    assert row.actual == row.expected + 1
    assert row.duplicate_identities == [inventory.sheets[0].cells[0].source_identity]


@pytest.mark.parametrize("failure", ["missing_part", "scan_error"])
def test_unreadable_source_saves_failed_handoff(tmp_path: Path, failure: str) -> None:
    path = _fixture(tmp_path / "source.xlsx")
    if failure == "missing_part":
        from io import BytesIO

        data = BytesIO()
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(data, "w") as broken:
            for member in source.infolist():
                if member.filename != "xl/worksheets/sheet1.xml":
                    broken.writestr(member, source.read(member.filename))
        path.write_bytes(data.getvalue())
        context = patch("excel_to_act.ingest.openpyxl_reader.scan_step1_source", wraps=scan_step1_source)
    else:
        context = patch("excel_to_act.ingest.openpyxl_reader.scan_step1_source", side_effect=ValueError("bad XML"))
    out = tmp_path / "out"
    with context:
        result = CliRunner().invoke(app, ["inspect", str(path), "--out", str(out)])
    assert result.exit_code == 1, result.output
    report = CompletenessReport.model_validate_json((out / "completeness.json").read_bytes())
    handoff = Handoff.model_validate_json((out / "handoff.json").read_bytes())
    assert report.status == handoff.status == "fail"
    assert report.blocking_reasons and handoff.blockers
    assert "missing" in result.output.lower() if failure == "missing_part" else "bad XML" in result.output


def test_failed_identity_gap_persists_and_cli_returns_nonzero(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "failed.xlsx")
    out = tmp_path / "out"
    original = OpenpyxlInventoryExtractor.extract

    def omit_cf(self, workbook_path, manifest):
        inventory = original(self, workbook_path, manifest)
        inventory.sheets[0].layout_objects = [item for item in inventory.sheets[0].layout_objects if item.kind != "conditional_formatting"]
        return inventory

    with patch.object(OpenpyxlInventoryExtractor, "extract", omit_cf):
        result = CliRunner().invoke(app, ["inspect", str(path), "--out", str(out)])
    assert result.exit_code == 1, result.output
    metadata = json.loads((out / "run_metadata.json").read_text(encoding="utf-8"))
    store = LocalArtifactStore(out)
    loaded = store.read_run(metadata["workbook_sha256"], metadata["run_id"])
    report = store.read_json(next(a.path for a in loaded.artifacts if a.name == "completeness.json"), CompletenessReport)
    handoff = store.read_json(next(a.path for a in loaded.artifacts if a.name == "handoff.json"), Handoff)
    assert report.status == handoff.status == "fail"
    assert any("conditional_formatting" in reason for reason in handoff.blockers)
    assert any(row.kind == "conditional_formatting" and row.missing_identities for row in report.object_coverage)
    downstream = build_index(out / "handoff.json", out, tmp_path / "step2")
    assert downstream["status"] == "blocked"


def test_opaque_part_is_unique_and_warnings_are_not_objects(tmp_path: Path) -> None:
    path = _fixture(tmp_path / "opaque.xlsx")
    with zipfile.ZipFile(path, "a") as package:
        package.writestr("xl/media/image1.png", b"opaque image")
    out = tmp_path / "out"
    original = OpenpyxlInventoryExtractor.extract

    def duplicate_diagnostics(self, workbook_path, manifest):
        inventory = original(self, workbook_path, manifest)
        part = next(part for part in manifest.package_parts if part.name == "xl/media/image1.png")
        warning = UnsupportedFeature(
            feature_type="warning_only", description="not another object",
            source_location=SourceLocation(workbook_path=str(path), object_type="workbook"),
            opaque=False,
        )
        duplicate = UnsupportedFeature(
            feature_type="same_media_part", description="duplicate evidence",
            source_location=part.source_location, opaque=True,
        )
        inventory.unsupported_features.extend([warning, duplicate, duplicate])
        return inventory

    with patch.object(OpenpyxlInventoryExtractor, "extract", duplicate_diagnostics):
        result = CliRunner().invoke(app, ["inspect", str(path), "--out", str(out)])
    assert result.exit_code == 0, result.output
    report = CompletenessReport.model_validate_json((out / "completeness.json").read_bytes())
    handoff = Handoff.model_validate_json((out / "handoff.json").read_bytes())
    package = next(row for row in report.object_coverage if row.kind == "package_part")
    assert package.sheet_name is None
    assert package.opaque_expected == package.opaque_actual == 1
    assert report.unsupported_or_opaque_objects == 0
    assert sum(entry["count"] for entry in handoff.opaque_summary) == 1
    assert not any(entry["feature_type"] == "warning_only" for entry in handoff.opaque_summary)


def test_legacy_report_and_exported_schemas_read_back(tmp_path: Path) -> None:
    report = CompletenessReport(workbook_sha256="a" * 64)
    data = report.model_dump(mode="json")
    data.pop("object_coverage", None)
    path = tmp_path / "legacy.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    assert LocalArtifactStore(tmp_path).read_json(path, CompletenessReport).object_coverage == []
    manifest = OpenpyxlWorkbookReader().read_manifest(_fixture(tmp_path / "legacy.xlsx"))
    old_manifest = manifest.model_dump(mode="json")
    for part in old_manifest["package_parts"]:
        part.pop("opaque_reason", None)
    manifest_path = tmp_path / "legacy-manifest.json"
    manifest_path.write_text(json.dumps(old_manifest), encoding="utf-8")
    assert all(part.opaque_reason is None for part in LocalArtifactStore(tmp_path).read_json(manifest_path, WorkbookManifest).package_parts)
    from excel_to_act.schemas import WorkbookInventory

    for name, model in (("workbook_manifest", WorkbookManifest), ("workbook_inventory", WorkbookInventory), ("completeness_report", CompletenessReport), ("handoff", Handoff)):
        assert json.loads(Path(f"schemas/{name}.schema.json").read_text(encoding="utf-8")) == model.model_json_schema()
