"""Completeness verification and the handoff artifact that ends Step 1."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.comments import Comment
from typer.testing import CliRunner

from excel_to_act.ingest.openpyxl_reader import OpenpyxlWorkbookReader
from excel_to_act.interfaces.cli import app
from excel_to_act.inventory.extractor import OpenpyxlInventoryExtractor
from excel_to_act.orchestrator.phase1 import Phase1Orchestrator
from excel_to_act.schemas import CompletenessReport, CompletenessStatus, FormulaGraph, Handoff, ModuleClassification
from excel_to_act.store.local_store import LocalArtifactStore
from excel_to_act.verify.completeness import verify_completeness


def make_fixture(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["B1"] = 0.05
    ws["B2"] = 100
    ws["C2"] = "=B1*B2"
    wb.save(path)
    wb.close()
    return path


def test_completeness_and_handoff_are_written(tmp_path: Path) -> None:
    fixture = make_fixture(tmp_path / "fixture.xlsx")
    out = tmp_path / "artifacts"
    metadata = Phase1Orchestrator().run(fixture, out)

    assert metadata.completeness_status in {"pass", "warn"}

    completeness = LocalArtifactStore(out).read_json(
        next(a.path for a in metadata.artifacts if a.name == "completeness.json"), CompletenessReport
    )
    assert completeness.status in {"pass", "warn"}
    check_names = {check.name for check in completeness.checks}
    assert {
        "sheets_accounted",
        "cells_accounted",
        "content_parts_accounted",
        "formulas_linked",
        "coverage_arithmetic",
    } <= check_names

    handoff_path = Path(next(a.path for a in metadata.artifacts if a.name == "handoff.json"))
    markdown_path = Path(next(a.path for a in metadata.artifacts if a.name == "handoff.md"))
    assert handoff_path.exists() and markdown_path.exists()
    markdown = markdown_path.read_text(encoding="utf-8")
    assert "Handoff" in markdown and "At a glance" in markdown and "Next steps" in markdown

    handoff = LocalArtifactStore(out).read_json(handoff_path, Handoff)
    assert handoff.summary.get("sheets") == 1
    assert handoff.summary.get("cells", 0) > 0
    assert handoff.summary.get("formula_cells") == 1
    assert any(artifact.kind == "inventory" for artifact in handoff.artifacts)
    assert (out / "handoff.md").exists()

    store = LocalArtifactStore(out)
    reloaded = store.read_run(metadata.workbook_sha256, metadata.run_id)
    assert reloaded.completeness_status == metadata.completeness_status


def test_missing_sheet_blocks_completeness(tmp_path: Path) -> None:
    fixture = make_fixture(tmp_path / "fixture.xlsx")
    manifest = OpenpyxlWorkbookReader().read_manifest(fixture)
    inventory = OpenpyxlInventoryExtractor().extract(fixture, manifest)
    inventory.sheets.clear()
    report = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())

    assert report.status == "fail"
    assert report.blocking_reasons
    assert any(gap.object_type == "sheet" for gap in report.gaps)
    assert report.status == CompletenessStatus.fail.value


def test_collected_comment_part_is_not_reported_as_dropped(tmp_path: Path) -> None:
    fixture = make_fixture(tmp_path / "comments.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(fixture)
    wb.active["A1"].comment = Comment("review me", "tester")
    wb.save(fixture)
    wb.close()
    manifest = OpenpyxlWorkbookReader().read_manifest(fixture)
    inventory = OpenpyxlInventoryExtractor().extract(fixture, manifest)
    report = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    check = next(check for check in report.checks if check.name == "content_parts_accounted")
    assert check.passed
    assert not any(gap.object_type == "comment" for gap in report.gaps)

    inventory.sheets[0].layout_objects = [
        item for item in inventory.sheets[0].layout_objects if item.kind != "comment"
    ]
    report = verify_completeness(manifest, inventory, FormulaGraph(), ModuleClassification())
    assert report.status == "fail"
    assert any(gap.object_type == "comment" for gap in report.gaps)


def test_cli_exits_nonzero_when_completeness_fails(tmp_path: Path) -> None:
    fixture = make_fixture(tmp_path / "fixture.xlsx")
    out = tmp_path / "artifacts"
    failing = CompletenessReport(
        workbook_sha256="a" * 64,
        status=CompletenessStatus.fail,
        blocking_reasons=["simulated gap"],
    )
    with patch("excel_to_act.orchestrator.phase1.verify_completeness", return_value=failing):
        result = CliRunner().invoke(app, ["inspect", str(fixture), "--out", str(out)])
    assert result.exit_code == 1
