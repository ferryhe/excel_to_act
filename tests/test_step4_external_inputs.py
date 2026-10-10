from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from typer.testing import CliRunner

from excel_to_act.interfaces.cli import app
from excel_to_act.steps.conversion_workflow import append_stage_artifact, create_workflow, record_decision
from excel_to_act.steps.step4 import external_inputs
from excel_to_act.steps.step4.discovery import discover_active_trace
from excel_to_act.steps.step4.implementation import create_implementation_plan
from excel_to_act.steps.step4.generator import generate_model
from excel_to_act.steps.step5 import validate_generated


_EXTENTS = {
    "external.ci_rate_by_age": "Qtable!W5:W110",
    "external.multiple_ci_rate_by_age": "Qtable!Y5:Y110",
    "external.minor_ci_rate_by_age": "Qtable!Z5:Z110",
    "external.special_ci_rate_by_age": "Qtable!AD5:AD110",
    "external.moderate_ci_rate_by_age": "Qtable!AH5:AH110",
}


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _json(path: Path, value: dict) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, sort_keys=True), encoding="utf-8")
    return _sha(path)


def _external_workflow(tmp_path: Path, extents: dict[str, str] | None = None,
                       age_axis: str = "Qtable!C5:C110", *,
                       generation_ready: bool = False) -> tuple[Path, Path]:
    extents = _EXTENTS if extents is None else extents
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    source_path = source_dir / "model.xlsx"
    workbook = Workbook()
    main = workbook.active
    main.title = "Main"
    if generation_ready:
        main["A1"] = 3
    sheets = {main.title: main}
    for address in [age_axis, *extents.values()]:
        sheet_name = address.split("!", 1)[0]
        if sheet_name not in sheets:
            sheets[sheet_name] = workbook.create_sheet(sheet_name)
    axis_sheet, axis_local = age_axis.split("!", 1)
    axis_min_col, axis_min_row, axis_max_col, axis_max_row = range_boundaries(axis_local)
    if axis_min_col != axis_max_col:
        raise ValueError("fixture external key axis must be vertical")
    for offset, row in enumerate(range(axis_min_row, axis_max_row + 1)):
        sheets[axis_sheet].cell(row, axis_min_col, offset)
    for extent in extents.values():
        sheet, local = extent.split("!", 1)
        min_col, min_row, max_col, max_row = range_boundaries(local)
        if min_col != max_col:
            raise ValueError("fixture external extent must be vertical")
        for row in range(min_row, max_row + 1):
            sheets[sheet].cell(row, min_col, "=UNSUPPORTED()")
    premium = workbook.create_sheet("Premium")
    premium["J1"] = "=Main!A1*2" if generation_ready else "=1"
    workbook.defined_names.add(__import__("openpyxl").workbook.defined_name.DefinedName(
        "GP", attr_text="'Premium'!$J$1"))
    workbook.save(source_path)
    workbook.close()
    source_hash = _sha(source_path)

    workflow = tmp_path / "workflow"
    create_workflow(workflow, {"workbook_path": str(source_path), "workbook_sha256": source_hash,
                               "source_id": "fixture-source", "run_id": "fixture-run"})
    for stage in (1, 2):
        append_stage_artifact(workflow, stage, "checkpoint.json", {"status": "ready_for_review"},
                              f"Synthetic Stage {stage}.")
        for reviewer in ("agent", "human"):
            record_decision(workflow, stage, reviewer, "approve", "Synthetic fixture approval.")

    scenario = {"scenario_id": "saved-pricing-no-overrides", "scenario_sha256": "d" * 64,
                "input_overrides": []}
    if generation_ready:
        scenario["primary_inputs"] = {
            "input.main": {"source_cell": "Main!A1", "source_literal": 3}}
    catalog_variables = ([{"variable_id": "input.main", "role": "raw_business_input",
                           "logical_name": "input.main", "shape": [], "axes": [],
                           "source_extents": [{"sheet": "Main", "range": "A1",
                                                "source_role": "stored_input"}]}]
                         if generation_ready else [])
    semantic_variables = ({
        "input.main": {"variable_id": "input.main", "role": "raw", "kind": "scalar",
                        "shape": [], "axes": [], "dependencies": [], "initial_condition": None,
                        "source_extents": [{"ref": "Main!A1", "role": "raw_input",
                                             "coordinate_mapping": {"kind": "scalar", "fixed_axes": []}}]},
        "model.answer": {"variable_id": "model.answer", "role": "derived", "kind": "scalar",
                          "shape": [], "axes": [], "dependencies": [{"variable_id": "input.main"}],
                          "initial_condition": None,
                          "source_extents": [{"ref": "Premium!J1", "role": "formula_output",
                                               "coordinate_mapping": {"kind": "scalar", "fixed_axes": []}}],
                          "equation_segments": [{"segment_id": "answer", "source_family_ids": ["source.answer"],
                                                 "equation_family_ids": ["equation.answer"], "index_axis": None,
                                                 "index_start": 0, "index_stop_exclusive": 1,
                                                 "execution": "scalar", "recurrence_group": None,
                                                 "snapshot_before_update": False}]},
    } if generation_ready else {})
    external_descriptors = []
    cuts = []
    axis_id = "fixture_age"
    for variable_id, extent in extents.items():
        sheet, local = extent.split("!", 1)
        min_col, min_row, max_col, max_row = range_boundaries(local)
        shape = [max_row - min_row + 1]
        if min_col != max_col:
            raise ValueError("fixture external extent must be vertical")
        catalog_variables.append({"variable_id": variable_id, "role": "formula_derived_external",
                                  "logical_name": variable_id, "shape": shape,
                                  "axes": [{"axis_id": axis_id, "role": "key"}],
                                  "source_extents": [{"sheet": sheet, "range": local, "source_role": "formula_output"}]})
        plan_extent = extent
        semantic_variables[variable_id] = {
            "variable_id": variable_id, "role": "external", "shape": shape,
            "axes": [{"axis_id": axis_id, "name": "age", "source": age_axis}],
            "source_extents": [{"ref": plan_extent, "role": "confirmed_external_cut",
                                "coordinate_mapping": {"kind": "row_ordinal", "axes": [
                                    {"axis": "age", "dimension": 0, "source_coordinate": "row",
                                     "origin": min_row, "index_origin": 0}]}}],
        }
        external_descriptors.append({"variable_id": variable_id, "boundary_variable_id": variable_id,
                                     "source_extents": [plan_extent], "shape": shape})
        members = []
        for row in range(min_row, max_row + 1):
            address = f"{sheet}!{get_column_letter(min_col)}{row}"
            formula = "=UNSUPPORTED()"
            members.append({"address": address, "formula": formula,
                            "formula_sha256": hashlib.sha256(formula.encode()).hexdigest(),
                            "source_workbook_sha256": source_hash})
        cuts.append({"variable_id": variable_id, "source_range": plan_extent,
                     "coordinate_count": shape[0], "traversal": "stopped_before_prerequisites",
                     "formula_members": members})

    boundary_payload = {
        "schema_version": "step3.input_boundary.v1", "phase": "input_boundary", "status": "ready_for_review",
        "boundary_sha256": "e" * 64,
        "source": {"workbook_sha256": source_hash},
        "target_selection": {"scenario": scenario},
        "variables": catalog_variables,
        "axis_metadata": [{"axis_id": axis_id, "shape": [axis_max_row - axis_min_row + 1],
                            "source_extents": [{"sheet": axis_sheet, "range": axis_local}]}],
    }
    boundary_revision = append_stage_artifact(workflow, 3, "input_boundary.json", boundary_payload,
                                              "Synthetic boundary.")
    boundary_reference = {
        "schema_version": "step3.input_boundary.reference.v1", "revision": boundary_revision["revision"],
        "boundary_sha256": boundary_payload["boundary_sha256"], "artifact": boundary_revision["artifact"],
        "review_receipts": {reviewer: {"receipt_id": f"{reviewer}-receipt", "stage": 3,
                                       "revision": boundary_revision["revision"], "decision": "approve",
                                       "artifact_sha256": boundary_revision["artifact"]["json_sha256"],
                                       "artifact_md_sha256": boundary_revision["artifact"]["md_sha256"]}
                            for reviewer in ("agent", "human")},
    }
    manifest_path = workflow / "workflow.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["input_boundary"] = {"revision": boundary_revision["revision"],
                                  "boundary_sha256": boundary_payload["boundary_sha256"],
                                  "artifact": boundary_revision["artifact"], "source_sha256": source_hash}
    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")

    evidence_dir = workflow / "evidence"
    model_scope_path = evidence_dir / "model_scope.json"
    candidate_trace_path = evidence_dir / "candidate_trace.json"
    candidate_profile_path = evidence_dir / "candidate_profile.json"
    _json(model_scope_path, {"source_sha256": source_hash, "allowed_regions": []})
    formula = "=Main!A1*2" if generation_ready else "=1"
    source_families = ([{"family_id": "source.answer", "source_members": ["Premium!J1"],
                         "member_count": 1}] if generation_ready else [])
    _json(candidate_profile_path, {"source_sha256": source_hash, "families": source_families})
    candidate_trace = {"source_sha256": source_hash, "external_formula_cuts": cuts,
                       "cells": [{"address": "Premium!J1", "role": "calculated_formula",
                                  "formula": formula}]}
    _json(candidate_trace_path, candidate_trace)
    plan_path = evidence_dir / "semantic_plan.json"
    semantic_map_path = evidence_dir / "semantic_map.json"
    semantic_map_hash = _json(semantic_map_path, {"source_mappings": ([
        {"source_family_id": "source.answer", "equation_family_id": "equation.answer",
         "variable_id": "model.answer", "module_id": "module.answer",
         "source_members": ["Premium!J1"]}] if generation_ready else [])})
    plan = {"status": "pass", "source": {"source_sha256": source_hash},
            "variables": semantic_variables, "external_boundary_variables": external_descriptors,
            "semantic_map_input": {"path": str(semantic_map_path), "sha256": semantic_map_hash},
            "evidence": {"model_scope": {"path": str(model_scope_path), "sha256": _sha(model_scope_path)},
                         "source_candidate_trace": {"path": str(candidate_trace_path), "sha256": _sha(candidate_trace_path)},
                         "source_family_profile": {"path": str(candidate_profile_path), "sha256": _sha(candidate_profile_path)}},
            "equation_families": ([{"family_id": "equation.answer", "function_name": "answer"}]
                                   if generation_ready else []),
            "source_mappings": ([{"source_family_id": "source.answer", "equation_family_id": "equation.answer",
                                   "variable_id": "model.answer", "module_id": "module.answer",
                                   "source_members": ["Premium!J1"], "member_count": 1}]
                                if generation_ready else []),
            "modules": ([{"module_id": "module.answer"}] if generation_ready else []),
            "execution_plan": ({"passes": [{"pass_id": "pass.answer", "module_id": "module.answer",
                                              "order": 0, "source_family_ids": ["source.answer"],
                                              "variable_ids": ["model.answer"]}]}
                               if generation_ready else {}),
            "coverage": ({"ordinary_formula_members_mapped": 1, "array_formula_followers_mapped": 0,
                          "active_formula_members_mapped": 1, "source_family_count": 1,
                          "semantic_family_count": 1, "array_instance_count": 0}
                         if generation_ready else {})}
    plan_hash = _json(plan_path, plan)
    check_path = evidence_dir / "semantic_check.json"
    check_hash = _json(check_path, {"status": "pass", "semantic_plan_sha256": plan_hash,
                                    "semantic_map_input_sha256": semantic_map_hash,
                                    "source_family_profile_sha256": _sha(candidate_profile_path)})
    stage3_payload = {
        "schema_version": "step3.analysis_design.v1", "status": "ready_for_review",
        "source": {"workbook_path": str(source_path), "workbook_sha256": source_hash,
                   "source_id": "fixture-source", "run_id": "fixture-run"},
        "input_boundary_reference": boundary_reference,
        "analysis": {"artifacts": {
            "semantic_plan": {"path": str(plan_path), "sha256": plan_hash},
            "semantic_check": {"path": str(check_path), "sha256": check_hash},
        }},
        "design": {"semantic_mapping_required": generation_ready, "scenarios": [scenario], "targets": [
            {"target_id": "T_GP", "selector": "GP", "shape": [], "result_order": 0}],
        },
    }
    append_stage_artifact(workflow, 3, "analysis_design.json", stage3_payload, "Synthetic approved design.")
    record_decision(workflow, 3, "agent", "approve", "Synthetic fixture approval.")
    record_decision(workflow, 3, "human", "approve", "Synthetic fixture approval.")
    return workflow, source_path


def _mock_oracle(source: Path, out: Path, *, targets: list[str], ranges: list[str]) -> dict:
    assert targets == []
    out.mkdir(parents=True)
    oracle_ranges = {}
    for selector in ranges:
        _sheet, local = selector.split("!", 1)
        min_col, min_row, max_col, max_row = range_boundaries(local)
        row_count, column_count = max_row - min_row + 1, max_col - min_col + 1
        formulas = [["=UNSUPPORTED()" for _ in range(column_count)] for _ in range(row_count)]
        oracle_ranges[selector] = {
            "qualified_address": selector,
            "values": [[float(row * column_count + col) for col in range(column_count)]
                       for row in range(row_count)],
            "excel_errors": [[None for _ in range(column_count)] for _ in range(row_count)],
            "formulas": formulas,
        }
    source_hash = _sha(source)
    payload = {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "pass",
               "engine": "Microsoft Excel", "engine_settings": {"full_rebuild": True},
               "source_sha256": source_hash, "source_copy_sha256": source_hash,
               "macros_executed": False, "input_overrides": [], "requested_ranges": ranges,
               "ranges": oracle_ranges}
    (out / "excel_oracle.json").write_text(json.dumps(payload), encoding="utf-8")
    (out / "excel_oracle.md").write_text("Synthetic native capture fixture.", encoding="utf-8")
    return {"status": "pass"}


def test_fixed_external_capture_binds_exact_catalog_and_formula_cut_set(tmp_path: Path, monkeypatch) -> None:
    workflow, source = _external_workflow(tmp_path)
    monkeypatch.setattr(external_inputs, "capture_excel_oracle", _mock_oracle)
    result = external_inputs.capture_external_inputs(workflow)
    assert result["status"] == "pass", result
    artifact = json.loads(Path(result["artifact"]).read_text(encoding="utf-8"))
    assert artifact["external_coordinate_count"] == 530
    assert artifact["formula_cut_member_count"] == 530
    assert set(artifact["external_values"]) == set(_EXTENTS)
    assert artifact["native_capture"]["source_copy_sha256"] == _sha(source)
    assert artifact["age_axis"]["values"] == list(range(106))
    assert artifact["native_capture"]["input_overrides"] == []
    assert artifact["native_capture"]["formula_cache_inputs"] is False
    current, _path, digest = external_inputs.load_current_external_inputs(workflow)
    assert digest == _sha(Path(result["artifact"]))
    assert current["external_values_sha256"] == artifact["external_values_sha256"]


def test_external_capture_uses_catalog_ids_and_declared_vector_length(tmp_path: Path, monkeypatch) -> None:
    extents = {"external.life_rate": "Rates!B2:B4", "external.accident_rate": "Rates!D2:D4"}
    workflow, _source = _external_workflow(tmp_path, extents, "Rates!A2:A4")
    monkeypatch.setattr(external_inputs, "capture_excel_oracle", _mock_oracle)
    result = external_inputs.capture_external_inputs(workflow)
    assert result["status"] == "pass", result
    artifact = json.loads(Path(result["artifact"]).read_text(encoding="utf-8"))
    assert artifact["external_coordinate_count"] == 6
    assert artifact["formula_cut_member_count"] == 6
    assert set(artifact["external_values"]) == set(extents)
    assert artifact["age_axis"]["axis_id"] == "fixture_age"
    assert artifact["age_axis"]["source_range"] == "Rates!A2:A4"
    assert artifact["age_axis"]["values"] == [0, 1, 2]


def test_step4_generation_runs_only_through_current_modular_preflight(tmp_path: Path, monkeypatch) -> None:
    from excel_to_act.steps.conversion_workflow import load_workflow

    workflow, source = _external_workflow(
        tmp_path, {"external.rate": "Rates!B2:B3"}, "Rates!A2:A3", generation_ready=True)
    monkeypatch.setattr(external_inputs, "capture_excel_oracle", _mock_oracle)
    capture = external_inputs.capture_external_inputs(workflow)
    assert capture["status"] == "pass", capture
    older_discovery = discover_active_trace(workflow)
    assert older_discovery["status"] == "pass", older_discovery
    older_trace_path = Path(older_discovery["active_trace"])
    discovery = discover_active_trace(workflow)
    assert discovery["status"] == "pass", discovery
    current_trace_path = Path(discovery["active_trace"])
    assert older_trace_path != current_trace_path
    assert discovery["native_excel_called"] is False
    assert discovery["formula_cache_inputs"] is False

    _, manifest = load_workflow(workflow)
    stage3_revision = manifest["stages"]["3"]["revisions"][-1]
    evidence = workflow / "evidence"
    plan_path = evidence / "semantic_plan.json"
    check_path = evidence / "semantic_check.json"
    profile_path = evidence / "candidate_profile.json"
    input_path = evidence / "implementation.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    cell_ranges = [
        {"variable_id": variable_id, "source_ref": extent["ref"], "role": extent["role"],
         "coordinate_mapping": extent["coordinate_mapping"]}
        for variable_id, variable in plan["variables"].items()
        for extent in variable.get("source_extents", [])
    ]
    implementation = {
        "schema_version": "step4.active_implementation.v1",
        "binding": {
            "active_trace_sha256": _sha(Path(discovery["active_trace"])),
            "candidate_trace_sha256": _sha(evidence / "candidate_trace.json"),
            "external_capture_sha256": _sha(Path(capture["artifact"])),
            "semantic_check_sha256": _sha(check_path),
            "semantic_plan_sha256": _sha(plan_path),
            "source_profile_sha256": _sha(profile_path),
            "source_sha256": _sha(source),
            "stage3_artifact_sha256": stage3_revision["artifact"]["json_sha256"],
            "stage3_revision": stage3_revision["revision"],
        },
        "reference_resolution": {
            "cell_ranges": cell_ranges,
            "named_references": [{"name": "GP", "scope": "workbook", "kind": "scalar",
                                  "source_ref": "Premium!J1", "variable_id": "model.answer"}],
        },
        "execution_plan": {"passes": [{"pass_id": "pass.answer", "module_id": "module.answer",
                                         "order": 0, "source_family_ids": ["source.answer"],
                                         "variable_ids": ["model.answer"]}]},
    }
    input_path.write_text(json.dumps(implementation, indent=2), encoding="utf-8")
    preflight = create_implementation_plan(workflow, input_path)
    assert preflight["status"] == "pass", preflight

    workflow_before_stale_trace = (workflow / "workflow.json").read_bytes()
    stale_generation = generate_model(workflow, older_trace_path)
    assert stale_generation["status"] == "blocked"
    assert "differs from the active trace bound by the current implementation preflight" in stale_generation["reason"]
    assert (workflow / "workflow.json").read_bytes() == workflow_before_stale_trace
    assert not (workflow / "stage4" / "revision-0001" / "generation_report.json").exists()
    assert not (workflow / "stage4" / "bundles" / "revision-0001" / "bundle").exists()

    generated = generate_model(workflow)
    assert generated["status"] == "pass", generated
    bundle = Path(generated["bundle"]["path"])
    assert (bundle / "modular_runtime.py").is_file()
    assert (bundle / "pricing.py").is_file()
    assert not (bundle / "runtime.py").exists()
    bundle_manifest = json.loads((bundle / "model_manifest.json").read_text(encoding="utf-8"))
    raw_values = json.loads((bundle / "source_values.json").read_text(encoding="utf-8"))
    assert bundle_manifest["execution_kind"] == "modular"
    assert raw_values == {"main!A1": 3}
    assert bundle_manifest["formula_cache_inputs"] is False
    assert _sha(source) == bundle_manifest["source_sha256"]
    report = json.loads((workflow / generated["artifact"]["json"]).read_text(encoding="utf-8"))
    assert report["smoke_execution"] == "pass"
    assert report["smoke"]["targets"]["GP"] == 6
    assert report["formula_cache_inputs"] is False

    for reviewer in ("agent", "human"):
        from excel_to_act.steps.conversion_workflow import record_decision
        record_decision(workflow, 4, reviewer, "approve", "Synthetic modular generation fixture.")
    validation = validate_generated(workflow)
    assert validation["status"] == "pass", validation
    assert validation["result_summary"]["targets"]["GP"] == 6


def test_external_capture_rejects_changed_catalog_pair_before_native_call(tmp_path: Path, monkeypatch) -> None:
    workflow, _source = _external_workflow(tmp_path)
    monkeypatch.setattr(external_inputs, "capture_excel_oracle",
                        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("must not call Excel")))
    catalog = workflow / "stage3" / "revision-0001" / "input_boundary.json"
    catalog.write_text(catalog.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    result = external_inputs.capture_external_inputs(workflow)
    assert result["status"] == "blocked"
    assert "catalog artifact pair" in result["reason"]


def test_step4_capture_external_cli_has_no_user_range_override(tmp_path: Path, monkeypatch) -> None:
    workflow, _source = _external_workflow(tmp_path)
    monkeypatch.setattr(external_inputs, "capture_excel_oracle", _mock_oracle)
    result = CliRunner().invoke(app, ["step4", "capture-external", "--workflow", str(workflow)])
    assert result.exit_code == 0, result.output
    payload = json.loads(result.output)
    assert payload["status"] == "pass"
    assert payload["external_coordinate_count"] == 530
