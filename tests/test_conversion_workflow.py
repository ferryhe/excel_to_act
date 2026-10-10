from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from typer.testing import CliRunner

from excel_to_act.interfaces.cli import app
from excel_to_act.steps import conversion_workflow as workflow
from excel_to_act.steps.step4.generator import _emit, generate_model
from excel_to_act.steps.step4.runtime import Runtime
from excel_to_act.steps.step3.calculation import _Parser
from excel_to_act.steps.step3 import input_boundary


def _delegated_stage3(root: Path, tmp_path: Path, *, mode: str = "production") -> tuple[Path, dict]:
    source = {"workbook_sha256": "a" * 64, "source_id": "source-1", "run_id": "run-1"}
    if mode == "synthetic_fixture":
        source["mode"] = mode
    workflow_root, _ = workflow.create_workflow(root, source)
    if mode == "synthetic_fixture":
        _, manifest = workflow.load_workflow(workflow_root)
        manifest["mode"] = mode
        workflow._write_manifest(workflow_root, manifest)
    stage1 = _approved_stage1(workflow_root)
    _approve(workflow_root, 1)
    stage2 = _append(workflow_root, 2, {"status": "pass", "stage": 2}, "Stage 2",
                     input_stages={1: stage1["artifact"]})
    _approve(workflow_root, 2)
    stage3 = _append(workflow_root, 3, {"status": "ready_for_review", "stage": 3}, "Stage 3",
                     input_stages={2: stage2["artifact"]})
    auth = {"schema_version": "workflow.review_delegation.authorization.v1",
            "authorized_by": "human", "authorized_at": "2026-10-09T04:47:19Z",
            "user_instruction_translation": "Use TypeSafe for this workflow.",
            "workflow": str(workflow_root), "source_sha256": "a" * 64,
            "source_id": "source-1", "run_id": "run-1", "delegate": "typesafe",
            "requested_model": "jev-latest", "stages": [3, 4, 5, 6],
            "scope": "Saved scenario only.", "minimum_approval_probability": 0.85,
            "default_policy_restored_for_other_workflows": True}
    auth_path = tmp_path / "authorization.json"
    auth_path.write_text(json.dumps(auth), encoding="utf-8")
    return workflow_root, {"stage3": stage3, "auth_path": auth_path}


def _typesafe_request(root: Path, *, binding_override: dict | None = None) -> dict:
    _, manifest = workflow.load_workflow(root)
    revision = workflow.current_revision(manifest, 3)
    artifact = revision["artifact"]
    binding = {"stage": 3, "revision": revision["revision"], "source_sha256": "a" * 64,
               "artifact": {"json_sha256": artifact["json_sha256"], "md_sha256": artifact["md_sha256"]}}
    if binding_override:
        binding.update(binding_override)
    request = {"model": "jev-latest", "state": {"artifact_binding": binding},
               "questions": {"stage_decision": {"type": "choice", "instructions": "Review this checkpoint.",
                                                   "criteria": {"approve": "Accept", "reject": "Return for revision",
                                                                "insufficient": "Evidence is insufficient"}}}}
    return request


def _write_typesafe_review(root: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, *,
                           choice: str = "approve", probability: float = 0.95,
                           binding_override: dict | None = None,
                           probabilities_override: dict[str, float] | None = None) -> Path:
    tmp_path.mkdir(parents=True, exist_ok=True)
    request = _typesafe_request(root, binding_override=binding_override)
    choices = request["questions"]["stage_decision"]["criteria"]
    other_probability = (1.0 - probability) / max(len(choices) - 1, 1)
    response = {"model": "jev-1.13.0", "answers": {"stage_decision": {
        "type": "choice", "choice": choice,
        "probabilities": probabilities_override or {
            option: probability if option == choice else other_probability for option in choices}}},
                "usage": {"input_tokens": 321, "output_tokens": 27}}
    request_path = tmp_path / "request-input.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")

    def fake_post(outbound: dict, api_key: str) -> tuple[int, bytes]:
        assert api_key == "test-typesafe-key"
        assert outbound["state"]["artifact_binding"] == request["state"]["artifact_binding"]
        assert outbound["state"]["client_correlation_id"]
        return 200, json.dumps(response).encode("utf-8")

    monkeypatch.setenv("TYPESAFE_API_KEY", "test-typesafe-key")
    monkeypatch.setattr(workflow, "_post_typesafe", fake_post)
    result = workflow.submit_typesafe_review(root, 3, request_path)
    return Path(result["review"])


def _rewrite_live_review_probabilities(root: Path, evidence_path: Path,
                                       probabilities: dict[str, float]) -> None:
    evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
    response_path = evidence_path.parent / evidence["response"]["path"]
    response = json.loads(response_path.read_text(encoding="utf-8"))
    response["answers"]["stage_decision"]["probabilities"] = probabilities
    response_path.write_text(json.dumps(response), encoding="utf-8")
    response_sha = workflow.hash_file(response_path)

    _, manifest = workflow.load_workflow(root)
    invocation = next(item for item in manifest["typesafe_invocations"]
                      if item["invocation_id"] == evidence["invocation_id"])
    call_path = root / invocation["call"]["path"]
    call = json.loads(call_path.read_text(encoding="utf-8"))
    call["response"]["sha256"] = response_sha
    call_path.write_text(json.dumps(call), encoding="utf-8")
    call_sha = workflow.hash_file(call_path)

    evidence["response"]["sha256"] = response_sha
    evidence["call"]["sha256"] = call_sha
    evidence_path.write_text(json.dumps(evidence), encoding="utf-8")
    invocation["response"]["sha256"] = response_sha
    invocation["call"]["sha256"] = call_sha
    invocation["review"]["sha256"] = workflow.hash_file(evidence_path)
    workflow._write_manifest(root, manifest)


def _write_handcrafted_typesafe_review(root: Path, tmp_path: Path) -> Path:
    _, manifest = workflow.load_workflow(root)
    revision = workflow.current_revision(manifest, 3)
    artifact = revision["artifact"]
    request = _typesafe_request(root)
    response = {"model": "jev-1.13.0", "answers": {"stage_decision": {
        "type": "choice", "choice": "approve", "probabilities": {"approve": 1.0}}}}
    request_path, response_path = tmp_path / "request.json", tmp_path / "response.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")
    response_path.write_text(json.dumps(response), encoding="utf-8")
    evidence = {"schema_version": "conversion.typesafe_review.v1", "stage": 3,
                "revision": revision["revision"], "source_sha256": "a" * 64,
                "artifact": {"json_sha256": artifact["json_sha256"], "md_sha256": artifact["md_sha256"]},
                "service": "TypeSafe", "requested_model": "jev-latest", "resolved_model": "jev-1.13.0",
                "question_id": "stage_decision", "decision": "approve",
                "request": {"path": request_path.name, "sha256": workflow.hash_file(request_path)},
                "response": {"path": response_path.name, "sha256": workflow.hash_file(response_path)}}
    evidence_path = tmp_path / "review.json"
    evidence_path.write_text(json.dumps(evidence), encoding="utf-8")
    return evidence_path


def _append(root: Path, stage: int, payload: dict, markdown: str, *, input_stages=None):
    return workflow.append_stage_artifact(
        root,
        stage,
        f"stage{stage}.json",
        payload,
        markdown,
        input_stages=input_stages,
    )


def _approve(root: Path, stage: int) -> None:
    for reviewer in ("agent", "human"):
        workflow.record_decision(root, stage, reviewer, "approve", f"Reviewed stage {stage}.")


def _approved_stage1(root: Path, payload: dict | None = None, markdown: str = "Stage 1") -> dict:
    return _append(root, 1, payload or {"status": "pass", "stage": 1}, markdown)


def test_human_approval_requires_matching_agent_first_at_every_stage(tmp_path: Path) -> None:
    root, _manifest = workflow.create_workflow(
        tmp_path / "workflow", {"workbook_sha256": "a" * 64, "source_id": "source-1", "run_id": "run-1"}
    )
    _, manifest = workflow.load_workflow(root)
    manifest["mode"] = "synthetic_fixture"
    workflow._write_manifest(root, manifest)

    prior = None
    for stage in workflow.STAGES:
        artifact = _append(
            root, stage, {"status": "ready_for_review", "stage": stage}, f"Stage {stage}",
            input_stages={stage - 1: prior["artifact"]} if prior is not None else None,
        )
        before = (root / "workflow.json").read_bytes()
        with pytest.raises(ValueError, match="Agent must approve this current JSON/Markdown revision"):
            workflow.record_decision(root, stage, "human", "approve", f"Human reviewed stage {stage}.")
        assert (root / "workflow.json").read_bytes() == before

        workflow.record_decision(root, stage, "agent", "approve", f"Agent reviewed stage {stage}.")
        workflow.record_decision(root, stage, "human", "approve", f"Human reviewed stage {stage}.")
        assert workflow.stage_status(*workflow.load_workflow(root), stage)["status"] in {
            "approved", "simulated_approved",
        }
        prior = artifact


def test_human_rejection_is_allowed_before_agent_review(tmp_path: Path) -> None:
    root, _manifest = workflow.create_workflow(
        tmp_path / "workflow", {"workbook_sha256": "a" * 64, "source_id": "source-1", "run_id": "run-1"}
    )
    _approved_stage1(root)
    receipt = workflow.record_decision(root, 1, "human", "reject", "Revise the checkpoint.", return_to=1)

    assert receipt["decision"] == "reject"
    assert workflow.stage_status(*workflow.load_workflow(root), 1)["status"] == "rejected"
    with pytest.raises(ValueError, match="already has a rejection"):
        workflow.record_decision(root, 1, "agent", "approve", "Approved despite rejection.")


def test_legacy_json_only_binding_rejects_identical_repackaging(tmp_path: Path) -> None:
    root = tmp_path / "workflow"
    workflow.create_workflow(root, {"workbook_sha256": "source-sha", "run_id": "run-1"})
    stage1 = _approved_stage1(root)
    _approve(root, 1)
    _append(root, 2, {"status": "pass", "stage": 2}, "Stage 2",
            input_stages={1: stage1["artifact"]["json_sha256"]})
    _approve(root, 2)

    # Reproduce the old manifest format while preserving both accepted report files.
    _, manifest = workflow.load_workflow(root)
    manifest["stages"]["2"]["revisions"][0]["input_stages"] = {
        "1": stage1["artifact"]["json_sha256"]
    }
    workflow._write_manifest(root, manifest)
    _approve(root, 2)  # The legacy record itself remains reviewable before repackaging.
    assert workflow.stage_status(*workflow.load_workflow(root), 2)["status"] == "approved"

    _append(root, 1, {"status": "pass", "stage": 1}, "Stage 1")
    state = workflow.stage_status(*workflow.load_workflow(root), 2)
    assert state["status"] == "stale"
    assert "stage 1 artifact changed" in state["reasons"]


def test_paired_binding_invalidates_downstream_on_markdown_revision(tmp_path: Path) -> None:
    root = tmp_path / "workflow"
    workflow.create_workflow(root, {"workbook_sha256": "source-sha", "run_id": "run-1"})
    stage1 = _approved_stage1(root)
    _approve(root, 1)
    _append(root, 2, {"status": "pass", "stage": 2}, "Stage 2",
            input_stages={1: stage1["artifact"]})
    _approve(root, 2)

    repackaged = _append(root, 1, {"status": "pass", "stage": 1}, "Stage 1 with clarified report")
    assert repackaged["artifact"]["json_sha256"] == stage1["artifact"]["json_sha256"]
    assert repackaged["artifact"]["md_sha256"] != stage1["artifact"]["md_sha256"]

    state = workflow.stage_status(*workflow.load_workflow(root), 2)
    assert state["status"] == "stale"
    assert "stage 1 artifact changed" in state["reasons"]


def test_legacy_binding_to_different_noncurrent_json_revision_is_stale(tmp_path: Path) -> None:
    root = tmp_path / "workflow"
    workflow.create_workflow(root, {"workbook_sha256": "source-sha", "run_id": "run-1"})
    first = _approved_stage1(root, {"status": "pass", "stage": 1, "scope": "old"})
    _approve(root, 1)
    _append(root, 2, {"status": "pass", "stage": 2}, "Stage 2",
            input_stages={1: first["artifact"]["json_sha256"]})
    _, manifest = workflow.load_workflow(root)
    manifest["stages"]["2"]["revisions"][0]["input_stages"] = {
        "1": first["artifact"]["json_sha256"]
    }
    workflow._write_manifest(root, manifest)
    _approve(root, 2)

    _append(root, 1, {"status": "pass", "stage": 1, "scope": "new"}, "Stage 1 revised")
    state = workflow.stage_status(*workflow.load_workflow(root), 2)
    assert state["status"] == "stale"
    assert "stage 1 artifact changed" in state["reasons"]


def test_step4_generation_never_deletes_preexisting_revision_directory(tmp_path: Path) -> None:
    root = tmp_path / "workflow"
    workflow.create_workflow(root, {"workbook_sha256": "source-sha", "run_id": "run-1"})
    stage1 = _approved_stage1(root)
    _approve(root, 1)
    stage2 = _append(root, 2, {"status": "pass", "stage": 2}, "Stage 2",
                     input_stages={1: stage1["artifact"]})
    _approve(root, 2)
    _append(root, 3, {"status": "ready_for_review", "stage": 3}, "Stage 3",
            input_stages={2: stage2["artifact"]})
    _approve(root, 3)

    existing = root / "stage4" / "revision-0001"
    existing.mkdir(parents=True)
    sentinel = existing / "keep.txt"
    sentinel.write_text("preserve this unrelated directory", encoding="utf-8")

    result = generate_model(root, tmp_path / "missing-trace.json")
    assert result["status"] == "blocked"
    assert "already exists" in result["reason"]
    assert sentinel.read_text(encoding="utf-8") == "preserve this unrelated directory"


def test_stage4_rejection_to_stage3_invalidates_old_stage4_approval(tmp_path: Path) -> None:
    root = tmp_path / "workflow"
    workflow.create_workflow(root, {"workbook_sha256": "source-sha", "run_id": "run-1"})
    stage1 = _approved_stage1(root)
    _approve(root, 1)
    stage2 = _append(root, 2, {"status": "pass", "stage": 2}, "Stage 2",
                     input_stages={1: stage1["artifact"]})
    _approve(root, 2)
    stage3 = _append(root, 3, {"status": "ready_for_review", "stage": 3}, "Stage 3",
                     input_stages={2: stage2["artifact"]})
    _approve(root, 3)
    _append(root, 4, {"status": "ready_for_review", "stage": 4}, "Stage 4",
            input_stages={3: stage3["artifact"]})
    _approve(root, 4)

    workflow.record_decision(root, 4, "agent", "reject", "Revise the design first.", return_to=3)
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["status"] == "stale"
    with pytest.raises(ValueError, match="stage 4 is stale"):
        workflow.require_stage_approved(root, 4)

    revised_stage3 = _append(root, 3, {"status": "ready_for_review", "stage": 3, "revision": 2},
                             "Revised Stage 3", input_stages={2: stage2["artifact"]})
    _approve(root, 3)
    assert revised_stage3["resolves_rejections"]
    state = workflow.stage_status(*workflow.load_workflow(root), 4)
    assert state["status"] == "stale"
    assert "stage 3 artifact changed" in state["reasons"]
    with pytest.raises(ValueError, match="stage 4 is stale"):
        workflow.require_stage_approved(root, 4)


def test_scoped_typesafe_delegation_replaces_only_stage3_human_receipt(tmp_path: Path,
                                                                      monkeypatch: pytest.MonkeyPatch) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["reviewer_pair"] == ["agent", "human"]
    with pytest.raises(ValueError, match="no active TypeSafe delegation"):
        workflow.record_decision(root, 3, "typesafe", "approve", "reviewed")

    receipt = workflow.register_review_delegation(root, paths["auth_path"])
    assert receipt["stages"] == [3, 4, 5, 6]
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["reviewer_pair"] == ["agent", "typesafe"]
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")
    evidence = _write_typesafe_review(root, tmp_path, monkeypatch)
    review = workflow.record_decision(root, 3, "typesafe", "approve", "Jev approval.", evidence_path=evidence)

    state = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert state["status"] == "approved"
    assert state["human"] is None
    assert state["typesafe"]["decision"] == "approve"
    assert review["reviewer"] == "typesafe"
    workflow.require_stage_approved(root, 3)
    with pytest.raises(ValueError, match="only available for delegated production stages 3..6"):
        workflow.record_decision(root, 2, "typesafe", "approve", "not permitted")


def test_typesafe_approval_requires_current_agent_receipt_without_writing(tmp_path: Path,
                                                                         monkeypatch: pytest.MonkeyPatch) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Agent reviewed Stage 3.")
    evidence = _write_typesafe_review(root, tmp_path / "typesafe", monkeypatch)

    _, manifest = workflow.load_workflow(root)
    revision = workflow.current_revision(manifest, 3)
    revision["decisions"] = [item for item in revision["decisions"] if item.get("reviewer") != "agent"]
    workflow._write_manifest(root, manifest)
    before = (root / "workflow.json").read_bytes()

    with pytest.raises(ValueError, match="Agent must approve this current JSON/Markdown revision"):
        workflow.record_decision(root, 3, "typesafe", "approve", "Jev approval.", evidence_path=evidence)
    assert (root / "workflow.json").read_bytes() == before


def test_human_override_is_bound_to_one_delegated_step5_revision(tmp_path: Path) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    _, before = workflow.load_workflow(root)
    delegation_record = before["review_delegation"]

    for stage in (3, 4):
        if stage == 4:
            prior = workflow.current_revision(workflow.load_workflow(root)[1], 3)
            _append(root, 4, {"status": "ready_for_review", "stage": 4}, "Stage 4",
                    input_stages={3: prior["artifact"]})
        workflow.record_decision(root, stage, "agent", "approve", f"Agent reviewed stage {stage}.")
        workflow.record_decision(root, stage, "human", "approve", f"Human reviewed stage {stage}.")

    stage4 = workflow.current_revision(workflow.load_workflow(root)[1], 4)
    stage5 = _append(root, 5, {"status": "ready_for_review", "stage": 5}, "Step 5 validation report",
                     input_stages={4: stage4["artifact"]})
    assert workflow.stage_status(*workflow.load_workflow(root), 5)["reviewer_pair"] == ["agent", "typesafe"]
    workflow_before_confirmation = (root / "workflow.json").read_bytes()
    confirmation = CliRunner().invoke(app, [
        "workflow", "confirm", "--workflow", str(root), "--stage", "5", "--reviewer", "human",
        "--decision", "approve", "--message", "I accept Step 5 report revision 1.",
    ])
    assert confirmation.exit_code != 0
    assert "Agent must approve this current JSON/Markdown revision" in confirmation.stdout
    assert (root / "workflow.json").read_bytes() == workflow_before_confirmation
    workflow.record_decision(root, 5, "agent", "approve", "Agent reviewed Step 5 revision 1.")
    confirmation = CliRunner().invoke(app, [
        "workflow", "confirm", "--workflow", str(root), "--stage", "5", "--reviewer", "human",
        "--decision", "approve", "--message", "I accept Step 5 report revision 1.",
    ])
    assert confirmation.exit_code == 0, confirmation.stdout
    confirmation_payload = json.loads(confirmation.stdout)
    assert confirmation_payload["status"] == "pass"
    assert confirmation_payload["receipt"]["reviewer_pair_override"]["authorization_sha256"] == delegation_record["authorization_sha256"]
    pending = workflow.stage_status(*workflow.load_workflow(root), 5)
    assert pending["status"] == "approved"
    assert pending["reviewer_pair"] == ["agent", "human"]
    assert pending["agent"]["decision"] == "approve"
    human_receipt = pending["human"]["reviewer_pair_override"]
    assert human_receipt == {
        "schema_version": "workflow.review_pair_override.v1",
        "authorization_sha256": delegation_record["authorization_sha256"],
        "reviewer_pair": ["agent", "human"],
    }
    _, after = workflow.load_workflow(root)
    assert after["review_delegation"] == delegation_record

    _append(root, 6, {"status": "ready_for_review", "stage": 6}, "Stage 6",
             input_stages={5: stage5["artifact"]})
    assert workflow.stage_status(*workflow.load_workflow(root), 6)["reviewer_pair"] == ["agent", "typesafe"]

    revised = _append(root, 5, {"status": "ready_for_review", "stage": 5, "revision": 2},
                      "Revised Step 5 validation report", input_stages={4: stage4["artifact"]})
    replacement = workflow.stage_status(*workflow.load_workflow(root), 5)
    assert replacement["revision"] == revised["revision"]
    assert replacement["status"] == "awaiting_confirmation"
    assert replacement["reviewer_pair"] == ["agent", "typesafe"]
    assert replacement["human"] is None

    _append(root, 4, {"status": "ready_for_review", "stage": 4, "revision": 2}, "Revised Stage 4",
            input_stages={3: workflow.current_revision(workflow.load_workflow(root)[1], 3)["artifact"]})
    stale = workflow.stage_status(*workflow.load_workflow(root), 5)
    assert stale["status"] == "stale"
    assert "stage 4 artifact changed" in stale["reasons"]
    assert workflow.stage_status(*workflow.load_workflow(root), 6)["status"] == "stale"


def test_pre_delegation_human_receipt_does_not_restore_pair(tmp_path: Path) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.record_decision(root, 3, "agent", "approve", "Agent reviewed Stage 3.")
    workflow.record_decision(root, 3, "human", "approve", "I reviewed Stage 3 before delegation.")
    workflow.register_review_delegation(root, paths["auth_path"])

    state = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert state["status"] == "awaiting_confirmation"
    assert state["reviewer_pair"] == ["agent", "typesafe"]
    assert "reviewer_pair_override" not in state["human"]


def test_delegated_human_rejection_requires_replacement_revision(tmp_path: Path) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Agent reviewed Stage 3.")
    workflow.record_decision(root, 3, "human", "reject", "Revise the design.", return_to=3)
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["status"] == "rejected"

    with pytest.raises(ValueError, match="already has a rejection"):
        workflow.record_decision(root, 3, "human", "approve", "I accept the same revision.")
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["status"] == "rejected"

    _, manifest = workflow.load_workflow(root)
    stage2 = workflow.current_revision(manifest, 2)
    revised = _append(root, 3, {"status": "ready_for_review", "stage": 3, "revision": 2},
                      "Replacement Stage 3", input_stages={2: stage2["artifact"]})
    state = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert state["revision"] == revised["revision"]
    assert state["status"] == "awaiting_confirmation"
    assert state["reviewer_pair"] == ["agent", "typesafe"]

    before_approval = (root / "workflow.json").read_bytes()
    with pytest.raises(ValueError, match="Agent must approve this current JSON/Markdown revision"):
        workflow.record_decision(root, 3, "human", "approve", "I accept the replacement revision.")
    assert (root / "workflow.json").read_bytes() == before_approval
    workflow.record_decision(root, 3, "agent", "approve", "Agent reviewed the replacement revision.")
    workflow.record_decision(root, 3, "human", "approve", "I accept the replacement revision.")
    approved = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert approved["status"] == "approved"
    assert approved["revision"] == revised["revision"]


def test_delegated_agent_rejection_requires_replacement_revision(tmp_path: Path) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "reject", "Revise the design.", return_to=3)
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["status"] == "rejected"

    with pytest.raises(ValueError, match="already has a rejection"):
        workflow.record_decision(root, 3, "agent", "approve", "I accept the same revision.")
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["status"] == "rejected"

    _, manifest = workflow.load_workflow(root)
    revised = _append(root, 3, {"status": "ready_for_review", "stage": 3, "revision": 2},
                      "Replacement Stage 3", input_stages={
                          2: workflow.current_revision(manifest, 2)["artifact"]})
    with pytest.raises(ValueError, match="Agent must approve this current JSON/Markdown revision"):
        workflow.record_decision(root, 3, "human", "approve", "I accept the replacement revision.")
    workflow.record_decision(root, 3, "agent", "approve", "Agent reviewed the replacement revision.")
    workflow.record_decision(root, 3, "human", "approve", "I accept the replacement revision.")
    approved = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert approved["status"] == "approved"
    assert approved["revision"] == revised["revision"]


def test_delegated_typesafe_rejection_requires_replacement_revision(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Agent reviewed Stage 3.")
    approval_evidence = _write_typesafe_review(root, tmp_path / "approve", monkeypatch, choice="approve")
    rejection_evidence = _write_typesafe_review(root, tmp_path / "reject", monkeypatch, choice="reject")
    workflow.record_decision(root, 3, "typesafe", "reject", "Revise the design.", return_to=3,
                             evidence_path=rejection_evidence)
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["status"] == "rejected"

    with pytest.raises(ValueError, match="already has a rejection"):
        workflow.record_decision(root, 3, "typesafe", "approve", "Approve the same revision.",
                                 evidence_path=approval_evidence)
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["status"] == "rejected"


def test_real_human_receipt_confirms_input_catalog_even_with_active_delegation(tmp_path: Path) -> None:
    root, _paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, tmp_path / "authorization.json")
    _, manifest = workflow.load_workflow(root)
    stage2 = workflow.current_revision(manifest, 2)
    catalog_revision = _append(root, 3,
                               {"schema_version": "step3.input_boundary.v1", "phase": "input_boundary",
                                "status": "ready_for_review"},
                               "Source-bound input catalog", input_stages={2: stage2["artifact"]})

    workflow.record_decision(root, 3, "agent", "approve", "Catalog structure reviewed.")
    pending = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert pending["reviewer_pair"] == ["agent", "human"]
    assert pending["status"] == "awaiting_confirmation"
    assert pending["human"] is None

    workflow.record_decision(root, 3, "human", "approve", "I confirm these input assumptions.")
    state = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert state["status"] == "input_boundary_confirmed"
    assert state["reviewer_pair"] == ["agent", "human"]
    assert state["human"]["decision"] == "approve"
    assert "reviewer_pair_override" not in state["human"]
    assert state["typesafe"] is None

    status = workflow.workflow_status(root)
    assert status["input_boundary"]["revision"] == catalog_revision["revision"]
    assert status["input_boundary"]["status"] == "input_boundary_confirmed"
    assert status["next_permitted_stage"] == 3
    with pytest.raises(ValueError, match="does not release Step 4"):
        workflow.require_stage_approved(root, 3)


def test_input_catalog_requires_human_even_with_valid_delegated_typesafe_recommendation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    _, manifest = workflow.load_workflow(root)
    stage2 = workflow.current_revision(manifest, 2)
    boundary_revision = _append(root, 3,
                                {"schema_version": "step3.input_boundary.v1", "phase": "input_boundary",
                                 "status": "ready_for_review", "boundary_sha256": "boundary-sha"},
                                "Source-bound input catalog", input_stages={2: stage2["artifact"]})
    _, manifest = workflow.load_workflow(root)
    manifest["input_boundary_required"] = True
    manifest["input_boundary"] = {"revision": boundary_revision["revision"],
                                  "artifact": boundary_revision["artifact"],
                                  "boundary_sha256": "boundary-sha"}
    workflow._write_manifest(root, manifest)
    workflow.record_decision(root, 3, "agent", "approve", "Catalog structure reviewed.")
    evidence = _write_typesafe_review(root, tmp_path / "typesafe", monkeypatch)
    workflow.record_decision(root, 3, "typesafe", "approve", "Optional delegated recommendation.",
                             evidence_path=evidence)

    state = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert state["reviewer_pair"] == ["agent", "human"]
    assert state["status"] == "awaiting_confirmation"
    assert state["agent"]["decision"] == "approve"
    assert state["human"] is None
    assert state["typesafe"]["reviewer"] == "typesafe"
    assert state["typesafe"]["decision"] == "approve"
    assert state["typesafe"]["evidence_status"] == "current"
    with pytest.raises(ValueError, match="input boundary is awaiting_confirmation"):
        input_boundary.require_input_boundary_confirmed(root, workflow.load_workflow(root)[1])
    with pytest.raises(ValueError, match="stage 3 is awaiting_confirmation"):
        workflow.require_stage_approved(root, 3)

    analysis = {"binding": {"workflow_path": str(root), "source": {
        "source_id": "source-1", "run_id": "run-1", "source_sha256": "a" * 64}}}
    with pytest.raises(ValueError, match="input boundary is awaiting_confirmation"):
        input_boundary.require_confirmed_boundary_for_analysis(tmp_path, analysis, analysis["binding"])

    workflow.record_decision(root, 3, "human", "approve", "I confirm this input catalog.")
    state = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert state["status"] == "input_boundary_confirmed"
    assert state["reviewer_pair"] == ["agent", "human"]
    input_boundary.require_input_boundary_confirmed(root, workflow.load_workflow(root)[1])
    with pytest.raises(ValueError, match="does not release Step 4"):
        workflow.require_stage_approved(root, 3)


def test_typesafe_evidence_must_match_current_report_and_mutation_stales_receipt(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")

    bad_dir = tmp_path / "bad"
    bad_dir.mkdir()
    # A changed report binding must not import, even when the packet and response hashes are current.
    with pytest.raises(ValueError, match="current source and artifact revision"):
        _write_typesafe_review(root, tmp_path / "bad", monkeypatch, binding_override={
            "artifact": {"json_sha256": "b" * 64, "md_sha256": "c" * 64}})

    good_dir = tmp_path / "good"
    good_dir.mkdir()
    evidence = _write_typesafe_review(root, good_dir, monkeypatch)
    workflow.record_decision(root, 3, "typesafe", "approve", "Jev approval.", evidence_path=evidence)
    response = evidence.parent / "response.json"
    response.write_text(response.read_text(encoding="utf-8").replace("0.95", "0.96"), encoding="utf-8")
    state = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert state["status"] == "stale"
    assert "response file changed" in state["reasons"][0]


def test_typesafe_insufficient_and_low_probability_do_not_approve(tmp_path: Path,
                                                                  monkeypatch: pytest.MonkeyPatch) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")

    insufficient_dir = tmp_path / "insufficient"
    insufficient_dir.mkdir()
    insufficient = _write_typesafe_review(root, insufficient_dir, monkeypatch, choice="insufficient")
    with pytest.raises(ValueError, match="marked the evidence insufficient"):
        workflow.record_decision(root, 3, "typesafe", "approve", "not sufficient", evidence_path=insufficient)

    low_dir = tmp_path / "low"
    low_dir.mkdir()
    low = _write_typesafe_review(root, low_dir, monkeypatch, probability=0.8)
    with pytest.raises(ValueError, match="below the delegated confidence threshold"):
        workflow.record_decision(root, 3, "typesafe", "approve", "low confidence", evidence_path=low)


def test_delegation_rejects_wrong_source_and_synthetic_workflows(tmp_path: Path) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    auth = json.loads(paths["auth_path"].read_text(encoding="utf-8"))
    auth["source_sha256"] = "f" * 64
    paths["auth_path"].write_text(json.dumps(auth), encoding="utf-8")
    with pytest.raises(ValueError, match="source_sha256"):
        workflow.register_review_delegation(root, paths["auth_path"])

    simulated, simulated_paths = _delegated_stage3(tmp_path / "synthetic", tmp_path, mode="synthetic_fixture")
    with pytest.raises(ValueError, match="synthetic workflow"):
        workflow.register_review_delegation(simulated, simulated_paths["auth_path"])


def test_delegation_cannot_lower_approval_floor_below_point_eighty_five(tmp_path: Path) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    authorization = json.loads(paths["auth_path"].read_text(encoding="utf-8"))
    authorization["minimum_approval_probability"] = 0.84
    paths["auth_path"].write_text(json.dumps(authorization), encoding="utf-8")
    with pytest.raises(ValueError, match="between 0.85 and 1"):
        workflow.register_review_delegation(root, paths["auth_path"])


def test_typesafe_rejection_is_recorded_with_its_return_route(tmp_path: Path,
                                                               monkeypatch: pytest.MonkeyPatch) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")
    evidence = _write_typesafe_review(root, tmp_path, monkeypatch, choice="reject")
    receipt = workflow.record_decision(root, 3, "typesafe", "reject", "Return to design.",
                                       return_to=2, evidence_path=evidence)
    state = workflow.stage_status(*workflow.load_workflow(root), 3)
    assert receipt["return_to"] == 2
    assert state["status"] == "stale"
    _, manifest = workflow.load_workflow(root)
    recorded = manifest["stages"]["3"]["revisions"][0]["decisions"][-1]
    assert recorded["reviewer"] == "typesafe"
    assert recorded["return_to"] == 2


def test_handwritten_typesafe_packets_cannot_be_imported_as_live_reviews(tmp_path: Path) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")
    evidence = _write_handcrafted_typesafe_review(root, tmp_path)
    with pytest.raises(ValueError, match="no matching live CLI invocation"):
        workflow.record_decision(root, 3, "typesafe", "approve", "handwritten", evidence_path=evidence)


def test_typesafe_submission_requires_current_agent_and_delegation(tmp_path: Path,
                                                                   monkeypatch: pytest.MonkeyPatch) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    request_path = tmp_path / "request.json"
    request_path.write_text(json.dumps(_typesafe_request(root)), encoding="utf-8")
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-typesafe-key")
    called = False

    def fake_post(_outbound: dict, _api_key: str) -> tuple[int, bytes]:
        nonlocal called
        called = True
        return 500, b"{}"

    monkeypatch.setattr(workflow, "_post_typesafe", fake_post)
    with pytest.raises(ValueError, match="no active TypeSafe delegation"):
        workflow.submit_typesafe_review(root, 3, request_path)
    assert not called

    workflow.register_review_delegation(root, paths["auth_path"])
    with pytest.raises(ValueError, match="Agent approval"):
        workflow.submit_typesafe_review(root, 3, request_path)
    assert not called


def test_typesafe_cli_makes_mocked_live_call_and_logs_compact_provenance(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")
    request_path = tmp_path / "request.json"
    request_path.write_text(json.dumps(_typesafe_request(root)), encoding="utf-8")
    response = {"model": "jev-1.13.0", "answers": {"stage_decision": {
        "type": "choice", "choice": "approve",
        "probabilities": {"approve": 0.95, "reject": 0.03, "insufficient": 0.02}}},
                "usage": {"input_tokens": 321, "output_tokens": 27}}
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-typesafe-key")
    monkeypatch.setattr(workflow, "_post_typesafe",
                        lambda outbound, key: (200, json.dumps(response).encode("utf-8")))
    runner = CliRunner()
    result = runner.invoke(app, ["workflow", "typesafe", "--workflow", str(root),
                                 "--stage", "3", "--request", str(request_path)])
    assert result.exit_code == 0, result.stdout
    payload = json.loads(result.stdout)
    assert payload["status"] == "pass"
    assert payload["correlation_id_origin"] == "client-generated; not provider-issued"
    _, manifest = workflow.load_workflow(root)
    assert len(manifest["typesafe_invocations"]) == 1
    evidence = Path(payload["review"])
    review_value = json.loads(evidence.read_text(encoding="utf-8"))
    assert review_value["usage"] == {"input_tokens": 321, "output_tokens": 27}
    call_path = root / manifest["typesafe_invocations"][0]["call"]["path"]
    call_value = json.loads(call_path.read_text(encoding="utf-8"))
    assert call_value["usage"] == review_value["usage"]
    assert manifest["typesafe_invocations"][0]["usage"] == review_value["usage"]
    workflow.record_decision(root, 3, "typesafe", "approve", "Live Jev call.", evidence_path=evidence)
    assert workflow.stage_status(*workflow.load_workflow(root), 3)["status"] == "approved"


@pytest.mark.parametrize("probabilities", [
    {"approve": 0.95, "reject": 0.03, "insufficient": 0.01, "other": 0.01},
    {"approve": 0.95, "reject": 0.05},
])
def test_typesafe_live_call_rejects_probability_keys_not_in_request(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, probabilities: dict[str, float],
) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")

    with pytest.raises(ValueError, match="do not match the submitted choice criteria"):
        _write_typesafe_review(root, tmp_path / "review", monkeypatch,
                               probabilities_override=probabilities)

    _, manifest = workflow.load_workflow(root)
    assert manifest.get("typesafe_invocations", []) == []
    assert not (root / "typesafe").exists()


@pytest.mark.parametrize("probabilities", [
    {"approve": 0.95, "reject": 0.03, "insufficient": 0.01, "other": 0.01},
    {"approve": 0.95, "reject": 0.05},
])
def test_typesafe_import_rejects_probability_keys_not_in_request(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, probabilities: dict[str, float],
) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")
    evidence = _write_typesafe_review(root, tmp_path / "valid", monkeypatch)
    _rewrite_live_review_probabilities(root, evidence, probabilities)

    with pytest.raises(ValueError, match="do not match the submitted choice criteria"):
        workflow.record_decision(root, 3, "typesafe", "approve", "Jev approval.", evidence_path=evidence)


def test_failed_typesafe_http_call_does_not_create_review_or_invocation_log(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")
    request_path = tmp_path / "request.json"
    request_path.write_text(json.dumps(_typesafe_request(root)), encoding="utf-8")
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-typesafe-key")
    _, before = workflow.load_workflow(root)
    monkeypatch.setattr(workflow, "_post_typesafe",
                        lambda _request, _key: (503, b"service unavailable"))

    with pytest.raises(ValueError, match="HTTP 503"):
        workflow.submit_typesafe_review(root, 3, request_path)

    _, after = workflow.load_workflow(root)
    assert after == before
    assert after.get("typesafe_invocations", []) == []
    assert not (root / "typesafe").exists()


def test_typesafe_response_requires_and_retains_documented_usage(tmp_path: Path,
                                                                monkeypatch: pytest.MonkeyPatch) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    workflow.register_review_delegation(root, paths["auth_path"])
    workflow.record_decision(root, 3, "agent", "approve", "Independent review passed.")
    request_path = tmp_path / "request.json"
    request_path.write_text(json.dumps(_typesafe_request(root)), encoding="utf-8")
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-typesafe-key")
    response = {"model": "jev-1.13.0", "answers": {"stage_decision": {
        "type": "choice", "choice": "approve", "probabilities": {"approve": 1.0}}}}
    monkeypatch.setattr(workflow, "_post_typesafe",
                        lambda _request, _key: (200, json.dumps(response).encode("utf-8")))

    with pytest.raises(ValueError, match="required usage object"):
        workflow.submit_typesafe_review(root, 3, request_path)

    _, manifest = workflow.load_workflow(root)
    assert manifest.get("typesafe_invocations", []) == []
    assert not (root / "typesafe").exists()


def test_submit_typesafe_review_has_no_production_post_override() -> None:
    import inspect

    assert "post" not in inspect.signature(workflow.submit_typesafe_review).parameters


def test_workflow_delegate_cli_emits_compact_json_and_rejects_unbound_auth(tmp_path: Path) -> None:
    root, paths = _delegated_stage3(tmp_path / "workflow", tmp_path)
    runner = CliRunner()
    result = runner.invoke(app, ["workflow", "delegate", "--workflow", str(root),
                                 "--authorization", str(paths["auth_path"])])
    assert result.exit_code == 0, result.stdout
    payload = json.loads(result.stdout)
    assert payload["status"] == "pass"
    assert payload["stages"] == [3, 4, 5, 6]

    wrong = json.loads(paths["auth_path"].read_text(encoding="utf-8"))
    wrong["workflow"] = str(tmp_path / "other")
    wrong_path = tmp_path / "wrong.json"
    wrong_path.write_text(json.dumps(wrong), encoding="utf-8")
    rejected = runner.invoke(app, ["workflow", "delegate", "--workflow", str(root),
                                   "--authorization", str(wrong_path)])
    assert rejected.exit_code == 1
    assert json.loads(rejected.stdout)["status"] == "blocked"

def test_emitted_formula_literals_use_python_boolean_blank_and_error_values(tmp_path: Path) -> None:
    source_path = tmp_path / "literals.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["A1"] = "=IF(FALSE,1/0,7)"
    sheet["A2"] = "=IF(TRUE,,#N/A)"
    sheet["A3"] = "=IFERROR(#N/A,3)"
    workbook.save(source_path)
    workbook.close()
    source = load_workbook(source_path, data_only=False, read_only=True)
    runtime = Runtime({}, {}, {}, {}, set(), [])
    try:
        cases = {"A1": 7, "A2": None, "A3": 3}
        for address, expected in cases.items():
            formula = source["Main"][address].value
            expression = _emit(_Parser(formula, "Main").parse(), "Main", f"Main!{address}")
            code = compile(f"def evaluate(env):\n    return {expression}\n", "generated_formula.py", "exec")
            namespace: dict[str, object] = {}
            exec(code, namespace)
            assert namespace["evaluate"](runtime) == expected
    finally:
        source.close()
