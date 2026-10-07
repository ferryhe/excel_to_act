from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableColumn
from typer.main import get_command
from typer.testing import CliRunner

from excel_to_act.interfaces.cli import app
from excel_to_act.schemas import (
    ActiveXControl,
    ActiveXEvents,
    CheckboxBinding,
    CheckboxBindings,
    Step2DependencyAudit,
    Step2DependencySnapshot,
    Step2ReadingManifest,
)
from excel_to_act.steps.step2.workflow import build_index


def _write_json(path: Path, value: object) -> str:
    data = (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode()
    path.write_bytes(data)
    return hashlib.sha256(data).hexdigest()


def _setup(
    tmp_path: Path, *, controls: bool = False,
    structured_selectors: tuple[str, str] | None = None,
    structured_table: dict | None = None,
    defined_names: bool = False, multi_area_name: bool = False,
    no_references: bool = False,
) -> tuple[Path, Path, Path, dict, dict]:
    raw = tmp_path / "raw"
    raw.mkdir()
    workbook_path = raw / "book.xlsx"
    book = Workbook()
    book.active.title = "Main"
    book.active["A1"] = "input"
    book.active["A2"] = "=Calc!A1"
    book.active["A3"] = "=SUM(Calc!B1:B2)"
    calc = book.create_sheet("Calc")
    calc["A1"] = 10
    calc["B1"] = 20
    calc["B2"] = 30
    if structured_selectors:
        if structured_table is None:
            calc["A1"], calc["B1"] = "Amount", "Rate"
            calc["A2"], calc["B2"] = 10, 20
            calc["A3"], calc["B3"] = 30, 40
            calc.add_table(Table(displayName="CalcTable", ref="A1:B3"))
        else:
            for row, values in enumerate(structured_table["rows"], 1):
                for column, value in enumerate(values, 1):
                    calc.cell(row, column, value)
            table = Table(
                displayName="CalcTable",
                ref=structured_table["ref"],
                headerRowCount=structured_table["header_row_count"],
                totalsRowCount=structured_table["totals_row_count"],
                totalsRowShown=structured_table["totals_row_shown"],
            )
            table.tableColumns = [
                TableColumn(
                    id=position,
                    name=name,
                    totalsRowFunction=structured_table.get("totals_functions", [None] * len(structured_table["columns"]))[position - 1],
                )
                for position, name in enumerate(structured_table["columns"], 1)
            ]
            calc.add_table(table)
        book.active["A2"] = f"=SUM(CalcTable[{structured_selectors[0]}])"
        book.active["A3"] = f"=SUM(CalcTable[{structured_selectors[1]}])"
    if defined_names:
        book.active["A2"] = "=RelativeName"
        book.active["A3"] = "=1"
        book.defined_names.add(DefinedName("RelativeName", attr_text="'Calc'!A2"))
        book.defined_names.add(DefinedName("UnusedAbsolute", attr_text="'Calc'!$A$1"))
    if multi_area_name:
        book.active["A2"] = "=SUM(Both)"
        book.active["A3"] = "=1"
        book.defined_names.add(DefinedName("Both", attr_text="'Calc'!$A$1,'Calc'!$B$1"))
    if no_references:
        book.active["A2"], book.active["A3"] = "=1", "=2"
    book.save(workbook_path)
    book.close()

    from excel_to_act.steps.step1.workflow import convert_directory

    step1_root = tmp_path / "step1"
    converted = convert_directory(raw, step1_root)
    source_run = converted["entries"][0]["run_path"]
    handoff_path = step1_root / source_run / "handoff.json"
    handoff = json.loads(handoff_path.read_text(encoding="utf-8"))
    source_sha = handoff["source"]["sha256"]
    run_id = handoff["run_id"]
    if controls:
        source_part = "xl/worksheets/sheet1.xml"
        checkbox = CheckboxBindings(
            workbook_sha256=source_sha,
            status="partial",
            bindings=[CheckboxBinding(
                sheet="Main", sheet_part=source_part, shape_id="7", control_name="Enabled",
                linked_cell_raw="Calc!$C$3", linked_cell="Calc!$C$3", linked_sheet="Calc",
                linked_address="C3", binding_status="resolved", sources=["xl/ctrlProps/ctrlProp1.xml"],
            )],
        )
        activex = ActiveXEvents(
            workbook_sha256=source_sha,
            status="complete",
            controls=[ActiveXControl(
                sheet="Main", sheet_part=source_part, sheet_code_name="Sheet1", shape_id="8",
                control_name="Submit", part="xl/activeX/activeX1.xml", binary_part="xl/activeX/activeX1.bin",
                event_procedures=["Submit_Click"], binding_status="resolved",
            )],
        )
        for name, model in (("checkbox_bindings.json", checkbox), ("activex_events.json", activex)):
            path = step1_root / source_run / name
            digest = _write_json(path, model.model_dump(mode="json"))
            handoff["artifacts"] = [ref for ref in handoff["artifacts"] if ref["name"] != name]
            handoff["artifacts"].append({"name": name, "path": name, "sha256": digest})
        _write_json(handoff_path, handoff)

    batch_path = step1_root / converted["artifacts"][0]["path"]
    index_dir = tmp_path / "step2"
    built = build_index(batch_path, step1_root, index_dir)
    assert built["status"] != "blocked", built.get("diagnostics")
    index_path = index_dir / "index.json"
    index = json.loads(index_path.read_text(encoding="utf-8"))
    scope = {
        "schema_version": "analysis.scope.v1",
        "workbook_sha256": source_sha,
        "source_run_id": run_id,
        "retained_sheets": ["Main"],
        "ignored_sheets": ["Calc"],
        "dependency_policy": {
            "mode": "retain_read_only_source_values_for_inbound_references",
            "snapshot": "retained_dependency_cells.json",
            "audit": "ignored_sheet_dependency_audit.json",
        },
        "confirmation": {
            "decisions": [{
                "question_id": "analysis_scope:ignore_sheet:Calc",
                "decision": "ignore_internal_logic_keep_dependency_values",
                "reviewer": "user",
            }],
        },
    }
    scope_path = tmp_path / "scope.json"
    _write_json(scope_path, scope)
    return step1_root, index_path, scope_path, index, scope


def _prepare_args(root: Path, index: Path, out: Path, scope: Path | None = None) -> list[str]:
    args = ["step2", "prepare", "--index", str(index), "--step1-root", str(root), "--out", str(out)]
    if scope is not None:
        args.extend(["--scope", str(scope)])
    return args


def test_prepare_help_and_catalog_advertise_only_implemented_reader() -> None:
    runner = CliRunner()
    prepare = get_command(app).commands["step2"].commands["prepare"]
    registered_options = {option for parameter in prepare.params for option in parameter.opts}
    assert {"--index", "--step1-root", "--out", "--scope", "--resume", "--dry-run"} <= registered_options

    catalog = json.loads(runner.invoke(app, ["step2", "tools"]).stdout)
    assert [command["name"] for command in catalog["reader_commands"]] == ["step2.prepare", "step2.query"]


def test_prepare_exports_json_schemas_for_runtime_contracts() -> None:
    for file_name, model in (
        ("step2_reading.schema.json", Step2ReadingManifest),
        ("step2_dependency_audit.schema.json", Step2DependencyAudit),
        ("step2_dependency_snapshot.schema.json", Step2DependencySnapshot),
    ):
        schema = json.loads((Path(__file__).parents[1] / "schemas" / file_name).read_text(encoding="utf-8"))
        assert schema == model.model_json_schema()


def test_prepare_package_binds_scope_controls_and_cached_dependency_facts(tmp_path: Path) -> None:
    root, index_path, scope_path, index, scope = _setup(tmp_path, controls=True)
    out = tmp_path / "reading"
    state = index_path.parent / "state.json"
    state_before = state.read_bytes()
    inventory_path = root / next(ref["path"] for ref in index["entries"][0]["artifacts"] if ref["name"] == "inventory.json")
    coverage_before = json.loads(inventory_path.read_text(encoding="utf-8"))["coverage"]

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 0, result.stdout
    response = json.loads(result.stdout)
    assert response["status"] == "prepared"
    assert response["source"]["source_sha256"] == scope["workbook_sha256"]
    assert response["run_id"] == scope["source_run_id"]
    assert state.read_bytes() == state_before
    assert json.loads(inventory_path.read_text(encoding="utf-8"))["coverage"] == coverage_before
    expected = {
        "manifest.json", "scope.json", "dependency_audit.json", "retained_dependency_cells.json",
        "step1/HANDOFF.json", "step1/HANDOFF.md", "step2/HANDOFF.json", "step2/HANDOFF.md",
    }
    assert expected <= {path.relative_to(out).as_posix() for path in out.rglob("*") if path.is_file()}

    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    audit = json.loads((out / "dependency_audit.json").read_text(encoding="utf-8"))
    snapshot = json.loads((out / "retained_dependency_cells.json").read_text(encoding="utf-8"))
    audit_source = audit["sources"][0]
    snapshot_source = snapshot["sources"][0]
    assert (out / "scope.json").read_bytes() == scope_path.read_bytes()
    assert manifest["scope"]["sha256"] == hashlib.sha256(scope_path.read_bytes()).hexdigest()
    assert manifest["roots"]["step1"] == str(root.resolve())
    assert manifest["roots"]["step2_index"] == str(index_path.parent.resolve())
    assert manifest["index"]["root"] == "step2_index" and not Path(manifest["index"]["path"]).is_absolute()
    assert manifest["scope"]["root"] == "reading" and manifest["scope"]["path"] == "scope.json"
    assert audit_source["source_sha256"] == scope["workbook_sha256"]
    assert audit_source["run_id"] == scope["source_run_id"]
    assert audit_source["scope_sha256"] == manifest["sources"][0]["scope_sha256"]
    assert audit_source["supported_static_inbound_reference_count"] == 2
    assert audit_source["retained_dependency_range_count"] == 2
    assert len(snapshot_source["cells"]) == 3
    assert all("cached_value_available" in cell and "cached_value" in cell for cell in snapshot_source["cells"])
    assert manifest["sources"][0]["source_sha256"] == scope["workbook_sha256"]
    assert manifest["sources"][0]["run_id"] == scope["source_run_id"]
    refs = manifest["sources"][0]["views"]
    view_refs = [ref for ref in refs if ref["root"] == "reading"]
    assert view_refs and all(not Path(ref["path"]).is_absolute() for ref in view_refs)
    view = next(json.loads((out / ref["path"]).read_text(encoding="utf-8")) for ref in view_refs if ref["scope"] == "sheet" and ref["sheet_name"] == "Main")
    checkbox = next(record for record in view["records"] if record["record_type"] == "checkbox_binding")
    assert checkbox["source_location"]["sheet_name"] == "Main"
    assert checkbox["source_location"]["ooxml_part"] == "xl/worksheets/sheet1.xml"
    assert checkbox["source_location"]["address"] is None
    assert checkbox["facts"]["record"]["linked_address"] == "C3"
    assert checkbox["facts"]["artifact_ref"]["sha256"]
    assert any(record["record_type"] == "activex_event" for record in view["records"])
    assert manifest["sources"][0]["lookup"]["sheet_views"]["Main"]["path"] in {
        ref["path"] for ref in refs
    }
    assert manifest["sources"][0]["lookup"]["workbook_view"]["sha256"]

    machine = json.loads((out / "step2/HANDOFF.json").read_text(encoding="utf-8"))
    human = (out / "step2/HANDOFF.md").read_text(encoding="utf-8")
    assert {"source_quality", "reference_integrity", "static_readiness", "runtime_and_numerical_behavior"} <= machine.keys()
    handoff_source = machine["static_readiness"]["sources"][0]
    assert machine["static_readiness"]["status"] == "partial"
    assert handoff_source["supported_static_inbound_reference_count"] == audit_source["supported_static_inbound_reference_count"]
    assert machine["static_readiness"]["audit"]["sha256"] == hashlib.sha256((out / "dependency_audit.json").read_bytes()).hexdigest()
    assert machine["static_readiness"]["snapshot"]["sha256"] == hashlib.sha256((out / "retained_dependency_cells.json").read_bytes()).hexdigest()
    assert machine["artifacts"]["manifest"] == {"root": "reading", "path": "manifest.json"}
    for name, filename in (("dependency_audit", "dependency_audit.json"), ("dependency_snapshot", "retained_dependency_cells.json")):
        ref = machine["artifacts"][name]
        data = (out / ref["path"]).read_bytes()
        assert ref["root"] == "reading" and ref["sha256"] == hashlib.sha256(data).hexdigest() and ref["bytes"] == len(data)
    assert all(label in human.lower() for label in ("source quality", "reference integrity", "static readiness", "runtime/numerical behavior"))
    relocated = tmp_path / "relocated-reading"
    shutil.copytree(out, relocated)
    for ref in manifest["sources"][0]["views"]:
        assert (relocated / ref["path"]).is_file()


def test_prepare_dry_run_is_read_only_for_absent_and_existing_outputs(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path)
    state = index_path.parent / "state.json"
    before_state = state.read_bytes()
    runner = CliRunner()
    absent = tmp_path / "absent-reading"
    result = runner.invoke(app, _prepare_args(root, index_path, absent, scope_path) + ["--dry-run"])
    assert result.exit_code == 0, result.stdout
    payload = json.loads(result.stdout)
    assert payload["status"] == "dry_run" and payload["planned_files"]
    assert not absent.exists()
    existing = tmp_path / "existing-reading"
    existing.mkdir()
    marker = existing / "keep.txt"
    marker.write_text("unchanged", encoding="utf-8")
    result = runner.invoke(app, _prepare_args(root, index_path, existing, scope_path) + ["--dry-run"])
    assert result.exit_code == 0, result.stdout
    assert list(existing.iterdir()) == [marker]
    assert marker.read_text(encoding="utf-8") == "unchanged"
    assert state.read_bytes() == before_state


def test_prepare_resume_reuses_hash_checked_outputs_without_rebuilding(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from excel_to_act.steps.step2 import prepare as prepare_module

    root, index_path, scope_path, _, _ = _setup(tmp_path)
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout
    first_payload = json.loads(first.stdout)
    assert first_payload["metrics"]["inventory_model_loads"] == 1
    assert first_payload["metrics"]["graph_builds"] == 1

    def unexpected_build(*args, **kwargs):
        raise AssertionError("unchanged resume rebuilt the dependency graph")

    monkeypatch.setattr(prepare_module.RegexFormulaGraphBuilder, "build", unexpected_build)
    second = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])
    assert second.exit_code == 0, second.stdout
    assert json.loads(second.stdout)["status"] == "reused"
    assert json.loads(second.stdout)["metrics"]["graph_builds"] == 0
    assert json.loads(second.stdout)["metrics"]["inventory_model_loads"] == 1
    assert json.loads(second.stdout)["metrics"]["raw_workbook_parses"] == 0


def test_prepare_resume_reuses_when_inventory_contains_region_views(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path, structured_selectors=("Amount", "Rate"))
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout
    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert any(ref["scope"] == "region" for ref in manifest["sources"][0]["views"])

    resumed = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])

    assert resumed.exit_code == 0, resumed.stdout
    payload = json.loads(resumed.stdout)
    assert payload["status"] == "reused"
    assert payload["metrics"]["graph_builds"] == 0


def test_prepare_resume_rebuilds_when_selected_index_filename_changes(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path)
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout
    first_manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))

    renamed_index = index_path.with_name("renamed-index.json")
    index_path.rename(renamed_index)
    resumed = runner.invoke(app, _prepare_args(root, renamed_index, out, scope_path) + ["--resume"])

    assert resumed.exit_code == 0, resumed.stdout
    assert json.loads(resumed.stdout)["status"] == "prepared"
    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["revision_id"] != first_manifest["revision_id"]
    assert manifest["index"]["root"] == "step2_index"
    assert manifest["index"]["path"] == renamed_index.name
    index_root = Path(manifest["roots"][manifest["index"]["root"]])
    assert (index_root / manifest["index"]["path"]).resolve() == renamed_index.resolve()


def test_prepare_resume_rebuilds_when_saved_index_binding_is_tampered(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path)
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout

    manifest_path = out / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["index"]["path"] = "missing-index.json"
    _write_json(manifest_path, manifest)

    resumed = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])

    assert resumed.exit_code == 0, resumed.stdout
    assert json.loads(resumed.stdout)["status"] == "prepared"
    repaired = Step2ReadingManifest.model_validate_json(manifest_path.read_bytes())
    assert repaired.index.root == "step2_index"
    assert repaired.index.path == index_path.name
    index_root = Path(repaired.roots[repaired.index.root])
    assert (index_root / repaired.index.path).resolve() == index_path.resolve()


@pytest.mark.parametrize("binding", ["step1_root", "scope_sha256", "source_run_id"])
def test_prepare_resume_rebuilds_when_saved_context_bindings_are_tampered(tmp_path: Path, binding: str) -> None:
    root, index_path, scope_path, index, _ = _setup(tmp_path)
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout

    manifest_path = out / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if binding == "step1_root":
        manifest["roots"]["step1"] = "C:/missing-step1-root"
    elif binding == "scope_sha256":
        manifest["scope"]["sha256"] = "0" * 64
    else:
        manifest["sources"][0]["run_id"] = "wrong-run"
    _write_json(manifest_path, manifest)

    resumed = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])

    assert resumed.exit_code == 0, resumed.stdout
    payload = json.loads(resumed.stdout)
    assert payload["status"] == "prepared"
    assert payload["metrics"]["graph_builds"] == 1
    repaired = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert repaired["roots"]["step1"] == str(root.resolve())
    assert repaired["scope"]["sha256"] == hashlib.sha256(scope_path.read_bytes()).hexdigest()
    assert repaired["sources"][0]["run_id"] == index["entries"][0]["run_id"]


@pytest.mark.parametrize("omitted", ["dependency_graph", "canonical_view"])
def test_prepare_resume_rebuilds_when_validated_source_output_is_omitted(tmp_path: Path, omitted: str) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path)
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout

    manifest_path = out / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source = manifest["sources"][0]
    if omitted == "dependency_graph":
        ref = source["dependency_graph"]
        source["dependency_graph"] = None
    else:
        ref = next(item for item in source["views"] if item["scope"] == "sheet" and item["sheet_name"] == "Main")
        source["views"].remove(ref)
    relative = ref["path"]
    del manifest["outputs"][relative]
    _write_json(manifest_path, manifest)
    (out / relative).unlink()

    resumed = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])

    assert resumed.exit_code == 0, resumed.stdout
    payload = json.loads(resumed.stdout)
    assert payload["status"] == "prepared"
    assert payload["metrics"]["graph_builds"] == 1
    repaired = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert relative in repaired["outputs"]
    assert (out / relative).is_file()
    if omitted == "dependency_graph":
        assert repaired["sources"][0]["dependency_graph"]["path"] == relative
    else:
        assert any(item["path"] == relative for item in repaired["sources"][0]["views"])


@pytest.mark.parametrize("field", ["allowed_dependency_ranges", "defined_names"])
def test_prepare_resume_rebuilds_when_dependency_metadata_is_omitted(tmp_path: Path, field: str) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path, defined_names=field == "defined_names")
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout

    manifest_path = out / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source = manifest["sources"][0]
    target = source if field == "allowed_dependency_ranges" else source["lookup"]
    assert len(target[field]) == 2
    expected = target[field]
    target[field] = []
    _write_json(manifest_path, manifest)

    resumed = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])

    assert resumed.exit_code == 0, resumed.stdout
    payload = json.loads(resumed.stdout)
    assert payload["status"] == "prepared"
    assert payload["metrics"]["graph_builds"] == 1
    repaired = json.loads(manifest_path.read_text(encoding="utf-8"))["sources"][0]
    restored = repaired[field] if field == "allowed_dependency_ranges" else repaired["lookup"][field]
    assert restored == expected


def test_prepare_metrics_count_inventory_deserialization_on_dry_run_and_resume(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from excel_to_act.schemas import WorkbookInventory

    root, index_path, scope_path, _, _ = _setup(tmp_path)
    original = WorkbookInventory.model_validate.__func__
    calls = []

    def counted(cls, value, *args, **kwargs):
        calls.append(cls)
        return original(cls, value, *args, **kwargs)

    monkeypatch.setattr(WorkbookInventory, "model_validate", classmethod(counted))
    runner = CliRunner()
    out = tmp_path / "reading"
    dry = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--dry-run"])
    assert dry.exit_code == 0, dry.stdout
    dry_metrics = json.loads(dry.stdout)["metrics"]
    assert len(calls) == 1 and dry_metrics["inventory_model_loads"] == 1
    assert dry_metrics["raw_workbook_parses"] == 0 and dry_metrics["graph_builds"] == 0

    calls.clear()
    prepared = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert prepared.exit_code == 0, prepared.stdout
    assert len(calls) == 1

    calls.clear()
    resumed = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])
    assert resumed.exit_code == 0, resumed.stdout
    resume_payload = json.loads(resumed.stdout)
    assert resume_payload["status"] == "reused" and len(calls) == 1
    assert resume_payload["metrics"]["inventory_model_loads"] == 1
    assert resume_payload["metrics"]["raw_workbook_parses"] == 0
    assert resume_payload["metrics"]["graph_builds"] == 0


def test_prepare_resume_invalidates_changed_scope_compiler_options_and_index(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from excel_to_act.steps.step2 import prepare as prepare_module

    root, index_path, scope_path, index, scope = _setup(tmp_path)
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout
    initial_revision = json.loads((out / "manifest.json").read_text())["revision_id"]

    monkeypatch.setattr(prepare_module, "PREPARE_OPTIONS", {**prepare_module.PREPARE_OPTIONS, "view_token_budget": 900})
    changed_options = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])
    assert changed_options.exit_code == 0, changed_options.stdout
    assert json.loads(changed_options.stdout)["status"] == "prepared"
    options_revision = json.loads((out / "manifest.json").read_text())["revision_id"]
    assert options_revision != initial_revision

    monkeypatch.setattr(prepare_module, "PREPARE_COMPILER_VERSION", "step2.prepare.test")
    changed_compiler = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])
    assert changed_compiler.exit_code == 0, changed_compiler.stdout
    assert json.loads(changed_compiler.stdout)["status"] == "prepared"
    compiler_revision = json.loads((out / "manifest.json").read_text())["revision_id"]
    assert compiler_revision != options_revision

    scope["confirmation"]["decisions"][0]["reviewer"] = "user-approved"
    _write_json(scope_path, scope)
    changed_scope = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])
    assert changed_scope.exit_code == 0, changed_scope.stdout
    assert json.loads(changed_scope.stdout)["status"] == "prepared"
    scope_revision = json.loads((out / "manifest.json").read_text())["revision_id"]
    assert scope_revision != compiler_revision

    # A changed native index reference invalidates the prepared revision while preserving failed rows.
    failed = {
        "source_id": "new-failed-source", "source_path": "broken.xlsx", "source_sha256": "bad",
        "run_id": "failed-run", "status": "error", "ready_for_next_step": False,
        "artifacts": [], "diagnostics": [{"code": "source_failed", "severity": "error", "message": "kept"}],
    }
    index["entries"].append(failed)
    _write_json(index_path, index)
    changed_index = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])
    assert changed_index.exit_code == 0, changed_index.stdout
    changed_payload = json.loads(changed_index.stdout)
    assert changed_payload["status"] == "prepared"
    assert changed_payload["metrics"]["graph_builds"] == 1
    last_manifest = json.loads((out / "manifest.json").read_text())
    assert len(last_manifest["sources"]) == 2

    moved_root = tmp_path / "moved-step1-root"
    shutil.copytree(root, moved_root)
    moved = runner.invoke(app, _prepare_args(moved_root, index_path, out, scope_path) + ["--resume"])
    assert moved.exit_code == 0, moved.stdout
    moved_payload = json.loads(moved.stdout)
    assert moved_payload["status"] == "prepared"
    assert moved_payload["metrics"]["graph_builds"] == 1
    moved_manifest = json.loads((out / "manifest.json").read_text())
    assert moved_manifest["revision_id"] != last_manifest["revision_id"]
    assert moved_manifest["roots"]["step1"] == str(moved_root.resolve())


@pytest.mark.parametrize(
    ("mutation", "code"),
    [
        ("wrong_source", "scope_source_mismatch"),
        ("wrong_run", "scope_run_mismatch"),
        ("unknown_sheet", "scope_sheet_unknown"),
        ("conflict", "scope_sheet_conflict"),
        ("bad_schema", "scope_invalid"),
        ("bad_dependency_policy", "scope_invalid"),
        ("missing_ignore_decision", "scope_decision_missing"),
    ],
)
def test_prepare_rejects_invalid_scope_identity_and_decisions(tmp_path: Path, mutation: str, code: str) -> None:
    root, index_path, scope_path, _, scope = _setup(tmp_path)
    if mutation == "wrong_source":
        scope["workbook_sha256"] = "different"
    elif mutation == "wrong_run":
        scope["source_run_id"] = "different"
    elif mutation == "unknown_sheet":
        scope["ignored_sheets"] = ["Missing"]
    elif mutation == "conflict":
        scope["retained_sheets"] = ["Main", "Calc"]
    elif mutation == "bad_dependency_policy":
        scope["dependency_policy"] = []
    elif mutation == "missing_ignore_decision":
        scope["confirmation"]["decisions"] = []
    else:
        scope["schema_version"] = "phase1.v1"
    _write_json(scope_path, scope)
    out = tmp_path / "reading"

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 1
    payload = json.loads(result.stdout)
    assert payload["status"] == "blocked"
    assert any(item["code"] == code for item in payload["diagnostics"])
    assert not out.exists()


def test_prepare_blocks_stale_ignore_confirmation_for_retained_sheet(tmp_path: Path) -> None:
    root, index_path, scope_path, _, scope = _setup(tmp_path)
    scope["retained_sheets"] = ["Main", "Calc"]
    scope["ignored_sheets"] = []
    _write_json(scope_path, scope)
    out = tmp_path / "reading"

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 1
    payload = json.loads(result.stdout)
    assert payload["status"] == "blocked"
    assert any(item["code"] == "scope_sheet_conflict" for item in payload["diagnostics"])
    assert not out.exists()


def test_prepare_blocks_conflicting_sheet_decisions_in_either_order(tmp_path: Path) -> None:
    root, index_path, scope_path, _, scope = _setup(tmp_path)
    ignore = scope["confirmation"]["decisions"][0]
    include = {**ignore, "decision": "include", "reviewer": "user"}
    messages = []
    runner = CliRunner()

    for order, decisions in (
        ("include-first", [include, ignore]),
        ("ignore-first", [ignore, include]),
    ):
        scope["confirmation"]["decisions"] = decisions
        _write_json(scope_path, scope)
        out = tmp_path / f"reading-{order}"
        result = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))

        assert result.exit_code == 1
        payload = json.loads(result.stdout)
        assert payload["status"] == "blocked"
        diagnostic = next(item for item in payload["diagnostics"] if item["code"] == "scope_sheet_conflict")
        messages.append(diagnostic["message"])
        assert not out.exists()

    assert messages[0] == messages[1]


def test_prepare_preserves_failed_batch_records_and_diagnostics(tmp_path: Path) -> None:
    root, index_path, scope_path, index, _ = _setup(tmp_path, no_references=True)
    failed = {
        "source_id": "failed-source", "source_path": "broken.xlsx", "source_sha256": "bad",
        "run_id": "failed-run", "status": "error", "ready_for_next_step": False,
        "artifacts": [], "diagnostics": [{"code": "source_failed", "severity": "error", "message": "kept"}],
    }
    index["entries"].append(failed)
    _write_json(index_path, index)

    result = CliRunner().invoke(app, _prepare_args(root, index_path, tmp_path / "reading", scope_path))

    assert result.exit_code == 0, result.stdout
    manifest = json.loads((tmp_path / "reading/manifest.json").read_text(encoding="utf-8"))
    saved = next(item for item in manifest["sources"] if item["source_id"] == "failed-source")
    assert saved["status"] == "error"
    assert saved["diagnostics"][0]["message"] == "kept"
    machine = json.loads((tmp_path / "reading/step2/HANDOFF.json").read_text(encoding="utf-8"))
    failed_row = next(item for item in machine["sources"] if item["source_id"] == "failed-source")
    assert failed_row["status"] == "error" and failed_row["ready_for_next_step"] is False
    assert failed_row["diagnostics"] == failed["diagnostics"]
    good_row = next(item for item in machine["sources"] if item["source_id"] != "failed-source")
    good_audit = next(item for item in machine["static_readiness"]["sources"] if item["source_id"] == good_row["source_id"])
    assert good_row["status"] in {"pass", "partial"}
    assert good_audit["supported_static_inbound_reference_count"] == 0
    assert good_audit["retained_dependency_range_count"] == 0
    assert good_audit["stored_cell_count"] == 0
    human = (tmp_path / "reading/step2/HANDOFF.md").read_text(encoding="utf-8")
    assert "`error`" in human and "source_failed: kept" in human
    good_line = next(line for line in human.splitlines() if f"`{good_row['source_path']}`" in line)
    assert f"| `{good_row['status']}` | {str(good_row['ready_for_next_step']).lower()} | — | 1 | 1 | 0 | 0 | 0 |" in good_line
    failed_line = next(line for line in human.splitlines() if "`broken.xlsx`" in line)
    assert "`error` | false | source_failed: kept" in failed_line


@pytest.mark.parametrize("selectors", [("#data", "#Data"), ("#Data", "#data")])
def test_prepare_resolves_structured_data_selectors_case_insensitively(tmp_path: Path, selectors: tuple[str, str]) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path, structured_selectors=selectors)
    out = tmp_path / "reading"

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 0, result.stdout
    audit = json.loads((out / "dependency_audit.json").read_text(encoding="utf-8"))["sources"][0]
    snapshot = json.loads((out / "retained_dependency_cells.json").read_text(encoding="utf-8"))["sources"][0]
    assert audit["supported_static_inbound_reference_count"] == 2
    assert audit["retained_dependency_range_count"] == 1
    assert audit["ranges"] == [{"sheet_name": "Calc", "address": "A2:B3"}]
    assert {(cell["address"], cell["value"]) for cell in snapshot["cells"]} == {
        ("A2", 10), ("B2", 20), ("A3", 30), ("B3", 40),
    }


@pytest.mark.parametrize(
    ("selectors", "expected_ranges", "expected_cells", "supported", "unresolved"),
    [
        (("@Amount", "@Rate"), ["A2", "B3"], ["A2", "B3"], 2, 0),
        (("[#Headers],[Amount]", "[#Data],[Amount]"), [], [], 0, 2),
        (("#Headers", "Amount"), ["A1:B1", "A2:A3"], ["A1", "B1", "A2", "A3"], 2, 0),
        (("#All", "#All"), ["A1:B3"], ["A1", "B1", "A2", "B2", "A3", "B3"], 2, 0),
    ],
)
def test_prepare_maps_only_exact_structured_table_selectors(
    tmp_path: Path,
    selectors: tuple[str, str],
    expected_ranges: list[str],
    expected_cells: list[str],
    supported: int,
    unresolved: int,
) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path, structured_selectors=selectors)
    out = tmp_path / "reading"

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 0, result.stdout
    audit = json.loads((out / "dependency_audit.json").read_text(encoding="utf-8"))["sources"][0]
    snapshot = json.loads((out / "retained_dependency_cells.json").read_text(encoding="utf-8"))["sources"][0]
    manifest_source = json.loads((out / "manifest.json").read_text(encoding="utf-8"))["sources"][0]
    assert audit["supported_static_inbound_reference_count"] == supported
    assert audit["unresolved_static_reference_count"] == unresolved
    assert [item["address"] for item in audit["ranges"]] == expected_ranges
    assert [item["address"] for item in snapshot["ranges"]] == expected_ranges
    assert [item["address"] for item in manifest_source["allowed_dependency_ranges"]] == expected_ranges
    assert sorted(cell["address"] for cell in snapshot["cells"]) == sorted(expected_cells)
    static_diagnostics = [item for item in audit["diagnostics"] if item["code"] == "static_reference_unresolved"]
    assert len(static_diagnostics) == unresolved


@pytest.mark.parametrize(
    ("table", "selectors", "expected_ranges", "expected_cells"),
    [
        (
            {
                "ref": "A1:A4",
                "rows": [["Amount"], [11], [22], ["=SUM(A2:A3)"]],
                "header_row_count": 1,
                "totals_row_count": 1,
                "totals_row_shown": True,
                "columns": ["Amount"],
                "totals_functions": ["sum"],
            },
            ("#Data", "#Data"),
            ["A2:A3"],
            ["A2", "A3"],
        ),
        (
            {
                "ref": "A1:A3",
                "rows": [[11], [22], [33]],
                "header_row_count": 0,
                "totals_row_count": 0,
                "totals_row_shown": False,
                "columns": ["Amount"],
            },
            ("#Data", "#Data"),
            ["A1:A3"],
            ["A1", "A2", "A3"],
        ),
        (
            {
                "ref": "A1:A3",
                "rows": [[11], [22], [33]],
                "header_row_count": 0,
                "totals_row_count": 0,
                "totals_row_shown": False,
                "columns": ["Amount"],
            },
            ("Amount", "@Amount"),
            ["A1:A3", "A3"],
            ["A1", "A2", "A3"],
        ),
    ],
)
def test_prepare_uses_validated_table_header_and_totals_bounds(
    tmp_path: Path,
    table: dict,
    selectors: tuple[str, str],
    expected_ranges: list[str],
    expected_cells: list[str],
) -> None:
    root, index_path, scope_path, index, _ = _setup(
        tmp_path, structured_selectors=selectors, structured_table=table,
    )
    inventory_ref = next(ref for ref in index["entries"][0]["artifacts"] if ref["name"] == "inventory.json")
    inventory = json.loads((root / inventory_ref["path"]).read_text(encoding="utf-8"))
    table_record = next(
        item for sheet in inventory["sheets"] for item in sheet["ranges"] if item["kind"] == "table"
    )
    assert table_record["metadata"]["header_row_count"] == table["header_row_count"]
    assert table_record["metadata"]["totals_row_count"] == table["totals_row_count"]

    out = tmp_path / "reading"
    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 0, result.stdout
    audit = json.loads((out / "dependency_audit.json").read_text(encoding="utf-8"))["sources"][0]
    snapshot = json.loads((out / "retained_dependency_cells.json").read_text(encoding="utf-8"))["sources"][0]
    manifest_source = json.loads((out / "manifest.json").read_text(encoding="utf-8"))["sources"][0]
    assert audit["supported_static_inbound_reference_count"] == 2
    assert audit["unresolved_static_reference_count"] == 0
    assert [item["address"] for item in audit["ranges"]] == expected_ranges
    assert [item["address"] for item in snapshot["ranges"]] == expected_ranges
    assert [item["address"] for item in manifest_source["allowed_dependency_ranges"]] == expected_ranges
    assert sorted(cell["address"] for cell in snapshot["cells"]) == sorted(expected_cells)


def test_prepare_reports_structured_destination_unresolved_without_legacy_table_bounds(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from excel_to_act.schemas import WorkbookInventory
    from excel_to_act.steps.step2 import prepare as prepare_module

    table = {
        "ref": "A1:A3",
        "rows": [[11], [22], [33]],
        "header_row_count": 0,
        "totals_row_count": 0,
        "totals_row_shown": False,
        "columns": ["Amount"],
    }
    root, index_path, scope_path, _, _ = _setup(
        tmp_path, structured_selectors=("#Data", "Amount"), structured_table=table,
    )
    validate_index = prepare_module._validate_saved_index

    def validate_legacy_inventory(*args, **kwargs):
        result = validate_index(*args, **kwargs)
        for _, model in kwargs["validated_artifacts"].values():
            if isinstance(model, WorkbookInventory):
                for sheet in model.sheets:
                    for item in sheet.ranges:
                        if item.kind == "table":
                            item.metadata.pop("header_row_count", None)
                            item.metadata.pop("totals_row_count", None)
                            item.metadata.pop("totals_row_shown", None)
        return result

    monkeypatch.setattr(prepare_module, "_validate_saved_index", validate_legacy_inventory)
    out = tmp_path / "reading"
    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 0, result.stdout
    audit = json.loads((out / "dependency_audit.json").read_text(encoding="utf-8"))["sources"][0]
    snapshot = json.loads((out / "retained_dependency_cells.json").read_text(encoding="utf-8"))["sources"][0]
    manifest_source = json.loads((out / "manifest.json").read_text(encoding="utf-8"))["sources"][0]
    assert audit["supported_static_inbound_reference_count"] == 0
    assert audit["unresolved_static_reference_count"] == 2
    unresolved = [item for item in audit["diagnostics"] if item["code"] == "static_reference_unresolved"]
    assert len(unresolved) == 2
    assert audit["ranges"] == snapshot["ranges"] == manifest_source["allowed_dependency_ranges"] == []
    assert snapshot["cells"] == []


def test_prepare_reports_relative_name_as_unresolved_and_lists_unused_absolute_name(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path, defined_names=True)
    out = tmp_path / "reading"

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 0, result.stdout
    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    lookup = manifest["sources"][0]["lookup"]["defined_names"]
    unused = next(item for item in lookup if item["name"] == "UnusedAbsolute")
    assert unused["scope"] == "workbook" and unused["address"] == "$A$1"
    assert unused["source_location"]["sheet_name"] == "Calc"
    relative = next(item for item in lookup if item["name"] == "RelativeName")
    assert relative["address"] == "A2" and relative["source_location"]["sheet_name"] == "Calc"
    audit = json.loads((out / "dependency_audit.json").read_text(encoding="utf-8"))["sources"][0]
    snapshot = json.loads((out / "retained_dependency_cells.json").read_text(encoding="utf-8"))["sources"][0]
    assert audit["supported_static_inbound_reference_count"] == 0
    assert audit["retained_dependency_range_count"] == 0
    assert audit["unresolved_static_reference_count"] == 1
    assert audit["diagnostics"][0]["code"] == "static_reference_unresolved"
    assert "not an absolute A1 reference" in audit["diagnostics"][0]["message"]
    assert snapshot["cells"] == []


def test_prepare_rejects_multi_area_defined_name_projection(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path, multi_area_name=True)
    out = tmp_path / "reading"

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 0, result.stdout
    audit = json.loads((out / "dependency_audit.json").read_text(encoding="utf-8"))["sources"][0]
    snapshot = json.loads((out / "retained_dependency_cells.json").read_text(encoding="utf-8"))["sources"][0]
    manifest_source = json.loads((out / "manifest.json").read_text(encoding="utf-8"))["sources"][0]
    both = [item for item in manifest_source["lookup"]["defined_names"] if item["name"] == "Both"]
    assert [item["address"] for item in both] == ["$A$1", "$B$1"]
    assert audit["supported_static_inbound_reference_count"] == 0
    assert audit["retained_dependency_range_count"] == 0
    assert audit["unresolved_static_reference_count"] == 1
    assert "2 declarations" in audit["diagnostics"][0]["message"]
    assert audit["ranges"] == snapshot["ranges"] == manifest_source["allowed_dependency_ranges"] == []
    assert snapshot["cells"] == []


def test_prepare_resume_rejects_required_output_omitted_from_manifest(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path)
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout
    manifest_path = out / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    del manifest["outputs"]["step2/HANDOFF.md"]
    _write_json(manifest_path, manifest)
    (out / "step2/HANDOFF.md").unlink()

    resumed = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])

    assert resumed.exit_code == 1
    payload = json.loads(resumed.stdout)
    assert payload["status"] == "blocked"
    diagnostic = next(item for item in payload["diagnostics"] if item["code"] == "prepared_output_manifest_incomplete")
    assert "step2/HANDOFF.md" in diagnostic["paths"]


def test_prepare_blocks_tampered_or_missing_generated_audit_and_snapshot_on_resume(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path)
    out = tmp_path / "reading"
    runner = CliRunner()
    first = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert first.exit_code == 0, first.stdout
    audit = out / "dependency_audit.json"
    audit.write_text("{}", encoding="utf-8")

    tampered = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])
    assert tampered.exit_code == 1
    assert any(item["code"] == "prepared_output_checksum_mismatch" for item in json.loads(tampered.stdout)["diagnostics"])
    # Recreate a valid bundle, then remove its source snapshot.
    rebuilt = runner.invoke(app, _prepare_args(root, index_path, out, scope_path))
    assert rebuilt.exit_code == 0, rebuilt.stdout
    (out / "retained_dependency_cells.json").unlink()
    missing = runner.invoke(app, _prepare_args(root, index_path, out, scope_path) + ["--resume"])
    assert missing.exit_code == 1
    assert any(item["code"] == "prepared_output_missing" for item in json.loads(missing.stdout)["diagnostics"])


@pytest.mark.parametrize("mutation", ["hash", "missing"])
def test_prepare_blocks_changed_or_missing_native_index_references_without_state_changes(tmp_path: Path, mutation: str) -> None:
    root, index_path, scope_path, index, _ = _setup(tmp_path)
    state = index_path.parent / "state.json"
    state_before = state.read_bytes()
    if mutation == "hash":
        index["entries"][0]["artifacts"][0]["sha256"] = "0" * 64
    else:
        index["entries"][0]["artifacts"] = [
            ref for ref in index["entries"][0]["artifacts"] if ref["name"] != "inventory.json"
        ]
    _write_json(index_path, index)
    index_before = index_path.read_bytes()
    out = tmp_path / "reading"

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 1
    payload = json.loads(result.stdout)
    assert payload["status"] == "blocked" and payload["diagnostics"]
    assert not out.exists()
    assert index_path.read_bytes() == index_before
    assert state.read_bytes() == state_before


def test_prepare_returns_structured_error_for_malformed_native_index_json(tmp_path: Path) -> None:
    root, index_path, scope_path, _, _ = _setup(tmp_path)
    index_path.write_text("{invalid", encoding="utf-8")
    out = tmp_path / "reading"

    result = CliRunner().invoke(app, _prepare_args(root, index_path, out, scope_path))

    assert result.exit_code == 1
    payload = json.loads(result.stdout)
    assert payload["status"] == "blocked"
    assert payload["diagnostics"][0]["code"] == "index_invalid"
    assert payload["artifacts"] == {}
    assert not out.exists()


def test_prepare_does_not_read_legacy_unhashed_sidecar_values(tmp_path: Path) -> None:
    root, index_path, scope_path, index, scope = _setup(tmp_path)
    inventory_ref = next(ref for ref in index["entries"][0]["artifacts"] if ref["name"] == "inventory.json")
    source_root = root / Path(inventory_ref["path"]).parent
    sidecars = [source_root / "retained_dependency_cells.json", source_root / "ignored_sheet_dependency_audit.json"]
    for path in sidecars:
        path.write_text('{"sources": [{"cells": [{"value": "POISON"}]}]}', encoding="utf-8")
    first_out = tmp_path / "first-reading"
    first = CliRunner().invoke(app, _prepare_args(root, index_path, first_out, scope_path))
    assert first.exit_code == 0, first.stdout
    first_snapshot = (first_out / "retained_dependency_cells.json").read_bytes()
    for path in sidecars:
        path.write_text('{"sources": [{"cells": [{"value": "CHANGED"}]}]}', encoding="utf-8")
    second_out = tmp_path / "second-reading"
    second = CliRunner().invoke(app, _prepare_args(root, index_path, second_out, scope_path))
    assert second.exit_code == 0, second.stdout
    assert (second_out / "retained_dependency_cells.json").read_bytes() == first_snapshot
    audit = json.loads((second_out / "dependency_audit.json").read_text(encoding="utf-8"))
    assert sorted(audit["ignored_legacy_sidecars"]) == sorted(path.name for path in sidecars)
