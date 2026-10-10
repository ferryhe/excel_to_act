"""Exercise the selected optional formula oracle on the Issue #7 fixture."""

from __future__ import annotations

import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from excel_to_act.ingest.openpyxl_reader import OpenpyxlWorkbookReader
from excel_to_act.inventory.extractor import OpenpyxlInventoryExtractor
from excel_to_act.verify.numerical import _unsupported_name_error, validate_workbook


def set_cache(path: Path, value: str, cell_type: str = "n") -> None:
    staged = path.with_suffix(".staged")
    with ZipFile(path) as source, ZipFile(staged, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet2.xml":
                root = ET.fromstring(data)
                cell = next(c for c in root.iter()
                            if c.tag.endswith("}c") and c.get("r") == "B1")
                cell.set("t", cell_type)
                cell.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v").text = value
                data = ET.tostring(root)
            target.writestr(item, data)
    staged.replace(path)


def main() -> None:
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "oracle-fixture.xlsx"
        workbook = Workbook()
        inputs = workbook.active
        inputs.title = "Inputs"
        inputs["A1"] = 2
        inputs["A2"] = 3
        calc = workbook.create_sheet("Calc")
        calc["B1"] = "=Inputs!A1+Inputs!A2*4"
        workbook.save(path)
        set_cache(path, "14")

        manifest = OpenpyxlWorkbookReader().read_manifest(path)
        inventory = OpenpyxlInventoryExtractor().extract(path, manifest)
        report = validate_workbook(path, inventory)
        assert report.status == "pass", report.model_dump()
        assert report.coverage.cached_cells == report.coverage.recalculated_cells == 1
        assert report.coverage.compared_cells == 1
        assert report.baseline_source == "saved_workbook_cache"
        assert report.actual_source == "formulas"
        set_cache(path, "15")
        manifest = OpenpyxlWorkbookReader().read_manifest(path)
        inventory = OpenpyxlInventoryExtractor().extract(path, manifest)
        failed = validate_workbook(path, inventory, absolute_tolerance=0.01)
        assert failed.status == "fail", failed.model_dump()
        difference, = failed.differences
        assert (difference.sheet, difference.address, difference.baseline,
                difference.actual, difference.tolerance) == ("Calc", "B1", 15, 14.0, 0.01)

        date_path = Path(directory) / "date-fixture.xlsx"
        dates = Workbook()
        dates.active.title = "Inputs"
        date_sheet = dates.create_sheet("Calc")
        date_sheet["B1"] = "=DATE(2024,1,1)"
        date_sheet["B1"].number_format = "yyyy-mm-dd"
        dates.save(date_path)
        set_cache(date_path, "45292")
        manifest = OpenpyxlWorkbookReader().read_manifest(date_path)
        inventory = OpenpyxlInventoryExtractor().extract(date_path, manifest)
        assert inventory.sheets[1].cells[0].ooxml_cell_type == "n"
        assert validate_workbook(date_path, inventory).status == "pass"
        set_cache(date_path, "45292.000000003")
        manifest = OpenpyxlWorkbookReader().read_manifest(date_path)
        inventory = OpenpyxlInventoryExtractor().extract(date_path, manifest)
        assert inventory.sheets[1].cells[0].date_serial_text == "45292.000000003"
        precision_mismatch = validate_workbook(date_path, inventory)
        assert precision_mismatch.status == "fail" and precision_mismatch.coverage.compared_cells == 1
        difference, = precision_mismatch.differences
        assert (difference.sheet, difference.address, difference.baseline,
                difference.actual, difference.tolerance) == ("Calc", "B1", 45292.000000003, 45292, 1e-9)
        set_cache(date_path, "2024-01-01", "str")
        manifest = OpenpyxlWorkbookReader().read_manifest(date_path)
        inventory = OpenpyxlInventoryExtractor().extract(date_path, manifest)
        assert inventory.sheets[1].cells[0].ooxml_cell_type == "str"
        date_mismatch = validate_workbook(date_path, inventory)
        assert date_mismatch.status == "fail" and date_mismatch.coverage.compared_cells == 1
        difference, = date_mismatch.differences
        assert (difference.sheet, difference.address, difference.baseline,
                difference.actual) == ("Calc", "B1", "2024-01-01", 45292)
        date_sheet["B1"] = '="2024-01-01"'
        dates.save(date_path)
        set_cache(date_path, "45292")
        manifest = OpenpyxlWorkbookReader().read_manifest(date_path)
        inventory = OpenpyxlInventoryExtractor().extract(date_path, manifest)
        text_recalculation = validate_workbook(date_path, inventory)
        assert text_recalculation.status == "fail" and text_recalculation.coverage.compared_cells == 1
        difference, = text_recalculation.differences
        assert (difference.sheet, difference.address, difference.baseline,
                difference.actual) == ("Calc", "B1", 45292.0, "2024-01-01")

        mac_path = Path(directory) / "date1904-fixture.xlsx"
        mac = Workbook()
        mac.epoch = CALENDAR_MAC_1904
        mac.active.title = "Inputs"
        mac_sheet = mac.create_sheet("Calc")
        mac_sheet["B1"] = "=DATE(2024,1,1)"
        mac_sheet["B1"].number_format = "yyyy-mm-dd"
        mac.save(mac_path)
        set_cache(mac_path, "43830")
        manifest = OpenpyxlWorkbookReader().read_manifest(mac_path)
        inventory = OpenpyxlInventoryExtractor().extract(mac_path, manifest)
        mac_report = validate_workbook(mac_path, inventory)
        assert mac_report.status == "incomplete" and not mac_report.differences
        assert mac_report.coverage.unsupported_formula == ["Calc!B1"]
        assert mac_report.coverage.compared_cells == 0
        assert any("untyped numeric serial" in message and "1904" in message
                   for message in mac_report.diagnostics)

        time_path = Path(directory) / "time-fixture.xlsx"
        times = Workbook()
        times.active.title = "Inputs"
        time_sheet = times.create_sheet("Calc")
        time_sheet["B1"] = "=TIME(12,0,0)"
        time_sheet["B1"].number_format = "h:mm:ss"
        times.save(time_path)
        set_cache(time_path, "0.5")
        manifest = OpenpyxlWorkbookReader().read_manifest(time_path)
        inventory = OpenpyxlInventoryExtractor().extract(time_path, manifest)
        assert inventory.sheets[1].cells[0].cached_value == "12:00:00"
        time_report = validate_workbook(time_path, inventory)
        assert time_report.status == "pass" and time_report.coverage.compared_cells == 1
        times.epoch = CALENDAR_MAC_1904
        times.save(time_path)
        set_cache(time_path, "0.5")
        manifest = OpenpyxlWorkbookReader().read_manifest(time_path)
        inventory = OpenpyxlInventoryExtractor().extract(time_path, manifest)
        mac_time_report = validate_workbook(time_path, inventory)
        assert mac_time_report.status == "pass" and mac_time_report.coverage.compared_cells == 1

        text_path = Path(directory) / "text-fixture.xlsx"
        texts = Workbook()
        texts.active.title = "Inputs"
        text_sheet = texts.create_sheet("Calc")
        text_sheet["B1"] = '="hello"'
        text_sheet["B1"].number_format = "yyyy-mm-dd"
        texts.save(text_path)
        set_cache(text_path, "hello", "str")
        manifest = OpenpyxlWorkbookReader().read_manifest(text_path)
        inventory = OpenpyxlInventoryExtractor().extract(text_path, manifest)
        assert inventory.sheets[1].cells[0].cached_value == "hello"
        text_report = validate_workbook(text_path, inventory)
        assert text_report.status == "pass" and text_report.coverage.compared_cells == 1

        error_path = Path(directory) / "error-fixture.xlsx"
        errors = Workbook()
        source = errors.active
        source.title = "Inputs"
        source["A1"] = "#NAME?"
        source["A1"].data_type = "e"
        result_sheet = errors.create_sheet("Calc")
        result_sheet["B1"] = "=Inputs!A1"
        errors.save(error_path)
        set_cache(error_path, "#VALUE!", "e")
        manifest = OpenpyxlWorkbookReader().read_manifest(error_path)
        inventory = OpenpyxlInventoryExtractor().extract(error_path, manifest)
        propagated = validate_workbook(error_path, inventory)
        assert propagated.status == "fail" and propagated.coverage.compared_cells == 1
        difference, = propagated.differences
        assert (difference.sheet, difference.address, difference.baseline,
                difference.actual) == ("Calc", "B1", "#VALUE!", "#NAME?")
        assert not _unsupported_name_error("=If(TRUE,Inputs!A1,0)")

        for formula, cached_type, status in (
            ("=1/0", "str", "fail"),
            ("=1/0", "e", "pass"),
            ('="#DIV/0!"', "e", "fail"),
            ('="#DIV/0!"', "str", "pass"),
        ):
            result_sheet["B1"] = formula
            errors.save(error_path)
            set_cache(error_path, "#DIV/0!", cached_type)
            manifest = OpenpyxlWorkbookReader().read_manifest(error_path)
            inventory = OpenpyxlInventoryExtractor().extract(error_path, manifest)
            assert inventory.sheets[1].cells[0].ooxml_cell_type == cached_type
            report = validate_workbook(error_path, inventory)
            assert report.status == status and report.coverage.compared_cells == 1, report.model_dump()
            if status == "fail":
                difference, = report.differences
                assert (difference.sheet, difference.address, difference.baseline,
                        difference.actual, difference.tolerance) == (
                            "Calc", "B1", "#DIV/0!", "#DIV/0!", 1e-9)
                assert type(difference.actual) is str

        result_sheet["B1"] = "=BOGUS(1)"
        errors.save(error_path)
        set_cache(error_path, "#NAME?", "e")
        manifest = OpenpyxlWorkbookReader().read_manifest(error_path)
        inventory = OpenpyxlInventoryExtractor().extract(error_path, manifest)
        unsupported = validate_workbook(error_path, inventory)
        assert unsupported.status == "incomplete" and unsupported.coverage.compared_cells == 0
        assert unsupported.coverage.unsupported_formula == ["Calc!B1"]
    print("formulas matched 1900 dates and 1904 times; 1904 dates stayed incomplete; precision, type, and error checks passed.")


if __name__ == "__main__":
    main()
