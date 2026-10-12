"""Build a standalone bundle from an accepted logical semantic plan."""

from __future__ import annotations

import copy
import hashlib
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from excel_to_act.steps.conversion_workflow import hash_bytes, hash_file, json_bytes
from excel_to_act.steps.step3.calculation import (
    CalculationBlocked, FormulaEvaluator, _excel_compare, _parse_reference,
)
from excel_to_act.steps.step4.modular_compiler import (
    SourceCell,
    build_source_cell_index,
    compile_mapped_families,
    render_grouped_driver,
    validate_family_coverage,
)


def _address_key(address: str) -> str:
    if "!" not in address:
        raise CalculationBlocked(f"source map contains an unqualified address: {address!r}")
    sheet, local = address.rsplit("!", 1)
    bounds = range_boundaries(local.replace("$", ""))
    if any(value is None for value in bounds) or bounds[0] != bounds[2] or bounds[1] != bounds[3]:
        raise CalculationBlocked(f"source member is not a single cell: {address}")
    return f"{sheet.casefold()}!{get_column_letter(bounds[0])}{bounds[1]}"


def _address(sheet: str, row: int, column: int) -> str:
    return f"{sheet}!{get_column_letter(column)}{row}"


def _addresses(address: str) -> list[str]:
    reference = _parse_reference(address, "Main")
    if reference is None or not reference.sheet:
        raise CalculationBlocked(f"source range is not a qualified finite range: {address}")
    return [_address(reference.sheet, row, column)
            for row in range(reference.min_row, reference.max_row + 1)
            for column in range(reference.min_col, reference.max_col + 1)]


def _source_index(plan: dict[str, Any]) -> tuple[dict[tuple[str, int, int], list[SourceCell]], dict[str, dict[str, Any]]]:
    index = build_source_cell_index(plan)
    variables = plan.get("variables", [])
    if isinstance(variables, dict):
        variable_records = variables
    elif isinstance(variables, list):
        variable_records = {item.get("variable_id"): item for item in variables}
    else:
        raise CalculationBlocked("semantic plan has no variable records")
    return index, variable_records


def _metadata_seed_addresses(
    source_index: dict[tuple[str, int, int], list[SourceCell]],
    variables: dict[str, dict[str, Any]],
) -> set[str]:
    """Return seed cells whose accepted source role places them in axis metadata."""
    addresses: set[str] = set()
    for (sheet, row, column), cells in source_index.items():
        accepted = []
        for cell in cells:
            variable = variables.get(cell.variable_id, {})
            initial = variable.get("initial_condition")
            if (cell.role == "initial_condition" and isinstance(initial, dict)
                    and initial.get("source_role") == "accepted_axis_metadata"):
                accepted.append(cell)
        if accepted:
            if len(accepted) != len(cells):
                raise CalculationBlocked("accepted metadata seed overlaps another logical source coordinate")
            addresses.add(_address_key(_address(sheet, row, column)))
    return addresses


def _metadata_seed_bindings(
    plan: dict[str, Any],
    source_index: dict[tuple[str, int, int], list[SourceCell]] | None = None,
    variables: dict[str, dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Bind accepted metadata seeds to checked group slots, separately from raw inputs."""
    if source_index is None or variables is None:
        source_index, variables = _source_index(plan)
    metadata_cells = plan.get("adapter_metadata_cells", [])
    metadata_groups = plan.get("adapter_metadata_groups", {})
    if not isinstance(metadata_cells, list) or not isinstance(metadata_groups, dict):
        raise CalculationBlocked("accepted metadata seed groups are malformed")
    by_address: dict[str, list[dict[str, Any]]] = {}
    for item in metadata_cells:
        if not isinstance(item, dict) or not isinstance(item.get("address"), str):
            raise CalculationBlocked("adapter metadata coordinate is malformed")
        by_address.setdefault(_address_key(item["address"]), []).append(item)

    result = []
    for (sheet, row, column), cells in source_index.items():
        for cell in cells:
            if cell.role != "initial_condition":
                continue
            variable = variables.get(cell.variable_id, {})
            initial = variable.get("initial_condition")
            if not isinstance(initial, dict) or initial.get("source_role") != "accepted_axis_metadata":
                continue
            if variable.get("role") != "derived":
                raise CalculationBlocked("accepted axis metadata may seed only a derived logical variable")
            shape = variable.get("shape")
            if (not isinstance(shape, list)
                    or any(isinstance(size, bool) or not isinstance(size, int) or size < 0
                           for size in shape)
                    or len(cell.indices) != len(shape)
                    or any(isinstance(index, bool) or not isinstance(index, int)
                           or index < 0 or index >= shape[dimension]
                           for dimension, index in enumerate(cell.indices))):
                raise CalculationBlocked(f"accepted metadata seed index is outside its variable: {cell.variable_id}")
            address = _address_key(_address(sheet, row, column))
            matches = by_address.get(address, [])
            if len(matches) != 1:
                raise CalculationBlocked(f"accepted metadata seed has no unique metadata slot: {address}")
            metadata_cell = matches[0]
            group_id = metadata_cell.get("group_id")
            indices = metadata_cell.get("indices")
            group = metadata_groups.get(group_id) if isinstance(group_id, str) else None
            if (not isinstance(group, dict) or group.get("classification") != "axis_metadata"
                    or not isinstance(indices, list) or len(indices) != 1
                    or isinstance(indices[0], bool) or not isinstance(indices[0], int)):
                raise CalculationBlocked(f"accepted metadata seed is not bound to an axis metadata group: {address}")
            group_addresses = group.get("addresses")
            if (not isinstance(group_addresses, list) or indices[0] < 0
                    or indices[0] >= len(group_addresses)
                    or _address_key(group_addresses[indices[0]]) != address):
                raise CalculationBlocked(f"accepted metadata seed slot does not match its source address: {address}")
            result.append({"variable_id": cell.variable_id, "indices": list(cell.indices),
                           "metadata_group": group_id, "metadata_indices": indices,
                           "source_address": address})
    result.sort(key=lambda item: (item["variable_id"], item["indices"], item["source_address"]))
    return result


def _source_axis_name(plan: dict[str, Any], addresses: list[str]) -> str | None:
    coordinates = []
    for address in addresses:
        parsed = _parse_reference(address, "Main")
        if parsed is None or not parsed.single:
            return None
        coordinates.append((parsed.sheet.casefold(), parsed.min_row, parsed.min_col))
    matches: set[str] = set()
    variables = plan.get("variables", {})
    records = variables.values() if isinstance(variables, dict) else variables
    for variable in records:
        for axis in variable.get("axes", []):
            source = _parse_reference(axis.get("source", ""), "Main")
            if source is None:
                continue
            if all(source.sheet.casefold() == sheet and source.min_row <= row <= source.max_row
                   and source.min_col <= col <= source.max_col for sheet, row, col in coordinates):
                name = axis.get("name")
                if isinstance(name, str):
                    matches.add(name)
    return next(iter(matches)) if len(matches) == 1 else None


def add_adapter_metadata(plan: dict[str, Any], catalog: dict[str, Any],
                         model_scope: dict[str, Any], trace: dict[str, Any],
                         source_path: Path, source_sha256: str,
                         candidate_trace: dict[str, Any]) -> dict[str, Any]:
    """Bind catalog metadata coordinates to stable group slots for the private adapter."""
    if plan.get("adapter_metadata_cells") or plan.get("adapter_metadata_groups"):
        raise CalculationBlocked("semantic plan already has adapter metadata bindings")
    source_index, variables = _source_index(plan)
    metadata_seed_addresses = _metadata_seed_addresses(source_index, variables)
    groups: dict[str, dict[str, Any]] = {}
    cells: list[dict[str, Any]] = []

    def add_group(group_id: str, addresses: list[str], axis: str | None,
                  classification: str) -> None:
        canonical = sorted({_address_key(address) for address in addresses},
                           key=lambda address: (address.split("!", 1)[0],
                                               range_boundaries(address.split("!", 1)[1])[1],
                                               range_boundaries(address.split("!", 1)[1])[0]))
        if not canonical:
            return
        # Model variables own overlapping coordinates. Their approved logical binding
        # remains authoritative; metadata groups cover only adapter-only source facts.
        available = []
        for address in canonical:
            sheet, local = address.split("!", 1)
            col, row, _max_col, _max_row = range_boundaries(local)
            if (not source_index.get((sheet, row, col))
                    or _address_key(address) in metadata_seed_addresses):
                available.append(address)
        if not available:
            return
        groups[group_id] = {"addresses": available, "classification": classification}
        for index, address in enumerate(available):
            cells.append({"address": address, "group_id": group_id,
                          "indices": [index], "axes": [axis]})

    # A confirmed catalog may identify a source-owned adapter blank that the
    # earlier candidate scope did not fence as a separate single cell.
    notes = catalog.get("source_scope_notes")
    reconciliation = notes.get("stored_coordinate_reconciliation", {}) if isinstance(notes, dict) else {}
    declared_blanks = {_address_key(item["address"])
                       for item in reconciliation.get("historical_empty_coordinates", [])
                       if isinstance(item, dict)
                       and item.get("source_value_status") == "stored_blank_no_literal"
                       and isinstance(item.get("address"), str)}
    catalog_blanks: set[str] = set()
    for record in catalog.get("source_records", []):
        extents = record.get("source_extents", [])
        if record.get("classification") != "adapter_metadata" or len(extents) != 1:
            continue
        extent = extents[0]
        if extent.get("source_role") != "raw_source":
            continue
        reference = _parse_reference(f"{extent.get('sheet', '')}!{extent.get('range', '')}", "Main")
        if reference is None or not reference.single:
            continue
        address = _address_key(_address(reference.sheet, reference.min_row, reference.min_col))
        if address in declared_blanks:
            provenance = record.get("provenance")
            if (reconciliation.get("source_copy_sha256") != source_sha256
                    or not isinstance(provenance, dict)
                    or provenance.get("source_workbook_sha256") != source_sha256):
                raise CalculationBlocked(f"catalog blank source binding is stale: {address}")
            catalog_blanks.add(address)

    from excel_to_act.steps.step4.discovery import _approved_stored_empty_addresses

    try:
        scoped_blanks = _approved_stored_empty_addresses(
            source_path, source_sha256, model_scope, candidate_trace)
        verified_catalog_blanks = _approved_stored_empty_addresses(
            source_path, source_sha256,
            {"allowed_regions": [{"source_ref": address} for address in sorted(catalog_blanks)]},
            candidate_trace) if catalog_blanks else set()
    except (OSError, ValueError) as exc:
        raise CalculationBlocked(str(exc)) from exc
    if verified_catalog_blanks != catalog_blanks:
        raise CalculationBlocked("catalog blank is absent from the source candidate trace")
    approved_blank_addresses = scoped_blanks | verified_catalog_blanks

    for record in catalog.get("axis_metadata", []):
        axis_id = record.get("axis_id")
        if not isinstance(axis_id, str):
            raise CalculationBlocked("confirmed catalog has malformed axis metadata")
        addresses: list[str] = []
        for extent in record.get("source_extents", []):
            parsed = _parse_reference(f"{extent.get('sheet', '')}!{extent.get('range', '')}", "Main")
            if parsed is None:
                raise CalculationBlocked(f"catalog axis extent is invalid: {axis_id}")
            addresses.extend(_address(parsed.sheet, row, col)
                             for row in range(parsed.min_row, parsed.max_row + 1)
                             for col in range(parsed.min_col, parsed.max_col + 1))
        add_group(f"axis.{axis_id}", addresses, _source_axis_name(plan, addresses), "axis_metadata")

    for record in catalog.get("source_records", []):
        record_id = record.get("record_id")
        classification = record.get("classification")
        if (not isinstance(record_id, str)
                or classification not in {"source_constant", "adapter_metadata"}):
            raise CalculationBlocked("confirmed catalog has an unsupported source metadata record")
        addresses = []
        for extent in record.get("source_extents", []):
            parsed = _parse_reference(f"{extent.get('sheet', '')}!{extent.get('range', '')}", "Main")
            if parsed is None:
                raise CalculationBlocked(f"catalog source record extent is invalid: {record_id}")
            addresses.extend(_address(parsed.sheet, row, col)
                             for row in range(parsed.min_row, parsed.max_row + 1)
                             for col in range(parsed.min_col, parsed.max_col + 1))
        audited = len(addresses) == 1 and _address_key(addresses[0]) in verified_catalog_blanks
        add_group(f"source.{record_id}", addresses, None,
                  "source_audited_blank" if audited else classification)

    active_trace_blanks = {
        _address_key(item["address"])
        for item in trace.get("cells", [])
        if isinstance(item, dict) and item.get("role") == "source_value"
        and item.get("value", object()) is None and isinstance(item.get("address"), str)
    }
    if not approved_blank_addresses <= active_trace_blanks:
        raise CalculationBlocked("approved stored-empty coordinate is not a current active-trace blank")
    add_group("source.verified_empty", sorted(scoped_blanks - verified_catalog_blanks),
              None, "source_audited_blank")
    projected_blank_addresses = {
        _address_key(address)
        for group_id, group in groups.items()
        if group["classification"] == "source_audited_blank"
        for address in group["addresses"]
    }
    if projected_blank_addresses != approved_blank_addresses:
        raise CalculationBlocked("a verified stored-empty coordinate overlaps an existing logical mapping")

    flat = [cell["address"] for cell in cells]
    if len(flat) != len(set(flat)):
        raise CalculationBlocked("catalog metadata groups overlap at an adapter-owned source coordinate")
    plan["adapter_metadata_groups"] = groups
    plan["adapter_metadata_cells"] = cells
    plan["metadata_seed_bindings"] = _metadata_seed_bindings(plan, source_index, variables)
    return {"groups": groups, "cells": cells,
            "source_coordinate_count": len(flat),
            "raw_source_coordinate_count": len({_address_key(item["address"])
                                                  for item in _raw_bindings(plan, source_index, variables)}),
            "source_metadata_coordinate_count": len(flat),
            "metadata_seed_binding_count": len(plan["metadata_seed_bindings"])}


def _raw_literal(value: Any, data_type: str) -> Any:
    if data_type == "e" and isinstance(value, str):
        return {"excel_error": value.upper()}
    if hasattr(value, "isoformat"):
        return {"excel_datetime": value.isoformat(), "date_only": type(value).__name__ == "date"}
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    raise CalculationBlocked(f"raw source value cannot be serialized: {type(value).__name__}")


def _canonical_hash(value: Any) -> str:
    payload = (json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")) + "\n")
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _verified_case_layout(cases: Any, mutable: dict[str, Any],
                          external_bindings: list[dict[str, Any]], source_path: Path,
                          plan: dict[str, Any], profile: dict[str, Any], base_trace: dict[str, Any],
                          target_names: list[str], names: dict[str, Any],
                          source_values: dict[str, Any], variant_sources: list[dict[str, Any]]) -> list[dict[str, str]]:
    """Replay exact requests from source formulas before admitting finite bundle variants."""
    if not isinstance(cases, list) or not cases:
        raise CalculationBlocked("exact-case support needs at least one bound request")
    source_hash = hash_file(source_path)
    def resolve_name(name: str, _sheet: str, _owner: str) -> str:
        if name.casefold() not in names:
            raise CalculationBlocked(f"named reference is not mapped: {name}")
        return f"runtime.lookup_name(values, {name.casefold()!r})"
    base_driver = render_grouped_driver(plan, compile_mapped_families(plan, base_trace, resolve_name))
    base_driver_sha = hashlib.sha256(base_driver.encode("utf-8")).hexdigest()
    target_refs = {item.get("name"): item.get("source_ref")
                   for item in plan.get("reference_resolution", {}).get("named_references", [])}
    proofs: list[dict[str, str]] = []
    for case in cases:
        required = {"trace", "raw_request", "public_request", "producer"}
        if not isinstance(case, dict) or set(case) not in (required, required | {"plan", "profile"}):
            raise CalculationBlocked("exact-case artifact references are incomplete")
        artifacts = {}
        for label in case:
            ref = case[label]
            if (not isinstance(ref, dict) or set(ref) != {"path", "sha256"}
                    or not isinstance(ref["path"], str)
                    or not isinstance(ref["sha256"], str)):
                raise CalculationBlocked(f"exact-case {label} reference is malformed")
            path = Path(ref["path"])
            if not path.is_file() or hash_file(path) != ref["sha256"]:
                raise CalculationBlocked(f"exact-case {label} artifact is missing or stale")
            artifacts[label] = json.loads(path.read_text(encoding="utf-8"))
        trace, raw_request = artifacts["trace"], artifacts["raw_request"]
        public, producer = artifacts["public_request"], artifacts["producer"]
        raw_records = raw_request.get("input_overrides") if isinstance(raw_request, dict) else None
        if not isinstance(raw_records, list):
            raise CalculationBlocked("exact-case raw request has no input records")
        raw = {_address_key(item["qualified_address"]): item["value"] for item in raw_records}
        if len(raw) != len(raw_records):
            raise CalculationBlocked("exact-case raw request duplicates a source address")
        raw_digest = hashlib.sha256(json.dumps(raw_records, sort_keys=True, ensure_ascii=False,
                                               separators=(",", ":")).encode("utf-8")).hexdigest()
        if (trace.get("source_sha256") != source_hash or trace.get("input_sha256") != raw_digest
                or trace.get("request_sha256") != case["raw_request"]["sha256"]
                or trace.get("producer_sha256") != case["producer"]["sha256"]
                or producer.get("source_sha256") != source_hash
                or producer.get("input_sha256") != raw_digest):
            raise CalculationBlocked("exact-case trace, producer, and raw request bindings differ")
        if (not isinstance(public, dict) or set(public) != {"inputs", "external_data"}
                or not isinstance(public["inputs"], dict)
                or set(public["inputs"]) != set(mutable)
                or not isinstance(public["external_data"], dict)):
            raise CalculationBlocked("exact-case public request is incomplete")
        input_digest = _canonical_hash(public["inputs"])
        mapped: dict[str, Any] = {}
        for name, spec in mutable.items():
            record = public["inputs"][name]
            if (not isinstance(record, dict) or set(record) != {"axes", "values"}
                    or _canonical_hash(record["axes"]) != _canonical_hash(spec["axes"])):
                raise CalculationBlocked(f"exact-case input axes differ: {name}")
            cells = spec.get("mutable_cells")
            if cells is not None:
                values = record["values"]
                if (not isinstance(values, list) or len(values) != len(cells)
                        or any(not isinstance(value, dict) or set(value) != {"indices", "value"}
                               or value["indices"] != cell["indices"]
                               for value, cell in zip(values, cells))):
                    raise CalculationBlocked(f"exact-case mutable cells differ: {name}")
                leaves = {tuple(value["indices"]): value["value"] for value in values}
            elif not spec["shape"]:
                leaves = {(): record["values"]}
            else:
                def flatten(value: Any, shape: list[int], indices: tuple[int, ...] = ()) -> dict[tuple[int, ...], Any]:
                    if not shape:
                        return {indices: value}
                    if not isinstance(value, list) or len(value) != shape[0]:
                        raise CalculationBlocked(f"exact-case input shape differs: {name}")
                    return {key: leaf for position, item in enumerate(value)
                            for key, leaf in flatten(item, shape[1:], (*indices, position)).items()}
                leaves = flatten(record["values"], spec["shape"])
            for binding in spec["bindings"]:
                mapped[binding["address"]] = leaves[tuple(binding["indices"])]
        if set(mapped) != set(raw) or any(_canonical_hash(raw[key]) != _canonical_hash(value)
                                              for key, value in mapped.items()):
            raise CalculationBlocked("exact-case public and source raw inputs differ")
        external_data = public["external_data"]
        if set(external_data) != {item["variable_id"] for item in external_bindings}:
            raise CalculationBlocked("exact-case external producer set differs")
        overrides = {}
        cuts = producer.get("cuts")
        if not isinstance(cuts, dict):
            raise CalculationBlocked("exact-case producer has no formula cuts")
        evaluator = FormulaEvaluator(source_path)
        try:
            for address, value in raw.items():
                ref = _parse_reference(address, "Main")
                sheet = evaluator.sheet_names[ref.sheet.casefold()]
                cell = evaluator.workbook[sheet].cell(ref.min_row, ref.min_col)
                if cell.data_type == "f":
                    raise CalculationBlocked(f"exact-case raw input is a formula: {address}")
                cell.value = value
                if type(cell.value) is not type(value) or cell.value != value:
                    raise CalculationBlocked(f"exact-case raw input readback differs: {address}")
            for binding in external_bindings:
                variable_id, source_range = binding["variable_id"], binding["source_range"]
                data = external_data[variable_id]
                cut = cuts.get(source_range)
                if (not isinstance(data, dict)
                        or set(data) != {"values", "axes", "source_sha256", "version", "input_sha256"}
                        or data["source_sha256"] != source_hash
                        or data["input_sha256"] != input_digest
                        or not isinstance(cut, dict)
                        or _canonical_hash(data["values"]) != _canonical_hash(cut.get("values"))
                        or len(data["values"]) != len(binding["addresses"])):
                    raise CalculationBlocked(f"exact-case external producer differs: {variable_id}")
                for address, value in zip(binding["addresses"], data["values"]):
                    ref = _parse_reference(address, "Main")
                    sheet = evaluator.sheet_names[ref.sheet.casefold()]
                    cell = evaluator.workbook[sheet].cell(ref.min_row, ref.min_col)
                    if cell.data_type != "f":
                        raise CalculationBlocked(f"exact-case external cut is not a formula: {address}")
                    formula = cell.value
                    overrides[address] = {"boundary_variable_id": source_range, "value": value,
                                          "source_formula": formula,
                                          "source_formula_sha256": hashlib.sha256(formula.encode()).hexdigest(),
                                          "capture_artifact_sha256": case["producer"]["sha256"]}
            evaluator.external_overrides = {
                evaluator._key(evaluator.sheet_names[_parse_reference(address, "Main").sheet.casefold()],
                               _parse_reference(address, "Main").min_row,
                               _parse_reference(address, "Main").min_col): value
                for address, value in overrides.items()}
            for name in target_names:
                ref = _parse_reference(target_refs.get(name, ""), "Main")
                if ref is None or not ref.single or not ref.sheet:
                    raise CalculationBlocked(f"exact-case target has no source cell: {name}")
                evaluator.evaluate_cell(ref.sheet, ref.min_row, ref.min_col, dependency=False)
            replay = {"cells": list(evaluator.cells.values()), "edges": list(evaluator.edges.values()),
                      "names": list(evaluator.names.values()), "lookups": evaluator.lookups,
                      "branches": evaluator.branches, "dynamic_ranges": evaluator.dynamic_ranges}
            for kind, actual in replay.items():
                expected = trace.get(kind)
                if (not isinstance(expected, list)
                        or _canonical_hash(sorted(actual, key=_canonical_hash))
                        != _canonical_hash(sorted(expected, key=_canonical_hash))):
                    raise CalculationBlocked(f"exact-case source {kind} closure differs from trace")
        finally:
            evaluator.close()
        case_plan, case_profile = artifacts.get("plan", plan), artifacts.get("profile", profile)
        case_names = names
        case_index = None
        if "plan" in artifacts:
            case_index, case_variables = _source_index(case_plan)
            base_index, base_variables = _source_index(plan)
            extra_variables = {key: value for key, value in case_variables.items()
                               if key not in base_variables}
            if (case_plan.get("source", {}).get("source_sha256") != source_hash
                    or case_profile.get("source_sha256") != source_hash
                    or any(case_variables.get(key, {}).get("shape") != value.get("shape")
                           for key, value in base_variables.items())
                    or any(value.get("role") != "derived" or value.get("initial_condition") is not None
                           for value in extra_variables.values())
                    or _target_bindings(target_names, case_plan) != _target_bindings(target_names, plan)
                    or case_plan.get("adapter_metadata_groups") != plan.get("adapter_metadata_groups")
                    or case_plan.get("adapter_metadata_cells") != plan.get("adapter_metadata_cells")):
                raise CalculationBlocked("exact-case logical shapes, metadata, or source differ from bundle")
            if _raw_bindings(case_plan, case_index, case_variables) != _raw_bindings(
                    plan, base_index, base_variables):
                raise CalculationBlocked("exact-case raw or fixed source bindings differ from bundle")
            case_values = _augment_literal_sources(source_path, source_hash, case_plan, trace,
                                                   source_values.copy())
            case_literals = {_address_key(address): case_values[_address_key(address)]
                             for address in case_plan.get("adapter_source_literals", [])
                             if _address_key(address) not in source_values}
            case_names = _named_bindings(case_plan, case_index,
                                         set(source_values) | set(case_literals))
            if not {item["name"].casefold() for item in trace.get("names", [])} <= set(case_names):
                raise CalculationBlocked("exact-case named source reference is unmapped")
        validate_family_coverage(case_plan, case_profile, trace)
        if "plan" not in artifacts:
            base_formulas = {cell["address"].casefold(): cell["formula"]
                             for cell in base_trace.get("cells", [])
                             if cell.get("role") == "calculated_formula"}
            case_formulas = {cell["address"].casefold(): cell["formula"]
                             for cell in trace["cells"] if cell.get("role") == "calculated_formula"}
            if base_formulas != case_formulas:
                raise CalculationBlocked("exact-case formula members differ from bundled source")
        def case_resolve_name(name: str, _sheet: str, _owner: str) -> str:
            if name.casefold() not in case_names:
                raise CalculationBlocked(f"named reference is not mapped: {name}")
            return f"runtime.lookup_name(values, {name.casefold()!r})"
        compiled = compile_mapped_families(case_plan, trace, case_resolve_name)
        driver = render_grouped_driver(case_plan, compiled)
        driver_sha = hashlib.sha256(driver.encode("utf-8")).hexdigest()
        if "plan" not in artifacts and driver_sha != base_driver_sha:
            raise CalculationBlocked("exact-case execution order differs from bundled driver")
        proof = {"input_sha256": input_digest,
                 "dependency_sha256": _canonical_hash(external_data),
                 "trace_sha256": case["trace"]["sha256"],
                 "raw_input_sha256": raw_digest,
                 "producer_sha256": case["producer"]["sha256"]}
        if "plan" in artifacts:
            variant_id = f"v{len(variant_sources) + 1:04d}"
            metadata = {"source_sha256": source_hash, "driver_sha256": driver_sha,
                        "trace_sha256": case["trace"]["sha256"],
                        "plan_sha256": case["plan"]["sha256"],
                        "profile_sha256": case["profile"]["sha256"],
                        "named_bindings": case_names,
                        "extra_variable_shapes": {key: value["shape"] for key, value in extra_variables.items()},
                        "source_literals": case_literals,
                        "formula_cells": _formula_cell_bindings(case_plan, trace, case_index)}
            proof.update({"variant_id": variant_id, "variant_metadata_sha256": _canonical_hash(metadata),
                          "driver_sha256": driver_sha})
            variant_sources.append({"variant_id": variant_id, "metadata": metadata,
                                    "driver": driver, "compiled": compiled})
        proofs.append(proof)
    if (len({(item["input_sha256"], item["dependency_sha256"]) for item in proofs}) != len(proofs)
            or hash_file(source_path) != source_hash):
        raise CalculationBlocked("exact-case proofs duplicate a request or source changed")
    return proofs


def _variable_contract_layout(definition: dict[str, Any], raw_bindings: list[dict[str, Any]],
                              external_bindings: list[dict[str, Any]],
                              trace_sha256: str, trace: dict[str, Any],
                              source_values: dict[str, Any], *, source_path: Path | None = None,
                              semantic_plan: dict[str, Any] | None = None,
                              source_profile: dict[str, Any] | None = None,
                              target_names: list[str] | None = None,
                              names: dict[str, Any] | None = None,
                              variant_sources: list[dict[str, Any]] | None = None,
                              capture_artifact_sha256: str | None = None) -> dict[str, Any]:
    domain = definition.get("support_domain")
    if (definition.get("schema_version") != "step3.tool_interface.v1"
            or definition.get("contract_readiness", {}).get("status") != "defined"
            or not isinstance(domain, dict) or domain.get("status") not in {"fixed_trace", "exact_cases"}
            or domain.get("active_trace_sha256") != trace_sha256):
        raise CalculationBlocked("variable-input interface or fixed-trace support domain is incomplete or stale")
    certificate = definition.get("support_certificate")
    if domain.get("trace_binding") == "post_discovery":
        approved = copy.deepcopy(definition)
        approved.pop("support_certificate", None)
        approved["support_domain"].pop("active_trace_sha256", None)
        if (not isinstance(certificate, dict)
                or certificate.get("schema_version") != "step4.support_certificate.v1"
                or certificate.get("approved_interface_sha256") != hash_bytes(json_bytes(approved))
                or certificate.get("design_sha256") != definition.get("bindings", {}).get(
                    "design", {}).get("artifact", {}).get("json_sha256")
                or certificate.get("source_sha256") != definition.get("bindings", {}).get(
                    "source", {}).get("workbook_sha256")
                or certificate.get("capture_artifact_sha256") != capture_artifact_sha256
                or certificate.get("active_trace_sha256") != trace_sha256
                or certificate.get("case_refs_sha256") != hash_bytes(json_bytes(domain.get("cases")))):
            raise CalculationBlocked("post-discovery support certificate is missing or stale")
    elif certificate is not None:
        raise CalculationBlocked("fixed-trace interface cannot carry a post-discovery certificate")
    objects = {item["variable_id"]: item for item in definition["objects"]}
    raw_ids = {item["variable_id"] for item in raw_bindings}
    if any(objects[item]["mutability"] == "unknown" for item in raw_ids & objects.keys()):
        raise CalculationBlocked("raw source bindings lack complete approved mutability policy")
    unlisted = raw_ids - objects.keys()
    if unlisted:
        records = (semantic_plan or {}).get("variables", {})
        variables = (records if isinstance(records, dict) else
                     {item.get("variable_id"): item for item in records if isinstance(item, dict)}
                     if isinstance(records, list) else {})
        approved_fixed = {(_address_key(binding["address"]), binding["variable_id"],
                           tuple(binding["indices"])) for binding in raw_bindings
                          if binding["variable_id"] in objects
                          and objects[binding["variable_id"]]["mutability"] == "immutable"}
        for binding in (item for item in raw_bindings if item["variable_id"] in unlisted):
            variable = variables.get(binding["variable_id"], {})
            initial = variable.get("initial_condition")
            if not isinstance(initial, dict):
                raise CalculationBlocked("raw source bindings lack an approved fixed seed alias")
            source_id = initial.get("input_variable_id")
            source_index = initial.get("input_index", [])
            seed_ref = initial.get("seed_ref")
            if (variable.get("role") != "derived" or not isinstance(seed_ref, str)
                    or not isinstance(source_index, list)
                    or binding["indices"] != [initial.get("logical_index")]
                    or source_id not in objects or objects[source_id]["mutability"] != "immutable"
                    or (_address_key(binding["address"]), source_id, tuple(source_index)) not in approved_fixed
                    or _address_key(seed_ref) != _address_key(binding["address"])
                    or _address_key(binding["address"]) not in source_values):
                raise CalculationBlocked("raw source bindings lack an approved fixed seed alias")
    external_ids = {item["variable_id"] for item in external_bindings}
    if set(domain.get("external_policy", {})) != external_ids:
        raise CalculationBlocked("each formula-derived external producer needs a same-request policy")
    if (not external_ids <= objects.keys() or any(
            definition["source_provenance"]["inputs"][variable_id].get("role")
            != "formula_derived_external" for variable_id in external_ids)):
        raise CalculationBlocked("external producer bindings lack approved interface provenance")
    if any(objects[variable_id]["required"] is not True or objects[variable_id]["mutability"] != "mutable"
           for variable_id in external_ids):
        raise CalculationBlocked("same-request external producers must be required mutable inputs")
    mutable: dict[str, Any] = {}
    for item in definition["objects"]:
        if item["mutability"] != "mutable":
            continue
        variable_id, name = item["variable_id"], item["logical_name"]
        if variable_id in external_ids:
            continue
        if variable_id not in raw_ids or definition["source_provenance"]["inputs"][variable_id].get("role") != "source_raw":
            raise CalculationBlocked(f"mutable input is not an approved raw source object: {name}")
        bindings = [binding for binding in raw_bindings if binding["variable_id"] == variable_id]
        if any(binding["address"] in {other["address"] for other in raw_bindings
                                      if other["variable_id"] != variable_id} for binding in bindings):
            raise CalculationBlocked(f"mutable input shares a raw source cell: {name}")
        shape = item["shape"]
        indices = [tuple(binding["indices"]) for binding in bindings]
        cells = item.get("mutable_cells")
        selected = ({tuple(cell["indices"]) for cell in cells} if cells is not None else set(indices))
        if (len(indices) != (math.prod(shape) if shape else 1) and cells is None
                or len(set(indices)) != len(indices)
                or any(len(index) != len(shape) or any(
                    coordinate < 0 or coordinate >= shape[axis]
                    for axis, coordinate in enumerate(index)) for index in indices)):
            raise CalculationBlocked(f"mutable input has incomplete or duplicate source-axis bindings: {name}")
        if cells is not None and (len(selected) != len(cells) or not selected <= set(indices)):
            raise CalculationBlocked(f"mutable input has missing or duplicate per-cell source bindings: {name}")
        bindings = [binding for binding in bindings if tuple(binding["indices"]) in selected]
        mutable[name] = {"variable_id": variable_id, "shape": shape, "axes": item["axes"],
                         "value_type": item["value_type"], "required": item["required"],
                         "default_policy": item["default_policy"], "allowed_range": item["allowed_range"],
                         "allow_blank": item.get("allow_blank", False), "bindings": bindings}
        if cells is not None:
            mutable[name]["mutable_cells"] = cells
    if not mutable:
        raise CalculationBlocked("approved interface has no mutable raw business input")
    verified_cases = None
    if domain["status"] == "exact_cases":
        if (source_path is None or semantic_plan is None or source_profile is None
                or target_names is None or names is None):
            raise CalculationBlocked("exact-case closure needs source, semantic, and target evidence")
        verified_cases = _verified_case_layout(
            domain.get("cases"), mutable, external_bindings, source_path,
            semantic_plan, source_profile, trace, target_names, names, source_values,
            variant_sources if variant_sources is not None else [],
        )
        if certificate is not None:
            certificate = {**certificate, "verified_cases_sha256": _canonical_hash(verified_cases)}
    selection_cells = {_address_key(item["formula_cell"])
                       for kind in ("lookups", "branches", "dynamic_ranges")
                       for item in trace.get(kind, []) if isinstance(item, dict)
                       and isinstance(item.get("formula_cell"), str) and "!" in item["formula_cell"]}
    control_cells = {_address_key(item["formula_cell"])
                     for kind in ("branches", "dynamic_ranges")
                     for item in trace.get(kind, []) if isinstance(item, dict)
                     and isinstance(item.get("formula_cell"), str) and "!" in item["formula_cell"]}
    lookups_by_owner: dict[str, list[dict[str, Any]]] = {}
    for lookup in trace.get("lookups", []):
        if isinstance(lookup, dict) and isinstance(lookup.get("formula_cell"), str):
            lookups_by_owner.setdefault(_address_key(lookup["formula_cell"]), []).append(lookup)
    trace_cells = {_address_key(cell["address"]): cell for cell in trace.get("cells", [])
                   if isinstance(cell, dict) and isinstance(cell.get("address"), str)}
    prerequisites: dict[str, list[dict[str, Any]]] = {}
    for edge in trace.get("edges", []):
        if isinstance(edge, dict) and isinstance(edge.get("consumer"), str):
            prerequisites.setdefault(_address_key(edge["consumer"]), []).append(edge)
    mutable_addresses = {binding["address"] for spec in mutable.values() for binding in spec["bindings"]}
    external_addresses = {address for binding in external_bindings for address in binding["addresses"]}
    fixed_trace_values: dict[str, Any] = {}
    pending, seen = ([(address, address) for address in selection_cells]
                     if verified_cases is None else []), set()
    verified_match_ranges: dict[str, set[str]] = {}
    for owner, lookups in lookups_by_owner.items():
        for lookup in lookups:
            if lookup.get("function") != "VLOOKUP":
                continue
            table = _parse_reference(lookup.get("table_range", ""), "Main")
            column = lookup.get("lookup_column")
            if (table is None or not table.sheet or type(column) is not int
                    or not 1 <= column <= table.max_col - table.min_col + 1):
                raise CalculationBlocked("variable-input lookup has no finite saved key column")
            pending.extend((_address_key(_address(table.sheet, row, table.min_col + column - 1)), owner)
                           for row in range(table.min_row, table.max_row + 1))
    while pending:
        address, owner = pending.pop()
        if (owner, address) in seen:
            continue
        seen.add((owner, address))
        if address in external_addresses:
            raise CalculationBlocked(
                f"variable-input selection depends on an unverified external producer: {address} for {owner}")
        if address in mutable_addresses:
            if address not in source_values:
                raise CalculationBlocked(f"fixed-trace selector has no saved raw value: {address}")
            fixed_trace_values[address] = source_values[address]
        if address in lookups_by_owner and (address != owner or owner in control_cells
                                            or len(lookups_by_owner[address]) > 1):
            for lookup in lookups_by_owner[address]:
                if lookup.get("function") == "MATCH":
                    reference = _parse_reference(lookup.get("lookup_range", ""), "Main")
                    formula = trace_cells.get(address, {})
                    if (trace.get("source_sha256") != definition["bindings"]["source"]["workbook_sha256"]
                            or formula.get("role") != "calculated_formula"
                            or not isinstance(formula.get("formula"), str)
                            or "MATCH(" not in formula["formula"].upper()
                            or lookup.get("match_type") != 0 or reference is None or not reference.sheet
                            or (reference.min_row != reference.max_row
                                and reference.min_col != reference.max_col)
                            or (reference.max_row - reference.min_row + 1)
                            * (reference.max_col - reference.min_col + 1) > 10_000):
                        raise CalculationBlocked(f"variable-input MATCH has no verified finite exact range: {address}")
                    keys = [_address_key(cell) for cell in _addresses(reference.address)]
                    if any(trace_cells.get(key, {}).get("role") != "source_value"
                           or type(trace_cells[key].get("value")) not in (str, int, float, bool, type(None))
                           for key in keys):
                        raise CalculationBlocked(f"variable-input MATCH keys are not source-bound: {address}")
                    matches = [index for index, key in enumerate(keys, 1)
                               if _excel_compare(trace_cells[key]["value"], lookup.get("lookup_value")) == 0]
                    vlookups = [item for item in lookups_by_owner[address]
                                if item.get("function") == "VLOOKUP"]
                    if (not matches or len(vlookups) != 1
                            or vlookups[0].get("return_column") != matches[0]):
                        raise CalculationBlocked(f"variable-input MATCH selection is unverified: {address}")
                    verified_match_ranges.setdefault(address, set()).add(reference.address.casefold())
                    pending.extend((key, owner) for key in keys)
                    continue
                if lookup.get("function") != "VLOOKUP":
                    raise CalculationBlocked("variable-input selection depends on an unverified lookup result")
                table = _parse_reference(lookup.get("table_range", ""), "Main")
                row, column = lookup.get("matched_row"), lookup.get("return_column")
                if (table is None or not table.sheet or type(row) is not int or type(column) is not int
                        or not table.min_row <= row <= table.max_row
                        or not 1 <= column <= table.max_col - table.min_col + 1):
                    raise CalculationBlocked("variable-input selection has no verified lookup return cell")
                pending.append((_address_key(_address(table.sheet, row, table.min_col + column - 1)), owner))
        for edge in prerequisites.get(address, []):
            if edge.get("relationship") == "lookup_table":
                continue  # The saved lookup footprint is handled through its key and selected return cells.
            if (edge.get("relationship") == "range_read"
                    and isinstance(edge.get("prerequisite"), str)
                    and edge["prerequisite"].replace("$", "").casefold()
                    in verified_match_ranges.get(address, set())):
                continue  # Every key in this exact MATCH range was checked above.
            if edge.get("relationship") not in {"cell_reference", "array_formula_anchor"}:
                raise CalculationBlocked(f"variable-input selection has an unsupported dependency: {address}")
            reference = _parse_reference(edge.get("prerequisite", ""), address.split("!", 1)[0])
            if reference is None or not reference.single or not reference.sheet:
                raise CalculationBlocked(f"variable-input selection has an unresolved dependency: {address}")
            pending.append((_address_key(reference.address), owner))
    return {"schema_version": "step4.variable_input_contract.v1", "interface_sha256": _canonical_hash(definition),
            "source_sha256": definition["bindings"]["source"]["workbook_sha256"],
            "design": definition["bindings"]["design"], "input_boundary": definition["bindings"]["input_boundary"],
            "support_domain": domain, "mutable": mutable,
            "fixed_trace_source_values": fixed_trace_values,
            **({"verified_cases": verified_cases} if verified_cases is not None else {}),
            **({"support_certificate": certificate} if certificate is not None else {}),
            "external_specs": {item["variable_id"]: {"shape": item["shape"], "axes": item["axes"],
                                "value_type": item["value_type"], "allowed_range": item["allowed_range"],
                                "allow_blank": item.get("allow_blank", False)}
                               for item in definition["objects"] if item["variable_id"] in external_ids},
            "outputs": definition["outputs"],
            "output_provenance": definition["source_provenance"].get("outputs", {})}


def _external_bindings(plan: dict[str, Any], capture: dict[str, Any] | None,
                       capture_artifact_sha256: str | None,
                       index: dict[tuple[str, int, int], list[SourceCell]]) -> list[dict[str, Any]]:
    variables = plan.get("variables", {})
    if isinstance(variables, list):
        variables = {item.get("variable_id"): item for item in variables}
    if not isinstance(variables, dict):
        raise CalculationBlocked("semantic plan has no logical-variable map")
    external = {key: item for key, item in variables.items()
                if isinstance(item, dict) and item.get("role") == "external"}
    if not external:
        if capture is not None:
            raise CalculationBlocked("external capture is present but the active plan declares no external variables")
        return []
    if (not isinstance(capture, dict) or not isinstance(capture_artifact_sha256, str)
            or capture.get("status") != "pass"
            or capture.get("schema_version") != "step4.external_inputs.v1"
            or capture.get("source_sha256") != plan.get("source", {}).get("source_sha256")):
        raise CalculationBlocked("formula-derived external variables require a current Step 4 capture artifact")
    values = capture.get("external_values")
    if not isinstance(values, dict) or _canonical_hash(values) != capture.get("external_values_sha256"):
        raise CalculationBlocked("external capture values are malformed or have a stale value hash")
    if set(values) != set(external):
        raise CalculationBlocked("external capture variable IDs differ from the approved plan")
    boundaries = {item.get("boundary_variable_id"): item for item in capture.get("boundaries", [])
                  if isinstance(item, dict)}
    if set(boundaries) != set(external):
        raise CalculationBlocked("external capture boundary records differ from the approved plan")
    result: list[dict[str, Any]] = []
    for variable_id, variable in sorted(external.items()):
        shape = variable.get("shape")
        extents = variable.get("source_extents", [])
        boundary = boundaries[variable_id]
        if not isinstance(shape, list) or len(shape) != 1 or not isinstance(extents, list) or len(extents) != 1:
            raise CalculationBlocked(f"external variable must be one declared vector extent: {variable_id}")
        extent = extents[0].get("ref") if isinstance(extents[0], dict) else None
        parsed = _parse_reference(extent or "", "Main")
        boundary_ref = _parse_reference(boundary.get("source_range", ""), "Main")
        if (parsed is None or boundary_ref is None or not parsed.sheet or not boundary_ref.sheet
                or (parsed.sheet.casefold(), parsed.min_row, parsed.max_row, parsed.min_col, parsed.max_col)
                != (boundary_ref.sheet.casefold(), boundary_ref.min_row, boundary_ref.max_row,
                    boundary_ref.min_col, boundary_ref.max_col)
                or shape != boundary.get("shape")):
            raise CalculationBlocked(f"external source extent or shape differs from capture: {variable_id}")
        vector = values[variable_id]
        addresses = [
            _address(parsed.sheet, row, col)
            for row in range(parsed.min_row, parsed.max_row + 1)
            for col in range(parsed.min_col, parsed.max_col + 1)
        ]
        if len(addresses) != shape[0] or not isinstance(vector, list) or len(vector) != shape[0]:
            raise CalculationBlocked(f"external vector length differs from its declared shape: {variable_id}")
        indices: list[list[int]] = []
        for index_in_vector, address in enumerate(addresses):
            key = _address_key(address)
            sheet, local = key.split("!", 1)
            col, row, _max_col, _max_row = range_boundaries(local)
            matches = [item for item in index.get((sheet, row, col), [])
                       if item.variable_id == variable_id and item.role == "external_boundary_input"]
            if len(matches) != 1 or matches[0].indices != (index_in_vector,):
                raise CalculationBlocked(f"external coordinate is not uniquely indexed by its approved axis: {address}")
            indices.append([index_in_vector])
            value = vector[index_in_vector]
            if isinstance(value, bool) or not (
                    value is None or isinstance(value, (str, int, float, dict))):
                raise CalculationBlocked(f"external vector contains an unsupported value at {address}")
            if isinstance(value, float) and not math.isfinite(value):
                raise CalculationBlocked(f"external vector contains a non-finite value at {address}")
            if isinstance(value, dict) and (set(value) != {"excel_error"} or not isinstance(value["excel_error"], str)):
                raise CalculationBlocked(f"external vector contains an invalid error value at {address}")
        result.append({"variable_id": variable_id, "source_range": boundary["source_range"],
                       "shape": shape, "addresses": [_address_key(address) for address in addresses],
                       "indices": indices})
    return result


def _augment_literal_sources(source_path: Path, source_hash: str, plan: dict[str, Any],
                             trace: dict[str, Any], values: dict[str, Any],
                             scenario_guard: dict[str, Any] | None = None) -> dict[str, Any]:
    """Reread declared raw inputs and table boundaries in formula mode."""
    if not source_path.is_file() or hash_file(source_path) != source_hash:
        raise CalculationBlocked("bound workbook is missing or changed before modular input capture")
    requested: set[str] = set()
    required_literals: set[str] = set()
    required_blanks: set[str] = set()
    source_index, variables = _source_index(plan)
    for record in variables.values():
        initial = record.get("initial_condition")
        if isinstance(initial, dict) and isinstance(initial.get("seed_ref"), str):
            if initial.get("source_role") == "accepted_axis_metadata":
                continue
            parsed = _parse_reference(initial["seed_ref"], "Main")
            if parsed is not None and parsed.single:
                address = _address(parsed.sheet, parsed.min_row, parsed.min_col)
                requested.add(address)
                required_literals.add(_address_key(address))
    for binding in _raw_bindings(plan, source_index, variables):
        address = binding["address"]
        requested.add(address)
        required_literals.add(_address_key(address))
    for address in plan.get("adapter_source_literals", []):
        requested.add(address)
        required_literals.add(_address_key(address))
    for record in plan.get("reference_resolution", {}).get("named_references", []):
        binding = record.get("table_binding", {})
        if binding.get("kind") == "header_vector":
            for field in ("source_ref",):
                reference = _parse_reference(binding.get(field, record.get("source_ref", "")), "Main")
                if reference is not None:
                    requested.update(_address(reference.sheet, row, column)
                                     for row in range(reference.min_row, reference.max_row + 1)
                                     for column in range(reference.min_col, reference.max_col + 1))
        for field in ("headers_ref", "row_key_ref", "key_ref"):
            reference = _parse_reference(binding.get(field, ""), "Main")
            if reference is not None:
                requested.update(_address(reference.sheet, row, column)
                                 for row in range(reference.min_row, reference.max_row + 1)
                                 for column in range(reference.min_col, reference.max_col + 1))
        for item in binding.get("source_blank_inputs", []):
            reference = _parse_reference(item.get("source_ref", ""), "Main")
            if reference is not None and reference.single:
                address = _address(reference.sheet, reference.min_row, reference.min_col)
                requested.add(address)
                required_blanks.add(_address_key(address))
    for address in (scenario_guard or {}):
        key = _address_key(address)
        requested.add(key)
        required_literals.add(key)
    for record in trace.get("cells", []):
        if isinstance(record, dict) and record.get("role") == "source_value" and isinstance(record.get("address"), str):
            requested.add(record["address"])
    metadata_groups = plan.get("adapter_metadata_groups", {})
    if not isinstance(metadata_groups, dict):
        raise CalculationBlocked("adapter metadata group map is malformed")
    metadata_addresses: set[str] = set()
    for group_id, record in metadata_groups.items():
        addresses = record.get("addresses") if isinstance(record, dict) else None
        if not isinstance(group_id, str) or not isinstance(addresses, list):
            raise CalculationBlocked("adapter metadata group has no coordinate list")
        for address in addresses:
            key = _address_key(address)
            metadata_addresses.add(key)
            requested.add(key)
            if record.get("classification") == "source_audited_blank":
                required_blanks.add(key)
            else:
                required_literals.add(key)
    formula_addresses = {
        _address_key(record["address"])
        for record in trace.get("cells", [])
        if isinstance(record, dict)
        and record.get("role") in {"calculated_formula", "calculated_array_formula"}
        and isinstance(record.get("address"), str)
    }
    overlap = sorted(required_literals & formula_addresses)
    if overlap:
        raise CalculationBlocked("declared raw source boundary overlaps a formula or array member: "
                                 + ", ".join(overlap[:5]))

    # Open formula mode once. The workbook is never saved; normal worksheet access
    # avoids repeatedly rescanning XML for thousands of sparse source coordinates.
    workbook = load_workbook(source_path, data_only=False, read_only=False, keep_vba=True)
    try:
        sheet_lookup = {sheet.title.casefold(): sheet.title for sheet in workbook.worksheets}
        for address in requested:
            key = _address_key(address)
            sheet_name, local = address.rsplit("!", 1)
            title = sheet_lookup.get(sheet_name.casefold())
            if title is None:
                raise CalculationBlocked(f"source support sheet is missing: {sheet_name}")
            cell = workbook[title][local.replace("$", "")]
            if cell.data_type == "f":
                # Formula cells are computed by the generated family driver, never captured as literals.
                if key in required_literals or key in required_blanks:
                    raise CalculationBlocked(f"declared raw source boundary is a formula cell: {address}")
                continue
            literal = _raw_literal(cell.value, cell.data_type)
            if key in required_blanks and literal is not None:
                raise CalculationBlocked(f"declared source blank is not blank: {address}")
            if key in metadata_addresses and key not in required_blanks and literal is None:
                raise CalculationBlocked(f"catalog metadata source coordinate is unexpectedly blank: {address}")
            if key in values and values[key] != literal:
                raise CalculationBlocked(f"trace source value differs from current workbook: {address}")
            values[key] = literal
    finally:
        workbook.close()
    if hash_file(source_path) != source_hash:
        raise CalculationBlocked("bound workbook changed during modular input capture")
    return values


def _unique_binding(index: dict[tuple[str, int, int], list[SourceCell]], address: str,
                    variable_id: str | None = None) -> SourceCell:
    key = _address_key(address)
    sheet, local = key.split("!", 1)
    bounds = range_boundaries(local)
    candidates = index.get((sheet, bounds[1], bounds[0]), [])
    if variable_id is not None:
        candidates = [item for item in candidates if item.variable_id == variable_id]
    unique = {(item.variable_id, item.indices, item.axes): item for item in candidates}
    if len(unique) != 1:
        raise CalculationBlocked(f"formula output coordinate is not uniquely mapped: {address}")
    return next(iter(unique.values()))


def _formula_cell_bindings(plan: dict[str, Any], trace: dict[str, Any],
                           index: dict[tuple[str, int, int], list[SourceCell]]) -> list[dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    for mapping in plan.get("source_mappings", []):
        variable_id = mapping.get("variable_id")
        for address in mapping.get("source_members", []):
            cell = _unique_binding(index, address, variable_id)
            key = _address_key(address)
            records[key] = {"address": key, "variable_id": cell.variable_id,
                            "indices": list(cell.indices), "role": "ordinary_formula"}

    trace_roles = {_address_key(item["address"]): item.get("role") for item in trace.get("cells", [])
                   if isinstance(item, dict) and isinstance(item.get("address"), str)}
    for array_family in plan.get("array_mappings", []):
        for instance in array_family.get("instances", []):
            variable_id = instance.get("variable_id")
            for address in instance.get("source_members", []):
                cell = _unique_binding(index, address, variable_id)
                key = _address_key(address)
                record = {"address": key, "variable_id": cell.variable_id,
                          "indices": list(cell.indices), "role": "array_follower",
                          "array_anchor": _address_key(instance["anchor"])}
                previous = records.get(key)
                if previous is not None and previous != record:
                    raise CalculationBlocked(f"formula cell has conflicting logical mappings: {address}")
                records[key] = record

    expected = {key for key, role in trace_roles.items()
                if role in {"calculated_formula", "calculated_array_formula"}}
    if set(records) != expected:
        raise CalculationBlocked("logical source map differs from active formula cells ("
                                  f"{len(expected - set(records))} unmapped, "
                                  f"{len(set(records) - expected)} extra)")
    return [records[key] for key in sorted(records)]


def _raw_bindings(plan: dict[str, Any], index: dict[tuple[str, int, int], list[SourceCell]],
                  variables: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    bindings: set[tuple[str, str, tuple[int, ...]]] = set()
    for (sheet, row, column), cells in index.items():
        address = _address(sheet, row, column)
        for cell in cells:
            variable = variables.get(cell.variable_id, {})
            initial = variable.get("initial_condition")
            metadata_seed = (cell.role == "initial_condition" and isinstance(initial, dict)
                             and initial.get("source_role") == "accepted_axis_metadata")
            if variable.get("role") == "raw" or (cell.role == "initial_condition" and not metadata_seed):
                bindings.add((address, cell.variable_id, cell.indices))
    return [{"address": address, "variable_id": variable_id, "indices": list(indices)}
            for address, variable_id, indices in sorted(bindings)]


def _coordinate_item(index: dict[tuple[str, int, int], list[SourceCell]],
                     raw_addresses: set[str], address: str,
                     expected_variable: str | None = None) -> dict[str, Any]:
    key = _address_key(address)
    sheet, local = key.split("!", 1)
    col, row, _max_col, _max_row = range_boundaries(local)
    candidates = index.get((sheet, row, col), [])
    if expected_variable:
        matches = [item for item in candidates if item.variable_id == expected_variable]
        if len({(item.variable_id, item.indices) for item in matches}) == 1:
            item = matches[0]
            return {"variable_id": item.variable_id, "indices": list(item.indices)}
    unique = {(item.variable_id, item.indices): item for item in candidates}
    if len(unique) == 1:
        item = next(iter(unique.values()))
        if item.role == "adapter_metadata" and item.metadata_group is not None:
            return {"adapter_metadata_group": item.metadata_group, "indices": list(item.indices)}
        return {"variable_id": item.variable_id, "indices": list(item.indices)}
    if key in raw_addresses:
        return {"address": key}
    return {"missing_address": key}


def _named_bindings(plan: dict[str, Any], index: dict[tuple[str, int, int], list[SourceCell]],
                    raw_addresses: set[str]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    references = plan.get("reference_resolution", {}).get("named_references", [])
    for item in references:
        name = item.get("name")
        if not isinstance(name, str) or not name:
            continue
        key = name.casefold()
        if key in result:
            raise CalculationBlocked(f"case-insensitive duplicate named binding: {name}")
        if item.get("kind") == "variable_coordinate":
            variable_id, indices = item.get("variable_id"), item.get("indices")
            if (not isinstance(variable_id, str) or not isinstance(indices, list)
                    or any(not isinstance(index, int) or index < 0 for index in indices)):
                raise CalculationBlocked(f"named coordinate {name} has an invalid logical binding")
            result[key] = {"kind": "coordinate", "variable_id": variable_id, "indices": indices}
            continue
        if item.get("kind") == "source_literal":
            address = item.get("source_ref")
            if not isinstance(address, str) or _address_key(address) not in raw_addresses:
                raise CalculationBlocked(f"named scalar {name} source literal is not in declared source values")
            result[key] = {"kind": "source_literal", "address": _address_key(address)}
            continue
        if item.get("kind") == "scalar":
            binding = item.get("binding")
            if binding is None:
                binding = {"kind": "variable", "variable_id": item.get("variable_id")}
            if not isinstance(binding, dict):
                raise CalculationBlocked(f"named scalar {name} has a malformed binding")
            binding_kind = binding.get("kind")
            if binding_kind in {"variable", "variable_coordinate"}:
                variable_id = binding.get("variable_id")
                indices = binding.get("indices", [])
                if not isinstance(variable_id, str) or not isinstance(indices, list):
                    raise CalculationBlocked(f"named scalar {name} has an invalid logical binding")
                if binding_kind == "variable_coordinate":
                    result[key] = {"kind": "coordinate", "variable_id": variable_id, "indices": indices}
                elif indices:
                    raise CalculationBlocked(f"named scalar {name} uses indices without variable_coordinate")
                else:
                    result[key] = {"kind": "variable", "variable_id": variable_id}
            elif binding_kind == "source_literal":
                address = binding.get("address", binding.get("source_ref"))
                if not isinstance(address, str) or _address_key(address) not in raw_addresses:
                    raise CalculationBlocked(f"named scalar {name} source literal is not in declared source values")
                result[key] = {"kind": "source_literal", "address": _address_key(address)}
            else:
                raise CalculationBlocked(f"named scalar {name} has unsupported binding kind: {binding_kind!r}")
            continue
        reference = item.get("source_ref")
        cell_range = _parse_reference(reference or "", "Main")
        if cell_range is None:
            result[key] = {"kind": "missing", "name": name}
            continue
        table_binding = item.get("table_binding", {})
        table_kind = table_binding.get("kind")
        if table_kind == "header_vector":
            addresses = [_address(cell_range.sheet, row, column)
                         for row in range(cell_range.min_row, cell_range.max_row + 1)
                         for column in range(cell_range.min_col, cell_range.max_col + 1)]
            result[key] = {"kind": "vector", "items": [
                _coordinate_item(index, raw_addresses, address) for address in addresses]}
            continue
        row_count = cell_range.max_row - cell_range.min_row + 1
        column_count = cell_range.max_col - cell_range.min_col + 1
        columns: dict[str, list[dict[str, Any]]] = {}
        formula_columns = table_binding.get("formula_columns", {})
        column_variables = table_binding.get("column_variables", {})
        value_column_index = table_binding.get("value_column_index")
        value_variable_id = table_binding.get("value_variable_id")
        key_column_index = table_binding.get("key_column_index", 1)
        key_variable_id = table_binding.get("key_variable_id")
        raw_variable_id = table_binding.get("raw_variable_id")
        for column_offset in range(column_count):
            absolute_column = cell_range.min_col + column_offset
            column_letter = get_column_letter(absolute_column)
            expected_variable = formula_columns.get(column_letter) or column_variables.get(column_letter)
            relative_index = column_offset + 1
            if relative_index == value_column_index and isinstance(value_variable_id, str):
                expected_variable = value_variable_id
            if relative_index == key_column_index and isinstance(key_variable_id, str):
                expected_variable = key_variable_id
            if expected_variable is None and isinstance(raw_variable_id, str):
                expected_variable = raw_variable_id
            cells = [_address(cell_range.sheet, row, absolute_column)
                     for row in range(cell_range.min_row, cell_range.max_row + 1)]
            columns[str(relative_index)] = [
                _coordinate_item(index, raw_addresses, address, expected_variable) for address in cells]
        result[key] = {"kind": "table", "row_count": row_count,
                       "column_count": column_count, "columns": columns,
                       "table_kind": table_kind}
    return result


def _target_bindings(target_names: list[str], plan: dict[str, Any]) -> dict[str, str]:
    refs = {item.get("name", "").casefold(): item
            for item in plan.get("reference_resolution", {}).get("named_references", [])
            if isinstance(item.get("name"), str)}
    result: dict[str, str] = {}
    for name in target_names:
        record = refs.get(name.casefold())
        if not isinstance(record, dict) or not isinstance(record.get("variable_id"), str):
            raise CalculationBlocked(f"approved target has no scalar logical binding: {name}")
        result[name] = record["variable_id"]
    return result


def _adapter_source() -> str:
    return '''"""Standalone source-address to logical-variable adapter."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import modular_runtime as runtime

def _canonical_hash(value):
    payload = json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")) + "\\n"
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()

def _allocate(shape, missing):
    if not shape:
        return missing
    if len(shape) == 1:
        return [runtime.MissingSourceValue(missing) for _ in range(shape[0])]
    if len(shape) == 2:
        return [[runtime.MissingSourceValue(missing) for _ in range(shape[1])]
                for _ in range(shape[0])]
    raise runtime.CalculationBlocked("standalone adapter supports at most two logical axes")

def _assign(values, variable_id, indices, value):
    if not indices:
        values[variable_id] = value
        return
    target = values[variable_id]
    for index in indices[:-1]:
        target = target[index]
    target[indices[-1]] = value

def build_values(source_values=None, scenario_guard_sha256=None, scenario_ids=None,
                 metadata_seed_bindings_sha256=None, raw_overrides=None, external_overrides=None):
    base = Path(__file__).resolve().parent
    if source_values is None:
        source_values = json.loads((base / "source_values.json").read_text(encoding="utf-8"))
    layout = json.loads((base / "input_layout.json").read_text(encoding="utf-8"))
    guard = layout.get("scenario_guard", {})
    if scenario_ids is not None and scenario_ids != guard.get("scenario_ids"):
        raise runtime.CalculationBlocked("bundle scenario identity differs from its generated guard")
    guard_values = guard.get("source_values")
    if (not isinstance(guard_values, dict) or not isinstance(scenario_guard_sha256, str)
            or _canonical_hash(guard_values) != scenario_guard_sha256):
        raise runtime.CalculationBlocked("bundle scenario guard differs from its generated contract")
    for address, expected in guard_values.items():
        if address not in source_values or source_values[address] != expected:
            raise runtime.CalculationBlocked(f"saved-scenario source selector changed: {address}")
    if raw_overrides:
        allowed = {item["address"] for item in layout["raw_bindings"]}
        if not set(raw_overrides) <= allowed:
            raise runtime.CalculationBlocked("variable input does not map to approved raw bindings")
        source_values = {**source_values, **raw_overrides}
    values = {"__source_values__": source_values}
    metadata_groups = layout.get("adapter_metadata_groups", {})
    metadata_values = {}
    if metadata_groups:
        metadata_payload = json.loads((base / "adapter_metadata.json").read_text(encoding="utf-8"))
        groups = metadata_payload.get("groups", {})
        if (metadata_payload.get("schema_version") != "step4.adapter_metadata.v1"
                or metadata_payload.get("source_sha256") != layout["source_sha256"]
                or not isinstance(groups, dict) or set(groups) != set(metadata_groups)):
            raise runtime.CalculationBlocked("adapter metadata file is stale or incomplete")
        metadata_values = {}
        for group_id, expected in metadata_groups.items():
            record = groups.get(group_id)
            items = record.get("values") if isinstance(record, dict) else None
            if (not isinstance(expected, dict) or not isinstance(record, dict)
                    or not isinstance(items, list) or len(items) != expected.get("length")
                    or record.get("addresses") != expected.get("addresses")
                    or _canonical_hash(items) != expected.get("sha256")
                    or record.get("sha256") != expected.get("sha256")):
                raise runtime.CalculationBlocked(f"adapter metadata group changed: {group_id}")
            metadata_values[group_id] = items
        values["__adapter_metadata__"] = metadata_values
    for variable_id, shape in layout["variable_shapes"].items():
        values[variable_id] = _allocate(shape, f"uninitialized logical value: {variable_id}")
    for binding in layout["raw_bindings"]:
        address = binding["address"]
        value = source_values.get(address, runtime.MissingSourceValue(address))
        _assign(values, binding["variable_id"], binding["indices"], value)
    metadata_seed_bindings = layout.get("metadata_seed_bindings")
    if (not isinstance(metadata_seed_bindings, list)
            or not isinstance(metadata_seed_bindings_sha256, str)
            or _canonical_hash(metadata_seed_bindings) != metadata_seed_bindings_sha256):
        raise runtime.CalculationBlocked("metadata seed bindings are missing or differ from the generated contract")
    metadata_layout = layout.get("adapter_metadata_groups", {})
    for binding in metadata_seed_bindings:
        if not isinstance(binding, dict):
            raise runtime.CalculationBlocked("metadata seed binding is malformed")
        variable_id = binding.get("variable_id")
        indices = binding.get("indices")
        metadata_group = binding.get("metadata_group")
        metadata_indices = binding.get("metadata_indices")
        source_address = binding.get("source_address")
        expected_group = metadata_layout.get(metadata_group) if isinstance(metadata_group, str) else None
        if (not isinstance(variable_id, str) or variable_id not in values
                or not isinstance(indices, list) or not isinstance(metadata_indices, list)
                or not isinstance(source_address, str) or not isinstance(expected_group, dict)
                or metadata_group not in metadata_values):
            raise runtime.CalculationBlocked("metadata seed binding has no checked value group")
        shape = layout["variable_shapes"].get(variable_id)
        if (not isinstance(shape, list) or len(indices) != len(shape)
                or any(not isinstance(index, int) or isinstance(index, bool)
                       or index < 0 or index >= shape[dimension]
                       for dimension, index in enumerate(indices))):
            raise runtime.CalculationBlocked(f"metadata seed index is outside its variable: {variable_id}")
        group_addresses = expected_group.get("addresses")
        if (not isinstance(group_addresses, list) or len(metadata_indices) != 1
                or not isinstance(metadata_indices[0], int) or isinstance(metadata_indices[0], bool)
                or metadata_indices[0] < 0 or metadata_indices[0] >= len(group_addresses)
                or group_addresses[metadata_indices[0]] != source_address):
            raise runtime.CalculationBlocked(f"metadata seed does not match its checked source slot: {source_address}")
        seed_value = metadata_values[metadata_group]
        for index in metadata_indices:
            if not isinstance(seed_value, list) or index >= len(seed_value):
                raise runtime.CalculationBlocked(f"metadata seed group is missing its source slot: {metadata_group}")
            seed_value = seed_value[index]
        _assign(values, variable_id, indices, seed_value)
    external_bindings = layout.get("external_bindings", [])
    if external_bindings and external_overrides is not None:
        if set(external_overrides) != {item["variable_id"] for item in external_bindings}:
            raise runtime.CalculationBlocked("same-request external input IDs differ from approved bindings")
        for binding in external_bindings:
            variable_id = binding["variable_id"]
            vector = external_overrides[variable_id]
            if not isinstance(vector, list) or len(vector) != len(binding["indices"]):
                raise runtime.CalculationBlocked(f"same-request external vector shape changed: {variable_id}")
            for indices, item in zip(binding["indices"], vector):
                _assign(values, variable_id, indices, item)
    elif external_bindings:
        external_path = base / "external_values.json"
        external_payload = json.loads(external_path.read_text(encoding="utf-8"))
        expected_capture = layout.get("external_capture")
        if (not isinstance(expected_capture, dict)
                or external_payload.get("schema_version") != "step4.external_values.v1"
                or external_payload.get("source_sha256") != layout["source_sha256"]
                or external_payload.get("capture_artifact_sha256") != expected_capture.get("artifact_sha256")
                or external_payload.get("capture_sha256") != expected_capture.get("capture_sha256")
                or external_payload.get("values_sha256") != expected_capture.get("values_sha256")):
            raise runtime.CalculationBlocked("external input file is stale or bound to another capture")
        external_values = external_payload.get("values")
        if not isinstance(external_values, dict) or _canonical_hash(external_values) != expected_capture.get("values_sha256"):
            raise runtime.CalculationBlocked("external input values are incomplete or changed")
        expected_ids = {binding["variable_id"] for binding in external_bindings}
        if set(external_values) != expected_ids:
            raise runtime.CalculationBlocked("external input IDs differ from the generated logical bindings")
        for binding in external_bindings:
            variable_id = binding["variable_id"]
            vector = external_values[variable_id]
            if (not isinstance(vector, list) or len(vector) != len(binding["indices"])
                    or len(vector) != binding["shape"][0]):
                raise runtime.CalculationBlocked(f"external vector shape changed: {variable_id}")
            for indices, item in zip(binding["indices"], vector):
                _assign(values, variable_id, indices, item)
    elif layout.get("external_capture") is not None:
        raise runtime.CalculationBlocked("bundle external capture has no external logical bindings")
    values["__names__"] = layout["named_bindings"]
    return values
'''


def _model_source(scenario_ids: list[str], scenario_guard_sha256: str,
                  metadata_seed_bindings_sha256: str, variable_contract_sha256: str | None = None) -> str:
    template = '''"""Run the standalone named-variable model through approved input adapters."""
from __future__ import annotations

import argparse
import hashlib
import importlib
import json
import math
from pathlib import Path
import modular_runtime as runtime
from input_adapter import build_values, _allocate, _canonical_hash
from pricing import run_pricing
from formula_families import FAMILY_FUNCTIONS

SCENARIO_IDS = __SCENARIO_IDS__
SCENARIO_GUARD_SHA256 = __SCENARIO_GUARD_SHA256__
METADATA_SEED_BINDINGS_SHA256 = __METADATA_SEED_BINDINGS_SHA256__
VARIABLE_CONTRACT_SHA256 = __VARIABLE_CONTRACT_SHA256__

def _json_value(value):
    if isinstance(value, runtime.ExcelError):
        return {"excel_error": value.code}
    if isinstance(value, runtime.MissingSourceValue):
        raise runtime.CalculationBlocked(f"active output is missing: {value.address}")
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if isinstance(value, dict) and "excel_error" in value:
        return value
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    raise runtime.CalculationBlocked(f"model output cannot serialize {type(value).__name__}")

def _leaves(value, shape, label):
    if not shape:
        if isinstance(value, list):
            raise runtime.CalculationBlocked(f"{label} has the wrong shape")
        return [((), value)]
    if not isinstance(value, list) or len(value) != shape[0]:
        raise runtime.CalculationBlocked(f"{label} has the wrong shape")
    return [((position, *indices), leaf)
            for position, item in enumerate(value)
            for indices, leaf in _leaves(item, shape[1:], label)]

def _checked_values(value, spec, label, allow_excel_error=False, dependent_value=None):
    cells = spec.get("mutable_cells")
    if cells is None:
        leaves = _leaves(value, spec["shape"], label)
        policies = [(indices, leaf, spec) for indices, leaf in leaves]
    else:
        if not isinstance(value, list) or len(value) != len(cells):
            raise runtime.CalculationBlocked(f"{label} needs exactly its declared mutable cells")
        policies = []
        for record, cell in zip(value, cells):
            if (not isinstance(record, dict) or set(record) != {"indices", "value"}
                    or _canonical_hash(record["indices"]) != _canonical_hash(cell["indices"])):
                raise runtime.CalculationBlocked(f"{label} mutable cell indices differ from the approved contract")
            policies.append((tuple(cell["indices"]), record["value"], cell))
    for _indices, leaf, leaf_spec in policies:
        policy = leaf_spec["allowed_range"]
        if (allow_excel_error and isinstance(leaf, dict) and set(leaf) == {"excel_error"}
                and isinstance(leaf["excel_error"], str) and leaf["excel_error"].startswith("#")):
            continue
        value_type = leaf_spec["value_type"]
        valid = (leaf is None and leaf_spec.get("allow_blank") is True
                 or value_type == "boolean" and type(leaf) is bool
                 or value_type == "integer" and type(leaf) is int
                 or value_type == "number" and type(leaf) is int and abs(leaf) < 10**308
                 or value_type == "number" and type(leaf) is float and math.isfinite(leaf)
                 or value_type == "string" and type(leaf) is str)
        if not valid:
            raise runtime.CalculationBlocked(f"{label} has an invalid {value_type} value")
        if policy["status"] == "declared" and leaf is not None and (
                "minimum" in policy and leaf < policy["minimum"] or
                "maximum" in policy and leaf > policy["maximum"]):
            raise runtime.CalculationBlocked(f"{label} is outside its approved range")
        choices = policy.get("values", [])
        if policy["status"] == "dependent_enum":
            choices = [*choices, dependent_value]
        if policy["status"] in ("enum", "dependent_enum") and leaf is not None and not any(
                type(leaf) is type(choice) and leaf == choice for choice in choices):
            raise runtime.CalculationBlocked(f"{label} is outside its approved values")
        if policy["status"] not in ("declared", "enum", "dependent_enum", "unbounded"):
            raise runtime.CalculationBlocked(f"{label} has no approved range policy")
    return {indices: leaf for indices, leaf, _spec in policies}

def _checked_axes(actual, spec, label):
    if not isinstance(actual, list) or len(actual) != len(spec["shape"]):
        raise runtime.CalculationBlocked(f"{label} axes or key order differ from the approved contract")
    for dimension, axis in enumerate(actual):
        keys = axis.get("keys") if isinstance(axis, dict) else None
        if keys is not None and (not isinstance(keys, list) or len(keys) != spec["shape"][dimension]
                or len({_canonical_hash(key) for key in keys}) != len(keys)):
            raise runtime.CalculationBlocked(f"{label} axis has duplicate or missing keys")
    if _canonical_hash(actual) != _canonical_hash(spec["axes"]):
        raise runtime.CalculationBlocked(f"{label} axes or key order differ from the approved contract")

def _variable_inputs(layout, inputs, external_data):
    contract = layout.get("variable_contract")
    if (not isinstance(contract, dict) or _canonical_hash(contract) != VARIABLE_CONTRACT_SHA256
            or contract.get("source_sha256") != layout.get("source_sha256")):
        raise runtime.CalculationBlocked("variable-input contract is absent or stale")
    if not isinstance(inputs, dict) or not isinstance(external_data, dict):
        raise runtime.CalculationBlocked("variable inputs and external data must be objects")
    specs = contract["mutable"]
    if set(inputs) - set(specs):
        raise runtime.CalculationBlocked("unknown or immutable variable input: " + ", ".join(sorted(set(inputs) - set(specs))))
    base = Path(__file__).resolve().parent
    manifest = json.loads((base / "model_manifest.json").read_text(encoding="utf-8"))
    files = manifest.get("files")
    if (not isinstance(files, dict) or not {"model.py", "input_layout.json", "source_values.json",
                                               "pricing.py", "formula_families.py"} <= set(files)
            or manifest.get("source_sha256") != contract["source_sha256"]):
        raise runtime.CalculationBlocked("variable-input bundle manifest is incomplete or stale")
    for filename, expected in files.items():
        path = base / filename
        if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            raise runtime.CalculationBlocked(f"variable-input bundle asset changed: {filename}")
    if manifest.get("variable_input_contract_sha256") != VARIABLE_CONTRACT_SHA256:
        raise runtime.CalculationBlocked("variable-input manifest is stale")
    certificate = contract.get("support_certificate")
    if certificate is not None and (manifest.get("support_certificate_sha256") != _canonical_hash(certificate)
            or certificate.get("verified_cases_sha256") != _canonical_hash(contract.get("verified_cases"))
            or certificate.get("active_trace_sha256") != manifest.get("trace_sha256")
            or certificate.get("capture_artifact_sha256") != manifest.get("external_capture_artifact_sha256")):
        raise runtime.CalculationBlocked("post-discovery support certificate is stale")
    source_values = json.loads((base / "source_values.json").read_text(encoding="utf-8"))
    effective = {}
    overrides = {}
    for name, spec in specs.items():
        if name not in inputs:
            if spec["required"] or spec["default_policy"]["status"] != "saved":
                raise runtime.CalculationBlocked(f"required variable input is missing: {name}")
            baseline = {tuple(binding["indices"]): source_values[binding["address"]]
                        for binding in spec["bindings"]}
            record = {"axes": spec["axes"], "values": (baseline[()] if not spec["shape"] else
                      [baseline[(index,)] for index in range(spec["shape"][0])] if len(spec["shape"]) == 1 else
                      [[baseline[(row, column)] for column in range(spec["shape"][1])]
                       for row in range(spec["shape"][0])])}
        else:
            record = inputs[name]
        if not isinstance(record, dict) or set(record) != {"values", "axes"}:
            raise runtime.CalculationBlocked(f"{name} must have values and axes")
        _checked_axes(record["axes"], spec, name)
        effective[name] = record
    by_id = {spec["variable_id"]: name for name, spec in specs.items()}
    for name, spec in specs.items():
        policy = spec["allowed_range"]
        dependent_value = None
        if policy["status"] == "dependent_enum":
            dependency = by_id.get(policy["variable_id"])
            if dependency is None or specs[dependency]["shape"]:
                raise runtime.CalculationBlocked(f"{name} has an unsupported dependent enum contract")
            dependent_value = effective[dependency]["values"]
        leaves = _checked_values(effective[name]["values"], spec, name,
                                 dependent_value=dependent_value)
        for binding in spec["bindings"]:
            overrides[binding["address"]] = leaves[tuple(binding["indices"])]
    input_digest = _canonical_hash(effective)
    for address, saved in contract["fixed_trace_source_values"].items():
        if _canonical_hash(overrides[address]) != _canonical_hash(saved):
            raise runtime.CalculationBlocked(f"unsupported fixed-trace formula or lookup selection: {address}")
    policies = contract["support_domain"]["external_policy"]
    if set(external_data) != set(policies):
        raise runtime.CalculationBlocked("same-request external producer inputs are missing or unknown")
    external_values = {}
    for variable_id, policy in policies.items():
        record = external_data[variable_id]
        if (not isinstance(record, dict) or set(record) != {"values", "axes", "source_sha256", "version", "input_sha256"}
                or record["source_sha256"] != policy["source_sha256"]
                or record["version"] != policy["version"] or record["input_sha256"] != input_digest):
            raise runtime.CalculationBlocked(f"external input is stale or bound to another request: {variable_id}")
        spec = contract["external_specs"][variable_id]
        _checked_axes(record["axes"], spec, variable_id)
        _checked_values(record["values"], spec, variable_id, allow_excel_error=True)
        external_values[variable_id] = record["values"]
    dependency_digest = _canonical_hash(external_data)
    cases = contract.get("verified_cases")
    if cases is not None:
        selected = next((case for case in cases
                         if case["input_sha256"] == input_digest
                         and case["dependency_sha256"] == dependency_digest), None)
        if selected is None:
            raise runtime.CalculationBlocked("request has no verified formula, lookup, and dependency closure")
        if "variant_id" in selected:
            contract = {**contract, "selected_variant": selected}
    return overrides, external_values, input_digest, dependency_digest, contract

def run_model(inputs=None, external_data=None):
    base = Path(__file__).resolve().parent
    layout = json.loads((base / "input_layout.json").read_text(encoding="utf-8"))
    if inputs is None:
        if external_data is not None:
            raise runtime.CalculationBlocked("external data requires a variable-input request")
        overrides, external_values, input_digest, dependency_digest, contract = None, None, None, None, None
    else:
        overrides, external_values, input_digest, dependency_digest, contract = _variable_inputs(
            layout, inputs, external_data if external_data is not None else {})
    values = build_values(scenario_guard_sha256=SCENARIO_GUARD_SHA256, scenario_ids=SCENARIO_IDS,
                          metadata_seed_bindings_sha256=METADATA_SEED_BINDINGS_SHA256,
                          raw_overrides=overrides, external_overrides=external_values)
    selected = contract.get("selected_variant") if contract is not None else None
    if selected is None:
        values = run_pricing(values, FAMILY_FUNCTIONS)
    else:
        variant_id = selected["variant_id"]
        metadata = json.loads((base / f"variant_{variant_id}.json").read_text(encoding="utf-8"))
        if (_canonical_hash(metadata) != selected["variant_metadata_sha256"]
                or metadata["source_sha256"] != layout["source_sha256"]
                or metadata["driver_sha256"] != selected["driver_sha256"]
                or metadata["trace_sha256"] != selected["trace_sha256"]):
            raise runtime.CalculationBlocked("selected formula variant metadata is stale")
        if set(metadata["source_literals"]) & set(values["__source_values__"]):
            raise runtime.CalculationBlocked("selected formula variant source literals overlap existing inputs")
        values["__source_values__"].update(metadata["source_literals"])
        for variable_id, shape in metadata["extra_variable_shapes"].items():
            if variable_id in values:
                raise runtime.CalculationBlocked("selected formula variant overwrites a logical variable")
            values[variable_id] = _allocate(shape, f"uninitialized logical value: {variable_id}")
        values["__names__"] = metadata["named_bindings"]
        # The manifest hashes every variant; import only the selected driver for this request.
        variant_families = importlib.import_module(f"formula_families_{variant_id}")
        variant_pricing = importlib.import_module(f"pricing_{variant_id}")
        values = variant_pricing.run_pricing(values, variant_families.FAMILY_FUNCTIONS)
        layout = {**layout, "formula_cells": metadata["formula_cells"],
                  "variable_shapes": {**layout["variable_shapes"], **metadata["extra_variable_shapes"]}}
    targets = {name: _json_value(runtime.at(values, variable_id, []))
               for name, variable_id in layout["target_bindings"].items()}
    calculated_variables = {}
    for item in layout["formula_cells"]:
        variable_id = item["variable_id"]
        record = calculated_variables.setdefault(variable_id, {
            "shape": layout["variable_shapes"][variable_id], "demanded_values": []})
        record["demanded_values"].append({
            "indices": item["indices"],
            "value": _json_value(runtime.at(values, variable_id, item["indices"])),
        })
    for record in calculated_variables.values():
        record["demanded_values"].sort(key=lambda item: tuple(item["indices"]))
    cells = {item["address"]: _json_value(runtime.at(values, item["variable_id"], item["indices"]))
             for item in layout["formula_cells"]}
    external_inputs = {binding["variable_id"]: _json_value(values[binding["variable_id"]])
                       for binding in layout.get("external_bindings", [])}
    age_axis = None
    external_capture = layout.get("external_capture")
    age_record = external_capture.get("age_axis") if isinstance(external_capture, dict) else None
    if isinstance(age_record, dict) and age_record.get("metadata_group"):
        age_axis = _json_value(values["__adapter_metadata__"][age_record["metadata_group"]])
    return {"schema_version": "step4.modular_model_result.v1", "status": "pass",
            "source_sha256": layout["source_sha256"], "targets": targets,
            "execution_mode": "variable_input" if contract is not None else "saved_scenario",
            "raw_inputs": overrides if contract is not None else {},
            "input_sha256": input_digest, "dependency_sha256": dependency_digest,
            "interface_sha256": contract["interface_sha256"] if contract is not None else None,
            "support_domain_sha256": _canonical_hash(contract["support_domain"]) if contract is not None else None,
            "target_axes": {provenance["selector"]: contract["outputs"]["properties"][name]["x-axes"]
                            for name, provenance in contract["output_provenance"].items()
                            if contract is not None and provenance.get("selector") in targets} if contract is not None else {},
            "calculated_variables": calculated_variables, "cells": cells,
            "external_inputs": external_inputs, "external_age_axis": age_axis,
            "calculated_variable_count": len(calculated_variables),
            "calculated_value_count": sum(len(item["demanded_values"])
                                           for item in calculated_variables.values()),
            "formula_count": len(cells), "cell_count": len(cells)}

def main():
    parser = argparse.ArgumentParser(description="Run the standalone logical-variable model")
    parser.add_argument("--out", type=Path, default=Path("model_result.json"))
    parser.add_argument("--inputs", type=Path, help="JSON request with inputs and same-request external_data")
    args = parser.parse_args()
    request = json.loads(args.inputs.read_text(encoding="utf-8")) if args.inputs is not None else None
    if args.inputs is not None and (not isinstance(request, dict) or set(request) != {"inputs", "external_data"}):
        raise runtime.CalculationBlocked("variable request must contain inputs and external_data")
    result = run_model(request["inputs"], request["external_data"]) if args.inputs is not None else run_model()
    args.out.write_text(json.dumps(result, indent=2, sort_keys=True) + "\\n", encoding="utf-8")

if __name__ == "__main__":
    main()
'''
    return (template.replace("__SCENARIO_IDS__", repr(scenario_ids))
            .replace("__SCENARIO_GUARD_SHA256__", repr(scenario_guard_sha256))
            .replace("__METADATA_SEED_BINDINGS_SHA256__", repr(metadata_seed_bindings_sha256))
            .replace("__VARIABLE_CONTRACT_SHA256__", repr(variable_contract_sha256)))


def build_modular_bundle_files(
    *,
    semantic_plan: dict[str, Any],
    source_profile: dict[str, Any],
    trace: dict[str, Any],
    raw_source_values: dict[str, Any],
    source_path: Path,
    semantic_plan_path: Path,
    source_profile_path: Path,
    semantic_map_path: Path,
    revision_dir: Path,
    target_names: list[str],
    path_coverage: dict[str, Any],
    source_sha256: str,
    design_sha256: str,
    semantic_plan_sha256: str,
    source_profile_sha256: str,
    trace_sha256: str,
    trace_path: Path,
    trace_source: str,
    scenario_ids: list[str],
    external_capture: dict[str, Any] | None = None,
    external_capture_artifact_sha256: str | None = None,
    scenario_guard: dict[str, Any] | None = None,
    implementation_plan_record: dict[str, Any] | None = None,
    implementation_plan_path: Path | None = None,
    implementation_plan_sha256: str | None = None,
    variable_interface: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Compile, package, and report a source-bound logical-variable bundle."""
    coverage = validate_family_coverage(semantic_plan, source_profile, trace)
    if semantic_plan.get("source", {}).get("source_sha256") != source_sha256:
        raise CalculationBlocked("semantic plan source hash differs from approved workbook")
    if not isinstance(raw_source_values, dict):
        raise CalculationBlocked("raw source-value ledger is not an object")
    source_values = {_address_key(key): value for key, value in raw_source_values.items()}
    scenario_guard = {_address_key(address): value for address, value in (scenario_guard or {}).items()}
    adapter_groups = semantic_plan.get("adapter_metadata_groups", {})
    if not isinstance(adapter_groups, dict):
        raise CalculationBlocked("adapter metadata groups are malformed")
    metadata_addresses = {_address_key(address) for group in adapter_groups.values()
                          for address in group.get("addresses", [])}
    all_source_values = _augment_literal_sources(source_path, source_sha256, semantic_plan, trace,
                                                 source_values, scenario_guard)
    changed_scenario = [address for address, expected in scenario_guard.items()
                        if all_source_values.get(address) != expected]
    if changed_scenario:
        raise CalculationBlocked("source inputs differ from the approved saved scenario: "
                                 + ", ".join(sorted(changed_scenario)[:8]))
    index, variables = _source_index(semantic_plan)
    raw_bindings = _raw_bindings(semantic_plan, index, variables)
    metadata_seed_bindings = _metadata_seed_bindings(semantic_plan, index, variables)
    declared_metadata_seeds = semantic_plan.get("metadata_seed_bindings")
    if declared_metadata_seeds is not None and declared_metadata_seeds != metadata_seed_bindings:
        raise CalculationBlocked("metadata seed bindings differ from the current source-role declarations")
    raw_binding_addresses = {_address_key(item["address"]) for item in raw_bindings}
    # Step 3 approves this saved scenario only. Bind every declared raw business
    # coordinate so changed selectors and configuration tables cannot activate
    # producer branches that were excluded from the reviewed model.
    scenario_guard = {_address_key(address): value for address, value in scenario_guard.items()}
    for binding in raw_bindings:
        address = _address_key(binding["address"])
        if address not in all_source_values:
            raise CalculationBlocked(f"saved-scenario raw input is missing: {address}")
        scenario_guard[address] = all_source_values[address]
    source_literals = {_address_key(address) for address in semantic_plan.get("adapter_source_literals", [])
                       if isinstance(address, str)}
    if metadata_addresses & raw_binding_addresses:
        raise CalculationBlocked("adapter metadata overlaps a declared raw business coordinate")
    unclassified = set(all_source_values) - raw_binding_addresses - metadata_addresses - source_literals
    if unclassified:
        raise CalculationBlocked("source values are outside the approved logical inputs and metadata: "
                                  + ", ".join(sorted(unclassified)[:8]))
    missing_raw = raw_binding_addresses - set(all_source_values)
    if missing_raw:
        raise CalculationBlocked("approved raw business coordinates are missing: "
                                  + ", ".join(sorted(missing_raw)[:8]))
    source_values = {address: value for address, value in all_source_values.items()
                     if address in raw_binding_addresses or address in source_literals}
    raw_addresses = set(source_values)
    metadata_payload_groups = {}
    for group_id, group in sorted(adapter_groups.items()):
        addresses = [_address_key(address) for address in group.get("addresses", [])]
        values = [all_source_values[address] for address in addresses]
        metadata_payload_groups[group_id] = {"addresses": addresses, "values": values,
                                             "sha256": _canonical_hash(values)}
    formula_cells = _formula_cell_bindings(semantic_plan, trace, index)
    names = _named_bindings(semantic_plan, index, raw_addresses)
    external_bindings = _external_bindings(semantic_plan, external_capture,
                                           external_capture_artifact_sha256, index)
    variant_sources: list[dict[str, Any]] = []
    variable_contract = (_variable_contract_layout(variable_interface, raw_bindings,
                                                    external_bindings, trace_sha256, trace, source_values,
                                                    source_path=source_path, semantic_plan=semantic_plan,
                                                    source_profile=source_profile, target_names=target_names,
                                                    names=names, variant_sources=variant_sources,
                                                    capture_artifact_sha256=external_capture_artifact_sha256)
                         if variable_interface is not None else None)
    if variable_contract is not None and variable_contract["source_sha256"] != source_sha256:
        raise CalculationBlocked("variable-input interface is bound to another source workbook")
    external_values = external_capture.get("external_values", {}) if external_capture else {}
    external_value_hash = external_capture.get("external_values_sha256") if external_capture else None
    external_capture_hash = external_capture.get("external_capture_sha256") if external_capture else None
    external_age_axis = None
    if external_capture is not None:
        age_axis = external_capture.get("age_axis")
        if (not isinstance(age_axis, dict) or not isinstance(age_axis.get("source_range"), str)
                or not isinstance(age_axis.get("values"), list)):
            raise CalculationBlocked("external capture has no source-bound age axis")
        age_addresses = {_address_key(address) for address in _addresses(age_axis["source_range"])}
        matching_groups = [group_id for group_id, group in adapter_groups.items()
                           if {_address_key(address) for address in group.get("addresses", [])} == age_addresses]
        if len(matching_groups) != 1:
            raise CalculationBlocked("external age axis does not map to one adapter metadata group")
        external_age_axis = {"source_range": age_axis["source_range"],
                             "values_sha256": age_axis.get("value_sha256"),
                             "metadata_group": matching_groups[0]}
    implementation_binding = None
    if any(value is not None for value in
           (implementation_plan_record, implementation_plan_path, implementation_plan_sha256)):
        if (implementation_plan_record is None or implementation_plan_path is None
                or not isinstance(implementation_plan_sha256, str)
                or not implementation_plan_path.is_file()
                or hash_file(implementation_plan_path) != implementation_plan_sha256):
            raise CalculationBlocked("modular bundle implementation preflight binding is incomplete or stale")
        implementation_binding = {"path": str(implementation_plan_path),
                                  "sha256": implementation_plan_sha256,
                                  "revision": implementation_plan_record.get("revision")}
    targets = _target_bindings(target_names, semantic_plan)
    def resolve_name(name: str, sheet: str, owner: str) -> str:
        if name.casefold() not in names:
            raise CalculationBlocked(f"named reference is not mapped: {name}")
        return f"runtime.lookup_name(values, {name.casefold()!r})"

    compiled = compile_mapped_families(semantic_plan, trace, resolve_name)
    driver = render_grouped_driver(semantic_plan, compiled)
    if compiled["compiled_member_count"] != coverage["ordinary_formula_count"]:
        raise CalculationBlocked("compiled ordinary formula count differs from source coverage")
    metadata_group_hashes = {
        group_id: {"addresses": record["addresses"], "length": len(record["values"]),
                   "sha256": record["sha256"]}
        for group_id, record in metadata_payload_groups.items()
    }
    layout = {
        "source_sha256": source_sha256,
        "variable_shapes": {key: value.get("shape", []) for key, value in sorted(variables.items())},
        "raw_bindings": raw_bindings,
        "metadata_seed_bindings": metadata_seed_bindings,
        "adapter_metadata_groups": metadata_group_hashes,
        "external_bindings": external_bindings,
        "external_capture": ({"artifact_sha256": external_capture_artifact_sha256,
                               "capture_sha256": external_capture_hash,
                               "values_sha256": external_value_hash,
                               "age_axis": external_age_axis}
                              if external_capture is not None else None),
        "implementation_plan": implementation_binding,
        "scenario_guard": {"scenario_ids": scenario_ids, "source_values": scenario_guard},
        "formula_cells": formula_cells,
        "named_bindings": names,
        "target_bindings": targets,
        "variable_contract": variable_contract,
    }
    equation_families = semantic_plan.get("equation_families", [])
    equation_group_count = (len(equation_families) if isinstance(equation_families, (dict, list)) else 0)
    family_manifest = {
        "schema_version": "step4.family_manifest.v1",
        "source_family_count": compiled["source_family_count"],
        "semantic_equation_group_count": equation_group_count,
        "compiled_variant_count": compiled["function_count"],
        "ordinary_formula_count": coverage["ordinary_formula_count"],
        "array_follower_count": coverage["array_follower_count"],
        "array_instance_count": coverage["array_instance_count"],
        "active_formula_member_count": coverage["active_formula_member_count"],
        "unmapped_count": 0,
        "source_family_to_function": compiled["source_family_to_function"],
        "source_family_variants": compiled["source_family_variants"],
        "member_function_exceptions": compiled["member_function_exceptions"],
        "families": semantic_plan.get("source_mappings", []),
        "arrays": semantic_plan.get("array_mappings", []),
    }
    bundle = revision_dir / "bundle"
    bundle.mkdir(parents=True)
    runtime_source = Path(__file__).with_name("modular_runtime.py").read_bytes()
    (bundle / "modular_runtime.py").write_bytes(runtime_source)
    (bundle / "input_adapter.py").write_text(_adapter_source(), encoding="utf-8")
    (bundle / "input_layout.json").write_text(
        json.dumps(layout, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    (bundle / "source_values.json").write_text(
        json.dumps(source_values, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    if metadata_payload_groups:
        metadata_payload = {"schema_version": "step4.adapter_metadata.v1",
                            "source_sha256": source_sha256,
                            "origin": "formula_mode_source", "formula_cache_inputs": False,
                            "groups": metadata_payload_groups}
        (bundle / "adapter_metadata.json").write_text(
            json.dumps(metadata_payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
            encoding="utf-8")
    if external_capture is not None:
        external_payload = {
            "schema_version": "step4.external_values.v1",
            "source_sha256": source_sha256,
            "capture_artifact_sha256": external_capture_artifact_sha256,
            "capture_sha256": external_capture_hash,
            "values_sha256": external_value_hash,
            "values": external_values,
            "boundaries": external_capture.get("boundaries", []),
            "formula_cut_members": external_capture.get("formula_cut_members", []),
            "formula_cut_formula_set_sha256": external_capture.get("formula_cut_formula_set_sha256"),
            "age_axis": external_capture.get("age_axis"),
        }
        (bundle / "external_values.json").write_text(
            json.dumps(external_payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    (bundle / "source_map.json").write_text(
        json.dumps({"semantic_plan_sha256": semantic_plan_sha256,
                    "formula_cells": formula_cells,
                    "external_bindings": external_bindings,
                    "adapter_metadata_cells": semantic_plan.get("adapter_metadata_cells", []),
                    "source_mappings": semantic_plan.get("source_mappings", []),
                    "array_mappings": semantic_plan.get("array_mappings", [])},
                   indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    family_source = "import modular_runtime as runtime\n\n" + "\n\n".join(
        compiled["functions"][name] for name in sorted(compiled["functions"]))
    family_source += "\n\nFAMILY_FUNCTIONS = {\n" + "\n".join(
        f"    {name!r}: {name}," for name in sorted(compiled["functions"])) + "\n}\n"
    (bundle / "formula_families.py").write_text(family_source, encoding="utf-8")
    pricing_source = ("from formula_families import FAMILY_FUNCTIONS\n\n" + driver +
                      "\ndef run(values):\n    return run_pricing(values, FAMILY_FUNCTIONS)\n")
    (bundle / "pricing.py").write_text(pricing_source, encoding="utf-8")
    for variant in variant_sources:
        variant_id = variant["variant_id"]
        functions = variant["compiled"]["functions"]
        variant_family_source = "import modular_runtime as runtime\n\n" + "\n\n".join(
            functions[name] for name in sorted(functions))
        variant_family_source += "\n\nFAMILY_FUNCTIONS = {\n" + "\n".join(
            f"    {name!r}: {name}," for name in sorted(functions)) + "\n}\n"
        (bundle / f"formula_families_{variant_id}.py").write_text(
            variant_family_source, encoding="utf-8")
        (bundle / f"pricing_{variant_id}.py").write_text(variant["driver"], encoding="utf-8")
        (bundle / f"variant_{variant_id}.json").write_text(
            json.dumps(variant["metadata"], indent=2, sort_keys=True, ensure_ascii=False) + "\n",
            encoding="utf-8")
    metadata_seed_bindings_sha256 = _canonical_hash(metadata_seed_bindings)
    (bundle / "model.py").write_text(
        _model_source(scenario_ids, _canonical_hash(scenario_guard),
                      metadata_seed_bindings_sha256,
                      _canonical_hash(variable_contract) if variable_contract is not None else None), encoding="utf-8")
    (bundle / "family_manifest.json").write_text(
        json.dumps(family_manifest, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    files = {path.name: hash_file(path) for path in sorted(bundle.iterdir()) if path.is_file()}
    raw_business_input_count = sum(item.get("role") == "raw" for item in variables.values())
    external_business_input_count = sum(item.get("role") == "external" for item in variables.values())
    derived_variable_count = sum(item.get("role") == "derived" for item in variables.values())
    manifest = {
        "schema_version": "step4.modular_model.v1", "execution_kind": "modular",
        "source_sha256": source_sha256,
        "design_sha256": design_sha256, "semantic_plan_sha256": semantic_plan_sha256,
        "source_profile_sha256": source_profile_sha256, "trace_sha256": trace_sha256,
        "trace_path": str(trace_path), "trace_source": trace_source, "scenario_ids": scenario_ids,
        "target_names": target_names,
        "formula_count": coverage["active_formula_member_count"],
        "array_member_count": coverage["array_follower_count"],
        "formula_addresses": [item["address"] for item in formula_cells],
        "raw_input_addresses": sorted(source_values),
        "external_input_addresses": sorted(address for binding in external_bindings
                                             for address in binding["addresses"]),
        "active_formula_addresses": [item["address"] for item in formula_cells],
        "ordinary_formula_count": coverage["ordinary_formula_count"],
        "array_follower_count": coverage["array_follower_count"],
        "array_instance_count": coverage["array_instance_count"],
        "active_formula_member_count": coverage["active_formula_member_count"],
        "source_family_count": compiled["source_family_count"],
        "semantic_equation_group_count": family_manifest["semantic_equation_group_count"],
        "compiled_variant_count": compiled["function_count"],
        "verified_request_variant_count": len(variant_sources),
        "business_input_count": raw_business_input_count + external_business_input_count,
        "raw_business_input_count": raw_business_input_count,
        "external_input_count": sum(len(binding["addresses"]) for binding in external_bindings),
        "external_business_input_count": external_business_input_count,
        "model_variable_count": len(variables), "derived_variable_count": derived_variable_count,
        "raw_source_coordinate_count": len(raw_binding_addresses),
        "source_metadata_coordinate_count": len(metadata_addresses),
        "metadata_seed_binding_count": len(metadata_seed_bindings),
        "raw_input_count": len(raw_binding_addresses),
        "external_capture_sha256": external_capture_hash,
        "external_capture_artifact_sha256": external_capture_artifact_sha256,
        "external_values_sha256": external_value_hash,
        "external_input_ranges": ([item["source_range"] for item in external_capture.get("boundaries", [])]
                                  if external_capture else []),
        "external_age_axis": (external_age_axis if external_capture else None),
        "external_formula_cut_member_count": (external_capture.get("formula_cut_member_count")
                                               if external_capture else 0),
        "external_formula_cut_formula_set_sha256": (external_capture.get("formula_cut_formula_set_sha256")
                                                       if external_capture else None),
        "external_age_axis_sha256": (external_capture.get("age_axis", {}).get("value_sha256")
                                      if external_capture else None),
        "implementation_plan_sha256": implementation_plan_sha256,
        "scenario_guard_sha256": _canonical_hash(scenario_guard),
        "scenario_guard_coordinate_count": len(scenario_guard),
        "unmapped_count": 0,
        "files": files, "runtime_dependencies": ["Python standard library only"],
        "source_workbook_at_runtime": False, "project_package_at_runtime": False,
        "formula_cache_inputs": False,
        "scope": ("saved scenario plus exact verified request variants"
                  if variant_sources else "approved saved scenario plus guarded fixed-trace variable inputs"
                  if variable_contract is not None else "approved saved scenario and source-bound logical plan"),
        "variable_input_contract_sha256": (_canonical_hash(variable_contract)
                                           if variable_contract is not None else None),
        "support_certificate_sha256": (_canonical_hash(variable_contract["support_certificate"])
                                       if variable_contract is not None
                                       and "support_certificate" in variable_contract else None),
    }
    (bundle / "model_manifest.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    files["model_manifest.json"] = hash_file(bundle / "model_manifest.json")
    return {"bundle": str(bundle), "bundle_hashes": files, "manifest": manifest,
            "coverage": coverage, "trace": trace, "trace_path": trace_path,
            "trace_hash": trace_sha256, "trace_source": trace_source,
            "source_hash": source_sha256, "source_path": str(source_path),
            "semantic_plan_path": str(semantic_plan_path),
            "source_profile_path": str(source_profile_path), "semantic_map_path": str(semantic_map_path),
            "target_names": target_names,
            "formula_count": coverage["active_formula_member_count"],
            "business_input_count": raw_business_input_count + external_business_input_count,
            "raw_business_input_count": raw_business_input_count,
            "external_business_input_count": external_business_input_count,
            "model_variable_count": len(variables), "derived_variable_count": derived_variable_count,
            "raw_source_coordinate_count": len(raw_binding_addresses),
            "source_metadata_coordinate_count": len(metadata_addresses),
            "raw_input_count": len(raw_binding_addresses), "array_member_count": coverage["array_follower_count"],
            "external_input_count": sum(len(binding["addresses"]) for binding in external_bindings),
            "function_count": compiled["function_count"], "source_family_count": compiled["source_family_count"],
            "path_coverage": path_coverage}
