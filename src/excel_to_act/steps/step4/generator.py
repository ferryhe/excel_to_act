"""Emit a bounded, standalone Python model from the approved source-bound design."""

from __future__ import annotations

import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from excel_to_act.steps.conversion_workflow import (
    append_stage_artifact,
    design_question_lines,
    display_report_value,
    hash_file,
    load_workflow,
    read_json,
    require_stage_approved,
)
from excel_to_act.steps.step3.calculation import _parse_reference
from excel_to_act.steps.step4.implementation import load_current_implementation_plan
from excel_to_act.steps.step4.modular_bundle import build_modular_bundle_files


def _key(address: str) -> str:
    sheet, cell = address.rsplit("!", 1)
    return f"{sheet.casefold()}!{cell.replace('$', '').upper()}"


def _require_modular_design(stage3_payload: dict[str, Any]) -> None:
    if not isinstance(stage3_payload, dict):
        design = None
        analysis = None
    else:
        design = stage3_payload.get("design")
        analysis = stage3_payload.get("analysis", {})
    artifacts = analysis.get("artifacts", {}) if isinstance(analysis, dict) else {}
    plan = artifacts.get("semantic_plan", {}) if isinstance(artifacts, dict) else {}
    check = artifacts.get("semantic_check", {}) if isinstance(artifacts, dict) else {}
    if (not isinstance(design, dict) or not isinstance(plan, dict) or not isinstance(check, dict)
            or design.get("semantic_mapping_required") is not True
            or not plan.get("path") or not plan.get("sha256")
            or not check.get("path") or not check.get("sha256")):
        raise ValueError(
            "Step 4 generates modular bundles only. Return to Step 3 and finish the semantic mapping and check, "
            "then run `step4 plan` with the current implementation input before `step4 generate`."
        )


def _path_implementation_evidence(design: dict[str, Any], trace: dict[str, Any],
                                  formula_addresses: set[str]) -> dict[str, Any]:
    selected = [item for item in design.get("options", [])
                if isinstance(item, dict)
                and (item.get("selected") is True or item.get("selected_for_draft") is True)]
    numerical_ids = {path_id for option in selected for path_id in option.get("numerical_paths", [])
                     if isinstance(path_id, str)}
    condition_only_ids = {path_id for option in selected for path_id in option.get("condition_only_paths", [])
                          if isinstance(path_id, str)}
    formula_cells = [item for item in trace.get("cells", [])
                     if isinstance(item, dict) and isinstance(item.get("formula"), str)]
    evidence: list[dict[str, Any]] = []
    for path in design.get("paths", []):
        if not isinstance(path, dict) or path.get("path_id") not in numerical_ids:
            continue
        formula = path.get("source_formula")
        matches = [item for item in formula_cells if item.get("formula") == formula]
        if len(matches) != 1:
            continue
        address = matches[0].get("address")
        if not isinstance(address, str) or _key(address) not in formula_addresses:
            continue
        evidence.append({"path_id": path["path_id"], "status": "generated",
                         "address": address, "source_formula": formula,
                         "evidence": "exact Stage 3 formula is present in the fresh active trace and generated formula manifest"})
    return {"known_path_ids": list(design.get("coverage", {}).get("known_path_ids", [])),
            "planned_path_ids": list(design.get("coverage", {}).get("planned_path_ids", [])),
            "numerical_path_ids": sorted(numerical_ids),
            "condition_only_path_ids": sorted(condition_only_ids),
            "implemented_path_ids": sorted(item["path_id"] for item in evidence),
            "implementation_evidence": evidence,
            "unproven_numerical_path_ids": sorted(numerical_ids - {item["path_id"] for item in evidence})}


def _saved_scenario_guard(design: dict[str, Any]) -> tuple[list[str], dict[str, Any]]:
    scenarios = design.get("scenarios", [])
    if (not isinstance(scenarios, list) or len(scenarios) != 1
            or not isinstance(scenarios[0], dict) or scenarios[0].get("input_overrides") != []):
        raise ValueError("modular generation requires one approved saved scenario with no overrides")
    scenario = scenarios[0]
    primary = scenario.get("primary_inputs")
    if not isinstance(primary, dict) or not primary:
        raise ValueError("approved saved scenario has no source-bound primary inputs")
    guard: dict[str, Any] = {}
    for name, record in primary.items():
        if not isinstance(name, str) or not isinstance(record, dict):
            raise ValueError("approved scenario primary input is malformed")
        address, value = record.get("source_cell"), record.get("source_literal")
        if not isinstance(address, str) or "!" not in address:
            raise ValueError(f"approved scenario input has no qualified source cell: {name}")
        key = _key(address)
        if key in guard:
            raise ValueError(f"approved scenario repeats a primary input source cell: {address}")
        guard[key] = value
    return [scenario["scenario_id"]], guard


def _modular_oracle_ranges(manifest: dict[str, Any]) -> list[str]:
    """Bound one fresh oracle to all active formulas and external boundary coordinates."""
    bounds_by_sheet: dict[str, tuple[str, int, int, int, int]] = {}
    for address in manifest.get("active_formula_addresses", []):
        parsed = _parse_reference(address, "")
        if parsed is None or not parsed.sheet or not parsed.single:
            raise ValueError(f"active formula address is not a qualified source cell: {address}")
        key = parsed.sheet.casefold()
        previous = bounds_by_sheet.get(key)
        if previous is None:
            bounds_by_sheet[key] = (parsed.sheet, parsed.min_row, parsed.min_col,
                                    parsed.max_row, parsed.max_col)
        else:
            bounds_by_sheet[key] = (previous[0], min(previous[1], parsed.min_row),
                                    min(previous[2], parsed.min_col), max(previous[3], parsed.max_row),
                                    max(previous[4], parsed.max_col))
    ranges = []
    for sheet, min_row, min_col, max_row, max_col in bounds_by_sheet.values():
        start = f"{get_column_letter(min_col)}{min_row}"
        end = f"{get_column_letter(max_col)}{max_row}"
        ranges.append(f"{sheet}!{start}" if start == end else f"{sheet}!{start}:{end}")
    ranges.extend(manifest.get("external_input_ranges", []))
    age_axis = manifest.get("external_age_axis", {})
    if isinstance(age_axis, dict) and isinstance(age_axis.get("source_range"), str):
        ranges.append(age_axis["source_range"])
    ranges = list(dict.fromkeys(ranges))
    requested_cells = set()
    for address in ranges:
        parsed = _parse_reference(address, "")
        if parsed is None or not parsed.sheet:
            raise ValueError(f"planned native oracle range is invalid: {address}")
        for row in range(parsed.min_row, parsed.max_row + 1):
            for column in range(parsed.min_col, parsed.max_col + 1):
                requested_cells.add((parsed.sheet.casefold(), row, column))
                if len(requested_cells) > 50_000:
                    raise ValueError("required Step 5 oracle ranges exceed the 50,000-cell capture limit")
    return ranges


def _bundle_files(stage3_payload: dict[str, Any], trace_path: Path,
                  revision_dir: Path, workflow_root: Path | None = None) -> dict[str, Any]:
    source = stage3_payload["source"]
    source_path = Path(source["workbook_path"]).expanduser().resolve()
    source_hash = source["workbook_sha256"]
    if not source_path.is_file() or hash_file(source_path) != source_hash:
        raise ValueError("bound source workbook is missing or changed")
    trace_path = trace_path.expanduser().resolve()
    trace_hash = hash_file(trace_path)
    historical = stage3_payload["design"].get("historical_evidence", [])
    matching = [item for item in historical if Path(item.get("path", "")).expanduser().resolve() == trace_path]
    discovery_path = trace_path.parent / "discovery.json"
    discovery: dict[str, Any] | None = None
    if matching:
        if any(item.get("sha256") != trace_hash for item in matching):
            raise ValueError("historical trace hash differs from the Stage 3 design evidence")
        trace_source = "historical_stage3_evidence"
    else:
        if not discovery_path.is_file():
            raise ValueError("trace must be the latest Stage 4 discovery bound by the current implementation preflight; "
                             "normally omit `--trace` to use that discovery")
        discovery = read_json(discovery_path)
        if (discovery.get("schema_version") != "step4.discovery.v1" or discovery.get("status") != "pass"
                or Path(discovery.get("active_trace_path", "")).resolve() != trace_path
                or discovery.get("active_trace_sha256") != trace_hash
                or discovery.get("source_sha256") != source_hash
                or discovery.get("stage3_artifact_sha256") != stage3_payload.get("_artifact_sha256")):
            raise ValueError("Stage 4 discovery trace is stale or bound to a different source/design")
        trace_source = "stage4_discovery"
    trace = read_json(trace_path)
    if (not isinstance(trace, dict) or trace.get("source_sha256") != source_hash
            or not isinstance(trace.get("cells"), list)):
        raise ValueError("trace source identity or cell list does not match the approved workbook")
    if trace_source == "stage4_discovery" and trace.get("design_sha256") != stage3_payload.get("_artifact_sha256"):
        raise ValueError("Stage 4 discovery trace was produced from a different approved design")

    design = stage3_payload["design"]
    target_names = {item["selector"] for item in design.get("targets", []) if isinstance(item.get("selector"), str)}
    for item in design.get("validation_plan", []):
        if isinstance(item, dict):
            target_names.update(value for value in item.get("compare", []) if isinstance(value, str) and "!" not in value and ":" not in value)
    trace_names = {(item.get("scope", "workbook"), item.get("name")) for item in trace.get("names", [])}
    target_names = {name for name in target_names
                    if any(record_name and record_name.casefold() == name.casefold() for _, record_name in trace_names)}
    if not target_names:
        raise ValueError("approved design does not identify a target name present in the active trace")

    _require_modular_design(stage3_payload)
    semantic_plan_path: Path | None = None
    source_profile_path: Path | None = None
    semantic_map_path: Path | None = None
    semantic_plan: dict[str, Any] | None = None
    source_profile: dict[str, Any] | None = None
    semantic_check: dict[str, Any] | None = None
    implementation_record: dict[str, Any] | None = None
    implementation_path: Path | None = None
    implementation_sha256: str | None = None
    external_capture: dict[str, Any] | None = None
    external_capture_path: Path | None = None
    external_capture_artifact_sha256: str | None = None
    scenario_ids: list[str] = []
    scenario_guard: dict[str, Any] = {}
    if workflow_root is None:
        raise ValueError("modular generation requires its workflow directory and a current implementation preflight")
    analysis = stage3_payload.get("analysis", {})
    artifacts = analysis.get("artifacts", {}) if isinstance(analysis, dict) else {}
    plan_artifact = artifacts.get("semantic_plan", {}) if isinstance(artifacts, dict) else {}
    check_artifact = artifacts.get("semantic_check", {}) if isinstance(artifacts, dict) else {}
    if not isinstance(plan_artifact, dict) or not isinstance(check_artifact, dict):
        raise ValueError("accepted modular design does not bind its semantic plan and check")
    semantic_plan_path = Path(plan_artifact.get("path", "")).expanduser().resolve()
    check_path = Path(check_artifact.get("path", "")).expanduser().resolve()
    if (not semantic_plan_path.is_file()
            or hash_file(semantic_plan_path) != plan_artifact.get("sha256")
            or not check_path.is_file()
            or hash_file(check_path) != check_artifact.get("sha256")):
        raise ValueError("bound semantic plan or source-only check is missing or changed")
    semantic_plan = read_json(semantic_plan_path)
    semantic_check = read_json(check_path)
    if semantic_plan.get("status") != "pass" or semantic_check.get("status") != "pass":
        raise ValueError("accepted modular design has no passing source-only semantic plan/check")
    if (semantic_check.get("semantic_plan_sha256") != plan_artifact.get("sha256")
            or semantic_check.get("source_family_profile_sha256") != semantic_plan.get("evidence", {}).get(
                "source_family_profile", {}).get("sha256")):
        raise ValueError("semantic check does not bind the current plan and source-family profile")
    map_record = semantic_plan.get("semantic_map_input", {})
    profile_record = semantic_plan.get("evidence", {}).get("source_family_profile", {})
    semantic_map_path = Path(map_record.get("path", "")).expanduser().resolve()
    source_profile_path = Path(profile_record.get("path", "")).expanduser().resolve()
    if (not semantic_map_path.is_file() or hash_file(semantic_map_path) != map_record.get("sha256")
            or not source_profile_path.is_file()
            or hash_file(source_profile_path) != profile_record.get("sha256")):
        raise ValueError("bound semantic map or source-family profile is missing or changed")
    if (semantic_plan.get("source", {}).get("source_sha256") != source_hash
            or semantic_check.get("semantic_map_input_sha256") != map_record.get("sha256")
            or semantic_check.get("source_family_profile_sha256") != profile_record.get("sha256")):
        raise ValueError("semantic plan/check source bindings differ from the approved workbook")
    source_profile = read_json(source_profile_path)
    implementation_record, implementation_path, implementation_sha256, validated = \
        load_current_implementation_plan(workflow_root)
    if Path(validated["trace_path"]).resolve() != trace_path:
        raise ValueError("generation trace differs from the active trace bound by the current implementation preflight")
    semantic_plan = validated["projected_plan"]
    source_profile = validated["projected_profile"]
    external_capture = validated["capture"]
    external_capture_path = validated["capture_path"]
    external_capture_artifact_sha256 = validated["capture_file_sha256"]
    scenario_ids, scenario_guard = _saved_scenario_guard(design)

    before = hash_file(source_path)
    workbook = load_workbook(source_path, data_only=False, read_only=False, keep_vba=False)
    try:
        raw_values: dict[str, Any] = {}
        formulas: dict[str, str] = {}
        arrays: dict[str, tuple[int, int, int, int]] = {}

        for record in trace["cells"]:
            address = record.get("address")
            sheet = record.get("sheet")
            role = record.get("role")
            if not isinstance(address, str) or not isinstance(sheet, str):
                raise ValueError("trace contains a cell without a source address")
            key = _key(address)
            if role == "source_value":
                cell = workbook[sheet][address.rsplit("!", 1)[1]]
                if cell.data_type == "f":
                    raise ValueError(f"formula cell was labeled as raw source input: {address}")
                value = cell.value
                if cell.data_type == "e" and isinstance(value, str):
                    value = {"excel_error": value.upper()}
                elif hasattr(value, "isoformat"):
                    value = {"excel_datetime": value.isoformat(), "date_only": type(value).__name__ == "date"}
                elif not (isinstance(value, (str, int, float, bool)) or value is None):
                    raise ValueError(f"unsupported raw source value at {address}: {type(value).__name__}")
                raw_values[key] = value
                continue
            if role == "external_boundary_input":
                continue
            if role not in {"calculated_formula", "calculated_array_formula"}:
                raise ValueError(f"unrecognized trace cell role at {address}: {role}")
            source_cell = workbook[sheet][address.rsplit("!", 1)[1]]
            source_array = source_cell.value if hasattr(source_cell.value, "text") and hasattr(source_cell.value, "ref") else None
            array_ref = record.get("array_ref") or (source_array.ref if source_array is not None else None)
            if array_ref:
                bounds = range_boundaries(array_ref)
                if any(value is None for value in bounds):
                    raise ValueError(f"array formula has incomplete bounds at {address}: {array_ref}")
                min_col, min_row, max_col, max_row = bounds
                anchor = f"{sheet}!{get_column_letter(min_col)}{min_row}"
                anchor_key = _key(anchor)
                arrays[anchor_key] = (min_row, min_col, max_row, max_col)

        for record in trace["cells"]:
            if record.get("role") != "calculated_formula" or record.get("array_ref"):
                continue
            address, sheet = record["address"], record["sheet"]
            if _key(address) in arrays:
                # The source workbook identifies the anchor's native array range even when
                # the active trace labels only its member cells with array_ref.
                continue
            cell = workbook[sheet][address.rsplit("!", 1)[1]]
            if cell.data_type != "f" or not isinstance(cell.value, str) or not cell.value.startswith("="):
                raise ValueError(f"active formula is missing from the current source: {address}")
            if record.get("formula") != cell.value:
                raise ValueError(f"active formula text differs from the current source: {address}")
            formulas[_key(address)] = cell.value

        for anchor_key, bounds in arrays.items():
            sheet, anchor_cell = anchor_key.split("!", 1)
            source_sheet = next((item.title for item in workbook.worksheets if item.title.casefold() == sheet), None)
            if source_sheet is None:
                raise ValueError(f"array formula sheet is missing: {sheet}")
            cell = workbook[source_sheet][anchor_cell]
            if not hasattr(cell.value, "text") or not hasattr(cell.value, "ref"):
                raise ValueError(f"array formula anchor is not present in source: {source_sheet}!{anchor_cell}")
            formula_text = cell.value.text
            ref = cell.value.ref
            actual_bounds = range_boundaries(ref)
            if tuple(actual_bounds) != (bounds[1], bounds[0], bounds[3], bounds[2]):
                raise ValueError(f"array range changed at {source_sheet}!{anchor_cell}")
            formulas[anchor_key] = formula_text
        path_coverage = _path_implementation_evidence(design, trace, set(formulas))


    finally:
        workbook.close()
    if hash_file(source_path) != before or before != source_hash:
        raise ValueError("source workbook changed while the generated bundle was being built")

    if (semantic_plan is None or source_profile is None or semantic_plan_path is None
            or source_profile_path is None or semantic_map_path is None
            or implementation_record is None or implementation_path is None
            or external_capture is None or external_capture_path is None):
        raise ValueError("current modular generation inputs were not loaded")
    built = build_modular_bundle_files(
        semantic_plan=semantic_plan, source_profile=source_profile, trace=trace,
        raw_source_values=raw_values, source_path=source_path,
        semantic_plan_path=semantic_plan_path, source_profile_path=source_profile_path,
        semantic_map_path=semantic_map_path, revision_dir=revision_dir,
        target_names=sorted(target_names), path_coverage=path_coverage,
        source_sha256=source_hash, design_sha256=stage3_payload["_artifact_sha256"],
        semantic_plan_sha256=hash_file(semantic_plan_path),
        source_profile_sha256=hash_file(source_profile_path), trace_sha256=trace_hash,
        trace_path=trace_path, trace_source=trace_source,
        scenario_ids=scenario_ids, scenario_guard=scenario_guard,
        external_capture=external_capture,
        external_capture_artifact_sha256=external_capture_artifact_sha256,
        implementation_plan_record=implementation_record,
        implementation_plan_path=implementation_path,
        implementation_plan_sha256=implementation_sha256,
    )
    built.update({"discovery_path": discovery_path if discovery is not None else None,
                  "input_files": {"semantic_plan": semantic_plan_path,
                                  "source_family_profile": source_profile_path,
                                  "semantic_map": semantic_map_path,
                                  "implementation_plan": implementation_path,
                                  "implementation_input": Path(implementation_record["implementation_input"]["path"]),
                                  "external_capture": external_capture_path}})
    built["oracle_ranges"] = _modular_oracle_ranges(built["manifest"])
    return built

def _latest_discovery_trace(root: Path, stage3_revision: dict[str, Any], source_hash: str) -> Path:
    discovery_root = root / "stage4" / "discovery"
    candidates = sorted(discovery_root.glob("revision-*/discovery.json"), reverse=True) if discovery_root.exists() else []
    for path in candidates:
        try:
            record = read_json(path)
            trace_path = Path(record["active_trace_path"]).resolve()
            if (record.get("status") == "pass"
                    and record.get("stage3_artifact_sha256") == stage3_revision["artifact"]["json_sha256"]
                    and record.get("source_sha256") == source_hash
                    and trace_path.is_file()
                    and hash_file(trace_path) == record.get("active_trace_sha256")):
                return trace_path
        except (KeyError, OSError, ValueError, TypeError):
            continue
    raise ValueError("no current Stage 4 discovery trace exists; run `step4 discover --workflow DIR` first")


def generate_model(workflow_dir: Path, trace_path: Path | None = None) -> dict[str, Any]:
    """Create Step 4 code only after an approved Step 3 design artifact."""
    bundle_revision_dir: Path | None = None
    created_bundle_dir = False
    stage4_root: Path | None = None
    try:
        root, workflow = load_workflow(workflow_dir)
        _manifest, stage3_revision = require_stage_approved(root, 3)
        stage3_json = root / stage3_revision["artifact"]["json"]
        stage3_payload = read_json(stage3_json)
        stage3_payload["_artifact_sha256"] = stage3_revision["artifact"]["json_sha256"]
        _require_modular_design(stage3_payload)
        load_current_implementation_plan(root)
        if trace_path is None:
            trace_path = _latest_discovery_trace(root, stage3_revision, workflow["source"]["workbook_sha256"])
        next_revision = len(workflow["stages"]["4"]["revisions"]) + 1
        stage4_root = (root / "stage4").resolve()
        revision_name = f"revision-{next_revision:04d}"
        report_revision_dir = stage4_root / revision_name
        bundle_revision_dir = stage4_root / "bundles" / revision_name
        if (not report_revision_dir.resolve().is_relative_to(stage4_root)
                or not bundle_revision_dir.resolve().is_relative_to(stage4_root)):
            raise ValueError("Step 4 output path resolves outside this workflow's Stage 4 folder")
        if report_revision_dir.exists() or bundle_revision_dir.exists():
            raise ValueError("Step 4 output directory already exists; preserve it and use a new workflow root")
        bundle_revision_dir.mkdir(parents=True)
        created_bundle_dir = True
        built = _bundle_files(stage3_payload, trace_path, bundle_revision_dir, root)
        smoke = _smoke_bundle(Path(built["bundle"]), bundle_revision_dir / "smoke", built["manifest"]["target_names"])
        manifest_path = Path(built["bundle"]) / "model_manifest.json"
        payload = {"schema_version": "step4.generation_report.v1", "tool": "step4.generate",
                   "status": "ready_for_review", "workflow": str(root),
                   "source_sha256": built["source_hash"], "design_sha256": stage3_revision["artifact"]["json_sha256"],
                   "scenario_ids": built["manifest"]["scenario_ids"],
                   "trace_sha256": built["trace_hash"],
                   "active_trace": {"path": str(built["trace_path"]), "sha256": hash_file(built["trace_path"]),
                                    "source": built["trace_source"]},
                   "bundle": {"path": str(Path(built["bundle"])), "manifest_path": str(manifest_path),
                              "files": built["bundle_hashes"]},
                   "formula_count": built["manifest"]["formula_count"],
                   "raw_input_count": built["manifest"]["raw_input_count"],
                   "array_member_count": built["manifest"]["array_member_count"],
                   "array_instance_count": built["manifest"].get("array_instance_count"),
                   "compiled_function_count": built.get("function_count"),
                   "source_family_count": built.get("source_family_count"),
                   "semantic_equation_group_count": built["manifest"].get("semantic_equation_group_count"),
                   "target_names": built["manifest"]["target_names"],
                   "oracle_ranges": built.get("oracle_ranges"),
                   "business_input_count": built["manifest"].get("business_input_count"),
                   "raw_business_input_count": built["manifest"].get("raw_business_input_count"),
                   "external_business_input_count": built["manifest"].get("external_business_input_count"),
                   "external_input_coordinate_count": built["manifest"].get("external_input_count"),
                   "external_capture": ({"artifact_sha256": built["manifest"].get("external_capture_artifact_sha256"),
                                         "capture_sha256": built["manifest"].get("external_capture_sha256"),
                                         "formula_cut_member_count": built["manifest"].get(
                                             "external_formula_cut_member_count"),
                                         "formula_cut_formula_set_sha256": built["manifest"].get(
                                             "external_formula_cut_formula_set_sha256"),
                                         "ranges": built["manifest"].get("external_input_ranges", []),
                                         "age_axis": built["manifest"].get("external_age_axis")}
                                        if built["manifest"].get("external_capture_sha256") else None),
                   "model_variable_count": built["manifest"].get("model_variable_count"),
                   "derived_variable_count": built["manifest"].get("derived_variable_count"),
                   "raw_source_coordinate_count": built["manifest"].get("raw_source_coordinate_count"),
                   "source_metadata_coordinate_count": built["manifest"].get("source_metadata_coordinate_count"),
                   "metadata_seed_binding_count": built["manifest"].get("metadata_seed_binding_count"),
                   "path_coverage": built["path_coverage"],
                   "runtime_dependencies": ["Python standard library only"],
                   "source_workbook_at_runtime": False, "project_package_at_runtime": False,
                   "formula_cache_inputs": False, "smoke_execution": "pass", "smoke": smoke}
        markdown = _generation_markdown(
            payload,
            design=stage3_payload.get("design", {}),
            design_link=(Path("../../stage3") / stage3_json.parent.name / stage3_json.name).as_posix(),
            entry_point_link=(Path("..") / "bundles" / revision_name / "bundle" / "model.py").as_posix(),
            manifest_link=(Path("..") / "bundles" / revision_name / "bundle" / "model_manifest.json").as_posix(),
        )
        inputs = {"step3_design": stage3_json, "source_workbook": Path(built["source_path"]),
                  "active_trace": built["trace_path"],
                  "generation_smoke": Path(smoke["summary_path"]),
                  "generation_smoke_result": Path(smoke["result_path"])}
        if built["discovery_path"] is not None:
            inputs["stage4_discovery"] = built["discovery_path"]
        inputs.update(built.get("input_files", {}))
        inputs.update({f"generated:{name}": Path(built["bundle"]) / name for name in built["bundle_hashes"]})
        revision = append_stage_artifact(root, 4, "generation_report.json", payload, markdown,
                                         input_files=inputs,
                                         input_stages={3: stage3_revision["artifact"]["json_sha256"]})
        return {"schema_version": payload["schema_version"], "tool": payload["tool"], "status": "pass",
                "workflow": str(root), "revision": revision["revision"], "artifact": revision["artifact"],
                "bundle": payload["bundle"], "counts": {"formula_count": payload["formula_count"],
                                                            "raw_input_count": payload["raw_input_count"],
                                                            "array_member_count": payload["array_member_count"],
                                                            "metadata_seed_binding_count": payload.get(
                                                                "metadata_seed_binding_count", 0)}}
    except Exception as exc:
        if (created_bundle_dir and bundle_revision_dir is not None and bundle_revision_dir.exists()
                and stage4_root is not None and bundle_revision_dir.resolve().is_relative_to(stage4_root)):
            shutil.rmtree(bundle_revision_dir, ignore_errors=True)
        return {"schema_version": "step4.generation_report.v1", "tool": "step4.generate",
                "status": "blocked", "reason": str(exc)}


def _generation_evidence_markdown(payload: dict[str, Any]) -> str:
    smoke = payload.get("smoke", {})
    coverage = payload.get("path_coverage", {})
    external = payload.get("external_capture")
    lines = ["# Step 4 Python Generation", "", "Status: **Ready for review**", "",
             f"Source SHA-256: `{payload['source_sha256']}`",
             f"Approved design SHA-256: `{payload['design_sha256']}`",
             f"Active trace source: `{payload['active_trace']['source']}`; SHA-256: `{payload['active_trace']['sha256']}`", "",
             f"Covered source formula positions: {payload['formula_count']}",
             f"Raw source coordinates: {payload.get('raw_source_coordinate_count', payload['raw_input_count'])}",
             (f"Business inputs: {payload.get('business_input_count')} ({payload.get('raw_business_input_count')} raw, "
              f"{payload.get('external_business_input_count')} formula-derived external); model variables: "
              f"{payload.get('model_variable_count')} ({payload.get('derived_variable_count')} derived)."
              if payload.get("business_input_count") is not None else ""),
             (f"Raw source coordinates: {payload.get('raw_source_coordinate_count')}; adapter metadata coordinates: "
              f"{payload.get('source_metadata_coordinate_count')}; external vector coordinates: "
              f"{payload.get('external_input_coordinate_count')}."
              if payload.get("raw_source_coordinate_count") is not None else ""),
             (f"Derived-variable initial conditions sourced from accepted adapter metadata: "
              f"{payload.get('metadata_seed_binding_count')}."
              if payload.get("metadata_seed_binding_count") is not None else ""),
             (f"External capture: {len(external.get('ranges', []))} declared vector ranges; "
              f"{external.get('formula_cut_member_count', 'Not recorded')} formula cut members; "
              f"age axis `{(external.get('age_axis') or {}).get('source_range', 'Not recorded')}`."
              if isinstance(external, dict) else "No external formula-derived inputs are declared."),
             (f"Compiled direct formula variants: {payload['compiled_function_count']} across "
              f"{payload['source_family_count']} source families"
              + (f" and {payload['semantic_equation_group_count']} semantic equation groups."
                 if payload["semantic_equation_group_count"] is not None else ".")
              if payload["compiled_function_count"] is not None else ""),
             f"Implemented numerical paths in the selected design: {len(coverage['implemented_path_ids'])} / {len(coverage['numerical_path_ids'])}.",
             "Implemented path IDs: " + (", ".join(
                 f"{item['path_id']} at {item['address']}" for item in coverage["implementation_evidence"])
                 or "none proven by an exact source-formula match"),
             "Condition-only paths are tracked separately and do not produce a solver: "
             + (", ".join(coverage["condition_only_path_ids"]) or "none"),
             (f"Array output members: {payload['array_member_count']}"
              + (f" across {payload['array_instance_count']} instances."
                 if payload["array_instance_count"] is not None else ".")),
             "Runtime imports: Python standard library only.",
             "The generated folder contains no workbook and never imports the project package.",
             "Formula caches are not included as constants.",
             f"Stage 4 standalone smoke passed for {smoke['formula_count']} formula cells and {len(smoke['targets'])} named targets.",
             "Stage 5 repeats isolated validation and performs a separate native Excel reconciliation.",
             "Review decisions and current receipts are recorded in workflow.json.", ""]
    return "\n".join(lines)


def _readable_value(value: Any, limit: int = 12) -> tuple[str, bool]:
    if not isinstance(value, (list, dict)):
        return display_report_value(value).replace("|", "\\|").replace("\r\n", " ").replace("\n", " "), False

    leaves: list[tuple[tuple[Any, ...], Any]] = []

    def visit(item: Any, path: tuple[Any, ...]) -> None:
        if len(leaves) > limit:
            return
        if isinstance(item, list):
            if not item:
                leaves.append((path, "[]"))
            else:
                for index, child in enumerate(item):
                    visit(child, (*path, index))
                    if len(leaves) > limit:
                        return
        elif isinstance(item, dict):
            if not item:
                leaves.append((path, "{}"))
            else:
                for key, child in item.items():
                    visit(child, (*path, key))
                    if len(leaves) > limit:
                        return
        else:
            leaves.append((path, item))

    visit(value, ())
    rows = []
    for path, item in leaves[:limit]:
        location = (f"position {list(path)}" if all(isinstance(part, int) for part in path)
                    else f"field path {list(path)}")
        shown = display_report_value(item).replace("|", "\\|").replace("\r\n", " ").replace("\n", " ")
        rows.append(f"{location} = {shown}")
    if len(leaves) > limit:
        rows.append("Additional values omitted; see the linked machine report")
    return "<br>".join(rows), True


def _axis_summary(target: dict[str, Any]) -> str:
    axes = target.get("axes", target.get("axis"))
    if axes is None:
        return "Not recorded"
    if not isinstance(axes, list):
        return display_report_value(axes)
    if not axes:
        return "[]"
    summaries = []
    for axis in axes:
        if not isinstance(axis, dict):
            summaries.append(display_report_value(axis))
            continue
        details = []
        for key in ("name", "axis_id", "role", "units", "source_range"):
            if key in axis:
                details.append(f"{key}: {display_report_value(axis[key])}")
        if "keys" in axis:
            details.append(f"keys: {_readable_value(axis['keys'], limit=6)[0]}")
        summaries.append("; ".join(details) or "Not recorded")
    return " / ".join(summaries).replace("|", "\\|")


def _generation_markdown(
    payload: dict[str, Any], *, design: dict[str, Any] | None = None,
    machine_link: str = "generation_report.json", workflow_link: str = "../../workflow.json",
    design_link: str | None = None, entry_point_link: str | None = None,
    manifest_link: str | None = None,
) -> str:
    design = design if isinstance(design, dict) else {}
    raw_targets = design.get("targets")
    design_targets = raw_targets if isinstance(raw_targets, list) else []
    ordered_targets = sorted(
        (item for item in design_targets if isinstance(item, dict)),
        key=lambda item: item.get("result_order") if isinstance(item.get("result_order"), int)
        and not isinstance(item.get("result_order"), bool) else 10**9,
    )
    requested = {item.get("selector") for item in ordered_targets if isinstance(item.get("selector"), str)}
    diagnostics = []
    for item in ordered_targets:
        names = item.get("diagnostic_intermediates")
        if isinstance(names, list):
            diagnostics.extend(name for name in names if isinstance(name, str) and name not in diagnostics)
    smoke = payload.get("smoke")
    smoke = smoke if isinstance(smoke, dict) else {}
    bundle_record = payload.get("bundle")
    bundle_path = bundle_record.get("path") if isinstance(bundle_record, dict) else None
    bundle = Path(bundle_path) if isinstance(bundle_path, str) else Path(".")
    raw_values = smoke.get("targets")
    values = raw_values if isinstance(raw_values, dict) else {}
    coverage_record = payload.get("path_coverage")
    coverage = coverage_record if isinstance(coverage_record, dict) else {}
    implemented = coverage.get("implemented_path_ids")
    implemented = implemented if isinstance(implemented, list) else None
    numerical = coverage.get("numerical_path_ids")
    numerical = numerical if isinstance(numerical, list) else None
    entry_point = entry_point_link or (bundle / "model.py").as_posix()
    bundle_manifest = manifest_link or (Path(entry_point).parent / "model_manifest.json").as_posix()
    lines = [
        "# Step 4 Python Generation", "",
        "## What this report asks you to accept", "",
        "The generated standalone Python bundle and its source-bound implementation evidence. This accepts the Step 4 code candidate only; Step 5 performs the separate native Excel reconciliation.", "",
        "## Requested targets", "",
    ]
    has_array_values = False
    if ordered_targets:
        for target in ordered_targets:
            selector = display_report_value(target.get("selector"))
            raw_result = values.get(target.get("selector")) if isinstance(target.get("selector"), str) else None
            result = raw_result if target.get("selector") in values else None
            readable, non_scalar = _readable_value(result)
            has_array_values = has_array_values or non_scalar
            shape = display_report_value(target.get("shape"))
            units = display_report_value(target.get("units"))
            axes = _axis_summary(target)
            lines.append(
                f"- **{selector}** — {display_report_value(target.get('result_kind'))}; "
                f"shape {shape}; units {units}; axes {axes}."
            )
            rendered = readable if non_scalar else f"`{readable}`"
            lines.append(f"  - Generated smoke result: {rendered} (standalone execution only).")
    else:
        lines.append("Requested target: Not recorded in the bound design.")
    lines.extend([
        "",
        "## Purpose and scope", "",
        f"Generated formula positions: {display_report_value(payload.get('formula_count'))}; "
        f"business input objects: {display_report_value(payload.get('business_input_count'))} "
        f"({display_report_value(payload.get('raw_business_input_count'))} raw, "
        f"{display_report_value(payload.get('external_business_input_count'))} formula-derived external).",
        "The bundle was generated from the bound Step 3 design and current active trace.", "",
        "## Verified results", "",
        f"- Generated bundle compile and isolated smoke: **{display_report_value(smoke.get('status'))}**; "
        f"{display_report_value(smoke.get('formula_count'))} formula cells and "
        f"{display_report_value(len(values) if isinstance(raw_values, dict) else None)} named outputs.",
        f"- Numerical routes implemented: {display_report_value(len(implemented) if implemented is not None else None)} / "
        f"{display_report_value(len(numerical) if numerical is not None else None)}; condition-only routes are tracked separately and do not produce a solver.",
    ])
    diagnostic_names = [name for name in diagnostics if name in values]
    unclassified_names = [name for name in values if name not in requested and name not in diagnostics]
    for heading, names in (("Generated smoke diagnostics (not requested business targets):", diagnostic_names),
                           ("Other named outputs (classification not recorded):", unclassified_names)):
        if names:
            lines.append(f"- {heading}")
            for name in names:
                readable, non_scalar = _readable_value(values[name])
                has_array_values = has_array_values or non_scalar
                rendered = readable if non_scalar else f"`{readable}`"
                lines.append(f"  - **{name}**: {rendered}.")
    lines.extend([
        "", "## Limits and unresolved items", "",
        "The smoke checks generated-code compilation and execution in isolation. It does not compare results with Excel. The bundle contains no source workbook, does not import the project package, and does not use formula-cache values as constants.",
        "Unknown classifications remain unclassified; this report does not infer business meaning from cell names or coordinates.",
    ])
    if has_array_values:
        lines.append("Array and table values above show explicit zero-based positions or recorded field paths; the full values remain in the linked machine report. Positions do not imply a time or business-axis origin.")
    lines.extend(design_question_lines(design.get("open_questions")))
    lines.extend([
        "", "## Handover and review", "",
        f"- Machine generation report: [generation_report.json](<{machine_link}>).",
        f"- Current acceptance ledger: [workflow.json](<{workflow_link}>).",
        f"- Generated entry point: [model.py](<{Path(entry_point).as_posix()}>); manifest: [model_manifest.json](<{Path(bundle_manifest).as_posix()}>).",
        "- Supported run command from the generated bundle directory: `python model.py --out RESULT.json`.",
    ])
    if design_link:
        lines.append(f"- Bound Step 3 design: [analysis_design.json](<{Path(design_link).as_posix()}>).")
    lines.extend([
        "- Next: Agent reviews this exact JSON/Markdown pair first, then the actual human reviews it unless a valid scoped TypeSafe delegation applies. A successful compile/smoke is technical evidence; it is not Agent PASS or human acceptance.",
        "- To reject: route the decision to the responsible current or earlier Step, retain this revision, refresh affected reports, and check the ledger before continuing.",
        "- Report status `ready_for_review` describes generated evidence. Current formal acceptance is read from the linked workflow ledger.", "",
        "## Detailed generation evidence", "",
        re.sub(r"(?m)^## ", "#### ", _generation_evidence_markdown(payload).replace(
            "# Step 4 Python Generation", "### Original generation details", 1)), "",
    ])
    return "\n".join(lines)

def _smoke_bundle(bundle: Path, smoke_dir: Path, target_names: list[str]) -> dict[str, Any]:
    """Compile and run the generated active trace without the workbook or project package."""
    python_files = sorted(bundle.glob("*.py"))
    if not python_files:
        raise ValueError("generated bundle contains no Python modules")
    for path in python_files:
        compile(path.read_text(encoding="utf-8"), path.name, "exec")
    smoke_dir.mkdir(parents=True)
    with tempfile.TemporaryDirectory(prefix="step4-generated-smoke-") as temporary:
        isolated = Path(temporary)
        for path in bundle.iterdir():
            if path.is_file():
                shutil.copy2(path, isolated / path.name)
        result_path = isolated / "model_result.json"
        runner = ("import runpy,sys; model_dir=sys.argv[1]; output=sys.argv[2]; "
                  "sys.path.insert(0,model_dir); sys.argv=[model_dir+'/model.py','--out',output]; "
                  "runpy.run_path(sys.argv[0],run_name='__main__')")
        completed = subprocess.run(
            [sys.executable, "-I", "-c", runner, str(isolated), str(result_path)],
            cwd=isolated, capture_output=True, text=True, timeout=600, check=False,
        )
        if completed.returncode != 0 or not result_path.is_file():
            detail = (completed.stderr or completed.stdout).strip()[-3000:]
            raise ValueError(f"generated standalone smoke failed with exit {completed.returncode}: {detail}")
        result_bytes = result_path.read_bytes()
        result = json.loads(result_bytes)
    actual_targets = result.get("targets", {})
    missing = sorted(set(target_names) - set(actual_targets))
    if missing:
        raise ValueError("generated standalone smoke omitted named targets: " + ", ".join(missing))
    saved_result = smoke_dir / "model_result.json"
    saved_result.write_bytes(result_bytes)
    summary = {"schema_version": "step4.generation_smoke.v1", "status": "pass",
               "compile_pass": True, "standalone_run_pass": True,
               "source_workbook_present": False, "project_package_required": False,
               "formula_cache_inputs": False, "targets": actual_targets,
               "cell_count": result.get("cell_count"), "formula_count": result.get("formula_count"),
               "result_path": str(saved_result), "result_sha256": hash_file(saved_result)}
    summary_path = smoke_dir / "generation_smoke.json"
    summary_path.write_text(json.dumps(summary, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
    return {**summary, "summary_path": str(summary_path)}


def tool_catalog() -> dict[str, Any]:
    from excel_to_act.steps.step4.tools import tool_catalog as catalog

    return catalog()
