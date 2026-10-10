"""Capture and validate the approved formula-derived Step 3 input boundaries."""

from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
import shutil
from typing import Any
from importlib.resources import files

from openpyxl import load_workbook

from excel_to_act.steps.conversion_workflow import (
    hash_file,
    load_workflow,
    read_json,
    require_stage_approved,
)
from excel_to_act.steps.step3.calculation import _parse_reference, capture_excel_oracle


def _json_bytes(value: Any) -> bytes:
    return (json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":")) + "\n").encode("utf-8")


def _json_hash(value: Any) -> str:
    return hashlib.sha256(_json_bytes(value)).hexdigest()


def _canonical_cell(address: str) -> str:
    reference = _parse_reference(address, "")
    if reference is None or not reference.single or not reference.sheet:
        raise ValueError(f"source address must be one qualified cell: {address}")
    from openpyxl.utils.cell import get_column_letter
    return f"{reference.sheet}!{get_column_letter(reference.min_col)}{reference.min_row}"


def _canonical_range(address: str) -> str:
    reference = _parse_reference(address, "")
    if reference is None or not reference.sheet:
        raise ValueError(f"source extent must be a qualified finite range: {address}")
    from openpyxl.utils.cell import get_column_letter
    start = f"{get_column_letter(reference.min_col)}{reference.min_row}"
    end = f"{get_column_letter(reference.max_col)}{reference.max_row}"
    return f"{reference.sheet}!{start}" if start == end else f"{reference.sheet}!{start}:{end}"


def _addresses(address: str) -> list[str]:
    reference = _parse_reference(address, "")
    if reference is None or not reference.sheet:
        raise ValueError(f"source extent must be a qualified finite range: {address}")
    from openpyxl.utils.cell import get_column_letter
    return [f"{reference.sheet}!{get_column_letter(column)}{row}"
            for row in range(reference.min_row, reference.max_row + 1)
            for column in range(reference.min_col, reference.max_col + 1)]


def _step3_context(workflow_dir: Path) -> dict[str, Any]:
    root, manifest = load_workflow(workflow_dir)
    _workflow, revision = require_stage_approved(root, 3)
    artifact = revision["artifact"]
    stage3_path = root / artifact["json"]
    stage3 = read_json(stage3_path)
    source = manifest.get("source", {})
    source_path = Path(source.get("workbook_path", "")).expanduser().resolve()
    source_hash = source.get("workbook_sha256")
    if not source_path.is_file() or not isinstance(source_hash, str) or hash_file(source_path) != source_hash:
        raise ValueError("approved source workbook is missing or changed")
    if stage3.get("source", {}).get("workbook_sha256") != source_hash:
        raise ValueError("approved Stage 3 design is bound to another source workbook")
    scenario_list = stage3.get("design", {}).get("scenarios", [])
    if len(scenario_list) != 1 or not isinstance(scenario_list[0], dict):
        raise ValueError("external capture requires exactly one approved saved scenario")
    scenario = scenario_list[0]
    if (scenario.get("input_overrides") != [] or not isinstance(scenario.get("scenario_id"), str)
            or not scenario["scenario_id"] or not isinstance(scenario.get("scenario_sha256"), str)):
        raise ValueError("external capture requires a hash-bound approved scenario with no overrides")

    boundary_reference = stage3.get("input_boundary_reference")
    pointer = manifest.get("input_boundary")
    if (not isinstance(boundary_reference, dict) or not isinstance(pointer, dict)
            or boundary_reference.get("schema_version") != "step3.input_boundary.reference.v1"
            or boundary_reference.get("revision") != pointer.get("revision")
            or boundary_reference.get("boundary_sha256") != pointer.get("boundary_sha256")
            or boundary_reference.get("artifact") != pointer.get("artifact")
            or pointer.get("source_sha256") != source_hash):
        raise ValueError("approved design does not bind the current confirmed input catalog")
    catalog_artifact = boundary_reference["artifact"]
    catalog_path = (root / catalog_artifact["json"]).resolve()
    catalog_md_path = (root / catalog_artifact["md"]).resolve()
    if (not catalog_path.is_file() or hash_file(catalog_path) != catalog_artifact.get("json_sha256")
            or not catalog_md_path.is_file() or hash_file(catalog_md_path) != catalog_artifact.get("md_sha256")):
        raise ValueError("confirmed input catalog artifact pair is stale or missing")
    catalog = read_json(catalog_path)
    receipts = boundary_reference.get("review_receipts")
    if not isinstance(receipts, dict):
        raise ValueError("approved design is missing the catalog review receipts")
    for reviewer in ("agent", "human"):
        receipt = receipts.get(reviewer)
        if (not isinstance(receipt, dict) or receipt.get("decision") != "approve"
                or receipt.get("stage") != 3 or receipt.get("revision") != boundary_reference.get("revision")
                or receipt.get("artifact_sha256") != catalog_artifact.get("json_sha256")
                or receipt.get("artifact_md_sha256") != catalog_artifact.get("md_sha256")):
            raise ValueError(f"confirmed input catalog is missing its current {reviewer} approval")
    if catalog.get("boundary_sha256") != boundary_reference.get("boundary_sha256"):
        raise ValueError("confirmed input catalog boundary hash differs from the approved design")
    if catalog.get("source", {}).get("workbook_sha256") != source_hash:
        raise ValueError("confirmed input catalog source differs from the approved source")
    target_selection = catalog.get("target_selection", {})
    if (target_selection.get("scenario", {}).get("scenario_sha256") != scenario.get("scenario_sha256")
            or target_selection.get("scenario", {}).get("scenario_id") != scenario.get("scenario_id")):
        raise ValueError("confirmed catalog scenario differs from the approved design scenario")

    analysis = stage3.get("analysis", {}).get("artifacts", {})
    semantic_record = analysis.get("semantic_plan")
    check_record = analysis.get("semantic_check")
    if not isinstance(semantic_record, dict) or not isinstance(check_record, dict):
        raise ValueError("approved modular design does not bind a semantic plan and check")
    semantic_path = Path(semantic_record.get("path", "")).expanduser().resolve()
    check_path = Path(check_record.get("path", "")).expanduser().resolve()
    if (not semantic_path.is_file() or hash_file(semantic_path) != semantic_record.get("sha256")
            or not check_path.is_file() or hash_file(check_path) != check_record.get("sha256")):
        raise ValueError("approved semantic plan or check is stale or missing")
    semantic_plan = read_json(semantic_path)
    semantic_check = read_json(check_path)
    if (semantic_plan.get("status") != "pass" or semantic_check.get("status") != "pass"
            or semantic_plan.get("source", {}).get("source_sha256") != source_hash
            or semantic_check.get("semantic_plan_sha256") != semantic_record.get("sha256")):
        raise ValueError("approved semantic plan/check is not current for this source")

    evidence = semantic_plan.get("evidence", {})
    evidence_paths: dict[str, tuple[Path, str]] = {}
    for name, record in {
        "model_scope": evidence.get("model_scope"),
        "candidate_trace": evidence.get("source_candidate_trace"),
        "candidate_profile": evidence.get("source_family_profile"),
    }.items():
        if not isinstance(record, dict):
            raise ValueError(f"semantic plan is missing bound {name} evidence")
        path = Path(record.get("path", "")).expanduser().resolve()
        digest = record.get("sha256")
        if not path.is_file() or not isinstance(digest, str) or hash_file(path) != digest:
            raise ValueError(f"bound {name} evidence is stale or missing")
        evidence_paths[name] = (path, digest)
    candidate_trace = read_json(evidence_paths["candidate_trace"][0])
    candidate_profile = read_json(evidence_paths["candidate_profile"][0])

    catalog_variables = {item.get("variable_id"): item for item in catalog.get("variables", [])
                         if isinstance(item, dict) and item.get("role") == "formula_derived_external"}
    plan_variables = semantic_plan.get("variables", {})
    if not isinstance(plan_variables, dict):
        raise ValueError("approved semantic plan has no logical-variable map")
    plan_external = {item.get("variable_id"): item for item in semantic_plan.get("external_boundary_variables", [])
                     if isinstance(item, dict)}
    boundaries: list[dict[str, Any]] = []
    cuts = {item.get("variable_id"): item for item in candidate_trace.get("external_formula_cuts", [])
            if isinstance(item, dict)}
    external_ids = set(catalog_variables)
    if not external_ids or set(plan_external) != external_ids or set(cuts) != external_ids:
        raise ValueError("confirmed catalog, approved design, and candidate trace must declare the same external inputs")

    axis_ids: set[str] = set()
    for item in catalog_variables.values():
        shape = item.get("shape")
        axes = item.get("axes")
        if (not isinstance(shape, list) or len(shape) != 1 or not isinstance(shape[0], int)
                or isinstance(shape[0], bool) or shape[0] < 1 or not isinstance(axes, list)):
            raise ValueError("external catalog inputs must be finite one-dimensional vectors")
        key_axes = [axis.get("axis_id") for axis in axes
                    if isinstance(axis, dict) and axis.get("role") == "key"]
        if len(key_axes) != 1 or not isinstance(key_axes[0], str):
            raise ValueError("each external vector must declare one source-key axis")
        axis_ids.add(key_axes[0])
    if len(axis_ids) != 1:
        raise ValueError("external vectors must share one declared key axis")
    axis_id = next(iter(axis_ids))
    age_axes = [item for item in catalog.get("axis_metadata", [])
                if isinstance(item, dict) and item.get("axis_id") == axis_id]
    if len(age_axes) != 1:
        raise ValueError("confirmed catalog must declare the shared external key axis")
    age_extents = age_axes[0].get("source_extents", [])
    if len(age_extents) != 1:
        raise ValueError("external key axis must have one finite source extent")
    age_axis_range = _canonical_range(
        f"{age_extents[0].get('sheet')}!{age_extents[0].get('range')}")
    axis_shape = age_axes[0].get("shape")
    axis_reference = _parse_reference(age_axis_range, "")
    if (axis_reference is None or not axis_reference.sheet or axis_reference.min_col != axis_reference.max_col
            or not isinstance(axis_shape, list) or axis_shape != [axis_reference.max_row - axis_reference.min_row + 1]):
        raise ValueError("external key axis must be a declared finite vertical vector")

    for variable_id in sorted(external_ids):
        catalog_item = catalog_variables[variable_id]
        variable = plan_variables.get(variable_id)
        boundary = plan_external[variable_id]
        if not isinstance(variable, dict) or variable.get("role") != "external":
            raise ValueError(f"external variable is missing from the approved semantic plan: {variable_id}")
        shape = catalog_item["shape"]
        if variable.get("shape") != shape or boundary.get("shape") != shape:
            raise ValueError(f"external variable shapes differ across the approved artifacts: {variable_id}")
        extents = catalog_item.get("source_extents", [])
        if len(extents) != 1:
            raise ValueError(f"external catalog variable must have one finite source extent: {variable_id}")
        extent_record = extents[0]
        extent = _canonical_range(f"{extent_record.get('sheet')}!{extent_record.get('range')}")
        semantic_extents = [_canonical_range(item) for item in boundary.get("source_extents", [])]
        variable_extents = [_canonical_range(item.get("ref", "")) for item in variable.get("source_extents", [])]
        cut = cuts[variable_id]
        normalized_extent = extent.casefold()
        if ([item.casefold() for item in semantic_extents] != [normalized_extent]
                or [item.casefold() for item in variable_extents] != [normalized_extent]
                or _canonical_range(cut.get("source_range", "")).casefold() != normalized_extent
                or cut.get("traversal") != "stopped_before_prerequisites"):
            raise ValueError(f"catalog, plan, and candidate cut disagree on external extent: {variable_id}")
        extent_ref = _parse_reference(extent, "")
        members = cut.get("formula_members", [])
        expected_addresses = set(_addresses(extent))
        if (extent_ref is None or extent_ref.min_col != extent_ref.max_col
                or extent_ref.max_row - extent_ref.min_row + 1 != shape[0]
                or cut.get("coordinate_count") != shape[0] or len(members) != shape[0]
                or {_canonical_cell(member.get("address", "")) for member in members} != expected_addresses):
            raise ValueError(f"external vector source extent and formula cut have different members: {variable_id}")
        axes = variable.get("axes", [])
        if not any(isinstance(axis, dict) and axis.get("source") == age_axis_range
                   for axis in axes):
            raise ValueError(f"external variable does not use the confirmed key axis: {variable_id}")
        boundaries.append({"boundary_variable_id": variable_id,
                           "logical_name": catalog_item.get("logical_name"),
                           "source_range": extent, "shape": shape})

    return {
        "root": root, "manifest": manifest, "stage3_revision": revision, "stage3": stage3,
        "source_path": source_path, "source_sha256": source_hash, "scenario": scenario,
        "catalog": catalog, "catalog_path": catalog_path, "catalog_md_path": catalog_md_path,
        "catalog_artifact": catalog_artifact, "boundary_reference": boundary_reference,
        "semantic_plan": semantic_plan, "semantic_plan_path": semantic_path,
        "semantic_plan_artifact": semantic_record, "semantic_check_path": check_path,
        "semantic_check_artifact": check_record, "candidate_trace": candidate_trace,
        "candidate_profile": candidate_profile, "evidence_paths": evidence_paths,
        "boundaries": boundaries, "external_ids": sorted(external_ids),
        "age_axis_id": axis_id, "age_axis_range": age_axis_range,
    }


def _read_source_cut_facts(context: dict[str, Any]) -> tuple[dict[str, dict[str, Any]], list[Any], str, str]:
    source_path = context["source_path"]
    source_hash = context["source_sha256"]
    source_before = hash_file(source_path)
    workbook = load_workbook(source_path, data_only=False, read_only=False, keep_vba=True)
    try:
        formula_records: dict[str, dict[str, Any]] = {}
        expected_cut_by_id = {
            item["variable_id"]: item for item in context["candidate_trace"].get("external_formula_cuts", [])
        }
        for boundary in context["boundaries"]:
            variable_id = boundary["boundary_variable_id"]
            cut = expected_cut_by_id[variable_id]
            members = cut.get("formula_members", [])
            extent_addresses = set(_addresses(boundary["source_range"]))
            expected_shape = boundary["shape"]
            if (cut.get("coordinate_count") != expected_shape[0] or len(members) != expected_shape[0]
                    or {_canonical_cell(item.get("address", "")) for item in members} != extent_addresses):
                raise ValueError(f"candidate cut member set does not cover the approved extent: {variable_id}")
            for member in members:
                address = _canonical_cell(member.get("address", ""))
                sheet, local = address.split("!", 1)
                cell = workbook[sheet][local]
                formula = cell.value if cell.data_type == "f" else None
                formula_hash = hashlib.sha256(formula.encode("utf-8")).hexdigest() if isinstance(formula, str) else None
                if (formula != member.get("formula") or formula_hash != member.get("formula_sha256")
                        or member.get("source_workbook_sha256") != source_hash):
                    raise ValueError(f"confirmed cut formula differs from current formula text: {address}")
                formula_records[address] = {"address": address, "boundary_variable_id": variable_id,
                                           "formula": formula, "formula_sha256": formula_hash}

        age_values: list[Any] = []
        for address in _addresses(context["age_axis_range"]):
            sheet, local = address.split("!", 1)
            cell = workbook[sheet][local]
            value = cell.value
            if (cell.data_type == "f" or isinstance(value, bool) or not isinstance(value, (int, float))
                    or not math.isfinite(float(value))):
                raise ValueError(f"approved age key is not a finite literal source value: {address}")
            age_values.append(value)
    finally:
        workbook.close()
    source_after = hash_file(source_path)
    if source_before != source_hash or source_after != source_hash:
        raise ValueError("source workbook changed during formula-mode cut and age-key inspection")
    formula_list = [formula_records[key] for key in sorted(formula_records)]
    expected_formula_count = sum(item["shape"][0] for item in context["boundaries"])
    if len(formula_list) != expected_formula_count:
        raise ValueError("approved external formula set does not match the declared vector shapes")
    return formula_records, age_values, _json_hash(formula_list), _json_hash(age_values)


def _validated_oracle_values(oracle: dict[str, Any], boundaries: list[dict[str, Any]],
                             formula_records: dict[str, dict[str, Any]]) -> dict[str, list[Any]]:
    if (oracle.get("schema_version") != "gp.excel_oracle.v1" or oracle.get("status") != "pass"
            or oracle.get("engine") != "Microsoft Excel" or oracle.get("macros_executed") is not False
            or oracle.get("input_overrides") != []):
        raise ValueError("external values require a successful macro-disabled native FullRebuild capture")
    ranges = oracle.get("ranges")
    if not isinstance(ranges, dict):
        raise ValueError("native capture has no range values")
    expected_ranges = {item["boundary_variable_id"]: item["source_range"] for item in boundaries}
    if set(oracle.get("requested_ranges", [])) != set(expected_ranges.values()):
        raise ValueError("native capture requested ranges differ from the confirmed external boundaries")
    result: dict[str, list[Any]] = {}
    for variable_id, selector in expected_ranges.items():
        record = ranges.get(selector)
        if not isinstance(record, dict):
            raise ValueError(f"native capture is missing the approved range: {selector}")
        values, errors, formulas = record.get("values"), record.get("excel_errors"), record.get("formulas")
        length = next(item["shape"][0] for item in boundaries if item["boundary_variable_id"] == variable_id)
        if not isinstance(values, list) or len(values) != length or any(not isinstance(row, list) or len(row) != 1 for row in values):
            raise ValueError(f"native external vector has the wrong shape: {selector}")
        if not isinstance(errors, list) or len(errors) != length or any(not isinstance(row, list) or len(row) != 1 for row in errors):
            raise ValueError(f"native external vector has no complete error ledger: {selector}")
        if not isinstance(formulas, list) or len(formulas) != length or any(not isinstance(row, list) or len(row) != 1 for row in formulas):
            raise ValueError(f"native external vector has no complete formula ledger: {selector}")
        parsed_values: list[Any] = []
        addresses = _addresses(selector)
        for index, address in enumerate(addresses):
            value, error, formula = values[index][0], errors[index][0], formulas[index][0]
            source_formula = formula_records[address]["formula"]
            if formula != source_formula:
                raise ValueError(f"native formula text differs from source formula at {address}")
            if error is not None:
                if not isinstance(error, str) or not error.startswith("#"):
                    raise ValueError(f"native error value is malformed at {address}")
                parsed_values.append({"excel_error": error.upper()})
            elif (isinstance(value, bool) or not isinstance(value, (int, float))
                  or not math.isfinite(float(value))):
                raise ValueError(f"native external output is not finite numeric or an explicit Excel error at {address}")
            else:
                parsed_values.append(value)
        result[variable_id] = parsed_values
    return result


def capture_external_inputs(workflow_dir: Path) -> dict[str, Any]:
    """Capture only the formula-derived vectors fixed by the approved catalog."""
    attempt_dir: Path | None = None
    created_attempt = False
    try:
        context = _step3_context(workflow_dir)
        root = context["root"]
        source_path = context["source_path"]
        source_hash = context["source_sha256"]
        formula_records, age_values, formula_set_hash, age_key_hash = _read_source_cut_facts(context)
        ranges = [item["source_range"] for item in sorted(context["boundaries"], key=lambda value: value["boundary_variable_id"])]
        stage4_root = (root / "stage4").resolve()
        if stage4_root.is_relative_to(source_path.parent) or source_path.is_relative_to(stage4_root):
            raise ValueError("Step 4 external capture must stay outside the source workbook directory")
        capture_root = stage4_root / "external_inputs"
        revisions = [path for path in capture_root.glob("revision-*") if path.is_dir()] if capture_root.exists() else []
        revision_number = max((int(path.name.removeprefix("revision-")) for path in revisions
                               if path.name.removeprefix("revision-").isdigit()), default=0) + 1
        attempt_dir = capture_root / f"revision-{revision_number:04d}"
        if attempt_dir.exists():
            raise ValueError("Step 4 external capture revision already exists; preserve it and retry in a clean workflow")
        attempt_dir.mkdir(parents=True)
        created_attempt = True
        native_dir = attempt_dir / "native_capture"
        native_result = capture_excel_oracle(source_path, native_dir, targets=[], ranges=ranges)
        if native_result.get("status") != "pass":
            raise ValueError(native_result.get("reason", "native Excel external capture failed"))
        oracle_path = native_dir / "excel_oracle.json"
        if not oracle_path.is_file():
            raise ValueError("native capture did not write excel_oracle.json")
        oracle = read_json(oracle_path)
        if oracle.get("source_sha256") != source_hash or oracle.get("source_copy_sha256") != source_hash:
            raise ValueError("native capture source/private-copy hashes differ from the approved source")
        values = _validated_oracle_values(oracle, context["boundaries"], formula_records)
        if hash_file(source_path) != source_hash:
            raise ValueError("source workbook changed during native external capture")

        evidence = context["evidence_paths"]
        stage3_revision = context["stage3_revision"]
        catalog_artifact = context["catalog_artifact"]
        receipts = context["boundary_reference"]["review_receipts"]
        capture_script = Path(str(files("excel_to_act.steps.step5").joinpath("native_excel_oracle.ps1"))).resolve()
        if not capture_script.is_file():
            raise ValueError("native Excel capture script is missing")
        payload = {
            "schema_version": "step4.external_inputs.v1", "tool": "step4.capture_external",
            "status": "pass", "workflow": str(root), "source_path": str(source_path),
            "source_sha256": source_hash,
            "stage3": {"revision": stage3_revision["revision"],
                       "artifact_json_sha256": stage3_revision["artifact"]["json_sha256"],
                       "artifact_md_sha256": stage3_revision["artifact"]["md_sha256"]},
            "scenario": {"scenario_id": context["scenario"]["scenario_id"],
                         "scenario_sha256": context["scenario"]["scenario_sha256"],
                         "input_overrides": []},
            "catalog": {"revision": context["boundary_reference"]["revision"],
                        "boundary_sha256": context["boundary_reference"]["boundary_sha256"],
                        "artifact_json_sha256": catalog_artifact["json_sha256"],
                        "artifact_md_sha256": catalog_artifact["md_sha256"],
                        "review_receipts": {reviewer: {"receipt_id": receipts[reviewer]["receipt_id"],
                                                        "decision": receipts[reviewer]["decision"],
                                                        "artifact_sha256": receipts[reviewer]["artifact_sha256"],
                                                        "artifact_md_sha256": receipts[reviewer]["artifact_md_sha256"]}
                                            for reviewer in ("agent", "human")}},
            "semantic_plan_sha256": context["semantic_plan_artifact"]["sha256"],
            "semantic_check_sha256": context["semantic_check_artifact"]["sha256"],
            "model_scope": {"sha256": evidence["model_scope"][1], "path": str(evidence["model_scope"][0])},
            "candidate_trace": {"sha256": evidence["candidate_trace"][1], "path": str(evidence["candidate_trace"][0])},
            "candidate_profile": {"sha256": evidence["candidate_profile"][1], "path": str(evidence["candidate_profile"][0])},
            "boundaries": context["boundaries"],
            "external_values": values, "external_values_sha256": _json_hash(values),
            "external_coordinate_count": sum(len(value) for value in values.values()),
            "formula_cut_members": [formula_records[key] for key in sorted(formula_records)],
            "formula_cut_member_count": len(formula_records), "formula_cut_formula_set_sha256": formula_set_hash,
            "age_axis": {"axis_id": context["age_axis_id"],
                         "source_range": context["age_axis_range"], "role": "source_adapter_metadata",
                         "values": age_values, "value_sha256": age_key_hash,
                         "business_input_count_increment": 0},
            "native_capture": {"path": str(oracle_path.resolve()), "sha256": hash_file(oracle_path),
                               "engine": oracle.get("engine"), "engine_settings": oracle.get("engine_settings"),
                               "full_rebuild_seconds": oracle.get("full_rebuild_seconds"),
                               "source_copy_sha256": oracle.get("source_copy_sha256"),
                               "capture_script_path": str(capture_script),
                               "capture_script_sha256": hash_file(capture_script),
                               "macros_executed": False, "input_overrides": [],
                               "formula_cache_inputs": False, "upstream_python_translation": False},
        }
        payload["external_capture_sha256"] = _json_hash(payload)
        artifact_path = attempt_dir / "external_inputs.json"
        artifact_path.write_bytes((json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n").encode("utf-8"))
        markdown = ["# Step 4 External Input Capture", "", "Status: **Pass**", "",
                    f"Source SHA-256: `{source_hash}`",
                    f"Step 3 revision: {stage3_revision['revision']} / `{stage3_revision['artifact']['json_sha256']}`",
                    f"Catalog revision: {context['boundary_reference']['revision']} / `{context['boundary_reference']['boundary_sha256']}`",
                    f"Native Excel: {oracle.get('engine')} FullRebuild; source copy `{oracle.get('source_copy_sha256')}`.",
                    "Macros were disabled; no input overrides, formula caches, or Python upstream producer were used.", "",
                    "| External variable | Source range | Values |", "|---|---|---:|"]
        markdown.extend(f"| `{item['boundary_variable_id']}` | `{item['source_range']}` | {len(values[item['boundary_variable_id']])} |"
                        for item in context["boundaries"])
        markdown.extend(["", f"Formula cuts bound: {len(formula_records)}; formula-set SHA-256 `{formula_set_hash}`.",
                    f"Shared key-axis metadata values: {len(age_values)}; SHA-256 `{age_key_hash}`.",
                         "This capture does not discover active dependencies or approve Stage 4.", ""])
        (attempt_dir / "external_inputs.md").write_text("\n".join(markdown), encoding="utf-8")
        artifact_hash = hash_file(artifact_path)
        return {"schema_version": payload["schema_version"], "tool": payload["tool"], "status": "pass",
                "workflow": str(root), "artifact": str(artifact_path), "artifact_sha256": artifact_hash,
                "markdown": str(attempt_dir / "external_inputs.md"),
                "external_capture_sha256": payload["external_capture_sha256"],
                "external_coordinate_count": payload["external_coordinate_count"],
                "formula_cut_member_count": len(formula_records), "stage_advanced": False}
    except Exception as exc:
        if created_attempt and attempt_dir is not None and attempt_dir.exists():
            shutil.rmtree(attempt_dir, ignore_errors=True)
        return {"schema_version": "step4.external_inputs.v1", "tool": "step4.capture_external",
                "status": "blocked", "reason": str(exc), "stage_advanced": False}


def load_current_external_bundle(workflow_dir: Path) -> tuple[dict[str, Any], dict[str, Any], Path, str]:
    """Return the current approved source/design context and its newest external capture."""
    context = _step3_context(workflow_dir)
    root = context["root"]
    capture_root = root / "stage4" / "external_inputs"
    paths = sorted(capture_root.glob("revision-*/external_inputs.json"), reverse=True) if capture_root.exists() else []
    if not paths:
        raise ValueError("no Step 4 external capture exists; run `step4 capture-external --workflow DIR` first")
    path = paths[0]
    capture = read_json(path)
    expected = {
        "source_sha256": context["source_sha256"],
        "stage3": {"revision": context["stage3_revision"]["revision"],
                   "artifact_json_sha256": context["stage3_revision"]["artifact"]["json_sha256"],
                   "artifact_md_sha256": context["stage3_revision"]["artifact"]["md_sha256"]},
        "scenario": {"scenario_id": context["scenario"]["scenario_id"],
                     "scenario_sha256": context["scenario"]["scenario_sha256"], "input_overrides": []},
        "catalog": {"revision": context["boundary_reference"]["revision"],
                    "boundary_sha256": context["boundary_reference"]["boundary_sha256"],
                    "artifact_json_sha256": context["catalog_artifact"]["json_sha256"],
                    "artifact_md_sha256": context["catalog_artifact"]["md_sha256"],
                    "review_receipts": {reviewer: {"receipt_id": context["boundary_reference"]["review_receipts"][reviewer]["receipt_id"],
                                                    "decision": context["boundary_reference"]["review_receipts"][reviewer]["decision"],
                                                    "artifact_sha256": context["boundary_reference"]["review_receipts"][reviewer]["artifact_sha256"],
                                                    "artifact_md_sha256": context["boundary_reference"]["review_receipts"][reviewer]["artifact_md_sha256"]}
                                        for reviewer in ("agent", "human")}},
        "semantic_plan_sha256": context["semantic_plan_artifact"]["sha256"],
        "semantic_check_sha256": context["semantic_check_artifact"]["sha256"],
        "external_values_sha256": _json_hash(capture.get("external_values")),
    }
    if (capture.get("schema_version") != "step4.external_inputs.v1" or capture.get("status") != "pass"
            or any(capture.get(key) != value for key, value in expected.items())
            or capture.get("model_scope") != {"sha256": context["evidence_paths"]["model_scope"][1],
                                               "path": str(context["evidence_paths"]["model_scope"][0])}
            or capture.get("candidate_trace") != {"sha256": context["evidence_paths"]["candidate_trace"][1],
                                                   "path": str(context["evidence_paths"]["candidate_trace"][0])}
            or capture.get("candidate_profile") != {"sha256": context["evidence_paths"]["candidate_profile"][1],
                                                     "path": str(context["evidence_paths"]["candidate_profile"][0])}):
        raise ValueError("latest Step 4 external capture is stale or bound to different evidence")
    oracle_path = Path(capture.get("native_capture", {}).get("path", "")).resolve()
    if (not oracle_path.is_file() or hash_file(oracle_path) != capture.get("native_capture", {}).get("sha256")
            or capture.get("native_capture", {}).get("source_copy_sha256") != context["source_sha256"]
            or capture.get("native_capture", {}).get("macros_executed") is not False
            or capture.get("native_capture", {}).get("input_overrides") != []):
        raise ValueError("native external capture file or source identity is stale")
    formula_records, age_values, formula_hash, age_hash = _read_source_cut_facts(context)
    if (capture.get("formula_cut_members") != [formula_records[key] for key in sorted(formula_records)]
            or capture.get("formula_cut_formula_set_sha256") != formula_hash
            or capture.get("age_axis", {}).get("source_range") != context["age_axis_range"]
            or capture.get("age_axis", {}).get("values") != age_values
            or capture.get("age_axis", {}).get("value_sha256") != age_hash):
        raise ValueError("current source formula cuts or age keys differ from the captured artifact")
    artifact_hash = hash_file(path)
    if capture.get("external_capture_sha256") != _capture_hash(capture):
        raise ValueError("external capture payload hash is invalid")
    return context, capture, path, artifact_hash


def load_current_external_inputs(workflow_dir: Path) -> tuple[dict[str, Any], Path, str]:
    """Return the newest current external value artifact and its file hash."""
    _context, capture, path, digest = load_current_external_bundle(workflow_dir)
    return capture, path, digest


def _capture_hash(capture: dict[str, Any]) -> str:
    return _json_hash({key: value for key, value in capture.items() if key != "external_capture_sha256"})
