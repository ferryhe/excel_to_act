"""Compile one mapped source formula into a direct logical-variable function."""

from __future__ import annotations

from bisect import bisect_left, bisect_right
import re
from dataclasses import dataclass
from typing import Any, Callable

from openpyxl.utils.cell import (
    column_index_from_string,
    get_column_letter,
    quote_sheetname,
    range_boundaries,
)

from excel_to_act.steps.step3.calculation import CalculationBlocked, CellRange, ExcelError, _Parser, _parse_reference


_CALLS = {
    "IF": ("excel_if", 2, 3),
    "IFERROR": ("excel_iferror", 2, 2),
    "IFNA": ("excel_ifna", 2, 2),
    "AND": ("excel_and", 1, 255),
    "OR": ("excel_or", 1, 255),
    "MIN": ("minimum", 1, 255),
    "MAX": ("maximum", 1, 255),
    "SUM": ("excel_sum", 1, 255),
    "SUMPRODUCT": ("sumproduct", 1, 255),
    "MATCH": ("match", 2, 3),
    "VLOOKUP": ("vlookup", 3, 4),
    "TRANSPOSE": ("transpose", 1, 1),
    "OFFSET": ("offset", 3, 5),
    "INDIRECT": ("indirect", 1, 1),
}
_FUNCTION_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


@dataclass(frozen=True)
class SourceCell:
    variable_id: str
    indices: tuple[int, ...]
    axes: tuple[str | None, ...]
    role: str
    source_address: str | None = None
    metadata_group: str | None = None


def _origin(value: Any) -> int:
    if isinstance(value, int) and not isinstance(value, bool):
        return value
    if isinstance(value, str) and value.isalpha():
        return column_index_from_string(value.upper())
    raise CalculationBlocked(f"semantic coordinate map has an invalid origin: {value!r}")


def _indices_for_coordinate(
    mapping: dict[str, Any], row: int, col: int, shape: list[int],
    extent: CellRange,
) -> tuple[tuple[int, ...], tuple[str | None, ...]]:
    if not shape:
        return (), ()
    indices: list[int | None] = [None] * len(shape)
    axes: list[str | None] = [None] * len(shape)
    array_ref = mapping.get("array_ref")
    if isinstance(array_ref, str):
        bounds = range_boundaries(array_ref.replace("$", ""))
        if any(value is None for value in bounds):
            raise CalculationBlocked(f"semantic array index map has incomplete bounds: {array_ref}")
        array_min_col, array_min_row, _array_max_col, _array_max_row = bounds
    else:
        array_min_col, array_min_row = extent.min_col, extent.min_row
    for axis in mapping.get("axes", []):
        dimension = axis.get("dimension")
        coordinate = axis.get("source_coordinate")
        origin = axis.get("origin", 0)
        offset = axis.get("index_origin", 0)
        if (not isinstance(dimension, int) or dimension < 0 or dimension >= len(shape)
                or not isinstance(offset, int)):
            raise CalculationBlocked("semantic coordinate map has an invalid axis dimension")
        if coordinate == "row":
            value = row - _origin(origin) + offset
        elif coordinate == "column":
            value = col - _origin(origin) + offset
        elif coordinate == "array_row":
            value = row - array_min_row + offset
        elif coordinate == "array_column":
            value = col - array_min_col + offset
        else:
            raise CalculationBlocked(f"semantic coordinate map uses an unsupported source coordinate: {coordinate}")
        indices[dimension] = value
        axes[dimension] = axis.get("axis")
    for fixed in mapping.get("fixed_axes", []):
        dimension, value = fixed.get("dimension"), fixed.get("index")
        if (not isinstance(dimension, int) or dimension < 0 or dimension >= len(shape)
                or not isinstance(value, int)):
            raise CalculationBlocked("semantic fixed-axis map has an invalid dimension or index")
        indices[dimension] = value
        if fixed.get("axis"):
            axes[dimension] = fixed["axis"]
    if mapping.get("kind") == "scalar" and isinstance(mapping.get("index"), int):
        indices[0] = mapping["index"]
    if any(value is None for value in indices):
        raise CalculationBlocked("semantic coordinate map does not define every variable axis")
    result = tuple(int(value) for value in indices)
    if any(value < 0 or value >= shape[dimension] for dimension, value in enumerate(result)):
        raise CalculationBlocked(f"semantic coordinate map points outside declared shape {shape}: {result}")
    return result, tuple(axes)


def build_source_cell_index(semantic_plan: dict[str, Any]) -> dict[tuple[str, int, int], list[SourceCell]]:
    """Index declared source extents as logical variable coordinates for compilation."""
    result: dict[tuple[str, int, int], list[SourceCell]] = {}
    formula_members: dict[str, set[tuple[str, int, int]]] = {}
    formula_family_members: dict[tuple[str, str], set[tuple[str, int, int]]] = {}

    def add_formula_member(variable_id: str, family_id: str | None, address: Any,
                           *, array: bool = False) -> None:
        if not isinstance(address, str):
            kind = "array formula" if array else "formula"
            raise CalculationBlocked(f"{kind} member address is not text")
        reference = _parse_reference(address, "Main")
        if reference is None or not reference.single:
            kind = "array formula" if array else "formula"
            raise CalculationBlocked(f"{kind} member is not a qualified cell: {address}")
        key = (reference.sheet.casefold(), reference.min_row, reference.min_col)
        formula_members.setdefault(variable_id, set()).add(key)
        if isinstance(family_id, str):
            formula_family_members.setdefault((variable_id, family_id), set()).add(key)

    for mapping in semantic_plan.get("source_mappings", []):
        variable_id = mapping.get("variable_id")
        if not isinstance(variable_id, str):
            continue
        for address in mapping.get("source_members", []):
            add_formula_member(variable_id, mapping.get("source_family_id"), address)
    for array_group in semantic_plan.get("array_mappings", []):
        for instance in array_group.get("instances", []):
            variable_id = instance.get("variable_id")
            if not isinstance(variable_id, str):
                continue
            family_id = instance.get("anchor_source_family_id")
            addresses = [instance.get("anchor"), *instance.get("source_members", [])]
            for address in addresses:
                add_formula_member(variable_id, family_id, address, array=True)
    variables = semantic_plan.get("variables")
    if isinstance(variables, dict):
        variable_records = [(key, value) for key, value in variables.items()]
    elif isinstance(variables, list):
        variable_records = [(item.get("variable_id"), item) for item in variables]
    else:
        raise CalculationBlocked("semantic plan has no logical variable list")
    for variable_id, variable in variable_records:
        shape = variable.get("shape")
        if not isinstance(variable_id, str) or not isinstance(shape, list):
            raise CalculationBlocked("semantic plan has an invalid logical variable")
        variable_role = variable.get("role")
        for source_extent in variable.get("source_extents", []):
            reference = _parse_reference(source_extent.get("ref", ""), "Main")
            mapping = source_extent.get("coordinate_mapping")
            if reference is None or not isinstance(mapping, dict):
                raise CalculationBlocked(f"variable extent has an invalid coordinate map: {variable_id}")
            # Some extents are labels or descriptive evidence, not addressable values
            # (for example a header-only axis on a matrix). They intentionally map only
            # one logical axis and must not be used to invent coordinates on the others.
            mapped_dimensions = {item.get("dimension") for item in mapping.get("axes", [])}
            mapped_dimensions.update(item.get("dimension") for item in mapping.get("fixed_axes", []))
            if len(mapped_dimensions) < len(shape):
                continue
            source_family_id = mapping.get("source_family_id")
            allowed_formula_members = (
                formula_family_members.get((variable_id, source_family_id), set())
                if isinstance(source_family_id, str) else formula_members.get(variable_id, set()))
            for row in range(reference.min_row, reference.max_row + 1):
                for col in range(reference.min_col, reference.max_col + 1):
                    if (variable_role != "external" and source_extent.get("role") == "formula_output"
                            and (reference.sheet.casefold(), row, col)
                            not in allowed_formula_members):
                        # Source extents can describe a sparse family bounding box.
                        # Only exact mapped members own formula-output coordinates;
                        # gaps may be raw literals or unselected source cells.
                        continue
                    indices, axes = _indices_for_coordinate(mapping, row, col, shape, reference)
                    key = (reference.sheet.casefold(), row, col)
                    cell_role = ("external_boundary_input" if variable_role == "external"
                                 else str(source_extent.get("role", "")))
                    cell = SourceCell(variable_id, indices, axes, cell_role)
                    bucket = result.setdefault(key, [])
                    if cell not in bucket:
                        bucket.append(cell)
        initial = variable.get("initial_condition")
        seed_address = initial.get("seed_ref") if isinstance(initial, dict) else None
        if isinstance(seed_address, str):
            seed_ref = _parse_reference(seed_address, "Main")
            if seed_ref is None:
                raise CalculationBlocked(f"logical variable has an invalid seed reference: {variable_id}")
            logical_index = initial.get("logical_index", 0)
            if not isinstance(logical_index, int):
                raise CalculationBlocked(f"logical variable has an invalid seed index: {variable_id}")
            axis_names = tuple(axis.get("name") for axis in variable.get("axes", []))
            if not shape and seed_ref.single:
                indices_by_cell = {(seed_ref.min_row, seed_ref.min_col): ((), ())}
            elif len(shape) == 1 and seed_ref.single:
                indices_by_cell = {(seed_ref.min_row, seed_ref.min_col): ((logical_index,), (axis_names[0],))}
            elif (len(shape) == 2 and seed_ref.min_row == seed_ref.max_row
                  and seed_ref.max_col - seed_ref.min_col + 1 == shape[1]):
                indices_by_cell = {
                    (seed_ref.min_row, col): ((logical_index, col - seed_ref.min_col), axis_names[:2])
                    for col in range(seed_ref.min_col, seed_ref.max_col + 1)
                }
            else:
                raise CalculationBlocked(f"logical seed extent has an unsupported shape: {variable_id}")
            for (row, col), (indices, axes) in indices_by_cell.items():
                if any(index < 0 or index >= shape[dimension] for dimension, index in enumerate(indices)):
                    raise CalculationBlocked(f"logical seed index is outside declared shape: {variable_id}")
                key = (seed_ref.sheet.casefold(), row, col)
                cell = SourceCell(variable_id, indices, axes, "initial_condition")
                bucket = result.setdefault(key, [])
                if cell not in bucket:
                    bucket.append(cell)
    for item in semantic_plan.get("adapter_metadata_cells", []):
        if not isinstance(item, dict):
            raise CalculationBlocked("adapter metadata coordinate is malformed")
        address, group_id = item.get("address"), item.get("group_id")
        indices, axes = item.get("indices"), item.get("axes", [])
        reference = _parse_reference(address or "", "Main")
        if (reference is None or not reference.single or not isinstance(group_id, str)
                or not isinstance(indices, list) or not indices
                or any(not isinstance(index, int) or index < 0 for index in indices)
                or not isinstance(axes, list) or len(axes) != len(indices)
                or any(axis is not None and not isinstance(axis, str) for axis in axes)):
            raise CalculationBlocked("adapter metadata coordinate has an invalid group or logical index")
        key = (reference.sheet.casefold(), reference.min_row, reference.min_col)
        if result.get(key):
            # A catalog coordinate already owned by a model variable must keep
            # that logical binding; it is not duplicated as adapter metadata.
            continue
        cell = SourceCell(f"__adapter_metadata__.{group_id}", tuple(indices), tuple(axes),
                          "adapter_metadata", metadata_group=group_id)
        result[key] = [cell]
    return result


def _index_expression(cell: SourceCell, owner_axes: tuple[str | None, ...],
                      owner_indices: tuple[int, ...],
                      explicit_alignment: dict[int, int] | None = None) -> str:
    owner_dimensions = {name: dimension for dimension, name in enumerate(owner_axes) if name is not None}
    result: list[str] = []
    for dimension, index in enumerate(cell.indices):
        axis = cell.axes[dimension] if dimension < len(cell.axes) else None
        owner_dimension = owner_dimensions.get(axis)
        if owner_dimension is None and explicit_alignment is not None:
            owner_dimension = explicit_alignment.get(dimension)
        if owner_dimension is None:
            result.append(str(index))
            continue
        delta = index - owner_indices[owner_dimension]
        result.append(f"i{owner_dimension}" if delta == 0 else f"(i{owner_dimension} {'+' if delta > 0 else '-'} {abs(delta)})")
    return "(" + ", ".join(result) + ("," if len(result) == 1 else "") + ")"


def _resolved_source_cells(
    reference: CellRange,
    owner_variable_id: str,
    semantic_plan: dict[str, Any],
    source_index: dict[tuple[str, int, int], list[SourceCell]],
) -> list[list[SourceCell]]:
    declared_variables = semantic_plan.get("variables", [])
    variables = declared_variables if isinstance(declared_variables, dict) else {
        item["variable_id"]: item for item in declared_variables
    }
    owner = variables.get(owner_variable_id)
    if owner is None:
        raise CalculationBlocked(f"formula output variable is not declared: {owner_variable_id}")
    dependencies = {item.get("variable_id") for item in owner.get("dependencies", [])}
    dependencies.add(owner_variable_id)
    source_literals = {
        _address_key(address) for address in semantic_plan.get("adapter_source_literals", [])
        if isinstance(address, str)
    }
    rows: list[list[SourceCell]] = []
    for row in range(reference.min_row, reference.max_row + 1):
        cells: list[SourceCell] = []
        for col in range(reference.min_col, reference.max_col + 1):
            declared = source_index.get((reference.sheet.casefold(), row, col), [])
            # Array anchors and some bridge formulas have complete coordinate
            # declarations but no expanded dependency list in the Stage 3 plan.
            candidates = [item for item in declared if item.variable_id in dependencies] or declared
            # A source literal may initialize a derived series at the same cell.
            # When a formula refers to that boundary, the declared seed owner is
            # the derived variable's initial_condition, while the raw variable
            # remains separately bound in the input adapter.
            seed_owners = [item for item in candidates if item.role == "initial_condition"]
            if seed_owners:
                candidates = seed_owners
            owner_axes = {axis.get("name") for axis in owner.get("axes", []) if axis.get("name")}
            axis_matches = [item for item in candidates if owner_axes.intersection(item.axes)]
            if axis_matches:
                candidates = axis_matches
            unique = {(item.variable_id, item.indices, item.axes): item for item in candidates}
            if len(unique) != 1:
                if not unique:
                    address = f"{reference.sheet}!{get_column_letter(col)}{row}"
                    key = _address_key(address)
                    if key in source_literals:
                        cells.append(SourceCell("__source_values__", (), (), "source_literal", key))
                        continue
                address = f"{reference.sheet}!{get_column_letter(col)}{row}"
                reason = "not mapped" if not unique else "ambiguous"
                raise CalculationBlocked(f"active source reference is {reason} for {owner_variable_id}: {address}")
            cells.append(next(iter(unique.values())))
        rows.append(cells)
    return rows


def _reference_nodes(node: Any) -> list[CellRange]:
    kind = node[0]
    if kind == "reference":
        return [node[1]]
    if kind == "range":
        return _reference_nodes(node[1]) + _reference_nodes(node[2])
    if kind == "unary":
        return _reference_nodes(node[2])
    if kind == "binary":
        return _reference_nodes(node[2]) + _reference_nodes(node[3])
    if kind == "call":
        return [reference for argument in node[2] for reference in _reference_nodes(argument)]
    return []


def _unsupported_nodes(node: Any) -> list[tuple[str, str]]:
    """Return external/invalid references and names requiring lazy resolution."""
    kind = node[0]
    if kind == "unsupported_reference":
        return [(kind, str(node[1]))]
    if kind == "name":
        return [(kind, str(node[1]))]
    if kind in {"range", "binary"}:
        children = node[1:] if kind == "range" else node[2:]
    elif kind == "unary":
        children = (node[2],)
    elif kind == "call":
        children = node[2]
    else:
        children = ()
    return [item for child in children for item in _unsupported_nodes(child)]


def _active_trace_reference(trace: dict[str, Any], formula_cell: str, raw_reference: str,
                            current_sheet: str) -> bool:
    """Check whether an unsupported reference has a matching active trace edge."""
    reference = _parse_reference(raw_reference, current_sheet)
    for edge in trace.get("edges", []):
        if (not isinstance(edge, dict) or not isinstance(edge.get("consumer"), str)
                or not isinstance(edge.get("prerequisite"), str)
                or edge["consumer"].casefold() != formula_cell.casefold()):
            continue
        prerequisite = edge["prerequisite"]
        if prerequisite.casefold() == raw_reference.casefold():
            return True
        cell = _parse_reference(prerequisite, current_sheet)
        if (reference is not None and cell is not None and cell.single
                and reference.sheet.casefold() == cell.sheet.casefold()
                and reference.min_row <= cell.min_row <= reference.max_row
                and reference.min_col <= cell.min_col <= reference.max_col):
            return True
    return False


def _active_trace_name(trace: dict[str, Any], name: str) -> bool:
    return any(isinstance(item, dict) and isinstance(item.get("name"), str)
               and item["name"].casefold() == name.casefold()
               for item in trace.get("names", []))


_VALUE_SLOT = re.compile(r"values\[(?P<quote>['\"])(?P<variable>[^'\"]+)\1\]")


def _parameterize_variable_ids(source: str) -> tuple[str, tuple[str, ...]]:
    """Replace concrete variable names with ordered logical input slots."""
    variables: list[str] = []
    slots: dict[str, int] = {}

    def replace(match: re.Match[str]) -> str:
        variable_id = match.group("variable")
        if variable_id not in slots:
            slots[variable_id] = len(variables)
            variables.append(variable_id)
        return f"values[refs[{slots[variable_id]}]]"

    return _VALUE_SLOT.sub(replace, source), tuple(variables)


def _ast_skeleton(node: Any) -> Any:
    kind = node[0]
    if kind in {"reference", "unsupported_reference"}:
        return (kind,)
    if kind == "literal":
        return (kind, repr(node[1]))
    if kind == "name":
        return (kind, str(node[1]).casefold())
    if kind == "range":
        return (kind, _ast_skeleton(node[1]), _ast_skeleton(node[2]))
    if kind == "unary":
        return (kind, node[1], _ast_skeleton(node[2]))
    if kind == "binary":
        return (kind, node[1], _ast_skeleton(node[2]), _ast_skeleton(node[3]))
    if kind == "call":
        return (kind, str(node[1]).casefold(), tuple(_ast_skeleton(arg) for arg in node[2]))
    return (kind,)


def _infer_group_axis_alignments(
    semantic_plan: dict[str, Any],
    formula_by_address: dict[str, str],
    source_index: dict[tuple[str, int, int], list[SourceCell]],
) -> dict[tuple[str, int, int, int, str, int], int]:
    """Infer index coupling only when translated formulas demonstrate a constant offset."""
    mappings = semantic_plan.get("source_mappings", [])
    groups: dict[tuple[str, Any], list[dict[str, Any]]] = {}
    variables = semantic_plan.get("variables", {})
    if isinstance(variables, list):
        variables = {item.get("variable_id"): item for item in variables}
    for mapping in mappings:
        equation_id = mapping.get("equation_family_id")
        variable_id = mapping.get("variable_id")
        for address in mapping.get("source_members", []):
            formula = formula_by_address.get(_address_key(address))
            target = _parse_reference(address, "Main")
            if not isinstance(formula, str) or target is None:
                continue
            try:
                skeleton = _ast_skeleton(_Parser(formula, target.sheet).parse())
            except CalculationBlocked:
                continue
            groups.setdefault((str(equation_id), skeleton), []).append({
                "address": address, "formula": formula, "target": target,
                "variable_id": variable_id,
            })

    observations: dict[tuple[str, int, int, int, str, int], list[tuple[tuple[int, ...], int]]] = {}
    for members in groups.values():
        for member in members:
            target = member["target"]
            variable_id = member["variable_id"]
            output_candidates = [cell for cell in source_index.get(
                (target.sheet.casefold(), target.min_row, target.min_col), [])
                if cell.variable_id == variable_id]
            output_unique = {(cell.indices, cell.axes): cell for cell in output_candidates}
            if len(output_unique) != 1:
                continue
            output = next(iter(output_unique.values()))
            parsed = _Parser(member["formula"], target.sheet).parse()
            for occurrence, reference in enumerate(_reference_nodes(parsed)):
                try:
                    referenced_rows = _resolved_source_cells(reference, variable_id, semantic_plan, source_index)
                except CalculationBlocked:
                    # An unmapped address can be a lazy branch boundary. It has no
                    # effect on index alignment observed from other family members.
                    continue
                for row_offset, row in enumerate(referenced_rows):
                    for col_offset, cell in enumerate(row):
                        for source_dimension, source_index_value in enumerate(cell.indices):
                            marker = (variable_id, occurrence, row_offset, col_offset,
                                      cell.variable_id, source_dimension)
                            observations.setdefault(marker, []).append((output.indices, source_index_value))

    alignments: dict[tuple[str, int, int, int, str, int], int] = {}
    for marker, pairs in observations.items():
        owner_dimension_count = len(pairs[0][0])
        choices = []
        for owner_dimension in range(owner_dimension_count):
            owner_values = [pair[0][owner_dimension] for pair in pairs]
            source_values = [pair[1] for pair in pairs]
            offsets = {source - owner for source, owner in zip(source_values, owner_values)}
            if len(set(owner_values)) > 1 and len(offsets) == 1:
                choices.append(owner_dimension)
        if len(choices) == 1:
            alignments[marker] = choices[0]
    return alignments


def lower_mapped_reference(
    reference: CellRange,
    *,
    owner_variable_id: str,
    owner_indices: tuple[int, ...],
    owner_axes: tuple[str | None, ...],
    semantic_plan: dict[str, Any],
    source_index: dict[tuple[str, int, int], list[SourceCell]],
    reference_occurrence: int | None = None,
    axis_alignments: dict[tuple[str, int, int, int, str, int], int] | None = None,
) -> str:
    """Lower a concrete source reference to declared named variables and logical indices."""
    rows = _resolved_source_cells(reference, owner_variable_id, semantic_plan, source_index)

    variable_ids = {cell.variable_id for row in rows for cell in row}
    def explicit_alignment(row_index: int, col_index: int, cell: SourceCell) -> dict[int, int]:
        if axis_alignments is None or reference_occurrence is None:
            return {}
        return {dimension: owner_dimension
                for (owner_id, occurrence, mapped_row, mapped_col, variable_id, dimension), owner_dimension
                in axis_alignments.items()
                if (owner_id, occurrence, mapped_row, mapped_col, variable_id) ==
                (owner_variable_id, reference_occurrence, row_index, col_index, cell.variable_id)}

    if len(rows) == 1 and len(rows[0]) == 1:
        cell = rows[0][0]
        if cell.role == "source_literal":
            return f"runtime.source_literal(values, {cell.source_address!r})"
        if cell.role == "adapter_metadata":
            return (f"runtime.adapter_metadata(values, {cell.metadata_group!r}, "
                    f"{_index_expression(cell, owner_axes, owner_indices, explicit_alignment(0, 0, cell))})")
        return (f"runtime.logical_reference(values[{cell.variable_id!r}], "
                f"{_index_expression(cell, owner_axes, owner_indices, explicit_alignment(0, 0, cell))}, ())")
    if len(variable_ids) == 1:
        variable_id = next(iter(variable_ids))
        first = rows[0][0]
        same_variable_grid = all(cell.variable_id == variable_id for row in rows for cell in row)
        if same_variable_grid:
            dimensions = len(first.indices)
            starts = tuple(min(cell.indices[dim] for row in rows for cell in row)
                           for dim in range(dimensions))
            ends = tuple(max(cell.indices[dim] for row in rows for cell in row)
                         for dim in range(dimensions))
            logical_size = 1
            for start, end in zip(starts, ends):
                logical_size *= end - start + 1
            if logical_size == len(rows) * len(rows[0]):
                start_expr = _index_expression(first, owner_axes, owner_indices,
                                               explicit_alignment(0, 0, first))
                logical_shape = tuple(end - start + 1 for start, end in zip(starts, ends))
                shape_expr = repr(logical_shape)
                source_shape_expr = repr((len(rows), len(rows[0])))
                if first.role == "adapter_metadata":
                    if any(cell.metadata_group != first.metadata_group or cell.role != "adapter_metadata"
                           for row in rows for cell in row):
                        raise CalculationBlocked("metadata reference crosses adapter groups")
                    return (f"runtime.adapter_metadata_reference(values, {first.metadata_group!r}, "
                            f"{start_expr}, {shape_expr}, source_shape={source_shape_expr})")
                return (f"runtime.logical_reference(values[{variable_id!r}], {start_expr}, "
                        f"{shape_expr}, source_shape={source_shape_expr})")
    matrix = [
        [((f"runtime.source_literal(values, {cell.source_address!r})")
          if cell.role == "source_literal" else
          (f"runtime.adapter_metadata(values, {cell.metadata_group!r}, "
           f"{_index_expression(cell, owner_axes, owner_indices, explicit_alignment(row_index, col_index, cell))})")
          if cell.role == "adapter_metadata" else
          (f"runtime.logical_reference(values[{cell.variable_id!r}], "
           f"{_index_expression(cell, owner_axes, owner_indices, explicit_alignment(row_index, col_index, cell))}, ())"))
         for col_index, cell in enumerate(row)]
        for row_index, row in enumerate(rows)
    ]
    return "runtime.table_view([" + ", ".join("[" + ", ".join(row) + "]" for row in matrix) + "])"


def compile_mapped_families(
    semantic_plan: dict[str, Any],
    trace: dict[str, Any],
    resolve_name: Callable[[str, str, str], str],
) -> dict[str, Any]:
    """Compile each reviewed source family once, deduplicating equal semantic bodies."""
    variables = semantic_plan.get("variables", {})
    if isinstance(variables, list):
        variables = {item.get("variable_id"): item for item in variables}
    equations = semantic_plan.get("equation_families", {})
    if isinstance(equations, list):
        equations = {item.get("family_id"): item for item in equations}
    if not isinstance(variables, dict) or not isinstance(equations, dict):
        raise CalculationBlocked("semantic plan has invalid variable or equation-family declarations")

    source_index = build_source_cell_index(semantic_plan)
    trace_cells = trace.get("cells")
    if not isinstance(trace_cells, list):
        raise CalculationBlocked("active trace has no formula records")
    formula_by_address = {
        _address_key(item.get("address")): item.get("formula")
        for item in trace_cells
        if isinstance(item, dict) and item.get("role") == "calculated_formula"
        and isinstance(item.get("formula"), str)
    }
    axis_alignments = _infer_group_axis_alignments(semantic_plan, formula_by_address, source_index)
    body_to_function: dict[tuple[str, str, str], str] = {}
    functions: dict[str, str] = {}
    mapping_to_function: dict[str, str] = {}
    family_exceptions: list[dict[str, Any]] = []
    family_variants: dict[str, list[str]] = {}
    source_bindings: dict[str, list[str]] = {}
    compiled_members = 0
    unsupported_references: list[dict[str, str]] = []
    unsupported_names: list[dict[str, str]] = []

    for mapping in semantic_plan.get("source_mappings", []):
        family_id = mapping.get("source_family_id")
        equation_id = mapping.get("equation_family_id")
        variable_id = mapping.get("variable_id")
        member_addresses = mapping.get("source_members")
        if (not isinstance(family_id, str) or not isinstance(equation_id, str)
                or not isinstance(variable_id, str) or not isinstance(member_addresses, list)
                or equation_id not in equations or variable_id not in variables):
            raise CalculationBlocked("semantic source-family mapping is incomplete")
        equation = equations[equation_id]
        base_name = equation.get("function_name")
        if not isinstance(base_name, str) or not _FUNCTION_NAME.fullmatch(base_name):
            raise CalculationBlocked(f"semantic equation family has an invalid function name: {equation_id}")
        member_implementations: list[tuple[str, str, str, tuple[str, ...]]] = []
        for address in member_addresses:
            key = _address_key(address)
            formula = formula_by_address.get(key)
            if not isinstance(formula, str):
                raise CalculationBlocked(f"active trace formula is missing for semantic member: {address}")
            parsed = _parse_reference(address, "Main")
            if parsed is None:
                raise CalculationBlocked(f"semantic formula member is not a qualified cell: {address}")
            candidates = [item for item in source_index.get(
                (parsed.sheet.casefold(), parsed.min_row, parsed.min_col), [])
                if item.variable_id == variable_id]
            unique = {(item.indices, item.axes): item for item in candidates}
            if len(unique) != 1:
                raise CalculationBlocked(f"formula output coordinate is not uniquely mapped: {address}")
            output = next(iter(unique.values()))
            index_parameters = tuple(f"i{dimension}" for dimension in range(len(output.indices)))
            reference_counter = 0
            parsed_formula = _Parser(formula, parsed.sheet).parse()
            for node_kind, token in _unsupported_nodes(parsed_formula):
                if node_kind == "unsupported_reference":
                    unsupported_references.append({"source_family_id": family_id,
                                                   "formula_cell": key,
                                                   "reference": token})

            def reference_resolver(reference: CellRange, sheet: str, owner: str) -> str:
                nonlocal reference_counter
                occurrence = reference_counter
                reference_counter += 1
                try:
                    return lower_mapped_reference(
                        reference, owner_variable_id=variable_id, owner_indices=output.indices,
                        owner_axes=output.axes, semantic_plan=semantic_plan, source_index=source_index,
                        reference_occurrence=occurrence, axis_alignments=axis_alignments,
                    )
                except CalculationBlocked as exc:
                    unsupported_references.append({"source_family_id": family_id,
                                                   "formula_cell": key,
                                                   "reference": reference.address,
                                                   "lowering_message": str(exc)})
                    return f"runtime.unsupported_reference({reference.address!r})"

            def family_name_resolver(name: str, sheet: str, owner: str) -> str:
                try:
                    return resolve_name(name, sheet, owner)
                except CalculationBlocked as exc:
                    unsupported_names.append({"source_family_id": family_id,
                                              "formula_cell": key,
                                              "name": name,
                                              "reason": str(exc)})
                    raise

            source = compile_formula_function(
                "compiled_formula", formula, parsed.sheet, address,
                reference_resolver, family_name_resolver, index_parameters=index_parameters,
            )
            header, original_body = source.split("\n", 1)
            original_parameters = header[len("def compiled_formula("):-2].split(", ")[1:]
            body, binding_variables = _parameterize_variable_ids(original_body)
            signature = ", ".join(("values", "refs", *original_parameters))
            member_implementations.append((address, signature, body, binding_variables))
            compiled_members += 1
        if not member_implementations:
            raise CalculationBlocked(f"source family has no formula members: {family_id}")
        counts: dict[tuple[str, str], int] = {}
        for _address, signature, body, _bindings in member_implementations:
            key = (signature, body)
            counts[key] = counts.get(key, 0) + 1
        ordered_implementations = sorted(counts, key=lambda item: (-counts[item], item[0], item[1]))
        primary_signature, primary_body = ordered_implementations[0]
        family_function_names: dict[tuple[str, str], str] = {}
        for signature, body in ordered_implementations:
            body_key = (equation_id, signature, body)
            function_name = body_to_function.get(body_key)
            if function_name is None:
                siblings = sum(key[0] == equation_id for key in body_to_function)
                function_name = base_name if siblings == 0 else f"{base_name}_variant_{siblings + 1}"
                if not _FUNCTION_NAME.fullmatch(function_name):
                    raise CalculationBlocked(f"generated formula function name is invalid: {function_name}")
                body_to_function[body_key] = function_name
                functions[function_name] = f"def {function_name}({signature}):\n{body}"
            family_function_names[(signature, body)] = function_name
        primary_name = family_function_names[(primary_signature, primary_body)]
        mapping_to_function[family_id] = primary_name
        family_variants[family_id] = [family_function_names[key] for key in ordered_implementations]
        primary_bindings = next(bindings for _address, signature, body, bindings in member_implementations
                                if (signature, body) == (primary_signature, primary_body))
        source_bindings[family_id] = list(primary_bindings)
        for address, signature, body, bindings in member_implementations:
            function_name = family_function_names[(signature, body)]
            if function_name != primary_name or bindings != primary_bindings:
                family_exceptions.append({"source_family_id": family_id,
                                          "address": address, "function_name": function_name,
                                          "binding_variables": list(bindings)})

    active_references = [item for item in unsupported_references
                         if _active_trace_reference(trace, item["formula_cell"], item["reference"],
                                                    item["formula_cell"].rsplit("!", 1)[0])]
    active_names = [item for item in unsupported_names if _active_trace_name(trace, item["name"])]
    if active_references or active_names:
        examples = [f"{item['formula_cell']} -> {item['reference']}" for item in active_references[:3]]
        examples.extend(f"{item['formula_cell']} -> name {item['name']}" for item in active_names[:3])
        raise CalculationBlocked("active source trace has unsupported emitted reference/name: "
                                 + "; ".join(examples))

    schedule = _build_execution_schedule(semantic_plan, trace, source_index)
    return {"functions": functions, "source_family_to_function": mapping_to_function,
            "source_family_bindings": source_bindings,
            "source_family_variants": family_variants, "member_function_exceptions": family_exceptions,
            "deferred_unsupported_references": unsupported_references,
            "deferred_unsupported_names": unsupported_names,
            "execution_schedule": schedule,
            "source_family_count": len(mapping_to_function), "compiled_member_count": compiled_members,
            "member_exception_count": len(family_exceptions), "function_count": len(functions)}


def _build_execution_schedule(
    semantic_plan: dict[str, Any], trace: dict[str, Any],
    source_index: dict[tuple[str, int, int], list[SourceCell]],
) -> dict[str, Any]:
    """Derive same-pass formula order from the checked cell-level trace edges."""
    variables = semantic_plan.get("variables", {})
    if isinstance(variables, list):
        variables = {item.get("variable_id"): item for item in variables}
    mappings = semantic_plan.get("source_mappings", [])
    passes = semantic_plan.get("execution_plan", {}).get("passes", [])
    if not isinstance(variables, dict) or not isinstance(mappings, list) or not isinstance(passes, list):
        raise CalculationBlocked("semantic plan lacks variables, mappings, or execution passes")
    if not passes:
        return {"family_contexts": {}, "family_pass": {}, "family_rank": {},
                "family_dependencies": {}, "same_index_dependencies": {},
                "lagged_edges": [], "same_pass_edge_count": 0,
                "cross_pass_edge_count": 0, "future_edge_count": 0,
                "lagged_edge_count": 0}

    family_owner: dict[str, str] = {}
    family_contexts: dict[str, dict[str, Any]] = {}
    family_pass: dict[str, str] = {}
    family_rank: dict[str, int] = {}
    pass_order = {record.get("pass_id"): record.get("order", position)
                  for position, record in enumerate(passes)}
    for pass_record in passes:
        for rank, family_id in enumerate(pass_record.get("source_family_ids", [])):
            family_pass[family_id] = pass_record["pass_id"]
            family_rank[family_id] = rank
    for mapping in mappings:
        family_id = mapping.get("source_family_id")
        variable_id = mapping.get("variable_id")
        variable = variables.get(variable_id, {})
        if not isinstance(family_id, str) or not isinstance(variable, dict):
            raise CalculationBlocked("semantic mapping has no execution-family owner")
        segments = [segment for segment in variable.get("equation_segments", [])
                    if family_id in segment.get("source_family_ids", [])]
        if len(segments) != 1:
            raise CalculationBlocked(f"family must have exactly one execution segment: {family_id}")
        segment = segments[0]
        axis_name = segment.get("index_axis")
        axes = [axis.get("name") for axis in variable.get("axes", [])]
        dimension = next((index for index, name in enumerate(axes) if name == axis_name), None)
        mapping_axes = mapping.get("index_mapping", {}).get("axes", [])
        axis_mapping = next((item for item in mapping_axes
                             if item.get("dimension") == dimension), {})
        if segment.get("execution") != "scalar" and dimension is None:
            raise CalculationBlocked(f"execution axis {axis_name!r} is not declared: {family_id}")
        family_contexts[family_id] = {
            "variable_id": variable_id,
            "axis": axis_name if dimension is not None else None,
            "dimension": dimension,
            "source_coordinate": axis_mapping.get("source_coordinate"),
            "execution": segment.get("execution"),
            "start": segment.get("index_start", 0),
            "stop": segment.get("index_stop_exclusive", 1),
            "recurrence_group": segment.get("recurrence_group"),
            "snapshot": bool(segment.get("snapshot_before_update")),
        }
        for address in mapping.get("source_members", []):
            key = _address_key(address)
            previous = family_owner.setdefault(key, family_id)
            if previous != family_id:
                raise CalculationBlocked(f"formula address belongs to multiple families: {address}")

    for array_group in semantic_plan.get("array_mappings", []):
        for instance in array_group.get("instances", []):
            family_id = instance.get("anchor_source_family_id")
            if not isinstance(family_id, str) or family_id not in family_contexts:
                raise CalculationBlocked("array instance has no mapped formula anchor family")
            for address in [instance.get("anchor"), *instance.get("source_members", [])]:
                if not isinstance(address, str):
                    continue
                key = _address_key(address)
                previous = family_owner.setdefault(key, family_id)
                if previous != family_id:
                    raise CalculationBlocked(f"array member belongs to multiple families: {address}")

    formula_members_by_row: dict[str, dict[int, list[tuple[int, str]]]] = {}
    for address in family_owner:
        reference = _parse_reference(address, "Main")
        if reference is None or not reference.single:
            continue
        formula_members_by_row.setdefault(reference.sheet.casefold(), {}).setdefault(
            reference.min_row, []).append((reference.min_col, address))
    formula_rows = {sheet: sorted(rows) for sheet, rows in formula_members_by_row.items()}
    for rows in formula_members_by_row.values():
        for members in rows.values():
            members.sort()

    def mapped_formula_members(reference_text: Any) -> list[str]:
        """Expand a finite trace range only to known mapped formula outputs."""
        reference = _parse_reference(reference_text, "Main")
        if reference is None or not reference.sheet:
            return []
        if reference.single:
            address = _address_key(reference_text)
            return [address] if address in family_owner else []
        if any(value is None for value in (
                reference.min_row, reference.max_row, reference.min_col, reference.max_col)):
            return []
        sheet = reference.sheet.casefold()
        rows = formula_rows.get(sheet, [])
        first = bisect_left(rows, reference.min_row)
        last = bisect_right(rows, reference.max_row)
        members_by_row = formula_members_by_row.get(sheet, {})
        return [address for row in rows[first:last]
                for column, address in members_by_row[row]
                if reference.min_col <= column <= reference.max_col]

    lookups_by_consumer: dict[str, list[dict[str, Any]]] = {}
    for lookup in trace.get("lookups", []):
        if isinstance(lookup, dict) and isinstance(lookup.get("formula_cell"), str):
            lookups_by_consumer.setdefault(lookup["formula_cell"].casefold(), []).append(lookup)

    def lookup_formula_members(consumer_address: str, reference_text: Any) -> list[str]:
        """Expand a VLOOKUP table edge to its searched key and selected return columns."""
        table = _parse_reference(reference_text, "Main")
        if (table is None or table.single or any(value is None for value in (
                table.min_row, table.max_row, table.min_col, table.max_col))):
            raise CalculationBlocked(
                f"active lookup-table edge has no finite table bounds: {consumer_address} -> {reference_text}"
            )

        def same_table(record: dict[str, Any]) -> bool:
            candidate = _parse_reference(record.get("table_range"), "Main")
            return (
                candidate is not None
                and candidate.sheet.casefold() == table.sheet.casefold()
                and (candidate.min_row, candidate.max_row, candidate.min_col, candidate.max_col)
                == (table.min_row, table.max_row, table.min_col, table.max_col)
            )

        records = [record for record in lookups_by_consumer.get(consumer_address.casefold(), [])
                   if record.get("function", "").upper() == "VLOOKUP" and same_table(record)]
        if not records:
            raise CalculationBlocked(
                f"active lookup-table edge lacks a matching VLOOKUP footprint: "
                f"{consumer_address} -> {reference_text}"
            )
        member_addresses: set[str] = set()
        for record in records:
            lookup_column = record.get("lookup_column")
            return_column = record.get("return_column")
            width = table.max_col - table.min_col + 1
            if (not isinstance(lookup_column, int) or not isinstance(return_column, int)
                    or not 1 <= lookup_column <= width or not 1 <= return_column <= width):
                raise CalculationBlocked(
                    f"active VLOOKUP footprint has an invalid column at {consumer_address}"
                )
            for relative_column in {lookup_column, return_column}:
                column = table.min_col + relative_column - 1
                column_range = (
                    f"{quote_sheetname(table.sheet)}!{get_column_letter(column)}{table.min_row}:"
                    f"{get_column_letter(column)}{table.max_row}"
                )
                member_addresses.update(mapped_formula_members(column_range))
        return sorted(member_addresses)

    def logical_axis_index(address: str, family_id: str) -> int | None:
        context = family_contexts[family_id]
        axis_name = context.get("axis")
        if not isinstance(axis_name, str):
            return None
        key = _address_key(address)
        sheet, local = key.split("!", 1)
        col, row, _max_col, _max_row = range_boundaries(local)
        candidates = [cell for cell in source_index.get((sheet, row, col), [])
                      if cell.variable_id == context["variable_id"]]
        coordinates = {
            cell.indices[dimension]
            for cell in candidates
            for dimension, name in enumerate(cell.axes)
            if name == axis_name and dimension < len(cell.indices)
        }
        if len(coordinates) > 1:
            raise CalculationBlocked(f"formula address has ambiguous logical time coordinate: {address}")
        return next(iter(coordinates)) if coordinates else None

    def schedule_coordinate(address: str, family_id: str, *, source_row: bool) -> int | None:
        if source_row:
            parsed = _parse_reference(address, "Main")
            return parsed.min_row if parsed is not None and parsed.single else None
        return logical_axis_index(address, family_id)

    family_dependencies: dict[str, set[str]] = {family_id: set() for family_id in family_contexts}
    same_index_dependencies: dict[str, set[str]] = {family_id: set() for family_id in family_contexts}
    lagged_edges: list[dict[str, Any]] = []
    same_pass_edges = 0
    cross_pass_edges = 0
    future_edges: list[dict[str, Any]] = []
    for edge in trace.get("edges", []):
        if not isinstance(edge, dict):
            continue
        if edge.get("relationship") == "array_formula_anchor":
            # Followers share the anchor's generated array value; this provenance edge
            # is not a calculation dependency between two logical formula families.
            continue
        try:
            consumer_address = _address_key(edge.get("consumer"))
        except CalculationBlocked:
            continue
        consumer_family = family_owner.get(consumer_address)
        if consumer_family is None:
            continue
        if edge.get("relationship") == "lookup_table":
            prerequisite_addresses = lookup_formula_members(
                consumer_address, edge.get("prerequisite"))
        else:
            prerequisite_addresses = mapped_formula_members(edge.get("prerequisite"))
        for prerequisite_address in prerequisite_addresses:
            prerequisite_family = family_owner.get(prerequisite_address)
            if prerequisite_family is None:
                continue
            consumer_pass = family_pass.get(consumer_family)
            prerequisite_pass = family_pass.get(prerequisite_family)
            if consumer_pass is None or prerequisite_pass is None:
                raise CalculationBlocked("active formula edge references a family outside the execution plan")
            if consumer_pass != prerequisite_pass:
                if pass_order.get(prerequisite_pass, -1) > pass_order.get(consumer_pass, -1):
                    raise CalculationBlocked(
                        "execution plan places a formula before its prerequisite: "
                        f"{consumer_address} needs {prerequisite_address}"
                    )
                cross_pass_edges += 1
                continue
            if consumer_family == prerequisite_family:
                # Same-family lagged recurrences are common; only a same-index edge is a cycle.
                source_row = family_contexts[consumer_family].get("source_coordinate") == "row"
                consumer_index = schedule_coordinate(consumer_address, consumer_family, source_row=source_row)
                prerequisite_index = schedule_coordinate(prerequisite_address, prerequisite_family,
                                                        source_row=source_row)
                if consumer_index is not None and prerequisite_index is not None:
                    if prerequisite_index == consumer_index:
                        same_index_dependencies[consumer_family].add(prerequisite_family)
                        family_dependencies[consumer_family].add(prerequisite_family)
                    elif prerequisite_index < consumer_index:
                        lagged_edges.append({"consumer": consumer_address,
                                             "prerequisite": prerequisite_address,
                                             "lag": consumer_index - prerequisite_index,
                                             "family_id": consumer_family})
                    else:
                        future_edges.append({"consumer": consumer_address,
                                             "prerequisite": prerequisite_address,
                                             "family_id": consumer_family})
                continue
            same_pass_edges += 1
            consumer_context = family_contexts[consumer_family]
            prerequisite_context = family_contexts[prerequisite_family]
            consumer_cluster = (consumer_context.get("axis"), consumer_context.get("dimension"))
            prerequisite_cluster = (prerequisite_context.get("axis"), prerequisite_context.get("dimension"))
            if consumer_cluster != prerequisite_cluster:
                family_dependencies[consumer_family].add(prerequisite_family)
                continue
            source_row = (consumer_context.get("source_coordinate") == "row"
                          and prerequisite_context.get("source_coordinate") == "row")
            consumer_index = schedule_coordinate(consumer_address, consumer_family, source_row=source_row)
            prerequisite_index = schedule_coordinate(prerequisite_address, prerequisite_family,
                                                     source_row=source_row)
            if consumer_index is None or prerequisite_index is None:
                family_dependencies[consumer_family].add(prerequisite_family)
                continue
            if prerequisite_index == consumer_index:
                same_index_dependencies[consumer_family].add(prerequisite_family)
                family_dependencies[consumer_family].add(prerequisite_family)
            elif prerequisite_index < consumer_index:
                lagged_edges.append({"consumer": consumer_address,
                                     "prerequisite": prerequisite_address,
                                     "lag": consumer_index - prerequisite_index,
                                     "family_id": consumer_family})
            else:
                future_edges.append({"consumer": consumer_address,
                                     "prerequisite": prerequisite_address,
                                     "family_id": consumer_family})
    if future_edges:
        sample = future_edges[0]
        raise CalculationBlocked(
            "ascending logical loop has a future-index formula prerequisite: "
            f"{sample['consumer']} needs {sample['prerequisite']}"
        )
    return {
        "family_contexts": family_contexts,
        "family_pass": family_pass,
        "family_rank": family_rank,
        "family_dependencies": {key: sorted(value) for key, value in family_dependencies.items()},
        "same_index_dependencies": {key: sorted(value) for key, value in same_index_dependencies.items()},
        "lagged_edges": lagged_edges,
        "same_pass_edge_count": same_pass_edges,
        "cross_pass_edge_count": cross_pass_edges,
        "future_edge_count": len(future_edges),
        "lagged_edge_count": len(lagged_edges),
    }


def render_grouped_driver(semantic_plan: dict[str, Any], compilation: dict[str, Any]) -> str:
    """Render logical-axis loops in trace-proven prerequisite order."""
    variables = semantic_plan.get("variables", {})
    if isinstance(variables, list):
        variables = {item.get("variable_id"): item for item in variables}
    mappings = semantic_plan.get("source_mappings", [])
    passes = semantic_plan.get("execution_plan", {}).get("passes", [])
    functions = compilation.get("source_family_to_function", {})
    schedule = compilation.get("execution_schedule", {})
    if (not isinstance(variables, dict) or not isinstance(mappings, list)
            or not isinstance(passes, list) or not isinstance(functions, dict)
            or not isinstance(schedule, dict)):
        raise CalculationBlocked("semantic plan has no valid source-family execution schedule")

    mappings_by_id = {item.get("source_family_id"): item for item in mappings
                      if isinstance(item, dict)}
    contexts = schedule.get("family_contexts", {})
    dependencies = schedule.get("family_dependencies", {})
    ranks = schedule.get("family_rank", {})
    same_index_dependencies = schedule.get("same_index_dependencies", {})
    if not all(isinstance(value, dict) for value in
               (contexts, dependencies, ranks, same_index_dependencies)):
        raise CalculationBlocked("trace-derived family execution schedule is incomplete")

    source_index = build_source_cell_index(semantic_plan)
    exceptions = {
        (item.get("source_family_id"), _address_key(item.get("address"))): item
        for item in compilation.get("member_function_exceptions", [])
        if isinstance(item, dict)
    }
    array_instances = {
        instance.get("anchor_source_family_id"): instance
        for group in semantic_plan.get("array_mappings", []) if isinstance(group, dict)
        for instance in group.get("instances", []) if isinstance(instance, dict)
    }

    emitted_schedule: dict[str, Any] = {}
    lines = [
        "def run_pricing(values, family_functions):",
        "    \"\"\"Run named-variable passes with yearly prerequisites in source order.\"\"\"",
        "    def _store(variable_id, indices, value):",
        "        target = values[variable_id]",
        "        if not indices:",
        "            values[variable_id] = value",
        "        elif len(indices) == 1:",
        "            target[indices[0]] = value",
        "        elif len(indices) == 2:",
        "            target[indices[0]][indices[1]] = value",
        "        else:",
        "            raise ValueError('logical output has more than two axes')",
        "",
    ]

    for pass_record in passes:
        pass_id = pass_record.get("pass_id")
        pass_families = [family_id for family_id in pass_record.get("source_family_ids", [])
                         if family_id in mappings_by_id]
        if not pass_families:
            continue
        cluster_by_family = {
            family_id: _schedule_cluster(family_id, contexts[family_id], array_instances)
            for family_id in pass_families
        }
        clusters: dict[tuple[Any, ...], list[str]] = {}
        for family_id in pass_families:
            clusters.setdefault(cluster_by_family[family_id], []).append(family_id)
        cluster_rank = {key: min(ranks.get(family_id, 0) for family_id in family_ids)
                        for key, family_ids in clusters.items()}
        cluster_dependencies: dict[tuple[Any, ...], set[tuple[Any, ...]]] = {
            key: set() for key in clusters
        }
        for family_id in pass_families:
            consumer_cluster = cluster_by_family[family_id]
            for prerequisite_family in dependencies.get(family_id, []):
                if prerequisite_family not in cluster_by_family:
                    continue
                prerequisite_cluster = cluster_by_family[prerequisite_family]
                if prerequisite_cluster != consumer_cluster:
                    cluster_dependencies[consumer_cluster].add(prerequisite_cluster)
        ordered_clusters = _topological_order(
            list(clusters), cluster_dependencies, cluster_rank,
            error_message=f"same-pass execution-cluster cycle in {pass_id}",
        )
        pass_lines: list[str] = []
        pass_report: list[dict[str, Any]] = []
        for cluster in ordered_clusters:
            family_ids = sorted(clusters[cluster], key=lambda item: ranks.get(item, 0))
            if cluster[0] == "axis":
                axis_lines, cluster_report = _render_axis_cluster(
                    family_ids, cluster, pass_id, variables, mappings_by_id,
                    contexts, dependencies, same_index_dependencies, ranks,
                    functions, compilation.get("source_family_bindings", {}),
                    exceptions, source_index, array_instances,
                )
            else:
                family_id = family_ids[0]
                axis_lines = _render_single_family(
                    family_id, variables, mappings_by_id, contexts, functions,
                    compilation.get("source_family_bindings", {}), exceptions,
                    source_index, array_instances,
                )
                cluster_report = {"families": family_ids, "loop_axis": None, "node_order": family_ids}
            pass_lines.extend(axis_lines)
            pass_report.append(cluster_report)
        if pass_lines:
            lines.append(f"    # Pass: {pass_id}")
            lines.extend("    " + line if line else "" for line in pass_lines)
            lines.append("")
        emitted_schedule[pass_id] = pass_report

    compilation["emitted_schedule"] = emitted_schedule
    rendered = "\n".join(lines + ["    return values", ""])
    compile(rendered, "pricing.py", "exec")
    _validate_emitted_schedule(emitted_schedule, schedule, passes)
    return rendered


def _schedule_cluster(family_id: str, context: dict[str, Any],
                      arrays: dict[str, dict[str, Any]]) -> tuple[Any, ...]:
    if family_id in arrays or context.get("axis") is None:
        return ("single", family_id)
    return ("axis", context["axis"], context["dimension"])


def _topological_order(nodes: list[Any], dependencies: dict[Any, set[Any]],
                       rank: dict[Any, int], *, priorities: dict[Any, int] | None = None,
                       error_message: str) -> list[Any]:
    remaining = {node: set(dependencies.get(node, set())) for node in nodes}
    dependents: dict[Any, set[Any]] = {node: set() for node in nodes}
    for consumer, prerequisites in remaining.items():
        for prerequisite in prerequisites:
            if prerequisite not in remaining:
                raise CalculationBlocked(f"execution dependency points outside its schedule: {prerequisite}")
            dependents[prerequisite].add(consumer)
    priorities = priorities or {}
    ordered: list[Any] = []
    ready = [node for node in nodes if not remaining[node]]
    while ready:
        ready.sort(key=lambda node: (priorities.get(node, 0), rank.get(node, 0), repr(node)))
        node = ready.pop(0)
        ordered.append(node)
        for consumer in dependents[node]:
            remaining[consumer].discard(node)
            if not remaining[consumer] and consumer not in ordered and consumer not in ready:
                ready.append(consumer)
    if len(ordered) != len(nodes):
        raise CalculationBlocked(error_message)
    return ordered


def _axis_schedule_nodes(
    family_ids: list[str], dependencies: dict[str, list[str]],
    same_index_dependencies: dict[str, list[str]], contexts: dict[str, dict[str, Any]],
    ranks: dict[str, int],
) -> tuple[list[tuple[str, Any]], dict[tuple[str, Any], list[str]]]:
    node_by_family: dict[str, tuple[str, Any]] = {}
    node_families: dict[tuple[str, Any], list[str]] = {}
    recurrence_contracts: dict[Any, tuple[Any, ...]] = {}
    for family_id in family_ids:
        context = contexts[family_id]
        if context.get("execution") == "ascending_recurrence":
            recurrence_group = context.get("recurrence_group")
            contract = (context.get("axis"), context.get("dimension"), context.get("start"),
                        context.get("stop"), context.get("snapshot"))
            previous = recurrence_contracts.setdefault(recurrence_group, contract)
            if (not recurrence_group or previous != contract or not context.get("snapshot")):
                raise CalculationBlocked(f"inconsistent recurrence contract: {recurrence_group}")
            node = ("recurrence", recurrence_group)
        else:
            node = ("family", family_id)
        node_by_family[family_id] = node
        node_families.setdefault(node, []).append(family_id)
    node_dependencies: dict[tuple[str, Any], set[tuple[str, Any]]] = {
        node: set() for node in node_families
    }
    for family_id in family_ids:
        consumer_node = node_by_family[family_id]
        for prerequisite_family in same_index_dependencies.get(family_id, []):
            if prerequisite_family not in node_by_family:
                continue
            prerequisite_node = node_by_family[prerequisite_family]
            if prerequisite_node == consumer_node:
                raise CalculationBlocked(
                    "same-year dependency inside an atomic recurrence group: "
                    f"{family_id} requires {prerequisite_family}"
                )
            node_dependencies[consumer_node].add(prerequisite_node)
    node_rank = {node: min(ranks.get(family_id, 0) for family_id in family_list)
                 for node, family_list in node_families.items()}
    recurrence_nodes = {node for node in node_families if node[0] == "recurrence"}
    ancestor_nodes: set[tuple[str, Any]] = set()
    pending = list(recurrence_nodes)
    while pending:
        node = pending.pop()
        for prerequisite in node_dependencies.get(node, set()):
            if prerequisite not in ancestor_nodes and prerequisite not in recurrence_nodes:
                ancestor_nodes.add(prerequisite)
                pending.append(prerequisite)
    priorities = {node: (0 if node in ancestor_nodes else 1 if node in recurrence_nodes else 2)
                  for node in node_families}
    order = _topological_order(
        list(node_families), node_dependencies, node_rank, priorities=priorities,
        error_message="same-index logical formula dependencies contain a cycle",
    )
    for family_list in node_families.values():
        family_list.sort(key=lambda family_id: ranks.get(family_id, 0))
    return order, node_families


def _member_calls(
    family_id: str, mappings: dict[str, dict[str, Any]], contexts: dict[str, dict[str, Any]],
    functions: dict[str, str], family_bindings: dict[str, list[str]],
    exceptions: dict[tuple[str, str], dict[str, Any]],
    source_index: dict[tuple[str, int, int], list[SourceCell]],
) -> list[dict[str, Any]]:
    mapping = mappings.get(family_id)
    context = contexts.get(family_id)
    if not isinstance(mapping, dict) or not isinstance(context, dict) or family_id not in functions:
        raise CalculationBlocked(f"execution pass references an unmapped family: {family_id}")
    variable_id = mapping.get("variable_id")
    result = []
    for member_address in mapping.get("source_members", []):
        key = _address_key(member_address)
        parsed = _parse_reference(member_address, "Main")
        if parsed is None or not parsed.single:
            raise CalculationBlocked(f"source family has a non-cell member: {member_address}")
        candidates = [cell for cell in source_index.get(
            (parsed.sheet.casefold(), parsed.min_row, parsed.min_col), [])
            if cell.variable_id == variable_id]
        unique = {(cell.indices, cell.axes): cell for cell in candidates}
        if len(unique) != 1:
            raise CalculationBlocked(f"family member has no unique logical output: {member_address}")
        indices = next(iter(unique))[0]
        override = exceptions.get((family_id, key), {})
        function_name = override.get("function_name", functions[family_id])
        bindings = override.get("binding_variables", family_bindings.get(family_id, []))
        if not isinstance(bindings, list):
            raise CalculationBlocked(f"family has invalid reference bindings: {family_id}")
        result.append({"address": key, "variable_id": variable_id, "indices": indices,
                       "function_name": function_name, "bindings": tuple(bindings)})
    return result


def _render_axis_cluster(
    family_ids: list[str], cluster: tuple[Any, ...], pass_id: str,
    variables: dict[str, dict[str, Any]], mappings: dict[str, dict[str, Any]],
    contexts: dict[str, dict[str, Any]], dependencies: dict[str, list[str]],
    same_index_dependencies: dict[str, list[str]], ranks: dict[str, int],
    functions: dict[str, str], family_bindings: dict[str, list[str]],
    exceptions: dict[tuple[str, str], dict[str, Any]],
    source_index: dict[tuple[str, int, int], list[SourceCell]],
    arrays: dict[str, dict[str, Any]],
) -> tuple[list[str], dict[str, Any]]:
    axis_name, dimension = cluster[1], cluster[2]
    if not isinstance(axis_name, str) or not isinstance(dimension, int):
        raise CalculationBlocked(f"invalid logical execution axis in {pass_id}")
    axis = _python_name(axis_name)
    node_order, node_families = _axis_schedule_nodes(
        family_ids, dependencies, same_index_dependencies, contexts, ranks)
    row_based = all(contexts[family_id].get("source_coordinate") == "row"
                    for family_id in family_ids)
    row_origin = None
    if row_based:
        formula_rows = [_parse_reference(call["address"], "Main").min_row
                        for family_id in family_ids
                        for call in _member_calls(family_id, mappings, contexts, functions,
                                                  family_bindings, exceptions, source_index)]
        if not formula_rows:
            raise CalculationBlocked(f"logical row loop has no formula members: {axis_name}")
        row_origin = min(formula_rows)
        start, stop = 0, max(formula_rows) - row_origin + 1
    else:
        start = min(contexts[family_id]["start"] for family_id in family_ids)
        stop = max(contexts[family_id]["stop"] for family_id in family_ids)
    if not isinstance(start, int) or not isinstance(stop, int) or stop <= start:
        raise CalculationBlocked(f"logical loop has invalid bounds: {axis_name}")

    calls_by_family: dict[str, list[dict[str, Any]]] = {}
    for family_id in family_ids:
        calls = _member_calls(family_id, mappings, contexts, functions, family_bindings,
                              exceptions, source_index)
        context = contexts[family_id]
        variable = variables[context["variable_id"]]
        shape = variable.get("shape", [])
        if not isinstance(shape, list) or dimension >= len(shape):
            raise CalculationBlocked(f"family logical rank differs from its loop axis: {family_id}")
        for call in calls:
            point = call["indices"][dimension]
            if point < context["start"] or point >= context["stop"]:
                raise CalculationBlocked(f"family member is outside its declared segment: {call['address']}")
            if row_based:
                source_row = _parse_reference(call["address"], "Main").min_row
                call["schedule_point"] = source_row - row_origin
                call["logical_offset"] = point - call["schedule_point"]
            else:
                call["schedule_point"] = point
                call["logical_offset"] = 0
        calls_by_family[family_id] = calls

    lines = [f"for {axis} in range({start}, {stop}):"]
    for node in node_order:
        family_list = node_families[node]
        recurrence = node[0] == "recurrence"
        if recurrence:
            lines.append("    _pending = []")
        for family_id in family_list:
            context = contexts[family_id]
            point_groups: dict[tuple[Any, ...], set[int]] = {}
            for call in calls_by_family[family_id]:
                indices = call["indices"]
                fixed_indices = tuple(value for index, value in enumerate(indices) if index != dimension)
                signature = (call["function_name"], call["bindings"], fixed_indices,
                             call["logical_offset"],
                             call["variable_id"])
                point_groups.setdefault(signature, set()).add(call["schedule_point"])
            for (function_name, bindings, fixed_indices, logical_offset, variable_id), points in sorted(
                    point_groups.items(), key=lambda item: (item[0][0], item[0][1], item[0][2], item[0][3])):
                ordered_points = sorted(points)
                contiguous = ordered_points == list(range(ordered_points[0], ordered_points[-1] + 1))
                if ordered_points == list(range(start, stop)):
                    condition = None
                elif contiguous:
                    condition = f"{ordered_points[0]} <= {axis} < {ordered_points[-1] + 1}"
                else:
                    condition = f"{axis} in {tuple(ordered_points)!r}"
                if condition:
                    lines.append(f"    if {condition}:")
                    indent = "        "
                else:
                    indent = "    "
                logical_axis_index = (axis if logical_offset == 0 else
                                      f"({axis} {'+' if logical_offset > 0 else '-'} {abs(logical_offset)})")
                indices = list(fixed_indices)
                indices.insert(dimension, logical_axis_index)
                index_args = ", ".join(str(value) for value in indices)
                call = (f"family_functions[{function_name!r}](values, {list(bindings)!r}"
                        f"{', ' + index_args if index_args else ''})")
                index_tuple = "(" + ", ".join(str(value) for value in indices) + ")"
                if len(indices) == 1:
                    index_tuple = f"({indices[0]},)"
                if recurrence:
                    lines.append(f"{indent}_pending.append(({variable_id!r}, {index_tuple}, {call}))")
                else:
                    target = f"values[{variable_id!r}]"
                    for value in indices:
                        target += f"[{value}]"
                    lines.append(f"{indent}{target} = {call}")
        if recurrence:
            lines.append("    for _variable_id, _indices, _value in _pending:")
            lines.append("        _store(_variable_id, _indices, _value)")
    lines.append("")
    family_order = [family_id for node in node_order for family_id in node_families[node]]
    return lines, {"families": family_ids, "loop_axis": axis_name,
                   "loop_dimension": dimension, "loop_start": start, "loop_stop": stop,
                   "node_order": [repr(node) for node in node_order],
                   "family_order": family_order}


def _render_single_family(
    family_id: str, variables: dict[str, dict[str, Any]], mappings: dict[str, dict[str, Any]],
    contexts: dict[str, dict[str, Any]], functions: dict[str, str],
    family_bindings: dict[str, list[str]], exceptions: dict[tuple[str, str], dict[str, Any]],
    source_index: dict[tuple[str, int, int], list[SourceCell]],
    arrays: dict[str, dict[str, Any]],
) -> list[str]:
    context = contexts[family_id]
    variable_id = context["variable_id"]
    variable = variables[variable_id]
    calls = _member_calls(family_id, mappings, contexts, functions, family_bindings,
                          exceptions, source_index)
    if family_id in arrays:
        shape = variable.get("shape", [])
        instance = arrays[family_id]
        if instance.get("variable_id") != variable_id:
            raise CalculationBlocked(f"array anchor target differs from its source mapping: {family_id}")
        anchor = _address_key(instance["anchor"])
        override = exceptions.get((family_id, anchor), {})
        function_name = override.get("function_name", functions[family_id])
        bindings = override.get("binding_variables", family_bindings.get(family_id, []))
        zero_args = ", ".join("0" for _ in shape)
        call = (f"family_functions[{function_name!r}](values, {list(bindings)!r}"
                f"{', ' + zero_args if zero_args else ''})")
        return [f"values[{variable_id!r}] = {call}"]
    if context.get("execution") == "scalar":
        lines = []
        for call in calls:
            if call["indices"]:
                raise CalculationBlocked(f"scalar output has logical indices: {family_id}")
            lines.append(f"values[{variable_id!r}] = family_functions[{call['function_name']!r}]"
                         f"(values, {list(call['bindings'])!r})")
        return lines
    raise CalculationBlocked(f"non-scalar family lacks a schedulable logical axis: {family_id}")


def _validate_emitted_schedule(emitted: dict[str, Any], schedule: dict[str, Any],
                               passes: list[dict[str, Any]]) -> None:
    family_pass = schedule["family_pass"]
    dependencies = schedule["family_dependencies"]
    positions: dict[str, tuple[str, int]] = {}
    clusters_by_family: dict[str, tuple[Any, ...]] = {}
    for pass_record in passes:
        pass_id = pass_record.get("pass_id")
        for cluster_index, cluster_report in enumerate(emitted.get(pass_id, [])):
            family_ids = cluster_report.get("family_order", cluster_report.get("families", []))
            cluster_key = (("axis", cluster_report["loop_axis"], cluster_report["loop_dimension"])
                           if cluster_report.get("loop_axis") is not None
                           else None)
            for position, family_id in enumerate(family_ids):
                positions[family_id] = (pass_id, cluster_index, position)
                clusters_by_family[family_id] = cluster_key or ("single", family_id)
    for consumer_family, prerequisites in dependencies.items():
        for prerequisite_family in prerequisites:
            if family_pass.get(consumer_family) != family_pass.get(prerequisite_family):
                continue
            if consumer_family not in positions or prerequisite_family not in positions:
                raise CalculationBlocked("emitted driver omitted a formula family with a prerequisite")
            consumer_cluster = clusters_by_family[consumer_family]
            prerequisite_cluster = clusters_by_family[prerequisite_family]
            if consumer_cluster == prerequisite_cluster:
                if positions[prerequisite_family][2] >= positions[consumer_family][2]:
                    raise CalculationBlocked(
                        f"emitted family order violates a source prerequisite: {consumer_family} needs {prerequisite_family}")
            elif positions[prerequisite_family][1] >= positions[consumer_family][1]:
                raise CalculationBlocked(
                    f"emitted pass-cluster order violates a source prerequisite: {consumer_family} needs {prerequisite_family}")


def _render_logical_groups(grouped: dict[tuple[Any, ...], list[dict[str, Any]]]) -> list[str]:
    """Render contiguous family members as axis loops, with atomic recurrence writes."""
    blocks: dict[tuple[Any, ...], list[dict[str, Any]]] = {}
    for key, members in grouped.items():
        (pass_id, execution_kind, axis_name, dimension, _segment_start, _segment_stop,
         recurrence_group, snapshot, variable_id, function_name, bindings, fixed_indices) = key
        ordered = sorted({member["index"] for member in members})
        if not ordered:
            continue
        if execution_kind == "ascending_recurrence":
            block_key = (pass_id, execution_kind, axis_name, dimension, _segment_start, _segment_stop,
                         recurrence_group, snapshot)
            blocks.setdefault(block_key, []).append({
                "variable_id": variable_id, "function_name": function_name,
                "bindings": bindings, "fixed_indices": fixed_indices,
                "active_indices": tuple(ordered),
            })
            continue
        start = previous = ordered[0]
        ranges: list[tuple[int, int]] = []
        for index in ordered[1:]:
            if index != previous + 1:
                ranges.append((start, previous + 1))
                start = index
            previous = index
        ranges.append((start, previous + 1))
        for start, stop in ranges:
            block_key = (pass_id, execution_kind, axis_name, dimension, start, stop,
                         recurrence_group, snapshot)
            blocks.setdefault(block_key, []).append({
                "variable_id": variable_id, "function_name": function_name,
                "bindings": bindings, "fixed_indices": fixed_indices,
                "active_indices": tuple(range(start, stop)),
            })

    rendered: list[str] = []
    for block_key, statements in sorted(blocks.items(), key=lambda item: item[0]):
        (_pass_id, execution_kind, axis_name, dimension, start, stop,
         recurrence_group, snapshot) = block_key
        loop_name = _python_name(axis_name or "index")
        if execution_kind == "scalar":
            continue
        if execution_kind not in {"vectorized", "ascending_recurrence"}:
            raise CalculationBlocked(f"unsupported execution segment kind: {execution_kind}")
        if execution_kind == "ascending_recurrence" and not recurrence_group:
            raise CalculationBlocked("ascending recurrence has no declared recurrence group")
        if execution_kind == "ascending_recurrence" and not snapshot:
            raise CalculationBlocked("ascending recurrence does not declare snapshot-before-update")
        rendered.append(f"for {loop_name} in range({start}, {stop}):")
        if execution_kind == "ascending_recurrence":
            rendered.append("    _pending = []")
        for item in statements:
            indices = list(item["fixed_indices"])
            indices.insert(dimension, loop_name)
            args = ", ".join(str(index) for index in indices)
            call = (f"family_functions[{item['function_name']!r}](values, {item['bindings']!r}"
                    f"{', ' + args if args else ''})")
            indices_literal = "(" + ", ".join(str(index) for index in indices) + ")"
            if len(indices) == 1:
                indices_literal = f"({indices[0]},)"
            if execution_kind == "ascending_recurrence":
                active = item["active_indices"]
                condition = ("" if active == tuple(range(start, stop)) else
                             f"    if {loop_name} in {active!r}:\n")
                if condition:
                    rendered.append(condition.rstrip("\n"))
                    rendered.append(
                        f"        _pending.append(({item['variable_id']!r}, {indices_literal}, {call}))")
                else:
                    rendered.append(
                        f"    _pending.append(({item['variable_id']!r}, {indices_literal}, {call}))")
            else:
                target = f"values[{item['variable_id']!r}]"
                for index in indices:
                    target += f"[{index}]"
                rendered.append(f"    {target} = {call}")
        if execution_kind == "ascending_recurrence":
            rendered.append("    for _variable_id, _indices, _value in _pending:")
            rendered.append("        _store(_variable_id, _indices, _value)")
        rendered.append("")
    return rendered


def _python_name(value: str) -> str:
    normalized = re.sub(r"\W+", "_", value)
    if not normalized or normalized[0].isdigit():
        normalized = "axis_" + normalized
    return "_" + normalized.casefold()


def _address_key(address: Any) -> str:
    if not isinstance(address, str) or "!" not in address:
        raise CalculationBlocked(f"source map contains an unqualified cell address: {address!r}")
    sheet, local = address.rsplit("!", 1)
    bounds = range_boundaries(local.replace("$", ""))
    if any(item is None for item in bounds) or bounds[0] != bounds[2] or bounds[1] != bounds[3]:
        raise CalculationBlocked(f"source map member is not a single cell: {address}")
    col, row, _max_col, _max_row = bounds
    return f"{sheet.casefold()}!{get_column_letter(col)}{row}"


def _array_key(sheet: str, array_ref: Any) -> tuple[str, str]:
    if not isinstance(array_ref, str) or not array_ref:
        raise CalculationBlocked("array instance is missing its source bounds")
    bounds = range_boundaries(array_ref.replace("$", ""))
    if any(item is None for item in bounds):
        raise CalculationBlocked(f"array instance has incomplete bounds: {array_ref}")
    min_col, min_row, max_col, max_row = bounds
    first = f"{get_column_letter(min_col)}{min_row}"
    last = f"{get_column_letter(max_col)}{max_row}"
    return sheet.casefold(), first if first == last else f"{first}:{last}"


def validate_family_coverage(
    semantic_plan: dict[str, Any], source_profile: dict[str, Any], trace: dict[str, Any],
) -> dict[str, int]:
    """Require the fresh source trace to match every reviewed ordinary and array member."""
    plan_source = semantic_plan.get("source", {})
    source_hash = plan_source.get("source_sha256")
    if (not isinstance(source_hash, str) or trace.get("source_sha256") != source_hash
            or source_profile.get("source_sha256") != source_hash):
        raise CalculationBlocked("semantic map, source profile, and active trace source hashes differ")

    profile_families = source_profile.get("families")
    mappings = semantic_plan.get("source_mappings")
    variables = semantic_plan.get("variables")
    equation_families = semantic_plan.get("equation_families")
    if not isinstance(profile_families, list) or not isinstance(mappings, list):
        raise CalculationBlocked("semantic plan or source profile has no family member list")
    variable_ids = ({item.get("variable_id") for item in variables}
                    if isinstance(variables, list) else set(variables or {}))
    equation_ids = ({item.get("family_id") for item in equation_families}
                    if isinstance(equation_families, list) else set(equation_families or {}))

    profiled: dict[str, set[str]] = {}
    for family in profile_families:
        family_id = family.get("family_id")
        members = family.get("source_members")
        if not isinstance(family_id, str) or family_id in profiled or not isinstance(members, list):
            raise CalculationBlocked("source profile has an invalid or duplicate family ID")
        normalized = {_address_key(address) for address in members}
        if len(normalized) != len(members) or family.get("member_count") != len(members):
            raise CalculationBlocked(f"source profile family has duplicate or inconsistent members: {family_id}")
        profiled[family_id] = normalized

    mapped: dict[str, set[str]] = {}
    mapped_members: set[str] = set()
    for item in mappings:
        family_id = item.get("source_family_id")
        members = item.get("source_members")
        if family_id not in profiled or family_id in mapped or not isinstance(members, list):
            raise CalculationBlocked(f"semantic source mapping is missing, duplicate, or unknown: {family_id}")
        normalized = {_address_key(address) for address in members}
        if (len(normalized) != len(members) or normalized != profiled[family_id]
                or item.get("member_count") != len(members)):
            raise CalculationBlocked(f"semantic source member mapping differs from its profile: {family_id}")
        if item.get("variable_id") not in variable_ids:
            raise CalculationBlocked(f"source family has an unknown target variable: {family_id}")
        if item.get("equation_family_id") not in equation_ids:
            raise CalculationBlocked(f"source family has an unknown equation family: {family_id}")
        duplicate = mapped_members & normalized
        if duplicate:
            raise CalculationBlocked("active formula member appears in multiple source mappings: "
                                      + ", ".join(sorted(duplicate)[:5]))
        mapped[family_id] = normalized
        mapped_members.update(normalized)
    if set(mapped) != set(profiled):
        raise CalculationBlocked("semantic source mappings do not cover every profiled formula family")

    cells = trace.get("cells")
    if not isinstance(cells, list):
        raise CalculationBlocked("active trace has no cell list")
    ordinary_trace: set[str] = set()
    follower_trace: dict[tuple[str, str], set[str]] = {}
    follower_records: dict[str, dict[str, Any]] = {}
    anchor_trace: dict[str, dict[str, Any]] = {}
    for cell in cells:
        role = cell.get("role")
        if role == "calculated_formula":
            address = _address_key(cell.get("address"))
            if address in ordinary_trace:
                raise CalculationBlocked(f"active trace duplicates formula member: {address}")
            ordinary_trace.add(address)
            anchor_trace[address] = cell
        elif role == "calculated_array_formula":
            address = _address_key(cell.get("address"))
            sheet = address.split("!", 1)[0]
            array_key = _array_key(sheet, cell.get("array_ref"))
            follower_trace.setdefault(array_key, set()).add(address)
            if address in follower_records:
                raise CalculationBlocked(f"active trace duplicates array follower: {address}")
            follower_records[address] = cell
    if ordinary_trace != mapped_members:
        raise CalculationBlocked(
            f"fresh trace and semantic map formula coverage differ ({len(ordinary_trace - mapped_members)} unmapped, "
            f"{len(mapped_members - ordinary_trace)} stale mapped members)")

    array_mappings = semantic_plan.get("array_mappings")
    if not isinstance(array_mappings, list):
        raise CalculationBlocked("semantic plan has no array mapping list")
    mapped_followers: dict[tuple[str, str], set[str]] = {}
    array_instances = 0
    for array_family in array_mappings:
        for instance in array_family.get("instances", []):
            array_instances += 1
            anchor = _address_key(instance.get("anchor"))
            sheet = anchor.split("!", 1)[0]
            key = _array_key(sheet, instance.get("array_ref"))
            if anchor not in anchor_trace:
                raise CalculationBlocked(f"array anchor is absent from the fresh formula trace: {anchor}")
            if anchor_trace[anchor].get("formula") != instance.get("source_formula"):
                raise CalculationBlocked(f"array anchor formula changed since semantic mapping: {anchor}")
            formula_family_id = instance.get("anchor_source_family_id")
            if formula_family_id not in mapped:
                raise CalculationBlocked(f"array anchor has no mapped formula family: {anchor}")
            members = instance.get("source_members")
            if not isinstance(members, list) or instance.get("member_count") != len(members):
                raise CalculationBlocked(f"array follower count is inconsistent at {anchor}")
            normalized = {_address_key(address) for address in members}
            if len(normalized) != len(members) or anchor in normalized:
                raise CalculationBlocked(f"array followers contain duplicate members or the anchor: {anchor}")
            min_col, min_row, max_col, max_row = range_boundaries(key[1])
            for member in normalized:
                member_sheet, local = member.split("!", 1)
                col, row, _end_col, _end_row = range_boundaries(local)
                if (member_sheet != sheet or not (min_row <= row <= max_row and min_col <= col <= max_col)):
                    raise CalculationBlocked(f"mapped array follower is outside its declared source bounds: {member}")
                record = follower_records.get(member)
                if (record is None or record.get("formula") != instance.get("source_formula")
                        or _array_key(member_sheet, record.get("array_ref")) != key):
                    raise CalculationBlocked(f"mapped array follower differs from its current anchor: {member}")
            existing = mapped_followers.setdefault(key, set())
            if existing & normalized:
                raise CalculationBlocked(f"array follower appears in multiple mappings: {anchor}")
            existing.update(normalized)
    if mapped_followers != follower_trace:
        expected_count = sum(map(len, mapped_followers.values()))
        actual_count = sum(map(len, follower_trace.values()))
        raise CalculationBlocked(f"fresh trace and semantic array-follower coverage differ ({actual_count} traced, "
                                  f"{expected_count} mapped)")

    coverage = semantic_plan.get("coverage", {})
    counts = {"ordinary_formula_count": len(ordinary_trace),
              "array_follower_count": sum(map(len, follower_trace.values())),
              "active_formula_member_count": len(ordinary_trace) + sum(map(len, follower_trace.values())),
              "source_family_count": len(profiled), "semantic_family_count": len(equation_ids),
              "array_instance_count": array_instances}
    expected_counts = {
        "ordinary_formula_count": coverage.get("ordinary_formula_members_mapped"),
        "array_follower_count": coverage.get("array_formula_followers_mapped"),
        "active_formula_member_count": coverage.get("active_formula_members_mapped"),
        "source_family_count": coverage.get("source_family_count"),
        "semantic_family_count": coverage.get("semantic_family_count"),
        "array_instance_count": coverage.get("array_instance_count"),
    }
    if any(counts[key] != expected_counts[key] for key in counts):
        raise CalculationBlocked("semantic coverage counts do not match the exact current source members")
    return counts


def _literal(value: Any) -> Any:
    if isinstance(value, ExcelError):
        return {"excel_error": value.code}
    if hasattr(value, "isoformat"):
        return {"excel_datetime": value.isoformat(), "date_only": type(value).__name__ == "date"}
    return value


def _emit_node(
    node: Any,
    *,
    sheet: str,
    owner: str,
    resolve_reference: Callable[[CellRange, str, str], str],
    resolve_name: Callable[[str, str, str], str],
) -> str:
    kind = node[0]
    if kind == "literal":
        return f"runtime.literal({_literal(node[1])!r})"
    if kind == "reference":
        reference = node[1]
        if not isinstance(reference, CellRange):
            raise CalculationBlocked(f"invalid source reference at {owner}")
        return resolve_reference(reference, sheet, owner)
    if kind == "name":
        try:
            return resolve_name(str(node[1]), sheet, owner)
        except CalculationBlocked:
            return f"runtime.unsupported_name({str(node[1])!r})"
    if kind == "unsupported_reference":
        return f"runtime.unsupported_reference({str(node[1])!r})"
    if kind == "range":
        left = _emit_node(node[1], sheet=sheet, owner=owner,
                          resolve_reference=resolve_reference, resolve_name=resolve_name)
        right = _emit_node(node[2], sheet=sheet, owner=owner,
                           resolve_reference=resolve_reference, resolve_name=resolve_name)
        return f"runtime.range_between(lambda: {left}, lambda: {right})"
    if kind == "unary":
        expression = _emit_node(node[2], sheet=sheet, owner=owner,
                                 resolve_reference=resolve_reference, resolve_name=resolve_name)
        return f"runtime.unary({node[1]!r}, lambda: {expression})"
    if kind == "binary":
        left = _emit_node(node[2], sheet=sheet, owner=owner,
                          resolve_reference=resolve_reference, resolve_name=resolve_name)
        right = _emit_node(node[3], sheet=sheet, owner=owner,
                           resolve_reference=resolve_reference, resolve_name=resolve_name)
        return f"runtime.binary({node[1]!r}, lambda: {left}, lambda: {right})"
    if kind == "call":
        name = str(node[1]).upper()
        if name.startswith("_XLFN."):
            name = name.removeprefix("_XLFN.")
        specification = _CALLS.get(name)
        if specification is None:
            raise CalculationBlocked(f"active function is not supported at {owner}: {name}")
        helper, minimum, maximum = specification
        args = node[2]
        if not minimum <= len(args) <= maximum:
            raise CalculationBlocked(f"invalid active {name} argument count at {owner}")
        emitted = [
            _emit_node(arg, sheet=sheet, owner=owner,
                       resolve_reference=resolve_reference, resolve_name=resolve_name)
            for arg in args
        ]
        if name == "IF" and len(emitted) == 2:
            emitted.append("runtime.literal(False)")
        if name == "INDIRECT":
            return (f"runtime.indirect(lambda: {emitted[0]}, "
                    "lambda target: runtime.lookup_name(values, target))")
        if name == "OFFSET":
            call_args = ", ".join(f"lambda: {arg}" for arg in emitted)
            return f"runtime.offset({call_args})"
        call_args = ", ".join(f"lambda: {arg}" for arg in emitted)
        return f"runtime.{helper}({call_args})"
    raise CalculationBlocked(f"formula node is not supported at {owner}: {kind}")


def compile_formula_function(
    function_name: str,
    formula: str,
    sheet: str,
    owner: str,
    resolve_reference: Callable[[CellRange, str, str], str],
    resolve_name: Callable[[str, str, str], str],
    index_parameters: tuple[str, ...] = (),
) -> str:
    """Return Python source for a direct formula expression over named variables.

    The callbacks must lower source coordinates to logical variable/index access. They
    are deliberately outside the generated pricing code, where workbook addresses
    are retained only for input and reconciliation adapters.
    """
    if not _FUNCTION_NAME.fullmatch(function_name):
        raise ValueError(f"invalid generated function name: {function_name!r}")
    if any(not _FUNCTION_NAME.fullmatch(name) for name in index_parameters):
        raise ValueError("generated index parameter is not a Python identifier")
    expression = _emit_node(_Parser(formula, sheet).parse(), sheet=sheet, owner=owner,
                            resolve_reference=resolve_reference, resolve_name=resolve_name)
    parameters = ", ".join(("values", *index_parameters))
    # A worksheet formula returns values, even when its selected branch is a
    # single-cell reference. Keep range references intact as lists for the
    # small array operations this runtime supports.
    return f"def {function_name}({parameters}):\n    return runtime.reference_value({expression})\n"
