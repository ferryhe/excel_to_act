"""Stage 3 design handoff built from checked static Step 3 artifacts only."""

from __future__ import annotations

import json
import math
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from excel_to_act.steps.conversion_workflow import (
    _delegation_state,
    append_stage_artifact,
    display_report_axes,
    display_report_value,
    hash_file,
    load_workflow,
    read_json,
    require_stage_approved,
)
from excel_to_act.steps.step3.workflow import _load_context, _read_stage
from excel_to_act.steps.step3.input_boundary import confirmed_boundary_reference, _require_bound_workflow
from excel_to_act.steps.step3.calculation import _parse_reference


_STATIC_ARTIFACTS = {
    "analysis": "analysis.json",
    "fields": "fields.json",
    "dependencies": "dependencies.json",
    "plan": "execution_plan.json",
    "validation": "validation.json",
}


def create_design_report(analysis_dir: Path, design_path: Path, workflow_dir: Path) -> dict[str, Any]:
    try:
        workflow_root, workflow = load_workflow(workflow_dir)
        _manifest, stage2 = require_stage_approved(workflow_root, 2)
        analysis_dir = analysis_dir.expanduser().resolve()
        design_path = design_path.expanduser().resolve()
        design = read_json(design_path)
        if not isinstance(design, dict) or design.get("schema_version") != "step3.analysis_design.input.v1":
            raise ValueError("design input must use schema_version step3.analysis_design.input.v1")
        context_dir, analysis, binding, step1_root, _inventory_path = _load_context(analysis_dir)
        boundary_reference = None
        boundary_json = None
        boundary_md = None
        if binding.get("workflow_path"):
            _require_bound_workflow(analysis, binding, workflow_root, workflow)
            boundary_reference, boundary_json, boundary_md = confirmed_boundary_reference(workflow_root, workflow)
        source = workflow["source"]
        source_binding = binding["source"]
        checked = {"source_id": source_binding.get("source_id"), "run_id": source_binding.get("run_id"),
                   "workbook_sha256": source_binding.get("source_sha256"),
                   "binding_sha256": analysis.get("binding_sha256")}
        declared_source = design.get("source")
        if not isinstance(declared_source, dict) or any(declared_source.get(key) != value for key, value in checked.items()):
            raise ValueError("design source identity must match the current Step 3 source and binding")
        if any(source.get(key) != checked.get(key) for key in ("source_id", "run_id", "workbook_sha256")):
            raise ValueError("Step 3 binding does not match the reviewed Step 1/2 source")
        source_copy = Path(source.get("step1_source_copy_path", "")).resolve()
        if not source_copy.is_file() or hash_file(source_copy) != source["workbook_sha256"]:
            raise ValueError("source-bound Step 3 workbook is missing or changed")
        if not source_copy.is_relative_to(Path(source["step1_run_path"]).resolve()):
            raise ValueError("Step 1 source copy is outside the reviewed finalized run")
        _validate_design_content(design)
        scenario_source_validation = _check_scenario_source_literals(
            design, source_copy, source["workbook_sha256"]
        )
        artifacts: dict[str, dict[str, str]] = {}
        declared_artifacts = design.get("analysis", {}).get("artifacts", {})
        for name, filename in _STATIC_ARTIFACTS.items():
            path = context_dir / filename
            if not path.is_file():
                raise ValueError(f"required static Step 3 artifact is missing: {filename}")
            digest = hash_file(path)
            declared = declared_artifacts.get(name, {})
            if not isinstance(declared, dict) or declared.get("sha256") != digest:
                raise ValueError(f"design input does not match current {filename} hash")
            if name in {"fields", "dependencies", "plan"} and analysis.get("stage_hashes", {}).get(name) != digest:
                raise ValueError(f"{filename} changed after the static analysis was validated")
            artifacts[name] = {"path": str(path), "sha256": digest}
        semantic_names = ("semantic_plan", "semantic_check")
        declared_semantic = [name for name in semantic_names if name in declared_artifacts]
        if design.get("semantic_mapping_required") is True and set(declared_semantic) != set(semantic_names):
            raise ValueError("semantic_mapping_required needs both semantic_plan and semantic_check artifact references")
        if declared_semantic and set(declared_semantic) != set(semantic_names):
            raise ValueError("semantic_plan and semantic_check must be declared together")
        semantic_values: dict[str, dict[str, Any]] = {}
        semantic_inputs: dict[str, Path] = {}
        static_semantic: dict[str, Any] | None = None
        for name in declared_semantic:
            value, digest = _read_stage(context_dir, analysis, name)
            json_path = context_dir / f"{name}.json"
            md_path = context_dir / f"{name}.md"
            declared = declared_artifacts[name]
            if not isinstance(declared, dict):
                raise ValueError(f"design input {name} reference must be an object")
            md_digest = hash_file(md_path) if md_path.is_file() else None
            expected_md = declared.get("md_sha256", declared.get("report_md_sha256"))
            if declared.get("sha256") != digest or not _declared_path_matches(declared.get("path"), json_path):
                raise ValueError(f"design input does not match current {name}.json hash or path")
            if (not md_path.is_file() or md_digest != value.get("report_md_sha256")
                    or expected_md != md_digest or not _declared_path_matches(declared.get("md_path"), md_path)):
                raise ValueError(f"design input does not match current {name}.md hash")
            semantic_values[name] = value
            artifacts[name] = {"path": str(json_path), "sha256": digest,
                               "md_path": str(md_path), "md_sha256": str(md_digest)}
            semantic_inputs[f"step3_{name}"] = json_path
            semantic_inputs[f"step3_{name}_markdown"] = md_path
        if declared_semantic:
            semantic_plan, semantic_check = semantic_values["semantic_plan"], semantic_values["semantic_check"]
            if semantic_plan.get("status") != "pass" or semantic_check.get("status") != "pass":
                raise ValueError("semantic plan and check must both have pass status")
            if semantic_plan.get("selection_basis") != semantic_check.get("selection_basis"):
                raise ValueError("semantic plan and check selection bases differ")
            coverage = semantic_plan.get("coverage", {})
            candidate = semantic_plan.get("selection_basis") == "static_candidate"
            complete_key = "unmapped_static_candidate_members" if candidate else "unmapped_active_formula_members"
            if (semantic_check.get("semantic_plan_sha256") != artifacts["semantic_plan"]["sha256"]
                    or coverage.get(complete_key) != 0):
                raise ValueError("semantic check is stale or its declared formula set is incompletely mapped")
            if candidate and (semantic_plan.get("readiness", {}).get("generation_ready") is not False
                              or semantic_plan.get("coverage", {}).get("candidate_closure") != "unknown"):
                raise ValueError("static candidate semantics cannot establish closure or release generation")
            static_semantic = _static_semantic_summary(semantic_plan, semantic_check)
            semantic_inputs.update(_semantic_evidence_inputs(context_dir, semantic_plan, semantic_check))
        validation = read_json(context_dir / "validation.json")
        if validation.get("status") != "pass" or validation.get("binding_sha256") != analysis["binding_sha256"]:
            raise ValueError("current Step 3 static validation is not a source-bound pass")
        for name in ("fields", "dependencies", "plan"):
            if validation.get(f"{name}_sha256") != artifacts[name]["sha256"]:
                raise ValueError(f"Step 3 validation does not match current {name} artifact")
        historical_inputs: dict[str, Path] = {}
        for index, evidence in enumerate(design.get("historical_evidence", [])):
            if not isinstance(evidence, dict) or not isinstance(evidence.get("path"), str):
                raise ValueError("each historical_evidence record needs a path and SHA-256")
            evidence_path = Path(evidence["path"]).expanduser().resolve()
            if not evidence_path.is_file() or hash_file(evidence_path) != evidence.get("sha256"):
                raise ValueError(f"historical evidence is missing or changed: {evidence['path']}")
            historical_inputs[f"historical_evidence_{index + 1}"] = evidence_path
        design["analysis"] = {"path": str(context_dir), "artifacts": artifacts}
        design["source"]["workbook_path"] = source["workbook_path"]
        design["source"]["workbook_sha256"] = source["workbook_sha256"]
        payload = {"schema_version": "step3.analysis_design.v1", "tool": "step3.report",
                   "status": "ready_for_review", "source": source,
                   "binding_sha256": analysis["binding_sha256"], "analysis": design["analysis"],
                   "scenario_source_validation": scenario_source_validation,
                   "static_semantic": static_semantic,
                   "design": {key: value for key, value in design.items() if key not in {"schema_version", "source", "analysis"}},
                   "review_policy": _review_policy(workflow_root, workflow),
                   "evidence_policy": {"static_artifacts_only": True, "formula_calculation_performed": False,
                                       "native_excel_called": False, "formula_cache_used": False}}
        if boundary_reference is not None:
            payload["input_boundary_reference"] = boundary_reference
        markdown = _design_markdown(
            payload,
            boundary_link=(
                (Path("..") / boundary_json.parent.name / boundary_json.name).as_posix()
                if boundary_json is not None else None
            ),
        )
        input_files = {"design_input": design_path, "source_workbook": Path(source["workbook_path"]),
                       "step1_source_copy": source_copy,
                       **historical_inputs,
                       **{f"step3_{key}": Path(ref["path"]) for key, ref in artifacts.items()},
                       **semantic_inputs}
        if boundary_json is not None and boundary_md is not None:
            input_files["confirmed_input_boundary_json"] = boundary_json
            input_files["confirmed_input_boundary_markdown"] = boundary_md
        revision = append_stage_artifact(workflow_root, 3, "analysis_design.json", payload, markdown,
                                         input_files=input_files,
                                         input_stages={2: stage2["artifact"]})
        return {"schema_version": payload["schema_version"], "tool": payload["tool"], "status": "pass",
                "workflow": str(workflow_root), "revision": revision["revision"], "artifact": revision["artifact"],
                "message": _review_policy_sentence(payload["review_policy"])}
    except Exception as exc:
        return {"schema_version": "step3.analysis_design.v1", "tool": "step3.report", "status": "blocked",
                "reason": str(exc)}


def _validate_design_content(design: dict[str, Any]) -> None:
    targets = design.get("targets")
    if not isinstance(targets, list) or not targets:
        raise ValueError("design must declare at least one target")
    target_ids: set[str] = set()
    orders: list[int] = []
    for target in targets:
        if not isinstance(target, dict) or not all(key in target for key in ("target_id", "selector", "result_order", "shape")):
            raise ValueError("each target needs target_id, selector, result_order, and shape")
        target_id = target["target_id"]
        if not isinstance(target_id, str) or not target_id or target_id in target_ids:
            raise ValueError("target_id values must be unique non-empty strings")
        if not isinstance(target["result_order"], int) or isinstance(target["result_order"], bool):
            raise ValueError("target result_order must be a zero-based integer")
        if not isinstance(target["shape"], list) or any(not isinstance(size, int) or size < 0 for size in target["shape"]):
            raise ValueError("target shape must be a list of non-negative dimensions; [] represents a scalar")
        target_ids.add(target_id)
        orders.append(target["result_order"])
    if sorted(orders) != list(range(len(targets))):
        raise ValueError("target result_order must be a stable zero-based sequence")
    scenarios = design.get("scenarios")
    if not isinstance(scenarios, list) or not scenarios:
        raise ValueError("design must declare at least one scenario")
    for scenario in scenarios:
        if not isinstance(scenario, dict) or not scenario.get("scenario_id") or not isinstance(scenario.get("primary_inputs"), dict):
            raise ValueError("each scenario needs an ID and primary_inputs object")
        for selector, declaration in scenario["primary_inputs"].items():
            if not isinstance(selector, str) or not selector:
                raise ValueError("scenario primary_input selectors must be non-empty strings")
            if not isinstance(declaration, dict):
                continue
            has_cell = "source_cell" in declaration
            has_literal = "source_literal" in declaration
            if has_cell != has_literal:
                raise ValueError(
                    f"scenario {scenario['scenario_id']} selector {selector} must declare both source_cell and source_literal"
                )
            if has_cell:
                if not isinstance(declaration["source_cell"], str) or not declaration["source_cell"].strip():
                    raise ValueError(f"scenario {scenario['scenario_id']} selector {selector} source_cell must be text")
                _require_json_scalar_literal(
                    declaration["source_literal"],
                    f"scenario {scenario['scenario_id']} selector {selector} source_literal",
                )
        if not isinstance(scenario.get("input_overrides", []), list):
            raise ValueError("scenario input_overrides must be a list")
    paths = design.get("paths", [])
    if not isinstance(paths, list) or any(not isinstance(item, dict) or not item.get("path_id") for item in paths):
        raise ValueError("paths must be a list of records with path_id")
    path_ids = [item["path_id"] for item in paths]
    if len(path_ids) != len(set(path_ids)):
        raise ValueError("path_id values must be unique")
    if any(item.get("target_id") not in target_ids or not item.get("scenario_id")
           or not isinstance(item.get("ordered_steps"), list) for item in paths):
        raise ValueError("each path needs a known target, scenario ID, and ordered_steps list")
    known = set(design.get("coverage", {}).get("known_path_ids", [])) if isinstance(design.get("coverage"), dict) else set()
    if set(path_ids) != known:
        raise ValueError("coverage.known_path_ids must exactly match the registered paths")
    for stage in ("documented_path_ids", "resolved_path_ids", "planned_path_ids", "implemented_path_ids", "runtime_verified_path_ids"):
        values = design.get("coverage", {}).get(stage, [])
        if not isinstance(values, list) or not set(values) <= known:
            raise ValueError(f"coverage.{stage} must cite IDs in coverage.known_path_ids")
    if any(not isinstance(design.get(key, []), list) for key in ("options", "open_questions", "validation_plan")):
        raise ValueError("options, open_questions, and validation_plan must be lists")
    for key in ("field_groups", "shared_modules"):
        if not isinstance(design.get(key, []), list):
            raise ValueError(f"{key} must be a list")
    if "tool_execution_ledger" in design:
        ledger = design["tool_execution_ledger"]
        if not isinstance(ledger, list):
            raise ValueError("tool_execution_ledger must be a list of descriptive records")
        required_ledger_keys = {"tool", "command", "purpose", "outputs", "status"}
        for record in ledger:
            if not isinstance(record, dict) or set(record) != required_ledger_keys:
                raise ValueError("each tool_execution_ledger record must contain only tool, command, purpose, outputs, and status")
            if any(not isinstance(record.get(key), str) or not record[key].strip()
                   for key in ("tool", "command", "purpose", "status")):
                raise ValueError("tool_execution_ledger tool, command, purpose, and status must be non-empty text")
            if (not isinstance(record.get("outputs"), list) or not record["outputs"]
                    or any(not isinstance(item, str) or not item.strip() for item in record["outputs"])):
                raise ValueError("tool_execution_ledger outputs must be a non-empty list of text")
    if "external_value_capture_plan" in design:
        capture = design["external_value_capture_plan"]
        required_capture_keys = {"phase", "ranges", "method", "binding", "validation",
                                 "current_values_available", "formula_cache_allowed"}
        optional_capture_keys = {"permitted_origins", "upstream_python_translation"}
        if (not isinstance(capture, dict) or required_capture_keys - set(capture)
                or set(capture) - required_capture_keys - optional_capture_keys):
            raise ValueError("external_value_capture_plan must use the documented capture-plan fields")
        if any(not isinstance(capture.get(key), str) or not capture[key].strip()
               for key in ("phase", "method", "binding", "validation")):
            raise ValueError("external_value_capture_plan phase, method, binding, and validation must be non-empty text")
        if (not isinstance(capture.get("ranges"), list) or not capture["ranges"]
                or any(not isinstance(item, str) or not item.strip() for item in capture["ranges"])):
            raise ValueError("external_value_capture_plan ranges must be a non-empty list of text")
        if any(not isinstance(capture.get(key), bool)
               for key in ("current_values_available", "formula_cache_allowed")):
            raise ValueError("external_value_capture_plan availability and cache flags must be booleans")
        if capture["current_values_available"] or capture["formula_cache_allowed"]:
            raise ValueError("Stage 3 external values must remain unavailable and formula caches prohibited")
        if ("permitted_origins" in capture
                and (not isinstance(capture["permitted_origins"], list)
                     or any(not isinstance(item, str) or not item.strip() for item in capture["permitted_origins"]))):
            raise ValueError("external_value_capture_plan permitted_origins must be a list of text")
        if ("upstream_python_translation" in capture
                and not isinstance(capture["upstream_python_translation"], bool)):
            raise ValueError("external_value_capture_plan upstream_python_translation must be a boolean")
    if any(key in design for key in ("native_oracle", "runtime_reconciliation", "computed_values")):
        raise ValueError("Stage 3 design cannot include Stage 5 runtime or native-oracle results")


def _require_json_scalar_literal(value: Any, label: str) -> None:
    if value is None or isinstance(value, (str, bool)):
        return
    if isinstance(value, int) and not isinstance(value, bool):
        return
    if isinstance(value, float) and math.isfinite(value):
        return
    raise ValueError(f"{label} must be a finite JSON scalar literal")


def _check_scenario_source_literals(design: dict[str, Any], workbook_path: Path,
                                   expected_source_sha256: str) -> dict[str, Any]:
    """Check only explicitly authored literal assertions against formula-mode source cells."""
    assertions = [
        (scenario["scenario_id"], selector, declaration["source_cell"], declaration["source_literal"])
        for scenario in design.get("scenarios", [])
        for selector, declaration in scenario.get("primary_inputs", {}).items()
        if isinstance(declaration, dict) and "source_cell" in declaration and "source_literal" in declaration
    ]
    if hash_file(workbook_path) != expected_source_sha256:
        raise ValueError("source-bound Step 3 workbook changed before scenario literal checks")
    workbook = load_workbook(workbook_path, data_only=False, read_only=False,
                             keep_vba=workbook_path.suffix.casefold() == ".xlsm")
    checked: list[dict[str, Any]] = []
    try:
        sheet_by_casefold = {name.casefold(): name for name in workbook.sheetnames}
        for scenario_id, selector, source_cell, expected in assertions:
            parsed = _parse_reference(source_cell, "")
            if "!" not in source_cell or parsed is None or not parsed.single:
                raise ValueError(
                    f"scenario {scenario_id} selector {selector} source_cell {source_cell!r} "
                    "must be a qualified single-cell address"
                )
            sheet_name = sheet_by_casefold.get(parsed.sheet.casefold())
            if sheet_name is None:
                raise ValueError(
                    f"scenario {scenario_id} selector {selector} source literal mismatch at {source_cell}: "
                    f"expected {expected!r}, actual <missing sheet>"
                )
            worksheet = workbook[sheet_name]
            cell = worksheet._cells.get((parsed.min_row, parsed.min_col))
            if cell is None:
                raise ValueError(
                    f"scenario {scenario_id} selector {selector} source literal mismatch at {source_cell}: "
                    f"expected {expected!r}, actual <missing physical cell>"
                )
            if cell.data_type == "f":
                raise ValueError(
                    f"scenario {scenario_id} selector {selector} source_cell {source_cell} is a formula, "
                    f"not a source literal (formula {cell.value!r})"
                )
            actual = cell.value
            if not _same_source_literal(expected, actual):
                raise ValueError(
                    f"scenario {scenario_id} selector {selector} source literal mismatch at {source_cell}: "
                    f"expected {expected!r}, actual {actual!r}"
                )
            checked.append({"scenario_id": scenario_id, "selector": selector,
                            "source_cell": source_cell, "source_literal": expected,
                            "formula_free": True})
    finally:
        workbook.close()
    if hash_file(workbook_path) != expected_source_sha256:
        raise ValueError("source-bound Step 3 workbook changed during scenario literal checks")
    return {"status": "pass", "source_sha256": expected_source_sha256,
            "assertion_count": len(checked), "assertions": checked,
            "formula_cache_used": False}


def _same_source_literal(expected: Any, actual: Any) -> bool:
    if expected is None or isinstance(expected, (str, bool)):
        return type(expected) is type(actual) and expected == actual
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and not isinstance(actual, bool) and expected == actual
    return False


def _design_evidence_markdown(value: dict[str, Any]) -> str:
    design = value["design"]
    lines = ["# Step 3 Analysis and Design", "", "Status: **Ready for review**", "",
             f"Source run: `{value['source']['run_id']}` · `{value['source']['source_id']}`",
             f"Workbook SHA-256: `{value['source']['workbook_sha256']}`",
             f"Step 3 binding SHA-256: `{value['binding_sha256']}`", "", "## Targets", ""]
    for target in sorted(design["targets"], key=lambda item: item["result_order"]):
        lines.append(
            f"- **{target['target_id']}** · `{target['selector']}` · order {target['result_order']} · "
            f"shape {display_report_value(target.get('shape'))} · units {display_report_value(target.get('units'))} · "
            f"axes {display_report_axes(target)}"
        )
        if target.get("location"):
            lines.append(f"  - Source location: `{target['location']}`")
    coverage = design.get("coverage", {})
    known_ids = coverage.get("known_path_ids")
    known_count = len(known_ids) if isinstance(known_ids, list) else None
    lines.extend(["", "## Known-path coverage", "",
                  f"Known-path denominator: {known_count if known_count is not None else 'Not recorded'} registered path units; all ratios are known-only.",
                  "Documented means each known route has a written ledger entry, including explicit unknowns; it does not mean the route is semantically resolved."])
    for label, key in (("Documented ledger", "documented_path_ids"), ("Resolved", "resolved_path_ids"),
                       ("Planned", "planned_path_ids"), ("Implemented", "implemented_path_ids"),
                       ("Runtime verified", "runtime_verified_path_ids")):
        values = coverage.get(key)
        numerator = len(values) if isinstance(values, list) else None
        lines.append(f"- {label}: {numerator if numerator is not None else 'Not recorded'} / "
                     f"{known_count if known_count is not None else 'Not recorded'} known path units")
    lines.extend(["", "## Scenario and scope", ""])
    for scenario in design.get("scenarios", []):
        lines.append(f"- **{scenario.get('scenario_id')}** · source SHA-256 `{scenario.get('source_sha256')}`")
        lines.append(f"  - Primary inputs: `{json.dumps(scenario.get('primary_inputs', {}), ensure_ascii=False, sort_keys=True)}`")
        lines.append(f"  - Input overrides: `{json.dumps(scenario.get('input_overrides', []), ensure_ascii=False)}`")
    literal_validation = value.get("scenario_source_validation", {})
    lines.extend(["", "## Saved scenario source-literal checks", "",
                  f"Source SHA-256: `{literal_validation.get('source_sha256', 'not checked')}`",
                  f"Assertions matched: {literal_validation.get('assertion_count', 'Not recorded')}; formula-cache inputs: `false`"])
    for assertion in literal_validation.get("assertions", []):
        lines.append(
            f"- **{assertion['scenario_id']} / {assertion['selector']}** · "
            f"`{assertion['source_cell']}` = `{json.dumps(assertion['source_literal'], ensure_ascii=False)}` · "
            "formula-free source literal"
        )
    scope = design.get("scope", {})
    lines.extend(["", f"Boundary: {scope.get('boundary', 'not stated')}", "", "Included:"])
    lines.extend(f"- {item}" for item in scope.get("included", []))
    lines.extend(["", "Excluded:"])
    lines.extend(f"- {item}" for item in scope.get("excluded", []))
    addendum = scope.get("scope_addendum")
    if addendum:
        lines.extend(["", "Scope addendum:", f"- Reason: {addendum.get('reason', 'not stated')}",
                      f"- Included sheets: {', '.join(addendum.get('include_sheets', []))}",
                      f"- Selected ranges: {addendum.get('selected_ranges', 'not stated')}",
                      f"- Policy: {addendum.get('policy', 'not stated')}"])
    lines.extend(["", "## Paths", ""])
    for path in design.get("paths", []):
        lines.append(f"- **{path['path_id']}** · {path.get('status', 'unknown')} · {path.get('condition', 'No condition supplied.')}")
        lines.append(f"  - Route: {' → '.join(path.get('ordered_steps', []))}")
        lines.append(f"  - Explanation: {path.get('explanation', 'No explanation supplied.')}")
        lines.append(f"  - Source formula: `{path.get('source_formula', 'not supplied')}`")
        lines.append(f"  - Error policy: {path.get('error_policy', 'not supplied')}")
        if path.get("evidence_ids"):
            lines.append(f"  - Evidence IDs: {', '.join(path['evidence_ids'])}")
        for unknown in path.get("unknowns", []):
            lines.append(f"  - Unknown: {unknown}")
    boundary = value.get("input_boundary_reference")
    if boundary:
        lines.extend(["", "## Confirmed input boundary", "",
                      f"- Boundary SHA-256: `{boundary.get('boundary_sha256')}`",
                      f"- Catalog revision: {boundary.get('revision')}",
                      f"- Catalog JSON: `{boundary.get('artifact', {}).get('json')}` · `{boundary.get('artifact', {}).get('json_sha256')}`",
                      f"- Catalog Markdown SHA-256: `{boundary.get('artifact', {}).get('md_sha256')}`",
                      f"- Reviewer pair: {', '.join(boundary.get('review_receipts', {}).get('reviewer_pair', []))}"])
        for reviewer in ("agent", "human", "typesafe"):
            receipt = boundary.get("review_receipts", {}).get(reviewer)
            if receipt:
                lines.append(f"- {reviewer.title()} receipt: `{receipt.get('receipt_id')}`")
    lines.extend(["", "## Field groups and shared modules", ""])
    for item in design.get("field_groups", []):
        lines.append(f"- **{item.get('id', item.get('group_id', item.get('name', 'field group')))}** · {item.get('shape', 'shape pending')} · {item.get('classification', item.get('status', 'status pending'))}")
        lines.append(f"  - Meaning: {item.get('meaning', item.get('description', item.get('purpose', 'No description supplied.')))}")
        if item.get("source_extents"):
            lines.append(f"  - Source extents: {', '.join(item['source_extents'])}")
        if item.get("execution_constraint"):
            lines.append(f"  - Execution constraint: {item['execution_constraint']}")
        if item.get("semantic_grouping_status"):
            lines.append(f"  - Grouping status: {item['semantic_grouping_status']}")
    for item in design.get("shared_modules", []):
        lines.append(f"- Shared module **{item.get('module_id', item.get('name', 'module'))}** · {item.get('purpose', item.get('description', 'No description supplied.'))}")
        lines.append(f"  - Implementation: {item.get('implementation_status', 'not stated')}; verification: {item.get('verification_status', 'not stated')}")
        if item.get("missing"):
            lines.append(f"  - Missing: {'; '.join(item['missing'])}")
    lines.extend(["", "## Options", ""])
    for option in design.get("options", []):
        lines.append(f"- **{option.get('option_id', 'option')}** · {option.get('scope', 'scope not specified')} · {option.get('status', 'status pending')}")
        if option.get("description") or option.get("tradeoff"):
            lines.append(f"  - {option.get('description', option.get('tradeoff'))}")
        if option.get("selected_for_draft") is not None:
            lines.append(f"  - Selected for draft: `{option['selected_for_draft']}`")
        planned = option.get("planned_known_path_coverage")
        if isinstance(planned, dict):
            lines.append(f"  - Planned known-path coverage: {planned.get('numerator', '?')} / {planned.get('denominator', '?')} ({planned.get('percent', '?')}%); {planned.get('basis', '')}")
        if option.get("numerical_paths"):
            lines.append(f"  - Numerical routes: {', '.join(option['numerical_paths'])}")
        if option.get("condition_only_paths"):
            lines.append(f"  - Condition-only routes: {', '.join(option['condition_only_paths'])}")
        if option.get("exclusions"):
            lines.append(f"  - Exclusions: {'; '.join(option['exclusions'])}")
        if option.get("error_policy"):
            lines.append(f"  - Error policy: {option['error_policy']}")
        if option.get("required_confirmations"):
            lines.append(f"  - Required confirmations: {'; '.join(option['required_confirmations'])}")
        for milestone in option.get("cumulative_path_coverage", []):
            lines.append(f"  - Cumulative milestone: {milestone.get('percent', '?')}% known paths · {', '.join(milestone.get('paths', []))}")
        if "all_configuration_coverage" in option:
            lines.append(f"  - All-configuration coverage: `{option['all_configuration_coverage']}`")
        if option.get("uncovered_paths"):
            lines.append(f"  - Uncovered paths: {', '.join(option['uncovered_paths'])}")
        if option.get("status"):
            lines.append(f"  - Status: {option['status']}")
    lines.extend(["", "## Validation plan", ""])
    for case in design.get("validation_plan", []):
        lines.append(f"- **{case.get('case_id', 'case')}** · {case.get('status', 'planned')}: {case.get('scope', 'scope not stated')}")
        if case.get("compare"):
            lines.append(f"  - Compare: {', '.join(case['compare'])}")
    execution = design.get("execution_design", {})
    if execution:
        lines.extend(["", "## Execution design", "",
                      f"- Backend: {execution.get('backend', 'not stated')}",
                      f"- Emission: {execution.get('emission', 'not stated')}",
                      f"- Ordering: {execution.get('ordering', 'not stated')}",
                      f"- Runtime contract: {execution.get('runtime_contract', 'not stated')}",
                      f"- Raw-data-only: `{execution.get('raw_data_only', 'not stated')}`",
                      f"- Cached formula outputs allowed: `{execution.get('cached_formula_outputs_allowed', 'not stated')}`"])
        lines.extend(f"- Source behavior to preserve: {item}" for item in execution.get("preserve_formulas", []))
    ledger = design.get("tool_execution_ledger")
    if ledger:
        lines.extend(["", "## Tool execution ledger", "",
                      "These are descriptive plan or evidence records; rendering this report does not execute the listed commands.",
                      "", "| Tool | Command | Purpose | Outputs | Status |", "| --- | --- | --- | --- | --- |"])
        lines.extend(
            "| " + " | ".join((
                _markdown_table_cell(item["tool"]), _markdown_table_cell(item["command"]),
                _markdown_table_cell(item["purpose"]),
                _markdown_table_cell("<br>".join(item["outputs"])),
                _markdown_table_cell(item["status"]),
            )) + " |"
            for item in ledger
        )
    capture = design.get("external_value_capture_plan")
    if capture:
        lines.extend(["", "## External-value capture prerequisite", "",
                      "This is a future-stage plan only. Stage 3 does not capture or include these values.",
                      f"- Phase: {capture['phase']}",
                      f"- Method: {capture['method']}",
                      f"- Binding: {capture['binding']}",
                      f"- Validation: {capture['validation']}",
                      f"- Current values available: `{str(capture['current_values_available']).lower()}`",
                      f"- Formula cache allowed: `{str(capture['formula_cache_allowed']).lower()}`",
                      "- Ranges:"])
        lines.extend(f"  - `{_markdown_table_cell(item)}`" for item in capture["ranges"])
        if "permitted_origins" in capture:
            lines.append(f"- Permitted origins: {', '.join(_markdown_table_cell(item) for item in capture['permitted_origins'])}")
        if "upstream_python_translation" in capture:
            lines.append(f"- Upstream Python translation: `{str(capture['upstream_python_translation']).lower()}`")
    historical = design.get("historical_evidence", [])
    if historical:
        lines.extend(["", "## Historical evidence", "",
                      "These hashes support source-bound path selection; they do not prove the new generated code or its run result."])
        lines.extend(f"- `{item['path']}` · SHA-256 `{item['sha256']}`" for item in historical)
    artifacts = value.get("analysis", {}).get("artifacts", {})
    if "semantic_plan" in artifacts and "semantic_check" in artifacts:
        plan = read_json(Path(artifacts["semantic_plan"]["path"]))
        coverage = plan.get("coverage", {})
        static_semantic = value.get("static_semantic")
        if not isinstance(static_semantic, dict):
            static_semantic = _static_semantic_summary(plan, plan)
        if plan.get("selection_basis") == "static_candidate":
            coverage_lines = [
                "- Selection basis: `static_candidate`; candidate closure remains **unknown**.",
                f"- Candidate formula positions mapped: {coverage.get('static_candidate_formula_members_mapped', '?')} / {coverage.get('static_candidate_formula_members', '?')}",
                f"- Candidate array followers mapped: {coverage.get('static_candidate_array_followers_mapped', '?')} / {coverage.get('static_candidate_array_followers', '?')}",
                f"- Candidate members mapped: {coverage.get('static_candidate_member_count_mapped', '?')} / {coverage.get('static_candidate_member_count', '?')}",
                "- This is exact source-member bookkeeping against a conservative candidate set; it does not prove active selection, candidate closure, or runtime behavior.",
            ]
            if plan.get("candidate_scope") is not None:
                coverage_lines.extend([
                    f"- Candidate scope: `{plan['candidate_scope']}`; frontiers: {coverage.get('scope_frontier_count', '?')}; proposed starts: {coverage.get('proposed_formula_start_count', '?')}.",
                    "- The scope fence is a proposed coordinate slice; omitted references remain unknown and are not treated as unused.",
                ])
        else:
            coverage_lines = [
                "- Selection basis: `active_trace`.",
                f"- Active formula members mapped: {coverage.get('active_formula_members_mapped', '?')} / {coverage.get('active_formula_members', '?')}",
                f"- Ordinary formula members: {coverage.get('ordinary_formula_members_mapped', '?')} / {coverage.get('ordinary_formula_members', '?')}",
                f"- Array followers: {coverage.get('array_formula_followers_mapped', '?')} / {coverage.get('array_formula_followers', '?')}",
                "- This is exact source-member bookkeeping against bound active-trace evidence; mapped meaning and equations still require review.",
            ]
        lines.extend(["", "## Static semantic mapping", "", *coverage_lines,
                      f"- Equation families: {coverage.get('semantic_family_count', '?')}; array families: {coverage.get('array_family_count', '?')}",
                      f"- Plan: `{artifacts['semantic_plan']['path']}` · SHA-256 `{artifacts['semantic_plan']['sha256']}`",
                      f"- Check: `{artifacts['semantic_check']['path']}` · SHA-256 `{artifacts['semantic_check']['sha256']}`"])
        external = plan.get("external_boundary_variables", [])
        if external:
            lines.append(f"- Formula-derived external boundary variables retained without values: {len(external)}.")
        classification = static_semantic["source_classification"]
        lines.extend(["", "Source role and kind counts:", "", "| Role | Kind | Variables |", "| --- | --- | ---: |"])
        lines.extend(f"| {row['role']} | {row['kind']} | {row['count']} |"
                     for row in static_semantic["role_kind_counts"])
        lines.extend(["", "Raw source extent checks:",
                      f"- Extents: {classification['raw_extent_count']}; distinct coordinates: {classification['distinct_raw_extent_coordinate_count']}",
                      f"- Raw seed variables/extents: {classification['raw_seed_variable_count']} / {classification['raw_seed_extent_count']}",
                      f"- Demanded formula overlaps: {classification['demanded_raw_formula_overlaps']}",
                      f"- Demanded array-follower overlaps: {classification['demanded_raw_array_follower_overlaps']}",
                      f"- Formula/follower coordinates outside raw demand: {classification['formula_coordinates_outside_raw_demand']} / {classification['array_follower_coordinates_outside_raw_demand']}",
                      f"- Evidence basis: {classification['source_kind_basis']}"])
    lines.extend(["", "## Open questions", ""])
    questions = design.get("open_questions", [])
    if questions:
        for item in questions:
            if isinstance(item, dict):
                question_id = item.get("question_id", "Not recorded")
                question_text = item.get("question", "Not recorded")
                disposition = item.get("disposition", item.get("status", "Not recorded"))
                lines.append(f"- **{question_id}** — {question_text} Recorded disposition: {str(disposition).rstrip('.')}.")
            else:
                lines.append(f"- {item}")
    else:
        lines.append("- None recorded.")
    lines.extend(["", "## Evidence", "", "This Step 3 command packages static analysis and design evidence. Declared historical evidence is reused by hash; this command performs no new formula calculation, native Excel rebuild, or formula-cache read.",
                  _review_policy_sentence(value.get("review_policy", {})), ""])
    return "\n".join(lines)


def _design_markdown(
    value: dict[str, Any], *, machine_link: str = "analysis_design.json",
    workflow_link: str = "../../workflow.json", boundary_link: str | None = None,
) -> str:
    design = value.get("design", {})
    coverage = design.get("coverage", {})
    known_ids = coverage.get("known_path_ids", [])
    targets = sorted(design.get("targets", []), key=lambda item: item.get("result_order", 0))
    scope = design.get("scope", {})
    principal_count = coverage.get("principal_design_coverage")
    numerical_ids = coverage.get("principal_numerical_path_ids", [])
    numerical_id_set = set(numerical_ids)
    condition_ids = [
        item.get("path_id") for item in design.get("paths", [])
        if isinstance(item, dict) and item.get("path_id") not in numerical_id_set
    ]
    lines = [
        "# Step 3 Analysis and Design", "",
        "## What this report asks you to accept", "",
        "The selected target, source-bound scope, input assumptions, path and module design, and planned validation. "
        "This is design evidence; it does not calculate formulas, generate code, or establish Excel equivalence.", "",
        "## Purpose and scope", "",
        f"Source run: `{display_report_value(value.get('source', {}).get('run_id'))}`.",
        f"Boundary: {display_report_value(scope.get('boundary'))}", "",
        "## Requested target", "",
    ]
    for target in targets:
        units_text = display_report_value(target.get("units")).rstrip(".")
        lines.append(
            f"- **{display_report_value(target.get('selector'))}** — "
            f"{display_report_value(target.get('result_kind'))}; "
            f"shape {display_report_value(target.get('shape'))}; units {units_text}; "
            f"axes {display_report_axes(target)}."
        )
        if target.get("location"):
            lines.append(f"  - Source location: `{target['location']}` (provenance coordinate).")
        for name in target.get("diagnostic_intermediates", []):
            lines.append(f"  - Diagnostic output: `{display_report_value(name)}`; not a requested business target.")
    field_groups = design.get("field_groups", [])
    modules = design.get("shared_modules", [])
    lines.extend(["", "## Field groups and calculation design", ""])
    if field_groups:
        lines.extend(["| Field group | Recorded shape | Purpose and source | Grouping status |",
                      "| --- | --- | --- | --- |"])
        for item in field_groups:
            group_id = display_report_value(item.get("id", item.get("group_id")))
            purpose = display_report_value(item.get("meaning", item.get("description", item.get("purpose"))))
            status = display_report_value(item.get("semantic_grouping_status", item.get("execution_constraint")))
            lines.append(
                f"| {_markdown_table_cell(group_id)} | "
                f"{_markdown_table_cell(display_report_value(item.get('shape')))} | "
                f"{_markdown_table_cell(purpose)} | {_markdown_table_cell(status)} |"
            )
    else:
        lines.append("Field groups: Not recorded.")
    execution = design.get("execution_design", {})
    ordering = execution.get("ordering") if isinstance(execution, dict) else None
    lines.append(f"\nRecorded calculation order: {display_report_value(ordering).rstrip('.')}.")
    if modules:
        lines.extend(["", "| Order | Module | Purpose | Implementation and verification |",
                      "| ---: | --- | --- | --- |"])
        for item in sorted(modules, key=lambda module: (
            module.get("order") if isinstance(module.get("order"), int) else 10**9,
            str(module.get("module_id", "")),
        )):
            order = display_report_value(item.get("order"))
            implementation = display_report_value(item.get("implementation_status"))
            verification = display_report_value(item.get("verification_status"))
            missing = item.get("missing", [])
            details = f"{implementation}; {verification}"
            if missing:
                details += "; missing: " + "; ".join(display_report_value(entry) for entry in missing)
            lines.append(
                f"| {order} | {_markdown_table_cell(display_report_value(item.get('module_id')))} | "
                f"{_markdown_table_cell(display_report_value(item.get('purpose')))} | "
                f"{_markdown_table_cell(details)} |"
            )
    else:
        lines.append("Shared modules: Not recorded.")
    preserve_formulas = execution.get("preserve_formulas", []) if isinstance(execution, dict) else []
    lines.extend(["", "Recorded recurrence and time-alignment constraints:"])
    front_constraints = [
        item for item in preserve_formulas
        if isinstance(item, str)
        and ("recurrence" in item.casefold() or "same-year" in item.casefold() or "t9 year-zero alignment" in item.casefold())
    ]
    if front_constraints:
        lines.extend(f"- {item}" for item in front_constraints)
    else:
        lines.append("- Not recorded.")
    lines.extend(["", "## Verified design coverage", ""])
    if isinstance(principal_count, dict):
        numerator, denominator = principal_count.get("numerator"), principal_count.get("denominator")
        lines.append(
            f"Principal numerical routes recorded: {numerator if numerator is not None else 'Not recorded'} / "
            f"{denominator if denominator is not None else 'Not recorded'}; these are plan entries, not executed checks."
        )
    else:
        lines.append(f"Known design-ledger units: {len(known_ids) if isinstance(known_ids, list) else 'Not recorded'}.")
    lines.append(f"Numerical route IDs: {', '.join(numerical_ids) if numerical_ids else 'Not recorded'}.")
    lines.append(f"Condition-only route IDs: {', '.join(condition_ids) if condition_ids else 'Not recorded'}.")
    lines.extend(["", "## Limits and unresolved items", ""])
    questions = design.get("open_questions")
    if isinstance(questions, list):
        lines.append(f"Open questions recorded: {len(questions)}.")
        for item in questions:
            if isinstance(item, dict):
                question_id = display_report_value(item.get("question_id"))
                question_text = display_report_value(item.get("question"))
                disposition = display_report_value(item.get("disposition", item.get("status")))
                answer = display_report_value(item.get("answer"))
                lines.append(
                    f"- **{question_id}** — {question_text} Recorded answer: {answer}; "
                    f"recorded disposition/status: {disposition.rstrip('.')}."
                )
            else:
                lines.append(f"- {display_report_value(item)}")
    else:
        lines.append("Open questions recorded: Not recorded.")
    lines.append(
        "Coverage is limited to the known source-bound path ledger. Planned routes are not implemented or runtime verified; "
        "the condition-only feedback guard is separate from numerical implementation."
    )
    boundary = value.get("input_boundary_reference")
    lines.extend(["", "## Handover and review", "",
                  f"- Machine design: [analysis_design.json](<{machine_link}>).",
                  f"- Current acceptance ledger: [workflow.json](<{workflow_link}>)."])
    if boundary:
        link = boundary_link or boundary.get("artifact", {}).get("json", "Not recorded")
        lines.append(f"- Confirmed input-boundary JSON: [input_boundary.json](<{Path(link).as_posix()}>).")
    lines.extend([
        "- Next: the Agent reviews this exact JSON/Markdown pair first. After Agent PASS, the actual human reviews it unless a valid, explicitly scoped TypeSafe delegation applies to final design.",
        "- A question that clarifies material target, scope, input kind, axis, units, or business meaning is not a final approval. To reject, route to the responsible current or earlier Step and refresh affected evidence.",
        "- Report creation status and Agent PASS are technical/review states; formal human or delegated acceptance remains in the current ledger.", "",
        "## Detailed design evidence", "",
        re.sub(r"(?m)^## ", "#### ", _design_evidence_markdown(value).replace(
            "# Step 3 Analysis and Design", "### Original design details", 1)), "",
    ])
    return "\n".join(lines)


def _markdown_table_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\r\n", "<br>").replace("\r", "<br>").replace("\n", "<br>")


def _review_policy(workflow_root: Path, workflow_manifest: dict[str, Any]) -> dict[str, Any]:
    delegation, error = _delegation_state(workflow_root, workflow_manifest)
    if delegation is not None and error is None and 3 in delegation.get("stages", []):
        return {"reviewers": ["agent", "typesafe"], "authorization_sha256": delegation["authorization_sha256"]}
    return {"reviewers": ["agent", "human"], "delegation_active": False}


def _review_policy_sentence(policy: dict[str, Any]) -> str:
    reviewers = policy.get("reviewers", ["agent", "human"])
    if reviewers == ["agent", "typesafe"]:
        return "Agent and delegated TypeSafe confirmation are pending for this exact source-bound workflow. TypeSafe remains a separate reviewer, not a human receipt."
    return "Agent and human confirmation are pending. Record the human receipt only after an explicit user response."


def _declared_path_matches(declared: Any, current: Path) -> bool:
    if not isinstance(declared, str):
        return False
    path = Path(declared).expanduser()
    if not path.is_absolute():
        path = current.parent / path
    try:
        return path.resolve() == current.resolve()
    except (OSError, ValueError):
        return False


def _semantic_evidence_inputs(
    context_dir: Path, semantic_plan: dict[str, Any], semantic_check: dict[str, Any],
) -> dict[str, Path]:
    """Bind the semantic plan's nested evidence directly to the Stage 3 revision."""
    evidence = semantic_plan.get("evidence")
    candidate = semantic_plan.get("selection_basis") == "static_candidate"
    trace_key = "source_candidate_trace" if candidate else "active_trace"
    references = {
        "semantic_map_input": semantic_plan.get("semantic_map_input"),
        "source_family_profile": evidence.get("source_family_profile") if isinstance(evidence, dict) else None,
        trace_key: evidence.get(trace_key) if isinstance(evidence, dict) else None,
    }
    check_hash_keys = {
        "semantic_map_input": "semantic_map_input_sha256",
        "source_family_profile": "source_family_profile_sha256",
        trace_key: f"{trace_key}_sha256",
    }
    if semantic_plan.get("candidate_scope") is not None:
        references["model_scope"] = evidence.get("model_scope") if isinstance(evidence, dict) else None
        check_hash_keys["model_scope"] = "model_scope_sha256"
    result: dict[str, Path] = {}
    for name, reference in references.items():
        if not isinstance(reference, dict) or not isinstance(reference.get("path"), str):
            raise ValueError(f"semantic plan is missing its {name} evidence path and SHA-256")
        path = Path(reference["path"]).expanduser()
        if not path.is_absolute():
            path = context_dir / path
        path = path.resolve()
        digest = reference.get("sha256")
        if not path.is_file() or not isinstance(digest, str) or hash_file(path) != digest:
            raise ValueError(f"semantic plan {name} evidence is missing or changed")
        if semantic_check.get(check_hash_keys[name]) != digest:
            raise ValueError(f"semantic check does not match the semantic plan {name} evidence")
        result[f"step3_{name}"] = path
    return result


def _static_semantic_summary(semantic_plan: dict[str, Any], semantic_check: dict[str, Any]) -> dict[str, Any]:
    plan_classification = semantic_plan.get("source_classification")
    check_classification = semantic_check.get("source_classification")
    if not isinstance(plan_classification, dict) or plan_classification != check_classification:
        raise ValueError("semantic plan and check have missing or inconsistent source-kind summaries")
    variables_by_role_and_kind = plan_classification.get("variables_by_role_and_kind")
    if not isinstance(variables_by_role_and_kind, dict):
        raise ValueError("semantic source-kind summary has no role/kind counts")
    candidate_scope = semantic_plan.get("candidate_scope")
    if candidate_scope is not None:
        evidence = semantic_plan.get("evidence", {})
        scope_ref = evidence.get("model_scope") if isinstance(evidence, dict) else None
        if (candidate_scope != semantic_check.get("candidate_scope")
                or not isinstance(scope_ref, dict)
                or semantic_check.get("model_scope_sha256") != scope_ref.get("sha256")):
            raise ValueError("semantic plan and check have missing or inconsistent candidate-scope evidence")
    role_kind_counts = [
        {"role": role, "kind": kind, "count": count}
        for role, kinds in sorted(variables_by_role_and_kind.items())
        if isinstance(kinds, dict)
        for kind, count in sorted(kinds.items())
    ]
    return {"coverage": semantic_plan.get("coverage", {}),
            "source_classification": plan_classification,
            "role_kind_counts": role_kind_counts,
            **({"candidate_scope": candidate_scope,
                "model_scope": semantic_plan["evidence"]["model_scope"]}
               if candidate_scope is not None else {})}
