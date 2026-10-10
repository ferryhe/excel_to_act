"""Source-bound semantic map checks for the modular Step 3 handoff."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

from excel_to_act.steps.step3 import workflow


def build_semantic_plan(analysis_dir: Path, semantic_map_path: Path) -> dict[str, Any]:
    """Validate and bind an agent-authored mapping without evaluating formulas."""
    try:
        analysis_dir, manifest, binding, root, inventory_path = workflow._load_context(analysis_dir)
        from excel_to_act.steps.step3.input_boundary import require_confirmed_boundary_for_analysis

        require_confirmed_boundary_for_analysis(analysis_dir, manifest, binding)
        fields, fields_sha = workflow._read_stage(analysis_dir, manifest, "fields")
        dependencies, dependencies_sha = workflow._read_stage(analysis_dir, manifest, "dependencies")
        inventory = workflow._read_json(inventory_path)
        semantic_map_path = semantic_map_path.expanduser().resolve()
        if not semantic_map_path.is_file():
            raise ValueError("semantic map input file does not exist")
        input_value = workflow._read_json(semantic_map_path)
        plan = _build_plan(input_value, semantic_map_path, analysis_dir, manifest, binding, fields, dependencies,
                           fields_sha, dependencies_sha, inventory, root)
        markdown = _plan_markdown(plan)
        plan["report_md_sha256"] = workflow._hash_bytes(markdown.encode("utf-8"))
        digest = workflow._record_stage(analysis_dir, manifest, "semantic_plan", plan)
        workflow._write_bytes(analysis_dir / "semantic_plan.md", markdown.encode("utf-8"))
        return {"tool": "step3.plan", "status": "pass", "analysis": str(analysis_dir),
                "semantic_plan": "semantic_plan.json", "sha256": digest,
                "report_md_sha256": plan["report_md_sha256"], "coverage": plan["coverage"],
                **({"candidate_scope": plan["candidate_scope"]} if plan.get("candidate_scope") else {}),
                "readiness": plan["readiness"]}
    except Exception as exc:
        return {"tool": "step3.plan", "status": "blocked",
                "diagnostics": [{"code": "semantic_map_invalid", "severity": "error", "message": str(exc)}]}


def validate_semantic_plan(analysis_dir: Path, semantic_plan_path: Path | None = None) -> dict[str, Any]:
    """Recheck the saved map and evidence references, then write a static check report."""
    try:
        analysis_dir, manifest, binding, root, inventory_path = workflow._load_context(analysis_dir)
        from excel_to_act.steps.step3.input_boundary import require_confirmed_boundary_for_analysis

        require_confirmed_boundary_for_analysis(analysis_dir, manifest, binding)
        fields, fields_sha = workflow._read_stage(analysis_dir, manifest, "fields")
        dependencies, dependencies_sha = workflow._read_stage(analysis_dir, manifest, "dependencies")
        inventory = workflow._read_json(inventory_path)
        plan, plan_sha = workflow._read_stage(analysis_dir, manifest, "semantic_plan")
        plan_path = (analysis_dir / "semantic_plan.json").resolve()
        if semantic_plan_path is not None and semantic_plan_path.expanduser().resolve() != plan_path:
            raise ValueError("--semantic-plan must name this analysis directory's current semantic_plan.json")
        expected_md = plan.get("report_md_sha256")
        md_path = analysis_dir / "semantic_plan.md"
        if not md_path.is_file() or workflow._hash_file(md_path) != expected_md:
            raise ValueError("semantic_plan.md changed or is missing")
        input_ref = plan.get("semantic_map_input", {})
        input_path = Path(input_ref.get("path", ""))
        if not input_path.is_file() or workflow._hash_file(input_path) != input_ref.get("sha256"):
            raise ValueError("semantic map input changed or is missing; rebuild the plan")
        expected = _build_plan(workflow._read_json(input_path), input_path, analysis_dir, manifest, binding,
                               fields, dependencies, fields_sha, dependencies_sha, inventory, root)
        if expected.get("plan_fingerprint") != plan.get("plan_fingerprint"):
            raise ValueError("semantic plan does not match its current source map and evidence")
        check = {"schema_version": "step3.semantic_check.v1", "tool": "step3.check",
                 "status": "pass", "binding_sha256": manifest["binding_sha256"],
                 "semantic_plan_sha256": plan_sha,
                 "semantic_map_input_sha256": input_ref["sha256"],
                 "source_family_profile_sha256": plan["evidence"]["source_family_profile"]["sha256"],
                 "selection_basis": plan["selection_basis"],
                 "coverage": plan["coverage"],
                 "source_classification": plan["source_classification"],
                 "readiness": {"static_mapping_complete": True, "semantic_review_required": True,
                               "generation_ready": False, "runtime_verified": False,
                               "native_reconciled": False}}
        trace_evidence_key = "source_candidate_trace" if plan["selection_basis"] == "static_candidate" else "active_trace"
        check[f"{trace_evidence_key}_sha256"] = plan["evidence"][trace_evidence_key]["sha256"]
        if plan["selection_basis"] == "static_candidate":
            check["candidate_closure"] = "unknown"
            if plan.get("candidate_scope") is not None:
                check["candidate_scope"] = plan["candidate_scope"]
                check["model_scope_sha256"] = plan["evidence"]["model_scope"]["sha256"]
        markdown = _check_markdown(check)
        check["report_md_sha256"] = workflow._hash_bytes(markdown.encode("utf-8"))
        digest = workflow._record_stage(analysis_dir, manifest, "semantic_check", check)
        workflow._write_bytes(analysis_dir / "semantic_check.md", markdown.encode("utf-8"))
        return {"tool": "step3.check", "status": "pass", "analysis": str(analysis_dir),
                "semantic_check": "semantic_check.json", "sha256": digest,
                "report_md_sha256": check["report_md_sha256"], "coverage": check["coverage"],
                "source_classification": check["source_classification"],
                **({"candidate_scope": check["candidate_scope"]} if check.get("candidate_scope") else {}),
                "readiness": check["readiness"]}
    except Exception as exc:
        return {"tool": "step3.check", "status": "blocked",
                "diagnostics": [{"code": "semantic_plan_invalid", "severity": "error", "message": str(exc)}]}


def _build_plan(value: Any, input_path: Path, analysis_dir: Path, manifest: dict[str, Any], binding: dict[str, Any],
                fields_artifact: dict[str, Any], dependencies: dict[str, Any],
                fields_sha: str, dependencies_sha: str, inventory: dict[str, Any],
                step1_root: Path) -> dict[str, Any]:
    if not isinstance(value, dict) or value.get("schema_version") != "step3.semantic_map.input.v1":
        raise ValueError("semantic map must use schema_version step3.semantic_map.input.v1")
    source_binding = binding["source"]
    declared_source = value.get("source")
    expected_source = {"source_id": source_binding["source_id"], "run_id": source_binding["run_id"],
                       "source_sha256": source_binding["source_sha256"],
                       "binding_sha256": manifest["binding_sha256"],
                       "fields_sha256": fields_sha, "dependencies_sha256": dependencies_sha}
    if not isinstance(declared_source, dict) or any(declared_source.get(key) != item for key, item in expected_source.items()):
        raise ValueError("semantic map source or current Step 3 artifact hashes do not match")

    selection_basis = value.get("selection_basis", "active_trace")
    if selection_basis not in {"active_trace", "static_candidate"}:
        raise ValueError("semantic map selection_basis must be active_trace or static_candidate")
    candidate = selection_basis == "static_candidate"
    profile_ref, profile = _read_evidence(value.get("source_family_profile"), "source family profile", input_path.parent)
    trace_field = "source_candidate_trace" if candidate else "active_trace"
    trace_ref, trace = _read_evidence(value.get(trace_field),
                                      "source-candidate formula trace" if candidate else "active formula trace",
                                      input_path.parent)
    expected_profile_schema = "gp.source_candidate_family_profile.v1" if candidate else "gp.source_family_profile.v1"
    if profile.get("schema_version") != expected_profile_schema or profile.get("status") != "pass":
        raise ValueError("source family profile is not a passing supported profile")
    if profile.get("source_sha256") != source_binding["source_sha256"] or profile.get("binding_sha256") != manifest["binding_sha256"]:
        raise ValueError("source family profile belongs to another source or Step 3 binding")
    if profile.get("inventory", {}).get("sha256") != binding["artifacts"]["inventory_sha256"]:
        raise ValueError("source family profile was built from a different inventory")
    if trace.get("source_sha256") != source_binding["source_sha256"]:
        raise ValueError("formula trace belongs to another workbook source")
    if candidate:
        from excel_to_act.steps.step3.profiling import _validate_candidate_trace_binding

        _validate_candidate_trace_binding(analysis_dir,
                                          manifest, binding, trace, Path(trace_ref["path"]), inventory)
        if trace.get("selection_basis") != "static_candidate" or profile.get("selection_basis") != "static_candidate":
            raise ValueError("candidate profile and trace must declare static_candidate selection")
        if profile.get("source_candidate_trace") != trace_ref:
            raise ValueError("candidate profile does not cite the current source-candidate trace")
        if (profile.get("candidate_scope") != trace.get("candidate_scope")
                or profile.get("model_scope") != trace.get("model_scope")):
            raise ValueError("candidate profile and trace model-scope evidence differ")
    else:
        if trace.get("native_excel_called") is not False or trace.get("formula_cache_inputs") is not False:
            raise ValueError("active trace must be source-only and must not use formula caches")
        if profile.get("selection_basis") != "active_trace":
            raise ValueError("source family profile must cite the active-trace selection basis")

    families_by_id: dict[str, dict[str, Any]] = {}
    profile_members: dict[str, str] = {}
    for family in profile.get("families", []):
        family_id = family.get("family_id")
        members = family.get("source_members")
        if not isinstance(family_id, str) or not family_id or family_id in families_by_id:
            raise ValueError("source family profile has missing or duplicate family IDs")
        if not isinstance(members, list) or family.get("member_count") != len(members):
            raise ValueError(f"source family {family_id} has inconsistent members")
        families_by_id[family_id] = family
        for address in members:
            key = _address_key(address)
            if key in profile_members:
                raise ValueError(f"source formula appears in multiple profile families: {address}")
            profile_members[key] = family_id

    trace_ordinary: dict[str, dict[str, Any]] = {}
    trace_arrays: dict[tuple[str, str], dict[str, Any]] = {}
    array_members: dict[tuple[str, str], list[str]] = defaultdict(list)
    for cell in trace.get("cells", []):
        role = cell.get("role")
        address = cell.get("address")
        if role == "calculated_formula":
            key = _address_key(address)
            if key in trace_ordinary:
                raise ValueError(f"active trace has duplicate formula address: {address}")
            trace_ordinary[key] = cell
        elif role == "calculated_array_formula":
            array_ref = cell.get("array_ref")
            if not isinstance(array_ref, str) or not array_ref:
                raise ValueError(f"array formula has no source range: {address}")
            sheet, _ = _qualified_address(address)
            key = (sheet.casefold(), array_ref.replace("$", "").upper())
            instance = trace_arrays.setdefault(key, {"sheet": sheet, "array_ref": array_ref.replace("$", "").upper(),
                                                      "formula": cell.get("formula")})
            if cell.get("formula") != instance["formula"]:
                raise ValueError(f"active array formula members disagree on formula: {address}")
            array_members[key].append(str(address))
    if candidate:
        for array in trace.get("arrays", []):
            if not isinstance(array, dict):
                raise ValueError("candidate trace has a malformed array instance")
            anchor = array.get("anchor")
            array_ref = array.get("array_ref")
            sheet = array.get("sheet")
            formula = array.get("formula")
            if not all(isinstance(item, str) and item for item in (anchor, array_ref, sheet, formula)):
                raise ValueError("candidate array instance needs a source anchor, finite range, sheet, and formula")
            key = (sheet.casefold(), array_ref.replace("$", "").upper())
            if key in trace_arrays:
                raise ValueError(f"candidate trace has a duplicate array instance: {sheet}!{array_ref}")
            trace_arrays[key] = {"sheet": sheet, "array_ref": array_ref.replace("$", "").upper(),
                                 "formula": formula, "anchor": anchor,
                                 "shape": array.get("shape"), "follower_count": array.get("follower_count")}
            followers = array.get("follower_addresses")
            if not isinstance(followers, list) or array.get("follower_count") != len(followers):
                raise ValueError(f"candidate array follower list is inconsistent: {sheet}!{array_ref}")
            array_members[key].extend(followers)
            if _address_key(anchor) not in trace_ordinary:
                raise ValueError(f"candidate array anchor is not an ordinary source formula candidate: {anchor}")
    if set(profile_members) != set(trace_ordinary):
        missing = len(set(trace_ordinary) - set(profile_members))
        extra = len(set(profile_members) - set(trace_ordinary))
        raise ValueError(f"formula profile/trace coverage differs ({missing} active formulas missing, {extra} extra)")
    fields, cell_fields = workflow._field_index(fields_artifact)
    for key, cell in trace_ordinary.items():
        sheet, local = _qualified_address(cell["address"])
        field_id = cell_fields.get((sheet.casefold(), local.upper()))
        if field_id is None or fields[field_id].get("role") != "calculated":
            raise ValueError(f"active formula is not represented by a current calculated primary field: {cell['address']}")
    for addresses in array_members.values():
        for address in addresses:
            sheet, local = _qualified_address(address)
            field_id = cell_fields.get((sheet.casefold(), local.upper()))
            if field_id is None or fields[field_id].get("role") != "calculated":
                raise ValueError(f"active array follower is not represented by a calculated primary field: {address}")

    modules = _validate_modules(value.get("modules"))
    variables = _validate_variables(value.get("variables"))
    if candidate:
        external_boundary_variables = _validate_external_boundary_variables(
            variables, binding, manifest, trace
        )
    else:
        if any(item.get("role") == "external" for item in variables.values()):
            raise ValueError("external boundary variables require a confirmed static-candidate trace")
        external_boundary_variables = []
    source_classification = validate_raw_source_extents(variables, inventory,
                                                        source_binding["source_sha256"])
    equation_families = _validate_equation_families(value.get("equation_families"), modules)
    reference_resolution = None
    execution_plan = None
    if "reference_resolution" in value:
        reference_resolution = _validate_reference_resolution(value["reference_resolution"], variables)
    source_mappings = _validate_source_mappings(value.get("source_mappings"), families_by_id,
                                               equation_families, variables, modules)
    if reference_resolution is not None:
        execution_plan = _validate_execution_plan(value.get("execution_plan"), modules,
                                                  set(families_by_id), variables, source_mappings)
    expected_family_ids = set(families_by_id)
    assigned_family_ids = {item["source_family_id"] for item in source_mappings}
    if assigned_family_ids != expected_family_ids:
        raise ValueError(f"source family mappings are incomplete ({len(expected_family_ids - assigned_family_ids)} unmapped, {len(assigned_family_ids - expected_family_ids)} unknown)")
    for item in source_mappings:
        if variables[item["variable_id"]].get("role") == "external":
            raise ValueError("formula source families cannot be mapped as external boundary values")
        source_family = families_by_id[item["source_family_id"]]
        item["source_formula_fingerprint"] = _source_formula_fingerprint(source_family)
        item["source_members"] = sorted(source_family["source_members"], key=_address_key)
        item["member_count"] = len(item["source_members"])
    for family_id, family in equation_families.items():
        family_mappings = [item for item in source_mappings if item["equation_family_id"] == family_id]
        fingerprint_input = {"family": family, "source_mappings": family_mappings}
        family["family_fingerprint"] = workflow._hash_bytes(workflow._json_bytes(fingerprint_input))
    array_maps = _validate_array_mappings(value.get("array_mappings"), trace_arrays, array_members,
                                          variables, modules, trace_ordinary, profile_members,
                                          {item["source_family_id"]: item for item in source_mappings})
    if reference_resolution is not None:
        _validate_execution_segments(variables, source_mappings, array_maps)
    _validate_mapped_demand(variables, source_mappings, array_maps)
    for group in array_maps:
        fingerprint_input = {"array_family_id": group["array_family_id"], "instances": group["instances"]}
        group["family_fingerprint"] = workflow._hash_bytes(workflow._json_bytes(fingerprint_input))

    ordinary_count = len(trace_ordinary)
    array_count = sum(len(addresses) for addresses in array_members.values())
    mapped_ordinary = sum(item["member_count"] for item in source_mappings)
    mapped_arrays = sum(item["member_count"] for group in array_maps for item in group["instances"])
    if candidate:
        coverage = {"static_candidate_formula_members": ordinary_count,
                    "static_candidate_formula_members_mapped": mapped_ordinary,
                    "static_candidate_array_followers": array_count,
                    "static_candidate_array_followers_mapped": mapped_arrays,
                    "static_candidate_member_count": ordinary_count + array_count,
                    "static_candidate_member_count_mapped": mapped_ordinary + mapped_arrays,
                    "unmapped_static_candidate_members": 0,
                    "source_family_count": len(families_by_id),
                    "semantic_family_count": len(equation_families),
                    "candidate_array_instance_count": len(trace_arrays),
                    "array_family_count": len(array_maps),
                    "candidate_closure": "unknown",
                    "active_formula_member_count": None}
    else:
        coverage = {"ordinary_formula_members": ordinary_count,
                    "ordinary_formula_members_mapped": mapped_ordinary,
                    "array_formula_followers": array_count,
                    "array_formula_followers_mapped": mapped_arrays,
                    "active_formula_members": ordinary_count + array_count,
                    "active_formula_members_mapped": mapped_ordinary + mapped_arrays,
                    "unmapped_active_formula_members": 0, "duplicate_active_formula_members": 0,
                    "source_family_count": len(families_by_id),
                    "semantic_family_count": len(equation_families),
                    "array_instance_count": len(trace_arrays),
                    "array_family_count": len(array_maps)}
    if mapped_ordinary != ordinary_count or mapped_arrays != array_count:
        raise ValueError("mapped source member counts do not equal the active trace")

    result = {"schema_version": "step3.semantic_plan.v1", "tool": "step3.plan", "status": "pass",
              "selection_basis": selection_basis,
              "binding_sha256": manifest["binding_sha256"], "source": expected_source,
              "semantic_map_input": {"path": str(input_path), "sha256": workflow._hash_file(input_path)},
              "evidence": {"source_family_profile": profile_ref, trace_field: trace_ref},
              "modules": modules, "variables": variables, "equation_families": equation_families,
              "source_mappings": source_mappings, "array_mappings": array_maps,
              "coverage": coverage, "source_classification": source_classification,
              "external_boundary_variables": external_boundary_variables,
              "readiness": {"static_mapping_complete": True, "semantic_review_required": True,
                            "generation_ready": False, "runtime_verified": False,
                            "native_reconciled": False, "gpu_executable": False}}
    if candidate and trace.get("model_scope") is not None:
        scope_ref = trace["model_scope"]
        result["candidate_scope"] = trace["candidate_scope"]
        result["evidence"]["model_scope"] = {"path": scope_ref["path"], "sha256": scope_ref["sha256"]}
        coverage["scope_frontier_count"] = trace["summary"]["scope_frontier_count"]
        coverage["proposed_formula_start_count"] = trace["summary"]["proposed_formula_start_count"]
    if reference_resolution is not None:
        source_blank_boundaries = _validate_source_blank_inputs(
            reference_resolution, variables, inventory, source_binding, step1_root
        )
        source_classification["verified_source_blank_boundary_count"] = len(source_blank_boundaries)
        result["reference_resolution"] = reference_resolution
        result["execution_plan"] = execution_plan
        result["source_blank_boundaries"] = source_blank_boundaries
        if isinstance(value.get("semantic_design_notes"), list):
            result["semantic_design_notes"] = list(value["semantic_design_notes"])
    result["plan_fingerprint"] = workflow._hash_bytes(workflow._json_bytes(result))
    return result


def validate_raw_source_extents(variables: dict[str, dict[str, Any]], inventory: dict[str, Any],
                               source_sha256: str) -> dict[str, Any]:
    """Reject demanded raw coordinates backed by source formulas, including array followers."""
    if not isinstance(inventory, dict) or inventory.get("workbook_sha256") != source_sha256:
        raise ValueError("source-kind check inventory does not match the bound workbook")
    sheets = inventory.get("sheets")
    if not isinstance(sheets, list):
        raise ValueError("source-kind check inventory has no sheet records")

    formula_cells: set[str] = set()
    array_ranges: list[tuple[str, tuple[int, int, int, int], str, str]] = []
    for sheet_record in sheets:
        if not isinstance(sheet_record, dict) or not isinstance(sheet_record.get("name"), str):
            raise ValueError("source-kind check inventory has a malformed sheet")
        sheet = sheet_record["name"]
        for cell in sheet_record.get("cells", []):
            if not isinstance(cell, dict) or not isinstance(cell.get("address"), str):
                continue
            address = f"{sheet}!{cell['address']}"
            direct_formula = (cell.get("formula_present") is True or cell.get("kind") == "formula"
                              or cell.get("data_type") == "f"
                              or isinstance(cell.get("formula"), str) and cell["formula"].startswith("=")
                              or isinstance(cell.get("raw_formula_text"), str))
            if direct_formula:
                formula_cells.add(_address_key(address))
            attributes = cell.get("raw_formula_attributes")
            if isinstance(attributes, dict) and attributes.get("t") == "array":
                array_ref = attributes.get("ref")
                if not isinstance(array_ref, str) or not array_ref:
                    raise ValueError(f"source array formula has no finite extent at {address}")
                try:
                    bounds = range_boundaries(array_ref.replace("$", ""))
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"source array formula has an invalid extent at {address}") from exc
                if None in bounds:
                    raise ValueError(f"source array formula has an unbounded extent at {address}")
                anchor = _address_key(f"{sheet}!{get_column_letter(bounds[0])}{bounds[1]}")
                array_ranges.append((sheet.casefold(), bounds, anchor, address))

    raw_extent_coordinates: set[str] = set()
    raw_formula_outside_demand = 0
    raw_array_follower_outside_demand = 0
    raw_extent_count = 0
    raw_seed_variable_count = 0
    raw_seed_extent_count = 0
    variables_by_role_and_kind: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    demanded_variables_by_role: dict[str, int] = defaultdict(int)
    for variable_id, variable in variables.items():
        role, kind = variable.get("role"), variable.get("kind")
        variables_by_role_and_kind[str(role)][str(kind)] += 1
        if _variable_has_demand(variable):
            demanded_variables_by_role[str(role)] += 1
        if role != "raw":
            continue
        extents = variable.get("source_extents", [])
        seeds = [extent for extent in extents if extent.get("role") == "raw_seed"]
        if seeds:
            raw_seed_variable_count += 1
            raw_seed_extent_count += len(seeds)
        for extent in extents:
            reference = extent.get("ref")
            coordinate_mapping = extent.get("coordinate_mapping")
            if not isinstance(coordinate_mapping, dict):
                raise ValueError(f"raw variable {variable_id} extent needs an explicit coordinate mapping")
            sheet, local, _dimensions = _validate_source_ref(reference, f"raw variable {variable_id}")
            try:
                min_col, min_row, max_col, max_row = range_boundaries(local.replace("$", ""))
            except (TypeError, ValueError) as exc:
                raise ValueError(f"raw variable {variable_id} has an invalid source extent {reference!r}") from exc
            if None in (min_col, min_row, max_col, max_row):
                raise ValueError(f"raw variable {variable_id} has an unbounded source extent {reference!r}")
            raw_extent_count += 1
            for row in range(min_row, max_row + 1):
                for column in range(min_col, max_col + 1):
                    address = f"{sheet}!{get_column_letter(column)}{row}"
                    address_key = _address_key(address)
                    raw_extent_coordinates.add(address_key)
                    is_formula = address_key in formula_cells
                    is_array_follower = any(
                        candidate_sheet == sheet.casefold()
                        and bounds[0] <= column <= bounds[2]
                        and bounds[1] <= row <= bounds[3]
                        and anchor != address_key
                        for candidate_sheet, bounds, anchor, _anchor_address in array_ranges
                    )
                    if not is_formula and not is_array_follower:
                        continue
                    try:
                        indices = _mapped_logical_indices(coordinate_mapping, address)
                    except ValueError as exc:
                        raise ValueError(
                            f"raw formula boundary cannot be checked against demand at {address} ({variable_id})"
                        ) from exc
                    if _is_demanded(variable, indices):
                        boundary = "array formula follower" if is_array_follower and not is_formula else "formula cell"
                        raise ValueError(
                            f"raw variable {variable_id} claims demanded {boundary} {address}; source formulas cannot be raw inputs"
                        )
                    if is_array_follower and not is_formula:
                        raw_array_follower_outside_demand += 1
                    else:
                        raw_formula_outside_demand += 1

    return {
        "source_kind_basis": "Promoted inventory formula markers and native array ranges; cached values are not consulted.",
        "source_formula_cell_count": len(formula_cells),
        "source_array_formula_range_count": len(array_ranges),
        "variables_by_role_and_kind": {
            role: dict(sorted(kinds.items())) for role, kinds in sorted(variables_by_role_and_kind.items())
        },
        "demanded_variables_by_role": dict(sorted(demanded_variables_by_role.items())),
        "raw_seed_variable_count": raw_seed_variable_count,
        "raw_seed_extent_count": raw_seed_extent_count,
        "raw_extent_count": raw_extent_count,
        "distinct_raw_extent_coordinate_count": len(raw_extent_coordinates),
        "demanded_raw_formula_overlaps": 0,
        "demanded_raw_array_follower_overlaps": 0,
        "formula_coordinates_outside_raw_demand": raw_formula_outside_demand,
        "array_follower_coordinates_outside_raw_demand": raw_array_follower_outside_demand,
    }


def _read_evidence(reference: Any, label: str, base_dir: Path) -> tuple[dict[str, str], dict[str, Any]]:
    if not isinstance(reference, dict) or not isinstance(reference.get("path"), str):
        raise ValueError(f"{label} must declare path and SHA-256")
    path = Path(reference["path"]).expanduser().resolve()
    if not Path(reference["path"]).expanduser().is_absolute():
        path = (base_dir / reference["path"]).resolve()
    if not path.is_file() or workflow._hash_file(path) != reference.get("sha256"):
        raise ValueError(f"{label} is missing or changed")
    value = workflow._read_json(path)
    if not isinstance(value, dict):
        raise ValueError(f"{label} must contain a JSON object")
    return {"path": str(path), "sha256": reference["sha256"]}, value


def _validate_source_blank_inputs(reference_resolution: dict[str, Any],
                                  variables: dict[str, dict[str, Any]], inventory: dict[str, Any],
                                  source_binding: dict[str, Any], step1_root: Path) -> list[dict[str, Any]]:
    """Verify explicitly declared blank table-boundary inputs from source bytes, never caches."""
    blank_inputs = [blank
                    for reference in reference_resolution["named_references"]
                    for table_binding in [reference.get("table_binding")]
                    if isinstance(table_binding, dict)
                    for blank in table_binding.get("source_blank_inputs", [])]
    if not blank_inputs:
        return []

    source_path = _bound_source_workbook(step1_root, source_binding)
    raw_extents: list[tuple[str, tuple[int, int, int, int], str, dict[str, Any]]] = []
    for variable_id, variable in variables.items():
        if variable["role"] != "raw":
            continue
        for extent in variable["source_extents"]:
            sheet, local, _ = _validate_source_ref(extent["ref"], f"raw source extent for {variable_id}")
            raw_extents.append((sheet.casefold(), range_boundaries(local.replace("$", "")),
                                variable_id, extent))

    formula_cells: set[str] = set()
    array_ranges: list[tuple[str, tuple[int, int, int, int]]] = []
    for sheet_record in inventory.get("sheets", []):
        if not isinstance(sheet_record, dict) or not isinstance(sheet_record.get("name"), str):
            continue
        sheet_name = sheet_record["name"]
        for cell in sheet_record.get("cells", []):
            if not isinstance(cell, dict) or not isinstance(cell.get("address"), str):
                continue
            address = f"{sheet_name}!{cell['address']}"
            if (cell.get("formula_present") is True or cell.get("kind") == "formula"
                    or cell.get("data_type") == "f"
                    or isinstance(cell.get("formula"), str) and cell["formula"].startswith("=")
                    or isinstance(cell.get("raw_formula_text"), str)):
                formula_cells.add(_address_key(address))
            attributes = cell.get("raw_formula_attributes")
            if isinstance(attributes, dict) and attributes.get("t") == "array":
                bounds = range_boundaries(str(attributes.get("ref", "")).replace("$", ""))
                if None in bounds:
                    raise ValueError(f"source blank overlaps an unbounded array formula at {address}")
                array_ranges.append((sheet_name.casefold(), bounds))

    verified: list[dict[str, Any]] = []
    workbook = load_workbook(source_path, data_only=False, read_only=True, keep_vba=True)
    try:
        for item in blank_inputs:
            address = item["source_ref"]
            sheet, local = _qualified_address(address)
            if sheet not in workbook.sheetnames:
                raise ValueError(f"source blank sheet is absent from the bound workbook: {address}")
            min_col, min_row, max_col, max_row = range_boundaries(local.replace("$", ""))
            if min_col != max_col or min_row != max_row:
                raise ValueError(f"source blank must be one cell: {address}")
            if _address_key(address) in formula_cells:
                raise ValueError(f"declared raw source blank is a formula cell: {address}")
            if any(candidate_sheet == sheet.casefold()
                   and bounds[0] <= min_col <= bounds[2] and bounds[1] <= min_row <= bounds[3]
                   for candidate_sheet, bounds in array_ranges):
                raise ValueError(f"declared raw source blank overlaps a source array formula: {address}")
            cell = workbook[sheet][local]
            if cell.value is not None or cell.data_type == "f":
                raise ValueError(f"declared source boundary is not a blank raw cell: {address}")
            covering = [record for record in raw_extents
                        if record[0] == sheet.casefold()
                        and record[1][0] <= min_col <= record[1][2]
                        and record[1][1] <= min_row <= record[1][3]]
            if not covering:
                raise ValueError(f"declared source blank is not represented by a raw input extent: {address}")
            if item["role"] == "formula_tail_blank":
                variable_id, expected_index = item["variable_id"], item["index"]
                matching = [record for record in covering if record[2] == variable_id]
                actual_indices = {_mapped_logical_indices(record[3]["coordinate_mapping"], address)
                                  for record in matching}
                if ((expected_index,) not in actual_indices
                        or not _is_demanded(variables[variable_id], (expected_index,))):
                    raise ValueError(f"formula tail blank has an invalid raw-variable coordinate: {address}")
            verified.append({"source_ref": address, "role": item["role"],
                             "variable_id": item.get("variable_id"), "index": item.get("index"),
                             "verified_as": "blank source cell; formula markers, array ranges, and workbook bytes checked"})
    finally:
        workbook.close()
    return verified


def _bound_source_workbook(step1_root: Path, source_binding: dict[str, Any]) -> Path:
    run_path = source_binding.get("run_path")
    if not isinstance(run_path, str) or not run_path:
        raise ValueError("source blank verification requires the checked Step 1 run path")
    run_dir = workflow._resolved_child(step1_root, run_path, "Step 1 run path")
    record_path = run_dir / "source.json"
    if not record_path.is_file():
        raise ValueError("source blank verification requires the Step 1 source.json record")
    record = workflow._read_json(record_path)
    if (not isinstance(record, dict) or record.get("run_id") != source_binding.get("run_id")
            or record.get("sha256") != source_binding.get("source_sha256")):
        raise ValueError("Step 1 source.json does not match the bound workbook")
    declared_path = record.get("source_path")
    if not isinstance(declared_path, str) or not declared_path:
        raise ValueError("Step 1 source.json has no source_path for source-blank verification")
    path = Path(declared_path).expanduser()
    if not path.is_absolute():
        path = (run_dir / path).resolve()
    if not path.is_file() or workflow._hash_file(path) != source_binding.get("source_sha256"):
        raise ValueError("bound source workbook is missing or its SHA-256 changed")
    return path.resolve()


def _validate_modules(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list) or not value:
        raise ValueError("modules must be a non-empty list")
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    orders: set[int] = set()
    for item in value:
        if not isinstance(item, dict):
            raise ValueError("each module must be an object")
        module_id, order, purpose = item.get("module_id"), item.get("order"), item.get("purpose")
        if not isinstance(module_id, str) or not module_id or module_id in seen:
            raise ValueError("module_id values must be unique non-empty strings")
        if not isinstance(order, int) or isinstance(order, bool) or order < 0 or order in orders:
            raise ValueError("module order values must be unique non-negative integers")
        if not isinstance(purpose, str) or not purpose.strip():
            raise ValueError(f"module {module_id} needs a purpose")
        seen.add(module_id)
        orders.add(order)
        result.append(dict(item))
    result.sort(key=lambda item: item["order"])
    return result


def _validate_variables(value: Any) -> dict[str, dict[str, Any]]:
    if not isinstance(value, list) or not value:
        raise ValueError("variables must be a non-empty list")
    result: dict[str, dict[str, Any]] = {}
    allowed = {"scalar", "projection_series", "lookup_vector", "matrix"}
    for item in value:
        if not isinstance(item, dict):
            raise ValueError("each variable must be an object")
        variable_id = item.get("variable_id")
        if not isinstance(variable_id, str) or not variable_id or variable_id in result:
            raise ValueError("variable_id values must be unique non-empty strings")
        if item.get("role") not in {"raw", "derived", "external"} or item.get("kind") not in allowed:
            raise ValueError(f"variable {variable_id} needs raw/derived/external role and supported kind")
        if item.get("role") == "external" and (not isinstance(item.get("boundary_variable_id"), str)
                                                 or not item["boundary_variable_id"]):
            raise ValueError(f"external variable {variable_id} needs a confirmed boundary_variable_id")
        shape, axes = item.get("shape"), item.get("axes")
        if (not isinstance(shape, list) or any(not isinstance(size, int) or isinstance(size, bool) or size < 0 for size in shape)
                or not isinstance(axes, list) or len(shape) != len(axes)):
            raise ValueError(f"variable {variable_id} needs matching shape and logical axes")
        expected_dimensions = {"scalar": 0, "projection_series": 1, "lookup_vector": 1, "matrix": 2}[item["kind"]]
        if len(shape) != expected_dimensions:
            raise ValueError(f"variable {variable_id} shape does not match kind {item['kind']}")
        if not isinstance(item.get("source_extents"), list) or not item["source_extents"]:
            raise ValueError(f"variable {variable_id} needs source_extents")
        demanded_indices = item.get("demanded_indices")
        if not (isinstance(demanded_indices, (list, dict))
                or isinstance(demanded_indices, str) and demanded_indices == "all"):
            raise ValueError(f"variable {variable_id} needs demanded_indices")
        if "initial_condition" not in item or not isinstance(item.get("equation_segments"), list):
            raise ValueError(f"variable {variable_id} needs an explicit initial_condition and equation_segments")
        if not isinstance(item.get("dependencies"), list):
            raise ValueError(f"variable {variable_id} dependencies must be a list")
        for key in ("error_policy", "output_usage"):
            if not isinstance(item.get(key), str) or not item[key].strip():
                raise ValueError(f"variable {variable_id} needs {key}")
        result[variable_id] = dict(item)
    return result


def _validate_external_boundary_variables(
    variables: dict[str, dict[str, Any]], binding: dict[str, Any], manifest: dict[str, Any],
    candidate_trace: dict[str, Any],
) -> list[dict[str, Any]]:
    """Bind formula-derived external variables to the separately confirmed catalog."""
    from excel_to_act.steps.conversion_workflow import load_workflow
    from excel_to_act.steps.step3.input_boundary import confirmed_boundary_reference

    workflow_path = binding.get("workflow_path")
    if not isinstance(workflow_path, str) or not workflow_path:
        raise ValueError("external boundary variables require a workflow-bound input catalog")
    workflow_root, workflow_manifest = load_workflow(Path(workflow_path))
    boundary_reference, boundary_path, _boundary_md = confirmed_boundary_reference(workflow_root, workflow_manifest)
    declared_reference = candidate_trace.get("input_boundary", {})
    for key in ("revision", "boundary_sha256", "artifact", "review_receipts"):
        if declared_reference.get(key) != boundary_reference.get(key):
            raise ValueError("external semantic variables cite a stale input-boundary revision")
    boundary = workflow._read_json(boundary_path)
    expected_source = {"source_id": binding["source"]["source_id"],
                       "run_id": binding["source"]["run_id"],
                       "workbook_sha256": binding["source"]["source_sha256"],
                       "analysis_binding_sha256": manifest["binding_sha256"]}
    if boundary.get("source") != expected_source:
        raise ValueError("confirmed external input catalog belongs to another source or analysis binding")
    catalog_variables = {
        item.get("variable_id"): item for item in boundary.get("variables", [])
        if isinstance(item, dict) and item.get("role") == "formula_derived_external"
    }
    semantic_by_boundary: dict[str, tuple[str, dict[str, Any]]] = {}
    for variable_id, variable in variables.items():
        if variable.get("role") != "external":
            continue
        boundary_id = variable["boundary_variable_id"]
        if boundary_id in semantic_by_boundary:
            raise ValueError(f"confirmed external boundary variable is represented more than once: {boundary_id}")
        semantic_by_boundary[boundary_id] = (variable_id, variable)
    if set(semantic_by_boundary) != set(catalog_variables):
        missing = sorted(set(catalog_variables) - set(semantic_by_boundary))
        extra = sorted(set(semantic_by_boundary) - set(catalog_variables))
        raise ValueError(f"external semantic variables must match every confirmed formula-derived input (missing={missing}, extra={extra})")

    checked: list[dict[str, Any]] = []
    for boundary_id, (variable_id, variable) in sorted(semantic_by_boundary.items()):
        declared = catalog_variables[boundary_id]
        if variable["shape"] != declared.get("shape"):
            raise ValueError(f"external variable {variable_id} shape differs from the confirmed input catalog")
        expected_extents = {
            _normalize_source_extent(f"{item['sheet']}!{item['range']}")
            for item in declared.get("source_extents", [])
            if isinstance(item, dict) and isinstance(item.get("sheet"), str) and isinstance(item.get("range"), str)
        }
        actual_extents = set()
        for extent in variable["source_extents"]:
            ref = extent.get("ref") if isinstance(extent, dict) else None
            if not isinstance(ref, str):
                raise ValueError(f"external variable {variable_id} has a malformed source extent")
            _validate_source_ref(ref, f"external variable {variable_id}")
            actual_extents.add(_normalize_source_extent(ref))
        if actual_extents != expected_extents:
            raise ValueError(f"external variable {variable_id} extents differ from the confirmed input catalog")
        checked.append({"variable_id": variable_id, "boundary_variable_id": boundary_id,
                        "logical_name": variable.get("logical_name"), "shape": variable["shape"],
                        "source_extents": sorted(actual_extents), "values_available_now": False,
                        "formula_cache_allowed": False})
    return checked


def _normalize_source_extent(reference: str) -> str:
    sheet, local, _dimensions = _validate_source_ref(reference, "source extent")
    return f"{sheet.casefold()}!{local.replace('$', '').upper()}"


def _validate_reference_resolution(value: Any, variables: dict[str, dict[str, Any]]) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise ValueError("reference_resolution must be an object")
    named = value.get("named_references")
    ranges = value.get("cell_ranges")
    if not isinstance(named, list) or not isinstance(ranges, list):
        raise ValueError("reference_resolution needs named_references and cell_ranges lists")
    seen_names: set[tuple[str, str]] = set()
    normalized_names: list[dict[str, Any]] = []
    for item in named:
        if not isinstance(item, dict) or not all(isinstance(item.get(key), str) and item[key].strip()
                                                 for key in ("name", "kind", "source_ref", "resolution")):
            raise ValueError("each named reference needs name, kind, source_ref, and resolution")
        scope = item.get("scope", "workbook")
        if not isinstance(scope, str) or not scope:
            raise ValueError(f"named reference {item['name']} has an invalid scope")
        key = (scope.casefold(), item["name"].casefold())
        if key in seen_names:
            raise ValueError(f"duplicate named reference: {scope}!{item['name']}")
        seen_names.add(key)
        _validate_source_ref(item["source_ref"], f"named reference {item['name']}")
        if item["kind"] == "scalar":
            if item.get("variable_id") not in variables:
                raise ValueError(f"named scalar {item['name']} needs a known variable_id")
        elif item["kind"] == "range":
            table_binding = item.get("table_binding")
            if not ((isinstance(table_binding, str) and table_binding.strip())
                    or (isinstance(table_binding, dict) and table_binding)):
                raise ValueError(f"named range {item['name']} needs an explicit table_binding")
            for key_name in ("row_key_ref", "column_headers_ref"):
                if item.get(key_name) is not None:
                    _validate_source_ref(item[key_name], f"named range {item['name']} {key_name}")
            if item.get("lookup_mode") is not None and not isinstance(item["lookup_mode"], str):
                raise ValueError(f"named range {item['name']} lookup_mode must be text")
            if isinstance(table_binding, dict) and "source_blank_inputs" in table_binding:
                blank_inputs = table_binding["source_blank_inputs"]
                if not isinstance(blank_inputs, list) or not blank_inputs:
                    raise ValueError(f"named range {item['name']} source_blank_inputs must be a non-empty list")
                table_sheet, table_local, _ = _validate_source_ref(item["source_ref"],
                                                                   f"named range {item['name']}")
                table_bounds = range_boundaries(table_local.replace("$", ""))
                seen_blanks: set[tuple[str, str]] = set()
                for blank in blank_inputs:
                    if not isinstance(blank, dict) or not isinstance(blank.get("source_ref"), str):
                        raise ValueError(f"named range {item['name']} has a malformed source blank input")
                    blank_sheet, blank_local, dimensions = _validate_source_ref(
                        blank["source_ref"], f"named range {item['name']} source blank"
                    )
                    if dimensions != 0 or blank_sheet.casefold() != table_sheet.casefold():
                        raise ValueError(f"named range {item['name']} source blank must be one cell on the table sheet")
                    blank_bounds = range_boundaries(blank_local.replace("$", ""))
                    if not (table_bounds[0] <= blank_bounds[0] <= table_bounds[2]
                            and table_bounds[1] <= blank_bounds[1] <= table_bounds[3]):
                        raise ValueError(f"named range {item['name']} source blank is outside its physical table")
                    blank_key = (blank_sheet.casefold(), blank_local.upper())
                    if blank_key in seen_blanks:
                        raise ValueError(f"named range {item['name']} repeats a source blank input")
                    seen_blanks.add(blank_key)
                    if blank.get("role") not in {"key_trailing_blank", "formula_tail_blank"}:
                        raise ValueError(f"named range {item['name']} source blank needs a supported role")
                    if blank["role"] == "formula_tail_blank":
                        variable_id, index = blank.get("variable_id"), blank.get("index")
                        if (variable_id not in variables or variables[variable_id]["role"] != "raw"
                                or not isinstance(index, int) or isinstance(index, bool) or index < 0):
                            raise ValueError(f"named range {item['name']} formula tail blank needs a raw variable and index")
        elif item["kind"] == "dynamic_binding":
            if not isinstance(item.get("bindings"), list) or not item["bindings"]:
                raise ValueError(f"dynamic name {item['name']} needs explicit source bindings")
        else:
            raise ValueError(f"named reference {item['name']} has unsupported kind {item['kind']!r}")
        normalized_names.append(dict(item))

    normalized_ranges: list[dict[str, Any]] = []
    seen_ranges: set[tuple[str, str, str, str]] = set()
    for item in ranges:
        if not isinstance(item, dict) or not isinstance(item.get("variable_id"), str):
            raise ValueError("each cell_ranges entry needs a source_ref and variable_id")
        variable_id = item["variable_id"]
        if variable_id not in variables:
            raise ValueError(f"cell range refers to unknown variable {variable_id}")
        source_ref = item.get("source_ref")
        sheet, local, _physical_dimensions = _validate_source_ref(source_ref, f"cell range for {variable_id}")
        key = (sheet.casefold(), local.upper(), variable_id, str(item.get("role", "")))
        if key in seen_ranges:
            raise ValueError(f"duplicate cell-range mapping: {source_ref}")
        seen_ranges.add(key)
        coordinate = item.get("coordinate_mapping")
        if not isinstance(coordinate, dict) or not isinstance(coordinate.get("kind"), str):
            raise ValueError(f"cell range {source_ref} needs an explicit coordinate_mapping")
        allowed_coordinate_kinds = {"scalar", "row_ordinal", "column_ordinal", "row_column_ordinal",
                                    "keyed_rows", "keyed_columns", "named_range", "lookup_vector",
                                    "projection_series", "matrix", "transpose_array"}
        if coordinate["kind"] not in allowed_coordinate_kinds:
            raise ValueError(f"cell range {source_ref} has unsupported coordinate mapping kind")
        axes = coordinate.get("axes", [])
        if not isinstance(axes, list):
            raise ValueError(f"cell range {source_ref} coordinate axes must be a list")
        expected_axis_names = {axis.get("name", axis.get("axis")) for axis in variables[variable_id].get("axes", [])
                               if isinstance(axis, dict)}
        seen_dimensions: set[int] = set()
        for axis in axes:
            if not isinstance(axis, dict) or not isinstance(axis.get("axis"), str):
                raise ValueError(f"cell range {source_ref} has a malformed coordinate axis")
            dimension = axis.get("dimension")
            if not isinstance(dimension, int) or isinstance(dimension, bool) or not 0 <= dimension < len(variables[variable_id]["shape"]):
                raise ValueError(f"cell range {source_ref} axis dimension is outside variable {variable_id}")
            if dimension in seen_dimensions:
                raise ValueError(f"cell range {source_ref} maps multiple axes to dimension {dimension}")
            seen_dimensions.add(dimension)
            if expected_axis_names and axis["axis"] not in expected_axis_names:
                raise ValueError(f"cell range {source_ref} uses undeclared logical axis {axis['axis']}")
            allowed_coordinates = {"row", "column"}
            if coordinate["kind"] == "transpose_array":
                allowed_coordinates.update({"array_row", "array_column"})
            if axis.get("source_coordinate") not in allowed_coordinates:
                raise ValueError(f"cell range {source_ref} axis needs an explicit source_coordinate")
            if not isinstance(axis.get("origin"), (str, int)) or isinstance(axis.get("origin"), bool):
                raise ValueError(f"cell range {source_ref} axis needs an explicit coordinate origin")
            if not isinstance(axis.get("index_origin"), int) or isinstance(axis.get("index_origin"), bool):
                raise ValueError(f"cell range {source_ref} axis needs an integer index_origin")
        logical_dimensions = len(variables[variable_id]["shape"])
        if coordinate["kind"] == "scalar" and (logical_dimensions != 0 or axes):
            raise ValueError(f"scalar cell range {source_ref} must map to a scalar variable")
        if coordinate["kind"] in {"row_ordinal", "column_ordinal", "keyed_rows", "keyed_columns",
                                   "lookup_vector", "projection_series"} and len(axes) != 1:
            raise ValueError(f"one-coordinate cell range {source_ref} must map exactly one logical axis")
        if coordinate["kind"] in {"row_column_ordinal", "matrix", "transpose_array"} and logical_dimensions != 2:
            raise ValueError(f"two-axis cell range {source_ref} must map to a matrix variable")
        if coordinate["kind"] in {"row_column_ordinal", "matrix", "transpose_array"} and seen_dimensions != set(range(logical_dimensions)):
            raise ValueError(f"cell range {source_ref} must map both axes of {variable_id}")
        normalized_ranges.append(dict(item))
    expected_extent_map: set[bytes] = set()
    for variable_id, variable in variables.items():
        for extent in variable["source_extents"]:
            if not isinstance(extent, dict) or not isinstance(extent.get("ref"), str) or not isinstance(extent.get("coordinate_mapping"), dict):
                raise ValueError(f"variable {variable_id} source_extents need explicit ref and coordinate_mapping objects")
            expected_extent_map.add(workflow._json_bytes({"source_ref": extent["ref"], "variable_id": variable_id,
                                                           "coordinate_mapping": extent["coordinate_mapping"],
                                                           "role": extent.get("role")}))
    actual_extent_map = {workflow._json_bytes({"source_ref": item["source_ref"], "variable_id": item["variable_id"],
                                               "coordinate_mapping": item["coordinate_mapping"],
                                               "role": item.get("role")}) for item in normalized_ranges}
    if actual_extent_map != expected_extent_map:
        raise ValueError("reference_resolution.cell_ranges must match every variable source extent exactly")
    return {**value, "named_references": normalized_names, "cell_ranges": normalized_ranges}


def _validate_source_ref(value: Any, label: str) -> tuple[str, str, int]:
    if not isinstance(value, str) or "!" not in value:
        raise ValueError(f"{label} needs a finite sheet-qualified source_ref")
    sheet, local = _qualified_address(value)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(local.replace("$", ""))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} has an unsupported source_ref {value!r}") from exc
    if None in (min_col, min_row, max_col, max_row):
        raise ValueError(f"{label} source_ref must have finite bounds")
    dimensions = 0 if min_col == max_col and min_row == max_row else 1 if min_col == max_col or min_row == max_row else 2
    return sheet, local.replace("$", ""), dimensions


def _validate_mapped_demand(variables: dict[str, dict[str, Any]],
                            source_mappings: list[dict[str, Any]],
                            array_mappings: list[dict[str, Any]]) -> None:
    """Ensure every mapped formula coordinate is included in its variable's demand."""
    mappings_by_id = {item["source_family_id"]: item for item in source_mappings}

    for mapping in source_mappings:
        variable_id = mapping["variable_id"]
        variable = variables[variable_id]
        index_mapping = mapping["index_mapping"]
        for address in mapping["source_members"]:
            indices = _mapped_logical_indices(index_mapping, address)
            if not _is_demanded(variable, indices):
                raise ValueError(
                    f"mapped source member {address} requires {variable_id}{_format_indices(indices)} "
                    "outside demanded_indices"
                )

    for group in array_mappings:
        for instance in group["instances"]:
            variable_id = instance["variable_id"]
            variable = variables[variable_id]
            anchor_mapping = mappings_by_id[instance["anchor_source_family_id"]]["index_mapping"]
            for address in instance["source_members"]:
                indices = _mapped_logical_indices(anchor_mapping, address, array_ref=instance["array_ref"])
                if not _is_demanded(variable, indices):
                    raise ValueError(
                        f"mapped array member {address} requires {variable_id}{_format_indices(indices)} "
                        "outside demanded_indices"
                    )


def _mapped_logical_indices(index_mapping: dict[str, Any], address: str,
                            *, array_ref: str | None = None) -> tuple[int, ...]:
    if array_ref is None and isinstance(index_mapping.get("array_ref"), str):
        array_ref = index_mapping["array_ref"]
    sheet, local = _qualified_address(address)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(local.replace("$", ""))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"cannot check demanded index for mapped source member {address}") from exc
    if None in (min_col, min_row, max_col, max_row) or min_col != max_col or min_row != max_row:
        raise ValueError(f"mapped source member must be one cell: {address}")

    index_by_dimension: dict[int, int] = {}
    axes = index_mapping.get("axes", [])
    if not isinstance(axes, list):
        raise ValueError(f"mapped source member {address} has malformed index axes")
    array_bounds = None
    if array_ref is not None:
        try:
            array_bounds = range_boundaries(array_ref.replace("$", ""))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"mapped source member {address} has invalid array_ref {array_ref!r}") from exc
        if None in array_bounds:
            raise ValueError(f"mapped source member {address} has incomplete array_ref {array_ref!r}")

    def coordinate_index(axis: dict[str, Any]) -> int:
        source_coordinate = axis.get("source_coordinate")
        origin = axis.get("origin")
        index_origin = axis.get("index_origin")
        if not isinstance(index_origin, int) or isinstance(index_origin, bool):
            raise ValueError(f"mapped source member {address} has an invalid index origin")
        if source_coordinate == "row":
            coordinate = min_row
            if not isinstance(origin, int) or isinstance(origin, bool):
                raise ValueError(f"mapped source member {address} needs an integer row origin")
            offset = coordinate - origin
        elif source_coordinate == "column":
            coordinate = min_col
            if not isinstance(origin, str):
                raise ValueError(f"mapped source member {address} needs a column-letter origin")
            offset = coordinate - column_index_from_string(origin)
        elif source_coordinate == "array_row":
            if array_bounds is None:
                raise ValueError(f"mapped source member {address} needs its array_ref for array_row")
            coordinate = min_row - array_bounds[1]
            if not isinstance(origin, int) or isinstance(origin, bool):
                raise ValueError(f"mapped source member {address} needs an integer array-row origin")
            offset = coordinate - origin
        elif source_coordinate == "array_column":
            coordinate = min_col
            if not isinstance(origin, str):
                raise ValueError(f"mapped source member {address} needs a column-letter array origin")
            offset = coordinate - column_index_from_string(origin)
        else:
            raise ValueError(f"mapped source member {address} has unsupported source coordinate {source_coordinate!r}")
        return offset + index_origin

    for axis in axes:
        if not isinstance(axis, dict):
            raise ValueError(f"mapped source member {address} has a malformed index axis")
        dimension = axis.get("dimension")
        if not isinstance(dimension, int) or isinstance(dimension, bool) or dimension < 0:
            raise ValueError(f"mapped source member {address} has an invalid index dimension")
        index = coordinate_index(axis)
        if dimension in index_by_dimension and index_by_dimension[dimension] != index:
            raise ValueError(f"mapped source member {address} has conflicting axis indices")
        index_by_dimension[dimension] = index

    fixed_axes = index_mapping.get("fixed_axes", [])
    if not isinstance(fixed_axes, list):
        raise ValueError(f"mapped source member {address} has malformed fixed axes")
    for axis in fixed_axes:
        if not isinstance(axis, dict):
            raise ValueError(f"mapped source member {address} has a malformed fixed axis")
        dimension, index = axis.get("dimension"), axis.get("index")
        if (not isinstance(dimension, int) or isinstance(dimension, bool) or dimension < 0
                or not isinstance(index, int) or isinstance(index, bool) or index < 0):
            raise ValueError(f"mapped source member {address} has an invalid fixed-axis index")
        if dimension in index_by_dimension and index_by_dimension[dimension] != index:
            raise ValueError(f"mapped source member {address} disagrees with its fixed-axis index")
        index_by_dimension[dimension] = index

    kind = index_mapping.get("kind")
    if kind == "scalar":
        return ()
    if not index_by_dimension or set(index_by_dimension) != set(range(len(index_by_dimension))):
        raise ValueError(f"mapped source member {address} does not define every logical index")
    return tuple(index_by_dimension[dimension] for dimension in range(len(index_by_dimension)))


def _format_indices(indices: tuple[int, ...]) -> str:
    return "" if not indices else "[" + ",".join(str(value) for value in indices) + "]"


def _is_demanded(variable: dict[str, Any], indices: tuple[int, ...]) -> bool:
    shape = variable["shape"]
    if len(indices) != len(shape) or any(index < 0 or index >= size for index, size in zip(indices, shape)):
        return False
    demand = variable["demanded_indices"]
    if demand == "all":
        return True
    if isinstance(demand, dict):
        if {"start", "stop_exclusive"}.issubset(demand):
            if len(indices) != 1:
                return False
            start, stop, step = demand.get("start"), demand.get("stop_exclusive"), demand.get("step", 1)
            return (all(isinstance(value, int) and not isinstance(value, bool) for value in (start, stop, step))
                    and step > 0 and start <= indices[0] < stop and (indices[0] - start) % step == 0)
        axes = variable["axes"]
        for dimension, axis in enumerate(axes):
            axis_name = axis.get("name") if isinstance(axis, dict) else None
            if axis_name not in demand or not _demand_contains(demand[axis_name], indices[dimension]):
                return False
        return True
    if isinstance(demand, list):
        if not indices:
            return not demand
        if len(indices) == 1:
            return indices[0] in demand
        return list(indices) in demand or tuple(indices) in demand
    return False


def _variable_has_demand(variable: dict[str, Any]) -> bool:
    shape = variable["shape"]
    demand = variable["demanded_indices"]
    if not shape:
        return _is_demanded(variable, ())
    if demand == "all":
        return all(size > 0 for size in shape)
    if isinstance(demand, list):
        return bool(demand)
    if isinstance(demand, dict) and {"start", "stop_exclusive"}.issubset(demand):
        start, stop = demand.get("start"), demand.get("stop_exclusive")
        return isinstance(start, int) and isinstance(stop, int) and stop > start
    if isinstance(demand, dict):
        for size, axis in zip(shape, variable["axes"]):
            axis_name = axis.get("name") if isinstance(axis, dict) else None
            axis_demand = demand.get(axis_name)
            if axis_name is None or not any(_demand_contains(axis_demand, index) for index in range(size)):
                return False
        return True
    return False


def _demand_contains(demand: Any, index: int) -> bool:
    if isinstance(demand, int) and not isinstance(demand, bool):
        return index == demand
    if isinstance(demand, str):
        if ".." in demand:
            start_text, stop_text = demand.split("..", 1)
            try:
                start, inclusive_stop = int(start_text), int(stop_text)
            except ValueError:
                return False
            return start <= index <= inclusive_stop
        try:
            return index == int(demand)
        except ValueError:
            return False
    if isinstance(demand, dict):
        start, stop, step = demand.get("start"), demand.get("stop_exclusive"), demand.get("step", 1)
        return (all(isinstance(value, int) and not isinstance(value, bool) for value in (start, stop, step))
                and step > 0 and start <= index < stop and (index - start) % step == 0)
    if isinstance(demand, list):
        return index in demand
    return False


def _validate_execution_segments(variables: dict[str, dict[str, Any]],
                                 source_mappings: list[dict[str, Any]],
                                 array_mappings: list[dict[str, Any]]) -> None:
    mappings_by_segment: dict[tuple[str, str], set[str]] = defaultdict(set)
    mappings_by_family: dict[str, dict[str, Any]] = {}
    for mapping in source_mappings:
        mappings_by_segment[(mapping["variable_id"], mapping["equation_segment"])].add(mapping["source_family_id"])
        mappings_by_family[mapping["source_family_id"]] = mapping
    segments_by_key: dict[tuple[str, str], dict[str, Any]] = {}
    for variable_id, variable in variables.items():
        seen: set[str] = set()
        for segment in variable["equation_segments"]:
            if not isinstance(segment, dict):
                raise ValueError(f"variable {variable_id} equation_segments must contain objects")
            segment_id = segment.get("segment_id")
            if not isinstance(segment_id, str) or not segment_id or segment_id in seen:
                raise ValueError(f"variable {variable_id} needs unique non-empty segment_id values")
            seen.add(segment_id)
            if segment.get("execution") not in {"vectorized", "ascending_recurrence", "scalar"}:
                raise ValueError(f"segment {variable_id}.{segment_id} has unsupported execution mode")
            start, stop = segment.get("index_start"), segment.get("index_stop_exclusive")
            if (not isinstance(start, int) or isinstance(start, bool) or start < 0
                    or not isinstance(stop, int) or isinstance(stop, bool) or stop <= start):
                raise ValueError(f"segment {variable_id}.{segment_id} needs a non-empty explicit index interval")
            index_axis = segment.get("index_axis")
            if index_axis == "scalar":
                if variable["shape"] or start != 0 or stop != 1:
                    raise ValueError(f"scalar segment {variable_id}.{segment_id} must use interval 0..1")
            else:
                dimensions = [position for position, axis in enumerate(variable["axes"])
                              if isinstance(axis, dict) and axis.get("name") == index_axis]
                if len(dimensions) != 1:
                    raise ValueError(f"segment {variable_id}.{segment_id} must name exactly one declared logical index axis")
                if stop > variable["shape"][dimensions[0]]:
                    raise ValueError(f"segment {variable_id}.{segment_id} index interval exceeds variable shape")
            row_start, row_stop = segment.get("source_row_start"), segment.get("source_row_stop_exclusive")
            if (not isinstance(row_start, int) or isinstance(row_start, bool) or row_start < 1
                    or not isinstance(row_stop, int) or isinstance(row_stop, bool) or row_stop <= row_start):
                raise ValueError(f"segment {variable_id}.{segment_id} needs explicit positive source-row bounds")
            declared_families = segment.get("source_family_ids")
            if not isinstance(declared_families, list) or not declared_families:
                raise ValueError(f"segment {variable_id}.{segment_id} needs source_family_ids")
            if len(declared_families) != len(set(declared_families)):
                raise ValueError(f"segment {variable_id}.{segment_id} repeats source_family_ids")
            assigned = mappings_by_segment.get((variable_id, segment_id), set())
            if set(declared_families) != assigned:
                raise ValueError(f"segment {variable_id}.{segment_id} source families do not match source_mappings")
            if segment["execution"] == "ascending_recurrence":
                if segment.get("update_order") != "ascending" or segment.get("snapshot_before_update") is not True:
                    raise ValueError(f"recurrence segment {variable_id}.{segment_id} must declare ascending snapshot updates")
                if not isinstance(segment.get("recurrence_group"), str) or not segment["recurrence_group"]:
                    raise ValueError(f"recurrence segment {variable_id}.{segment_id} needs a recurrence_group")
            elif segment.get("snapshot_before_update") not in {True, False}:
                raise ValueError(f"segment {variable_id}.{segment_id} must state snapshot_before_update explicitly")
            lagged = segment.get("lagged_dependencies", [])
            if not isinstance(lagged, list):
                raise ValueError(f"segment {variable_id}.{segment_id} lagged_dependencies must be a list")
            declared_equations = segment.get("equation_family_ids")
            if not isinstance(declared_equations, list) or not declared_equations:
                raise ValueError(f"segment {variable_id}.{segment_id} needs equation_family_ids")
            mapped_equations = {item["equation_family_id"] for item in source_mappings
                                if item["variable_id"] == variable_id and item["equation_segment"] == segment_id}
            if set(declared_equations) != mapped_equations:
                raise ValueError(f"segment {variable_id}.{segment_id} equation families do not match source_mappings")
            axis_intervals = segment.get("axis_intervals")
            if axis_intervals is not None:
                if not isinstance(axis_intervals, list):
                    raise ValueError(f"segment {variable_id}.{segment_id} axis_intervals must be a list")
                interval_dimensions: set[int] = set()
                axis_names = {axis.get("name"): position for position, axis in enumerate(variable["axes"])
                              if isinstance(axis, dict)}
                for interval in axis_intervals:
                    if not isinstance(interval, dict):
                        raise ValueError(f"segment {variable_id}.{segment_id} has a malformed axis interval")
                    dimension, axis_name = interval.get("dimension"), interval.get("axis")
                    interval_start, interval_stop = interval.get("start"), interval.get("stop_exclusive")
                    if (not isinstance(dimension, int) or isinstance(dimension, bool)
                            or dimension < 0 or dimension >= len(variable["shape"])
                            or axis_names.get(axis_name) != dimension
                            or not isinstance(interval_start, int) or isinstance(interval_start, bool)
                            or interval_start < 0
                            or not isinstance(interval_stop, int) or isinstance(interval_stop, bool)
                            or interval_stop <= interval_start or interval_stop > variable["shape"][dimension]):
                        raise ValueError(f"segment {variable_id}.{segment_id} has an invalid logical-axis interval")
                    if dimension in interval_dimensions:
                        raise ValueError(f"segment {variable_id}.{segment_id} repeats a logical-axis interval")
                    interval_dimensions.add(dimension)
                if variable["shape"] and interval_dimensions != set(range(len(variable["shape"]))):
                    raise ValueError(f"segment {variable_id}.{segment_id} must bound every logical axis")
                if not variable["shape"] and axis_intervals:
                    raise ValueError(f"scalar segment {variable_id}.{segment_id} cannot declare axis intervals")
                if variable["shape"]:
                    primary_dimension = next(position for position, axis in enumerate(variable["axes"])
                                             if isinstance(axis, dict) and axis.get("name") == index_axis)
                    primary = next(item for item in axis_intervals if item["dimension"] == primary_dimension)
                    if (primary["start"], primary["stop_exclusive"]) != (start, stop):
                        raise ValueError(f"segment {variable_id}.{segment_id} primary interval disagrees with axis_intervals")
            fixed_axis_indices = segment.get("fixed_axis_indices", [])
            if not isinstance(fixed_axis_indices, list):
                raise ValueError(f"segment {variable_id}.{segment_id} fixed_axis_indices must be a list")
            for fixed in fixed_axis_indices:
                if (not isinstance(fixed, dict) or not isinstance(fixed.get("dimension"), int)
                        or isinstance(fixed.get("dimension"), bool)
                        or not isinstance(fixed.get("index"), int) or isinstance(fixed.get("index"), bool)
                        or fixed["dimension"] < 0 or fixed["dimension"] >= len(variable["shape"])
                        or fixed["index"] < 0 or fixed["index"] >= variable["shape"][fixed["dimension"]]):
                    raise ValueError(f"segment {variable_id}.{segment_id} has an invalid fixed-axis index")
            segments_by_key[(variable_id, segment_id)] = segment

    def validate_member(variable_id: str, segment: dict[str, Any], mapping: dict[str, Any],
                        address: str, *, array_ref: str | None = None) -> None:
        family_id = mapping["source_family_id"]
        equation_id = mapping["equation_family_id"]
        segment_id = segment["segment_id"]
        if family_id not in segment["source_family_ids"]:
            raise ValueError(f"mapped member {address} family {family_id} is outside segment {variable_id}.{segment_id}")
        if equation_id not in segment["equation_family_ids"]:
            raise ValueError(f"mapped member {address} equation {equation_id} is outside segment {variable_id}.{segment_id}")
        indices = _mapped_logical_indices(mapping["index_mapping"], address, array_ref=array_ref)
        if segment["index_axis"] != "scalar":
            axis_positions = [position for position, axis in enumerate(variables[variable_id]["axes"])
                              if isinstance(axis, dict) and axis.get("name") == segment["index_axis"]]
            if len(axis_positions) != 1:
                raise ValueError(f"segment {variable_id}.{segment_id} index_axis is not unique")
            index = indices[axis_positions[0]]
            if not segment["index_start"] <= index < segment["index_stop_exclusive"]:
                raise ValueError(
                    f"mapped member {address} index {index} is outside segment {variable_id}.{segment_id} "
                    f"interval {segment['index_start']}..{segment['index_stop_exclusive']}"
                )
        elif indices:
            raise ValueError(f"scalar segment {variable_id}.{segment_id} cannot contain indexed member {address}")
        for fixed in segment.get("fixed_axis_indices", []):
            dimension, fixed_index = fixed.get("dimension"), fixed.get("index")
            if not isinstance(dimension, int) or dimension >= len(indices) or indices[dimension] != fixed_index:
                raise ValueError(f"mapped member {address} violates a fixed-axis constraint in {variable_id}.{segment_id}")
        for interval in segment.get("axis_intervals", []):
            dimension, start, stop = interval["dimension"], interval["start"], interval["stop_exclusive"]
            if not start <= indices[dimension] < stop:
                raise ValueError(
                    f"mapped member {address} index {indices[dimension]} is outside declared axis "
                    f"{interval['axis']} interval {start}..{stop} in {variable_id}.{segment_id}"
                )

    for mapping in source_mappings:
        variable_id, segment_id = mapping["variable_id"], mapping["equation_segment"]
        segment = segments_by_key.get((variable_id, segment_id))
        if segment is None:
            raise ValueError(f"source family {mapping['source_family_id']} references an unknown variable segment")
        for address in mapping["source_members"]:
            validate_member(variable_id, segment, mapping, address)

    for group in array_mappings:
        for instance in group["instances"]:
            mapping = mappings_by_family.get(instance["anchor_source_family_id"])
            if mapping is None:
                raise ValueError(f"array anchor family is not mapped: {instance['anchor_source_family_id']}")
            variable_id = instance["variable_id"]
            if mapping["variable_id"] != variable_id:
                raise ValueError(f"array anchor maps to a different variable than {instance['anchor']}")
            segment = segments_by_key.get((variable_id, mapping["equation_segment"]))
            if segment is None:
                raise ValueError(f"array anchor references an unknown variable segment: {instance['anchor']}")
            for address in instance["source_members"]:
                validate_member(variable_id, segment, mapping, address, array_ref=instance["array_ref"])


def _validate_execution_plan(value: Any, modules: list[dict[str, Any]], source_family_ids: set[str],
                             variables: dict[str, dict[str, Any]],
                             source_mappings: list[dict[str, Any]]) -> dict[str, Any]:
    if not isinstance(value, dict) or not isinstance(value.get("passes"), list):
        raise ValueError("modular semantic map needs an execution_plan with ordered passes")
    passes = value["passes"]
    module_ids = {item["module_id"] for item in modules}
    mapped_modules = {item["source_family_id"]: item["module_id"] for item in source_mappings}
    seen_passes: set[str] = set()
    seen_families: set[str] = set()
    seen_modules: set[str] = set()
    orders: list[int] = []
    for item in passes:
        if (not isinstance(item, dict) or not isinstance(item.get("pass_id"), str) or not item["pass_id"].strip()
                or not isinstance(item.get("order"), int) or isinstance(item.get("order"), bool)
                or not isinstance(item.get("module_id"), str)
                or not isinstance(item.get("purpose"), str) or not item["purpose"].strip()):
            raise ValueError("each execution pass needs pass_id, integer order, module_id, and purpose")
        if item["pass_id"] in seen_passes:
            raise ValueError(f"duplicate execution pass_id {item['pass_id']}")
        seen_passes.add(item["pass_id"])
        if item["module_id"] not in module_ids:
            raise ValueError(f"execution pass references unknown module {item['module_id']}")
        orders.append(item["order"])
        seen_modules.add(item["module_id"])
        pass_families = item.get("source_family_ids")
        if not isinstance(pass_families, list) or any(not isinstance(family_id, str) for family_id in pass_families):
            raise ValueError(f"execution pass {item['pass_id']} needs a source_family_ids list")
        if len(pass_families) != len(set(pass_families)):
            raise ValueError(f"execution pass {item['pass_id']} repeats a source family")
        for family_id in pass_families:
            if family_id not in source_family_ids:
                raise ValueError(f"execution pass {item['pass_id']} references unknown source family {family_id}")
            if family_id in seen_families:
                raise ValueError(f"source family {family_id} appears in more than one execution pass")
            if mapped_modules.get(family_id) != item["module_id"]:
                raise ValueError(f"source family {family_id} execution module differs from its source mapping")
            seen_families.add(family_id)
        pass_variables = item.get("variable_ids")
        if not isinstance(pass_variables, list) or any(not isinstance(variable_id, str) for variable_id in pass_variables):
            raise ValueError(f"execution pass {item['pass_id']} needs a variable_ids list")
        if len(pass_variables) != len(set(pass_variables)):
            raise ValueError(f"execution pass {item['pass_id']} repeats a variable")
        if set(pass_variables) - set(variables):
            raise ValueError(f"execution pass {item['pass_id']} references an unknown variable")
    if orders != sorted(orders) or len(set(orders)) != len(orders):
        raise ValueError("execution passes must have unique increasing order values")
    if seen_modules != module_ids:
        raise ValueError("execution passes must include every declared module at least once")
    if seen_families != source_family_ids:
        raise ValueError("execution passes must assign every source family exactly once")
    constraints = value.get("constraints")
    if not isinstance(constraints, list) or any(not isinstance(item, str) or not item.strip() for item in constraints):
        raise ValueError("execution_plan must state its execution constraints")
    return dict(value)


def _validate_equation_families(value: Any, modules: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    if not isinstance(value, list) or not value:
        raise ValueError("equation_families must be a non-empty list")
    module_ids = {item["module_id"] for item in modules}
    result: dict[str, dict[str, Any]] = {}
    for item in value:
        if not isinstance(item, dict):
            raise ValueError("each equation family must be an object")
        family_id = item.get("family_id")
        if not isinstance(family_id, str) or not family_id or family_id in result:
            raise ValueError("equation family IDs must be unique non-empty strings")
        if item.get("module_id") not in module_ids or not isinstance(item.get("function_name"), str) or not item["function_name"].strip():
            raise ValueError(f"equation family {family_id} needs a known module and function_name")
        if not isinstance(item.get("description"), str) or not item["description"].strip():
            raise ValueError(f"equation family {family_id} needs a readable description")
        if not isinstance(item.get("source_family_ids"), list) or not item["source_family_ids"]:
            raise ValueError(f"equation family {family_id} needs source_family_ids")
        result[family_id] = dict(item)
    return result


def _validate_source_mappings(value: Any, source_families: dict[str, dict[str, Any]],
                              equation_families: dict[str, dict[str, Any]],
                              variables: dict[str, dict[str, Any]],
                              modules: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        raise ValueError("source_mappings must be a list")
    module_ids = {item["module_id"] for item in modules}
    by_family: dict[str, dict[str, Any]] = {}
    for item in value:
        if not isinstance(item, dict):
            raise ValueError("each source mapping must be an object")
        source_id = item.get("source_family_id")
        if source_id not in source_families or source_id in by_family:
            raise ValueError(f"missing, unknown, or duplicate source_family_id: {source_id}")
        if item.get("equation_family_id") not in equation_families:
            raise ValueError(f"source family {source_id} refers to an unknown equation family")
        if item.get("variable_id") not in variables or item.get("module_id") not in module_ids:
            raise ValueError(f"source family {source_id} refers to an unknown variable or module")
        if item["module_id"] != equation_families[item["equation_family_id"]]["module_id"]:
            raise ValueError(f"source family {source_id} and its equation family use different modules")
        if not isinstance(item.get("equation_segment"), str) or not item["equation_segment"].strip():
            raise ValueError(f"source family {source_id} needs an equation_segment")
        if not isinstance(item.get("index_mapping"), dict) or not item["index_mapping"]:
            raise ValueError(f"source family {source_id} needs an explicit index_mapping object")
        by_family[source_id] = dict(item)
    for family_id, family in equation_families.items():
        assigned = sorted(source_id for source_id, item in by_family.items()
                          if item["equation_family_id"] == family_id)
        declared = family["source_family_ids"]
        if len(declared) != len(set(declared)) or sorted(declared) != assigned:
            raise ValueError(f"equation family {family_id} source_family_ids do not match source_mappings")
    return [by_family[key] for key in sorted(by_family)]


def _validate_array_mappings(value: Any, trace_arrays: dict[tuple[str, str], dict[str, Any]],
                             array_members: dict[tuple[str, str], list[str]],
                             variables: dict[str, dict[str, Any]],
                             modules: list[dict[str, Any]],
                             trace_ordinary: dict[str, dict[str, Any]],
                             profile_members: dict[str, str],
                             source_mappings: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        raise ValueError("array_mappings must be a list")
    module_ids = {item["module_id"] for item in modules}
    seen_instances: set[tuple[str, str]] = set()
    family_ids: set[str] = set()
    result: list[dict[str, Any]] = []
    for group in value:
        if not isinstance(group, dict):
            raise ValueError("each array mapping must be an object")
        family_id = group.get("array_family_id")
        if not isinstance(family_id, str) or not family_id or family_id in family_ids:
            raise ValueError("array_family_id values must be unique non-empty strings")
        family_ids.add(family_id)
        instances = group.get("instances")
        if not isinstance(instances, list) or not instances:
            raise ValueError(f"array family {family_id} needs at least one instance")
        normalized_instances: list[dict[str, Any]] = []
        for instance in instances:
            if not isinstance(instance, dict):
                raise ValueError(f"array family {family_id} contains a non-object instance")
            anchor = instance.get("anchor")
            sheet, local_anchor = _qualified_address(anchor)
            array_ref = str(instance.get("array_ref", "")).replace("$", "").upper()
            key = (sheet.casefold(), array_ref)
            if key not in trace_arrays or key in seen_instances:
                raise ValueError(f"missing, unknown, or duplicate active array instance: {anchor} {array_ref}")
            expected_instance = trace_arrays[key]
            expected_anchor, expected_shape = _array_geometry(expected_instance["sheet"], expected_instance["array_ref"])
            if local_anchor.upper() != expected_anchor.split("!", 1)[1].upper():
                raise ValueError(f"array anchor does not match the active trace: {anchor}")
            anchor_key = _address_key(expected_anchor)
            anchor_family_id = profile_members.get(anchor_key)
            anchor_mapping = source_mappings.get(anchor_family_id)
            if anchor_key not in trace_ordinary or anchor_mapping is None:
                raise ValueError(f"array anchor is not mapped as an active ordinary formula: {anchor}")
            if instance.get("anchor_source_family_id") != anchor_family_id:
                raise ValueError(f"array instance {anchor} must cite its mapped anchor source_family_id")
            if instance.get("shape") != expected_shape:
                raise ValueError(f"array shape does not match the active trace: {anchor}")
            if not isinstance(instance.get("orientation"), str) or not instance["orientation"].strip():
                raise ValueError(f"array instance {anchor} needs an orientation")
            if instance.get("variable_id") not in variables or instance.get("module_id") not in module_ids:
                raise ValueError(f"array instance {anchor} refers to an unknown variable or module")
            if anchor_mapping["variable_id"] != instance["variable_id"]:
                raise ValueError(f"array anchor {anchor} and its follower mapping use different variables")
            if anchor_mapping["module_id"] != instance["module_id"]:
                raise ValueError(f"array anchor {anchor} and its follower mapping use different modules")
            axes = instance.get("axes")
            if not isinstance(axes, list) or len(axes) != len(expected_shape):
                raise ValueError(f"array instance {anchor} needs one logical axis per dimension")
            members = sorted(array_members[key], key=_address_key)
            area = expected_shape[0] * expected_shape[1]
            if not members or len(members) >= area:
                raise ValueError(f"active array instance {anchor} has an invalid active follower subset")
            normalized_instances.append({**instance, "anchor": expected_anchor,
                                         "array_ref": expected_instance["array_ref"],
                                         "shape": expected_shape, "source_formula": expected_instance["formula"],
                                         "source_formula_fingerprint": workflow._hash_bytes(
                                             str(expected_instance["formula"]).encode("utf-8")),
                                         "source_members": members, "member_count": len(members)})
            seen_instances.add(key)
        result.append({**group, "instances": normalized_instances})
    if seen_instances != set(trace_arrays):
        raise ValueError(f"array mapping coverage is incomplete ({len(set(trace_arrays) - seen_instances)} active instances unmapped)")
    return sorted(result, key=lambda item: item["array_family_id"])


def _source_formula_fingerprint(family: dict[str, Any]) -> str:
    value = {key: family.get(key) for key in ("sheet", "column", "normalized_formula", "representative_formula")}
    return workflow._hash_bytes(workflow._json_bytes(value))


def _address_key(address: Any) -> str:
    if not isinstance(address, str) or "!" not in address:
        raise ValueError(f"source address must be sheet-qualified: {address!r}")
    sheet, local = _qualified_address(address)
    return f"{sheet.casefold()}!{local.replace('$', '').upper()}"


def _qualified_address(address: Any) -> tuple[str, str]:
    if not isinstance(address, str) or "!" not in address:
        raise ValueError(f"source address must be sheet-qualified: {address!r}")
    sheet, local = address.rsplit("!", 1)
    sheet = sheet.strip()
    if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, local.replace("$", "")


def _array_geometry(sheet: str, address: str) -> tuple[str, list[int]]:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(address)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"unsupported array range {sheet}!{address}") from exc
    if None in (min_col, min_row, max_col, max_row):
        raise ValueError(f"array range must be finite: {sheet}!{address}")
    anchor = f"{sheet}!{get_column_letter(min_col)}{min_row}"
    return anchor, [max_row - min_row + 1, max_col - min_col + 1]


def _plan_markdown(plan: dict[str, Any]) -> str:
    coverage = plan["coverage"]
    classification = plan["source_classification"]
    candidate = plan.get("selection_basis") == "static_candidate"
    if candidate:
        coverage_lines = [
            f"- Static candidate formula positions mapped: {coverage['static_candidate_formula_members_mapped']} / {coverage['static_candidate_formula_members']}",
            f"- Static candidate array followers mapped: {coverage['static_candidate_array_followers_mapped']} / {coverage['static_candidate_array_followers']}",
            f"- Static candidate members mapped: {coverage['static_candidate_member_count_mapped']} / {coverage['static_candidate_member_count']}",
            "- Candidate closure: **unknown**; this profile is a static branch-union, not an active trace.",
            f"- Semantic equation families: {coverage['semantic_family_count']} from {coverage['source_family_count']} source candidates",
            f"- Candidate array instances: {coverage['candidate_array_instance_count']} in {coverage['array_family_count']} array families",
        ]
        summary = "This artifact maps the checked static source-candidate formula set. Both readable branches and lookup candidate ranges may be present; it does not establish active formulas, closure, or semantic correctness."
    else:
        coverage_lines = [
            f"- Ordinary formulas mapped: {coverage['ordinary_formula_members_mapped']} / {coverage['ordinary_formula_members']}",
            f"- Array followers mapped: {coverage['array_formula_followers_mapped']} / {coverage['array_formula_followers']}",
            f"- Active formula members mapped: {coverage['active_formula_members_mapped']} / {coverage['active_formula_members']}",
            f"- Semantic equation families: {coverage['semantic_family_count']} from {coverage['source_family_count']} source candidates",
            f"- Array instances: {coverage['array_instance_count']} in {coverage['array_family_count']} array families",
        ]
        summary = "This artifact maps the checked historical active formula set to named variables, modules, and equation families. It does not execute formulas or prove semantic correctness."
    lines = ["# Step 3 Semantic Map", "", "Status: **Complete static mapping draft; review required**", "",
             f"Source `{plan['source']['source_id']}` · run `{plan['source']['run_id']}` · SHA-256 `{plan['source']['source_sha256']}`",
             f"Binding SHA-256 `{plan['source']['binding_sha256']}`", "",
             f"Selection basis: `{plan['selection_basis']}`. {summary}", "",
             "## Coverage", "", *coverage_lines, "",
             ]
    if plan.get("candidate_scope") is not None:
        lines.extend([f"Candidate scope: `{plan['candidate_scope']}`; omitted source references remain frontiers and closure stays **unknown**.",
                      f"Scope frontier records: {coverage['scope_frontier_count']}; proposed formula-start members: {coverage['proposed_formula_start_count']}.", ""])
    lines.extend([
             "## Ordered modules", ""]
    )
    lines.extend(f"{item['order']}. **{item['module_id']}** — {item['purpose']}" for item in plan["modules"])
    lines.extend(["", "## Variable catalog", ""])
    lines.extend(f"- `{item['variable_id']}` · {item['role']} {item['kind']} · shape {item['shape']} · axes {', '.join(str(axis.get('name', axis)) if isinstance(axis, dict) else str(axis) for axis in item['axes'])}"
                 for item in plan["variables"].values())
    lines.extend(["", "## Source kinds and demand", "",
                  f"- Raw source extents checked: {classification['raw_extent_count']} extents / {classification['distinct_raw_extent_coordinate_count']} distinct coordinates",
                  f"- Raw seed variables/extents: {classification['raw_seed_variable_count']} / {classification['raw_seed_extent_count']}",
                  f"- Source-verified blank table-boundary inputs: {classification.get('verified_source_blank_boundary_count', 0)} (separately declared, never formula caches)",
                  f"- Formula and array-follower coordinates outside raw demand: {classification['formula_coordinates_outside_raw_demand']} / {classification['array_follower_coordinates_outside_raw_demand']}",
                  "- Demanded formula or array-follower coordinates declared raw: 0",
                  f"- Role/kind counts: `{workflow._json_bytes(classification['variables_by_role_and_kind']).decode('utf-8')}`",
                  "- Formula classification uses promoted source formula markers and native array ranges; cached results are not consulted."])
    lines.extend(["", "## Review boundary", "",
                  "A full mapping is a bookkeeping result only. Independent Agent review and the workflow's authorized acceptance reviewer must confirm the family equations, variable/index rules, recurrence segments, axes, and array orientation before generation. The default acceptance reviewer is the human; a source-bound TypeSafe delegation applies only when explicitly registered.",
                  "Generation-ready, runtime-verified, native-reconciled, and GPU-executable are all false.", ""])
    return "\n".join(lines)


def _check_markdown(check: dict[str, Any]) -> str:
    coverage = check["coverage"]
    classification = check["source_classification"]
    if check.get("selection_basis") == "static_candidate":
        lines = [
            "# Step 3 Semantic Map Check", "", "Status: **Pass — static candidate mapping coverage only**", "",
            "Selection basis: `static_candidate`; candidate closure remains **unknown**.",
            f"- Candidate formula positions: {coverage['static_candidate_formula_members_mapped']} / {coverage['static_candidate_formula_members']}",
            f"- Candidate array followers: {coverage['static_candidate_array_followers_mapped']} / {coverage['static_candidate_array_followers']}",
            f"- Unmapped candidate members: {coverage['unmapped_static_candidate_members']}",
            f"- Raw extents: {classification['raw_extent_count']} checked; demanded formula/array-follower overlaps: 0",
            f"- Raw seed variables/extents: {classification['raw_seed_variable_count']} / {classification['raw_seed_extent_count']}",
            f"- Source-verified blank table-boundary inputs: {classification.get('verified_source_blank_boundary_count', 0)}",
            f"- Formula/follower coordinates outside raw demand: {classification['formula_coordinates_outside_raw_demand']} / {classification['array_follower_coordinates_outside_raw_demand']}",
            "- Active formula selection or closure proven: no", "- Semantic correctness reviewed: no",
            "- Generation ready: no", "- Runtime verified: no", "- Native reconciled: no", "",
        ]
        if check.get("candidate_scope") is not None:
            lines.extend([f"- Candidate scope: `{check['candidate_scope']}`",
                          f"- Scope SHA-256: `{check['model_scope_sha256']}`",
                          "- Scope frontiers remain unresolved; no active closure is claimed."])
        return "\n".join(lines)
    return "\n".join(["# Step 3 Semantic Map Check", "", "Status: **Pass — static mapping coverage only**", "",
                      f"- Ordinary formulas: {coverage['ordinary_formula_members_mapped']} / {coverage['ordinary_formula_members']}",
                      f"- Array followers: {coverage['array_formula_followers_mapped']} / {coverage['array_formula_followers']}",
                      f"- Active members: {coverage['active_formula_members_mapped']} / {coverage['active_formula_members']}",
                      f"- Unmapped members: {coverage['unmapped_active_formula_members']}",
                      f"- Raw extents: {classification['raw_extent_count']} checked; demanded formula/array-follower overlaps: 0",
                      f"- Raw seed variables/extents: {classification['raw_seed_variable_count']} / {classification['raw_seed_extent_count']}",
                      f"- Source-verified blank table-boundary inputs: {classification.get('verified_source_blank_boundary_count', 0)}",
                      f"- Formula/follower coordinates outside raw demand: {classification['formula_coordinates_outside_raw_demand']} / {classification['array_follower_coordinates_outside_raw_demand']}",
                      "- Semantic correctness reviewed: no", "- Generation ready: no", "- Runtime verified: no", ""])
