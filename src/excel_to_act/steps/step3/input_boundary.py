"""Reviewable Step 3 input boundary and source-only dependency cut proof."""

from __future__ import annotations

from collections import defaultdict, deque
from itertools import product
import hashlib
import json
from math import prod
from pathlib import Path
import re
from typing import Any
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from excel_to_act.steps import conversion_workflow as conversion
from excel_to_act.steps.step3 import workflow as step3


_TARGETS_SCHEMA = "step3.input_targets.v1"
_CATALOG_SCHEMA = "step3.input_boundary.input.v1"
_BOUNDARY_SCHEMA = "step3.input_boundary.v1"
_SAMPLE_SCHEMA = "step3.sample_boundary.v1"
_VALUE_KEYS = {"value", "values", "cached_value", "cached_values", "formula_cache_values"}


def _reject_unlisted_fields(value: Any, allowed: set[str], label: str) -> None:
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be an object")
    extra = sorted(set(value) - allowed)
    if extra:
        raise ValueError(f"{label} contains unsupported fields: {', '.join(extra)}")


def _validate_scope_notes(notes: Any, source_sha256: str) -> None:
    _reject_unlisted_fields(notes, {
        "note", "non_retained_historical_candidates", "stored_coordinate_reconciliation",
        "historical_raw_candidate_sha256", "historical_trace_candidate", "excluded_formula_extents",
    }, "source_scope_notes")
    if "note" in notes and not isinstance(notes["note"], str):
        raise ValueError("source_scope_notes.note must be text")
    candidates = notes.get("non_retained_historical_candidates", [])
    if not isinstance(candidates, list) or any(not isinstance(item, str) for item in candidates):
        raise ValueError("source_scope_notes.non_retained_historical_candidates must be a list of addresses")
    if "historical_raw_candidate_sha256" in notes and not isinstance(notes["historical_raw_candidate_sha256"], str):
        raise ValueError("source_scope_notes.historical_raw_candidate_sha256 must be text")
    if "historical_trace_candidate" in notes:
        trace = notes["historical_trace_candidate"]
        _reject_unlisted_fields(trace, {"path", "sha256", "use"}, "historical_trace_candidate")
        if any(not isinstance(trace.get(key), str) for key in ("path", "sha256", "use")):
            raise ValueError("historical_trace_candidate path, sha256, and use must be text")
    excluded = notes.get("excluded_formula_extents", [])
    if not isinstance(excluded, list):
        raise ValueError("source_scope_notes.excluded_formula_extents must be a list")
    for extent in excluded:
        _reject_unlisted_fields(extent, {"sheet", "range", "reason"}, "excluded formula extent")
        if any(not isinstance(extent.get(key), str) for key in ("sheet", "range", "reason")):
            raise ValueError("excluded formula extent sheet, range, and reason must be text")
    if "stored_coordinate_reconciliation" in notes:
        reconciliation = notes["stored_coordinate_reconciliation"]
        _reject_unlisted_fields(reconciliation, {
            "historical_empty_coordinates", "inspection", "interpretation", "source_copy_sha256",
            "stored_coordinate_audit",
        }, "stored_coordinate_reconciliation")
        for key in ("inspection", "interpretation", "source_copy_sha256"):
            if key in reconciliation and not isinstance(reconciliation[key], str):
                raise ValueError(f"stored_coordinate_reconciliation.{key} must be text")
        empty_coordinates = reconciliation.get("historical_empty_coordinates", [])
        if not isinstance(empty_coordinates, list):
            raise ValueError("historical_empty_coordinates must be a list")
        for coordinate in empty_coordinates:
            _reject_unlisted_fields(coordinate, {"address", "data_type", "source_value_status"},
                                    "historical empty coordinate")
            if any(not isinstance(coordinate.get(key), str)
                   for key in ("address", "data_type", "source_value_status")):
                raise ValueError("historical empty coordinate fields must be text")
        if "stored_coordinate_audit" in reconciliation:
            audit = reconciliation["stored_coordinate_audit"]
            _reject_unlisted_fields(audit, {
                "path", "sha256", "source_sha256", "current_retained_raw_coordinate_count",
                "old_trace_raw_coordinate_count", "empty_source_coordinate_count", "coordinates",
            }, "stored_coordinate_audit")
            for key in ("path", "sha256", "source_sha256"):
                if not isinstance(audit.get(key), str) or not audit[key].strip():
                    raise ValueError(f"stored_coordinate_audit.{key} must be non-empty text")
            for key in ("current_retained_raw_coordinate_count", "old_trace_raw_coordinate_count",
                        "empty_source_coordinate_count"):
                value = audit.get(key)
                if isinstance(value, bool) or not isinstance(value, int) or value < 0:
                    raise ValueError(f"stored_coordinate_audit.{key} must be a non-negative integer")
            if not isinstance(audit.get("coordinates"), list):
                raise ValueError("stored_coordinate_audit.coordinates must be a list")
            for coordinate in audit["coordinates"]:
                _reject_unlisted_fields(coordinate, {
                    "address", "classification", "package_part", "later_design_requirement",
                    "formula_present", "literal_content_present", "stored_cell_element_present",
                }, "stored coordinate audit entry")
                for key in ("address", "classification", "package_part", "later_design_requirement"):
                    if not isinstance(coordinate.get(key), str):
                        raise ValueError(f"stored coordinate audit {key} must be text")
                for key in ("formula_present", "literal_content_present", "stored_cell_element_present"):
                    if not isinstance(coordinate.get(key), bool):
                        raise ValueError(f"stored coordinate audit {key} must be boolean")
            audit_path = Path(audit["path"]).expanduser().resolve()
            if (not _is_sha256(audit["sha256"]) or audit["source_sha256"] != source_sha256
                    or not audit_path.is_file() or conversion.hash_file(audit_path) != audit["sha256"]):
                raise ValueError("stored_coordinate_audit file or source identity does not match")
            audit_file = conversion.read_json(audit_path)
            if (not isinstance(audit_file, dict) or any(
                    audit_file.get(key) != audit[key] for key in (
                        "source_sha256", "current_retained_raw_coordinate_count",
                        "old_trace_raw_coordinate_count", "empty_source_coordinate_count", "coordinates"))):
                raise ValueError("stored_coordinate_audit content does not match the declared audit")


def _with_verified_stored_blanks(catalog: dict[str, Any], fields: dict[str, Any],
                                 inventory: dict[str, Any], source_path: Path,
                                 source_sha256: str) -> dict[str, Any]:
    reconciliation = catalog.get("source_scope_notes", {}).get("stored_coordinate_reconciliation", {})
    blanks = [item for item in reconciliation.get("historical_empty_coordinates", [])
              if item.get("source_value_status") == "stored_blank_no_literal"]
    if not blanks:
        return fields
    if (reconciliation.get("source_copy_sha256") != source_sha256 or not source_path.is_file()
            or conversion.hash_file(source_path) != source_sha256):
        raise ValueError("stored blank source checksum does not match the bound workbook")
    existing = {_address_key(f"{field['sheet']}!{member['address']}")
                for field in fields["fields"] for member in field["members"]}
    parts = {str(sheet["name"]).casefold(): next((
        (cell.get("source_location") or {}).get("ooxml_part") for cell in sheet.get("cells", [])
        if (cell.get("source_location") or {}).get("ooxml_part")), None)
        for sheet in inventory["sheets"]}
    added: list[dict[str, Any]] = []
    seen: set[str] = set()
    with ZipFile(source_path) as package:
        for item in blanks:
            address = item.get("address")
            if not isinstance(address, str) or "!" not in address:
                raise ValueError("stored blank needs a qualified source address")
            sheet, cell = address.rsplit("!", 1)
            key = _address_key(address)
            part = parts.get(sheet.casefold())
            if key in seen or key in existing or not part or not re.fullmatch(r"[A-Za-z]+[1-9][0-9]*", cell):
                raise ValueError(f"stored blank has duplicate or unknown source coordinates: {address}")
            seen.add(key)
            element = ET.fromstring(package.read(part)).find(
                f".//{{http://schemas.openxmlformats.org/spreadsheetml/2006/main}}c[@r='{cell.upper()}']")
            if (element is None or list(element) or (element.text or "").strip()
                    or item.get("data_type") != element.get("t", "n")):
                raise ValueError(f"source cell is not a stored blank: {address}")
            added.append({"field_id": f"stored-blank:{key}", "sheet": sheet, "role": "source",
                          "members": [{"address": cell.upper(), "source_kind": "stored_blank"}]})
    return {**fields, "fields": [*fields["fields"], *added]}


def create_input_boundary(
    analysis_dir: Path,
    workflow_dir: Path,
    targets_path: Path,
    catalog_path: Path,
) -> dict[str, Any]:
    """Validate a grouped catalog and submit it as a separate Stage 3 checkpoint."""
    try:
        workflow_root, manifest = conversion.load_workflow(workflow_dir)
        conversion._require_upstream_writable(manifest)
        analysis_dir, analysis, binding, _step1_root, _inventory_path = step3._load_context(analysis_dir)
        _require_bound_workflow(analysis, binding, workflow_root, manifest)
        _stage2 = conversion.require_stage_approved(workflow_root, 2)[1]

        targets_path = targets_path.expanduser().resolve()
        catalog_path = catalog_path.expanduser().resolve()
        if not targets_path.is_file() or not catalog_path.is_file():
            raise ValueError("targets and catalog input JSON files must exist")
        targets = conversion.read_json(targets_path)
        catalog = conversion.read_json(catalog_path)
        _assert_no_embedded_values(targets, "targets")
        _assert_no_embedded_values(catalog, "catalog")
        target_selection = _validate_targets(targets, binding, analysis["binding_sha256"])
        _validate_catalog_identity(catalog, binding, target_selection)
        if "source_scope_notes" in catalog:
            _validate_scope_notes(catalog["source_scope_notes"], binding["source"]["source_sha256"])

        fields, fields_sha = step3._read_stage(analysis_dir, analysis, "fields")
        dependencies, dependencies_sha = step3._read_stage(analysis_dir, analysis, "dependencies")
        _plan, plan_sha = step3._read_stage(analysis_dir, analysis, "plan")
        source_copy = manifest.get("source", {}).get("step1_source_copy_path")
        fields = _with_verified_stored_blanks(
            catalog, fields, step3._load_inventory(_inventory_path),
            Path(source_copy) if isinstance(source_copy, str) else Path(),
            binding["source"]["source_sha256"],
        )
        topology_evidence = catalog.get("topology_evidence")
        trace_path: Path | None = None
        trace_sha: str | None = None
        if topology_evidence is not None:
            if not isinstance(topology_evidence, dict):
                raise ValueError("topology_evidence must be an object when supplied")
            trace_path = Path(topology_evidence.get("path", "")).expanduser().resolve()
            trace_sha = topology_evidence.get("sha256")
            if not trace_path.is_file() or conversion.hash_file(trace_path) != trace_sha:
                raise ValueError("topology trace is missing or its declared SHA-256 does not match")
            trace = conversion.read_json(trace_path)
            cut_report = _cut_report(trace, target_selection, catalog, fields, binding)
        else:
            cut_report = _source_candidate_report(target_selection, catalog, fields)
        groups, variables, axis_metadata, source_records, exclusions, policy = _validate_catalog(
            catalog, target_selection, fields, cut_report)
        sample_boundary = None
        if "sample_boundary" in catalog:
            sample_boundary = _validate_sample_boundary(
                catalog["sample_boundary"], variables, axis_metadata, Path(source_copy),
                binding["source"]["source_sha256"])
        if cut_report["topology_available"]:
            cut_report["input_coverage"].update({
                "assigned_raw_coordinate_count": len(cut_report["_reachable_raw_keys"]),
                "unassigned_raw_coordinate_count": 0,
                "status": "complete_for_historical_topology_candidate",
            })
        else:
            assigned_raw = sum(1 for item in variables if item["role"] == "source_raw"
                               for _ in _variable_coordinates(item))
            assigned_raw += sum(len(_extent_coordinates(extent)) for item in axis_metadata
                                for extent in item["source_extents"])
            assigned_raw += sum(len(_extent_coordinates(extent)) for item in source_records
                                for extent in item["source_extents"])
            cut_report["input_coverage"].update({
                "assigned_raw_coordinate_count": assigned_raw,
                "unassigned_raw_coordinate_count": None,
                "status": "target_closure_unproven_without_topology",
            })

        source = binding["source"]
        scenario = targets["scenario"]
        boundary_basis = {
            "source": {"source_id": source["source_id"], "run_id": source["run_id"],
                       "workbook_sha256": source["source_sha256"],
                       "analysis_binding_sha256": analysis["binding_sha256"]},
            "scenario": scenario,
            "target_selection": target_selection,
            "groups": groups,
            "variables": variables,
            "axis_metadata": axis_metadata,
            "source_records": source_records,
            "upstream_exclusions": exclusions,
            "value_source_policy": policy,
            "source_scope_notes": catalog.get("source_scope_notes", {}),
            "sample_boundary": sample_boundary,
            "topology_evidence": ({"use": "source_topology_only", "sha256": trace_sha}
                                  if trace_path is not None else None),
        }
        boundary_sha = conversion.hash_bytes(conversion.json_bytes(boundary_basis))
        counts = _input_counts(variables, source_records, axis_metadata)
        counts["retained_source_raw_coordinate_count"] = cut_report["retained_raw_source_coordinates"]
        counts["formula_derived_external_element_count"] = sum(
            prod(item["shape"]) for item in variables if item["role"] == "formula_derived_external")
        payload = {
            "schema_version": _BOUNDARY_SCHEMA,
            "tool": "step3.input-catalog",
            "phase": "input_boundary",
            "status": "ready_for_review",
            "source": boundary_basis["source"],
            "scenario": scenario,
            "target_selection": target_selection,
            "groups": groups,
            "variables": variables,
            "axis_metadata": axis_metadata,
            "source_records": source_records,
            "upstream_exclusions": exclusions,
            "open_questions": catalog.get("open_questions", []),
            "value_source_policy": policy,
            "source_scope_notes": catalog.get("source_scope_notes", {}),
            "sample_boundary": sample_boundary,
            "boundary_sha256": boundary_sha,
            "counts": counts,
            "input_coverage": cut_report["input_coverage"],
            "pruned_source_analysis": {key: value for key, value in cut_report.items() if not key.startswith("_")},
            "input_artifacts": {
                "analysis_binding_sha256": analysis["binding_sha256"],
                "fields_sha256": fields_sha,
                "dependencies_sha256": dependencies_sha,
                "plan_sha256": plan_sha,
                "targets_path": str(targets_path),
                "targets_sha256": conversion.hash_file(targets_path),
                "catalog_input_path": str(catalog_path),
                "catalog_input_sha256": conversion.hash_file(catalog_path),
                "topology_trace_path": str(trace_path) if trace_path is not None else None,
                "topology_trace_sha256": trace_sha,
            },
            "evidence_policy": {"formula_calculation_performed": False,
                                "native_excel_called": False,
                                "formula_cache_used": False,
                                "trace_values_used": False,
                                "topology_status": cut_report["status"]},
        }
        markdown = _boundary_markdown(payload)
        audit = catalog.get("source_scope_notes", {}).get("stored_coordinate_reconciliation", {}).get(
            "stored_coordinate_audit")
        revision = conversion.append_stage_artifact(
            workflow_root, 3, "input_boundary.json", payload, markdown,
            input_files={
                "analysis_json": analysis_dir / "analysis.json",
                "fields_json": analysis_dir / "fields.json",
                "dependencies_json": analysis_dir / "dependencies.json",
                "execution_plan_json": analysis_dir / "execution_plan.json",
                "targets_input": targets_path,
                "catalog_input": catalog_path,
                **({"stored_coordinate_audit": Path(audit["path"])} if audit is not None else {}),
                **({"topology_trace": trace_path} if trace_path is not None else {}),
                **({f"sample_request_{index:02d}": Path(case["path"])
                    for index, case in enumerate(sample_boundary["request_cases"])}
                   if sample_boundary is not None else {}),
                **({"sample_fixed_literal_contract": Path(sample_boundary["fixed_literal_contract"]["path"])}
                   if sample_boundary is not None else {}),
            },
            input_stages={2: _stage2["artifact"]},
        )
        return {"schema_version": _BOUNDARY_SCHEMA, "tool": "step3.input-catalog", "status": "pass",
                "workflow": str(workflow_root), "revision": revision["revision"],
                "artifact": revision["artifact"], "boundary_sha256": boundary_sha,
                "counts": counts, "pruned_source_analysis": payload["pruned_source_analysis"]}
    except Exception as exc:
        return {"schema_version": _BOUNDARY_SCHEMA, "tool": "step3.input-catalog", "status": "blocked",
                "diagnostics": [{"code": "input_boundary_invalid", "severity": "error", "message": str(exc)}]}


def _validate_sample_boundary(sample: Any, variables: list[dict[str, Any]],
                              axis_metadata: list[dict[str, Any]],
                              workbook_path: Path, source_sha256: str) -> dict[str, Any]:
    """Check an external, source-bound supported request subset of the broader catalog."""
    _reject_unlisted_fields(sample, {"schema_version", "source_workbook_sha256", "request_cases",
                                    "case_axes", "request_fields", "external_vectors",
                                    "fixed_source_fields", "fixed_source_literals",
                                    "fixed_literal_contract"}, "sample_boundary")
    if sample.get("schema_version") != _SAMPLE_SCHEMA or sample.get("source_workbook_sha256") != source_sha256:
        raise ValueError("sample boundary schema or source workbook SHA-256 does not match")
    if not workbook_path.is_file() or conversion.hash_file(workbook_path) != source_sha256:
        raise ValueError("sample boundary source workbook is missing or changed")
    cases, fields, vectors = (sample.get(key) for key in
                            ("request_cases", "request_fields", "external_vectors"))
    fixed_fields, literals = sample.get("fixed_source_fields"), sample.get("fixed_source_literals")
    if any(not isinstance(items, list) or not items for items in
           (cases, fields, vectors, fixed_fields, literals)):
        raise ValueError("sample boundary requires nonempty request, vector, and fixed-source lists")
    contract_ref = sample.get("fixed_literal_contract")
    _reject_unlisted_fields(contract_ref, {"path", "sha256"}, "fixed literal contract")
    contract_path = Path(contract_ref.get("path", "")).expanduser().resolve()
    if (not _is_sha256(contract_ref.get("sha256")) or not contract_path.is_file()
            or conversion.hash_file(contract_path) != contract_ref["sha256"]):
        raise ValueError("fixed literal contract is missing or its SHA-256 changed")
    contract = conversion.read_json(contract_path)
    if (not isinstance(contract, dict)
            or contract.get("schema_version") != "step4.private_fixed_literals.v1"
            or contract.get("source_sha256") != source_sha256
            or not _is_sha256(contract.get("implementation_plan_sha256"))
            or not _is_sha256(contract.get("interface_sha256"))
            or not isinstance(contract.get("records"), list)):
        raise ValueError("fixed literal contract has wrong schema or source binding")
    expected_literals: set[str] = set()
    for record in contract["records"]:
        if (not isinstance(record, dict) or set(record) != {"address", "indices", "value", "variable_id"}
                or not isinstance(record["address"], str)
                or _reference_bounds(record["address"]) is None
                or not isinstance(record["indices"], list)
                or not isinstance(record["variable_id"], str) or not record["variable_id"]):
            raise ValueError("fixed literal contract contains a malformed source record")
        key = _address_key(record["address"])
        if key in expected_literals:
            raise ValueError("fixed literal contract contains duplicate source addresses")
        expected_literals.add(key)
    if (any(not isinstance(item, dict) or not isinstance(item.get("address"), str) for item in literals)
            or {_address_key(item["address"]) for item in literals} != expected_literals):
        raise ValueError("fixed source literal addresses differ from the bound source closure")
    variables_by_id = {item["variable_id"]: item for item in variables}
    field_names: set[str] = set()
    field_cells: set[str] = set()
    request_ids: set[str] = set()
    for field in fields:
        _reject_unlisted_fields(field, {"name", "variable_id", "source_cells", "fixed_source_literal"},
                                "sample request field")
        name, variable_id, cells = field.get("name"), field.get("variable_id"), field.get("source_cells")
        variable = variables_by_id.get(variable_id)
        if not isinstance(name, str) or not name or name in field_names or variable is None:
            raise ValueError("sample request field name or source variable is missing or duplicated")
        if variable.get("logical_name") != name or variable.get("role") != "source_raw":
            raise ValueError(f"sample request field {name} does not match a raw source variable")
        if ("fixed_source_literal" in field) != (variable.get("logical_role") == "fixed_source"):
            raise ValueError(f"sample request field {name} has inconsistent fixed-source classification")
        if not isinstance(cells, list) or not cells or any(not isinstance(cell, str) for cell in cells):
            raise ValueError(f"sample request field {name} needs exact source cells")
        keys = [_address_key(cell) for cell in cells]
        catalog_keys = {key for key, _ in _variable_coordinates(variable)}
        if len(set(keys)) != len(keys) or not set(keys) <= catalog_keys or field_cells.intersection(keys):
            raise ValueError(f"sample request field {name} has duplicate or unbound source cells")
        if variable.get("kind") == "scalar" and set(keys) != catalog_keys:
            raise ValueError(f"sample scalar field {name} must bind its exact source cell")
        field_names.add(name)
        field_cells.update(keys)
        request_ids.add(variable_id)
    axes = sample.get("case_axes")
    if not isinstance(axes, dict) or not axes or not set(axes) <= field_names:
        raise ValueError("sample case axes must name request fields")
    for name, values in axes.items():
        if not isinstance(values, list) or not values or len({json.dumps(v, sort_keys=True) for v in values}) != len(values):
            raise ValueError(f"sample case axis {name} has invalid or duplicate values")
    expected_axes = set(product(*([json.dumps(v, sort_keys=True) for v in axes[name]] for name in axes)))
    if len(cases) != len(expected_axes):
        raise ValueError("sample requests do not cover the declared Cartesian case axes")
    seen_paths: set[str] = set()
    actual_axes: set[tuple[str, ...]] = set()
    for case in cases:
        _reject_unlisted_fields(case, {"path", "sha256"}, "sample request case")
        path_value, digest = case.get("path"), case.get("sha256")
        if not isinstance(path_value, str) or not _is_sha256(digest):
            raise ValueError("sample request path or SHA-256 is invalid")
        path = Path(path_value).expanduser().resolve()
        if str(path) in seen_paths or not path.is_file() or conversion.hash_file(path) != digest:
            raise ValueError("sample request is duplicate, missing, or its SHA-256 changed")
        seen_paths.add(str(path))
        request = conversion.read_json(path)
        inputs, external = request.get("inputs"), request.get("external_data")
        if not isinstance(inputs, dict) or set(inputs) != field_names or not isinstance(external, dict):
            raise ValueError("sample request fields or external vectors do not match the declared boundary")
        if any(not isinstance(inputs[name], dict) for name in axes):
            raise ValueError("sample case axis has a malformed request field")
        actual_axes.add(tuple(json.dumps(inputs[name].get("values"), sort_keys=True) for name in axes))
        selected: set[str] = set()
        for field in fields:
            name, variable = field["name"], variables_by_id[field["variable_id"]]
            item = inputs[name]
            if not isinstance(item, dict):
                raise ValueError(f"sample request field {name} is malformed")
            if set(item) != {"axes", "values"} or not isinstance(item["axes"], list):
                raise ValueError(f"sample request field {name} has malformed values or axes")
            if variable.get("kind") == "table":
                shape = variable.get("shape")
                declared_axes = variable.get("axes")
                if (not isinstance(shape, list) or not isinstance(declared_axes, list)
                        or len(item["axes"]) != len(shape) or len(declared_axes) != len(shape)
                        or any(not isinstance(axis, dict) or set(axis) != {"name", "role", "keys"}
                               or {key: axis[key] for key in ("name", "role")} != declared_axes[index]
                               or not isinstance(axis["keys"], list) or len(axis["keys"]) != size
                               or len({json.dumps(key, sort_keys=True) for key in axis["keys"]}) != size
                               for index, (axis, size) in enumerate(zip(item["axes"], shape)))):
                    raise ValueError(f"sample table {name} has incompatible sparse-table axes or shape")
                mapping = {tuple(entry["logical_index"]): _address_key(entry["source_cell"])
                           for entry in variable.get("coordinate_map", [])}
                entries = item.get("values")
                if (not isinstance(entries, list) or any(not isinstance(entry, dict)
                        or set(entry) != {"indices", "value"} or not isinstance(entry["indices"], list)
                        or len(entry["indices"]) != len(shape)
                        or any(type(index) is not int or index < 0 or index >= size
                               for index, size in zip(entry["indices"], shape)) for entry in entries)):
                    raise ValueError(f"sample table {name} lacks sparse entries")
                selected_cells = [mapping.get(tuple(entry.get("indices", []))) for entry in entries
                                  if isinstance(entry, dict)]
                if (len(selected_cells) != len(entries) or None in selected_cells
                        or len(set(selected_cells)) != len(selected_cells)):
                    raise ValueError(f"sample table {name} has unbound sparse indices")
                selected.update(selected_cells)
            else:
                if item["axes"] != [] or isinstance(item.get("values"), (list, dict)):
                    raise ValueError(f"sample scalar {name} has a non-scalar request value")
                selected.update(_address_key(cell) for cell in field["source_cells"])
            if "fixed_source_literal" in field and item.get("values") != field["fixed_source_literal"]:
                raise ValueError(f"sample fixed request field {name} changed across cases")
        if selected != field_cells:
            raise ValueError("sample request sparse cells do not equal declared request cells")
        if set(external) != {item["name"] for item in vectors}:
            raise ValueError("sample external vector names do not match")
        input_digest = hashlib.sha256((json.dumps(
            inputs, sort_keys=True, ensure_ascii=False, separators=(",", ":")) + "\n").encode("utf-8")).hexdigest()
        for vector in vectors:
            item = external[vector["name"]]
            variable = variables_by_id.get(vector.get("variable_id"))
            if (not isinstance(item, dict) or item.get("source_sha256") != source_sha256
                    or variable is None
                    or set(item) != {"source_sha256", "input_sha256", "version", "axes", "values"}
                    or item.get("input_sha256") != input_digest
                    or item.get("version") != vector.get("producer_version")
                    or not isinstance(item.get("values"), list)
                    or len(item["values"]) != len(vector["age_keys"])
                    or item.get("axes") != [{**axis, "keys": vector["age_keys"]}
                                            for axis in variable.get("axes", [])]):
                raise ValueError(f"sample external vector {vector['name']} has wrong source or age keys")
    if actual_axes != expected_axes:
        raise ValueError("sample requests do not exercise the declared Cartesian case axes")
    vector_names: set[str] = set()
    for vector in vectors:
        _reject_unlisted_fields(vector, {"name", "variable_id", "source_extent", "age_extent",
                                         "age_keys", "producer_version"},
                                "sample external vector")
        name, variable = vector.get("name"), variables_by_id.get(vector.get("variable_id"))
        keys = vector.get("age_keys")
        key_axes = [axis for axis in variable.get("axes", []) if axis.get("role") == "key"] if variable else []
        catalog_axes = [axis for axis in axis_metadata if key_axes and axis.get("axis_id") == key_axes[0].get("axis_id")]
        expected_extent = (f"{catalog_axes[0]['source_extents'][0]['sheet']}!"
                           f"{catalog_axes[0]['source_extents'][0]['range']}") if (len(catalog_axes) == 1
                           and len(catalog_axes[0].get("source_extents", [])) == 1) else None
        if (not isinstance(name, str) or name in vector_names or variable is None
                or variable.get("variable_id") != name or variable.get("role") != "formula_derived_external"
                or not isinstance(vector.get("producer_version"), str) or not vector["producer_version"]
                or len(key_axes) != 1 or len(variable.get("axes", [])) != 1
                or not isinstance(keys, list) or not keys or expected_extent is None
                or _address_key(vector.get("age_extent", "")) != _address_key(expected_extent)
                or catalog_axes[0].get("shape") != [len(keys)]
                or [_address_key(cell) for cell in [vector.get("source_extent", "")]] !=
                [_address_key(f"{item['sheet']}!{item['range']}") for item in variable["source_extents"]]
                or len(_extent_coordinates({"sheet": vector["age_extent"].rsplit("!", 1)[0],
                                            "range": vector["age_extent"].rsplit("!", 1)[1]})) != len(keys)):
            raise ValueError("sample external vector does not match the source catalog or age extent")
        vector_names.add(name)
    workbook = load_workbook(workbook_path, data_only=False, read_only=False,
                             keep_vba=workbook_path.suffix.casefold() == ".xlsm")
    try:
        checked_literals: set[str] = set()
        for vector in vectors:
            bounds = _reference_bounds(vector["age_extent"])
            if bounds is None:
                raise ValueError("sample vector age extent is malformed")
            min_col, min_row, max_col, max_row, sheet_name = bounds
            sheet = next((workbook[name] for name in workbook.sheetnames if name.casefold() == sheet_name.casefold()), None)
            if (sheet is None or min_col != max_col or
                    [sheet.cell(row, min_col).value for row in range(min_row, max_row + 1)] != vector["age_keys"] or
                    any(sheet.cell(row, min_col).data_type == "f" for row in range(min_row, max_row + 1))):
                raise ValueError("sample vector age keys differ from the source workbook axis")
        for entry in [*fixed_fields, *literals]:
            allowed = {"name", "address", "source_literal"} if entry in fixed_fields else {
                "address", "source_literal", "present"}
            _reject_unlisted_fields(entry, allowed, "fixed source literal")
            address = entry.get("address")
            bounds = _reference_bounds(address) if isinstance(address, str) else None
            if bounds is None or bounds[0] != bounds[2] or bounds[1] != bounds[3]:
                raise ValueError("fixed source literal needs a qualified single cell")
            col, row, _, _, sheet_name = bounds
            sheet = next((workbook[name] for name in workbook.sheetnames if name.casefold() == sheet_name.casefold()), None)
            cell = sheet._cells.get((row, col)) if sheet is not None else None
            actual = cell.value if cell is not None else None
            expected = entry.get("source_literal")
            key = _address_key(address)
            if (key in checked_literals or cell is not None and cell.data_type == "f"
                    or type(actual) is not type(expected) and not (
                        isinstance(actual, (int, float)) and not isinstance(actual, bool)
                        and isinstance(expected, (int, float)) and not isinstance(expected, bool))
                    or actual != expected or "present" in entry and entry["present"] is not (cell is not None)):
                raise ValueError(f"fixed source literal mismatch at {address}")
            checked_literals.add(key)
            if entry in literals and not any(
                    _address_key(record["address"]) == key and record["value"] == expected
                    for record in contract["records"]):
                raise ValueError(f"fixed source literal differs from the bound source closure at {address}")
            if entry in fixed_fields:
                matching = [item for item in variables if item.get("logical_name") == entry.get("name")]
                if (len(matching) != 1 or matching[0].get("logical_role") != "fixed_source"
                        or key not in {cell_key for cell_key, _ in _variable_coordinates(matching[0])}):
                    raise ValueError(f"fixed source field {entry.get('name')} is not catalog-bound")
        for field in fields:
            if "fixed_source_literal" in field and not any(
                entry.get("name") == field["name"]
                and _address_key(entry["address"]) in {_address_key(cell) for cell in field["source_cells"]}
                and entry["source_literal"] == field["fixed_source_literal"]
                for entry in fixed_fields
            ):
                raise ValueError("fixed request field lacks a matching checked source literal")
    finally:
        workbook.close()
    if conversion.hash_file(workbook_path) != source_sha256:
        raise ValueError("sample boundary source workbook changed during literal checks")
    return {**sample, "request_case_count": len(cases), "request_field_count": len(fields),
            "request_raw_cell_count": len(field_cells), "external_vector_count": len(vectors),
            "fixed_source_literal_count": len(literals), "request_variable_ids": sorted(request_ids),
            "fixed_variable_ids": sorted({item["variable_id"] for item in variables
                                          if any(_address_key(address) in checked_literals
                                                 for _key, address in _variable_coordinates(item))} - request_ids),
            "internal_variable_ids": sorted(set(variables_by_id) - request_ids - {
                item["variable_id"] for item in variables
                if any(_address_key(address) in checked_literals
                       for _key, address in _variable_coordinates(item))})}


def _require_bound_workflow(analysis: dict[str, Any], binding: dict[str, Any],
                            workflow_root: Path, manifest: dict[str, Any]) -> None:
    declared = binding.get("workflow_path")
    if not isinstance(declared, str) or Path(declared).resolve() != workflow_root:
        raise ValueError("analysis was not prepared for this workflow; run step3 prepare with --workflow")
    source = manifest.get("source", {})
    bound = binding.get("source", {})
    if (source.get("source_id") != bound.get("source_id")
            or source.get("run_id") != bound.get("run_id")
            or source.get("workbook_sha256") != bound.get("source_sha256")):
        raise ValueError("analysis source identity does not match the workflow")
    if manifest.get("input_boundary_required") is not True:
        raise ValueError("workflow is not bound to the Step 3 input-boundary process")


def _validate_targets(targets: Any, binding: dict[str, Any], analysis_binding_sha256: str) -> dict[str, Any]:
    if not isinstance(targets, dict) or targets.get("schema_version") != _TARGETS_SCHEMA:
        raise ValueError(f"targets input must use schema_version {_TARGETS_SCHEMA}")
    _reject_unlisted_fields(targets, {"schema_version", "source", "scenario", "targets"}, "targets input")
    _reject_unlisted_fields(targets.get("source"), {
        "source_id", "run_id", "workbook_sha256", "analysis_binding_sha256",
    }, "targets source")
    expected_source = {"source_id": binding["source"]["source_id"],
                       "run_id": binding["source"]["run_id"],
                       "workbook_sha256": binding["source"]["source_sha256"],
                       "analysis_binding_sha256": analysis_binding_sha256}
    if targets.get("source") != expected_source:
        raise ValueError("targets input does not match the prepared source and analysis binding")
    scenario = targets.get("scenario")
    _reject_unlisted_fields(scenario, {
        "scenario_id", "selectors", "overrides", "description", "scope", "scenario_sha256",
    }, "scenario metadata")
    if not isinstance(scenario.get("scenario_id"), str) or not scenario["scenario_id"].strip():
        raise ValueError("targets input must identify one explicit scenario")
    selectors = scenario.get("selectors")
    if (not isinstance(selectors, dict) or not selectors
            or any(not isinstance(key, str) or not key.strip() or not isinstance(value, str) or not value.strip()
                   for key, value in selectors.items())):
        raise ValueError("scenario selectors must be a non-empty map of names to text metadata")
    overrides = scenario.get("overrides")
    if not isinstance(overrides, dict) or overrides:
        raise ValueError("scenario overrides must be an empty object in the input-metadata phase")
    for key in ("description", "scope"):
        if key in scenario and not isinstance(scenario[key], str):
            raise ValueError(f"scenario {key} must be text metadata")
    scenario_sha = scenario.get("scenario_sha256")
    scenario_basis = {key: value for key, value in scenario.items() if key != "scenario_sha256"}
    if not _is_sha256(scenario_sha) or conversion.hash_bytes(conversion.json_bytes(scenario_basis)) != scenario_sha:
        raise ValueError("scenario_sha256 must hash the declared scenario metadata")
    raw_targets = targets.get("targets")
    if not isinstance(raw_targets, list) or not raw_targets:
        raise ValueError("targets input must declare at least one ordered target")
    ids: set[str] = set()
    orders: list[int] = []
    normalized: list[dict[str, Any]] = []
    for item in raw_targets:
        _reject_unlisted_fields(item, {
            "target_id", "selector", "result_order", "shape", "start_cells",
            "result_kind", "units", "axes",
        },
                                "each target")
        target_id, selector = item.get("target_id"), item.get("selector")
        order, shape, starts = item.get("result_order"), item.get("shape"), item.get("start_cells")
        if not isinstance(target_id, str) or not target_id or target_id in ids:
            raise ValueError("target_id values must be unique non-empty strings")
        if not isinstance(selector, str) or not selector.strip():
            raise ValueError(f"target {target_id} needs an exact selector")
        if isinstance(order, bool) or not isinstance(order, int) or order < 0:
            raise ValueError(f"target {target_id} result_order must be a zero-based integer")
        if not isinstance(shape, list) or any(isinstance(x, bool) or not isinstance(x, int) or x < 0 for x in shape):
            raise ValueError(f"target {target_id} shape must be a list of non-negative dimensions")
        if not isinstance(starts, list) or not starts or any(not isinstance(x, str) for x in starts):
            raise ValueError(f"target {target_id} needs exact start_cells")
        for key in ("result_kind", "units"):
            if key in item and item[key] is not None and (not isinstance(item[key], str) or not item[key].strip()):
                raise ValueError(f"target {target_id} {key} must be non-empty text or null")
        if "axes" in item and item["axes"] is not None:
            if not isinstance(item["axes"], list) or any(not isinstance(axis, dict) for axis in item["axes"]):
                raise ValueError(f"target {target_id} axes must be a list of axis objects or null")
        start_keys = [_address_key(value) for value in starts]
        if len(set(start_keys)) != len(start_keys):
            raise ValueError(f"target {target_id} has duplicate start cells")
        ids.add(target_id)
        orders.append(order)
        normalized_target = {"target_id": target_id, "selector": selector, "result_order": order,
                             "shape": shape, "start_cells": starts}
        normalized_target.update({key: item[key] for key in ("result_kind", "units", "axes") if key in item})
        normalized.append(normalized_target)
    if sorted(orders) != list(range(len(normalized))):
        raise ValueError("target result_order must be a complete zero-based sequence")
    return {"source": expected_source, "scenario": scenario, "targets": normalized}


def _validate_catalog_identity(catalog: Any, binding: dict[str, Any], target_selection: dict[str, Any]) -> None:
    if not isinstance(catalog, dict) or catalog.get("schema_version") != _CATALOG_SCHEMA:
        raise ValueError(f"catalog input must use schema_version {_CATALOG_SCHEMA}")
    if catalog.get("source") != target_selection["source"]:
        raise ValueError("catalog source identity does not match the targets input")
    if catalog.get("scenario") != target_selection["scenario"]:
        raise ValueError("catalog scenario does not match the targets input")
    if catalog.get("target_selection") != target_selection:
        raise ValueError("catalog target order and selectors must exactly match the targets input")


def _validate_catalog(catalog: dict[str, Any], targets: dict[str, Any], fields: dict[str, Any],
                      cut_report: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]],
                                                           list[dict[str, Any]], list[dict[str, Any]],
                                                           list[dict[str, Any]], dict[str, Any]]:
    _reject_unlisted_fields(catalog, {
        "schema_version", "source", "scenario", "target_selection", "groups", "variables",
        "axis_metadata", "source_records", "upstream_exclusions", "value_source_policy",
        "open_questions", "source_scope_notes", "topology_evidence", "draft_status", "scope_note",
        "sample_boundary",
    }, "catalog")
    groups = catalog.get("groups")
    variables = catalog.get("variables")
    axis_metadata = catalog.get("axis_metadata")
    records = catalog.get("source_records")
    exclusions = catalog.get("upstream_exclusions")
    policy = catalog.get("value_source_policy")
    _reject_unlisted_fields(policy, {
        "formula_cache_allowed", "raw_source_origin", "formula_derived_external_values_available_now",
        "allowed_future_external_sources", "external_values_must_bind",
    }, "catalog value_source_policy")
    questions = catalog.get("open_questions", [])
    if not isinstance(groups, list) or not groups:
        raise ValueError("catalog must define at least one logical input group")
    if not isinstance(variables, list) or not variables:
        raise ValueError("catalog must define logical scalar, vector, or table variables")
    if not isinstance(axis_metadata, list):
        raise ValueError("catalog axis_metadata must explicitly preserve key vectors and index coordinates")
    if not isinstance(records, list):
        raise ValueError("catalog source_records must explicitly classify non-business stored inputs")
    if not isinstance(exclusions, list):
        raise ValueError("catalog upstream_exclusions must be a list")
    if not isinstance(policy, dict) or policy.get("formula_cache_allowed") is not False:
        raise ValueError("value_source_policy must explicitly prohibit formula-cache inputs")
    if "raw_source_origin" in policy and policy["raw_source_origin"] != "formula_mode workbook source":
        raise ValueError("value_source_policy.raw_source_origin must describe the bound formula-mode workbook source")
    if ("formula_derived_external_values_available_now" in policy
            and policy["formula_derived_external_values_available_now"] is not False):
        raise ValueError("formula-derived external values must remain unavailable in the input-metadata phase")
    policy_lists = {
        "allowed_future_external_sources": ["approved fresh Excel capture", "explicitly supplied rate file"],
        "external_values_must_bind": ["source SHA-256", "scenario SHA-256", "confirmed boundary SHA-256"],
    }
    for key, allowed_values in policy_lists.items():
        if key in policy:
            value = policy[key]
            if not isinstance(value, list) or not value or any(not isinstance(item, str) or not item.strip()
                                                               for item in value):
                raise ValueError(f"value_source_policy.{key} must be a non-empty list of text metadata")
            if value != allowed_values:
                raise ValueError(f"value_source_policy.{key} must use the declared input-boundary policy values")
    if not isinstance(questions, list) or any(not isinstance(item, (str, dict)) for item in questions):
        raise ValueError("open_questions must be a list of unresolved questions")
    for index, question in enumerate(questions):
        if isinstance(question, dict):
            _reject_unlisted_fields(question, {"question_id", "question"}, f"open_questions[{index}]")
            if (not isinstance(question.get("question_id"), str) or not question["question_id"].strip()
                    or not isinstance(question.get("question"), str) or not question["question"].strip()):
                raise ValueError(f"open_questions[{index}] question_id and question must be non-empty text")
        elif not question.strip():
            raise ValueError(f"open_questions[{index}] must be non-empty text")
    if "source_scope_notes" in catalog:
        _validate_scope_notes(catalog["source_scope_notes"], targets["source"]["workbook_sha256"])
    for key in ("draft_status", "scope_note"):
        if key in catalog and not isinstance(catalog[key], str):
            raise ValueError(f"catalog {key} must be text")
    if "topology_evidence" in catalog:
        evidence = catalog["topology_evidence"]
        _reject_unlisted_fields(evidence, {"path", "sha256", "use"}, "topology_evidence")
        if (not isinstance(evidence.get("path"), str) or not evidence["path"].strip()
                or not _is_sha256(evidence.get("sha256"))
                or evidence.get("use") != "source_topology_only"):
            raise ValueError("topology_evidence needs a text path, SHA-256, and source_topology_only use")
    question_ids = {item.get("question_id") for item in questions if isinstance(item, dict)
                    and isinstance(item.get("question_id"), str)}

    group_ids: set[str] = set()
    clean_groups: list[dict[str, Any]] = []
    for group in groups:
        _reject_unlisted_fields(group, {"group_id", "name", "classification"}, "each group")
        group_id = group.get("group_id")
        if not isinstance(group_id, str) or not group_id.strip() or group_id in group_ids:
            raise ValueError("group_id values must be unique non-empty strings")
        if not isinstance(group.get("name"), str) or not group["name"].strip():
            raise ValueError(f"group {group_id} needs a reviewable name")
        if not isinstance(group.get("classification"), str) or not group["classification"].strip():
            raise ValueError(f"group {group_id} needs an explicit classification")
        group_ids.add(group_id)
        clean_groups.append(group)

    fields_by_cell: dict[str, dict[str, Any]] = {}
    name_descriptors = {str(item.get("name", "")).casefold(): item
                        for item in fields.get("descriptors", {}).get("defined_names", [])}
    for field in fields.get("fields", []):
        for member in field.get("members", []):
            key = _address_key(f"{field.get('sheet')}!{member.get('address')}")
            fields_by_cell[key] = {"field_id": field.get("field_id"), "role": field.get("role"),
                                   "address": f"{field.get('sheet')}!{member.get('address')}"}
    if not fields_by_cell:
        raise ValueError("current fields artifact contains no retained source coordinates")

    target_ids = {item["target_id"] for item in targets["targets"]}
    assigned_raw: dict[str, str] = {}
    assigned_boundary: dict[str, str] = {}
    normalized_variables: list[dict[str, Any]] = []
    variable_ids: set[str] = set()
    formula_trace = cut_report["_formula_by_address"]
    cut_keys = set(cut_report["_declared_boundary_keys"])
    topology_available = cut_report["topology_available"]
    target_starts = {_address_key(cell) for target in targets["targets"] for cell in target["start_cells"]}
    for variable in variables:
        _validate_variable(variable, group_ids, target_ids)
        variable_id = variable["variable_id"]
        if variable_id in variable_ids:
            raise ValueError(f"duplicate variable_id {variable_id}")
        variable_ids.add(variable_id)
        coords = _variable_coordinates(variable)
        if not coords:
            raise ValueError(f"variable {variable_id} has no exact source coordinates")
        for key, _address in coords:
            if key in target_starts:
                raise ValueError(f"target output cell cannot be promoted to an input: {_address}")
            cell = fields_by_cell.get(key)
            if cell is None:
                raise ValueError(f"catalog extent includes an unstored or out-of-scope cell: {_address}")
            declared_role = variable["role"]
            actual_formula = cell["role"] == "calculated"
            if declared_role == "source_raw" and actual_formula:
                raise ValueError(f"source_raw variable {variable_id} includes a formula cell: {_address}")
            if declared_role == "formula_derived_external" and not actual_formula:
                raise ValueError(f"formula-derived external variable {variable_id} includes a raw cell: {_address}")
            if declared_role == "source_raw":
                previous = assigned_raw.get(key)
                if previous:
                    raise ValueError(f"raw source coordinate overlaps {previous} and {variable_id}: {_address}")
                assigned_raw[key] = variable_id
            else:
                if key not in cut_keys or (topology_available and key not in formula_trace):
                    raise ValueError(f"external formula coordinate is not in the selected trace: {_address}")
                previous = assigned_boundary.get(key)
                if previous:
                    raise ValueError(f"external boundary coordinate overlaps {previous} and {variable_id}: {_address}")
                assigned_boundary[key] = variable_id
        provenance = variable.get("provenance", {})
        for name in provenance.get("source_defined_names", []):
            descriptor = name_descriptors.get(str(name).casefold())
            if descriptor is None:
                raise ValueError(f"source-defined name {name} is not present in current canonical name descriptors")
            descriptor_key = _address_key(_canonical_ref(str(descriptor.get("address", "")), descriptor.get("sheet")))
            if descriptor_key not in {key for key, _address in coords}:
                raise ValueError(f"source-defined name {name} does not resolve to a declared variable coordinate")
        for header in variable.get("source_column_labels", []):
            if (not isinstance(header, dict) or not isinstance(header.get("source_cell"), str)
                    or not isinstance(header.get("label"), str)):
                raise ValueError(f"variable {variable_id} has invalid source column header metadata")
            header_cell = fields_by_cell.get(_address_key(header["source_cell"]))
            if header_cell is None or header_cell["role"] != "source":
                raise ValueError(f"source column label is not backed by a retained raw source cell: {header['source_cell']}")
        normalized_variables.append(variable)

    axis_ids: set[str] = set()
    normalized_axes: list[dict[str, Any]] = []
    for axis in axis_metadata:
        _reject_unlisted_fields(axis, {
            "axis_id", "group_id", "name", "shape", "source_role", "source_extents", "provenance",
        }, "each axis_metadata item")
        axis_id, axis_name = axis.get("axis_id"), axis.get("name")
        if not isinstance(axis_id, str) or not axis_id or axis_id in axis_ids:
            raise ValueError("axis_id values must be unique non-empty strings")
        if not isinstance(axis_name, str) or not axis_name.strip() or axis.get("group_id") not in group_ids:
            raise ValueError(f"axis metadata {axis_id} needs a name and known group")
        shape, extents = axis.get("shape"), axis.get("source_extents")
        if (not isinstance(shape, list) or len(shape) != 1 or isinstance(shape[0], bool)
                or not isinstance(shape[0], int) or shape[0] <= 0
                or not isinstance(extents, list) or not extents):
            raise ValueError(f"axis metadata {axis_id} must declare a non-empty one-dimensional source extent")
        coords = [coordinate for extent in extents for coordinate in _extent_coordinates(extent)]
        if len(coords) != shape[0] or axis.get("source_role") != "raw_source":
            raise ValueError(f"axis metadata {axis_id} shape or source role does not match its coordinates")
        _reject_unlisted_fields(axis.get("provenance"), {"description", "source_workbook_sha256"},
                                f"axis metadata {axis_id} provenance")
        if (not isinstance(axis.get("provenance"), dict)
                or not isinstance(axis["provenance"].get("description"), str)
                or not axis["provenance"]["description"].strip()):
            raise ValueError(f"axis metadata {axis_id} needs source provenance")
        if ("source_workbook_sha256" in axis["provenance"]
                and not isinstance(axis["provenance"]["source_workbook_sha256"], str)):
            raise ValueError(f"axis metadata {axis_id} provenance source_workbook_sha256 must be text")
        for extent in extents:
            _reject_unlisted_fields(extent, {"sheet", "range", "source_role"}, f"axis metadata {axis_id} extent")
            if extent.get("source_role") != "raw_source":
                raise ValueError(f"axis metadata {axis_id} extent must be raw_source")
        for key, address in coords:
            cell = fields_by_cell.get(key)
            if cell is None or cell["role"] != "source":
                raise ValueError(f"axis metadata {axis_id} is not backed by retained raw source: {address}")
            if key in assigned_raw or key in assigned_boundary:
                raise ValueError(f"axis metadata overlaps another source extent: {address}")
            assigned_raw[key] = axis_id
        axis_ids.add(axis_id)
        normalized_axes.append(axis)

    for variable in normalized_variables:
        for axis in variable["axes"]:
            if axis.get("role") == "key" and axis.get("axis_id") not in axis_ids:
                raise ValueError(f"variable {variable['variable_id']} references unknown axis metadata")

    normalized_records: list[dict[str, Any]] = []
    for record in records:
        _reject_unlisted_fields(record, {"record_id", "group_id", "classification", "source_extents",
                                         "provenance", "open_question_id"}, "each source record")
        if record.get("classification") not in {
            "adapter_metadata", "source_constant", "pending_classification"
        }:
            raise ValueError("source_records must classify each grouped extent as adapter_metadata, source_constant, or pending_classification")
        record_id, group_id = record.get("record_id"), record.get("group_id")
        if not isinstance(record_id, str) or not record_id or group_id not in group_ids:
            raise ValueError("each source record needs a record_id and known group_id")
        if record["classification"] == "pending_classification" and (
            not isinstance(record.get("open_question_id"), str) or not record["open_question_id"].strip()
        ):
            raise ValueError(f"pending source record {record_id} needs an open_question_id")
        if "open_question_id" in record:
            question_id = record["open_question_id"]
            if not isinstance(question_id, str) or not question_id.strip():
                raise ValueError(f"source record {record_id} open_question_id must be non-empty text")
            if question_id not in question_ids:
                raise ValueError(f"source record {record_id} cites an undeclared open question")
        extents = record.get("source_extents")
        if not isinstance(extents, list) or not extents:
            raise ValueError(f"source record {record_id} needs exact source_extents")
        provenance = record.get("provenance")
        if provenance is not None:
            _reject_unlisted_fields(provenance, {"description", "source_workbook_sha256", "source_check"},
                                    f"source record {record_id} provenance")
            if not isinstance(provenance.get("description"), str) or not provenance["description"].strip():
                raise ValueError(f"source record {record_id} provenance needs a description")
            if "source_workbook_sha256" in provenance and not isinstance(provenance["source_workbook_sha256"], str):
                raise ValueError(f"source record {record_id} provenance source_workbook_sha256 must be text")
        if provenance is not None and "source_check" in provenance:
            source_check = provenance["source_check"]
            _reject_unlisted_fields(source_check, {
                "source_copy_sha256", "literal_zero_coordinate_count", "formula_coordinates", "inspection",
            }, f"source record {record_id} source_check")
            if not isinstance(source_check.get("source_copy_sha256"), str):
                raise ValueError(f"source record {record_id} source_check source_copy_sha256 must be text")
            zero_count = source_check.get("literal_zero_coordinate_count")
            if isinstance(zero_count, bool) or not isinstance(zero_count, int) or zero_count < 0:
                raise ValueError(f"source record {record_id} source_check literal_zero_coordinate_count must be non-negative")
            formula_count = source_check.get("formula_coordinates")
            if isinstance(formula_count, bool) or not isinstance(formula_count, int) or formula_count < 0:
                raise ValueError(f"source record {record_id} source_check formula_coordinates must be a non-negative count")
            if not isinstance(source_check.get("inspection"), str):
                raise ValueError(f"source record {record_id} source_check inspection must be text")
        for extent in extents:
            _reject_unlisted_fields(extent, {"sheet", "range", "source_role"}, f"source record {record_id} extent")
            if extent.get("source_role") != "raw_source":
                raise ValueError(f"source record {record_id} extents must be raw_source")
            for key, address in _extent_coordinates(extent):
                cell = fields_by_cell.get(key)
                if cell is None:
                    raise ValueError(f"source record extent includes an unstored or out-of-scope cell: {address}")
                if cell["role"] == "calculated":
                    raise ValueError(f"source record extent includes a formula cell: {address}")
                if key in assigned_raw:
                    raise ValueError(f"raw source coordinate overlaps {assigned_raw[key]} and {record_id}: {address}")
                if key in assigned_boundary:
                    raise ValueError(f"raw source coordinate overlaps an external formula boundary: {address}")
                assigned_raw[key] = record_id
        normalized_records.append(record)

    if topology_available:
        expected_raw = set(cut_report["_reachable_raw_keys"])
        actual_raw = set(assigned_raw)
        missing_raw = sorted(expected_raw - actual_raw)
        extra_raw = sorted(actual_raw - expected_raw)
        if missing_raw or extra_raw:
            raise ValueError(f"catalog must account for every retained raw coordinate in the declared topology candidate (missing={len(missing_raw)}, extra={len(extra_raw)})")
    missing_boundary = sorted(cut_keys - set(assigned_boundary))
    extra_boundary = sorted(set(assigned_boundary) - cut_keys)
    if missing_boundary or extra_boundary:
        raise ValueError(f"catalog formula-derived external variables must exactly cover the proposed cut (missing={len(missing_boundary)}, extra={len(extra_boundary)})")

    boundary_variables = {item["variable_id"]: item for item in normalized_variables
                          if item["role"] == "formula_derived_external"}
    seen_exclusions: set[tuple[str, str]] = set()
    clean_exclusions: list[dict[str, Any]] = []
    for item in exclusions:
        _reject_unlisted_fields(item, {"boundary_variable_id", "named_source", "source_range", "excluded_range", "reason"},
                                "each upstream exclusion")
        if item.get("boundary_variable_id") not in boundary_variables:
            raise ValueError("each upstream exclusion must cite a formula-derived boundary variable")
        name, ref = item.get("named_source"), item.get("source_range")
        if not isinstance(name, str) or not isinstance(ref, str) or not isinstance(item.get("reason"), str) or not item["reason"].strip():
            raise ValueError("upstream exclusions need named_source, exact source_range, and reason")
        descriptor = name_descriptors.get(name.casefold())
        if descriptor is None or _canonical_ref(str(descriptor.get("address", "")), descriptor.get("sheet")) != _canonical_ref(ref):
            raise ValueError(f"upstream exclusion {name} does not match the canonical defined-name descriptor")
        if "excluded_range" in item:
            named_bounds = _reference_bounds(ref)
            excluded_bounds = (_reference_bounds(item["excluded_range"])
                               if isinstance(item["excluded_range"], str) else None)
            if excluded_bounds is not None and (excluded_bounds[0] > excluded_bounds[2]
                                                or excluded_bounds[1] > excluded_bounds[3]):
                raise ValueError(f"upstream exclusion {name} excluded_range has reversed endpoints")
            if (named_bounds is None or excluded_bounds is None
                    or excluded_bounds[4].casefold() != named_bounds[4].casefold()
                    or excluded_bounds[0] < named_bounds[0] or excluded_bounds[1] < named_bounds[1]
                    or excluded_bounds[2] > named_bounds[2] or excluded_bounds[3] > named_bounds[3]):
                raise ValueError(f"upstream exclusion {name} excluded_range must be within the named source range")
        key = (item["boundary_variable_id"], name.casefold())
        if key in seen_exclusions:
            raise ValueError(f"duplicate upstream exclusion for {name}")
        seen_exclusions.add(key)
        clean_exclusions.append(item)
    excluded_for = {item["boundary_variable_id"] for item in clean_exclusions}
    if set(boundary_variables) - excluded_for:
        raise ValueError("each formula-derived external variable must identify its upstream exclusion")
    return clean_groups, normalized_variables, normalized_axes, normalized_records, clean_exclusions, policy


def _validate_variable(variable: Any, group_ids: set[str], target_ids: set[str]) -> None:
    _reject_unlisted_fields(variable, {
        "variable_id", "group_id", "logical_name", "role", "kind", "logical_role", "shape", "axes",
        "source_extents", "consumer_targets", "provenance", "value_source_policy", "coordinate_map",
        "curve_axis", "curve_count", "source_column_labels",
    }, "each variable")
    variable_id = variable.get("variable_id")
    if not isinstance(variable_id, str) or not variable_id:
        raise ValueError("each variable needs a non-empty variable_id")
    if variable.get("group_id") not in group_ids:
        raise ValueError(f"variable {variable_id} references an unknown group")
    if not isinstance(variable.get("logical_name"), str) or not variable["logical_name"].strip():
        raise ValueError(f"variable {variable_id} needs a logical_name")
    role, kind = variable.get("role"), variable.get("kind")
    if role not in {"source_raw", "formula_derived_external"}:
        raise ValueError(f"variable {variable_id} role must be source_raw or formula_derived_external")
    if kind not in {"scalar", "vector", "table"}:
        raise ValueError(f"variable {variable_id} kind must be scalar, vector, or table")
    if variable.get("logical_role") not in {"business_input", "series", "configuration_table", "fixed_source", "other"}:
        raise ValueError(f"variable {variable_id} needs an explicit logical_role")
    shape, axes = variable.get("shape"), variable.get("axes")
    if not isinstance(shape, list) or any(isinstance(size, bool) or not isinstance(size, int) or size <= 0 for size in shape):
        raise ValueError(f"variable {variable_id} shape must contain positive dimensions")
    expected_rank = {"scalar": 0, "vector": 1, "table": 2}[kind]
    if len(shape) != expected_rank or not isinstance(axes, list) or len(axes) != expected_rank:
        raise ValueError(f"variable {variable_id} kind, shape, and axes rank do not match")
    for axis in axes:
        _reject_unlisted_fields(axis, {"name", "role", "axis_id"}, f"variable {variable_id} axis")
        if not isinstance(axis.get("name"), str) or not axis["name"].strip():
            raise ValueError(f"variable {variable_id} has an axis without a name")
        if axis.get("role") not in {"key", "observation", "series", "category", "record"}:
            raise ValueError(f"variable {variable_id} axis needs a declared role")
        if "axis_id" in axis and (not isinstance(axis["axis_id"], str) or not axis["axis_id"].strip()):
            raise ValueError(f"variable {variable_id} axis_id must be non-empty text when supplied")
    extents = variable.get("source_extents")
    if not isinstance(extents, list) or not extents:
        raise ValueError(f"variable {variable_id} needs source_extents")
    for extent in extents:
        _reject_unlisted_fields(extent, {"sheet", "range", "source_role"}, f"variable {variable_id} extent")
    expected_source_role = "raw_source" if role == "source_raw" else "formula_output"
    coordinates = [coordinate for extent in extents for coordinate in _extent_coordinates(extent)]
    if kind == "scalar" and len(coordinates) != 1:
        raise ValueError(f"scalar {variable_id} must cover one source coordinate")
    if kind == "vector" and len(coordinates) != shape[0]:
        raise ValueError(f"vector {variable_id} source coordinates do not match its declared length")
    if kind == "table":
        coordinate_map = variable.get("coordinate_map")
        if coordinate_map is None:
            if len(extents) != 1:
                raise ValueError(f"table {variable_id} with one rectangular extent must declare exactly one source extent")
            extent_shape = _extent_shape(extents[0])
            if extent_shape != tuple(shape):
                raise ValueError(f"table {variable_id} source extent does not match its declared shape")
        else:
            _validate_sparse_table_map(variable_id, coordinate_map, coordinates, shape)
        curve_axis = variable.get("curve_axis")
        curve_count = variable.get("curve_count", 0)
        if curve_axis is not None and not isinstance(curve_axis, str):
            raise ValueError(f"table {variable_id} curve_axis must be text")
        if isinstance(curve_count, bool) or not isinstance(curve_count, int) or curve_count < 0:
            raise ValueError(f"table {variable_id} curve_count must be a non-negative integer")
        if curve_axis is not None:
            matching = [index for index, axis in enumerate(axes) if axis.get("name") == curve_axis]
            if len(matching) != 1 or curve_count != shape[matching[0]]:
                raise ValueError(f"table {variable_id} curve_axis and curve_count do not match its declared axes")
        elif curve_count != 0:
            raise ValueError(f"table {variable_id} must declare curve_axis when curve_count is non-zero")
    elif any(key in variable for key in ("coordinate_map", "curve_axis", "curve_count")):
        raise ValueError(f"non-table variable {variable_id} cannot contain table mapping fields")
    if len({item[0] for item in coordinates}) != len(coordinates):
        raise ValueError(f"variable {variable_id} source_extents overlap")
    for extent in extents:
        if extent.get("source_role") != expected_source_role:
            raise ValueError(f"variable {variable_id} extent source_role must be {expected_source_role}")
    consumers = variable.get("consumer_targets")
    if not isinstance(consumers, list) or not consumers or any(item not in target_ids for item in consumers):
        raise ValueError(f"variable {variable_id} needs known consumer_targets")
    provenance = variable.get("provenance")
    _reject_unlisted_fields(provenance, {
        "description", "source_workbook_sha256", "coordinates_are_source_facts", "source_defined_names",
    }, f"variable {variable_id} provenance")
    if not isinstance(provenance, dict) or not isinstance(provenance.get("description"), str) or not provenance["description"].strip():
        raise ValueError(f"variable {variable_id} needs explicit provenance")
    if "source_workbook_sha256" in provenance and not isinstance(provenance["source_workbook_sha256"], str):
        raise ValueError(f"variable {variable_id} provenance source_workbook_sha256 must be text")
    if "coordinates_are_source_facts" in provenance and not isinstance(provenance["coordinates_are_source_facts"], bool):
        raise ValueError(f"variable {variable_id} provenance coordinates_are_source_facts must be boolean")
    if "source_defined_names" in provenance and (
            not isinstance(provenance["source_defined_names"], list)
            or any(not isinstance(name, str) for name in provenance["source_defined_names"])):
        raise ValueError(f"variable {variable_id} provenance source_defined_names must be a list of names")
    policy = variable.get("value_source_policy")
    policy_fields = ({"origin", "formula_cache_allowed"} if role == "source_raw" else
                     {"allowed_future_sources", "formula_cache_allowed", "available_now"})
    _reject_unlisted_fields(policy, policy_fields, f"variable {variable_id} value_source_policy")
    if role == "source_raw" and policy.get("origin") != "workbook_formula_mode_source":
        raise ValueError(f"raw variable {variable_id} must use formula-mode workbook source values")
    if role == "source_raw" and policy.get("formula_cache_allowed") is not False:
        raise ValueError(f"raw variable {variable_id} must explicitly prohibit formula-cache inputs")
    if role == "formula_derived_external":
        allowed = {"approved_excel_capture", "user_rate_file"}
        if policy.get("allowed_future_sources") != sorted(allowed):
            raise ValueError(f"external variable {variable_id} must require a future approved capture or user rate file")
        if policy.get("formula_cache_allowed") is not False or policy.get("available_now") is not False:
            raise ValueError(f"external variable {variable_id} cannot use formula caches or claim values are available now")
    labels = variable.get("source_column_labels", [])
    if not isinstance(labels, list):
        raise ValueError(f"variable {variable_id} source_column_labels must be a list")
    for label in labels:
        _reject_unlisted_fields(label, {"source_cell", "label"}, f"variable {variable_id} source column label")


def _source_candidate_report(targets: dict[str, Any], catalog: dict[str, Any],
                             fields: dict[str, Any]) -> dict[str, Any]:
    """Validate source roles while making no target-reachability claim."""
    cells: dict[str, dict[str, Any]] = {}
    for field in fields.get("fields", []):
        for member in field.get("members", []):
            address = f"{field.get('sheet')}!{member.get('address')}"
            cells[_address_key(address)] = {"address": address, "role": field.get("role")}
    for target in targets["targets"]:
        for start in target["start_cells"]:
            cell = cells.get(_address_key(start))
            if cell is None or cell["role"] != "calculated":
                raise ValueError(f"target start is not a retained calculated field in the current source: {start}")

    raw_keys: set[str] = set()
    boundary_keys: set[str] = set()
    for variable in catalog.get("variables", []):
        if not isinstance(variable, dict):
            continue
        destination = boundary_keys if variable.get("role") == "formula_derived_external" else raw_keys
        for key, address in _variable_coordinates(variable):
            cell = cells.get(key)
            expected_role = "calculated" if destination is boundary_keys else "source"
            if cell is None or cell["role"] != expected_role:
                role_label = "calculated" if expected_role == "calculated" else "raw"
                raise ValueError(f"catalog candidate does not match a retained {role_label} source field: {address}")
            destination.add(key)
    for item in [*catalog.get("axis_metadata", []), *catalog.get("source_records", [])]:
        if not isinstance(item, dict):
            continue
        for extent in item.get("source_extents", []):
            for key, address in _extent_coordinates(extent):
                cell = cells.get(key)
                if cell is None or cell["role"] != "source":
                    raise ValueError(f"catalog metadata is not backed by a retained raw source field: {address}")
                raw_keys.add(key)

    source_count = sum(field.get("role") == "source" for field in cells.values())
    return {
        "status": "source_candidate_only_target_closure_unavailable",
        "source_topology_only": False,
        "topology_available": False,
        "trace_source_sha256": None,
        "trace_targets": [],
        "candidate_boundary_coordinate_count": len(boundary_keys),
        "candidate_boundary_reached_count": None,
        "candidate_boundary_unreached_count": None,
        "retained_formula_positions": None,
        "retained_raw_source_coordinates": len(raw_keys),
        "source_field_raw_coordinate_count": source_count,
        "all_source_raw_coordinates_across_fields": source_count,
        "unknown_count": None,
        "unknown_examples": ["No target-bound topology was supplied; formula closure, dynamic branches, and lookup-selected dependencies remain unproven."],
        "unrepresented_source_coordinate_count": None,
        "unrepresented_source_coordinate_examples": [],
        "unrepresented_semantics": "No source topology was supplied, so absent coordinates and blank-read behavior have not been measured.",
        "input_coverage": {
            "retained_raw_coordinate_count": len(raw_keys),
            "source_field_raw_coordinate_count": source_count,
            "assigned_raw_coordinate_count": None,
            "unassigned_raw_coordinate_count": None,
            "status": "target_closure_unproven_without_topology",
        },
        "_reachable_raw_keys": set(),
        "_declared_boundary_keys": boundary_keys,
        "_formula_by_address": {},
    }


def _validate_sparse_table_map(
    variable_id: str,
    coordinate_map: Any,
    extent_coordinates: list[tuple[str, str]],
    shape: list[int],
) -> None:
    if not isinstance(coordinate_map, list) or not coordinate_map:
        raise ValueError(f"sparse table {variable_id} needs an explicit coordinate_map")
    expected = {key for key, _address in extent_coordinates}
    mapped: set[str] = set()
    logical_indices: set[tuple[int, int]] = set()
    for item in coordinate_map:
        _reject_unlisted_fields(item, {"source_cell", "logical_index"}, f"sparse table {variable_id} coordinate")
        if not isinstance(item.get("source_cell"), str):
            raise ValueError(f"sparse table {variable_id} has an invalid coordinate_map entry")
        indices = item.get("logical_index")
        if (not isinstance(indices, list) or len(indices) != 2
                or any(isinstance(index, bool) or not isinstance(index, int) for index in indices)):
            raise ValueError(f"sparse table {variable_id} logical_index must contain two integers")
        if any(index < 0 or index >= shape[dimension] for dimension, index in enumerate(indices)):
            raise ValueError(f"sparse table {variable_id} logical_index is outside its declared shape")
        source_key = _address_key(item["source_cell"])
        if source_key not in expected:
            raise ValueError(f"sparse table {variable_id} coordinate_map cites a cell outside its extents: {item['source_cell']}")
        index_key = (indices[0], indices[1])
        if source_key in mapped or index_key in logical_indices:
            raise ValueError(f"sparse table {variable_id} coordinate_map has a duplicate source or logical slot")
        mapped.add(source_key)
        logical_indices.add(index_key)
    if mapped != expected:
        raise ValueError(f"sparse table {variable_id} coordinate_map must cover every exact source coordinate")


def _cut_report(trace: Any, targets: dict[str, Any], catalog: dict[str, Any],
                fields: dict[str, Any], binding: dict[str, Any]) -> dict[str, Any]:
    evidence = catalog.get("topology_evidence")
    if not isinstance(trace, dict) or trace.get("schema_version") not in {"step4.discovery_trace.v1", "gp.active_trace.v1"}:
        raise ValueError("topology evidence must use a supported source trace schema")
    if trace.get("source_sha256") != binding["source"]["source_sha256"]:
        raise ValueError("topology trace belongs to a different source workbook")
    if trace.get("native_excel_called") is not False or trace.get("formula_cache_inputs") is not False:
        raise ValueError("topology trace must be source-only and formula-cache free")
    if not isinstance(evidence, dict) or evidence.get("use") != "source_topology_only":
        raise ValueError("topology_evidence must explicitly be source_topology_only")
    if not isinstance(trace.get("cells"), list) or not isinstance(trace.get("edges"), list):
        raise ValueError("topology trace is missing cells or dependency edges")
    trace_targets = trace.get("targets")
    if not isinstance(trace_targets, dict):
        raise ValueError("topology trace is missing target selectors")
    trace_cells: dict[str, dict[str, Any]] = {}
    for item in trace["cells"]:
        if isinstance(item, dict) and isinstance(item.get("address"), str):
            # Copy only the non-value topology fields; no trace value is used.
            trace_cells[_address_key(item["address"])] = {
                "address": item["address"], "role": item.get("role"), "formula": item.get("formula")}
    for target in targets["targets"]:
        target_evidence = trace_targets.get(target["target_id"])
        if not isinstance(target_evidence, dict):
            raise ValueError(f"topology trace does not contain target {target['target_id']}")
        for key in ("selector", "result_order", "shape"):
            if target_evidence.get(key) != target.get(key):
                raise ValueError(f"topology trace {key} does not match target {target['target_id']}")
        for start in target["start_cells"]:
            if _address_key(start) not in trace_cells:
                raise ValueError(f"target start cell is absent from the topology trace: {start}")

    cuts = [item for item in catalog.get("variables", [])
            if isinstance(item, dict) and item.get("role") == "formula_derived_external"]
    cut_keys: set[str] = set()
    cut_addresses: dict[str, str] = {}
    for variable in cuts:
        for key, address in _variable_coordinates(variable):
            cut_keys.add(key)
            cut_addresses[key] = address
    field_cells: dict[str, dict[str, Any]] = {}
    for field in fields.get("fields", []):
        for member in field.get("members", []):
            address = f"{field.get('sheet')}!{member.get('address')}"
            field_cells[_address_key(address)] = {"address": address, "role": field.get("role")}
    for key, address in cut_addresses.items():
        item = trace_cells.get(key)
        if item is None or item.get("role") != "calculated_formula" or not isinstance(item.get("formula"), str):
            raise ValueError(f"proposed external boundary is not an ordinary traced formula: {address}")

    adjacency: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for edge in trace["edges"]:
        if not isinstance(edge, dict) or not isinstance(edge.get("consumer"), str) or not isinstance(edge.get("prerequisite"), str):
            raise ValueError("topology trace contains an invalid edge")
        adjacency[_address_key(edge["consumer"])].append((edge["prerequisite"], str(edge.get("relationship", ""))))
    lookups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for lookup in trace.get("lookups", []):
        if isinstance(lookup, dict) and isinstance(lookup.get("formula_cell"), str):
            # Exclude lookup_value: it is data, not topology, and is never consumed below.
            lookups[_address_key(lookup["formula_cell"])].append({key: value for key, value in lookup.items()
                                                                   if key not in _VALUE_KEYS | {"lookup_value"}})

    raw_keys: set[str] = set()
    formula_keys: set[str] = set()
    reached_cuts: set[str] = set()
    unrepresented: set[str] = set()
    unknown: set[str] = set()
    pending = deque(_address_key(cell) for target in targets["targets"] for cell in target["start_cells"])
    seen: set[str] = set()
    while pending:
        key = pending.popleft()
        if key in seen:
            continue
        seen.add(key)
        if key in cut_keys:
            reached_cuts.add(key)
            continue
        cell = trace_cells.get(key)
        field = field_cells.get(key)
        if field is None:
            # A trace or expanded finite range may name a physical coordinate
            # absent from retained fields; do not imply blank semantics.
            unrepresented.add(key)
        if cell is not None and cell.get("role") == "calculated_formula":
            formula_keys.add(key)
        elif field is not None and field.get("role") == "source":
            raw_keys.add(key)

        for prerequisite, relationship in adjacency.get(key, []):
            if relationship == "lookup_table":
                continue
            if relationship == "range_read":
                expanded = _expand_reference(prerequisite)
                if expanded is None:
                    unknown.add(f"range_read:{prerequisite}")
                else:
                    pending.extend(_address_key(address) for address in expanded)
            elif relationship == "array_formula_anchor":
                pending.append(_address_key(prerequisite))
            elif relationship == "cell_reference":
                pending.append(_address_key(prerequisite))
            else:
                unknown.add(f"{relationship}:{prerequisite}")

        for lookup in lookups.get(key, []):
            function = str(lookup.get("function", "")).upper()
            if function == "MATCH":
                ref = lookup.get("lookup_range")
                expanded = _expand_reference(ref) if isinstance(ref, str) else None
                if expanded is None:
                    unknown.add(f"MATCH:{ref}")
                else:
                    pending.extend(_address_key(address) for address in expanded)
            elif function in {"VLOOKUP", "HLOOKUP"}:
                ref = lookup.get("table_range")
                matched_row = lookup.get("matched_row")
                return_column = lookup.get("return_column")
                bounds = _reference_bounds(ref) if isinstance(ref, str) else None
                if (bounds is None or isinstance(matched_row, bool) or not isinstance(matched_row, int)
                        or isinstance(return_column, bool) or not isinstance(return_column, int) or return_column < 1):
                    unknown.add(f"{function}:{ref}:selected_return_unresolved")
                    continue
                min_col, min_row, max_col, max_row, sheet = bounds
                if function == "VLOOKUP":
                    key_col = min_col + int(lookup.get("lookup_column", 1)) - 1
                    if key_col > max_col or return_column > max_col - min_col + 1:
                        unknown.add(f"{function}:{ref}:column_out_of_range")
                        continue
                    pending.extend(_address_key(f"{sheet}!{get_column_letter(key_col)}{row}")
                                   for row in range(min_row, max_row + 1))
                    pending.append(_address_key(f"{sheet}!{get_column_letter(min_col + return_column - 1)}{matched_row}"))
                else:
                    key_row = min_row + int(lookup.get("lookup_row", 1)) - 1
                    if key_row > max_row or return_column > max_row - min_row + 1:
                        unknown.add(f"{function}:{ref}:row_out_of_range")
                        continue
                    pending.extend(_address_key(f"{sheet}!{get_column_letter(col)}{key_row}")
                                   for col in range(min_col, max_col + 1))
                    pending.append(_address_key(f"{sheet}!{get_column_letter(min_col + return_column - 1)}{matched_row}"))
            else:
                unknown.add(f"lookup:{function}")

    if unknown:
        status = "historical_topology_candidate_with_unknowns"
    elif unrepresented:
        status = "historical_topology_candidate_with_unrepresented_coordinates"
    else:
        status = "historical_topology_candidate"
    reachable_raw = raw_keys
    # Metadata consumed by the validator is private and removed before serialization.
    report = {
        "status": status,
        "source_topology_only": True,
        "trace_source_sha256": trace.get("source_sha256"),
        "trace_targets": sorted(target["target_id"] for target in targets["targets"]),
        "candidate_boundary_coordinate_count": len(cut_keys),
        "candidate_boundary_reached_count": len(reached_cuts),
        "candidate_boundary_unreached_count": len(cut_keys - reached_cuts),
        "retained_formula_positions": len(formula_keys),
        "retained_raw_source_coordinates": len(reachable_raw),
        "unknown_count": len(unknown),
        "unknown_examples": sorted(unknown)[:20],
        "unrepresented_source_coordinate_count": len(unrepresented),
        "unrepresented_source_coordinate_examples": sorted(unrepresented)[:20],
        "unrepresented_semantics": "Coordinates are absent from current retained fields; this does not establish whether a formula reads a physical blank or what blank-read behavior Excel applies.",
        "input_coverage": {"retained_raw_coordinate_count": len(reachable_raw),
                           "assigned_raw_coordinate_count": None,
                           "unassigned_raw_coordinate_count": None,
                           "status": "validated_after_catalog_partition"},
        "_reachable_raw_keys": reachable_raw,
        "_declared_boundary_keys": cut_keys,
        "_formula_by_address": {key: trace_cells[key]["formula"] for key in cut_keys if key in trace_cells},
        "topology_available": True,
    }
    return report


def _input_counts(variables: list[dict[str, Any]], source_records: list[dict[str, Any]],
                  axis_metadata: list[dict[str, Any]]) -> dict[str, int]:
    def mapped_elements(item: dict[str, Any]) -> int:
        coordinate_map = item.get("coordinate_map")
        if item.get("kind") == "table" and isinstance(coordinate_map, list):
            return len(coordinate_map)
        return prod(item["shape"]) if item["shape"] else 1

    logical_count = len(variables)
    scalar_count = sum(item["kind"] == "scalar" for item in variables)
    vector_count = sum(item["kind"] == "vector" for item in variables)
    vector_curves = sum(item["kind"] == "vector" and item.get("logical_role") == "series" for item in variables)
    table_count = sum(item["kind"] == "table" for item in variables)
    table_curves = sum(int(item.get("curve_count", 0)) for item in variables if item["kind"] == "table")
    elements = sum(mapped_elements(item) for item in variables)
    raw_variables = [item for item in variables if item["role"] == "source_raw"]
    external_variables = [item for item in variables if item["role"] == "formula_derived_external"]
    raw_elements = sum(mapped_elements(item) for item in raw_variables)
    external_elements = sum(mapped_elements(item) for item in external_variables)
    sparse_table_mapped_coordinates = sum(len(item.get("coordinate_map", []))
                                          for item in variables if item["kind"] == "table")
    sparse_table_bounding_slots = sum(prod(item["shape"]) for item in variables
                                      if item["kind"] == "table" and item.get("coordinate_map") is not None)
    variable_coordinates = sum(sum(len(_extent_coordinates(extent))
                                   for extent in item.get("source_extents", [])) for item in variables)
    record_coordinates = {
        classification: sum(
            len(_extent_coordinates(extent))
            for item in source_records if item.get("classification") == classification
            for extent in item.get("source_extents", [])
        )
        for classification in ("adapter_metadata", "source_constant", "pending_classification")
    }
    axis_coordinates = sum(len(_extent_coordinates(extent))
                           for item in axis_metadata for extent in item.get("source_extents", []))
    record_coordinate_count = sum(record_coordinates.values())
    return {"logical_input_object_count": logical_count,
            "scalar_field_count": scalar_count,
            "vector_object_count": vector_count,
            "non_curve_vector_count": vector_count - vector_curves,
            "vector_curve_count": vector_curves,
            "table_count": table_count,
            "table_curve_count": table_curves,
            "sparse_table_mapped_coordinate_count": sparse_table_mapped_coordinates,
            "sparse_table_bounding_slot_count": sparse_table_bounding_slots,
            "field_or_curve_count": scalar_count + vector_curves + table_curves,
            "source_raw_variable_count": len(raw_variables),
            "formula_derived_external_variable_count": len(external_variables),
            "declared_logical_element_count": elements,
            "source_raw_declared_element_count": raw_elements,
            "formula_derived_external_element_count": external_elements,
            "source_coordinate_count": variable_coordinates + record_coordinate_count + axis_coordinates,
            "axis_metadata_element_count": axis_coordinates,
            "source_record_coordinate_count": record_coordinate_count,
            "adapter_metadata_coordinate_count": record_coordinates["adapter_metadata"],
            "source_constant_coordinate_count": record_coordinates["source_constant"],
            "pending_classification_coordinate_count": record_coordinates["pending_classification"],
            "source_record_group_count": len(source_records)}


def input_boundary_state(root: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    """Return freshness and receipts for the current catalog pointer, even after final design replaces current Stage 3."""
    pointer = manifest.get("input_boundary")
    if not isinstance(pointer, dict):
        return {"status": "pending", "revision": None, "reasons": ["input boundary has not been submitted"]}
    revisions = manifest.get("stages", {}).get("3", {}).get("revisions", [])
    revision = next((item for item in revisions if item.get("revision") == pointer.get("revision")), None)
    if revision is None or revision.get("artifact") != pointer.get("artifact"):
        return {"status": "stale", "revision": pointer.get("revision"), "reasons": ["input boundary pointer does not match its saved revision"]}
    state = stage_status_for_boundary(root, manifest, revision)
    if state.get("status") in {"approved", "simulated_approved"}:
        state["status"] = "input_boundary_confirmed" if state["status"] == "approved" else "simulated_input_boundary_confirmed"
    state["boundary_sha256"] = pointer.get("boundary_sha256")
    return state


def stage_status_for_boundary(root: Path, manifest: dict[str, Any], revision: dict[str, Any]) -> dict[str, Any]:
    return conversion.stage_status(root, manifest, 3, revision_override=revision)


def require_input_boundary_confirmed(root: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    state = input_boundary_state(root, manifest)
    allowed = {"input_boundary_confirmed"}
    if manifest.get("mode") == "synthetic_fixture":
        allowed.add("simulated_input_boundary_confirmed")
    if state.get("status") not in allowed:
        raise ValueError(f"input boundary is {state.get('status')}; final Step 3 design requires its confirmed catalog")
    return state


def require_confirmed_boundary_for_analysis(analysis_dir: Path, analysis: dict[str, Any],
                                            binding: dict[str, Any]) -> tuple[Path, dict[str, Any]] | None:
    workflow_path = binding.get("workflow_path")
    if not workflow_path:
        return None
    root, manifest = conversion.load_workflow(Path(workflow_path))
    _require_bound_workflow(analysis, binding, root, manifest)
    state = require_input_boundary_confirmed(root, manifest)
    pointer = manifest["input_boundary"]
    return root, {**pointer, "state": state}


def _markdown_table_cell(value: Any) -> str:
    return str(value).replace("\r\n", "<br>").replace("\r", "<br>").replace("\n", "<br>").replace("|", "\\|")


def _boundary_evidence_markdown(payload: dict[str, Any]) -> str:
    counts = payload["counts"]
    cut = payload["pruned_source_analysis"]
    lines = ["# Step 3 Input Boundary", "", f"Status: **{payload['status']}**",
             f"Boundary SHA-256: `{payload['boundary_sha256']}`", "",
             "This catalog records inputs and source topology. It does not provide external values, calculate formulas, use formula caches, or release Step 4.",
             "", "## Binding", "",
             f"Source: `{payload['source']['source_id']}` / run `{payload['source']['run_id']}` / workbook `{payload['source']['workbook_sha256']}`",
             f"Scenario: `{payload['scenario']['scenario_id']}` ({payload['scenario']['scenario_sha256']})",
             "Targets:", ""]
    for target in sorted(payload["target_selection"]["targets"], key=lambda item: item["result_order"]):
        lines.append(
            f"- `{conversion.display_report_value(target.get('target_id'))}` · order "
            f"{conversion.display_report_value(target.get('result_order'))} · selector "
            f"{conversion.display_report_value(target.get('selector'))} · kind "
            f"{conversion.display_report_value(target.get('result_kind'))} · shape "
            f"{conversion.display_report_value(target.get('shape'))} · declared axes "
            f"{conversion.display_report_axes(target)} · units "
            f"{conversion.display_report_value(target.get('units'))}"
        )
    lines.extend(["", "## Counts", "", "| Measure | Count |", "|---|---:|"])
    for key in ("logical_input_object_count", "field_or_curve_count", "scalar_field_count", "vector_object_count",
                "non_curve_vector_count", "vector_curve_count",
                "table_count", "table_curve_count", "source_raw_variable_count",
                "sparse_table_mapped_coordinate_count", "sparse_table_bounding_slot_count",
                "formula_derived_external_variable_count", "declared_logical_element_count",
                "source_raw_declared_element_count", "formula_derived_external_element_count", "source_coordinate_count",
                "axis_metadata_element_count", "adapter_metadata_coordinate_count",
                "source_constant_coordinate_count", "pending_classification_coordinate_count",
                "source_record_group_count"):
        lines.append(f"| {key.replace('_', ' ')} | {counts[key]} |")
    lines.extend(["", "## Source catalog groups", "", "| Group | Classification | Source objects |", "|---|---|---|"])
    for group in payload["groups"]:
        members = [item["logical_name"] for item in payload["variables"] if item["group_id"] == group["group_id"]]
        lines.append(f"| {_markdown_table_cell(group['name'])} (`{_markdown_table_cell(group['group_id'])}`) | "
                     f"{_markdown_table_cell(group['classification'])} | "
                     f"{_markdown_table_cell(', '.join(members) or 'source records only')} |")
    lines.extend(["", "## Variables", "", "| ID | Kind and shape | Axes | Source role | Source extents | Source header metadata | Consumers |", "|---|---|---|---|---|---|---|"])
    for variable in payload["variables"]:
        extents = ", ".join(f"{item['sheet']}!{item['range']}" for item in variable["source_extents"])
        labels = "; ".join(f"{item['source_cell']}={item['label']}" for item in variable.get("source_column_labels", []))
        axes = ("scalar" if not variable["axes"] else
                conversion.display_report_axes(variable))
        lines.append(f"| `{_markdown_table_cell(variable['variable_id'])}` — {_markdown_table_cell(variable['logical_name'])} | "
                     f"{_markdown_table_cell(variable['kind'])} {_markdown_table_cell(variable['shape'])} | "
                     f"{_markdown_table_cell(axes)} | {_markdown_table_cell(variable['role'])} | "
                     f"{_markdown_table_cell(extents)} | {_markdown_table_cell(labels)} | "
                     f"{_markdown_table_cell(', '.join(variable['consumer_targets']))} |")
    if payload.get("axis_metadata"):
        lines.extend(["", "## Source axes and index metadata", "", "| Axis | Shape | Source extent | Meaning recorded in this draft |", "|---|---:|---|---|"])
        for axis in payload["axis_metadata"]:
            extents = ", ".join(f"{item['sheet']}!{item['range']}" for item in axis["source_extents"])
            lines.append(f"| `{_markdown_table_cell(axis['axis_id'])}` — {_markdown_table_cell(axis['name'])} | "
                         f"{_markdown_table_cell(axis['shape'])} | {_markdown_table_cell(extents)} | "
                         f"{_markdown_table_cell(axis['provenance']['description'])} |")
    lines.extend(["", "## Source topology and closure", ""])
    if cut.get("topology_available"):
        lines.extend([f"Status: **{cut['status']}**. Formula positions retained: {cut['retained_formula_positions']}; retained raw coordinates: {cut['retained_raw_source_coordinates']}; candidate boundary coordinates reached: {cut['candidate_boundary_reached_count']}/{cut['candidate_boundary_coordinate_count']}; unresolved topology: {cut['unknown_count']}.",
                      f"Coordinates referenced by the trace but absent from retained fields: {cut['unrepresented_source_coordinate_count']}; examples: {', '.join(cut['unrepresented_source_coordinate_examples']) or 'none'}.",
                      cut["unrepresented_semantics"],
                      "This is a source-topology candidate based on the declared trace. A later approved run must rediscover the boundary before runtime work."])
    else:
        lines.extend([f"Status: **{cut['status']}**. Catalog source coordinates listed: {cut['retained_raw_source_coordinates']} raw plus {cut['candidate_boundary_coordinate_count']} formula-derived candidates; the full fields artifact has {cut['all_source_raw_coordinates_across_fields']} retained raw source coordinates across all sheets. Target reachability, formula closure, dynamic branches, and lookup-selected dependencies remain unknown.",
                      "No trace was supplied, so unrepresented coordinates and blank-read semantics have not been measured. Unknown coverage is not zero.",
                      "This catalog can be reviewed before Step 4 discovery. Confirmation records proposed inputs only; it does not certify completeness or release Step 4."])
    notes = payload.get("source_scope_notes", {})
    if isinstance(notes, dict):
        if notes.get("note"):
            lines.extend(["", f"Source scope note: {notes['note']}"])
        if notes.get("non_retained_historical_candidates"):
            lines.append("Historical candidate coordinates absent from current retained fields: " + ", ".join(f"`{item}`" for item in notes["non_retained_historical_candidates"]))
        absence = notes.get("stored_coordinate_reconciliation", notes.get("physical_absence_evidence", {}))
        if isinstance(absence, dict):
            audit = absence.get("stored_coordinate_audit")
            empty_coordinates = absence.get("historical_empty_coordinates")
            if audit is not None:
                lines.append(f"The source-coordinate audit `{audit['path']}` (SHA-256 `{audit['sha256']}`) is bound to the same workbook SHA. It reconciles {audit['old_trace_raw_coordinate_count']} historical raw candidates against {audit['current_retained_raw_coordinate_count']} current retained raw coordinates and identifies {audit['empty_source_coordinate_count']} stored-but-empty source cells:")
                if audit.get("coordinates"):
                    lines.append(", ".join(f"`{item['address']}` (cell element present; no formula or literal content)" for item in audit["coordinates"]))
            elif empty_coordinates:
                lines.append("No separate source-coordinate audit or reconciliation counts are attached. "
                             "Current-source verification records these stored blank cells: " +
                             ", ".join(f"`{item['address']}` (cell element present; no formula or literal content)"
                                       for item in empty_coordinates))
            if empty_coordinates:
                lines.append(absence.get("interpretation") or
                             "No interpretation was supplied. Later blank-read semantics remain unresolved.")
        historical = notes.get("historical_trace_candidate")
        if isinstance(historical, dict):
            lines.append(f"Historical trace `{historical['path']}` (SHA-256 `{historical['sha256']}`) is retained as a lead only: {historical['use']}.")
        if notes.get("excluded_formula_extents"):
            lines.append("Excluded formula extents: " + "; ".join(f"`{item['sheet']}!{item['range']}` — {item['reason']}" for item in notes["excluded_formula_extents"]))
    lines.extend(["", "## Upstream exclusions", "", "| Input | Named source | Excluded range | Reason |", "|---|---|---|---|"])
    for item in payload["upstream_exclusions"]:
        lines.append(f"| `{_markdown_table_cell(item['boundary_variable_id'])}` | "
                     f"`{_markdown_table_cell(item['named_source'])}` | "
                     f"`{_markdown_table_cell(item.get('excluded_range', item['source_range']))}` | {_markdown_table_cell(item['reason'])} |")
    if payload.get("source_records"):
        lines.extend(["", "## Non-business and pending source records", "", "These records account for retained source coordinates without treating headers, labels, or literals as separate business variables. Pending classifications remain visible and are not asserted as understood inputs.", ""])
        for item in payload["source_records"]:
            extents = ", ".join(f"{extent['sheet']}!{extent['range']}" for extent in item["source_extents"])
            question = f"; open question `{item['open_question_id']}`" if item.get("open_question_id") else ""
            description = item.get("provenance", {}).get("description", "")
            lines.append(f"- `{item['record_id']}` ({item['classification']}; group `{item['group_id']}`{question}): {extents}. {description}")
    questions = payload.get("open_questions")
    questions = questions if isinstance(questions, list) else []
    if questions:
        lines.extend(["", "## Open questions", ""])
        for item in questions:
            lines.append(f"- {item if isinstance(item, str) else json.dumps(item, ensure_ascii=False, sort_keys=True)}")
    lines.extend(["", "## Review gate", "",
                  "This checkpoint always requires an independent Agent decision and an explicit human confirmation on this JSON/Markdown pair, even when a TypeSafe delegation is active. A TypeSafe receipt may appear as a separate recommendation, but it cannot approve or reject this checkpoint or replace either required decision. Catalogs accept only documented metadata fields; unknown fields such as `sample_rates` or `rates` block submission. Approval sets `input_boundary_confirmed` only; Stage 3 remains the next permitted checkpoint.", ""])
    return "\n".join(lines)


def _boundary_markdown(
    payload: dict[str, Any], *, machine_link: str = "input_boundary.json",
    workflow_link: str = "../../workflow.json",
) -> str:
    counts = payload.get("counts")
    counts = counts if isinstance(counts, dict) else {}
    selection = payload.get("target_selection")
    selection = selection if isinstance(selection, dict) else {}
    source = payload.get("source")
    source = source if isinstance(source, dict) else {}
    scenario = selection.get("scenario")
    scenario = scenario if isinstance(scenario, dict) else {}
    cut = payload.get("pruned_source_analysis")
    cut = cut if isinstance(cut, dict) else {}
    variables = payload.get("variables")
    variables = variables if isinstance(variables, list) else []
    objects = [
        ("Source catalog objects", counts.get("logical_input_object_count")),
        ("Scalar objects", counts.get("scalar_field_count")),
        ("Vector objects", counts.get("vector_object_count")),
        ("Table objects", counts.get("table_count")),
        ("Raw source objects", counts.get("source_raw_variable_count")),
        ("Formula-derived external objects", counts.get("formula_derived_external_variable_count")),
    ]
    lines = [
        "# Step 3 Input Boundary", "",
        "## What this checkpoint asks you to confirm", "",
        "The supported request fields, broader source roles and shapes, and current source-bound evidence. "
        "Confirmation records these input assumptions; it does not approve the final Step 3 design or release Step 4.", "",
        "## Purpose and scope", "",
        f"Bound source run: `{conversion.display_report_value(source.get('run_id'))}`; source ID: "
        f"`{conversion.display_report_value(source.get('source_id'))}`. Scenario: `{conversion.display_report_value(scenario.get('scenario_id'))}`. "
        "The source scenario and requested targets are listed below; formula caches are not inputs.", "",
        "## Verified results", "",
        "Counts below describe the broader source catalog, not caller-mutable request fields. "
        "A cell coordinate or mapped table slot is source provenance, not a separate business input.", "",
        "| Measure | Count |", "|---|---:|"]
    lines.extend(f"| {label} | {conversion.display_report_value(value)} |" for label, value in objects)
    sample = payload.get("sample_boundary")
    if isinstance(sample, dict):
        lines.extend(["", "## Supported request and fixed source boundary", "",
                      f"Current workbook SHA-256: `{sample['source_workbook_sha256']}`. "
                      f"Exact request cases: {sample['request_case_count']}; named request fields: "
                      f"{sample['request_field_count']}; selected raw request cells: "
                      f"{sample['request_raw_cell_count']}; external vectors: "
                      f"{sample['external_vector_count']}; separately checked fixed source literals: "
                      f"{sample['fixed_source_literal_count']}.", "",
                      "Request axes and complete Cartesian membership are recorded in the machine checkpoint. "
                      "Fixed contract fields carried by requests are not extra mutable inputs. "
                      f"Separate fixed-literal source ledger: `{sample['fixed_literal_contract']['path']}` "
                      f"(SHA-256 `{sample['fixed_literal_contract']['sha256']}`).", "",
                      "| Request field | Exact source cells | Scope |", "|---|---|---|"])
        for field in sample["request_fields"]:
            scope = (f"fixed to `{field['fixed_source_literal']}`" if "fixed_source_literal" in field
                     else "request-selected within the stated case axes")
            lines.append(f"| {_markdown_table_cell(field['name'])} | "
                         f"{_markdown_table_cell(', '.join(field['source_cells']))} | {scope} |")
        lines.extend(["", "External vector source/age extents:", ""])
        for vector in sample["external_vectors"]:
            lines.append(f"- `{vector['name']}`: {vector['source_extent']}; age keys "
                         f"{vector['age_extent']} ({len(vector['age_keys'])} values per request); "
                         f"producer version `{vector['producer_version']}`.")
        lines.extend(["", "Fixed source fields:", ""])
        for entry in sample["fixed_source_fields"]:
            lines.append(f"- `{entry['name']}` at `{entry['address']}` = "
                         f"`{json.dumps(entry['source_literal'], ensure_ascii=False)}` (formula-mode source literal).")
        lines.extend(["", "Exact fixed source literals (including physical presence):", ""])
        for entry in sample["fixed_source_literals"]:
            lines.append(f"- `{entry['address']}` = `{json.dumps(entry['source_literal'], ensure_ascii=False)}`; "
                         f"physical cell present: `{str(entry['present']).lower()}`.")
        lines.extend(["", f"Broader internal source catalog objects: "
                      f"{len(sample['internal_variable_ids'])}; their IDs remain in the machine checkpoint.", ""])
    lines.extend(["", "### Source catalog groups and roles", ""])
    groups = payload.get("groups")
    groups = groups if isinstance(groups, list) else []
    for group in groups:
        if not isinstance(group, dict):
            continue
        group_id = group.get("group_id")
        members = [item for item in variables if isinstance(item, dict) and item.get("group_id") == group_id]
        source_records = payload.get("source_records")
        source_records = source_records if isinstance(source_records, list) else []
        records = [item for item in source_records if isinstance(item, dict) and item.get("group_id") == group_id]
        lines.append(
            f"**{conversion.display_report_value(group.get('name', group_id))}** — {conversion.display_report_value(group.get('classification'))}; "
            f"{len(members)} source objects, {len(records)} source-record groups."
        )
        for variable in members:
            extents = [f"{conversion.display_report_value(item.get('sheet'))}!{conversion.display_report_value(item.get('range'))}"
                       for item in variable.get("source_extents", []) if isinstance(item, dict)]
            extent_text = "; ".join(extents[:4]) or "source location not recorded"
            if len(extents) > 4:
                extent_text += f"; {len(extents) - 4} more recorded extents in the detailed table below"
            shape = variable.get("shape")
            shape_text = "scalar" if variable.get("kind") == "scalar" else f"{conversion.display_report_value(variable.get('kind'))} {conversion.display_report_value(shape)}"
            logical_name = conversion.display_report_value(variable.get("logical_name", variable.get("variable_id")))
            variable_id = variable.get("variable_id")
            name_text = f"`{logical_name}`" + (f" (`{variable_id}`)" if variable_id and variable_id != logical_name else "")
            lines.append(
                f"- {name_text} — "
                f"{shape_text}, role `{conversion.display_report_value(variable.get('role'))}`; "
                f"input axes {conversion.display_report_axes(variable)}; {extent_text}."
            )
        for record in records:
            record_extents = [f"{conversion.display_report_value(item.get('sheet'))}!{conversion.display_report_value(item.get('range'))}"
                              for item in record.get("source_extents", []) if isinstance(item, dict)]
            extent_text = "; ".join(record_extents[:3]) or "source locations not recorded"
            if len(record_extents) > 3:
                extent_text += f"; {len(record_extents) - 3} more recorded locations in the detailed table below"
            lines.append(
                f"- Source record `{conversion.display_report_value(record.get('record_id'))}` — "
                f"role `{conversion.display_report_value(record.get('classification'))}`; {extent_text}."
            )
    targets = selection.get("targets", [])
    targets = targets if isinstance(targets, list) else []
    if targets:
        lines.extend(["", "Requested targets:", ""])
        for target in sorted(
            (item for item in targets if isinstance(item, dict)),
            key=lambda item: item.get("result_order") if isinstance(item.get("result_order"), int)
            and not isinstance(item.get("result_order"), bool) else 10**9,
        ):
            lines.append(
                f"- `{conversion.display_report_value(target.get('target_id'))}`: {conversion.display_report_value(target.get('selector'))} "
                f"(order {conversion.display_report_value(target.get('result_order'))}; "
                f"kind {conversion.display_report_value(target.get('result_kind'))}; "
                f"shape {conversion.display_report_value(target.get('shape'))}; "
                f"units {conversion.display_report_value(target.get('units'))}; "
                f"declared axes {conversion.display_report_axes(target)}); source start "
                f"{', '.join(target.get('start_cells', [])) if isinstance(target.get('start_cells'), list) else 'Not recorded'}."
            )
    if variables:
        kinds = sorted({conversion.display_report_value(item.get("kind")) for item in variables if isinstance(item, dict)})
        lines.extend(["", f"Recorded source catalog kinds: {', '.join(kinds) if kinds else 'Not recorded'}."])
    topology_available = cut.get("topology_available")
    if topology_available is True:
        lines.append(
            f"Trace-bound topology status: {conversion.display_report_value(cut.get('status'))}; unresolved topology count: "
            f"{conversion.display_report_value(cut.get('unknown_count'))}."
        )
    elif topology_available is False:
        lines.append("No current trace-bound topology is available; target reachability and formula closure remain unknown.")
    else:
        lines.append("Trace-bound topology availability: Not recorded; target reachability and formula closure remain unknown.")
    questions = payload.get("open_questions")
    questions = questions if isinstance(questions, list) else []
    lines.extend(["", "## Limits and unresolved items", ""])
    lines.append(f"Open questions recorded: {conversion.display_report_value(len(questions) if isinstance(payload.get('open_questions'), list) else None)}.")
    for question in questions:
        if isinstance(question, dict):
            question_id = conversion.display_report_value(question.get("question_id"))
            question_text = conversion.display_report_value(question.get("question"))
            status = conversion.display_report_value(question.get("disposition", question.get("status")))
            lines.append(f"- **{question_id}** — {question_text} Recorded disposition/status: {status.rstrip('.')}.")
        else:
            lines.append(f"- {conversion.display_report_value(question)}")
    lines.extend([
        "A catalogued coordinate is not automatically a logical input. Unknown reachability, closure, dynamic branches, lookup-selected dependencies, and blank-read semantics remain unknown where the bound trace does not settle them.", "",
        "## Handover and review", "",
        f"- Machine checkpoint: [input_boundary.json](<{machine_link}>).",
        f"- Current acceptance ledger: [workflow.json](<{workflow_link}>).",
        "- The Agent reviews the exact JSON/Markdown pair first. This input-boundary checkpoint always requires a separate actual human confirmation; clarification of a material question is not approval.",
        "- To reject: route the decision to the current or earlier responsible Step, retain this revision, refresh affected artifacts and reviews, then check the ledger before continuing.",
        "- `ready_for_review` means the report was assembled from checked evidence. Agent PASS, human confirmation, and final-design approval are separate ledger states.", "",
        "## Detailed source and boundary evidence", "",
        re.sub(r"(?m)^## ", "#### ", _boundary_evidence_markdown(payload).replace(
            "# Step 3 Input Boundary", "### Original boundary evidence", 1)), "",
    ])
    return "\n".join(lines)


def _extent_coordinates(extent: Any) -> list[tuple[str, str]]:
    if not isinstance(extent, dict):
        raise ValueError("each source extent must be an object")
    sheet, cell_range = extent.get("sheet"), extent.get("range")
    if not isinstance(sheet, str) or not sheet.strip() or not isinstance(cell_range, str) or not cell_range.strip():
        raise ValueError("source extent needs a sheet name and A1 range")
    bounds = _reference_bounds(f"{sheet}!{cell_range}")
    if bounds is None:
        raise ValueError(f"invalid finite source extent: {sheet}!{cell_range}")
    min_col, min_row, max_col, max_row, canonical_sheet = bounds
    area = (max_col - min_col + 1) * (max_row - min_row + 1)
    if area > 100_000:
        raise ValueError(f"source extent is too large for a reviewed input object: {sheet}!{cell_range}")
    return [(_address_key(f"{canonical_sheet}!{get_column_letter(col)}{row}"),
             f"{canonical_sheet}!{get_column_letter(col)}{row}")
            for row in range(min_row, max_row + 1) for col in range(min_col, max_col + 1)]


def _variable_coordinates(variable: dict[str, Any]) -> list[tuple[str, str]]:
    return [coord for extent in variable.get("source_extents", []) for coord in _extent_coordinates(extent)]


def _extent_shape(extent: dict[str, Any]) -> tuple[int, int]:
    bounds = _reference_bounds(f"{extent['sheet']}!{extent['range']}")
    if bounds is None:
        raise ValueError("source extent is not a finite A1 range")
    min_col, min_row, max_col, max_row, _sheet = bounds
    return max_row - min_row + 1, max_col - min_col + 1


def _expand_reference(reference: Any) -> list[str] | None:
    if not isinstance(reference, str):
        return None
    bounds = _reference_bounds(reference)
    if bounds is None:
        return None
    min_col, min_row, max_col, max_row, sheet = bounds
    if (max_col - min_col + 1) * (max_row - min_row + 1) > 100_000:
        return None
    return [f"{sheet}!{get_column_letter(col)}{row}"
            for row in range(min_row, max_row + 1) for col in range(min_col, max_col + 1)]


def _reference_bounds(reference: str) -> tuple[int, int, int, int, str] | None:
    if "!" not in reference:
        return None
    sheet, address = reference.rsplit("!", 1)
    sheet = sheet.strip()
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    address = address.replace("$", "")
    # Whole-row/whole-column, named, external, and 3-D references remain unsupported.
    if not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]*(?::[A-Za-z]{1,3}[1-9][0-9]*)?", address):
        return None
    try:
        min_col, min_row, max_col, max_row = range_boundaries(address)
    except ValueError:
        return None
    if None in {min_col, min_row, max_col, max_row}:
        return None
    return int(min_col), int(min_row), int(max_col), int(max_row), sheet


def _canonical_ref(reference: str, fallback_sheet: str | None = None) -> str:
    if "!" not in reference and fallback_sheet:
        reference = f"{fallback_sheet}!{reference}"
    bounds = _reference_bounds(reference)
    if bounds is None:
        return reference.replace("$", "").casefold()
    min_col, min_row, max_col, max_row, sheet = bounds
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return f"{sheet}!{start}" + (f":{end}" if start != end else "")


def _address_key(address: str) -> str:
    return _canonical_ref(address).casefold()


def _assert_no_embedded_values(value: Any, label: str) -> None:
    if isinstance(value, dict):
        for key, item in value.items():
            if str(key).casefold() in _VALUE_KEYS:
                raise ValueError(f"{label} must describe value provenance, not embed source values or caches")
            _assert_no_embedded_values(item, label)
    elif isinstance(value, list):
        for item in value:
            _assert_no_embedded_values(item, label)


def _is_sha256(value: Any) -> bool:
    return isinstance(value, str) and re.fullmatch(r"[0-9a-f]{64}", value) is not None


def _boundary_revision_reference(root: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    state = require_input_boundary_confirmed(root, manifest)
    pointer = manifest["input_boundary"]
    return {"schema_version": "step3.input_boundary.reference.v1",
            "revision": pointer["revision"], "boundary_sha256": pointer["boundary_sha256"],
            "artifact": pointer["artifact"], "review_receipts": {
                key: state.get(key) for key in ("reviewer_pair", "agent", "human", "typesafe")}}


def confirmed_boundary_reference(root: Path, manifest: dict[str, Any]) -> tuple[dict[str, Any], Path, Path]:
    reference = _boundary_revision_reference(root, manifest)
    json_path = (root / reference["artifact"]["json"]).resolve()
    md_path = (root / reference["artifact"]["md"]).resolve()
    if (not json_path.is_file() or conversion.hash_file(json_path) != reference["artifact"]["json_sha256"]
            or not md_path.is_file() or conversion.hash_file(md_path) != reference["artifact"]["md_sha256"]):
        raise ValueError("confirmed input-boundary artifact pair is missing or changed")
    return reference, json_path, md_path
