"""Source-bound native evidence and formula traces for a finite Step 3 case set."""

from __future__ import annotations

import hashlib
import itertools
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from excel_to_act.steps.conversion_workflow import hash_file, load_workflow, read_json
from excel_to_act.steps.step3.calculation import FormulaEvaluator, _parse_reference
from excel_to_act.steps.step3.input_boundary import confirmed_boundary_reference


def _digest(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(",", ":")).encode("utf-8")).hexdigest()


def _public_digest(value: Any) -> str:
    return hashlib.sha256((json.dumps(value, sort_keys=True, ensure_ascii=False,
                                      separators=(",", ":")) + "\n").encode("utf-8")).hexdigest()


def _address(address: str) -> str:
    ref = _parse_reference(address, "")
    if ref is None or not ref.single or not ref.sheet:
        raise ValueError(f"case source address is not a qualified single cell: {address}")
    return f"{ref.sheet.casefold()}!{ref.min_col}:{ref.min_row}"


def _ref(record: Any, label: str) -> tuple[Path, str]:
    if (not isinstance(record, dict) or set(record) != {"path", "sha256"}
            or not isinstance(record["path"], str)
            or not isinstance(record["sha256"], str) or len(record["sha256"]) != 64):
        raise ValueError(f"{label} reference needs an exact path and SHA-256")
    path = Path(record["path"]).expanduser().resolve()
    if not path.is_file() or hash_file(path) != record["sha256"]:
        raise ValueError(f"{label} is missing or changed: {path}")
    return path, record["sha256"]


def _numeric_equal(native: Any, public: Any) -> bool:
    return (type(native) in {int, float} and type(public) in {int, float}
            and math.isfinite(native) and math.isfinite(public) and native == public)


def _case_axes(boundary: dict[str, Any], public: dict[str, Any]) -> tuple[Any, ...]:
    axes = boundary["sample_boundary"]["case_axes"]
    result = []
    for name, choices in axes.items():
        item = public["inputs"].get(name)
        if not isinstance(item, dict) or item.get("axes") != [] or item.get("values") not in choices:
            raise ValueError(f"case axis {name} is outside the approved request domain")
        result.append(item["values"])
    return tuple(result)


def _source_cells(boundary: dict[str, Any], public: dict[str, Any]) -> dict[str, Any]:
    sample = boundary["sample_boundary"]
    fields = sample["request_fields"]
    if set(public.get("inputs", {})) != {item["name"] for item in fields}:
        raise ValueError("public request fields differ from the confirmed boundary")
    variables = {item["variable_id"]: item for item in boundary["variables"]}
    mapped: dict[str, Any] = {}
    for field in fields:
        name, addresses = field["name"], field["source_cells"]
        variable = variables[field["variable_id"]]
        record = public["inputs"][name]
        if not isinstance(record, dict) or set(record) != {"axes", "values"}:
            raise ValueError(f"public request field is malformed: {name}")
        if (not isinstance(record["axes"], list) or len(record["axes"]) != len(variable["shape"])
                or any(not isinstance(axis, dict) or
                       {key: value for key, value in axis.items() if key != "keys"} != declared
                       or not isinstance(axis.get("keys"), list)
                       or len(axis["keys"]) != variable["shape"][index]
                       for index, (axis, declared) in enumerate(zip(record["axes"], variable["axes"])))):
            raise ValueError(f"public request axes differ from the source boundary: {name}")
        if len(addresses) == 1:
            if record["axes"] != [] or isinstance(record["values"], (list, dict)):
                raise ValueError(f"public scalar request is malformed: {name}")
            pairs = [(addresses[0], record["values"])]
        else:
            values = record["values"]
            selected = {entry["source_cell"]: tuple(entry["logical_index"])
                        for entry in variable.get("coordinate_map", [])
                        if entry.get("source_cell") in addresses}
            if (set(selected) != set(addresses) or not isinstance(values, list)
                    or len(values) != len(addresses)
                    or any(not isinstance(item, dict) or set(item) != {"indices", "value"}
                           for item in values)):
                raise ValueError(f"public sparse request is malformed: {name}")
            indexed = {tuple(item["indices"]): item["value"] for item in values}
            if len(indexed) != len(values) or set(indexed) != set(selected.values()):
                raise ValueError(f"public sparse indices differ from the source map: {name}")
            pairs = [(address, indexed[selected[address]]) for address in addresses]
        for address, value in pairs:
            key = _address(address)
            if key in mapped:
                raise ValueError(f"public request repeats a source cell: {address}")
            mapped[key] = value
    if len(mapped) != sample["request_raw_cell_count"]:
        raise ValueError("public request does not map the complete raw-cell boundary")
    return mapped


def _oracle_cuts(boundary: dict[str, Any], oracle: dict[str, Any], public: dict[str, Any],
                 workbook: Any, source_hash: str) -> tuple[dict[str, Any], int]:
    vectors = boundary["sample_boundary"]["external_vectors"]
    external = public.get("external_data")
    if not isinstance(external, dict) or set(external) != {item["variable_id"] for item in vectors}:
        raise ValueError("public external vectors differ from the confirmed boundary")
    cuts: dict[str, Any] = {}
    normalized = 0
    for spec in vectors:
        name, extent = spec["variable_id"], spec["source_extent"]
        data = external[name]
        variable = next(item for item in boundary["variables"] if item["variable_id"] == name)
        expected_axes = [{**variable["axes"][0], "keys": spec["age_keys"]}]
        if (not isinstance(data, dict)
                or set(data) != {"axes", "values", "source_sha256", "version", "input_sha256"}
                or data["source_sha256"] != source_hash
                or data["version"] != spec["producer_version"]
                or data["input_sha256"] != _public_digest(public["inputs"])
                or data["axes"] != expected_axes):
            raise ValueError(f"public external vector has wrong source, version, input, or axis: {name}")
        ref = _parse_reference(extent, "")
        if ref is None or not ref.sheet or ref.min_col != ref.max_col:
            raise ValueError(f"external source extent is not a finite vector: {extent}")
        age_ref = _parse_reference(spec["age_extent"], "")
        if (age_ref is None or not age_ref.sheet or age_ref.min_col != age_ref.max_col
                or [workbook[age_ref.sheet].cell(row, age_ref.min_col).value
                    for row in range(age_ref.min_row, age_ref.max_row + 1)] != spec["age_keys"]):
            raise ValueError(f"external age axis differs from the source workbook: {name}")
        native = oracle.get("ranges", {}).get(extent)
        length = ref.max_row - ref.min_row + 1
        if (not isinstance(native, dict) or native.get("address") != extent
                or len(data["values"]) != length or length != len(spec["age_keys"])
                or any(not isinstance(native.get(key), list) or len(native[key]) != length
                       for key in ("values", "formulas", "excel_errors"))):
            raise ValueError(f"native external vector is missing or has the wrong shape: {name}")
        sheet = workbook[ref.sheet]
        for index, value in enumerate(data["values"]):
            cell = sheet.cell(ref.min_row + index, ref.min_col)
            rows = [native[key][index] for key in ("values", "formulas", "excel_errors")]
            if any(not isinstance(row, list) or len(row) != 1 for row in rows):
                raise ValueError(f"native cut has an incomplete row: {extent}")
            native_value, formula, error = (row[0] for row in rows)
            if (cell.data_type != "f" or formula != cell.value or error is not None
                    or not _numeric_equal(native_value, value)):
                raise ValueError(f"native cut differs from source/public request: {name} row {index}")
            if type(native_value) is not type(value):
                normalized += 1
        cuts[extent] = {"values": data["values"], "native_range_sha256": _digest(native),
                        "numeric_equivalence": "finite numeric equality; 0 and 0.0 compare equal"}
    return cuts, normalized


def _verify_fixed_literals(boundary: dict[str, Any], workbook: Any, source_hash: str) -> None:
    sample = boundary["sample_boundary"]
    sheets = {name.casefold(): name for name in workbook.sheetnames}
    contract = sample["fixed_literal_contract"]
    _ref(contract, "fixed source literal contract")
    literals = sample["fixed_source_literals"]
    if len(literals) != sample["fixed_source_literal_count"]:
        raise ValueError("fixed source literal set is incomplete")
    for item in literals:
        ref = _parse_reference(item["address"], "")
        if ref is None or not ref.single or not ref.sheet:
            raise ValueError("fixed source literal address is malformed")
        cell = workbook[sheets[ref.sheet.casefold()]]._cells.get((ref.min_row, ref.min_col))
        if ((cell is not None) != item["present"]
                or (cell.value if cell is not None else None) != item["source_literal"]
                or (cell is not None and cell.data_type == "f")):
            raise ValueError(f"fixed source literal differs from workbook: {item['address']}")
    for item in sample["fixed_source_fields"]:
        ref = _parse_reference(item["address"], "")
        cell = workbook[sheets[ref.sheet.casefold()]].cell(ref.min_row, ref.min_col)
        if cell.data_type == "f" or cell.value != item["source_literal"]:
            raise ValueError(f"fixed source field differs from workbook: {item['address']}")
    if sample["source_workbook_sha256"] != source_hash:
        raise ValueError("fixed source literals refer to another workbook")


def verify_native_case_producer(producer: dict[str, Any], case: dict[str, Any],
                                source: Path, raw: dict[str, Any], public: dict[str, Any]) -> None:
    """Independently recheck the sealed native readback and source formula cuts."""
    source_hash = hash_file(source)
    if producer.get("source_sha256") != source_hash or producer.get("status") != "pass":
        raise ValueError("exact-case producer source is stale")
    boundary_path, _ = _ref(producer.get("boundary_json"), "exact-case boundary")
    boundary = read_json(boundary_path)
    if (boundary.get("boundary_sha256") != producer.get("boundary_sha256")
            or boundary.get("sample_boundary", {}).get("source_workbook_sha256") != source_hash):
        raise ValueError("exact-case producer boundary differs from the current source")
    evidence = producer.get("evidence")
    if not isinstance(evidence, dict):
        raise ValueError("exact-case producer has no bound evidence")
    for label in ("raw_request", "public_request"):
        if evidence.get(label) != case.get(label):
            raise ValueError(f"exact-case producer {label} differs from the case reference")
        _ref(evidence[label], label)
    oracle_path, _ = _ref(evidence.get("native_oracle"), "native oracle")
    oracle = read_json(oracle_path)
    records = raw.get("input_overrides")
    if (not isinstance(records, list) or len(records) != boundary["sample_boundary"]["request_raw_cell_count"]
            or producer.get("input_sha256") != _digest(records)
            or raw.get("input_sha256") != producer["input_sha256"]):
        raise ValueError("exact-case native raw input digest differs")
    mapped = _source_cells(boundary, public)
    actual = {_address(item["qualified_address"]): item["value"] for item in records}
    if len(actual) != len(records) or _digest(actual) != _digest(mapped):
        raise ValueError("exact-case native public/raw mapping differs")
    if (oracle.get("schema_version") != "gp.excel_oracle.v1"
            or oracle.get("tool") != "step5.oracle"
            or oracle.get("status") != "pass" or oracle.get("engine") != "Microsoft Excel"
            or oracle.get("engine_settings", {}).get("version") is None
            or oracle.get("engine_settings", {}).get("calculation") != -4105
            or oracle.get("calculation_state") != 0 or oracle.get("macros_executed") is not False
            or type(oracle.get("full_rebuild_seconds")) not in {int, float}
            or not math.isfinite(oracle["full_rebuild_seconds"])
            or oracle["full_rebuild_seconds"] < 0
            or oracle.get("source_sha256") != source_hash
            or oracle.get("source_copy_sha256") != source_hash
            or oracle.get("input_sha256") != producer["input_sha256"]
            or oracle.get("scenario_id") != raw.get("scenario_id")
            or oracle.get("input_overrides") != records
            or not {item["source_extent"] for item in boundary["sample_boundary"]["external_vectors"]}
            <= set(oracle.get("requested_ranges", []))):
        raise ValueError("exact-case native FullRebuild provenance differs")
    applied = oracle.get("applied_inputs")
    if not isinstance(applied, dict) or set(applied) != {item["qualified_address"] for item in records}:
        raise ValueError("exact-case native readback is incomplete")
    for item in records:
        value = applied[item["qualified_address"]]
        if (not isinstance(value, dict) or value.get("excel_error") is not None
                or value.get("value") != item["value"]
                or str(value.get("formula", "")).startswith("=")):
            raise ValueError(f"exact-case native readback differs: {item['qualified_address']}")
    workbook = load_workbook(source, data_only=False, read_only=False, keep_vba=False)
    try:
        cuts, count = _oracle_cuts(boundary, oracle, public, workbook, source_hash)
    finally:
        workbook.close()
    if (_digest(cuts) != _digest(producer.get("cuts"))
            or count != producer.get("numeric_equivalence", {}).get("representation_only_count")):
        raise ValueError("exact-case native cuts differ from sealed producer")


def load_bound_case(manifest_path: Path, case_id: str, source: Path,
                    boundary_path: Path) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    """Read one sealed case with its current source, request and boundary lineage."""
    manifest_path = manifest_path.expanduser().resolve()
    manifest = read_json(manifest_path)
    if (manifest.get("schema_version") != "step3.exact_cases.v1"
            or manifest.get("status") != "pass" or not isinstance(manifest.get("cases"), list)):
        raise ValueError("exact-case manifest is not a passing supported case set")
    matches = [case for case in manifest["cases"] if case.get("case_id") == case_id]
    if len(matches) != 1 or len({case.get("case_id") for case in manifest["cases"]}) != len(manifest["cases"]):
        raise ValueError(f"exact-case ID is missing or repeated: {case_id}")
    case = matches[0]
    labels = ("trace", "raw_request", "public_request", "producer")
    artifacts = {name: read_json(_ref(case.get(name), f"exact-case {case_id} {name}")[0])
                 for name in labels}
    producer, trace = artifacts["producer"], artifacts["trace"]
    source_hash = hash_file(source)
    boundary_hash = hash_file(boundary_path)
    if (producer.get("boundary_json", {}).get("sha256") != boundary_hash
            or producer.get("source_sha256") != source_hash
            or trace.get("schema_version") != "step3.exact_case_trace.v1"
            or trace.get("case_id") != case_id
            or trace.get("source_sha256") != source_hash
            or trace.get("input_sha256") != producer.get("input_sha256")
            or trace.get("request_sha256") != case["raw_request"]["sha256"]
            or trace.get("producer_sha256") != case["producer"]["sha256"]):
        raise ValueError(f"exact-case {case_id} source, trace, request or boundary differs")
    verify_native_case_producer(producer, case, source, artifacts["raw_request"], artifacts["public_request"])
    identity = {"case_id": case_id,
                "case_manifest": {"path": str(manifest_path), "sha256": hash_file(manifest_path)},
                "boundary_sha256": boundary_hash, "source_sha256": source_hash,
                **{name: case[name] for name in labels}}
    return identity, trace, artifacts


def verify_case_semantic_pair(case: dict[str, Any], source_sha256: str,
                              boundary_sha256: str) -> tuple[dict[str, Any], dict[str, Any], dict[str, Path]]:
    """Check one distinct profile/plan and every nested input before semantic replay."""
    if set(case) != {"trace", "raw_request", "public_request", "producer", "profile", "plan"}:
        raise ValueError("exact-case requires a distinct plan and profile")
    profile_path, _ = _ref(case["profile"], "exact-case profile")
    plan_path, _ = _ref(case["plan"], "exact-case plan")
    profile, plan = read_json(profile_path), read_json(plan_path)
    identity = plan.get("case_identity")
    if (not isinstance(identity, dict) or not isinstance(identity.get("case_id"), str)
            or identity.get("source_sha256") != source_sha256
            or identity.get("boundary_sha256") != boundary_sha256
            or any(identity.get(name) != case[name]
                   for name in ("trace", "raw_request", "public_request", "producer"))
            or profile.get("case_identity") != identity
            or profile.get("schema_version") != "step3.exact_case_family_profile.v1"
            or profile.get("selection_basis") != "exact_case"
            or plan.get("schema_version") != "step3.semantic_plan.v1"
            or plan.get("selection_basis") != "exact_case"
            or plan.get("status") != "pass" or profile.get("status") != "pass"
            or plan.get("source", {}).get("source_sha256") != source_sha256
            or not isinstance(plan.get("reference_resolution"), dict)
            or not isinstance(plan.get("execution_plan"), dict)
            or not isinstance(plan["execution_plan"].get("passes"), list)
            or plan.get("evidence", {}).get("source_family_profile") != case["profile"]
            or plan.get("evidence", {}).get("exact_case_trace") != case["trace"]
            or profile.get("exact_case_trace") != case["trace"]):
        raise ValueError("exact-case semantic pair belongs to another source, boundary, trace or request")
    nested = {"case_manifest": identity.get("case_manifest"),
              "semantic_map_input": plan.get("semantic_map_input"),
              "inventory": profile.get("inventory")}
    paths = {name: _ref(ref, f"exact-case {name}")[0] for name, ref in nested.items()}
    semantic_map = read_json(paths["semantic_map_input"])
    manifest = read_json(paths["case_manifest"])
    if (semantic_map.get("selection_basis") != "exact_case"
            or semantic_map.get("case_id") != identity["case_id"]
            or semantic_map.get("case_manifest") != nested["case_manifest"]
            or semantic_map.get("source_family_profile") != case["profile"]
            or semantic_map.get("exact_case_trace") != case["trace"]
            or semantic_map.get("source") != plan.get("source")
            or not isinstance(manifest.get("cases"), list)
            or len([item for item in manifest["cases"]
                    if item.get("case_id") == identity["case_id"]]) != 1
            or any(manifest_case.get(name) != case[name]
                   for manifest_case in manifest["cases"]
                   if manifest_case.get("case_id") == identity["case_id"]
                   for name in ("trace", "raw_request", "public_request", "producer"))):
        raise ValueError("exact-case semantic map or case manifest belongs to another request")
    return plan, profile, paths


def _produce_case(case: dict[str, Any], boundary: dict[str, Any], source: Path,
                  source_hash: str, out: Path, workbook: Any) -> dict[str, Any]:
    case_id = case["case_id"]
    paths = {label: _ref(case[label], f"{case_id} {label}")
             for label in ("raw_request", "public_request", "native_oracle")}
    raw, public, oracle = (read_json(paths[label][0]) for label in
                           ("raw_request", "public_request", "native_oracle"))
    records = raw.get("input_overrides")
    if not isinstance(records, list) or len(records) != boundary["sample_boundary"]["request_raw_cell_count"]:
        raise ValueError(f"{case_id}: raw request has the wrong input count")
    raw_cells = {_address(item["qualified_address"]): item["value"] for item in records}
    raw_digest = _digest(records)
    if len(raw_cells) != len(records) or raw.get("input_sha256") != raw_digest:
        raise ValueError(f"{case_id}: raw input digest or unique address set differs")
    public_cells = _source_cells(boundary, public)
    if _digest(raw_cells) != _digest(public_cells):
        raise ValueError(f"{case_id}: raw/public request mapping differs")
    if (raw.get("source_sha256") != source_hash or oracle.get("source_sha256") != source_hash
            or oracle.get("source_copy_sha256") != source_hash
            or oracle.get("input_sha256") != raw_digest
            or oracle.get("input_overrides") != records
            or oracle.get("scenario_id") != raw.get("scenario_id")
            or oracle.get("schema_version") != "gp.excel_oracle.v1"
            or oracle.get("tool") != "step5.oracle"
            or oracle.get("status") != "pass" or oracle.get("engine") != "Microsoft Excel"
            or oracle.get("macros_executed") is not False
            or oracle.get("engine_settings", {}).get("version") is None
            or oracle.get("engine_settings", {}).get("calculation") != -4105
            or oracle.get("calculation_state") != 0):
        raise ValueError(f"{case_id}: native FullRebuild provenance is incomplete or mismatched")
    applied = oracle.get("applied_inputs")
    if not isinstance(applied, dict) or set(applied) != {item["qualified_address"] for item in records}:
        raise ValueError(f"{case_id}: native input readback is incomplete")
    for item in records:
        readback = applied[item["qualified_address"]]
        if (readback.get("excel_error") is not None or readback.get("value") != item["value"]
                or readback.get("formula", "").startswith("=")):
            raise ValueError(f"{case_id}: native input readback differs at {item['qualified_address']}")
    cuts, normalized = _oracle_cuts(boundary, oracle, public, workbook, source_hash)
    case_out = out / case_id
    if case_out.exists():
        raise ValueError(f"case output already exists: {case_out}")
    case_out.mkdir()
    refs = {label: {"path": str(path), "sha256": digest}
            for label, (path, digest) in paths.items()}
    producer = {"schema_version": "step3.exact_case_producer.v1", "status": "pass",
                "case_id": case_id, "source_sha256": source_hash,
                "boundary_sha256": boundary["boundary_sha256"],
                "boundary_json": {"path": str(boundary["_artifact_path"]),
                                  "sha256": hash_file(boundary["_artifact_path"])},
                "input_sha256": raw_digest, "evidence": refs, "cuts": cuts,
                "numeric_equivalence": {"absolute": 0, "relative": 0,
                                        "representation_only_count": normalized}}
    producer_path = case_out / "producer.json"
    producer_path.write_text(json.dumps(producer, indent=2, sort_keys=True,
                                        ensure_ascii=False) + "\n", encoding="utf-8")
    producer_hash = hash_file(producer_path)
    verify_native_case_producer(producer, refs, source, raw, public)
    overrides = {}
    for extent, cut in cuts.items():
        ref = _parse_reference(extent, "")
        for index, value in enumerate(cut["values"]):
            cell = workbook[ref.sheet].cell(ref.min_row + index, ref.min_col)
            address = f"{ref.sheet}!{cell.coordinate}"
            formula = cell.value
            overrides[address] = {"boundary_variable_id": extent, "value": value,
                                  "source_formula": formula,
                                  "source_formula_sha256": hashlib.sha256(formula.encode()).hexdigest(),
                                  "capture_artifact_sha256": producer_hash}
    evaluator = FormulaEvaluator(source, external_overrides=overrides)
    try:
        for item in records:
            ref = _parse_reference(item["qualified_address"], "")
            cell = evaluator.workbook[ref.sheet].cell(ref.min_row, ref.min_col)
            if cell.data_type == "f":
                raise ValueError(f"{case_id}: raw override is a formula")
            cell.value = item["value"]
            if type(cell.value) is not type(item["value"]) or cell.value != item["value"]:
                raise ValueError(f"{case_id}: raw override readback differs")
        for target in boundary["target_selection"]["targets"]:
            for address in target["start_cells"]:
                ref = _parse_reference(address, "")
                evaluator.evaluate_cell(ref.sheet, ref.min_row, ref.min_col, dependency=False)
        trace = {"schema_version": "step3.exact_case_trace.v1", "case_id": case_id,
                 "source_sha256": source_hash, "input_sha256": raw_digest,
                 "request_sha256": paths["raw_request"][1], "producer_sha256": producer_hash,
                 "cells": list(evaluator.cells.values()), "edges": list(evaluator.edges.values()),
                 "names": list(evaluator.names.values()), "lookups": evaluator.lookups,
                 "branches": evaluator.branches, "dynamic_ranges": evaluator.dynamic_ranges}
    finally:
        evaluator.close()
    trace_path = case_out / "trace.json"
    trace_path.write_text(json.dumps(trace, indent=2, sort_keys=True,
                                     ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    if hash_file(source) != source_hash:
        raise ValueError(f"{case_id}: source workbook changed during trace")
    return {"case_id": case_id, "axes": _case_axes(boundary, public),
            "trace": {"path": str(trace_path), "sha256": hash_file(trace_path)},
            "producer": {"path": str(producer_path), "sha256": producer_hash},
            "raw_request": refs["raw_request"], "public_request": refs["public_request"],
            "native_oracle": refs["native_oracle"], "representation_only_count": normalized}


def produce_exact_cases(workflow_dir: Path, manifest_path: Path, out_dir: Path) -> dict[str, Any]:
    """Verify an explicit finite manifest and seal native cuts before tracing formulas."""
    try:
        root, workflow = load_workflow(workflow_dir)
        reference, boundary_path, _ = confirmed_boundary_reference(root, workflow)
        if reference is None or boundary_path is None:
            raise ValueError("a confirmed Stage 3 input boundary is required")
        boundary = read_json(boundary_path)
        boundary["_artifact_path"] = boundary_path
        source = Path(workflow["source"]["step1_source_copy_path"]).resolve()
        source_hash = workflow["source"]["workbook_sha256"]
        if not source.is_file() or hash_file(source) != source_hash:
            raise ValueError("current bound source copy is missing or changed")
        manifest = read_json(manifest_path)
        cases = manifest.get("cases") if isinstance(manifest, dict) else None
        sample = boundary["sample_boundary"]
        if (not isinstance(manifest, dict)
                or set(manifest) != {"schema_version", "cases"}
                or manifest["schema_version"] != "step3.exact_cases.manifest.v1"
                or not isinstance(cases, list) or len(cases) != sample["request_case_count"]
                or any(not isinstance(case, dict) or set(case) != {
                    "case_id", "raw_request", "public_request", "native_oracle"}
                       for case in cases)):
            raise ValueError("exact-case manifest must contain the declared complete case set")
        approved = {item["sha256"] for item in sample["request_cases"]}
        if len(approved) != len(cases) or {case["public_request"]["sha256"] for case in cases} != approved:
            raise ValueError("exact-case public requests differ from the confirmed boundary")
        ids = [case["case_id"] for case in cases]
        if len(set(ids)) != len(ids) or any(not isinstance(name, str) or not name
                                          or Path(name).name != name for name in ids):
            raise ValueError("exact-case IDs must be unique safe directory names")
        # Validate the complete Cartesian set before producing any output.
        tuples = [_case_axes(boundary, read_json(_ref(case["public_request"],
                                                "public request")[0])) for case in cases]
        axes = sample["case_axes"]
        expected = set(itertools.product(*axes.values()))
        if len(set(tuples)) != len(tuples) or set(tuples) != expected:
            raise ValueError("exact-case requests do not cover each declared scenario once")
        out = out_dir.expanduser().resolve()
        if out.exists():
            raise ValueError("exact-case output already exists")
        out.mkdir(parents=True)
        workbook = load_workbook(source, data_only=False, read_only=False, keep_vba=False)
        try:
            _verify_fixed_literals(boundary, workbook, source_hash)
            results = [_produce_case(case, boundary, source, source_hash, out, workbook)
                       for case in cases]
        finally:
            workbook.close()
        payload = {"schema_version": "step3.exact_cases.v1", "status": "pass",
                   "source_sha256": source_hash, "boundary_json_sha256": hash_file(boundary_path),
                   "manifest_sha256": hash_file(manifest_path), "cases": results}
        result_path = out / "exact_cases.json"
        result_path.write_text(json.dumps(payload, indent=2, sort_keys=True,
                                          ensure_ascii=False) + "\n", encoding="utf-8")
        return {"tool": "step3.produce_exact_cases", "status": "pass",
                "case_count": len(results), "artifact": str(result_path),
                "artifact_sha256": hash_file(result_path)}
    except (OSError, KeyError, TypeError, ValueError) as exc:
        return {"tool": "step3.produce_exact_cases", "status": "blocked", "reason": str(exc)}
