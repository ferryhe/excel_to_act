"""Build a reusable, source-bound syntactic profile of selected formula cells."""

from __future__ import annotations

from collections import Counter, defaultdict, deque
import hashlib
from pathlib import Path
import re
from types import SimpleNamespace
from typing import Any

from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils.cell import get_column_letter, range_boundaries

from excel_to_act.steps.step3 import workflow


_CANDIDATE_TRACE_SCHEMA = "gp.source_candidate_trace.v1"
_CANDIDATE_PROFILE_SCHEMA = "gp.source_candidate_family_profile.v1"
_MAX_RANGE_AREA = 100_000
_MAX_EXPANDED_COORDINATES = 1_000_000


def profile_source_families(analysis_dir: Path, trace_path: Path, out_dir: Path) -> dict[str, Any]:
    """Group trace-selected ordinary formulas by translated source syntax only."""
    return _profile_source_families(analysis_dir, trace_path, out_dir, candidate=False)


def profile_source_candidates(analysis_dir: Path, trace_path: Path, out_dir: Path) -> dict[str, Any]:
    """Profile formula-text candidates from a Step 3 source-only trace."""
    return _profile_source_families(analysis_dir, trace_path, out_dir, candidate=True)


def _profile_source_families(
    analysis_dir: Path, trace_path: Path, out_dir: Path, *, candidate: bool,
) -> dict[str, Any]:
    try:
        analysis_dir, manifest, binding, step1_root, inventory_path = workflow._load_context(analysis_dir)
        if binding.get("workflow_path"):
            from excel_to_act.steps.step3.input_boundary import require_confirmed_boundary_for_analysis

            require_confirmed_boundary_for_analysis(analysis_dir, manifest, binding)
        trace_path = trace_path.expanduser().resolve()
        out_dir = out_dir.expanduser().resolve()
        if not trace_path.is_file():
            raise ValueError("active trace does not exist")
        if out_dir == analysis_dir or out_dir.is_relative_to(analysis_dir):
            raise ValueError("profile output must be outside the bound Step 3 analysis directory")
        if out_dir == step1_root or out_dir.is_relative_to(step1_root):
            raise ValueError("profile output must be outside the read-only Step 1 source root")
        if candidate:
            from excel_to_act.steps.conversion_workflow import load_workflow

            workflow_root, _workflow_manifest = load_workflow(Path(binding["workflow_path"]))
            protected = [workflow_root]
            if out_dir in protected or any(out_dir.is_relative_to(item) for item in protected):
                raise ValueError("candidate profile output must be outside its bound workflow")

        inventory = workflow._load_inventory(inventory_path)
        trace = workflow._read_json(trace_path)
        expected_trace_schema = _CANDIDATE_TRACE_SCHEMA if candidate else None
        accepted_trace_schemas = ({expected_trace_schema} if candidate else
                                  {"step4.discovery_trace.v1", "gp.active_trace.v1"})
        if not isinstance(trace, dict) or trace.get("schema_version") not in accepted_trace_schemas:
            label = "source-candidate trace" if candidate else "source-bound active trace"
            raise ValueError(f"trace must use a supported {label} schema")
        source_sha = manifest["binding"]["source"]["source_sha256"]
        if trace.get("source_sha256") != source_sha or inventory.get("workbook_sha256") != source_sha:
            raise ValueError("trace, inventory, and prepared analysis do not bind the same workbook")
        if candidate:
            _validate_candidate_trace_binding(analysis_dir, manifest, binding, trace, trace_path, inventory)
            scope_reference = trace.get("model_scope")
            if isinstance(scope_reference, dict):
                protected_scope_inputs = [Path(scope_reference["path"]).resolve()]
                full_trace = scope_reference.get("full_candidate_trace")
                if isinstance(full_trace, dict) and isinstance(full_trace.get("path"), str):
                    protected_scope_inputs.append(Path(full_trace["path"]).resolve())
                if any(out_dir == item or out_dir.is_relative_to(item) or item.is_relative_to(out_dir)
                       for item in protected_scope_inputs):
                    raise ValueError("candidate profile output must be outside its model-scope and full-trace inputs")
        elif trace.get("native_excel_called") is not False or trace.get("formula_cache_inputs") is not False:
            raise ValueError("trace must be source-only and must not use formula caches")
        records = trace.get("cells")
        if not isinstance(records, list):
            raise ValueError("trace cells must be a list")

        source_cells = {
            _address_key(sheet["name"], cell["address"]): cell
            for sheet in inventory["sheets"]
            for cell in sheet.get("cells", [])
        }
        families: dict[tuple[str, str, str], dict[str, Any]] = {}
        seen: set[str] = set()
        for record in records:
            if not isinstance(record, dict) or record.get("role") != "calculated_formula":
                continue
            sheet, address = record.get("sheet"), record.get("address")
            if not isinstance(sheet, str) or not isinstance(address, str):
                raise ValueError("trace formula record is missing its sheet or address")
            key = _address_key(sheet, address)
            if key in seen:
                raise ValueError(f"trace contains a duplicate ordinary formula address: {sheet}!{address}")
            seen.add(key)
            cell = source_cells.get(key)
            formula = cell.get("formula") if cell else None
            if not isinstance(formula, str) or record.get("formula") != formula:
                raise ValueError(f"trace formula differs from the current canonical inventory at {key}")
            if candidate and record.get("formula_sha256") != _formula_sha256(formula):
                raise ValueError(f"trace formula hash differs from the current canonical inventory at {key}")
            local_address = _local_address(sheet, address)
            column = "".join(character for character in local_address if character.isalpha()).upper()
            try:
                compact = "=" + "".join(token.value for token in Tokenizer(formula).items
                                         if token.type != "WHITE-SPACE")
                normalized = Translator(compact, origin=local_address).translate_formula(f"{column}1000")
            except (TokenizerError, TranslatorError, ValueError, TypeError, IndexError) as exc:
                raise ValueError(f"cannot normalize source formula at {key}: {type(exc).__name__}") from exc
            family_key = (sheet.casefold(), column, normalized)
            family = families.setdefault(family_key, {
                "sheet": sheet, "column": column, "normalized_formula": normalized,
                "representative_address": f"{sheet}!{local_address}", "representative_formula": formula,
                "source_members": [],
            })
            family["source_members"].append(f"{sheet}!{local_address}")
        if not seen and not (candidate and trace.get("arrays")):
            raise ValueError("trace contains no ordinary formula cells to profile")

        ordered: list[dict[str, Any]] = []
        for key, family in sorted(families.items()):
            family_id = "syntax-" + workflow._hash_bytes(
                workflow._json_bytes([family["sheet"], family["column"], family["normalized_formula"]])
            )[:16]
            ordered.append({**family, "family_id": family_id,
                            "source_members": sorted(family["source_members"], key=str.casefold),
                            "member_count": len(family["source_members"])})
        profiled_members = [_address_key(*_split_address(address))
                            for family in ordered for address in family["source_members"]]
        if len(profiled_members) != len(seen) or set(profiled_members) != seen:
            raise ValueError("profile family membership does not exactly match the ordinary formula trace")

        sheet_formulas: Counter[str] = Counter()
        sheet_families: Counter[str] = Counter()
        for family in ordered:
            sheet_formulas[family["sheet"]] += family["member_count"]
            sheet_families[family["sheet"]] += 1
        trace_sha = workflow._hash_file(trace_path)
        schema = _CANDIDATE_PROFILE_SCHEMA if candidate else "gp.source_family_profile.v1"
        trace_key = "source_candidate_trace" if candidate else "historical_selection_trace"
        payload = {
            "schema_version": schema, "status": "pass",
            "source_sha256": source_sha, "binding_sha256": manifest["binding_sha256"],
            "inventory": {"path": str(inventory_path), "sha256": workflow._hash_file(inventory_path)},
            trace_key: {"path": str(trace_path), "sha256": trace_sha},
            "selection_basis": "static_candidate" if candidate else "active_trace",
            "policy": {
                "formula_calculation_performed": False, "native_excel_called": False,
                "formula_cache_used": False, "trace_computed_values_read": False,
                "selection_scope": ("Ordinary formula addresses in the source-only static candidate trace" if candidate
                                    else "Ordinary formula addresses in the supplied source-bound active trace only"),
                "grouping": "Same source sheet/column and identical tokenized formula after row translation",
                "limitation": ("Static branch-union candidates only; unresolved paths and arrays are not active selection" if candidate
                               else "Syntactic candidates only; no business grouping, algebraic proof, or new active closure"),
            },
            "summary": {"ordinary_formula_members": len(seen), "syntactic_families": len(ordered),
                        "array_member_handling": "Array followers require separate semantic shape mapping",
                        **({"candidate_array_instance_count": len(trace.get("arrays", [])),
                            "unresolved_path_count": trace.get("summary", {}).get("unresolved_path_count", 0),
                            "truncated_path_count": trace.get("summary", {}).get("truncated_path_count", 0),
                            "candidate_closure": "unknown"} if candidate else {})},
            "sheets": [{"sheet": sheet, "ordinary_formula_members": sheet_formulas[sheet],
                        "syntactic_families": sheet_families[sheet]} for sheet in sorted(sheet_formulas)],
            "families": ordered,
        }
        if candidate:
            payload["arrays"] = trace.get("arrays", [])
            if trace.get("model_scope") is not None:
                payload["candidate_scope"] = trace["candidate_scope"]
                payload["model_scope"] = trace["model_scope"]
                payload["summary"]["scope_frontier_count"] = trace["summary"].get("scope_frontier_count", 0)
                payload["summary"]["proposed_formula_start_count"] = trace["summary"].get(
                    "proposed_formula_start_count", 0)
        markdown = _profile_markdown(payload)
        basename = "source_candidate_family_profile" if candidate else "source_family_profile"
        json_path, md_path = out_dir / f"{basename}.json", out_dir / f"{basename}.md"
        json_bytes = workflow._json_bytes(payload)
        md_bytes = markdown.encode("utf-8")
        out_dir.mkdir(parents=True, exist_ok=True)
        workflow._write_bytes(json_path, json_bytes)
        workflow._write_bytes(md_path, md_bytes)
        return {"schema_version": payload["schema_version"], "tool": "step3.profile", "status": "pass",
                "profile": str(json_path), "profile_sha256": workflow._hash_bytes(json_bytes),
                "markdown": str(md_path), "markdown_sha256": workflow._hash_bytes(md_bytes),
                "summary": payload["summary"], "selection_basis": payload["selection_basis"],
                **({"candidate_scope": payload["candidate_scope"]} if payload.get("candidate_scope") else {}),
                "formula_cache_used": False, "native_excel_called": False}
    except Exception as exc:
        return {"tool": "step3.profile", "status": "blocked",
                "diagnostics": [{"code": "profile_input_invalid", "severity": "error", "message": str(exc)}]}


def build_source_candidate_trace(
    analysis_dir: Path, workflow_dir: Path, out_dir: Path, model_scope_path: Path | None = None,
) -> dict[str, Any]:
    """Trace static formula-reference candidates without values or Step 4 history."""
    try:
        from excel_to_act.steps.conversion_workflow import load_workflow
        from excel_to_act.steps.step3.input_boundary import (
            _require_bound_workflow,
            _validate_catalog_identity,
            _validate_targets,
            confirmed_boundary_reference,
        )

        analysis_dir, manifest, binding, step1_root, inventory_path = workflow._load_context(analysis_dir)
        workflow_root, workflow_manifest = load_workflow(workflow_dir)
        _require_bound_workflow(analysis_dir, binding, workflow_root, workflow_manifest)
        boundary_reference, boundary_path, _boundary_md = confirmed_boundary_reference(workflow_root, workflow_manifest)
        fields, fields_sha = workflow._read_stage(analysis_dir, manifest, "fields")
        _dependencies, dependencies_sha = workflow._read_stage(analysis_dir, manifest, "dependencies")
        _plan, plan_sha = workflow._read_stage(analysis_dir, manifest, "plan")
        inventory = workflow._load_inventory(inventory_path)
        boundary = workflow._read_json(boundary_path)
        target_selection, targets_sha, catalog_sha, targets_path, catalog_path = _read_confirmed_selection(
            boundary, boundary_path, binding, manifest, inventory, boundary_reference,
            _validate_targets, _validate_catalog_identity,
        )
        out_dir = out_dir.expanduser().resolve()
        model_scope_path = model_scope_path.expanduser().resolve() if model_scope_path is not None else None
        protected = [analysis_dir, step1_root, workflow_root, boundary_path,
                     targets_path, catalog_path]
        source_path_value = binding["source"].get("source_path")
        if isinstance(source_path_value, str) and source_path_value:
            source_path = Path(source_path_value).expanduser()
            if not source_path.is_absolute():
                run_path = workflow._resolved_child(step1_root, binding["source"]["run_path"], "Step 1 run path")
                source_path = (run_path / source_path).resolve()
            if source_path.is_file():
                protected.append(source_path)
        for item in protected:
            if (out_dir == item or out_dir.is_relative_to(item)
                    or item.is_relative_to(out_dir)):
                raise ValueError("source-trace output must be outside bound analysis, workflow, source, and Step 1 roots")

        source_sha = binding["source"]["source_sha256"]
        ignored = {str(item).casefold() for item in binding["scope"].get("ignored_sheets", [])}
        all_cells: dict[str, tuple[str, dict[str, Any]]] = {}
        formulas: dict[str, tuple[str, str, dict[str, Any]]] = {}
        sheet_names: dict[str, str] = {}
        for sheet in inventory["sheets"]:
            name = str(sheet["name"])
            sheet_names[name.casefold()] = name
            for cell in sheet.get("cells", []):
                address = _local_address(name, str(cell.get("address", "")))
                key = _address_key(name, address)
                all_cells[key] = (name, cell)
                formula = cell.get("formula")
                if isinstance(formula, str) and formula.startswith("="):
                    formulas[key] = (name, address.replace("$", "").upper(), cell)

        raw_candidates = _catalog_raw_coordinates(boundary)
        raw_candidate_coordinates = set().union(*raw_candidates.values()) if raw_candidates else set()
        cut_coordinates, cut_extents = _catalog_formula_cuts(boundary, formulas, source_sha)
        exclusion_coordinates, exclusions = _catalog_upstream_exclusions(boundary)
        analysis_reference = {
            "path": str(analysis_dir), "binding_sha256": manifest["binding_sha256"],
            "fields_sha256": fields_sha, "dependencies_sha256": dependencies_sha,
            "structural_plan_sha256": plan_sha, "inventory_path": str(inventory_path),
            "inventory_sha256": workflow._hash_file(inventory_path),
        }
        boundary_reference_with_inputs = {**boundary_reference,
                                          "catalog_input_sha256": catalog_sha,
                                          "targets_input_sha256": targets_sha}
        model_scope = None
        scope_regions: dict[str, list[tuple[int, int, int, int]]] = {}
        if model_scope_path is not None:
            if not model_scope_path.is_file():
                raise ValueError("model-scope JSON does not exist")
            model_scope, scope_regions = _validate_model_scope(
                model_scope_path, analysis_dir, manifest, binding, inventory,
                {key: binding["source"][key] for key in ("source_id", "run_id", "source_sha256")},
                target_selection["scenario"], target_selection["targets"], analysis_reference,
                boundary_reference_with_inputs, formulas, cut_coordinates, exclusion_coordinates,
            )
            protected.extend([model_scope_path, Path(model_scope["full_candidate_trace"]["path"])])
            for item in protected:
                if (out_dir == item or out_dir.is_relative_to(item)
                        or item.is_relative_to(out_dir)):
                    raise ValueError("source-trace output must be outside its model-scope input and bound full trace")
        graph_payload = dict(inventory)
        graph_payload["sheets"] = [
            ({**sheet, "cells": []} if str(sheet["name"]).casefold() in ignored else sheet)
            for sheet in inventory["sheets"]
        ]
        from excel_to_act.graph.builder import RegexFormulaGraphBuilder
        from excel_to_act.schemas import WorkbookInventory

        graph = RegexFormulaGraphBuilder().build(WorkbookInventory.model_validate(graph_payload))
        nodes = {node.id: node for node in graph.nodes}
        edges_by_formula: dict[str, list[Any]] = defaultdict(list)
        for edge in graph.edges:
            location = edge.source_location
            if location and location.sheet_name and location.address:
                edges_by_formula[_address_key(location.sheet_name, location.address)].append(nodes[edge.target])
        unsupported_by_formula: dict[str, list[str]] = defaultdict(list)
        for item in graph.unsupported_features:
            location = item.source_location
            if location and location.sheet_name and location.address:
                unsupported_by_formula[_address_key(location.sheet_name, location.address)].append(item.description)

        array_anchors, array_followers = _inventory_arrays(inventory, all_cells)
        formula_records: dict[str, dict[str, Any]] = {}
        raw_reached: set[str] = set()
        other_sources: set[str] = set()
        unknown_by_key: dict[tuple[str, str, str], dict[str, Any]] = {}
        expanded_coordinates = 0
        expanded_coordinate_keys: set[str] = set()
        expanded_references: set[str] = set()
        scope_clipping_cache: dict[
            str, tuple[list[tuple[int, int, int, int]], list[tuple[int, int, int, int]], int]
        ] = {}
        truncated = False
        cut_hits: set[str] = set()
        visited: set[str] = set()
        scope_frontiers_by_key: dict[tuple[str, str, str], dict[str, Any]] = {}
        fields_by_id, fields_cell_index = workflow._field_index(fields)
        pending: deque[str] = deque()
        target_starts = []
        for target in target_selection["targets"]:
            for address in target["start_cells"]:
                target_starts.append({"target_id": target["target_id"], "source_ref": address})
                pending.append(_address_key(*_split_address(address)))
        if model_scope is not None:
            pending.extend(model_scope["proposed_formula_start_keys"])

        def unknown(address: str, category: str, evidence: str) -> None:
            key = (address, category, evidence)
            unknown_by_key.setdefault(key, {"address": address, "category": category, "evidence": evidence})

        def visit_coordinate(address_key: str, *, from_formula: str, reference: str) -> None:
            nonlocal expanded_coordinates, truncated
            if model_scope is not None and not _scope_contains_key(scope_regions, address_key):
                _record_scope_frontier(scope_frontiers_by_key, unknown, formulas, from_formula,
                                       reference, address_key, "coordinate")
                return
            if address_key in expanded_coordinate_keys:
                return
            expanded_coordinate_keys.add(address_key)
            expanded_coordinates += 1
            if expanded_coordinates > _MAX_EXPANDED_COORDINATES:
                truncated = True
                unknown(from_formula, "trace_truncated", f"Expanded-coordinate cap reached while reading {reference}")
                return
            if address_key in cut_coordinates:
                cut_hits.add(address_key)
                return
            if address_key in exclusion_coordinates:
                unknown(from_formula, "unapproved_upstream_exclusion", reference)
                return
            member = formulas.get(address_key)
            if member is not None:
                if member[0].casefold() in ignored:
                    unknown(from_formula, "ignored_sheet_boundary", f"{member[0]}!{member[1]}")
                else:
                    pending.append(address_key)
                return
            if address_key in array_followers:
                anchor = array_followers[address_key]
                unknown(from_formula, "array_follower_reference", f"{address_key} maps to anchor {anchor}")
                if anchor in formulas:
                    if model_scope is not None and not _scope_contains_key(scope_regions, anchor):
                        _record_scope_frontier(scope_frontiers_by_key, unknown, formulas, from_formula,
                                               reference, anchor, "array_anchor")
                    else:
                        pending.append(anchor)
                return
            if address_key in raw_candidate_coordinates:
                raw_reached.add(address_key)
            cell_entry = all_cells.get(address_key)
            if cell_entry is None:
                unknown(from_formula, "unrepresented_source_coordinate", address_key)
            elif cell_entry[1].get("kind") == "blank":
                unknown(from_formula, "blank_source_read_semantics", address_key)
            else:
                other_sources.add(address_key)

        def visit_reference(node: Any, formula_address: str) -> None:
            nonlocal truncated
            kind = getattr(node.kind, "value", str(node.kind))
            reference_text = node.label
            if kind in {"cell", "range"}:
                formula_key = _address_key(*_split_address(formula_address))
                formula_record = formulas.get(formula_key)
                if formula_record is not None:
                    reference_text = _source_reference_text(formula_record[2].get("formula", ""),
                                                             node.label, formula_record[0])
            if node.metadata.get("structured_reference"):
                unknown(formula_address, "structured_reference_selection", str(node.metadata["structured_reference"]))
                return
            if kind == "external":
                unknown(formula_address, "external_reference", node.label)
                return
            if kind == "name":
                local = node.metadata.get("address")
                source_location = node.source_location
                sheet = source_location.sheet_name if source_location else None
                if not isinstance(local, str) or not local or not sheet:
                    unknown(formula_address, "unresolved_defined_name", node.label)
                    return
                if local.startswith("=") or local.startswith("#"):
                    unknown(formula_address, "formula_defined_name", node.label)
                    return
                reference = local if "!" in local else f"{sheet}!{local}"
            elif kind in {"cell", "range"}:
                reference = node.label
                if "!" not in reference:
                    referring_sheet, _ = _split_address(formula_address)
                    reference = f"{referring_sheet}!{reference}"
            else:
                unknown(formula_address, "unsupported_reference", node.label)
                return
            try:
                sheet, local = _split_address(reference)
                canonical_sheet = sheet_names.get(sheet.casefold())
                if canonical_sheet is None:
                    unknown(formula_address, "unresolved_reference_sheet", reference)
                    return
                bounds = range_boundaries(local.replace("$", ""))
            except (ValueError, TypeError):
                unknown(formula_address, "unresolved_reference", reference)
                return
            if None in bounds:
                if model_scope is None:
                    unknown(formula_address, "unbounded_range", reference)
                    return
                # Excel's whole-row/whole-column references are bounded by the
                # worksheet grid. Clipping that finite grid to the small
                # Agent-proposed regions avoids expanding a million-row span.
                min_col, min_row, max_col, max_row = bounds
                bounds = (1 if min_col is None else min_col,
                          1 if min_row is None else min_row,
                          16_384 if max_col is None else max_col,
                          1_048_576 if max_row is None else max_row)
            min_col, min_row, max_col, max_row = bounds
            area = (max_col - min_col + 1) * (max_row - min_row + 1)
            if model_scope is None and area > _MAX_RANGE_AREA:
                truncated = True
                unknown(formula_address, "range_expansion_truncated", f"{reference} has {area} cells")
                return
            reference_key = f"{canonical_sheet.casefold()}!{local.replace('$', '').upper()}"
            if model_scope is None:
                if reference_key in expanded_references:
                    return
                expanded_references.add(reference_key)
                regions_to_expand = [(min_col, min_row, max_col, max_row)]
            else:
                clipped = scope_clipping_cache.get(reference_key)
                if clipped is None:
                    regions_to_expand, excluded_regions = _clip_to_scope(
                        (min_col, min_row, max_col, max_row), scope_regions.get(canonical_sheet.casefold(), []))
                    included_area = sum(_rect_area(region) for region in regions_to_expand)
                    clipped = (regions_to_expand, excluded_regions, included_area)
                    scope_clipping_cache[reference_key] = clipped
                regions_to_expand, excluded_regions, included_area = clipped
                for excluded_region in excluded_regions:
                    _record_scope_frontier(scope_frontiers_by_key, unknown, formulas, formula_address,
                                           reference_text, f"{canonical_sheet}!{_rect_to_a1(excluded_region)}",
                                           "region")
                if included_area > _MAX_RANGE_AREA:
                    truncated = True
                    unknown(formula_address, "range_expansion_truncated",
                            f"{reference} has {included_area} in-scope cells")
                    regions_to_expand = []
                if reference_key in expanded_references:
                    return
                expanded_references.add(reference_key)
            for region in regions_to_expand:
                region_min_col, region_min_row, region_max_col, region_max_row = region
                for row in range(region_min_row, region_max_row + 1):
                    for column in range(region_min_col, region_max_col + 1):
                        address = f"{canonical_sheet}!{get_column_letter(column)}{row}"
                        visit_coordinate(_address_key(canonical_sheet, address),
                                         from_formula=formula_address, reference=reference_text)

        while pending:
            address_key = pending.popleft()
            if address_key in visited:
                continue
            visited.add(address_key)
            if address_key in cut_coordinates:
                cut_hits.add(address_key)
                continue
            if address_key in exclusion_coordinates:
                unknown(address_key, "unapproved_upstream_exclusion", address_key)
                continue
            item = formulas.get(address_key)
            if item is None:
                visit_coordinate(address_key, from_formula=address_key, reference=address_key)
                continue
            sheet, local, cell = item
            if sheet.casefold() in ignored:
                unknown(address_key, "ignored_sheet_boundary", address_key)
                continue
            formula = cell["formula"]
            field_id = fields_cell_index.get((sheet.casefold(), local.upper()))
            if field_id is None or fields_by_id[field_id].get("role") != "calculated":
                raise ValueError(f"candidate formula is not represented by a current calculated field: {sheet}!{local}")
            formula_records[address_key] = {
                "sheet": sheet, "address": f"{sheet}!{local}", "formula": formula,
                "formula_sha256": _formula_sha256(formula), "role": "calculated_formula",
            }
            _record_formula_semantic_unknowns(formula, f"{sheet}!{local}", unknown)
            scoped_unbounded_operands = _unbounded_a1_range_operands(formula) if model_scope is not None else []
            for diagnostic in unsupported_by_formula.get(address_key, []):
                if model_scope is not None and any(operand in diagnostic for operand in scoped_unbounded_operands):
                    continue
                unknown(f"{sheet}!{local}", _unknown_category(diagnostic), diagnostic)
            if cell.get("raw_formula_attributes", {}).get("t") == "array":
                array = array_anchors.get(address_key)
                if array is None:
                    unknown(f"{sheet}!{local}", "array_geometry_unresolved", "No finite canonical array extent")
                elif array["missing_extent_coordinates"]:
                    unknown(f"{sheet}!{local}", "array_geometry_incomplete", array["array_ref"])
            for node in edges_by_formula.get(address_key, []):
                visit_reference(node, f"{sheet}!{local}")
            for operand in scoped_unbounded_operands:
                visit_reference(SimpleNamespace(kind="range", label=operand, metadata={}), f"{sheet}!{local}")

        if truncated:
            unknown("", "trace_truncated", "At least one configured expansion cap was reached")
        candidate_cells = sorted(formula_records.values(), key=lambda item: item["address"].casefold())
        arrays = [array_anchors[key] for key in sorted(array_anchors) if key in formula_records]
        raw_coverage = _classify_raw_candidates(boundary, raw_candidates, raw_reached,
                                                bool(unknown_by_key) or bool(cut_coordinates))
        cut_records = [cut_extents[key] for key in sorted(cut_extents)]
        unknowns = [unknown_by_key[key] for key in sorted(unknown_by_key)]
        scope_frontiers = sorted(scope_frontiers_by_key.values(),
                                 key=lambda item: (item["referring_formula"].casefold(),
                                                   item["reference"].casefold(),
                                                   item["excluded_source_ref"].casefold()))
        payload = {
            "schema_version": _CANDIDATE_TRACE_SCHEMA, "tool": "step3.source-trace", "status": "candidate_trace",
            "selection_basis": "static_candidate", "source_sha256": source_sha,
            "binding_sha256": manifest["binding_sha256"],
            "source": {key: binding["source"][key] for key in ("source_id", "run_id", "source_sha256")},
            "scenario": target_selection["scenario"],
            "scenario_sha256": target_selection["scenario"]["scenario_sha256"],
            "targets": target_selection["targets"],
            "analysis": analysis_reference,
            "input_boundary": boundary_reference_with_inputs,
            "formula_calculation_performed": False, "formula_cache_used": False,
            "formula_cache_inputs": False, "native_excel_called": False, "trace_computed_values_read": False,
            "cells": candidate_cells, "arrays": arrays,
            "external_formula_cuts": cut_records,
            "upstream_exclusions": exclusions,
            "raw_candidate_reachability": raw_coverage,
            "other_referenced_source_coordinates": sorted(other_sources, key=str.casefold),
            "unknowns": unknowns,
            "truncated": truncated,
            "summary": {
                "target_count": len(target_selection["targets"]),
                "candidate_formula_member_count": len(candidate_cells),
                "candidate_array_instance_count": len(arrays),
                "candidate_array_follower_count": sum(item["follower_count"] for item in arrays),
                "declared_cut_formula_coordinate_count": len(cut_coordinates),
                "reached_cut_formula_coordinate_count": len(cut_hits),
                **({"scope_frontier_count": len(scope_frontiers),
                    "scope_frontier_coordinate_count": sum(item["coordinate_count"] for item in scope_frontiers),
                    "proposed_formula_start_count": sum(len(item["formula_members"])
                                                        for item in model_scope["proposed_formula_starts"])}
                   if model_scope is not None else {}),
                "raw_candidate_object_count": len(raw_coverage["objects"]),
                "raw_reached_coordinate_count": len(raw_reached),
                "unresolved_path_count": len(unknowns),
                "truncated_path_count": int(truncated),
                "candidate_formula_selection_complete_for_known_references": not truncated and not scope_frontiers,
                "candidate_closure": "unknown",
                "fresh_step4_active_formula_members": None,
                "numerically_reconciled_members": None,
            },
            "limitations": [
                "Both statically readable conditional arms are included; branch selection is unknown.",
                "Lookup ranges are candidate unions; selected rows and return values are unknown.",
                "Dynamic, broken, unsupported, excluded, and upstream-excluded references remain explicit unknowns.",
                "Formula-derived external extents are provenance-only cuts; their prerequisite formulas are not traversed.",
                "A static candidate trace is not a Step 4 active trace, formula evaluation, or numerical result.",
            ],
        }
        if model_scope is not None:
            payload["candidate_scope"] = model_scope["candidate_scope"]
            payload["model_scope"] = _public_model_scope(model_scope)
            payload["candidate_starts"] = {
                "target_starts": target_starts,
                "proposed_formula_starts": model_scope["proposed_formula_starts"],
            }
            payload["scope_frontiers"] = scope_frontiers
            payload["limitations"].append(
                "The model-scope fence is an Agent-proposed source-coordinate slice; omitted references remain explicit frontiers, not evidence that they are unused."
            )
        markdown = _source_candidate_markdown(payload)
        json_bytes, md_bytes = workflow._json_bytes(payload), markdown.encode("utf-8")
        out_dir.mkdir(parents=True, exist_ok=True)
        json_path, md_path = out_dir / "source_candidate_trace.json", out_dir / "source_candidate_trace.md"
        workflow._write_bytes(json_path, json_bytes)
        workflow._write_bytes(md_path, md_bytes)
        return {"schema_version": _CANDIDATE_TRACE_SCHEMA, "tool": "step3.source-trace", "status": "pass",
                "trace": str(json_path), "trace_sha256": workflow._hash_bytes(json_bytes),
                "markdown": str(md_path), "markdown_sha256": workflow._hash_bytes(md_bytes),
                "summary": payload["summary"],
                **({"candidate_scope": payload["candidate_scope"]} if payload.get("candidate_scope") else {}),
                "formula_cache_used": False, "native_excel_called": False}
    except Exception as exc:
        return {"schema_version": _CANDIDATE_TRACE_SCHEMA, "tool": "step3.source-trace", "status": "blocked",
                "diagnostics": [{"code": "source_trace_invalid", "severity": "error", "message": str(exc)}]}


def _read_confirmed_selection(
    boundary: dict[str, Any], boundary_path: Path, binding: dict[str, Any], manifest: dict[str, Any],
    inventory: dict[str, Any], boundary_reference: dict[str, Any], validate_targets: Any,
    validate_catalog_identity: Any,
) -> tuple[dict[str, Any], str, str, Path, Path]:
    from excel_to_act.steps import conversion_workflow

    source = binding["source"]
    expected_source = {"source_id": source["source_id"], "run_id": source["run_id"],
                       "workbook_sha256": source["source_sha256"],
                       "analysis_binding_sha256": manifest["binding_sha256"]}
    if (boundary.get("schema_version") != "step3.input_boundary.v1"
            or boundary.get("source") != expected_source
            or boundary.get("boundary_sha256") != boundary_reference.get("boundary_sha256")):
        raise ValueError("confirmed catalog is not bound to the current source and analysis")
    if inventory.get("workbook_sha256") != source["source_sha256"]:
        raise ValueError("canonical inventory does not match the confirmed catalog source")
    input_artifacts = boundary.get("input_artifacts")
    if not isinstance(input_artifacts, dict):
        raise ValueError("confirmed boundary has no target/catalog input hashes")
    paths: dict[str, Path] = {}
    hashes: dict[str, str] = {}
    for key in ("targets", "catalog_input"):
        path_value = input_artifacts.get(f"{key}_path")
        sha_value = input_artifacts.get(f"{key}_sha256")
        if not isinstance(path_value, str) or not isinstance(sha_value, str):
            raise ValueError(f"confirmed boundary is missing its {key} input reference")
        path = Path(path_value).expanduser().resolve()
        if not path.is_file() or conversion_workflow.hash_file(path) != sha_value:
            raise ValueError(f"confirmed boundary {key} input is missing or changed")
        paths[key], hashes[key] = path, sha_value
    targets_value = conversion_workflow.read_json(paths["targets"])
    catalog_value = conversion_workflow.read_json(paths["catalog_input"])
    target_selection = validate_targets(targets_value, binding, manifest["binding_sha256"])
    validate_catalog_identity(catalog_value, binding, target_selection)
    if boundary.get("target_selection") != target_selection:
        raise ValueError("confirmed catalog target/scenario metadata differs from its reviewed target input")
    scenario = target_selection["scenario"]
    scenario_basis = {key: value for key, value in scenario.items() if key != "scenario_sha256"}
    if workflow._hash_bytes(workflow._json_bytes(scenario_basis)) != scenario.get("scenario_sha256"):
        raise ValueError("confirmed scenario hash does not match its declared metadata")
    return target_selection, hashes["targets"], hashes["catalog_input"], paths["targets"], paths["catalog_input"]


def _validate_model_scope(
    scope_path: Path, analysis_dir: Path, manifest: dict[str, Any], binding: dict[str, Any],
    inventory: dict[str, Any], source: dict[str, Any], scenario: dict[str, Any],
    targets: list[dict[str, Any]], analysis_reference: dict[str, Any],
    boundary_reference: dict[str, Any], formulas: dict[str, tuple[str, str, dict[str, Any]]],
    cut_coordinates: set[str], exclusion_coordinates: set[str],
) -> tuple[dict[str, Any], dict[str, list[tuple[int, int, int, int]]]]:
    value = workflow._read_json(scope_path)
    allowed_keys = {"schema_version", "status", "source", "scenario", "targets", "analysis",
                    "input_boundary", "full_candidate_trace", "allowed_regions", "proposed_formula_starts"}
    if not isinstance(value, dict) or set(value) - allowed_keys:
        raise ValueError("model-scope input has an unsupported schema or unknown fields")
    if value.get("schema_version") != "gp.source_model_scope.input.v1" or value.get("status") != "agent_proposed":
        raise ValueError("model-scope input must be an Agent-proposed gp.source_model_scope.input.v1")
    for key, expected in (("source", source), ("scenario", scenario), ("targets", targets),
                          ("analysis", analysis_reference), ("input_boundary", boundary_reference)):
        if value.get(key) != expected:
            raise ValueError(f"model-scope {key} binding differs from the current confirmed source and analysis")

    full_ref = value.get("full_candidate_trace")
    if (not isinstance(full_ref, dict) or set(full_ref) != {"path", "sha256"}
            or not isinstance(full_ref.get("path"), str) or not isinstance(full_ref.get("sha256"), str)):
        raise ValueError("model-scope input needs a full_candidate_trace path and SHA-256")
    full_path = Path(full_ref["path"]).expanduser()
    if not full_path.is_absolute():
        full_path = (scope_path.parent / full_path).resolve()
    else:
        full_path = full_path.resolve()
    if not full_path.is_file() or workflow._hash_file(full_path) != full_ref["sha256"]:
        raise ValueError("bound full candidate trace is missing or changed")
    full_trace = workflow._read_json(full_path)
    if not isinstance(full_trace, dict) or full_trace.get("schema_version") != _CANDIDATE_TRACE_SCHEMA:
        raise ValueError("bound full candidate trace has an unsupported schema")
    if full_trace.get("model_scope") is not None or full_trace.get("candidate_scope") is not None:
        raise ValueError("model-scope reference must point to the immutable unscoped candidate trace")
    _validate_candidate_trace_binding(analysis_dir, manifest, binding, full_trace, full_path, inventory)
    for key, expected in (("source", source), ("scenario", scenario), ("targets", targets),
                          ("analysis", analysis_reference), ("input_boundary", boundary_reference)):
        if full_trace.get(key) != expected:
            raise ValueError(f"full candidate trace {key} differs from the current model-scope binding")

    raw_regions = value.get("allowed_regions")
    if not isinstance(raw_regions, list) or not raw_regions:
        raise ValueError("model-scope input needs at least one allowed source region")
    sheet_names = {str(item["name"]).casefold(): str(item["name"]) for item in inventory.get("sheets", [])}
    scope_regions: dict[str, list[tuple[int, int, int, int]]] = defaultdict(list)
    normalized_regions: list[dict[str, str]] = []
    for region in raw_regions:
        if (not isinstance(region, dict) or set(region) != {"source_ref", "rationale"}
                or not isinstance(region.get("source_ref"), str) or not region["source_ref"].strip()
                or not isinstance(region.get("rationale"), str) or not region["rationale"].strip()):
            raise ValueError("each allowed region needs only a finite source_ref and non-empty rationale")
        sheet, local = _split_address(region["source_ref"])
        canonical_sheet = sheet_names.get(sheet.casefold())
        if canonical_sheet is None:
            raise ValueError(f"model-scope region names an unknown worksheet: {sheet}")
        try:
            bounds = range_boundaries(local.replace("$", ""))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"model-scope region has an unsupported source_ref: {region['source_ref']}") from exc
        if None in bounds or _rect_area(bounds) > _MAX_RANGE_AREA:
            raise ValueError(f"model-scope region must be finite and at most {_MAX_RANGE_AREA} cells: {region['source_ref']}")
        scope_regions[canonical_sheet.casefold()].append(bounds)
        normalized_regions.append({"source_ref": f"{canonical_sheet}!{_rect_to_a1(bounds)}",
                                  "rationale": region["rationale"]})
    for regions in scope_regions.values():
        regions.sort()

    def is_allowed(key: str) -> bool:
        return _scope_contains_key(scope_regions, key)

    if any(not is_allowed(key) for key in cut_coordinates):
        raise ValueError("model-scope regions must contain every confirmed formula-derived input cut coordinate")
    for target in targets:
        for address in target["start_cells"]:
            if not is_allowed(_address_key(*_split_address(address))):
                raise ValueError(f"model-scope regions omit target start {address}")
    target_start_keys = {
        _address_key(*_split_address(address))
        for target in targets for address in target["start_cells"]
    }
    if not set(formulas):
        raise ValueError("canonical inventory has no formula coordinates for model-scope validation")

    raw_starts = value.get("proposed_formula_starts", [])
    if not isinstance(raw_starts, list):
        raise ValueError("proposed_formula_starts must be a list")
    proposed_formula_starts: list[dict[str, Any]] = []
    proposed_formula_start_keys: list[str] = []
    seen_starts: set[str] = set()
    for start in raw_starts:
        if (not isinstance(start, dict) or set(start) != {"source_ref", "rationale"}
                or not isinstance(start.get("source_ref"), str) or not start["source_ref"].strip()
                or not isinstance(start.get("rationale"), str) or not start["rationale"].strip()):
            raise ValueError("each proposed formula start needs only a finite source_ref and non-empty rationale")
        sheet, local = _split_address(start["source_ref"])
        canonical_sheet = sheet_names.get(sheet.casefold())
        if canonical_sheet is None:
            raise ValueError(f"proposed formula start names an unknown worksheet: {sheet}")
        try:
            bounds = range_boundaries(local.replace("$", ""))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"proposed formula start has an unsupported source_ref: {start['source_ref']}") from exc
        if None in bounds or _rect_area(bounds) > _MAX_RANGE_AREA:
            raise ValueError(f"proposed formula start must be finite and at most {_MAX_RANGE_AREA} cells: {start['source_ref']}")
        members = []
        for row in range(bounds[1], bounds[3] + 1):
            for column in range(bounds[0], bounds[2] + 1):
                address = f"{canonical_sheet}!{get_column_letter(column)}{row}"
                key = _address_key(canonical_sheet, address)
                formula_record = formulas.get(key)
                if formula_record is None:
                    raise ValueError(f"proposed formula start includes a non-formula source coordinate: {address}")
                if not is_allowed(key):
                    raise ValueError(f"proposed formula start is outside allowed regions: {address}")
                if key in cut_coordinates or key in exclusion_coordinates:
                    raise ValueError(f"proposed formula start cannot cross an external cut or excluded region: {address}")
                if key in target_start_keys:
                    raise ValueError(f"proposed formula starts must be distinct from confirmed target starts: {address}")
                if key in seen_starts:
                    raise ValueError(f"proposed formula start is repeated: {address}")
                seen_starts.add(key)
                proposed_formula_start_keys.append(key)
                formula = formula_record[2]["formula"]
                members.append({"address": f"{canonical_sheet}!{formula_record[1]}",
                                "formula": formula, "formula_sha256": _formula_sha256(formula)})
        proposed_formula_starts.append({"source_ref": start["source_ref"], "rationale": start["rationale"],
                                        "formula_members": members})

    model_scope = {
        "path": str(scope_path.resolve()), "sha256": workflow._hash_file(scope_path),
        "candidate_scope": "scope_bounded_static_candidate_design",
        "full_candidate_trace": {"path": str(full_path), "sha256": full_ref["sha256"]},
        "allowed_regions": normalized_regions,
        "proposed_formula_starts": proposed_formula_starts,
        "proposed_formula_start_keys": proposed_formula_start_keys,
    }
    return model_scope, dict(scope_regions)


def _public_model_scope(model_scope: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in model_scope.items() if key != "proposed_formula_start_keys"}


def _rect_area(bounds: tuple[int, int, int, int]) -> int:
    return (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1)


def _rect_to_a1(bounds: tuple[int, int, int, int]) -> str:
    first = f"{get_column_letter(bounds[0])}{bounds[1]}"
    last = f"{get_column_letter(bounds[2])}{bounds[3]}"
    return first if first == last else f"{first}:{last}"


def _intersect_rect(
    first: tuple[int, int, int, int], second: tuple[int, int, int, int],
) -> tuple[int, int, int, int] | None:
    result = (max(first[0], second[0]), max(first[1], second[1]),
              min(first[2], second[2]), min(first[3], second[3]))
    return result if result[0] <= result[2] and result[1] <= result[3] else None


def _subtract_rect(
    whole: tuple[int, int, int, int], cut: tuple[int, int, int, int],
) -> list[tuple[int, int, int, int]]:
    intersection = _intersect_rect(whole, cut)
    if intersection is None:
        return [whole]
    min_col, min_row, max_col, max_row = whole
    cut_min_col, cut_min_row, cut_max_col, cut_max_row = intersection
    pieces = []
    if min_row < cut_min_row:
        pieces.append((min_col, min_row, max_col, cut_min_row - 1))
    if cut_max_row < max_row:
        pieces.append((min_col, cut_max_row + 1, max_col, max_row))
    if min_col < cut_min_col:
        pieces.append((min_col, cut_min_row, cut_min_col - 1, cut_max_row))
    if cut_max_col < max_col:
        pieces.append((cut_max_col + 1, cut_min_row, max_col, cut_max_row))
    return pieces


def _clip_to_scope(
    reference: tuple[int, int, int, int], allowed: list[tuple[int, int, int, int]],
) -> tuple[list[tuple[int, int, int, int]], list[tuple[int, int, int, int]]]:
    included: list[tuple[int, int, int, int]] = []
    excluded = [reference]
    for region in allowed:
        remaining = []
        for piece in excluded:
            intersection = _intersect_rect(piece, region)
            if intersection is None:
                remaining.append(piece)
            else:
                included.append(intersection)
                remaining.extend(_subtract_rect(piece, intersection))
        excluded = remaining
        if not excluded:
            break
    return included, excluded


def _scope_contains_key(
    regions: dict[str, list[tuple[int, int, int, int]]], address_key: str,
) -> bool:
    try:
        sheet, local = address_key.rsplit("!", 1)
        bounds = range_boundaries(local.replace("$", ""))
    except (ValueError, TypeError):
        return False
    if None in bounds or bounds[0] != bounds[2] or bounds[1] != bounds[3]:
        return False
    return any(_intersect_rect(bounds, region) is not None for region in regions.get(sheet.casefold(), []))


def _record_scope_frontier(
    frontiers: dict[tuple[str, str, str], dict[str, Any]], unknown: Any,
    formulas: dict[str, tuple[str, str, dict[str, Any]]], referring_formula: str,
    reference_text: str, excluded_source_ref: str, extent_kind: str,
) -> None:
    formula_key = _address_key(*_split_address(referring_formula))
    formula_record = formulas.get(formula_key)
    if formula_record is None:
        raise ValueError(f"scope frontier has no canonical referring formula: {referring_formula}")
    sheet, local = _split_address(excluded_source_ref)
    bounds = range_boundaries(local.replace("$", ""))
    if None in bounds:
        raise ValueError(f"scope frontier must be a finite excluded extent: {excluded_source_ref}")
    canonical_extent = f"{sheet}!{_rect_to_a1(bounds)}"
    key = (formula_key, reference_text, canonical_extent)
    frontiers.setdefault(key, {
        "referring_formula": f"{formula_record[0]}!{formula_record[1]}",
        "referring_formula_sha256": _formula_sha256(formula_record[2]["formula"]),
        "reference": reference_text,
        "excluded_source_ref": canonical_extent,
        "coordinate_count": _rect_area(bounds),
        "extent_kind": extent_kind,
    })
    unknown(f"{formula_record[0]}!{formula_record[1]}", "model_scope_frontier",
            f"{reference_text} leaves the scope at {canonical_extent}")


def _validate_candidate_trace_binding(
    analysis_dir: Path, manifest: dict[str, Any], binding: dict[str, Any],
    trace: dict[str, Any], trace_path: Path, inventory: dict[str, Any],
) -> None:
    from excel_to_act.steps.conversion_workflow import load_workflow
    from excel_to_act.steps.step3.input_boundary import (
        _require_bound_workflow, _validate_catalog_identity, _validate_targets,
        confirmed_boundary_reference,
    )

    if trace.get("selection_basis") != "static_candidate":
        raise ValueError("source-candidate trace has no static-candidate selection basis")
    for key in ("formula_calculation_performed", "formula_cache_used", "formula_cache_inputs",
                "native_excel_called", "trace_computed_values_read"):
        if trace.get(key) is not False:
            raise ValueError(f"source-candidate trace violates source-only policy: {key}")
    if trace.get("binding_sha256") != manifest.get("binding_sha256"):
        raise ValueError("source-candidate trace belongs to another analysis binding")
    if trace.get("source_sha256") != binding["source"]["source_sha256"]:
        raise ValueError("source-candidate trace belongs to another workbook")
    current_analysis, current_manifest, current_binding, _step1_root, inventory_path = workflow._load_context(analysis_dir)
    if trace.get("analysis", {}).get("fields_sha256") != current_manifest.get("stage_hashes", {}).get("fields"):
        raise ValueError("source-candidate trace fields are stale")
    if trace.get("analysis", {}).get("dependencies_sha256") != current_manifest.get("stage_hashes", {}).get("dependencies"):
        raise ValueError("source-candidate trace dependencies are stale")
    if trace.get("analysis", {}).get("structural_plan_sha256") != current_manifest.get("stage_hashes", {}).get("plan"):
        raise ValueError("source-candidate trace structural plan is stale")
    if (trace.get("analysis", {}).get("inventory_sha256") != current_binding["artifacts"]["inventory_sha256"]
            or Path(trace.get("analysis", {}).get("inventory_path", "")).resolve() != inventory_path):
        raise ValueError("source-candidate trace canonical inventory is stale")
    if Path(trace.get("analysis", {}).get("path", "")).resolve() != current_analysis:
        raise ValueError("source-candidate trace belongs to a different continuation analysis directory")
    workflow_value, workflow_manifest = load_workflow(Path(current_binding["workflow_path"]))
    _require_bound_workflow(current_manifest, current_binding, workflow_value, workflow_manifest)
    current_ref, boundary_path, _boundary_md = confirmed_boundary_reference(workflow_value, workflow_manifest)
    boundary = workflow._read_json(boundary_path)
    target_selection, targets_sha, catalog_sha, _targets_path, _catalog_path = _read_confirmed_selection(
        boundary, boundary_path, current_binding, current_manifest,
        workflow._load_inventory(inventory_path), current_ref, _validate_targets, _validate_catalog_identity,
    )
    declared_ref = trace.get("input_boundary", {})
    for key in ("revision", "boundary_sha256", "artifact", "review_receipts"):
        if declared_ref.get(key) != current_ref.get(key):
            raise ValueError("source-candidate trace confirmed input-boundary reference is stale")
    if declared_ref.get("targets_input_sha256") != targets_sha or declared_ref.get("catalog_input_sha256") != catalog_sha:
        raise ValueError("source-candidate trace target/catalog input hashes are stale")
    if trace.get("scenario") != target_selection["scenario"] or trace.get("targets") != target_selection["targets"]:
        raise ValueError("source-candidate trace target or scenario metadata changed")
    if inventory.get("workbook_sha256") != binding["source"]["source_sha256"]:
        raise ValueError("source-candidate validation inventory is not bound to the workbook")

    source_cells: dict[str, tuple[str, dict[str, Any]]] = {}
    source_formulas: dict[str, tuple[str, str, dict[str, Any]]] = {}
    for sheet_record in inventory.get("sheets", []):
        sheet = str(sheet_record["name"])
        for cell in sheet_record.get("cells", []):
            local = _local_address(sheet, str(cell.get("address", "")))
            key = _address_key(sheet, local)
            source_cells[key] = (sheet, cell)
            formula = cell.get("formula")
            if isinstance(formula, str) and formula.startswith("="):
                source_formulas[key] = (sheet, local.replace("$", "").upper(), cell)

    trace_formula_keys: set[str] = set()
    for record in trace.get("cells", []):
        if not isinstance(record, dict) or record.get("role") != "calculated_formula":
            raise ValueError("source-candidate trace contains a non-formula cell record")
        key = _address_key(*_split_address(record.get("address", "")))
        formula_record = source_formulas.get(key)
        if (formula_record is None or record.get("formula") != formula_record[2].get("formula")
                or record.get("formula_sha256") != _formula_sha256(formula_record[2]["formula"])):
            raise ValueError(f"source-candidate trace formula differs from current source at {key}")
        if key in trace_formula_keys:
            raise ValueError(f"source-candidate trace repeats a formula address: {key}")
        trace_formula_keys.add(key)

    expected_arrays, _followers = _inventory_arrays(inventory, source_cells)
    expected_candidate_arrays = [expected_arrays[key] for key in sorted(expected_arrays)
                                 if key in trace_formula_keys]
    if trace.get("arrays") != expected_candidate_arrays:
        raise ValueError("source-candidate array geometry differs from the canonical inventory")

    _expected_cut_coordinates, expected_cuts = _catalog_formula_cuts(boundary, source_formulas,
                                                                      binding["source"]["source_sha256"])
    expected_cut_records = [expected_cuts[key] for key in sorted(expected_cuts)]
    if trace.get("external_formula_cuts") != expected_cut_records:
        raise ValueError("source-candidate external formula cuts differ from the confirmed catalog or current source")
    _expected_exclusion_coordinates, expected_exclusions = _catalog_upstream_exclusions(boundary)
    if trace.get("upstream_exclusions") != expected_exclusions:
        raise ValueError("source-candidate upstream exclusions differ from the confirmed catalog")

    scope_ref = trace.get("model_scope")
    candidate_scope = trace.get("candidate_scope")
    if scope_ref is None:
        if (candidate_scope is not None or "scope_frontiers" in trace or "candidate_starts" in trace):
            raise ValueError("unscoped source-candidate trace contains model-scope metadata")
    else:
        if candidate_scope != "scope_bounded_static_candidate_design":
            raise ValueError("source-candidate trace has an unsupported scope label")
        if (not isinstance(scope_ref, dict) or not isinstance(scope_ref.get("path"), str)
                or not isinstance(scope_ref.get("sha256"), str)):
            raise ValueError("scope-bounded candidate trace is missing its model-scope reference")
        scope_path = Path(scope_ref["path"]).expanduser().resolve()
        if not scope_path.is_file() or workflow._hash_file(scope_path) != scope_ref["sha256"]:
            raise ValueError("source-candidate model-scope input is missing or changed")
        expected_source = {key: current_binding["source"][key]
                           for key in ("source_id", "run_id", "source_sha256")}
        checked_scope, scope_regions = _validate_model_scope(
            scope_path, current_analysis, current_manifest, current_binding,
            workflow._load_inventory(inventory_path), expected_source,
            target_selection["scenario"], target_selection["targets"],
            trace.get("analysis", {}), declared_ref, source_formulas,
            _expected_cut_coordinates, _expected_exclusion_coordinates,
        )
        if _public_model_scope(checked_scope) != scope_ref:
            raise ValueError("source-candidate trace model-scope metadata differs from the current scope input")
        expected_starts = {
            "target_starts": [
                {"target_id": target["target_id"], "source_ref": address}
                for target in target_selection["targets"] for address in target["start_cells"]
            ],
            "proposed_formula_starts": checked_scope["proposed_formula_starts"],
        }
        if trace.get("candidate_starts") != expected_starts:
            raise ValueError("scope-bounded candidate trace target/proposed starts changed")
        if not isinstance(trace.get("scope_frontiers"), list):
            raise ValueError("scope-bounded candidate trace has no explicit scope-frontier list")
        frontier_keys: set[tuple[str, str, str]] = set()
        for item in trace["scope_frontiers"]:
            if (not isinstance(item, dict)
                    or not all(isinstance(item.get(key), str) and item[key]
                               for key in ("referring_formula", "referring_formula_sha256", "reference",
                                           "excluded_source_ref", "extent_kind"))
                    or not isinstance(item.get("coordinate_count"), int)
                    or isinstance(item.get("coordinate_count"), bool)):
                raise ValueError("scope-bounded candidate trace has a malformed scope frontier")
            formula_key = _address_key(*_split_address(item["referring_formula"]))
            source_formula = source_formulas.get(formula_key)
            if (source_formula is None or source_formula[2].get("formula") is None
                    or item["referring_formula_sha256"] != _formula_sha256(source_formula[2]["formula"])):
                raise ValueError("scope frontier does not cite a current source formula")
            if item["reference"].strip("'").casefold() not in source_formula[2]["formula"].casefold():
                raise ValueError("scope frontier reference text is not present in its source formula")
            excluded_sheet, excluded_local = _split_address(item["excluded_source_ref"])
            excluded_bounds = range_boundaries(excluded_local.replace("$", ""))
            excluded_intersections, _remaining = _clip_to_scope(
                excluded_bounds, scope_regions.get(excluded_sheet.casefold(), [])) if None not in excluded_bounds else ([], [])
            if (None in excluded_bounds or _rect_area(excluded_bounds) != item["coordinate_count"]
                    or excluded_intersections or item["extent_kind"] not in {"coordinate", "region", "array_anchor"}):
                raise ValueError("scope frontier extent/count is inconsistent with the declared model scope")
            key = (formula_key, item["reference"], item["excluded_source_ref"])
            if key in frontier_keys:
                raise ValueError("source-candidate trace repeats a scope frontier")
            frontier_keys.add(key)
        summary = trace.get("summary", {})
        if (summary.get("scope_frontier_count") != len(frontier_keys)
                or summary.get("scope_frontier_coordinate_count") != sum(
                    item["coordinate_count"] for item in trace["scope_frontiers"])):
            raise ValueError("scope-frontier summary differs from its explicit records")
        required_proposed = set(checked_scope["proposed_formula_start_keys"])
        if not required_proposed <= trace_formula_keys:
            raise ValueError("scope-bounded candidate trace omits a proposed formula start")
        if any(not _scope_contains_key(scope_regions, key) for key in trace_formula_keys):
            raise ValueError("scope-bounded candidate trace includes a formula outside allowed regions")
        expected_frontier_unknowns = {
            (item["referring_formula"], "model_scope_frontier",
             f"{item['reference']} leaves the scope at {item['excluded_source_ref']}")
            for item in trace["scope_frontiers"]
        }
        actual_frontier_unknowns = {
            (item.get("address"), item.get("category"), item.get("evidence"))
            for item in trace.get("unknowns", [])
            if isinstance(item, dict) and item.get("category") == "model_scope_frontier"
        }
        if actual_frontier_unknowns != expected_frontier_unknowns:
            raise ValueError("scope-frontier unknown records differ from their explicit extents")
    _ = trace_path


def _catalog_raw_coordinates(boundary: dict[str, Any]) -> dict[str, set[str]]:
    result: dict[str, set[str]] = {}
    for variable in boundary.get("variables", []):
        if not isinstance(variable, dict) or variable.get("role") != "source_raw":
            continue
        variable_id = variable.get("variable_id")
        if not isinstance(variable_id, str) or variable_id in result:
            raise ValueError("confirmed raw input catalog has a missing or duplicate variable ID")
        coordinates: set[str] = set()
        for extent in variable.get("source_extents", []):
            if not isinstance(extent, dict) or not isinstance(extent.get("sheet"), str) or not isinstance(extent.get("range"), str):
                raise ValueError(f"raw input {variable_id} has an invalid source extent")
            coordinates.update(_coordinates_for_ref(f"{extent['sheet']}!{extent['range']}", _MAX_RANGE_AREA))
        result[variable_id] = coordinates
    return result


def _catalog_formula_cuts(
    boundary: dict[str, Any], formulas: dict[str, tuple[str, str, dict[str, Any]]], source_sha: str,
) -> tuple[set[str], dict[str, dict[str, Any]]]:
    coordinate_owners: dict[str, str] = {}
    extents: dict[str, dict[str, Any]] = {}
    for variable in boundary.get("variables", []):
        if not isinstance(variable, dict) or variable.get("role") != "formula_derived_external":
            continue
        variable_id = str(variable.get("variable_id", ""))
        for extent in variable.get("source_extents", []):
            reference = f"{extent.get('sheet')}!{extent.get('range')}"
            coordinates = _coordinates_for_ref(reference, _MAX_RANGE_AREA)
            members = []
            for key in coordinates:
                formula_cell = formulas.get(key)
                if formula_cell is None:
                    raise ValueError(f"confirmed external formula cut is not a source formula at {key}")
                prior_owner = coordinate_owners.get(key)
                if prior_owner is not None and prior_owner != variable_id:
                    raise ValueError(f"confirmed formula cuts overlap at {key}")
                coordinate_owners[key] = variable_id
                sheet, address, cell = formula_cell
                formula = cell["formula"]
                members.append({"address": f"{sheet}!{address}", "formula": formula,
                                "formula_sha256": _formula_sha256(formula),
                                "source_workbook_sha256": source_sha})
            extent_key = f"{variable_id}:{reference}"
            extents[extent_key] = {"variable_id": variable_id, "source_range": reference,
                                   "coordinate_count": len(coordinates),
                                   "formula_members": sorted(members, key=lambda item: item["address"].casefold()),
                                   "traversal": "stopped_before_prerequisites"}
    return set(coordinate_owners), extents


def _catalog_upstream_exclusions(boundary: dict[str, Any]) -> tuple[set[str], list[dict[str, Any]]]:
    coordinates: set[str] = set()
    exclusions = []
    for item in boundary.get("upstream_exclusions", []):
        if not isinstance(item, dict):
            continue
        reference = item.get("source_range")
        if not isinstance(reference, str):
            continue
        try:
            members = _coordinates_for_ref(reference, _MAX_RANGE_AREA)
        except ValueError:
            members = set()
            exclusions.append({**item, "expansion": "unresolved_symbolic_or_unbounded"})
        else:
            coordinates.update(members)
            exclusions.append({**item, "coordinate_count": len(members),
                               "expansion": "coordinate_boundary_only"})
    return coordinates, exclusions


def _inventory_arrays(
    inventory: dict[str, Any], all_cells: dict[str, tuple[str, dict[str, Any]]],
) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    arrays: dict[str, dict[str, Any]] = {}
    followers: dict[str, str] = {}
    for sheet in inventory.get("sheets", []):
        name = str(sheet["name"])
        for cell in sheet.get("cells", []):
            attributes = cell.get("raw_formula_attributes")
            if not isinstance(attributes, dict) or attributes.get("t") != "array":
                continue
            formula = cell.get("formula")
            address = _local_address(name, str(cell.get("address", ""))).replace("$", "").upper()
            if not isinstance(formula, str):
                continue
            array_ref = attributes.get("ref")
            if not isinstance(array_ref, str) or not array_ref:
                continue
            try:
                bounds = range_boundaries(array_ref.replace("$", ""))
            except (TypeError, ValueError):
                continue
            if None in bounds:
                continue
            min_col, min_row, max_col, max_row = bounds
            shape = [max_row - min_row + 1, max_col - min_col + 1]
            anchor = _address_key(name, address)
            members = []
            missing = 0
            for row in range(min_row, max_row + 1):
                for column in range(min_col, max_col + 1):
                    coordinate = f"{name}!{get_column_letter(column)}{row}"
                    key = _address_key(name, coordinate)
                    if key == anchor:
                        continue
                    if key in all_cells:
                        members.append(coordinate)
                        followers[key] = anchor
                    else:
                        missing += 1
            arrays[anchor] = {"sheet": name, "anchor": f"{name}!{address}",
                              "array_ref": array_ref.replace("$", "").upper(), "formula": formula,
                              "formula_sha256": _formula_sha256(formula), "shape": shape,
                              "follower_addresses": sorted(members, key=str.casefold),
                              "follower_count": len(members), "missing_extent_coordinates": missing}
    return arrays, followers


def _coordinates_for_ref(reference: str, max_area: int) -> set[str]:
    sheet, local = _split_address(reference)
    try:
        bounds = range_boundaries(local.replace("$", ""))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"unsupported source range {reference}") from exc
    if None in bounds:
        raise ValueError(f"unbounded source range {reference}")
    min_col, min_row, max_col, max_row = bounds
    area = (max_col - min_col + 1) * (max_row - min_row + 1)
    if area > max_area:
        raise ValueError(f"source range exceeds the finite expansion limit: {reference} ({area} cells)")
    return {f"{sheet.casefold()}!{get_column_letter(column)}{row}"
            for row in range(min_row, max_row + 1)
            for column in range(min_col, max_col + 1)}


def _classify_raw_candidates(
    boundary: dict[str, Any], raw_coordinates: dict[str, set[str]], reached: set[str],
    unresolved_scope: bool,
) -> dict[str, Any]:
    objects = []
    for variable in boundary.get("variables", []):
        if not isinstance(variable, dict) or variable.get("role") != "source_raw":
            continue
        variable_id = str(variable["variable_id"])
        coordinates = raw_coordinates.get(variable_id, set())
        reached_coordinates = sorted(coordinates & reached)
        if reached_coordinates:
            status = "reached"
        elif unresolved_scope:
            status = "unresolved_possibly_reached"
        else:
            status = "not_reached_in_known_graph"
        objects.append({"variable_id": variable_id, "logical_name": variable.get("logical_name"),
                        "group_id": variable.get("group_id"), "kind": variable.get("kind"),
                        "coordinate_count": len(coordinates), "reached_coordinate_count": len(reached_coordinates),
                        "reached_examples": reached_coordinates[:5], "status": status})
    counts = Counter(item["status"] for item in objects)
    coordinate_union = set().union(*raw_coordinates.values()) if raw_coordinates else set()
    return {"object_count": len(objects), "coordinate_count": len(coordinate_union),
            "status_counts": {key: counts.get(key, 0) for key in
                              ("reached", "not_reached_in_known_graph", "unresolved_possibly_reached")},
            "not_reached_is_known_only": not unresolved_scope, "objects": objects}


def _record_formula_semantic_unknowns(formula: str, address: str, record: Any) -> None:
    branches = {"IF", "IFS", "CHOOSE", "SWITCH"}
    lookups = {"VLOOKUP", "HLOOKUP", "XLOOKUP", "LOOKUP", "INDEX", "MATCH"}
    dynamics = {"INDIRECT", "OFFSET"}
    try:
        tokens = Tokenizer(formula).items
    except (TokenizerError, IndexError, ValueError) as exc:
        record(address, "formula_tokenization_unknown", type(exc).__name__)
        return
    for token in tokens:
        if token.type != "FUNC" or token.subtype != "OPEN":
            continue
        function = token.value[:-1].rsplit(".", 1)[-1].upper()
        if function in branches:
            record(address, "conditional_branch_selection", function)
        if function in lookups:
            record(address, "lookup_selection_unknown", function)
        if function in dynamics:
            record(address, "dynamic_reference", function)


def _unbounded_a1_range_operands(formula: str) -> list[str]:
    """Return tokenizer range operands that are whole-row or whole-column A1 refs."""
    try:
        tokens = Tokenizer(formula).items
    except (TokenizerError, IndexError, ValueError):
        return []
    found: list[str] = []
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        reference = token.value.rsplit("!", 1)[-1].replace("$", "")
        if re.fullmatch(r"[A-Z]{1,3}:[A-Z]{1,3}|[1-9][0-9]*:[1-9][0-9]*", reference, re.IGNORECASE):
            found.append(token.value)
    return found


def _source_reference_text(formula: str, graph_label: str, referring_sheet: str) -> str:
    """Recover the literal A1 operand; graph labels may qualify local references."""
    try:
        target_sheet, target_local = (_split_address(graph_label) if "!" in graph_label
                                      else (referring_sheet, graph_label))
        target_key = (target_sheet.casefold(), target_local.replace("$", "").casefold())
        for token in Tokenizer(formula).items:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            raw = token.value
            try:
                token_sheet, token_local = (_split_address(raw) if "!" in raw
                                            else (referring_sheet, raw))
            except ValueError:
                continue
            if (token_sheet.casefold(), token_local.replace("$", "").casefold()) == target_key:
                return raw
    except (TokenizerError, IndexError, ValueError, TypeError):
        pass
    return graph_label


def _unknown_category(description: str) -> str:
    text = description.casefold()
    if "dynamic reference" in text:
        return "dynamic_reference"
    if "#ref!" in text:
        return "broken_reference"
    if "could not resolve" in text or "unresolved" in text:
        return "unsupported_reference"
    return "formula_reference_unknown"


def _formula_sha256(formula: str) -> str:
    return hashlib.sha256(formula.encode("utf-8")).hexdigest()


def _source_candidate_markdown(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    raw = payload["raw_candidate_reachability"]
    lines = ["# Step 3 Source Candidate Trace", "", "Status: **Static source candidates; closure unknown**", "",
             f"Source SHA-256: `{payload['source_sha256']}`",
             f"Analysis binding SHA-256: `{payload['binding_sha256']}`",
             f"Scenario: `{payload['scenario']['scenario_id']}` · `{payload['scenario_sha256']}`",
             f"Input-boundary revision: `{payload['input_boundary']['revision']}` · `{payload['input_boundary']['boundary_sha256']}`", "",
             "No Step 4 history is required or relabeled. The trace reads canonical formula text and structural references only; it does not read formula caches, evaluate formulas, call Excel, or emit numeric values.", "",
             "## Candidate selection", "",
             f"Targets: {len(payload['targets'])}; ordinary formula candidates: {summary['candidate_formula_member_count']}; array instances: {summary['candidate_array_instance_count']}; stored followers: {summary['candidate_array_follower_count']}.",
             f"Declared external formula cuts: {summary['declared_cut_formula_coordinate_count']} coordinates; reached cut coordinates: {summary['reached_cut_formula_coordinate_count']}.",
             f"Unknown paths: {summary['unresolved_path_count']}; truncated: {summary['truncated_path_count']}; candidate closure: **unknown**.", "",
             "## Raw catalog candidates", "",
             f"Raw input objects: {raw['object_count']}; source coordinates: {raw['coordinate_count']}. A not-reached label applies only to the known reference graph.", "",
             "| Classification | Objects |", "| --- | ---: |"]
    lines.extend(f"| {label.replace('_', ' ')} | {count} |" for label, count in raw["status_counts"].items())
    lines.extend(["", "## Unresolved or unapproved boundaries", ""])
    lines.extend(f"- `{item['address']}` · {item['category']} · {item['evidence']}"
                 for item in payload["unknowns"][:80])
    if len(payload["unknowns"]) > 80:
        lines.append(f"- {len(payload['unknowns']) - 80} additional explicit records are in the JSON artifact.")
    if payload.get("model_scope"):
        model_scope = payload["model_scope"]
        lines.extend(["", "## Agent-proposed model scope", "",
                      f"Scope: `{payload['candidate_scope']}` · `{model_scope['path']}` · SHA-256 `{model_scope['sha256']}`.",
                      f"Allowed regions: {len(model_scope['allowed_regions'])}; proposed formula-start members: {summary['proposed_formula_start_count']}; explicit scope frontiers: {summary['scope_frontier_count']} ({summary['scope_frontier_coordinate_count']} coordinates).",
                      "The target starts and proposed formula starts are separate. Proposed starts are candidate roots only; scope frontiers remain untraced and do not establish that omitted formulas are unused.", ""])
        lines.extend(f"- `{item['source_ref']}` — {item['rationale']} ({len(item['formula_members'])} source formulas)"
                     for item in model_scope["proposed_formula_starts"])
        for item in payload.get("scope_frontiers", [])[:40]:
            lines.append(f"- Frontier from `{item['referring_formula']}` (`{item['referring_formula_sha256'][:12]}`): `{item['reference']}` omits `{item['excluded_source_ref']}` ({item['coordinate_count']} coordinates).")
        if len(payload.get("scope_frontiers", [])) > 40:
            lines.append(f"- {len(payload['scope_frontiers']) - 40} additional scope frontiers are in the JSON artifact.")
    lines.extend(["", "The JSON artifact preserves exact formula text and SHA-256 for candidate formula cells and formula-derived cut coordinates. Conditional arms are candidate unions; lookup-selected rows remain unknown. External formula cuts stop traversal before their prerequisites.", ""])
    return "\n".join(lines)


def _address_key(sheet: str, address: str) -> str:
    local = _local_address(sheet, address)
    return f"{sheet.casefold()}!{local.replace('$', '').upper()}"


def _split_address(address: str) -> tuple[str, str]:
    sheet, local = address.rsplit("!", 1)
    return sheet.strip("'").replace("''", "'"), local


def _local_address(sheet: str, address: str) -> str:
    if "!" not in address:
        return address
    qualified_sheet, local = address.rsplit("!", 1)
    qualified_sheet = qualified_sheet.strip("'").replace("''", "'")
    if qualified_sheet.casefold() != sheet.casefold():
        raise ValueError(f"trace address sheet does not match its record: {address}")
    return local


def _profile_markdown(payload: dict[str, Any]) -> str:
    summary = payload["summary"]
    candidate = payload.get("selection_basis") == "static_candidate"
    trace_key = "source_candidate_trace" if candidate else "historical_selection_trace"
    title = "Source Candidate Formula-Family Profile" if candidate else "Source Formula-Family Profile"
    selection = "Static candidate formulas" if candidate else "Selected ordinary formulas"
    closure = (f" Unknown paths: {summary['unresolved_path_count']}; candidate closure remains unknown."
               if candidate else "")
    lines = [f"# {title}", "", "Status: **Static syntactic profile**", "",
             f"Source SHA-256: `{payload['source_sha256']}`",
             f"Step 3 binding SHA-256: `{payload['binding_sha256']}`",
             f"Trace: `{payload[trace_key]['path']}` · SHA-256 `{payload[trace_key]['sha256']}`", "",
             f"{selection}: {summary['ordinary_formula_members']}; syntactic families: {summary['syntactic_families']}.{closure}", "",
             ("The source-only trace contributes static branch-union candidates; it does not prove active selection. "
              if candidate else "The supplied trace selects source addresses. ") +
             "Formula text comes from the checked canonical inventory; trace values and formula caches are not read.",
             "These are syntactic candidates, not business equations or a new active closure. Array followers require separate shape mapping.", "",
             "| Sheet | Ordinary formulas | Syntactic families |", "| --- | ---: | ---: |"]
    lines.extend(f"| {item['sheet']} | {item['ordinary_formula_members']} | {item['syntactic_families']} |"
                 for item in payload["sheets"])
    if payload.get("model_scope"):
        scope_ref = payload["model_scope"]
        lines.extend(["", f"Candidate scope: `{payload['candidate_scope']}`; model-scope SHA-256 `{scope_ref['sha256']}`; scope-frontier records: {summary['scope_frontier_count']}; proposed formula-start members: {summary['proposed_formula_start_count']}. Candidate closure and active coverage remain unknown."])
    profile_name = "source_candidate_family_profile.json" if candidate else "source_family_profile.json"
    lines.extend(["", f"Full formulas, membership, and IDs are in `{profile_name}`.", ""])
    return "\n".join(lines)
