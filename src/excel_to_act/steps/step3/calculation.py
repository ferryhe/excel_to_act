"""Narrow lazy formula evaluation and native Excel comparison for the GP target."""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from datetime import date, datetime
from importlib.resources import files
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

_FORMULA_CACHE_POLICY = "Formula cells are evaluated from formula text; data_only formula caches are never loaded."
_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"}
_CELL = re.compile(r"^\$?([A-Z]{1,3})\$?([1-9][0-9]*)$", re.IGNORECASE)
_SHEET_PREFIX = re.compile(r"^(?:'((?:[^']|'')+)'|([^!]+))!(.+)$")


class CalculationBlocked(ValueError):
    """An unsupported or cyclic active calculation path."""


@dataclass(frozen=True)
class ExcelError:
    code: str


@dataclass(frozen=True)
class CellRange:
    sheet: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def address(self) -> str:
        start = f"{get_column_letter(self.min_col)}{self.min_row}"
        end = f"{get_column_letter(self.max_col)}{self.max_row}"
        return f"{self.sheet}!{start}" if start == end else f"{self.sheet}!{start}:{end}"

    @property
    def single(self) -> bool:
        return self.min_row == self.max_row and self.min_col == self.max_col


def _parse_reference(text: str, current_sheet: str) -> CellRange | None:
    sheet = current_sheet
    local = text
    match = _SHEET_PREFIX.fullmatch(text)
    if match:
        sheet = (match.group(1).replace("''", "'") if match.group(1) is not None else match.group(2)).strip()
        local = match.group(3)
    elif "!" in text or "[" in text or "]" in text:
        return None
    try:
        bounds = range_boundaries(local.replace("$", ""))
    except (TypeError, ValueError):
        return None
    if any(value is None for value in bounds):
        return None
    min_col, min_row, max_col, max_row = bounds
    if min_col < 1 or min_row < 1 or max_col > 16_384 or max_row > 1_048_576:
        return None
    return CellRange(sheet, min_row, min_col, max_row, max_col)


class _Parser:
    _PRECEDENCE = {"=": 10, "<>": 10, "<": 10, ">": 10, "<=": 10, ">=": 10,
                   "&": 20, "+": 30, "-": 30, "*": 40, "/": 40, "^": 50}

    def __init__(self, formula: str, sheet: str):
        try:
            self.tokens = [token for token in Tokenizer(formula).items if token.type not in {"WSPACE", "WHITE-SPACE"}]
        except (TokenizerError, IndexError) as exc:
            raise CalculationBlocked(f"unsupported formula tokenization: {exc}") from exc
        self.position = 0
        self.sheet = sheet

    def parse(self) -> Any:
        if not self.tokens:
            raise CalculationBlocked("empty formula")
        result = self._expression(0)
        if self.position != len(self.tokens):
            token = self.tokens[self.position]
            raise CalculationBlocked(f"unsupported formula token {token.value!r}")
        return result

    def _peek(self) -> Any | None:
        return self.tokens[self.position] if self.position < len(self.tokens) else None

    def _take(self) -> Any:
        token = self._peek()
        if token is None:
            raise CalculationBlocked("unexpected end of formula")
        self.position += 1
        return token

    def _expression(self, minimum: int) -> Any:
        token = self._take()
        if token.type == "OPERAND" and token.subtype == "RANGE":
            left = token.value
            if left.startswith("#") and left in _ERRORS:
                left_node = ("literal", ExcelError(left))
            elif "[" in left or "]" in left:
                left_node = ("unsupported_reference", left)
            else:
                reference = _parse_reference(left, self.sheet)
                if reference is None and any(character in left for character in "[]!:"):
                    left_node = ("unsupported_reference", left)
                else:
                    left_node = ("reference", reference) if reference else ("name", left)
        elif token.type == "OPERAND" and token.subtype == "NUMBER":
            left_node = ("literal", float(token.value))
        elif token.type == "OPERAND" and token.subtype == "TEXT":
            left_node = ("literal", token.value[1:-1].replace('""', '"'))
        elif token.type == "OPERAND" and token.subtype == "LOGICAL":
            left_node = ("literal", token.value.upper() == "TRUE")
        elif token.type == "OPERAND" and token.subtype == "ERROR":
            left_node = ("literal", ExcelError(token.value.upper()))
        elif token.type == "OPERATOR-PREFIX" and token.value in {"+", "-"}:
            left_node = ("unary", token.value, self._expression(60))
        elif token.type == "PAREN" and token.subtype == "OPEN":
            left_node = self._expression(0)
            close = self._take()
            if close.type != "PAREN" or close.subtype != "CLOSE":
                raise CalculationBlocked("unbalanced formula parentheses")
        elif token.type == "FUNC" and token.subtype == "OPEN":
            raw_name = token.value[:-1]
            range_left, separator, function_name = raw_name.partition(":")
            args: list[Any] = []
            if self._peek() and self._peek().type != "FUNC" or self._peek() and self._peek().subtype != "CLOSE":
                while True:
                    next_token = self._peek()
                    if next_token is None:
                        raise CalculationBlocked(f"unclosed function {raw_name}")
                    if next_token.type == "SEP" and next_token.subtype == "ARG":
                        args.append(("literal", None))
                        self._take()
                        continue
                    if next_token.type == "FUNC" and next_token.subtype == "CLOSE":
                        break
                    args.append(self._expression(0))
                    next_token = self._peek()
                    if next_token and next_token.type == "SEP" and next_token.subtype == "ARG":
                        self._take()
                        continue
                    if next_token and next_token.type == "FUNC" and next_token.subtype == "CLOSE":
                        break
                    raise CalculationBlocked(f"unsupported argument separator in {raw_name}")
            close = self._take()
            if close.type != "FUNC" or close.subtype != "CLOSE":
                raise CalculationBlocked(f"unclosed function {raw_name}")
            name = function_name if separator else raw_name
            left_node = ("call", name.upper(), args)
            if separator:
                reference = _parse_reference(range_left, self.sheet)
                if reference is None or not reference.single:
                    raise CalculationBlocked(f"unsupported range endpoint {range_left!r}")
                left_node = ("range", ("reference", reference), left_node)
        else:
            raise CalculationBlocked(f"unsupported formula token {token.value!r}")

        while self._peek() is not None:
            token = self._peek()
            if token.type == "OPERATOR-POSTFIX" and token.value == "%":
                self._take()
                left_node = ("unary", "%", left_node)
                continue
            if token.type != "OPERATOR-INFIX" or token.value not in self._PRECEDENCE:
                break
            precedence = self._PRECEDENCE[token.value]
            if precedence < minimum:
                break
            self._take()
            right_minimum = precedence if token.value == "^" else precedence + 1
            right_node = self._expression(right_minimum)
            left_node = ("binary", token.value, left_node, right_node)
        return left_node


class FormulaEvaluator:
    """Evaluate only cells reached from a requested formula, without loading caches."""

    def __init__(self, workbook_path: Path,
                 external_overrides: dict[str, dict[str, Any]] | None = None):
        self.workbook_path = workbook_path.resolve()
        self.workbook = load_workbook(self.workbook_path, data_only=False, read_only=False, keep_vba=False)
        self.sheet_names = {sheet.casefold(): sheet for sheet in self.workbook.sheetnames}
        self.cache: dict[tuple[str, int, int], Any] = {}
        self.stack: list[tuple[str, int, int]] = []
        self.cells: dict[str, dict[str, Any]] = {}
        self.edges: dict[tuple[str, str], dict[str, Any]] = {}
        self.branches: list[dict[str, Any]] = []
        self.lookups: list[dict[str, Any]] = []
        self.names: dict[str, dict[str, Any]] = {}
        self.dynamic_ranges: list[dict[str, Any]] = []
        self.name_stack: list[str] = []
        self.array_cells: dict[tuple[str, int, int], dict[str, Any]] = {}
        self.array_results: dict[tuple[str, int, int], list[list[Any]]] = {}
        self.external_overrides: dict[tuple[str, int, int], dict[str, Any]] = {}
        for address, record in (external_overrides or {}).items():
            reference = _parse_reference(address, "")
            if (reference is None or not reference.single or not reference.sheet
                    or not isinstance(record, dict)):
                raise CalculationBlocked(f"external override must bind a qualified source cell: {address!r}")
            key = self._key(reference.sheet, reference.min_row, reference.min_col)
            if key in self.external_overrides:
                raise CalculationBlocked(f"duplicate external override address: {address}")
            required = ("boundary_variable_id", "source_formula", "source_formula_sha256",
                        "capture_artifact_sha256")
            if (not all(isinstance(record.get(field), str) and record[field] for field in required)
                    or "value" not in record):
                raise CalculationBlocked(f"external override provenance is incomplete: {address}")
            self.external_overrides[key] = dict(record)
        for worksheet in self.workbook.worksheets:
            for cell in worksheet._cells.values():
                value = cell.value
                if not hasattr(value, "text") or not hasattr(value, "ref"):
                    continue
                bounds = range_boundaries(value.ref)
                min_col, min_row, max_col, max_row = bounds
                anchor = self._key(worksheet.title, cell.row, cell.column)
                descriptor = {"anchor": anchor, "text": value.text, "ref": value.ref,
                              "min_row": min_row, "min_col": min_col,
                              "max_row": max_row, "max_col": max_col}
                for row in range(min_row, max_row + 1):
                    for col in range(min_col, max_col + 1):
                        self.array_cells[self._key(worksheet.title, row, col)] = descriptor

    def close(self) -> None:
        self.workbook.close()

    @staticmethod
    def _key(sheet: str, row: int, col: int) -> tuple[str, int, int]:
        return sheet.casefold(), row, col

    @staticmethod
    def _address(sheet: str, row: int, col: int) -> str:
        return f"{sheet}!{get_column_letter(col)}{row}"

    def _record_edge(self, prerequisite: str, relationship: str = "cell_reference") -> None:
        if not self.stack:
            return
        active = self.stack[-1]
        source = self._address(self.sheet_names.get(active[0], active[0]), active[1], active[2])
        prerequisite_match = _SHEET_PREFIX.fullmatch(prerequisite)
        if prerequisite_match:
            sheet = (prerequisite_match.group(1).replace("''", "'")
                     if prerequisite_match.group(1) is not None else prerequisite_match.group(2)).strip()
            canonical_sheet = self.sheet_names.get(sheet.casefold(), sheet)
            prerequisite = f"{canonical_sheet}!{prerequisite_match.group(3)}"
        key = (source, prerequisite)
        edge = self.edges.setdefault(key, {"consumer": source, "prerequisite": prerequisite, "relationship": relationship})
        if relationship != "cell_reference":
            edge["relationship"] = relationship

    def evaluate_cell(self, sheet: str, row: int, col: int, *, dependency: bool = True) -> Any:
        canonical_sheet = self.sheet_names.get(sheet.casefold())
        if canonical_sheet is None:
            raise CalculationBlocked(f"worksheet not found in source workbook: {sheet}")
        sheet = canonical_sheet
        key = self._key(sheet, row, col)
        address = self._address(sheet, row, col)
        if dependency:
            self._record_edge(address)
        external = self.external_overrides.get(key)
        if external is not None:
            if key in self.array_cells:
                raise CalculationBlocked(f"external boundary overlaps a source array formula: {address}")
            cell = self.workbook[sheet].cell(row=row, column=col)
            formula = cell.value if cell.data_type == "f" else None
            if not isinstance(formula, str):
                raise CalculationBlocked(f"external boundary is not a source formula cell: {address}")
            formula_hash = hashlib.sha256(formula.encode("utf-8")).hexdigest()
            if (formula != external["source_formula"]
                    or formula_hash != external["source_formula_sha256"]):
                raise CalculationBlocked(f"external boundary source formula changed: {address}")
            value = external["value"]
            if isinstance(value, dict) and set(value) == {"excel_error"}:
                value = ExcelError(str(value["excel_error"]).upper())
            elif not (value is None or isinstance(value, (str, int, float, bool))):
                raise CalculationBlocked(f"external boundary value is not a scalar: {address}")
            self.cache[key] = value
            self.cells[address] = {
                "sheet": sheet, "address": address, "role": "external_boundary_input",
                "formula": formula, "formula_sha256": formula_hash,
                "boundary_variable_id": external["boundary_variable_id"],
                "capture_artifact_sha256": external["capture_artifact_sha256"],
                "value": _json_value(value),
            }
            return value
        if key in self.cache:
            return self.cache[key]
        array = self.array_cells.get(key)
        if array is not None and key != array["anchor"]:
            anchor_key = array["anchor"]
            anchor_address = self._address(self.sheet_names[anchor_key[0]], anchor_key[1], anchor_key[2])
            member_edge = (address, anchor_address)
            self.edges.setdefault(member_edge, {"consumer": address, "prerequisite": anchor_address,
                                                "relationship": "array_formula_anchor"})
            self.evaluate_cell(sheet, anchor_key[1], anchor_key[2], dependency=False)
            matrix = self.array_results.get(anchor_key)
            if matrix is None:
                raise CalculationBlocked(f"array formula at {self._address(*anchor_key)} did not return a matrix")
            value = matrix[row - array["min_row"]][col - array["min_col"]]
            self.cache[key] = value
            self.cells[address] = {"sheet": sheet, "address": address, "role": "calculated_array_formula",
                                   "formula": array["text"], "array_ref": array["ref"], "value": _json_value(value)}
            return value
        if key in self.stack:
            cycle = [self._address(*item) for item in self.stack[self.stack.index(key):]] + [address]
            raise CalculationBlocked(f"active formula cycle: {' -> '.join(cycle)}")
        cell = self.workbook[sheet].cell(row=row, column=col)
        array = self.array_cells.get(key)
        formula = (array["text"] if array is not None else cell.value) if cell.data_type == "f" else None
        record = {"sheet": sheet, "address": address, "role": "calculated_formula" if formula is not None else "source_value",
                  "formula": formula if isinstance(formula, str) else None}
        self.cells[address] = record
        if formula is None:
            value = self._source_value(cell.value, cell.data_type)
            self.cache[key] = value
            record["value"] = _json_value(value)
            return value
        if not isinstance(formula, str):
            raise CalculationBlocked(f"unsupported formula object at {address}: {type(formula).__name__}")
        self.stack.append(key)
        try:
            result = self.evaluate_formula(formula, sheet, address)
            if array is not None and isinstance(result, list):
                self.array_results[key] = result
                result = result[row - array["min_row"]][col - array["min_col"]]
            self.cache[key] = result
            record["value"] = _json_value(result)
            return result
        finally:
            self.stack.pop()

    @staticmethod
    def _source_value(value: Any, data_type: str) -> Any:
        if data_type == "e" and isinstance(value, str):
            return ExcelError(value.upper())
        return value

    def evaluate_formula(self, formula: str, sheet: str, owner: str) -> Any:
        try:
            tree = _Parser(formula, sheet).parse()
        except CalculationBlocked as exc:
            raise CalculationBlocked(f"{owner} {formula}: {exc}") from exc
        value = self._evaluate(tree, sheet, owner)
        if isinstance(value, CellRange):
            return self._scalar(value)
        return value

    def _scalar(self, value: Any) -> Any:
        if isinstance(value, CellRange):
            if not value.single:
                raise CalculationBlocked(f"multi-cell reference used as a scalar: {value.address}")
            return self.evaluate_cell(value.sheet, value.min_row, value.min_col)
        return value

    def _range_values(self, reference: CellRange) -> list[list[Any]]:
        self._record_edge(reference.address, "range_read")
        return [
            [self.evaluate_cell(reference.sheet, row, col, dependency=False)
             for col in range(reference.min_col, reference.max_col + 1)]
            for row in range(reference.min_row, reference.max_row + 1)
        ]

    def _flatten(self, value: Any) -> list[Any]:
        if isinstance(value, CellRange):
            return [item for row in self._range_values(value) for item in row]
        if isinstance(value, list):
            return [item for row in value for item in (row if isinstance(row, list) else [row])]
        return [value]

    def _name(self, name: str, current_sheet: str) -> Any:
        qualified_local = None
        if current_sheet in self.workbook.sheetnames:
            qualified_local = self.workbook[current_sheet].defined_names.get(name)
        definition = qualified_local or self.workbook.defined_names.get(name)
        scope = current_sheet if qualified_local else "workbook"
        if definition is None:
            return ExcelError("#NAME?")
        marker = f"{scope}!{name}" if scope != "workbook" else name
        if marker.casefold() in {item.casefold() for item in self.name_stack}:
            raise CalculationBlocked(f"defined-name cycle: {' -> '.join(self.name_stack + [marker])}")
        text = definition.attr_text
        if not isinstance(text, str) or not text:
            raise CalculationBlocked(f"defined name {marker} has no supported definition")
        if "[" in text or "]" in text:
            raise CalculationBlocked(f"defined name {marker} refers to an external workbook")
        definition_text = text[1:] if text.startswith("=") else text
        reference = _parse_reference(definition_text, current_sheet)
        self.names.setdefault(marker, {"name": name, "scope": scope, "definition": text,
                                       "destination": reference.address if reference else None})
        if reference is not None:
            return reference
        self.name_stack.append(marker)
        try:
            return self.evaluate_formula("=" + definition_text, current_sheet, f"defined name {marker}")
        finally:
            self.name_stack.pop()

    def _evaluate(self, node: Any, sheet: str, owner: str) -> Any:
        kind = node[0]
        if kind == "literal":
            return node[1]
        if kind == "reference":
            return node[1]
        if kind == "name":
            return self._name(node[1], sheet)
        if kind == "unsupported_reference":
            raise CalculationBlocked(f"unsupported or external reference syntax {node[1]!r}")
        if kind == "range":
            left = self._evaluate(node[1], sheet, owner)
            right = self._evaluate(node[2], sheet, owner)
            if not isinstance(left, CellRange) or not left.single or not isinstance(right, CellRange):
                raise CalculationBlocked(f"unsupported range expression in {owner}")
            if left.sheet.casefold() != right.sheet.casefold():
                raise CalculationBlocked(f"cross-sheet range operator is unsupported in {owner}")
            result = CellRange(left.sheet, left.min_row, left.min_col, right.max_row, right.max_col)
            self.dynamic_ranges.append({"formula_cell": owner, "left": left.address, "right": right.address,
                                        "selected_range": result.address})
            return result
        if kind == "call":
            return self._call(node[1], node[2], sheet, owner)
        if kind == "unary":
            value = self._scalar(self._evaluate(node[2], sheet, owner))
            if isinstance(value, ExcelError):
                return value
            if node[1] == "%":
                return _number(value) / 100.0
            number = _number(value)
            return number if node[1] == "+" else -number
        if kind == "binary":
            left = self._evaluate(node[2], sheet, owner)
            right = self._evaluate(node[3], sheet, owner)
            return _array_binary(node[1], left, right, self)
        raise CalculationBlocked(f"unknown parsed expression {kind!r} at {owner}")

    def _arguments(self, args: list[Any], sheet: str, owner: str) -> list[Any]:
        return [self._evaluate(arg, sheet, owner) for arg in args]

    def _call(self, name: str, args: list[Any], sheet: str, owner: str) -> Any:
        if name == "IF":
            if len(args) not in {2, 3}:
                raise CalculationBlocked(f"IF requires two or three arguments at {owner}")
            condition = self._scalar(self._evaluate(args[0], sheet, owner))
            if isinstance(condition, ExcelError):
                return condition
            selected = 1 if _truthy(condition) else (2 if len(args) == 3 else None)
            self.branches.append({"formula_cell": owner, "function": "IF", "condition": _json_value(condition),
                                  "selected_argument": selected})
            return False if selected is None else self._evaluate(args[selected], sheet, owner)
        if name == "IFERROR":
            if len(args) != 2:
                raise CalculationBlocked(f"IFERROR requires two arguments at {owner}")
            value = self._evaluate(args[0], sheet, owner)
            if isinstance(value, ExcelError):
                self.branches.append({"formula_cell": owner, "function": "IFERROR", "caught_error": value.code,
                                      "selected_argument": 2})
                return self._evaluate(args[1], sheet, owner)
            self.branches.append({"formula_cell": owner, "function": "IFERROR", "caught_error": None,
                                  "selected_argument": 1})
            return value
        if name in {"IFNA", "_XLFN.IFNA"}:
            if len(args) != 2:
                raise CalculationBlocked(f"IFNA requires two arguments at {owner}")
            value = self._evaluate(args[0], sheet, owner)
            if isinstance(value, ExcelError) and value.code == "#N/A":
                self.branches.append({"formula_cell": owner, "function": "IFNA", "caught_error": value.code,
                                      "selected_argument": 2})
                return self._evaluate(args[1], sheet, owner)
            self.branches.append({"formula_cell": owner, "function": "IFNA",
                                  "caught_error": value.code if isinstance(value, ExcelError) else None,
                                  "selected_argument": 1})
            return value
        if name in {"AND", "OR"}:
            values = self._arguments(args, sheet, owner)
            if error := next((value for value in values if isinstance(value, ExcelError)), None):
                return error
            truth = [_truthy(self._scalar(value)) for value in values]
            return all(truth) if name == "AND" else any(truth)
        if name in {"SUM", "MIN", "MAX"}:
            values = self._arguments(args, sheet, owner)
            numbers: list[float] = []
            for value in values:
                for item in self._flatten(value):
                    if isinstance(item, ExcelError):
                        return item
                    if isinstance(item, bool):
                        if not isinstance(value, CellRange):
                            numbers.append(float(item))
                    elif isinstance(item, (int, float)):
                        numbers.append(float(item))
            if name == "SUM":
                return sum(numbers)
            if not numbers:
                return 0.0
            return min(numbers) if name == "MIN" else max(numbers)
        if name == "SUMPRODUCT":
            values = self._arguments(args, sheet, owner)
            arrays = [self._matrix(value) for value in values]
            if not arrays:
                return 0.0
            shape = (len(arrays[0]), len(arrays[0][0]) if arrays[0] else 0)
            if any((len(array), len(array[0]) if array else 0) != shape for array in arrays):
                return ExcelError("#VALUE!")
            total = 0.0
            for row in range(shape[0]):
                for col in range(shape[1]):
                    product = 1.0
                    for array in arrays:
                        item = array[row][col]
                        if isinstance(item, ExcelError):
                            return item
                        product *= _number(item, text_as_zero=True)
                    total += product
            return total
        if name == "TRANSPOSE":
            if len(args) != 1:
                raise CalculationBlocked(f"TRANSPOSE requires one argument at {owner}")
            matrix = self._matrix(self._evaluate(args[0], sheet, owner))
            return [list(row) for row in zip(*matrix)]
        if name == "MATCH":
            return self._match(args, sheet, owner)
        if name == "VLOOKUP":
            return self._vlookup(args, sheet, owner)
        if name == "OFFSET":
            return self._offset(args, sheet, owner)
        if name == "INDIRECT":
            return self._indirect(args, sheet, owner)
        raise CalculationBlocked(f"unsupported active function {name} at {owner}")

    def _matrix(self, value: Any) -> list[list[Any]]:
        if isinstance(value, CellRange):
            return self._range_values(value)
        if isinstance(value, list):
            return value
        return [[self._scalar(value)]]

    def _match(self, args: list[Any], sheet: str, owner: str) -> Any:
        if len(args) not in {2, 3}:
            raise CalculationBlocked(f"MATCH requires two or three arguments at {owner}")
        needle = self._scalar(self._evaluate(args[0], sheet, owner))
        lookup_range = self._evaluate(args[1], sheet, owner)
        if isinstance(needle, ExcelError):
            return needle
        if not isinstance(lookup_range, CellRange) or (lookup_range.min_row != lookup_range.max_row and lookup_range.min_col != lookup_range.max_col):
            return ExcelError("#N/A")
        mode_value = self._scalar(self._evaluate(args[2], sheet, owner)) if len(args) == 3 else 1
        if isinstance(mode_value, ExcelError):
            return mode_value
        mode = int(_number(mode_value))
        if mode not in {-1, 0, 1}:
            return ExcelError("#N/A")
        values = [item for row in self._range_values(lookup_range) for item in row]
        self.lookups.append({"formula_cell": owner, "function": "MATCH", "lookup_value": _json_value(needle),
                             "lookup_range": lookup_range.address, "match_type": mode})
        if mode == 0:
            index = next((i for i, value in enumerate(values) if _excel_compare(value, needle) == 0), None)
        elif mode == 1:
            index = next((i for i in range(len(values) - 1, -1, -1)
                          if not isinstance(values[i], ExcelError) and _excel_compare(values[i], needle) <= 0), None)
            if index is not None and any(_excel_compare(values[i], values[i + 1]) > 0 for i in range(len(values) - 1)):
                raise CalculationBlocked(f"MATCH approximate lookup range is not ascending at {owner}: {lookup_range.address}")
        else:
            index = next((i for i, value in enumerate(values) if not isinstance(value, ExcelError) and _excel_compare(value, needle) >= 0), None)
            if index is not None and any(_excel_compare(values[i], values[i + 1]) < 0 for i in range(len(values) - 1)):
                raise CalculationBlocked(f"MATCH reverse approximate lookup range is not descending at {owner}: {lookup_range.address}")
        return index + 1 if index is not None else ExcelError("#N/A")

    def _vlookup(self, args: list[Any], sheet: str, owner: str) -> Any:
        if len(args) not in {3, 4}:
            raise CalculationBlocked(f"VLOOKUP requires three or four arguments at {owner}")
        needle = self._scalar(self._evaluate(args[0], sheet, owner))
        table = self._evaluate(args[1], sheet, owner)
        column_value = self._scalar(self._evaluate(args[2], sheet, owner))
        exact = False
        if len(args) == 4:
            mode_value = self._scalar(self._evaluate(args[3], sheet, owner))
            if isinstance(mode_value, ExcelError):
                return mode_value
            exact = not _truthy(mode_value)
        if isinstance(needle, ExcelError):
            return needle
        if not isinstance(table, CellRange) or not isinstance(column_value, (int, float)):
            return ExcelError("#VALUE!")
        column = int(column_value)
        width = table.max_col - table.min_col + 1
        if column < 1:
            return ExcelError("#VALUE!")
        if column > width:
            return ExcelError("#REF!")
        keys = [self.evaluate_cell(table.sheet, row, table.min_col, dependency=False)
                for row in range(table.min_row, table.max_row + 1)]
        match_index: int | None = None
        if exact:
            match_index = next((i for i, key in enumerate(keys) if _excel_compare(key, needle) == 0), None)
        else:
            candidate_start = 0
            if (isinstance(needle, (int, float)) and keys and isinstance(keys[0], str)
                    and len(keys) > 1 and all(isinstance(key, (int, float)) for key in keys[1:])):
                # Native approximate VLOOKUP ignores a leading text header for numeric keys.
                candidate_start = 1
            if any(isinstance(key, ExcelError) for key in keys):
                return next(key for key in keys if isinstance(key, ExcelError))
            sorted_keys = keys[candidate_start:]
            if any(_excel_compare(sorted_keys[i], sorted_keys[i + 1]) > 0 for i in range(len(sorted_keys) - 1)):
                raise CalculationBlocked(f"VLOOKUP approximate lookup range is not ascending at {owner}: {table.address}")
            selected = next((i for i in range(len(sorted_keys) - 1, -1, -1)
                             if _excel_compare(sorted_keys[i], needle) <= 0), None)
            match_index = selected + candidate_start if selected is not None else None
        self._record_edge(table.address, "lookup_table")
        self.lookups.append({"formula_cell": owner, "function": "VLOOKUP", "lookup_value": _json_value(needle),
                             "table_range": table.address, "lookup_column": 1,
                             "return_column": column, "match_mode": "exact" if exact else "approximate",
                             "matched_row": table.min_row + match_index if match_index is not None else None})
        if match_index is None:
            return ExcelError("#N/A")
        # Only the selected return cell is evaluated; unselected columns may contain GP feedback.
        result = self.evaluate_cell(table.sheet, table.min_row + match_index, table.min_col + column - 1, dependency=False)
        return 0 if result is None else result

    def _offset(self, args: list[Any], sheet: str, owner: str) -> Any:
        if len(args) not in {3, 4, 5}:
            raise CalculationBlocked(f"OFFSET requires three to five arguments at {owner}")
        base = self._evaluate(args[0], sheet, owner)
        if isinstance(base, ExcelError):
            return base
        if not isinstance(base, CellRange):
            return ExcelError("#VALUE!")
        values = [self._scalar(self._evaluate(arg, sheet, owner)) for arg in args[1:]]
        if error := next((value for value in values if isinstance(value, ExcelError)), None):
            return error
        row_offset, col_offset = int(_number(values[0])), int(_number(values[1]))
        height = int(_number(values[2])) if len(values) >= 3 else base.max_row - base.min_row + 1
        width = int(_number(values[3])) if len(values) >= 4 else base.max_col - base.min_col + 1
        start_row, start_col = base.min_row + row_offset, base.min_col + col_offset
        end_row, end_col = start_row + height - 1, start_col + width - 1
        if height < 1 or width < 1 or start_row < 1 or start_col < 1 or end_row > 1_048_576 or end_col > 16_384:
            return ExcelError("#REF!")
        result = CellRange(base.sheet, start_row, start_col, end_row, end_col)
        self.dynamic_ranges.append({"formula_cell": owner, "function": "OFFSET", "base": base.address,
                                    "row_offset": row_offset, "column_offset": col_offset,
                                    "selected_range": result.address})
        return result

    def _indirect(self, args: list[Any], sheet: str, owner: str) -> Any:
        if len(args) not in {1, 2}:
            raise CalculationBlocked(f"INDIRECT requires one or two arguments at {owner}")
        text = self._scalar(self._evaluate(args[0], sheet, owner))
        if isinstance(text, ExcelError):
            return text
        a1 = self._scalar(self._evaluate(args[1], sheet, owner)) if len(args) == 2 else True
        if isinstance(a1, ExcelError):
            return a1
        if not _truthy(a1):
            raise CalculationBlocked(f"INDIRECT R1C1 syntax is unsupported at {owner}")
        if not isinstance(text, str):
            return ExcelError("#REF!")
        reference = _parse_reference(text, sheet)
        if reference is not None:
            self.dynamic_ranges.append({"formula_cell": owner, "function": "INDIRECT", "text": text,
                                        "selected_range": reference.address})
            return reference
        if "!" in text or "[" in text or "]" in text:
            raise CalculationBlocked(f"unsupported or external INDIRECT reference {text!r} at {owner}")
        value = self._name(text, sheet)
        self.dynamic_ranges.append({"formula_cell": owner, "function": "INDIRECT", "text": text,
                                    "selected_range": value.address if isinstance(value, CellRange) else None})
        return value


def _number(value: Any, *, text_as_zero: bool = False) -> float:
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        if value == "" or text_as_zero:
            return 0.0
        try:
            return float(value)
        except ValueError as exc:
            raise CalculationBlocked(f"unsupported text-to-number coercion: {value!r}") from exc
    raise CalculationBlocked(f"unsupported numeric value type: {type(value).__name__}")


def _truthy(value: Any) -> bool:
    if isinstance(value, ExcelError):
        return False
    if value is None:
        return False
    if isinstance(value, str):
        return value != ""
    return bool(value)


def _excel_compare(left: Any, right: Any) -> int:
    if isinstance(left, ExcelError) or isinstance(right, ExcelError):
        if isinstance(left, ExcelError) and isinstance(right, ExcelError) and left.code == right.code:
            return 0
        raise CalculationBlocked("Excel error encountered in lookup comparison")
    if left is None:
        left = 0
    if right is None:
        right = 0
    if isinstance(left, str) and isinstance(right, str):
        left_value, right_value = left.casefold(), right.casefold()
    elif isinstance(left, (int, float, bool)) and isinstance(right, (int, float, bool)):
        left_value, right_value = float(left), float(right)
    elif isinstance(left, str) and isinstance(right, (int, float, bool)):
        return -1
    elif isinstance(right, str) and isinstance(left, (int, float, bool)):
        return 1
    else:
        type_order = {bool: 0, int: 0, float: 0, str: 1, datetime: 2, date: 2}
        left_rank = type_order.get(type(left), 3)
        right_rank = type_order.get(type(right), 3)
        if left_rank != right_rank:
            return (left_rank > right_rank) - (left_rank < right_rank)
        left_value, right_value = str(left).casefold(), str(right).casefold()
    return (left_value > right_value) - (left_value < right_value)


def _binary(operator: str, left: Any, right: Any) -> Any:
    if isinstance(left, ExcelError):
        return left
    if isinstance(right, ExcelError):
        return right
    if operator == "&":
        return ("" if left is None else str(left)) + ("" if right is None else str(right))
    if operator in {"=", "<>", "<", ">", "<=", ">="}:
        comparison = _excel_compare(left, right)
        return {"=": comparison == 0, "<>": comparison != 0, "<": comparison < 0,
                ">": comparison > 0, "<=": comparison <= 0, ">=": comparison >= 0}[operator]
    try:
        first, second = _number(left), _number(right)
    except CalculationBlocked:
        return ExcelError("#VALUE!")
    if operator == "+":
        return first + second
    if operator == "-":
        return first - second
    if operator == "*":
        return first * second
    if operator == "/":
        return ExcelError("#DIV/0!") if second == 0 else first / second
    if operator == "^":
        try:
            result = first ** second
        except (OverflowError, ValueError, ZeroDivisionError):
            return ExcelError("#NUM!")
        return result if math.isfinite(result) else ExcelError("#NUM!")
    raise CalculationBlocked(f"unsupported operator {operator!r}")


def _matrix(value: Any, evaluator: FormulaEvaluator) -> list[list[Any]]:
    return evaluator._matrix(value)


def _array_binary(operator: str, left: Any, right: Any, evaluator: FormulaEvaluator) -> Any:
    left_multi = isinstance(left, CellRange) and not left.single or isinstance(left, list)
    right_multi = isinstance(right, CellRange) and not right.single or isinstance(right, list)
    if not left_multi and not right_multi:
        return _binary(operator, evaluator._scalar(left), evaluator._scalar(right))
    left_matrix, right_matrix = _matrix(left, evaluator), _matrix(right, evaluator)
    rows = len(left_matrix) if left_multi else len(right_matrix)
    cols = (len(left_matrix[0]) if left_matrix else 0) if left_multi else (len(right_matrix[0]) if right_matrix else 0)
    if left_multi and right_multi and (len(left_matrix), len(left_matrix[0]) if left_matrix else 0) != (len(right_matrix), len(right_matrix[0]) if right_matrix else 0):
        return ExcelError("#VALUE!")
    result: list[list[Any]] = []
    for row in range(rows):
        values: list[Any] = []
        for col in range(cols):
            first = left_matrix[row][col] if left_multi else left_matrix[0][0]
            second = right_matrix[row][col] if right_multi else right_matrix[0][0]
            values.append(_binary(operator, first, second))
        result.append(values)
    return result


def _json_value(value: Any) -> Any:
    if isinstance(value, ExcelError):
        return {"excel_error": value.code}
    if isinstance(value, CellRange):
        return {"range": value.address}
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, list):
        return [[_json_value(item) for item in row] if isinstance(row, list) else _json_value(row) for row in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _within(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def _validate_output_path(out_dir: Path, workbook_path: Path, *bound_files: Path) -> Path:
    out = out_dir.expanduser().resolve()
    source = workbook_path.expanduser().resolve()
    if _within(out, source.parent) or _within(source, out):
        raise ValueError("output directory must stay outside the source workbook directory")
    for bound in bound_files:
        resolved = bound.expanduser().resolve()
        if _within(out, resolved.parent) or _within(resolved, out):
            raise ValueError(f"output directory overlaps a bound input directory: {resolved.parent}")
    return out


def _atomic_write(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_bytes(content)
    temporary.replace(path)


def _json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True, default=str) + "\n").encode("utf-8")


def _formula_signature(formula: Any) -> tuple[tuple[str, str, str], ...] | None:
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    try:
        return tuple((token.type, token.subtype, token.value) for token in Tokenizer(formula).items
                     if token.type not in {"WSPACE", "WHITE-SPACE"})
    except (TokenizerError, IndexError):
        return None


def _matches_value(actual: Any, expected: Any, *, abs_tol: float, rel_tol: float) -> bool:
    if isinstance(actual, dict) and "excel_error" in actual:
        actual = actual["excel_error"]
    if isinstance(expected, dict) and "excel_error" in expected:
        expected = expected["excel_error"]
    if isinstance(actual, bool) or isinstance(expected, bool):
        return actual is expected
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return math.isclose(float(actual), float(expected), rel_tol=rel_tol, abs_tol=abs_tol)
    return actual == expected


def _read_primary_inputs(evaluator: FormulaEvaluator, expected: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    worksheet = evaluator.workbook["Main"]
    actual: dict[str, Any] = {}
    mismatches: list[dict[str, Any]] = []
    for address, expected_record in expected.items():
        cell = worksheet[address]
        value = _json_value(evaluator._source_value(cell.value, cell.data_type))
        actual[address] = {"value": value, "formula": cell.value if cell.data_type == "f" else None}
        if cell.data_type == "f" or not _matches_value(value, expected_record.get("value"), abs_tol=0.0, rel_tol=0.0):
            mismatches.append({"cell": f"Main!{address}", "expected": expected_record.get("value"),
                               "actual": value, "formula_cell": cell.data_type == "f"})
    return actual, mismatches


def _baseline_matrices(baseline: dict[str, Any]) -> tuple[list[list[Any]], list[list[Any]]]:
    values = baseline.get("premium_values_B9_CD115")
    formulas = baseline.get("premium_formulas_B9_CD115")
    ranges = baseline.get("ranges", {})
    record = ranges.get("Premium!B9:CD115") if isinstance(ranges, dict) else None
    if isinstance(values, list) and isinstance(formulas, list):
        errors = baseline.get("premium_errors_B9_CD115")
        if not isinstance(errors, list) and isinstance(record, dict):
            errors = record.get("excel_errors")
        if isinstance(errors, list) and len(errors) == len(values) and all(len(error_row) == len(value_row)
                                                                           for error_row, value_row in zip(errors, values)):
            values = [[{"excel_error": error} if error else value for value, error in zip(value_row, error_row)]
                      for value_row, error_row in zip(values, errors)]
        return values, formulas
    if isinstance(record, dict) and isinstance(record.get("values"), list) and isinstance(record.get("formulas"), list):
        values = record["values"]
        errors = record.get("excel_errors")
        if isinstance(errors, list) and len(errors) == len(values) and all(len(error_row) == len(value_row)
                                                                           for error_row, value_row in zip(errors, values)):
            values = [[{"excel_error": error} if error else value for value, error in zip(value_row, error_row)]
                      for value_row, error_row in zip(values, errors)]
        return values, record["formulas"]
    raise ValueError("baseline must include the finite Premium!B9:CD115 comparison range")


def _trace_markdown(trace: dict[str, Any]) -> str:
    lines = ["# Active GP dependency trace", "", f"Source SHA-256: `{trace['source_sha256']}`", "",
             f"Evaluated cells: {len(trace['cells'])}; dependency records: {len(trace['edges'])}; "
             f"lookups: {len(trace['lookups'])}; branches: {len(trace['branches'])}.", "",
             "## Selected names", ""]
    lines.extend(f"- `{item['name']}` ({item['scope']}): `{item.get('destination') or item['definition']}`"
                 for item in trace["names"][:100])
    if len(trace["names"]) > 100:
        lines.append(f"- {len(trace['names']) - 100} more names are listed in the JSON trace.")
    lines.extend(["", "## Dynamic ranges", ""])
    lines.extend(f"- `{item.get('formula_cell')}` selected `{item.get('selected_range')}` via {item.get('function', 'range expression')}"
                 for item in trace["dynamic_ranges"][:100])
    if len(trace["dynamic_ranges"]) > 100:
        lines.append(f"- {len(trace['dynamic_ranges']) - 100} more ranges are listed in the JSON trace.")
    lines.extend(["", "## Lookup results", ""])
    lines.extend(f"- `{item['formula_cell']}` {item['function']} `{item.get('table_range', item.get('lookup_range'))}`, "
                 f"return column {item.get('return_column', 'n/a')} ({item.get('return_header') or 'unlabeled'}), "
                 f"{item.get('match_mode', item.get('match_type'))}; row {item.get('matched_row', 'n/a')}"
                 for item in trace["lookups"][:100])
    if len(trace["lookups"]) > 100:
        lines.append(f"- {len(trace['lookups']) - 100} more lookups are listed in the JSON trace.")
    lines.extend(["", "## Branch outcomes", ""])
    lines.extend(f"- `{item['formula_cell']}` {item['function']}: argument {item.get('selected_argument')}; "
                 f"condition/error `{item.get('condition', item.get('caught_error'))}`"
                 for item in trace["branches"][:100])
    if len(trace["branches"]) > 100:
        lines.append(f"- {len(trace['branches']) - 100} more branch records are listed in the JSON trace.")
    lines.extend(["", "## Evaluated-cell examples", "", "| Address | Role | Value |", "|---|---|---|"])
    lines.extend(f"| `{item['address']}` | {item['role']} | `{json.dumps(item.get('value'), ensure_ascii=False, default=str)}` |"
                 for item in trace["cells"][:100])
    if len(trace["cells"]) > 100:
        lines.append(f"\n{len(trace['cells']) - 100} more cell records are listed in the JSON trace.")
    return "\n".join(lines) + "\n"


def _result_markdown(result: dict[str, Any]) -> str:
    lines = ["# GP calculation result", "", f"Status: **{result['status']}**", "",
             f"Source: `{result.get('source_path')}`", f"Source SHA-256: `{result.get('source_sha256')}`", ""]
    if result.get("status") in {"pass", "fail"} and result.get("computed"):
        lines.extend(["## Computed values", "", "| Item | Python result | Excel baseline |", "|---|---:|---:|"])
        for name, actual in result["computed"].items():
            expected = result["reconciliation"]["root_values"].get(name, {}).get("expected")
            lines.append(f"| {name} | {actual} | {expected} |")
        lines.extend(["", f"Compared {result['reconciliation']['compared_premium_formulas']} reached Premium formulas "
                      f"and {result['reconciliation']['annual_value_count']} annual values against native Excel.", "",
                      f"Tolerance: absolute {result['reconciliation']['abs_tol']}; relative {result['reconciliation']['rel_tol']}.", "",
                      "The evaluator used source literals and formula text. Stored formula caches were not loaded."])
        if result["reconciliation"]["mismatches"]:
            lines.extend(["", "## Mismatches", "", "```json",
                          json.dumps(result["reconciliation"]["mismatches"][:50], ensure_ascii=False, indent=2), "```"])
    elif result.get("reason"):
        lines.extend(["## Blocker", "", result["reason"]])
    return "\n".join(lines) + "\n"


def _lookup_headers(evaluator: FormulaEvaluator) -> None:
    for lookup in evaluator.lookups:
        table = _parse_reference(lookup.get("table_range", ""), "") if lookup.get("table_range") else None
        if table is None or "return_column" not in lookup:
            continue
        if table.sheet == "Benefit_Table":
            header_row = 3
        else:
            header_row = table.min_row
        lookup["return_header"] = evaluator.workbook[table.sheet].cell(
            header_row, table.min_col + lookup["return_column"] - 1).value


def calculate_gp(workbook_path: Path, baseline_path: Path, out_dir: Path, *, target: str = "GP",
                 abs_tol: float = 1e-10, rel_tol: float = 1e-9) -> dict[str, Any]:
    """Evaluate the active saved-scenario GP path and reconcile it to native Excel."""
    source = workbook_path.expanduser().resolve()
    baseline_file = baseline_path.expanduser().resolve()
    try:
        out = _validate_output_path(out_dir, source, baseline_file)
    except ValueError as exc:
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": str(exc)}
    if target != "GP":
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": "the current runtime supports only the exact GP target at Premium!J1"}
    if not source.is_file() or not baseline_file.is_file():
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": "source workbook or native Excel baseline does not exist"}
    source_hash = _hash_file(source)
    try:
        baseline = json.loads(baseline_file.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": f"cannot read native Excel baseline: {exc}"}
    if not isinstance(baseline, dict):
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": "native Excel baseline must be a JSON object"}
    if (baseline.get("schema_version") != "gp.excel_oracle.v1" or baseline.get("tool") != "step5.oracle"
            or baseline.get("status") != "pass" or baseline.get("engine") != "Microsoft Excel"):
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": "baseline must be a successful gp.excel_oracle.v1 step5.oracle result from Microsoft Excel"}
    if baseline.get("source_sha256") != source_hash or baseline.get("source_copy_sha256") != source_hash:
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": "native Excel baseline source hash and source-copy hash must both match the workbook"}
    if baseline.get("macros_executed") is not False or baseline.get("input_overrides") != []:
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": "baseline must use a macro-disabled full rebuild with no input overrides"}
    expected_inputs = baseline.get("primary_inputs")
    required_inputs = {"C3", "C4", "C6", "C7", "C8", "J48", "O26"}
    if (not isinstance(expected_inputs, dict) or set(expected_inputs) != required_inputs
            or any(not isinstance(record, dict) or "value" not in record for record in expected_inputs.values())):
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": "native Excel baseline must include values for all seven saved primary inputs"}
    try:
        expected_values, expected_formulas = _baseline_matrices(baseline)
        if len(expected_values) != 107 or any(len(row) != 81 for row in expected_values) or len(expected_formulas) != 107 or any(len(row) != 81 for row in expected_formulas):
            raise ValueError("baseline Premium!B9:CD115 must be a 107 by 81 finite matrix")
    except ValueError as exc:
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": str(exc)}

    evaluator: FormulaEvaluator | None = None
    input_mismatches: list[dict[str, Any]] = []
    source_inputs: dict[str, Any] = {}
    trace: dict[str, Any] = {}
    try:
        evaluator = FormulaEvaluator(source)
        source_inputs, input_mismatches = _read_primary_inputs(evaluator, expected_inputs)
        if input_mismatches:
            raise ValueError(f"saved scenario inputs differ from native baseline: {input_mismatches}")
        computed_values = {
            "AnnuityDue": evaluator.evaluate_cell("Premium", 1, 8, dependency=False),
            "PVLoading": evaluator.evaluate_cell("Premium", 2, 8, dependency=False),
            "PVFB": evaluator.evaluate_cell("Premium", 3, 8, dependency=False),
            "GP": evaluator.evaluate_cell("Premium", 1, 10, dependency=False),
        }
        _lookup_headers(evaluator)
        mismatches: list[dict[str, Any]] = []
        root_values: dict[str, Any] = {}
        root_cells = {"AnnuityDue": "H1", "PVLoading": "H2", "PVFB": "H3", "GP": "J1"}
        after = baseline.get("after_full_rebuild", {})
        for name, actual in computed_values.items():
            cell_address = root_cells[name]
            expected_record = after.get(cell_address) if isinstance(after, dict) else None
            expected = expected_record.get("value") if isinstance(expected_record, dict) else None
            if isinstance(expected_record, dict) and expected_record.get("excel_error"):
                expected = {"excel_error": expected_record["excel_error"]}
            formula_record = evaluator.cells.get(f"Premium!{cell_address}", {})
            expected_formula = expected_record.get("formula") if isinstance(expected_record, dict) else None
            formula_ok = expected_formula is None or _formula_signature(formula_record.get("formula")) == _formula_signature(expected_formula)
            value_ok = _matches_value(_json_value(actual), expected, abs_tol=abs_tol, rel_tol=rel_tol)
            root_values[name] = {"cell": cell_address, "actual": _json_value(actual), "expected": expected,
                                 "formula_matches": formula_ok, "matches": value_ok and formula_ok}
            if not value_ok or not formula_ok:
                mismatches.append({"cell": f"Premium!{cell_address}", "actual": _json_value(actual),
                                   "expected": expected, "formula_matches": formula_ok})

        compared = 0
        annual_compared = 0
        for address, record in evaluator.cells.items():
            if record.get("sheet") != "Premium":
                continue
            match = re.fullmatch(r"Premium!([A-Z]+)([0-9]+)", address)
            if match is None:
                continue
            column, row = column_index_from_string(match.group(1)), int(match.group(2))
            if not (9 <= row <= 115 and 2 <= column <= 82) or record.get("role") == "source_value":
                continue
            expected_formula = expected_formulas[row - 9][column - 2]
            expected_value = expected_values[row - 9][column - 2]
            formula_ok = _formula_signature(record.get("formula")) == _formula_signature(expected_formula)
            value_ok = _matches_value(record.get("value"), expected_value, abs_tol=abs_tol, rel_tol=rel_tol)
            compared += 1
            if 10 <= row <= 115 and (8 <= column <= 17 or 20 <= column <= 22 or 74 <= column <= 79):
                annual_compared += 1
            if not formula_ok or not value_ok:
                mismatches.append({"cell": address, "actual": record.get("value"), "expected": expected_value,
                                   "formula_matches": formula_ok})
        reconciliation = {"status": "pass" if not mismatches else "fail", "abs_tol": abs_tol,
                          "rel_tol": rel_tol, "root_values": root_values,
                          "compared_premium_formulas": compared, "annual_value_count": annual_compared,
                          "mismatches": mismatches}
        result = {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp",
                  "status": reconciliation["status"], "target": target, "target_location": "Premium!J1",
                  "source_path": str(source), "source_sha256": source_hash,
                  "scenario": {"id": "saved-workbook-configuration", "primary_inputs": source_inputs,
                               "input_overrides": []},
                  "baseline": {"path": str(baseline_file), "source_sha256": baseline.get("source_sha256"),
                               "engine": baseline.get("engine"), "engine_settings": baseline.get("engine_settings")},
                  "formula_cache_policy": _FORMULA_CACHE_POLICY,
                  "computed": {name: _json_value(value) for name, value in computed_values.items()},
                  "reconciliation": reconciliation,
                  "counts": {"evaluated_cells": len(evaluator.cells), "dependency_records": len(evaluator.edges),
                             "lookups": len(evaluator.lookups), "branches": len(evaluator.branches),
                             "dynamic_ranges": len(evaluator.dynamic_ranges),
                             "source_value_cells": sum(item["role"] == "source_value" for item in evaluator.cells.values()),
                             "calculated_formula_cells": sum(item["role"] != "source_value" for item in evaluator.cells.values())}}
        trace = {"schema_version": "gp.active_trace.v1", "source_path": str(source), "source_sha256": source_hash,
                 "cells": sorted(evaluator.cells.values(), key=lambda item: item["address"]),
                 "discovery_order": list(evaluator.cells),
                 "completion_order": [evaluator._address(evaluator.sheet_names[key[0]], key[1], key[2])
                                      for key in evaluator.cache],
                 "edges": sorted(evaluator.edges.values(), key=lambda item: (item["consumer"], item["prerequisite"])),
                 "names": sorted(evaluator.names.values(), key=lambda item: (item["scope"], item["name"])),
                 "lookups": evaluator.lookups, "branches": evaluator.branches,
                 "dynamic_ranges": evaluator.dynamic_ranges,
                 "source_error_values": [item for item in evaluator.cells.values()
                                         if isinstance(item.get("value"), dict) and "excel_error" in item["value"]]}
    except (CalculationBlocked, ValueError, KeyError, TypeError) as exc:
        if input_mismatches:
            return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                    "target": target, "source_path": str(source), "source_sha256": source_hash,
                    "reason": f"saved scenario does not match native baseline: {input_mismatches}",
                    "input_mismatches": input_mismatches}
        result = {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                  "target": target, "target_location": "Premium!J1", "source_path": str(source),
                  "source_sha256": source_hash, "reason": str(exc),
                  "scenario": {"primary_inputs": source_inputs},
                  "formula_cache_policy": _FORMULA_CACHE_POLICY, "input_mismatches": input_mismatches,
                  "reconciliation": {"status": "blocked", "abs_tol": abs_tol, "rel_tol": rel_tol}}
        if evaluator is not None:
            _lookup_headers(evaluator)
            trace = {"schema_version": "gp.active_trace.v1", "source_path": str(source), "source_sha256": source_hash,
                     "cells": sorted(evaluator.cells.values(), key=lambda item: item["address"]),
                     "discovery_order": list(evaluator.cells),
                     "completion_order": [evaluator._address(evaluator.sheet_names[key[0]], key[1], key[2])
                                          for key in evaluator.cache],
                     "edges": sorted(evaluator.edges.values(), key=lambda item: (item["consumer"], item["prerequisite"])),
                     "names": sorted(evaluator.names.values(), key=lambda item: (item["scope"], item["name"])),
                     "lookups": evaluator.lookups, "branches": evaluator.branches,
                     "dynamic_ranges": evaluator.dynamic_ranges}
    finally:
        if evaluator is not None:
            evaluator.close()
    if _hash_file(source) != source_hash:
        return {"schema_version": "gp.calculation_result.v1", "tool": "step3.gp", "status": "blocked",
                "reason": "source workbook changed during calculation"}
    result_path = out / "gp_result.json"
    trace_path = out / "active_trace.json"
    result["artifacts"] = [str(result_path), str(out / "gp_result.md")]
    if trace:
        result["artifacts"].extend([str(trace_path), str(out / "active_trace.md")])
    _atomic_write(result_path, _json_bytes(result))
    _atomic_write(out / "gp_result.md", _result_markdown(result).encode("utf-8"))
    if trace:
        _atomic_write(trace_path, _json_bytes(trace))
        _atomic_write(out / "active_trace.md", _trace_markdown(trace).encode("utf-8"))
    return result


def _oracle_markdown(result: dict[str, Any]) -> str:
    lines = ["# Native Excel oracle", "", f"Status: **{result['status']}**", "",
             f"Source: `{result['source_path']}`", f"Source SHA-256: `{result['source_sha256']}`", "",
             f"Engine: {result['engine']} {result['engine_settings'].get('version', '')}; "
             f"full rebuild: {result['full_rebuild_seconds']:.3f} s.", "",
             "## Named targets", ""]
    lines.extend(f"- `{name}`: `{record.get('value')}` (formula `{record.get('formula')}`)"
                 for name, record in result.get("named_targets", {}).items())
    lines.extend(["", "## Requested ranges", ""])
    for address, record in result.get("ranges", {}).items():
        rows = len(record.get("values", []))
        columns = len(record["values"][0]) if rows else 0
        lines.append(f"- `{address}`: {rows} rows by {columns} columns, with values and formulas in JSON.")
    lines.extend(["", "Macros executed: `false`; input overrides: `none`; workbook saved: `false`."])
    return "\n".join(lines) + "\n"


def capture_excel_oracle(workbook_path: Path, out_dir: Path, *, targets: list[str],
                         ranges: list[str]) -> dict[str, Any]:
    """Capture requested cells from a private read-only copy using native Excel on Windows."""
    source = workbook_path.expanduser().resolve()
    try:
        out = _validate_output_path(out_dir, source)
    except ValueError as exc:
        return {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "blocked",
                "reason": str(exc)}
    if not source.is_file():
        return {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "blocked",
                "reason": "source workbook does not exist"}
    if not targets and not ranges:
        return {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "blocked",
                "reason": "request at least one defined name or finite range"}
    if os.name != "nt":
        return {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "blocked",
                "reason": "native Excel oracle is available only on Windows with Microsoft Excel installed"}
    powershell = shutil.which("powershell.exe") or shutil.which("pwsh.exe")
    if not powershell:
        return {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "blocked",
                "reason": "PowerShell is unavailable; native Excel oracle could not start"}

    source_hash = _hash_file(source)
    canonical_targets = list(dict.fromkeys(targets))
    canonical_ranges: list[dict[str, Any]] = []
    total_cells = 0
    try:
        workbook = load_workbook(source, data_only=False, read_only=True, keep_vba=False)
        try:
            sheet_titles = {title.casefold(): title for title in workbook.sheetnames}
            for requested in list(dict.fromkeys(ranges)):
                reference = _parse_reference(requested, "")
                if reference is None or "!" not in requested:
                    raise ValueError(f"range must be a qualified finite A1 range: {requested}")
                sheet_title = sheet_titles.get(reference.sheet.casefold())
                if sheet_title is None:
                    raise ValueError(f"worksheet does not exist: {reference.sheet}")
                size = (reference.max_row - reference.min_row + 1) * (reference.max_col - reference.min_col + 1)
                total_cells += size
                if total_cells > 50_000:
                    raise ValueError("requested oracle ranges exceed the 50,000-cell limit")
                start = f"{get_column_letter(reference.min_col)}{reference.min_row}"
                end = f"{get_column_letter(reference.max_col)}{reference.max_row}"
                address = start if reference.single else f"{start}:{end}"
                canonical_ranges.append({"sheet": sheet_title, "address": address,
                                         "qualified_address": f"{sheet_title}!{address}"})
            available_names = set(workbook.defined_names)
            available_names.update(name for sheet in workbook.worksheets for name in sheet.defined_names)
            unknown = [name for name in canonical_targets if name not in available_names]
            if unknown:
                raise ValueError(f"unknown defined name(s): {', '.join(unknown)}")
        finally:
            workbook.close()
    except (OSError, ValueError) as exc:
        return {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "blocked",
                "reason": str(exc)}

    resource = Path(str(files("excel_to_act.steps.step5").joinpath("native_excel_oracle.ps1")))
    primary_addresses = ["C3", "C4", "C6", "C7", "C8", "J48", "O26"]
    try:
        with tempfile.TemporaryDirectory(prefix="step3-excel-oracle-") as temp_name:
            temp_dir = Path(temp_name)
            copy_path = temp_dir / source.name
            request_path = temp_dir / "request.json"
            output_path = temp_dir / "oracle.json"
            shutil.copy2(source, copy_path)
            copy_hash = _hash_file(copy_path)
            request = {"source_path": str(source), "source_sha256": source_hash, "copy_path": str(copy_path),
                       "targets": canonical_targets, "ranges": canonical_ranges,
                       "primary_input_addresses": primary_addresses}
            request_path.write_bytes(_json_bytes(request))
            completed = subprocess.run(
                [powershell, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File",
                 str(resource), "-RequestPath", str(request_path), "-OutputPath", str(output_path)],
                capture_output=True, text=True, timeout=300, check=False)
            if completed.returncode != 0:
                detail = (completed.stderr or completed.stdout).strip()
                raise CalculationBlocked(f"native Excel oracle failed: {detail[-2000:]}")
            payload = json.loads(output_path.read_text(encoding="utf-8"))
            if _hash_file(copy_path) != copy_hash or _hash_file(source) != source_hash:
                raise CalculationBlocked("source workbook or private oracle copy changed during native recalculation")
    except (OSError, ValueError, json.JSONDecodeError, subprocess.TimeoutExpired, CalculationBlocked) as exc:
        return {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "blocked",
                "source_path": str(source), "source_sha256": source_hash, "reason": str(exc)}

    named_targets = payload.get("named_targets", {})
    after_full_rebuild = payload.get("after_full_rebuild", {})
    payload["source_copy_sha256"] = copy_hash
    payload["source_path"] = str(source)
    payload["source_sha256"] = source_hash
    payload["schema_version"] = "gp.excel_oracle.v1"
    payload["tool"] = "step5.oracle"
    payload["status"] = "pass"
    payload["requested_targets"] = canonical_targets
    payload["requested_ranges"] = [item["qualified_address"] for item in canonical_ranges]
    payload["input_overrides"] = []
    payload["macros_executed"] = False
    target_cells = {"AnnuityDue": "H1", "PVLoading": "H2", "PVFB": "H3", "GP": "J1"}
    for name, cell in target_cells.items():
        if name in named_targets:
            after_full_rebuild[cell] = named_targets[name]
    payload["after_full_rebuild"] = after_full_rebuild
    if "Premium!B9:CD115" in payload.get("ranges", {}):
        premium = payload["ranges"]["Premium!B9:CD115"]
        payload["premium_values_B9_CD115"] = premium.get("values", [])
        payload["premium_formulas_B9_CD115"] = premium.get("formulas", [])
    result_path = out / "excel_oracle.json"
    _atomic_write(result_path, _json_bytes(payload))
    _atomic_write(out / "excel_oracle.md", _oracle_markdown(payload).encode("utf-8"))
    range_summaries = []
    for address, record in payload.get("ranges", {}).items():
        values = record.get("values", [])
        range_summaries.append({"address": address, "rows": len(values),
                                "columns": len(values[0]) if values else 0,
                                "excel_error_count": sum(error is not None for row in record.get("excel_errors", []) for error in row)})
    return {"schema_version": payload["schema_version"], "tool": payload["tool"], "status": payload["status"],
            "source_path": payload["source_path"], "source_sha256": payload["source_sha256"],
            "engine": payload["engine"], "engine_settings": payload["engine_settings"],
            "full_rebuild_seconds": payload["full_rebuild_seconds"],
            "macros_executed": payload["macros_executed"], "input_overrides": payload["input_overrides"],
            "named_targets": payload.get("named_targets", {}), "ranges": range_summaries,
            "artifacts": [str(result_path), str(out / "excel_oracle.md")]}

