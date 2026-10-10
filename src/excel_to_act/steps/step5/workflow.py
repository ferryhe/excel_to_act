"""Run the standalone generated model and reconcile it with a fresh Excel oracle."""

from __future__ import annotations

import ast
import hashlib
import json
import math
from pathlib import Path, PurePosixPath
import re
import shutil
import subprocess
import sys
import tempfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

from excel_to_act.steps.conversion_workflow import (
    append_stage_artifact,
    current_revision,
    design_question_lines,
    display_report_value,
    hash_file,
    load_workflow,
    read_json,
    require_stage_approved,
)
from excel_to_act.steps.step3.calculation import _formula_signature, _parse_reference, capture_excel_oracle


_PRIMARY_ADDRESSES = {"C3", "C4", "C6", "C7", "C8", "J48", "O26"}
_STDLIB = {"__future__", "argparse", "collections", "copy", "dataclasses", "datetime", "functools",
           "hashlib", "itertools", "json", "math", "operator", "pathlib", "re", "statistics", "sys", "typing"}


def _inside(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def _bundle_file_paths(bundle: Path, expected_files: Any) -> list[tuple[str, Path]]:
    if not isinstance(expected_files, dict) or not expected_files:
        raise ValueError("generated bundle must declare its complete file/hash ledger")
    result: list[tuple[str, Path]] = []
    for name, expected_hash in sorted(expected_files.items()):
        if not isinstance(name, str) or not name or not isinstance(expected_hash, str):
            raise ValueError("generated bundle file ledger entries must map relative paths to SHA-256 strings")
        relative = PurePosixPath(name.replace("\\", "/"))
        if relative.is_absolute() or any(part in {"", ".", ".."} for part in relative.parts):
            raise ValueError(f"generated bundle file path is not a safe relative path: {name}")
        path = (bundle / Path(*relative.parts)).resolve()
        if not _inside(path, bundle) or not path.is_file() or hash_file(path) != expected_hash:
            raise ValueError(f"generated bundle file changed or is missing: {name}")
        result.append((name, path))
    return result


def _copy_bundle_files(bundle: Path, expected_files: Any, destination: Path) -> list[str]:
    copied: list[str] = []
    for name, source_path in _bundle_file_paths(bundle, expected_files):
        target_path = destination / Path(*PurePosixPath(name.replace("\\", "/")).parts)
        target_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source_path, target_path)
        copied.append(name)
    return copied


def _stage4(root: Path) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any], Path]:
    workflow, revision = require_stage_approved(root, 4)
    revision_root = root / revision["artifact"]["json"]
    report = read_json(revision_root)
    if not isinstance(report, dict) or report.get("schema_version") != "step4.generation_report.v1":
        raise ValueError("current Step 4 artifact is not a generation report")
    bundle = Path(report["bundle"]["path"]).resolve()
    if not bundle.is_dir() or not _inside(bundle, root / "stage4"):
        raise ValueError("generated bundle is missing or outside this workflow's Stage 4 folder")
    manifest_path = bundle / "model_manifest.json"
    manifest = read_json(manifest_path)
    if manifest.get("schema_version") not in {"step4.generated_model.v1", "step4.modular_model.v1"}:
        raise ValueError("generated bundle manifest is missing or unsupported")
    report_files = report["bundle"].get("files", {})
    _bundle_file_paths(bundle, report_files)
    manifest_files = manifest.get("files", {})
    if not isinstance(manifest_files, dict):
        raise ValueError("generated model manifest must declare the bundle file/hash ledger")
    for name, expected in manifest_files.items():
        if report_files.get(name) != expected:
            raise ValueError(f"generation report and model manifest disagree on bundle file: {name}")
    if manifest.get("source_sha256") != workflow["source"].get("workbook_sha256"):
        raise ValueError("generated bundle source identity differs from this workflow")
    return workflow, revision, report, bundle


def capture_oracle(workflow_dir: Path, targets: list[str] | None = None,
                   ranges: list[str] | None = None) -> dict[str, Any]:
    """Capture new native Excel evidence without advancing the Stage 5 checkpoint."""
    try:
        root, _ = load_workflow(workflow_dir)
        workflow, stage4_revision, report, _bundle = _stage4(root)
        source = Path(workflow["source"]["workbook_path"]).resolve()
        requested_targets = list(dict.fromkeys(targets if targets is not None else report.get("target_names", [])))
        if ranges is not None:
            planned_ranges = ranges
        else:
            planned_ranges = (report.get("oracle_ranges") or report.get("comparison_ranges")
                              or (["Premium!B9:CD115"] if "GP" in requested_targets else []))
        requested_ranges = list(dict.fromkeys(planned_ranges))
        if not requested_targets and not requested_ranges:
            raise ValueError("request a target name or a finite range")
        stage5_root = root / "stage5"
        stage5_root.mkdir(parents=True, exist_ok=True)
        attempt = 1
        while (stage5_root / f"oracle-attempt-{attempt:04d}").exists():
            attempt += 1
        output_dir = stage5_root / f"oracle-attempt-{attempt:04d}"
        result = capture_excel_oracle(source, output_dir, targets=requested_targets, ranges=requested_ranges)
        if result.get("status") != "pass":
            return {**result, "tool": "step5.oracle", "status": "blocked"}
        oracle_path = output_dir / "excel_oracle.json"
        oracle = read_json(oracle_path)
        oracle["tool"] = "step5.oracle"
        oracle["stage4_artifact_sha256"] = stage4_revision["artifact"]["json_sha256"]
        oracle["scenario_ids"] = report.get("scenario_ids", [])
        oracle_path.write_text(json.dumps(oracle, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
        result["tool"] = "step5.oracle"
        result["artifact_sha256"] = hash_file(oracle_path)
        result["artifact"] = str(oracle_path)
        result["markdown"] = str(output_dir / "excel_oracle.md")
        result["stage_advanced"] = False
        return result
    except Exception as exc:
        return {"schema_version": "gp.excel_oracle.v1", "tool": "step5.oracle", "status": "blocked",
                "reason": str(exc), "stage_advanced": False}


def _verify_python(bundle: Path, manifest: dict[str, Any]) -> list[str]:
    failures: list[str] = []
    try:
        files = _bundle_file_paths(bundle, manifest.get("files"))
    except (OSError, ValueError) as exc:
        return [str(exc)]
    python_files = [(name, path) for name, path in files if path.suffix == ".py"]
    if not any(name == "model.py" for name, _path in python_files):
        failures.append("generated bundle is missing model.py")
    local_modules = {PurePosixPath(name.replace("\\", "/")).stem for name, _path in python_files}
    for name, path in python_files:
        source = path.read_text(encoding="utf-8")
        try:
            compile(source, name, "exec")
            tree = ast.parse(source, filename=name)
        except (SyntaxError, UnicodeError) as exc:
            failures.append(f"{name} does not compile: {exc}")
            continue
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                for alias in node.names:
                    root_name = alias.name.split(".", 1)[0]
                    if root_name not in _STDLIB and root_name not in local_modules:
                        failures.append(f"{name} imports unsupported module {root_name}")
            elif isinstance(node, ast.ImportFrom) and node.module:
                root_name = node.module.split(".", 1)[0]
                if node.level or (root_name not in _STDLIB and root_name not in local_modules):
                    failures.append(f"{name} imports unsupported module {root_name}")
    for name, _path in files:
        if Path(name).suffix.casefold() in {".xls", ".xlsx", ".xlsm"}:
            failures.append(f"source workbook must not be included in the generated bundle: {name}")
    python_names = {name for name, _path in python_files}
    modular_bundle = (manifest.get("execution_kind") == "modular"
                      or manifest.get("schema_version") == "step4.modular_model.v1")
    runtime_name = "modular_runtime.py" if modular_bundle else "runtime.py"
    if runtime_name not in python_names:
        failures.append(f"generated bundle is missing {runtime_name}")
    raw_values = read_json(bundle / "source_values.json")
    if not isinstance(raw_values, dict):
        failures.append("source_values.json is not an object")
        return failures
    raw_addresses = set(manifest.get("raw_input_addresses", []))
    formula_addresses = set(manifest.get("formula_addresses", []))
    if set(raw_values) != raw_addresses:
        failures.append("source_values.json keys do not match the generated manifest raw-input ledger")
    if raw_addresses & formula_addresses:
        failures.append("a formula address is also listed as a raw input")
    if modular_bundle:
        layout_path = bundle / "input_layout.json"
        layout = read_json(layout_path) if layout_path.is_file() else {}
        bindings = layout.get("external_bindings", []) if isinstance(layout, dict) else None
        if not isinstance(bindings, list):
            failures.append("modular input_layout.json has no external-binding list")
            bindings = []
        external_ids: list[str] = []
        bound_external_addresses: list[str] = []
        for binding in bindings:
            if not isinstance(binding, dict) or not isinstance(binding.get("variable_id"), str):
                failures.append("modular external binding is malformed")
                continue
            variable_id = binding["variable_id"]
            external_ids.append(variable_id)
            shape = binding.get("shape")
            addresses = binding.get("addresses")
            indices = binding.get("indices")
            if (not isinstance(shape, list) or len(shape) != 1 or not isinstance(shape[0], int)
                    or isinstance(shape[0], bool) or shape[0] < 1
                    or not isinstance(addresses, list) or len(addresses) != shape[0]
                    or not isinstance(indices, list) or len(indices) != shape[0]):
                failures.append(f"modular external binding has an invalid shape or coordinate ledger: {variable_id}")
                continue
            try:
                normalized = [_address_key(address) for address in addresses]
            except (TypeError, ValueError) as exc:
                failures.append(f"modular external binding has an invalid source address: {variable_id}: {exc}")
                continue
            if len(set(normalized)) != len(normalized):
                failures.append(f"modular external binding repeats a source coordinate: {variable_id}")
            bound_external_addresses.extend(normalized)
            if indices != [[index] for index in range(shape[0])]:
                failures.append(f"modular external binding indices are not a complete vector: {variable_id}")
        if len(set(external_ids)) != len(external_ids):
            failures.append("modular external bindings repeat a variable ID")
        try:
            declared_external_addresses = [_address_key(address)
                                           for address in manifest.get("external_input_addresses", [])]
        except (TypeError, ValueError) as exc:
            failures.append(f"generated manifest has an invalid external input address: {exc}")
            declared_external_addresses = []
        if (len(set(declared_external_addresses)) != len(declared_external_addresses)
                or set(declared_external_addresses) != set(bound_external_addresses)
                or len(declared_external_addresses) != manifest.get("external_input_count", 0)):
            failures.append("generated external address ledger differs from its logical bindings")
        try:
            normalized_raw_addresses = {_address_key(address) for address in raw_addresses}
            normalized_formula_addresses = {_address_key(address) for address in formula_addresses}
        except (TypeError, ValueError) as exc:
            failures.append(f"generated source/formula address ledger is invalid: {exc}")
            normalized_raw_addresses, normalized_formula_addresses = set(), set()
        if set(bound_external_addresses) & (normalized_raw_addresses | normalized_formula_addresses):
            failures.append("an external boundary coordinate overlaps a raw or calculated formula address")

        if external_ids:
            external_name = "external_values.json"
            if external_name not in manifest.get("files", {}) or not (bundle / external_name).is_file():
                failures.append("generated bundle is missing its declared external_values.json ledger")
            else:
                try:
                    external = read_json(bundle / external_name)
                    values = external.get("values") if isinstance(external, dict) else None
                    capture_binding = layout.get("external_capture", {})
                    values_sha = hashlib.sha256(
                        (json.dumps(values, sort_keys=True, ensure_ascii=False, separators=(",", ":")) + "\n")
                        .encode("utf-8")).hexdigest()
                    if (not isinstance(external, dict)
                            or external.get("schema_version") != "step4.external_values.v1"
                            or external.get("source_sha256") != manifest.get("source_sha256")
                            or external.get("capture_artifact_sha256") != capture_binding.get("artifact_sha256")
                            or external.get("capture_artifact_sha256") != manifest.get("external_capture_artifact_sha256")
                            or external.get("capture_sha256") != manifest.get("external_capture_sha256")
                            or external.get("values_sha256") != manifest.get("external_values_sha256")
                            or external.get("values_sha256") != values_sha
                            or not isinstance(values, dict) or set(values) != set(external_ids)):
                        failures.append("external_values.json is stale or differs from the generated external bindings")
                    else:
                        for binding in bindings:
                            vector = values.get(binding["variable_id"])
                            if not isinstance(vector, list) or len(vector) != binding["shape"][0]:
                                failures.append("external_values.json has an invalid vector shape: "
                                                + binding["variable_id"])
                except (OSError, ValueError, TypeError) as exc:
                    failures.append(f"external_values.json cannot be validated: {exc}")
        elif declared_external_addresses or manifest.get("external_input_count", 0):
            failures.append("generated manifest declares external coordinates without logical bindings")
    if manifest.get("formula_cache_inputs") is not False:
        failures.append("generated manifest permits formula-cache inputs")
    return failures


def validate_generated(workflow_dir: Path) -> dict[str, Any]:
    """Compile and run the generated bundle in an isolated directory with no workbook."""
    try:
        root, workflow = load_workflow(workflow_dir)
        _workflow, stage4_revision, report, bundle = _stage4(root)
        manifest = read_json(bundle / "model_manifest.json")
        failures = _verify_python(bundle, manifest)
        if failures:
            raise ValueError("generated-code checks failed: " + "; ".join(failures[:12]))
        stage5_root = root / "stage5" / "execution"
        stage5_root.mkdir(parents=True, exist_ok=True)
        attempt = 1
        while (stage5_root / f"attempt-{attempt:04d}").exists():
            attempt += 1
        attempt_dir = stage5_root / f"attempt-{attempt:04d}"
        attempt_dir.mkdir()
        with tempfile.TemporaryDirectory(prefix="step5-standalone-") as temp_name:
            isolated = Path(temp_name)
            copied_files = _copy_bundle_files(bundle, report["bundle"].get("files", {}), isolated)
            if set(copied_files) != set(report["bundle"].get("files", {})):
                raise ValueError("isolated validation did not receive every declared bundle file")
            if (isolated / Path(workflow["source"]["workbook_path"]).name).exists():
                raise ValueError("source workbook was copied into the standalone validation directory")
            result_path = isolated / "model_result.json"
            runner = ("import runpy,sys; model_dir=sys.argv[1]; output=sys.argv[2]; "
                      "sys.path.insert(0,model_dir); sys.argv=[model_dir+'/model.py','--out',output]; "
                      "runpy.run_path(sys.argv[0],run_name='__main__')")
            completed = subprocess.run([sys.executable, "-I", "-c", runner, str(isolated), str(result_path)],
                                       cwd=isolated, capture_output=True, text=True, timeout=600, check=False)
            if completed.returncode != 0 or not result_path.is_file():
                detail = (completed.stderr or completed.stdout).strip()[-3000:]
                raise ValueError(f"generated standalone run failed with exit {completed.returncode}: {detail}")
            result_bytes = result_path.read_bytes()
            result = json.loads(result_bytes)
            saved_result = attempt_dir / "model_result.json"
            saved_result.write_bytes(result_bytes)
        stage4_json = root / stage4_revision["artifact"]["json"]
        evidence = {"schema_version": "step5.execution_validation.v1", "tool": "step5.validate",
                    "status": "pass", "source_sha256": workflow["source"]["workbook_sha256"],
                    "stage4_artifact_sha256": stage4_revision["artifact"]["json_sha256"],
                    "generation_report_sha256": hash_file(stage4_json),
                    "bundle_path": str(bundle), "bundle_files": report["bundle"]["files"],
                    "result_path": str(saved_result), "result_sha256": hash_file(saved_result),
                    "compile_pass": True, "standalone_run_pass": True,
                    "runtime_dependencies": ["Python standard library only"],
                    "source_workbook_present_at_runtime": False, "project_package_required_at_runtime": False,
                    "result_summary": {"cell_count": result.get("cell_count"),
                                       "formula_count": result.get("formula_count"),
                                       "targets": result.get("targets", {})}}
        evidence_path = attempt_dir / "execution_validation.json"
        evidence_path.write_text(json.dumps(evidence, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8")
        markdown_path = attempt_dir / "execution_validation.md"
        markdown_path.write_text("\n".join(["# Step 5 Generated-Code Validation", "", "Status: **Pass**", "",
                                              f"Source SHA-256: `{evidence['source_sha256']}`",
                                              f"Step 4 report SHA-256: `{evidence['stage4_artifact_sha256']}`",
                                              f"Generated cells: {result.get('cell_count')}; formulas: {result.get('formula_count')}.",
                                              "The generated model compiled and ran in an isolated folder without the source workbook or project package.",
                                              "This report contains no native Excel comparison; run `step5 reconcile` separately.", ""]), encoding="utf-8")
        return {"schema_version": evidence["schema_version"], "tool": evidence["tool"], "status": "pass",
                "validation": str(evidence_path), "markdown": str(markdown_path),
                "result": str(saved_result), "result_sha256": evidence["result_sha256"],
                "result_summary": evidence["result_summary"]}
    except Exception as exc:
        return {"schema_version": "step5.execution_validation.v1", "tool": "step5.validate",
                "status": "blocked", "reason": str(exc)}


def _matches(actual: Any, expected: Any, abs_tol: float, rel_tol: float) -> bool:
    if isinstance(actual, dict) and "excel_error" in actual:
        actual = actual["excel_error"]
    if isinstance(expected, dict) and "excel_error" in expected:
        expected = expected["excel_error"]
    if isinstance(actual, bool) or isinstance(expected, bool):
        return actual is expected
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return math.isclose(float(actual), float(expected), abs_tol=abs_tol, rel_tol=rel_tol)
    return actual == expected


def _preflight_oracle(oracle: Any, source_hash: str, stage4_hash: str) -> dict[str, Any]:
    if not isinstance(oracle, dict) or oracle.get("schema_version") != "gp.excel_oracle.v1":
        raise ValueError("baseline must be a gp.excel_oracle.v1 JSON object from step5.oracle")
    if oracle.get("tool") != "step5.oracle" or oracle.get("status") != "pass" or oracle.get("engine") != "Microsoft Excel":
        raise ValueError("baseline must be a successful Step 5 Microsoft Excel full-rebuild result")
    if type(oracle.get("calculation_state")) is not int or oracle["calculation_state"] != 0:
        raise ValueError("native Excel calculation must be complete before its oracle can be reconciled")
    if oracle.get("source_sha256") != source_hash or oracle.get("source_copy_sha256") != source_hash:
        raise ValueError("oracle source and private-copy hashes must match the approved workbook")
    if oracle.get("stage4_artifact_sha256") != stage4_hash:
        raise ValueError("oracle was captured for a different approved Step 4 artifact")
    if oracle.get("macros_executed") is not False or oracle.get("input_overrides") != []:
        raise ValueError("oracle must disable macros and use no input overrides")
    primary = oracle.get("primary_inputs")
    if not isinstance(primary, dict) or set(primary) != _PRIMARY_ADDRESSES:
        raise ValueError("oracle must include all seven saved primary input addresses")
    if any(not isinstance(record, dict) or "value" not in record for record in primary.values()):
        raise ValueError("each oracle primary input must be an object containing a value")
    return oracle


def _address_key(address: str) -> str:
    reference = _parse_reference(address, "")
    if (reference is None or not reference.single or not reference.sheet):
        raise ValueError(f"formula address must be a qualified single cell: {address}")
    return f"{reference.sheet.casefold()}!{get_column_letter(reference.min_col)}{reference.min_row}"


def _oracle_range_cells(oracle: dict[str, Any]) -> tuple[dict[str, dict[str, Any]], list[str]]:
    ranges = oracle.get("ranges")
    if not isinstance(ranges, dict) or not ranges:
        raise ValueError("native oracle must include the requested formula ranges")
    cells: dict[str, dict[str, Any]] = {}
    selectors: list[str] = []
    for selector, record in ranges.items():
        if not isinstance(record, dict):
            raise ValueError(f"native oracle range record is invalid: {selector}")
        qualified = record.get("qualified_address", selector)
        bounds = _parse_reference(qualified, "")
        if bounds is None or not bounds.sheet:
            raise ValueError(f"native oracle range is not a qualified finite address: {qualified}")
        values, formulas = record.get("values"), record.get("formulas")
        row_count, column_count = bounds.max_row - bounds.min_row + 1, bounds.max_col - bounds.min_col + 1
        if (not isinstance(values, list) or len(values) != row_count
                or not isinstance(formulas, list) or len(formulas) != row_count
                or any(not isinstance(row, list) or len(row) != column_count for row in values)
                or any(not isinstance(row, list) or len(row) != column_count for row in formulas)):
            raise ValueError(f"native oracle values/formulas do not match range shape: {qualified}")
        errors = record.get("excel_errors")
        if errors is not None and (not isinstance(errors, list) or len(errors) != row_count
                                   or any(not isinstance(row, list) or len(row) != column_count for row in errors)):
            raise ValueError(f"native oracle error matrix does not match range shape: {qualified}")
        selectors.append(qualified)
        for row_offset, row_values in enumerate(values):
            for column_offset, value in enumerate(row_values):
                row = bounds.min_row + row_offset
                column = bounds.min_col + column_offset
                address = f"{bounds.sheet}!{get_column_letter(column)}{row}"
                key = _address_key(address)
                error = errors[row_offset][column_offset] if errors is not None else None
                if error:
                    value = {"excel_error": str(error)}
                record_value = {"address": address, "value": value,
                                "formula": formulas[row_offset][column_offset], "range": qualified}
                previous = cells.get(key)
                if previous is not None:
                    if previous["value"] != value or previous["formula"] != record_value["formula"]:
                        raise ValueError(f"overlapping native oracle ranges disagree at {address}")
                    continue
                cells[key] = record_value
    return cells, sorted(selectors)


def _active_formula_details(report: dict[str, Any], manifest: dict[str, Any],
                            source_path: Path | None = None,
                            source_hash: str | None = None) -> tuple[list[str], dict[str, str]]:
    trace_ref = report.get("active_trace", {})
    trace_path_value = trace_ref.get("path") if isinstance(trace_ref, dict) else None
    formula_texts: dict[str, str] = {}
    if isinstance(trace_path_value, str):
        trace_path = Path(trace_path_value).expanduser().resolve()
        if not trace_path.is_file() or hash_file(trace_path) != trace_ref.get("sha256"):
            raise ValueError("approved Step 4 active trace is missing or changed")
        trace = read_json(trace_path)
        if (not isinstance(trace, dict) or trace.get("source_sha256") != report.get("source_sha256")
                or not isinstance(trace.get("cells"), list)):
            raise ValueError("approved Step 4 active trace has the wrong source or cell list")
        formula_cells = [cell for cell in trace["cells"]
                         if isinstance(cell, dict)
                         and cell.get("role") in {"calculated_formula", "calculated_array_formula"}]
        addresses = [cell.get("address") for cell in formula_cells]
        for cell in formula_cells:
            address, formula = cell.get("address"), cell.get("formula")
            if not isinstance(address, str) or not isinstance(formula, str) or not formula.startswith("="):
                raise ValueError("approved active trace must include formula text for every active formula member")
            key = _address_key(address)
            if key in formula_texts:
                raise ValueError(f"approved active trace duplicates formula address: {address}")
            formula_texts[key] = formula
    else:
        addresses = manifest.get("active_formula_addresses")
        if not isinstance(addresses, list):
            addresses = manifest.get("active_addresses")
    if not isinstance(addresses, list) or not addresses or any(not isinstance(item, str) for item in addresses):
        raise ValueError("Step 4 does not provide the complete active formula-member address set")
    normalized = [_address_key(item) for item in addresses]
    if len(normalized) != len(set(normalized)):
        raise ValueError("Step 4 active formula-member address set contains duplicates")
    declared = manifest.get("active_formula_addresses")
    if isinstance(declared, list) and set(_address_key(item) for item in declared) != set(normalized):
        raise ValueError("model manifest active formula addresses differ from the bound trace")
    if not formula_texts and source_path is not None:
        if not isinstance(source_hash, str) or not source_path.is_file() or hash_file(source_path) != source_hash:
            raise ValueError("legacy active formula source is missing or differs from the bound workbook")
        workbook = load_workbook(source_path, data_only=False, read_only=False)
        try:
            sheet_names = {name.casefold(): name for name in workbook.sheetnames}
            for address in addresses:
                reference = _parse_reference(address, "")
                if reference is None or not reference.sheet or not reference.single:
                    raise ValueError(f"active formula address is not a qualified source cell: {address}")
                sheet_title = sheet_names.get(reference.sheet.casefold())
                if sheet_title is None:
                    raise ValueError(f"active formula source sheet is missing: {reference.sheet}")
                cell = workbook[sheet_title].cell(reference.min_row, reference.min_col)
                if cell.data_type != "f" or not isinstance(cell.value, str) or not cell.value.startswith("="):
                    raise ValueError(f"active formula source text is missing: {address}")
                formula_texts[_address_key(address)] = cell.value
        finally:
            workbook.close()
        if hash_file(source_path) != source_hash:
            raise ValueError("source workbook changed while reading active formula text")
    if formula_texts and set(formula_texts) != set(normalized):
        raise ValueError("bound active formula text set differs from the active address set")
    return addresses, formula_texts


def _active_formula_addresses(report: dict[str, Any], manifest: dict[str, Any]) -> list[str]:
    addresses, _formula_texts = _active_formula_details(report, manifest)
    return addresses


def _compare_active_formula_members(addresses: list[str], native_cells: dict[str, dict[str, Any]],
                                    model_cells: dict[str, Any], expected_formulas: dict[str, str],
                                    abs_tol: float, rel_tol: float
                                    ) -> tuple[int, int, float | None, set[str], list[dict[str, Any]]]:
    compared = 0
    formula_identity_matches = 0
    max_abs_difference: float | None = None
    compared_addresses: set[str] = set()
    mismatches: list[dict[str, Any]] = []
    for address in addresses:
        key = _address_key(address)
        native = native_cells.get(key)
        if native is None:
            mismatches.append({"kind": "active_formula_missing_from_oracle", "address": address})
            continue
        if key not in model_cells:
            mismatches.append({"kind": "active_formula_missing_from_model", "address": address})
            continue
        compared_addresses.add(key)
        compared += 1
        expected_formula = expected_formulas.get(key)
        actual_formula = native.get("formula")
        expected_signature = _formula_signature(expected_formula)
        actual_signature = _formula_signature(actual_formula)
        if expected_signature is None or actual_signature is None or expected_signature != actual_signature:
            mismatches.append({"kind": "active_formula_identity", "address": address,
                               "source_formula": expected_formula, "excel_formula": actual_formula})
        else:
            formula_identity_matches += 1
        if not _matches(model_cells[key], native["value"], abs_tol, rel_tol):
            mismatches.append({"kind": "formula_value", "address": address,
                               "formula": native.get("formula"), "model": model_cells[key],
                               "excel": native["value"]})
        model_value = model_cells[key]
        excel_value = native["value"]
        if (isinstance(model_value, (int, float)) and not isinstance(model_value, bool)
                and isinstance(excel_value, (int, float)) and not isinstance(excel_value, bool)
                and math.isfinite(float(model_value)) and math.isfinite(float(excel_value))):
            difference = abs(float(model_value) - float(excel_value))
            max_abs_difference = difference if max_abs_difference is None else max(max_abs_difference, difference)
    return compared, formula_identity_matches, max_abs_difference, compared_addresses, mismatches


def _read_literal_source_axis(source_path: Path, expected_source_sha256: Any,
                              axis_ref: Any, expected_values: list[Any]
                              ) -> tuple[list[Any], list[dict[str, Any]]]:
    """Verify the captured age axis is a literal source range, without using cached formula values."""
    mismatches: list[dict[str, Any]] = []
    if not source_path.is_file():
        return [], [{"kind": "external_age_axis_source_missing", "path": str(source_path)}]
    source_before = hash_file(source_path)
    if source_before != expected_source_sha256:
        return [], [{"kind": "external_age_axis_source_binding",
                     "expected": expected_source_sha256, "actual": source_before}]
    try:
        # This workbook is inspected only; VBA parts are neither executed nor saved.
        workbook = load_workbook(source_path, data_only=False, read_only=False)
    except Exception as exc:
        return [], [{"kind": "external_age_axis_source_unreadable", "error": str(exc)}]
    source_values: list[Any] = []
    try:
        sheet_title = next((name for name in workbook.sheetnames
                            if name.casefold() == axis_ref.sheet.casefold()), None)
        if sheet_title is None:
            mismatches.append({"kind": "external_age_axis_source_sheet_missing", "sheet": axis_ref.sheet})
        else:
            worksheet = workbook[sheet_title]
            rows = worksheet.iter_rows(min_row=axis_ref.min_row, max_row=axis_ref.max_row,
                                       min_col=axis_ref.min_col, max_col=axis_ref.max_col)
            for offset, row in enumerate(rows):
                cell = row[0]
                address = f"{sheet_title}!{get_column_letter(axis_ref.min_col)}{axis_ref.min_row + offset}"
                source_values.append(cell.value)
                if cell.data_type == "f":
                    mismatches.append({"kind": "external_age_axis_source_formula", "address": address,
                                       "formula": cell.value})
                elif offset >= len(expected_values) or not _matches(cell.value, expected_values[offset], 0, 0):
                    mismatches.append({"kind": "external_age_axis_source_value", "address": address,
                                       "source": cell.value,
                                       "captured": expected_values[offset] if offset < len(expected_values) else None})
            if len(source_values) != len(expected_values):
                mismatches.append({"kind": "external_age_axis_source_coverage",
                                   "expected": len(expected_values), "actual": len(source_values)})
    finally:
        workbook.close()
    source_after = hash_file(source_path)
    if source_after != expected_source_sha256:
        mismatches.append({"kind": "external_age_axis_source_binding",
                           "expected": expected_source_sha256, "actual": source_after})
    return source_values, mismatches


def _compare_external_boundary(manifest: dict[str, Any], bundle: Path,
                               model_result: dict[str, Any], native_cells: dict[str, dict[str, Any]],
                               abs_tol: float, rel_tol: float,
                               source_path: Path) -> tuple[dict[str, int], list[dict[str, Any]]]:
    """Compare separately captured external vectors and their source formula cuts."""
    layout_path = bundle / "input_layout.json"
    if not layout_path.is_file():
        return ({"external_input_coordinates": 0, "external_age_axis_keys": 0,
                 "external_formula_cut_members": 0},
                [{"kind": "external_boundary_missing", "reason": "input layout file is missing"}])
    layout = read_json(layout_path)
    expected_ids = {item.get("variable_id") for item in layout.get("external_bindings", [])
                    if isinstance(item, dict) and isinstance(item.get("variable_id"), str)}
    expected_addresses = manifest.get("external_input_addresses", [])
    metrics = {"external_input_coordinates": 0, "external_age_axis_keys": 0,
               "external_formula_cut_members": 0}
    if not expected_ids:
        if expected_addresses or manifest.get("external_input_ranges"):
            return metrics, [{"kind": "external_boundary_manifest", "reason": "manifest declares coordinates without external bindings"}]
        return metrics, []

    payload_path = bundle / "external_values.json"
    if not payload_path.is_file():
        return metrics, [{"kind": "external_boundary_missing", "reason": "external value or layout file is missing"}]
    external_payload = read_json(payload_path)
    values = external_payload.get("values")
    model_values = model_result.get("external_inputs")
    mismatches: list[dict[str, Any]] = []
    if (external_payload.get("schema_version") != "step4.external_values.v1"
            or external_payload.get("source_sha256") != manifest.get("source_sha256")
            or external_payload.get("capture_artifact_sha256") != manifest.get("external_capture_artifact_sha256")
            or external_payload.get("capture_artifact_sha256") != layout.get("external_capture", {}).get("artifact_sha256")
            or external_payload.get("capture_sha256") != manifest.get("external_capture_sha256")
            or external_payload.get("values_sha256") != manifest.get("external_values_sha256")
            or not isinstance(values, dict)
            or not isinstance(model_values, dict)
            or set(values) != expected_ids
            or set(model_values) != expected_ids):
        mismatches.append({"kind": "external_boundary_binding",
                           "reason": "external values, model result, and manifest bindings do not agree"})
        return metrics, mismatches

    boundaries = external_payload.get("boundaries")
    if (not isinstance(boundaries, list)
            or {item.get("boundary_variable_id") for item in boundaries if isinstance(item, dict)} != expected_ids
            or external_payload.get("formula_cut_formula_set_sha256")
            != manifest.get("external_formula_cut_formula_set_sha256")):
        mismatches.append({"kind": "external_boundary_binding",
                           "reason": "captured boundaries or formula-set hash differ from the generated manifest"})
        return metrics, mismatches

    cut_records = external_payload.get("formula_cut_members")
    cut_by_address = {_address_key(item.get("address")): item for item in cut_records
                      if isinstance(item, dict) and isinstance(item.get("address"), str)} \
        if isinstance(cut_records, list) else {}
    expected_cut_count = manifest.get("external_formula_cut_member_count")
    if (not isinstance(cut_records, list) or len(cut_by_address) != len(cut_records)
            or len(cut_records) != expected_cut_count):
        mismatches.append({"kind": "external_formula_cut_ledger",
                           "expected": expected_cut_count,
                           "actual": len(cut_records) if isinstance(cut_records, list) else None})

    actual_addresses: set[str] = set()
    for boundary in boundaries:
        if not isinstance(boundary, dict):
            mismatches.append({"kind": "external_boundary_record", "reason": "boundary record is malformed"})
            continue
        variable_id, extent, shape = (boundary.get("boundary_variable_id"), boundary.get("source_range"),
                                      boundary.get("shape"))
        if (variable_id not in expected_ids or not isinstance(extent, str)
                or not isinstance(shape, list) or len(shape) != 1
                or not isinstance(shape[0], int) or shape[0] < 1):
            mismatches.append({"kind": "external_boundary_record", "variable_id": variable_id})
            continue
        try:
            parsed = _parse_reference(extent, "")
        except (TypeError, ValueError):
            parsed = None
        if (parsed is None or not parsed.sheet or parsed.min_col != parsed.max_col
                or parsed.max_row - parsed.min_row + 1 != shape[0]):
            mismatches.append({"kind": "external_boundary_shape", "variable_id": variable_id,
                               "source_range": extent, "shape": shape})
            continue
        captured_vector, result_vector = values.get(variable_id), model_values.get(variable_id)
        if (not isinstance(captured_vector, list) or len(captured_vector) != shape[0]
                or not isinstance(result_vector, list) or len(result_vector) != shape[0]):
            mismatches.append({"kind": "external_boundary_shape", "variable_id": variable_id,
                               "expected_length": shape[0]})
            continue
        for index, row in enumerate(range(parsed.min_row, parsed.max_row + 1)):
            address = f"{parsed.sheet}!{get_column_letter(parsed.min_col)}{row}"
            key = _address_key(address)
            actual_addresses.add(key)
            native = native_cells.get(key)
            if native is None:
                mismatches.append({"kind": "external_input_missing_from_oracle", "address": address})
                continue
            cut = cut_by_address.get(key)
            formula = native.get("formula")
            expected_formula = cut.get("formula") if isinstance(cut, dict) else None
            expected_formula_sha = (hashlib.sha256(expected_formula.encode("utf-8")).hexdigest()
                                    if isinstance(expected_formula, str) else None)
            if (not isinstance(expected_formula, str) or formula != expected_formula
                    or cut.get("boundary_variable_id") != variable_id
                    or cut.get("formula_sha256") != expected_formula_sha):
                mismatches.append({"kind": "external_formula_cut_changed", "address": address,
                                   "captured_formula": expected_formula, "excel_formula": formula})
                continue
            metrics["external_formula_cut_members"] += 1
            metrics["external_input_coordinates"] += 1
            captured, generated, excel_value = captured_vector[index], result_vector[index], native.get("value")
            if not _matches(generated, captured, 0, 0):
                mismatches.append({"kind": "external_input_bundle_value", "variable_id": variable_id,
                                   "index": index, "address": address,
                                   "model": generated, "captured": captured})
            if not _matches(generated, excel_value, abs_tol, rel_tol):
                mismatches.append({"kind": "external_input_value", "variable_id": variable_id,
                                   "index": index, "address": address,
                                   "model": generated, "excel": excel_value})

    declared_addresses = {_address_key(address) for address in expected_addresses if isinstance(address, str)}
    if actual_addresses != declared_addresses:
        mismatches.append({"kind": "external_input_address_coverage",
                           "missing": sorted(actual_addresses - declared_addresses)[:10],
                           "extra": sorted(declared_addresses - actual_addresses)[:10]})

    age_axis = external_payload.get("age_axis")
    layout_age = layout.get("external_capture", {}).get("age_axis")
    manifest_age = manifest.get("external_age_axis")
    age_values = age_axis.get("values") if isinstance(age_axis, dict) else None
    model_axis = model_result.get("external_age_axis")
    if (not isinstance(age_axis, dict) or not isinstance(layout_age, dict)
            or age_axis.get("source_range") != layout_age.get("source_range")
            or not isinstance(manifest_age, dict)
            or age_axis.get("source_range") != manifest_age.get("source_range")
            or age_axis.get("value_sha256") != manifest.get("external_age_axis_sha256")
            or age_axis.get("value_sha256") != manifest_age.get("values_sha256")
            or not isinstance(layout_age.get("metadata_group"), str)
            or layout_age.get("metadata_group") != manifest_age.get("metadata_group")
            or not isinstance(age_values, list) or not isinstance(model_axis, list)
            or model_axis != age_values):
        mismatches.append({"kind": "external_age_axis_binding",
                           "reason": "captured, generated, and manifest age-key metadata do not agree"})
    else:
        try:
            axis_ref = _parse_reference(age_axis["source_range"], "")
        except (TypeError, ValueError):
            axis_ref = None
        if (axis_ref is None or not axis_ref.sheet or axis_ref.min_col != axis_ref.max_col
                or axis_ref.max_row - axis_ref.min_row + 1 != len(age_values)):
            mismatches.append({"kind": "external_age_axis_shape", "source_range": age_axis.get("source_range")})
        else:
            metadata_file = bundle / "adapter_metadata.json"
            metadata = read_json(metadata_file) if metadata_file.is_file() else {}
            metadata_group = layout_age.get("metadata_group")
            metadata_record = metadata.get("groups", {}).get(metadata_group, {})
            if (metadata_record.get("values") != age_values
                    or [_address_key(address) for address in metadata_record.get("addresses", [])]
                    != [_address_key(f"{axis_ref.sheet}!{get_column_letter(axis_ref.min_col)}{row}")
                        for row in range(axis_ref.min_row, axis_ref.max_row + 1)]):
                mismatches.append({"kind": "external_age_axis_bundle_value",
                                   "reason": "captured age keys differ from bundled source metadata"})
            source_values, source_errors = _read_literal_source_axis(
                source_path, manifest.get("source_sha256"), axis_ref, age_values)
            mismatches.extend(source_errors)
            for index, row in enumerate(range(axis_ref.min_row, axis_ref.max_row + 1)):
                address = f"{axis_ref.sheet}!{get_column_letter(axis_ref.min_col)}{row}"
                native = native_cells.get(_address_key(address))
                if native is None:
                    mismatches.append({"kind": "external_age_axis_missing_from_oracle", "address": address})
                elif ((isinstance(native.get("formula"), str) and native["formula"].startswith("="))
                      or not _matches(age_values[index], native.get("value"), 0, 0)):
                    mismatches.append({"kind": "external_age_axis_value", "address": address,
                                       "model": age_values[index], "excel": native.get("value")})
                elif index >= len(source_values) or not _matches(source_values[index], age_values[index], 0, 0):
                    mismatches.append({"kind": "external_age_axis_source_value", "address": address,
                                       "source": source_values[index] if index < len(source_values) else None,
                                       "captured": age_values[index]})
                else:
                    metrics["external_age_axis_keys"] += 1
    expected_coordinates = manifest.get("external_input_count", 0)
    expected_age_keys = len(age_values) if isinstance(age_values, list) else 0
    if metrics["external_input_coordinates"] != expected_coordinates:
        mismatches.append({"kind": "external_input_coordinate_coverage",
                           "expected": expected_coordinates,
                           "actual": metrics["external_input_coordinates"]})
    if metrics["external_formula_cut_members"] != expected_cut_count:
        mismatches.append({"kind": "external_formula_cut_coverage",
                           "expected": expected_cut_count,
                           "actual": metrics["external_formula_cut_members"]})
    if metrics["external_age_axis_keys"] != expected_age_keys:
        mismatches.append({"kind": "external_age_axis_coverage",
                           "expected": expected_age_keys,
                           "actual": metrics["external_age_axis_keys"]})
    return metrics, mismatches


def reconcile(workflow_dir: Path, validation_path: Path, oracle_path: Path,
              abs_tol: float = 1e-12, rel_tol: float = 1e-12) -> dict[str, Any]:
    """Write the Stage 5 comparison report from standalone output and a valid oracle."""
    try:
        root, workflow = load_workflow(workflow_dir)
        workflow_manifest, stage4_revision, report, bundle = _stage4(root)
        stage5_root = (root / "stage5").resolve()
        validation_path, oracle_path = validation_path.resolve(), oracle_path.resolve()
        if not _inside(validation_path, stage5_root) or not _inside(oracle_path, stage5_root):
            raise ValueError("validation and oracle evidence must be inside this workflow's Stage 5 directory")
        validation = read_json(validation_path)
        if (validation.get("schema_version") != "step5.execution_validation.v1" or validation.get("status") != "pass"
                or validation.get("stage4_artifact_sha256") != stage4_revision["artifact"]["json_sha256"]
                or validation.get("source_sha256") != workflow["source"].get("workbook_sha256")):
            raise ValueError("execution validation is incomplete or bound to another source/Step 4 artifact")
        result_path = Path(validation["result_path"]).resolve()
        if not _inside(result_path, stage5_root) or not result_path.is_file() or hash_file(result_path) != validation.get("result_sha256"):
            raise ValueError("standalone model result is missing or changed")
        model_result = read_json(result_path)
        oracle = _preflight_oracle(read_json(oracle_path), workflow["source"]["workbook_sha256"],
                                  stage4_revision["artifact"]["json_sha256"])
        mismatches: list[dict[str, Any]] = []
        compared_targets = 0
        target_results: dict[str, Any] = {}
        expected_targets = set(report.get("target_names", []))
        native_targets = oracle.get("named_targets", {})
        actual_targets = model_result.get("targets", {})
        for name in sorted(expected_targets):
            native = native_targets.get(name)
            if not isinstance(native, dict) or "value" not in native:
                mismatches.append({"kind": "target_missing_from_oracle", "target": name})
                continue
            if name not in actual_targets:
                mismatches.append({"kind": "target_missing_from_model", "target": name})
                continue
            compared_targets += 1
            excel_value = native.get("excel_error") if native.get("excel_error") is not None else native.get("value")
            matched = _matches(actual_targets[name], excel_value, abs_tol, rel_tol)
            target_results[name] = {"model": actual_targets[name], "excel": excel_value, "matched": matched}
            if not matched:
                mismatches.append({"kind": "target_value", "target": name,
                                   "model": actual_targets[name], "excel": excel_value})
        primary_mismatches = []
        raw = read_json(bundle / "source_values.json")
        for address in sorted(_PRIMARY_ADDRESSES):
            key = f"main!{address}"
            native_value = oracle["primary_inputs"][address]["value"]
            if key not in raw or not _matches(raw.get(key), native_value, 0, 0):
                primary_mismatches.append({"address": f"Main!{address}", "model_source": raw.get(key), "excel": native_value})
        if primary_mismatches:
            mismatches.extend({"kind": "primary_input", **item} for item in primary_mismatches)

        generated_manifest = read_json(bundle / "model_manifest.json")
        native_formula_values, comparison_ranges = _oracle_range_cells(oracle)
        source_workbook = Path(workflow["source"]["workbook_path"]).resolve()
        active_formula_addresses, expected_formula_texts = _active_formula_details(
            report, generated_manifest, source_workbook, workflow["source"]["workbook_sha256"])
        external_metrics, external_mismatches = _compare_external_boundary(
            generated_manifest, bundle, model_result, native_formula_values, abs_tol, rel_tol,
            source_workbook)
        mismatches.extend(external_mismatches)
        formula_values = model_result.get("cells", {})
        if not isinstance(formula_values, dict):
            raise ValueError("standalone model result must include a source-addressed cells object")
        model_formula_values: dict[str, Any] = {}
        for address, value in formula_values.items():
            if not isinstance(address, str):
                raise ValueError("standalone model cells must use qualified source addresses")
            key = _address_key(address)
            if key in model_formula_values:
                raise ValueError(f"standalone model result contains duplicate cell addresses: {address}")
            model_formula_values[key] = value
        (compared_formulas, formula_identity_matches, max_abs_formula_difference,
         compared_formula_addresses, formula_mismatches) = _compare_active_formula_members(
            active_formula_addresses, native_formula_values, model_formula_values,
            expected_formula_texts, abs_tol, rel_tol)
        mismatches.extend(formula_mismatches)
        status = "ready_for_review" if not mismatches and compared_targets == len(expected_targets) else "not_ready"
        stage3_revision = current_revision(workflow_manifest, 3)
        if stage3_revision is None:
            raise ValueError("approved Stage 3 design artifact is missing")
        stage3_payload = read_json(root / stage3_revision["artifact"]["json"])
        design = stage3_payload.get("design", {})
        trace_path = Path(report.get("active_trace", {}).get("path", ""))
        active_trace = read_json(trace_path) if trace_path.is_file() else {}
        stage4_paths = report.get("path_coverage", {})
        implementation_evidence = stage4_paths.get("implementation_evidence", [])
        runtime_evidence: list[dict[str, Any]] = []
        for item in implementation_evidence:
            path_id, address, source_formula = item.get("path_id"), item.get("address"), item.get("source_formula")
            if (not isinstance(path_id, str) or not isinstance(address, str) or not isinstance(source_formula, str)
                    or path_id not in stage4_paths.get("numerical_path_ids", [])):
                continue
            address_key = _address_key(address)
            trace_match = any(cell.get("address", "").casefold() == address.casefold()
                              and cell.get("formula") == source_formula
                              for cell in active_trace.get("cells", []) if isinstance(cell, dict))
            generated_addresses = generated_manifest.get("active_formula_addresses",
                                                        generated_manifest.get("formula_addresses", []))
            generated_match = isinstance(generated_addresses, list) and address_key in {
                _address_key(value) for value in generated_addresses if isinstance(value, str)}
            if not trace_match or not generated_match or address_key not in model_formula_values:
                continue
            if address_key in compared_formula_addresses:
                comparison = {"kind": "native_formula_range", "address": address}
            else:
                aliases = [entry.get("name") for entry in active_trace.get("names", [])
                           if isinstance(entry, dict) and isinstance(entry.get("destination"), str)
                           and entry["destination"].casefold() == address.casefold()]
                target_alias = next((name for name in aliases
                                     if isinstance(name, str) and name in target_results
                                     and target_results[name].get("matched") is True), None)
                if target_alias is None:
                    continue
                comparison = {"kind": "native_named_target", "target": target_alias, "address": address}
            if status == "ready_for_review":
                runtime_evidence.append({"path_id": path_id, "status": "matched", "address": address,
                                         "source_formula": source_formula, "comparison": comparison})
        condition_ids = set(stage4_paths.get("condition_only_path_ids", []))
        condition_evidence = []
        if status == "ready_for_review" and report.get("active_trace", {}).get("source") == "stage4_discovery":
            for path in design.get("paths", []):
                if isinstance(path, dict) and path.get("path_id") in condition_ids:
                    condition_evidence.append({"path_id": path["path_id"],
                                               "status": "inactive_condition_checked_no_solver",
                                               "evidence": "fresh active trace completed and the approved GP targets reconciled; no feedback solver was generated"})
        path_coverage = {"known_path_ids": stage4_paths.get("known_path_ids", []),
                         "planned_path_ids": stage4_paths.get("planned_path_ids", []),
                         "numerical_path_ids": stage4_paths.get("numerical_path_ids", []),
                         "condition_only_path_ids": sorted(condition_ids),
                         "implemented_path_ids": stage4_paths.get("implemented_path_ids", []),
                         "runtime_verified_path_ids": sorted({item["path_id"] for item in runtime_evidence}),
                         "implementation_evidence": implementation_evidence,
                         "runtime_evidence": runtime_evidence,
                         "condition_evidence": condition_evidence}
        delegation = workflow_manifest.get("review_delegation")
        delegated = (isinstance(delegation, dict)
                     and isinstance(delegation.get("stages"), list)
                     and 5 in delegation["stages"])
        review_policy = {"required_reviewer_pair": ["agent", "typesafe"] if delegated else ["agent", "human"],
                         "delegated_typesafe_active": delegated,
                         "delegation": ({"authorization_sha256": delegation.get("authorization_sha256"),
                                         "source_sha256": delegation.get("source_sha256"),
                                         "stages": delegation.get("stages"),
                                         "requested_model": delegation.get("requested_model"),
                                         "minimum_approval_probability": delegation.get("minimum_approval_probability")}
                                        if delegated else None)}
        payload = {"schema_version": "step5.validation_report.v1", "tool": "step5.reconcile",
                   "status": status, "source_sha256": workflow["source"]["workbook_sha256"],
                   "stage4_artifact_sha256": stage4_revision["artifact"]["json_sha256"],
                   "review_policy": review_policy,
                   "standalone_validation": {"path": str(validation_path), "sha256": hash_file(validation_path),
                                             "result_sha256": validation["result_sha256"]},
                   "oracle": {"path": str(oracle_path), "sha256": hash_file(oracle_path),
                              "engine": oracle.get("engine"), "engine_settings": oracle.get("engine_settings"),
                              "source_copy_sha256": oracle.get("source_copy_sha256"),
                              "macros_executed": oracle.get("macros_executed"),
                              "input_overrides": oracle.get("input_overrides")},
                   "tolerances": {"absolute": abs_tol, "relative": rel_tol},
                   "target_results": target_results,
                   "path_coverage": path_coverage,
                   "compared": {"named_targets": compared_targets, "required_targets": len(expected_targets),
                                "primary_inputs": len(_PRIMARY_ADDRESSES) - len(primary_mismatches),
                                "active_formula_members": compared_formulas,
                                "required_active_formula_members": len(active_formula_addresses),
                                "active_formula_identity_matches": formula_identity_matches,
                                "required_active_formula_identities": len(active_formula_addresses),
                                "active_formula_max_abs_difference": max_abs_formula_difference,
                                **external_metrics,
                                "required_external_input_coordinates": generated_manifest.get("external_input_count", 0),
                                "required_external_formula_cut_members": generated_manifest.get(
                                    "external_formula_cut_member_count", 0),
                                "required_external_age_axis_keys": len(
                                    read_json(bundle / "external_values.json").get("age_axis", {}).get("values", []))
                                    if (bundle / "external_values.json").is_file() else 0,
                                "active_premium_formula_cells": sum(
                                    1 for item in active_formula_addresses if item.split("!", 1)[0].casefold() == "premium"),
                                "comparison_ranges": comparison_ranges},
                   "mismatch_count": len(mismatches), "mismatch_examples": mismatches[:50],
                   "formula_cache_inputs": False, "standalone_project_imports": False}
        next_revision = len(workflow["stages"]["5"]["revisions"]) + 1
        revision_dir = root / "stage5" / f"revision-{next_revision:04d}"
        if revision_dir.exists():
            raise ValueError("Stage 5 report output already exists; preserve it and create a new workflow revision")
        markdown = _reconciliation_markdown(payload, design=design)
        revision = append_stage_artifact(root, 5, "validation_report.json", payload, markdown,
                                         input_files={"standalone_validation": validation_path,
                                                      "standalone_result": result_path,
                                                      "native_oracle": oracle_path,
                                                      "source_workbook": source_workbook,
                                                      **{f"generated:{name}": bundle / name for name in report["bundle"]["files"]}},
                                         input_stages={4: stage4_revision["artifact"]["json_sha256"]})
        return {"schema_version": payload["schema_version"], "tool": payload["tool"],
                "status": "pass" if status == "ready_for_review" else "fail",
                "workflow": str(root), "revision": revision["revision"], "artifact": revision["artifact"],
                "compared": payload["compared"], "mismatch_count": payload["mismatch_count"]}
    except Exception as exc:
        return {"schema_version": "step5.validation_report.v1", "tool": "step5.reconcile",
                "status": "blocked", "reason": str(exc)}


def _reconciliation_evidence_markdown(value: dict[str, Any], design: dict[str, Any] | None = None) -> str:
    compared = value["compared"]
    lines = ["# Step 5 Validation and Reconciliation", "",
             f"Status: **{'Ready for review' if value['status'] == 'ready_for_review' else 'Not ready'}**", "",
             f"Source SHA-256: `{value['source_sha256']}`",
             f"Approved Step 4 artifact: `{value['stage4_artifact_sha256']}`", "",
             "## Standalone execution and native Excel comparison", "",
             "- Generated code compile and isolated execution: passed.",
             f"- Microsoft Excel engine: `{value['oracle']['engine']}`",
             f"- Private copy SHA-256: `{value['oracle']['source_copy_sha256']}`",
             f"- Macros executed: `{value['oracle']['macros_executed']}`; input overrides: `{value['oracle']['input_overrides']}`",
             f"- Primary inputs matched: {compared['primary_inputs']} / 7",
             f"- Named targets matched: {compared['named_targets']} / {compared['required_targets']}",
             f"- Active source formula members compared: {compared['active_formula_members']} / {compared['required_active_formula_members']}",
             f"- Active source formula identities matched: {compared.get('active_formula_identity_matches', 'unknown')} / {compared.get('required_active_formula_identities', 'unknown')}",
             f"- Maximum absolute difference across numeric active formula values: {compared.get('active_formula_max_abs_difference') if compared.get('active_formula_max_abs_difference') is not None else 'unknown'}",
             f"- External captured coordinates compared: {compared.get('external_input_coordinates', 0)} / {compared.get('required_external_input_coordinates', 0)}; source formula cuts checked: {compared.get('external_formula_cut_members', 0)} / {compared.get('required_external_formula_cut_members', 0)}; key-axis values checked: {compared.get('external_age_axis_keys', 0)} / {compared.get('required_external_age_axis_keys', 0)}.",
             f"- Active Premium formula cells compared: {compared['active_premium_formula_cells']}",
             f"- Tolerances: abs={value['tolerances']['absolute']}, rel={value['tolerances']['relative']}",
             f"- Mismatches: {value['mismatch_count']}", ""]
    if value.get("target_results"):
        requested = {item.get("selector") for item in (design or {}).get("targets", []) if isinstance(item, dict)}
        diagnostics = {name for item in (design or {}).get("targets", []) if isinstance(item, dict)
                       for name in item.get("diagnostic_intermediates", []) if isinstance(name, str)}
        lines.extend(["## Native named result checks", ""])
        for name, record in sorted(value["target_results"].items()):
            kind = ("requested design target" if name in requested else
                    "diagnostic value" if name in diagnostics else
                    "classification not recorded")
            lines.append(f"- `{name}` ({kind}): model `{record.get('model')}`, Excel `{record.get('excel')}`, matched `{record.get('matched')}`")
    coverage = value.get("path_coverage", {})
    lines.extend(["", "## Generated path evidence", "",
                  f"- Implemented numerical routes: {len(coverage.get('implemented_path_ids', []))} / {len(coverage.get('numerical_path_ids', []))}.",
                  f"- Runtime-verified numerical routes: {len(coverage.get('runtime_verified_path_ids', []))} / {len(coverage.get('numerical_path_ids', []))}.",
                  f"- Inactive condition-only checks: {len(coverage.get('condition_evidence', []))} / {len(coverage.get('condition_only_path_ids', []))}; no feedback solver is included."])
    for item in coverage.get("runtime_evidence", []):
        lines.append(f"- `{item['path_id']}`: `{item['address']}` · {item['comparison']['kind']} · `{item['source_formula']}`.")
    for item in coverage.get("condition_evidence", []):
        lines.append(f"- `{item['path_id']}`: {item['evidence']}.")
    if value["mismatch_examples"]:
        lines.extend(["## Mismatch examples", ""])
        lines.extend(f"- `{json.dumps(item, ensure_ascii=False, sort_keys=True)}`" for item in value["mismatch_examples"])
    pair = value.get("review_policy", {}).get("required_reviewer_pair", ["agent", "human"])
    reviewer_text = "Agent and delegated TypeSafe" if pair == ["agent", "typesafe"] else "Agent and human"
    delegation = value.get("review_policy", {}).get("delegation")
    if isinstance(delegation, dict):
        lines.append(f"TypeSafe delegation: stages {delegation.get('stages')}; model `{delegation.get('requested_model')}`; authorization SHA-256 `{delegation.get('authorization_sha256')}`.")
    lines.extend(["", f"Review pair for this revision: {reviewer_text}. Current decisions are in workflow.json.", ""])
    return "\n".join(lines)


def _reconciliation_markdown(
    value: dict[str, Any], *, design: dict[str, Any] | None = None,
    machine_link: str = "validation_report.json", workflow_link: str = "../../workflow.json",
) -> str:
    design = design if isinstance(design, dict) else {}
    compared = value.get("compared")
    compared = compared if isinstance(compared, dict) else {}
    raw_results = value.get("target_results")
    target_results = raw_results if isinstance(raw_results, dict) else {}
    raw_design_targets = design.get("targets")
    design_targets = raw_design_targets if isinstance(raw_design_targets, list) else []
    declared_targets = [item for item in design_targets if isinstance(item, dict)]
    target_by_name = {
        item.get("selector"): item for item in declared_targets if isinstance(item.get("selector"), str)
    }
    requested_names = set(target_by_name)
    diagnostic_names = {
        name for item in declared_targets
        for name in (item.get("diagnostic_intermediates") if isinstance(item.get("diagnostic_intermediates"), list) else [])
        if isinstance(name, str)
    }
    requested_results = {name: record for name, record in target_results.items() if name in requested_names}
    diagnostic_results = {name: record for name, record in target_results.items() if name in diagnostic_names}
    unclassified_results = {name: record for name, record in target_results.items()
                            if name not in requested_names and name not in diagnostic_names}

    def result_count(records: dict[str, Any]) -> str:
        if not records:
            return "Not recorded" if not target_results else "0 / 0"
        matched = sum(isinstance(record, dict) and record.get("matched") is True for record in records.values())
        return f"{matched} / {len(records)} matched"

    def cell(value_item: Any) -> str:
        text = display_report_value(value_item).replace("|", "\\|")
        return text.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")

    def readable_value(item: Any, limit: int = 12) -> tuple[str, bool]:
        if not isinstance(item, (list, dict)):
            return cell(item), False
        leaves: list[tuple[tuple[Any, ...], Any]] = []

        def visit(current: Any, path: tuple[Any, ...]) -> None:
            if len(leaves) > limit:
                return
            if isinstance(current, list):
                if not current:
                    leaves.append((path, "[]"))
                else:
                    for index, child in enumerate(current):
                        visit(child, (*path, index))
                        if len(leaves) > limit:
                            return
            elif isinstance(current, dict):
                if not current:
                    leaves.append((path, "{}"))
                else:
                    for key, child in current.items():
                        visit(child, (*path, key))
                        if len(leaves) > limit:
                            return
            else:
                leaves.append((path, current))

        visit(item, ())
        rows = []
        for path, leaf in leaves[:limit]:
            label = f"position {list(path)}" if all(isinstance(part, int) for part in path) else f"field path {list(path)}"
            rows.append(f"{label} = {cell(leaf)}")
        if len(leaves) > limit:
            rows.append("Additional values omitted; see linked machine report")
        return "<br>".join(rows), True

    def axis_summary(target: dict[str, Any] | None) -> str:
        if target is None:
            return "Not recorded"
        axes = target.get("axes", target.get("axis"))
        if axes is None:
            return "Not recorded"
        if not isinstance(axes, list):
            return cell(axes)
        if not axes:
            return "[]"
        summaries = []
        for axis in axes:
            if not isinstance(axis, dict):
                summaries.append(cell(axis))
                continue
            details = [f"{key}: {cell(axis[key])}" for key in ("name", "axis_id", "role", "units", "source_range") if key in axis]
            if "keys" in axis:
                keys_text, _ = readable_value(axis["keys"], limit=6)
                details.append(f"keys: {keys_text}")
            summaries.append("; ".join(details) or "Not recorded")
        return " / ".join(summaries)

    path_coverage = value.get("path_coverage")
    path_coverage = path_coverage if isinstance(path_coverage, dict) else {}
    numerical_ids = path_coverage.get("numerical_path_ids")
    runtime_ids = path_coverage.get("runtime_verified_path_ids")
    condition_ids = path_coverage.get("condition_only_path_ids")
    condition_evidence = path_coverage.get("condition_evidence")
    status = value.get("status")
    is_ready = status == "ready_for_review"
    readiness = "ready for review" if is_ready else "not ready" if isinstance(status, str) else "Not recorded"
    maximum_difference = compared.get("active_formula_max_abs_difference")
    review_policy = value.get("review_policy")
    review_policy = review_policy if isinstance(review_policy, dict) else {}
    required_reviewers = review_policy.get("required_reviewer_pair", ["agent", "human"])
    if required_reviewers == ["agent", "typesafe"]:
        next_review = (
            "- Next: Agent PASS is followed by the explicitly authorized TypeSafe decision on this exact pair. "
            "TypeSafe is not a human receipt. Agent and delegated TypeSafe confirmation are pending."
        )
    else:
        next_review = (
            "- Next: the Agent reviews this exact JSON/Markdown pair first. After Agent PASS, the actual human reviews. "
            "Technical reconciliation readiness and Agent PASS are not human acceptance."
        )
    lines = [
        "# Step 5 Validation and Reconciliation", "",
        "## What this report asks you to accept", "",
        "The standalone model's comparison with the bound native Excel calculation for the declared source and scenario, including the reported input, formula and result checks.", "",
        "## Purpose and scope", "",
        f"Technical reconciliation: **{readiness}**. "
        f"Mismatch count: {cell(value.get('mismatch_count'))}.",
        "The Excel oracle used a private copy, disabled macros, and recorded its calculation settings and input overrides in the evidence appendix.", "",
        "## Verified results", "",
        f"- Requested design targets: {result_count(requested_results)}.",
        f"- Diagnostic values: {result_count(diagnostic_results)}.",
    ]
    if unclassified_results:
        lines.append(f"- Other named checks (classification not recorded): {result_count(unclassified_results)}.")

    result_order: list[str] = []
    for target in sorted(
        declared_targets,
        key=lambda item: item.get("result_order") if isinstance(item.get("result_order"), int)
        and not isinstance(item.get("result_order"), bool) else 10**9,
    ):
        selector = target.get("selector")
        if isinstance(selector, str) and selector in target_results and selector not in result_order:
            result_order.append(selector)
    for target in declared_targets:
        names = target.get("diagnostic_intermediates")
        if isinstance(names, list):
            for name in names:
                if isinstance(name, str) and name in target_results and name not in result_order:
                    result_order.append(name)
    result_order.extend(name for name in target_results if name not in result_order)
    has_non_scalar = False
    if target_results:
        lines.extend(["", "### Named result comparison", "",
                      "Requested targets follow the Step 3 result order. Diagnostic and unclassified checks follow them.", "",
                      "| Result | Classification | Kind | Declared shape | Units | Axes | Model | Excel | Matched |",
                      "| --- | --- | --- | --- | --- | --- | --- | --- | :---: |"])
        for name in result_order:
            record = target_results[name]
            record = record if isinstance(record, dict) else {}
            target = target_by_name.get(name)
            classification = ("Requested target" if name in requested_names else
                              "Diagnostic" if name in diagnostic_names else
                              "Not classified in design")
            metadata = target if isinstance(target, dict) else None
            kind = cell(metadata.get("result_kind")) if metadata is not None else "Not recorded"
            shape = cell(metadata.get("shape")) if metadata is not None else "Not recorded"
            units = cell(metadata.get("units")) if metadata is not None else "Not recorded"
            axes = axis_summary(metadata)
            model, model_array = readable_value(record.get("model"))
            excel, excel_array = readable_value(record.get("excel"))
            has_non_scalar = has_non_scalar or model_array or excel_array
            lines.append(
                f"| {cell(name)} | {classification} | {kind} | {shape} | {units} | {axes} | "
                f"{model} | {excel} | {cell(record.get('matched'))} |"
            )
        if has_non_scalar:
            lines.extend(["", "Array and table values show zero-based positions or recorded field paths. Position labels do not imply a quarter, year, or other business-axis origin; only explicitly recorded axis keys provide semantic labels. The full Model and Excel values remain in the linked machine report."])
        else:
            lines.append("")
    tolerances = value.get("tolerances")
    if isinstance(tolerances, dict):
        lines.append(
            f"- Recorded result tolerances: absolute={cell(tolerances.get('absolute'))}, relative={cell(tolerances.get('relative'))}."
        )
    else:
        lines.append("- Recorded result tolerances: Not recorded.")
    formula_identities = compared.get("active_formula_identity_matches")
    required_identities = compared.get("required_active_formula_identities")
    lines.extend([
        f"- Primary inputs matched: {cell(compared.get('primary_inputs'))} / 7.",
        f"- Active formula identities matched: {cell(formula_identities)} / {cell(required_identities)}.",
        f"- External coordinates and formula cuts checked: {cell(compared.get('external_input_coordinates'))} / "
        f"{cell(compared.get('required_external_input_coordinates'))} coordinates; "
        f"{cell(compared.get('external_formula_cut_members'))} / "
        f"{cell(compared.get('required_external_formula_cut_members'))} formula cut members.",
        f"- Numerical routes runtime-verified: {cell(len(runtime_ids) if isinstance(runtime_ids, list) else None)} / "
        f"{cell(len(numerical_ids) if isinstance(numerical_ids, list) else None)}.",
    ])
    if isinstance(condition_ids, list) and condition_ids:
        if isinstance(condition_evidence, list) and condition_evidence:
            lines.append(f"- Condition-only inactive checks: {len(condition_evidence)} / {len(condition_ids)}; no feedback solver is included.")
        else:
            lines.append(f"- Condition-only routes: {len(condition_ids)}; inactive-condition evidence is not recorded here and they are not counted as numerical routes.")
    lines.append(
        f"- Maximum absolute difference across compared active formula values: {cell(maximum_difference)}."
    )
    if value.get("mismatch_count"):
        examples = value.get("mismatch_examples")
        examples = examples if isinstance(examples, list) else []
        lines.append(f"- Mismatch examples: {', '.join(str(item.get('kind', 'unclassified')) for item in examples[:5] if isinstance(item, dict)) or 'Not recorded'}.")
    lines.extend(["", "## Limits and unresolved items", "",
                  "This report describes one bound workbook and scenario. The Stage 4 generated-code smoke is not an Excel comparison; these native Excel checks are separate. Route and result coverage does not establish all-configuration behavior or an unimplemented solver."])
    lines.extend(design_question_lines(design.get("open_questions")))
    lines.extend([
        "", "## Handover and review", "",
        f"- Machine reconciliation: [validation_report.json](<{machine_link}>).",
        f"- Current acceptance ledger: [workflow.json](<{workflow_link}>).",
        next_review,
        "- To reject: route the decision to the responsible current or earlier Step, preserve this revision, refresh affected evidence and reports, then inspect the ledger before continuing.",
        "- The linked ledger records current formal acceptance separately from `ready_for_review` technical status.", "",
        "## Detailed comparison evidence", "",
        re.sub(r"(?m)^## ", "#### ", _reconciliation_evidence_markdown(value, design).replace(
            "# Step 5 Validation and Reconciliation", "### Original reconciliation details", 1)), "",
    ])
    return "\n".join(lines)

def tool_catalog() -> dict[str, Any]:
    from excel_to_act.steps.step5.tools import tool_catalog as catalog

    return catalog()
