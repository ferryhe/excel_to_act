"""Compile a checked Step 2 index into a reusable source reading package."""

from __future__ import annotations

import hashlib
import json
import re
import time
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter, range_boundaries

from excel_to_act.graph.builder import RegexFormulaGraphBuilder
from excel_to_act.schemas import (
    ActiveXEvents,
    CheckboxBindings,
    DependencyAuditSource,
    DependencySnapshotSource,
    DefinedNameLookup,
    FormulaGraph,
    GraphNode,
    GraphNodeKind,
    ReadingFileRef,
    ReadingLookup,
    ReadingSource,
    SourceLocation,
    Step2DependencyAudit,
    Step2DependencySnapshot,
    Step2Index,
    Step2ReadingManifest,
    WorkbookInventory,
    WorkbookManifest,
    ViewRecord,
)
from excel_to_act.steps.step2.workflow import _validate_saved_index
from excel_to_act.views import _range_bounds, _range_facts, _record, _stable_id, compile_views

PREPARE_COMPILER_VERSION = "step2.prepare.v4"
PREPARE_OPTIONS: dict[str, Any] = {"view_token_budget": 1200, "static_dependency_policy": "inbound-v1"}
_DEPENDENCY_POLICY = "retain_read_only_source_values_for_inbound_references"
_CELL_RANGE = re.compile(r"^(?:(?:'((?:[^']|'')+)'|([^!]+))!)?(\$?[A-Z]{1,3}\$?[1-9]\d*(?::\$?[A-Z]{1,3}\$?[1-9]\d*)?)$", re.IGNORECASE)
_ABSOLUTE_CELL_RANGE = re.compile(r"^\$[A-Z]{1,3}\$[1-9]\d*(?::\$[A-Z]{1,3}\$[1-9]\d*)?$", re.IGNORECASE)
_SIMPLE_STRUCTURED_REF = re.compile(r"^(?P<table>[A-Za-z_][A-Za-z0-9_.]*)\[(?P<selector>[^\[\]]+)\]$")


def _json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _file_sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _diagnostic(code: str, message: str, **details: Any) -> dict[str, Any]:
    return {"code": code, "severity": "error", "message": message, **details}


def _blocked(
    diagnostics: list[dict[str, Any]], *, source: dict[str, Any] | None = None,
    run_id: str | None = None, inventory_model_loads: int = 0,
) -> dict[str, Any]:
    return {
        "tool": "step2.prepare", "status": "blocked", "source": source, "run_id": run_id,
        "artifacts": {}, "planned_files": [], "diagnostics": diagnostics,
        "next_action": {"action": "review_diagnostics"},
        "metrics": {"inventory_model_loads": inventory_model_loads, "raw_workbook_parses": 0, "graph_builds": 0},
    }


def _source_ref(path: str, sha256: str, root: str) -> ReadingFileRef:
    return ReadingFileRef(root=root, path=path, sha256=sha256)


def _expected_view_bindings(source_id: str, inventory: WorkbookInventory) -> list[tuple[str, str, str, str | None]]:
    source_dir = f"sources/{source_id}/views"
    view_id = _stable_id(source_id, "workbook", "", "")
    result = [(f"{source_dir}/{view_id}.json", view_id, "workbook", None)]
    for sheet in inventory.sheets:
        view_id = _stable_id(source_id, "sheet", sheet.name, "")
        result.append((f"{source_dir}/{view_id}.json", view_id, "sheet", sheet.name))
        for region in sheet.ranges:
            if region.source_location.sheet_name != sheet.name or _range_bounds(region.address) is None:
                continue
            record = _record(source_id, "range", region.source_location, _range_facts(region))
            view_id = _stable_id(source_id, "region", sheet.name, record.record_id)
            result.append((f"{source_dir}/{view_id}.json", view_id, "region", sheet.name))
    return result


def _defined_name_lookup(inventory: WorkbookInventory, graph: FormulaGraph) -> list[DefinedNameLookup]:
    result = []
    for item in inventory.workbook_ranges:
        if item.kind != "defined_name" or not item.name:
            continue
        scope = str(item.metadata.get("scope", "workbook"))
        graph_node = next((
            node for node in graph.nodes
            if node.kind == GraphNodeKind.name
            and node.metadata.get("name") == item.name
            and node.metadata.get("scope", "workbook") == scope
            and node.source_location == item.source_location
        ), None)
        result.append(DefinedNameLookup(
            name=item.name, scope=scope, node_id=graph_node.id if graph_node else None,
            address=item.address, source_location=item.source_location,
        ))
    return result


def _scope_parts(scope: dict[str, Any], inventory: WorkbookInventory) -> tuple[list[str], list[str], list[dict[str, Any]]]:
    names = [sheet.name for sheet in inventory.sheets]
    ignored = scope.get("ignored_sheets", [])
    retained = scope.get("retained_sheets")
    if not isinstance(ignored, list) or any(not isinstance(name, str) for name in ignored):
        raise ValueError("scope_invalid: ignored_sheets must be a list of sheet names")
    if retained is not None and (not isinstance(retained, list) or any(not isinstance(name, str) for name in retained)):
        raise ValueError("scope_invalid: retained_sheets must be a list of sheet names")
    retained = list(retained) if retained is not None else [name for name in names if name not in ignored]
    if len(set(ignored)) != len(ignored) or len(set(retained)) != len(retained):
        raise ValueError("scope_sheet_conflict: sheet decisions contain duplicate names")
    unknown = sorted((set(ignored) | set(retained)) - set(names))
    if unknown:
        raise ValueError(f"scope_sheet_unknown: unknown workbook sheet(s): {', '.join(unknown)}")
    overlap = sorted(set(ignored) & set(retained))
    if overlap:
        raise ValueError(f"scope_sheet_conflict: sheets are both retained and excluded: {', '.join(overlap)}")
    if set(ignored) | set(retained) != set(names):
        raise ValueError("scope_sheet_conflict: retained_sheets and ignored_sheets must decide every source sheet")
    dependency_policy = scope.get("dependency_policy", {})
    if not isinstance(dependency_policy, dict):
        raise ValueError("scope_invalid: dependency_policy must be an object")
    if ignored and dependency_policy.get("mode") != _DEPENDENCY_POLICY:
        raise ValueError(f"scope_dependency_policy_invalid: excluded sheets require dependency mode {_DEPENDENCY_POLICY!r}")
    confirmation = scope.get("confirmation", {})
    decisions = confirmation.get("decisions", []) if isinstance(confirmation, dict) else []
    if not isinstance(decisions, list):
        raise ValueError("scope_invalid: confirmation.decisions must be a list")
    sheet_decisions: dict[str, list[Any]] = {}
    for item in decisions:
        if not isinstance(item, dict):
            continue
        question = item.get("question_id")
        if isinstance(question, str) and question.startswith("analysis_scope:ignore_sheet:"):
            sheet_decisions.setdefault(question, []).append(item.get("decision"))
    for question, values in sheet_decisions.items():
        if any(value != values[0] for value in values[1:]):
            raise ValueError(f"scope_sheet_conflict: confirmation question {question!r} has conflicting decisions")
        sheet = question.removeprefix("analysis_scope:ignore_sheet:")
        if sheet not in names:
            raise ValueError(f"scope_sheet_unknown: confirmation references unknown sheet {sheet!r}")
        decision = values[0]
        if sheet in retained and decision == "ignore_internal_logic_keep_dependency_values":
            raise ValueError(f"scope_sheet_conflict: retained sheet {sheet!r} has an ignore confirmation decision")
    for sheet in ignored:
        question = f"analysis_scope:ignore_sheet:{sheet}"
        values = sheet_decisions.get(question, [])
        if not values or values[0] != "ignore_internal_logic_keep_dependency_values":
            raise ValueError(f"scope_decision_missing: no matching ignore decision for sheet {sheet!r}")
    return retained, ignored, decisions


def _default_scope(entry: Any, inventory: WorkbookInventory) -> dict[str, Any]:
    return {
        "schema_version": "analysis.scope.v1",
        "decision_source": "default_all_sheets",
        "workbook_sha256": entry.source_sha256,
        "source_run_id": entry.run_id,
        "retained_sheets": [sheet.name for sheet in inventory.sheets],
        "ignored_sheets": [],
        "confirmation": {"decisions": []},
    }


def _scope_error(scope: Any, entries: list[Any], inventories: dict[str, WorkbookInventory]) -> tuple[str | None, list[str], list[str], list[dict[str, Any]]]:
    if not isinstance(scope, dict) or scope.get("schema_version") != "analysis.scope.v1":
        return None, [], [], [_diagnostic("scope_invalid", "Scope must be a JSON analysis.scope.v1 object.")]
    source_sha = scope.get("workbook_sha256")
    run_id = scope.get("source_run_id")
    sha_matches = [entry for entry in entries if entry.source_sha256 == source_sha]
    if not sha_matches:
        return None, [], [], [_diagnostic("scope_source_mismatch", "Scope workbook_sha256 does not match any source in the native index.")]
    matches = [entry for entry in sha_matches if entry.run_id == run_id]
    if not matches:
        return None, [], [], [_diagnostic("scope_run_mismatch", "Scope source_run_id does not match the selected source run.")]
    if len(matches) != 1:
        return None, [], [], [_diagnostic("scope_identity_ambiguous", "Scope identity matches more than one native index entry.")]
    entry = matches[0]
    inventory = inventories.get(entry.source_id)
    if inventory is None:
        return entry.source_id, [], [], [_diagnostic("scope_source_unavailable", "The scoped source has no validated workbook inventory.", source_id=entry.source_id)]
    try:
        retained, ignored, _ = _scope_parts(scope, inventory)
    except ValueError as exc:
        code, _, message = str(exc).partition(": ")
        return entry.source_id, [], [], [_diagnostic(code or "scope_invalid", message or str(exc), source_id=entry.source_id)]
    return entry.source_id, retained, ignored, []


def _parse_destination(address: str, default_sheet: str | None) -> tuple[str | None, str | None]:
    text = address.strip().lstrip("=").strip()
    match = _CELL_RANGE.fullmatch(text)
    if not match or not _ABSOLUTE_CELL_RANGE.fullmatch(match.group(3)):
        return None, None
    sheet = match.group(1).replace("''", "'") if match.group(1) is not None else match.group(2)
    return (sheet or default_sheet), match.group(3).replace("$", "").upper()


def _destination(node: GraphNode, source_location: SourceLocation | None = None) -> tuple[str | None, str | None]:
    location = node.source_location
    if node.kind in {GraphNodeKind.cell, GraphNodeKind.range} and location:
        if node.metadata.get("structured_reference"):
            table_address = node.metadata.get("table_ref")
            sheet = node.metadata.get("sheet") or location.sheet_name
            operand = node.metadata.get("structured_reference")
            table_name = node.metadata.get("table")
            if not isinstance(operand, str) or not isinstance(table_name, str) or not isinstance(table_address, str) or not isinstance(sheet, str):
                return None, None
            match = _SIMPLE_STRUCTURED_REF.fullmatch(operand)
            if not match or match.group("table").casefold() != table_name.casefold():
                return None, None
            selector = match.group("selector").strip()
            selector_key = selector.casefold()
            try:
                min_col, min_row, max_col, max_row = range_boundaries(table_address)
            except ValueError:
                return None, None
            header_rows = node.metadata.get("table_header_row_count")
            totals_rows = node.metadata.get("table_totals_row_count")
            if type(header_rows) is not int or type(totals_rows) is not int or header_rows not in {0, 1} or totals_rows not in {0, 1}:
                return None, None
            if node.metadata.get("table_totals_row_shown") is True and totals_rows == 0:
                return None, None
            data_start = min_row + header_rows
            data_end = max_row - totals_rows
            headers = node.metadata.get("table_columns")
            metadata_column = node.metadata.get("column")
            if selector_key in {"#data", "#headers"}:
                if metadata_column is not None:
                    return None, None
                if selector_key == "#headers":
                    if header_rows == 0:
                        return None, None
                    header_end = min_row + header_rows - 1
                    return sheet, f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{header_end}"
                if data_start > data_end:
                    return None, None
                return sheet, f"{get_column_letter(min_col)}{data_start}:{get_column_letter(max_col)}{data_end}"
            current_row = selector.startswith("@")
            column = selector[1:] if current_row else selector
            if not column or column.startswith("#") or (metadata_column is not None and str(metadata_column).casefold() != column.casefold()):
                return None, None
            if not isinstance(headers, list) or len(headers) != max_col - min_col + 1 or not headers:
                return None, None
            try:
                offset = next(i for i, header in enumerate(headers) if str(header).casefold() == column.casefold())
            except StopIteration:
                return None, None
            col = get_column_letter(min_col + offset)
            if current_row:
                if source_location is None or not source_location.address:
                    return None, None
                try:
                    source_min_col, source_min_row, source_max_col, source_max_row = range_boundaries(source_location.address)
                except ValueError:
                    return None, None
                if source_min_col != source_max_col or source_min_row != source_max_row or not data_start <= source_min_row <= data_end:
                    return None, None
                return sheet, f"{col}{source_min_row}"
            if data_start > data_end:
                return None, None
            return sheet, f"{col}{data_start}:{col}{data_end}"
        return location.sheet_name, location.address
    if node.kind == GraphNodeKind.name and isinstance(node.metadata.get("address"), str):
        return _parse_destination(node.metadata["address"], location.sheet_name if location else None)
    return None, None


def _dependency_data(
    entry: Any,
    inventory: WorkbookInventory,
    graph: FormulaGraph,
    retained: list[str],
    excluded: list[str],
    scope_sha256: str,
) -> tuple[DependencyAuditSource, DependencySnapshotSource]:
    nodes = {node.id: node for node in graph.nodes}
    retained_keys, excluded_keys = {name.casefold() for name in retained}, {name.casefold() for name in excluded}
    references: list[dict[str, Any]] = []
    range_keys: set[tuple[str, str]] = set()
    unresolved = [
        (feature.source_location, feature.description)
        for feature in graph.unsupported_features
        if feature.source_location.sheet_name and feature.source_location.sheet_name.casefold() in retained_keys
    ]
    for edge in graph.edges:
        source, target = nodes.get(edge.source), nodes.get(edge.target)
        source_sheet = source.source_location.sheet_name if source and source.source_location else None
        if not source_sheet or source_sheet.casefold() not in retained_keys or target is None:
            continue
        if target.kind == GraphNodeKind.name:
            name = target.metadata.get("name")
            name_scope = str(target.metadata.get("scope", "workbook"))
            declarations = [
                item for item in inventory.workbook_ranges
                if item.kind == "defined_name"
                and (item.name or "").casefold() == str(name or "").casefold()
                and str(item.metadata.get("scope", "workbook")).casefold() == name_scope.casefold()
            ]
            if len(declarations) != 1:
                if edge.source_location is not None:
                    unresolved.append((edge.source_location, f"Defined name {name!r} has {len(declarations)} declarations in scope {name_scope!r}; its full destination is unresolved"))
                continue
            declaration = declarations[0]
            if declaration.address != target.metadata.get("address") or declaration.source_location != target.source_location:
                if edge.source_location is not None:
                    unresolved.append((edge.source_location, f"Defined name {name!r} graph projection does not match its validated declaration"))
                continue
        target_sheet, address = _destination(target, edge.source_location)
        if target.metadata.get("structured_reference") and (not target_sheet or not address):
            if edge.source_location is not None:
                unresolved.append((edge.source_location, f"Structured reference {target.metadata['structured_reference']!r} is unsupported or cannot be mapped exactly"))
            continue
        if target.kind == GraphNodeKind.name and (not target_sheet or not address):
            if edge.source_location is not None:
                unresolved.append((edge.source_location, f"Defined-name destination {target.metadata.get('address', '')!r} is not an absolute A1 reference"))
            continue
        if not target_sheet or target_sheet.casefold() not in excluded_keys or not address:
            continue
        try:
            if any(value is None for value in range_boundaries(address)):
                continue
        except ValueError:
            if edge.source_location is not None:
                unresolved.append((edge.source_location, f"Cannot resolve static destination {address!r}"))
            continue
        key = (target_sheet, address)
        range_keys.add(key)
        references.append({
            "source": edge.source,
            "target": edge.target,
            "source_location": edge.source_location.model_dump(mode="json") if edge.source_location else None,
            "target_sheet": target_sheet,
            "target_address": address,
            "target_kind": target.kind.value,
            "formula": edge.formula,
        })
    ranges = [{"sheet_name": sheet, "address": address} for sheet, address in sorted(range_keys, key=lambda row: (row[0].casefold(), row[1]))]
    selected_cells: dict[tuple[str, str], Any] = {}
    sheet_lookup = {sheet.name.casefold(): sheet for sheet in inventory.sheets}
    for item in ranges:
        sheet = sheet_lookup[item["sheet_name"].casefold()]
        min_col, min_row, max_col, max_row = range_boundaries(item["address"])
        for cell in sheet.cells:
            if min_row <= cell.row <= max_row and min_col <= cell.column <= max_col:
                selected_cells[(sheet.name, cell.address)] = cell
    cells = [cell for _, cell in sorted(selected_cells.items(), key=lambda row: (row[1].source_location.sheet_index or 0, row[1].row, row[1].column))]
    diag = [{
        "code": "static_reference_unresolved",
        "severity": "warning",
        "message": description,
        "source_location": location.model_dump(mode="json"),
    } for location, description in unresolved]
    if excluded:
        diag.append({"code": "dynamic_reference_coverage_unknown", "severity": "warning", "message": "Static audit does not establish that dynamic or indirect dependencies are absent."})
    audit = DependencyAuditSource(
        source_id=entry.source_id,
        source_sha256=entry.source_sha256,
        run_id=entry.run_id,
        scope_sha256=scope_sha256,
        supported_static_inbound_reference_count=len(references),
        retained_dependency_range_count=len(ranges),
        unresolved_static_reference_count=len(unresolved),
        ranges=ranges,
        references=references,
        diagnostics=diag,
    )
    snapshot = DependencySnapshotSource(
        source_sha256=entry.source_sha256, run_id=entry.run_id, scope_sha256=scope_sha256,
        ranges=ranges, cells=cells,
    )
    return audit, snapshot


def _control_records(entry: Any, inventory: WorkbookInventory, artifact_models: dict[tuple[str, str], tuple[Path, Any]]) -> dict[str, list[ViewRecord]]:
    result: dict[str, list[ViewRecord]] = {}
    sheets = {sheet.name: sheet for sheet in inventory.sheets}
    refs = {Path(ref.name).name: ref for ref in entry.artifacts}
    for name, model_type, records, kind in (
        ("checkbox_bindings.json", CheckboxBindings, "bindings", "checkbox_binding"),
        ("activex_events.json", ActiveXEvents, "controls", "activex_event"),
    ):
        checked = artifact_models.get((entry.source_id, name))
        if checked is None or not isinstance(checked[1], model_type):
            continue
        artifact = checked[1]
        ref = refs[name]
        artifact_ref = {"root": "step1", "path": ref.path, "sha256": ref.sha256}
        for row in getattr(artifact, records):
            sheet = sheets.get(row.sheet)
            if sheet is None:
                raise ValueError(f"{name} has a control for unknown sheet {row.sheet!r}")
            identity = row.shape_id if getattr(row, "shape_id", None) else getattr(row, "control_name", "")
            location = SourceLocation(
                workbook_path=sheet.source_location.workbook_path,
                sheet_name=row.sheet,
                sheet_index=sheet.index,
                object_type=kind,
                object_id=f"{row.sheet_part}#{identity}",
                ooxml_part=row.sheet_part,
                source_identity=entry.source_id,
            )
            result.setdefault(row.sheet, []).append(_record(entry.source_id, kind, location, {
                "record": row.model_dump(mode="json"), "artifact_ref": artifact_ref,
            }))
    return result


def _md(value: Any) -> str:
    return str(value).replace("|", r"\|").replace("\r", " ").replace("\n", " ")


def _human_step1(data: dict[str, Any]) -> str:
    lines = ["# Step 1 source handoff", "", f"**Index status:** `{data['source_quality']['status']}`", "", "## Sources", "", "| Source | Run | Status | Diagnostics |", "| --- | --- | --- | --- |"]
    for item in data["sources"]:
        lines.append(f"| `{_md(item.get('source_path') or item['source_id'])}` | `{_md(item.get('run_id') or 'unavailable')}` | `{_md(item['status'])}` | {len(item['diagnostics'])} |")
    lines.extend(["", "Step 1 source quality and failure diagnostics are preserved from the validated native index.", ""])
    return "\n".join(lines)


def _human_step2(data: dict[str, Any]) -> str:
    quality, integrity, readiness, runtime = (data[key] for key in ("source_quality", "reference_integrity", "static_readiness", "runtime_and_numerical_behavior"))
    lines = [
        "# Step 2 prepared reading handoff", "",
        f"**Source quality:** `{quality['status']}` · **Reference integrity:** `{integrity['status']}` · **Static readiness:** `{readiness['status']}` · **Runtime/numerical behavior:** `{runtime['status']}`", "",
        "Source facts and stored values remain bound to the identified Step 1 run. Static references are incomplete where dynamic or unsupported formula references remain.", "",
        "## Prepared sources", "", "| Source | Run | Status | Ready | Diagnostics | Retained sheets | Excluded sheets | Static inbound refs | Ranges | Stored cells |", "| --- | --- | --- | --- | --- | ---: | ---: | ---: | ---: | ---: |",
    ]
    audit_by_id = {item["source_id"]: item for item in data["static_readiness"]["sources"]}
    for source in data["sources"]:
        audit = audit_by_id.get(source["source_id"], {})
        details = "; ".join(f"{item.get('code', 'diagnostic')}: {item.get('message', '')}" for item in source.get("diagnostics", [])) or "—"
        lines.append(f"| `{_md(source.get('source_path') or source['source_id'])}` | `{_md(source.get('run_id') or 'unavailable')}` | `{_md(source['status'])}` | {str(source['ready_for_next_step']).lower()} | {_md(details)} | {len(source['retained_sheets'])} | {len(source['excluded_sheets'])} | {audit.get('supported_static_inbound_reference_count', 0)} | {audit.get('retained_dependency_range_count', 0)} | {audit.get('stored_cell_count', 0)} |")
    lines.extend(["", "## Scope and limitations", "", "- Excluded sheet dependencies include only saved source cells reached by supported static inbound references.", "- Empty or missing cells are not synthesized; cached values are reported as stored and are not recalculated.", "- Runtime, macro, and changed-scenario numerical behavior have not been verified.", "", "## Next action", "", "Read this package's manifest and select only paths allowed by the source scope.", ""])
    return "\n".join(lines)


def _result_artifacts(out: Path, outputs: dict[str, str]) -> dict[str, dict[str, Any]]:
    names = set(outputs) | {"manifest.json"}
    result: dict[str, dict[str, Any]] = {}
    for name in sorted(names):
        path = out / name
        if path.is_file():
            data = path.read_bytes()
            result[name] = {"root": "reading", "path": name, "sha256": _sha(data), "bytes": len(data)}
    return result


def prepare(
    index_path: Path,
    step1_root: Path,
    out: Path,
    *,
    scope_path: Path | None = None,
    resume: bool = False,
    dry_run: bool = False,
) -> dict[str, Any]:
    """Prepare one validated Step 2 index without dispatching recovery tools."""
    started = time.perf_counter()
    index_path, root, out = index_path.expanduser().resolve(), step1_root.expanduser().resolve(), out.expanduser().resolve()
    try:
        index_bytes = index_path.read_bytes()
        index = Step2Index.model_validate_json(index_bytes)
    except (OSError, ValueError) as exc:
        return _blocked([_diagnostic("index_invalid", f"Cannot load native Step 2 index: {exc}")])
    artifact_models: dict[tuple[str, str], tuple[Path, Any]] = {}
    validation = _validate_saved_index(index_path, root, validated_artifacts=artifact_models)
    inventory_model_loads = validation["metrics"]["inventory_models_deserialized"]
    if validation["status"] != "pass":
        return _blocked(validation["diagnostics"], run_id=index.batch_id, inventory_model_loads=inventory_model_loads)

    inventories: dict[str, WorkbookInventory] = {}
    manifests: dict[str, WorkbookManifest] = {}
    usable: list[Any] = []
    diagnostics: list[dict[str, Any]] = []
    for entry in index.entries:
        if entry.status not in {"pass", "partial"}:
            continue
        inventory_item = artifact_models.get((entry.source_id, "inventory.json"))
        manifest_item = artifact_models.get((entry.source_id, "workbook_manifest.json"))
        if not inventory_item or not manifest_item or not isinstance(inventory_item[1], WorkbookInventory) or not isinstance(manifest_item[1], WorkbookManifest):
            diagnostics.append(_diagnostic("source_artifacts_missing", "Usable index entry needs validated inventory.json and workbook_manifest.json.", source_id=entry.source_id))
            continue
        inventory, manifest = inventory_item[1], manifest_item[1]
        if not entry.source_sha256 or not entry.run_id or inventory.workbook_sha256 != entry.source_sha256 or manifest.sha256 != entry.source_sha256:
            diagnostics.append(_diagnostic("source_identity_mismatch", "Source SHA/run identity is incomplete or differs from validated workbook artifacts.", source_id=entry.source_id))
            continue
        inventories[entry.source_id], manifests[entry.source_id] = inventory, manifest
        usable.append(entry)
    if diagnostics:
        return _blocked(diagnostics, run_id=index.batch_id, inventory_model_loads=inventory_model_loads)
    supplied_scope: dict[str, Any] | None = None
    supplied_scope_bytes: bytes | None = None
    external_scope_sha: str | None = None
    if scope_path is not None:
        try:
            supplied_scope_bytes = scope_path.expanduser().resolve().read_bytes()
            supplied_scope = json.loads(supplied_scope_bytes)
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            return _blocked([_diagnostic("scope_invalid", f"Cannot load scope JSON: {exc}")], run_id=index.batch_id, inventory_model_loads=inventory_model_loads)
        scope_id, _retained, _ignored, scope_diags = _scope_error(supplied_scope, index.entries, inventories)
        if scope_diags:
            return _blocked(scope_diags, run_id=index.batch_id, inventory_model_loads=inventory_model_loads)
        external_scope_sha = _sha(supplied_scope_bytes)
    else:
        scope_id = None

    effective_scopes: dict[str, tuple[dict[str, Any], str]] = {}
    for entry in usable:
        if supplied_scope is not None and entry.source_id == scope_id:
            scope = supplied_scope
            scope_sha256 = external_scope_sha or ""
        else:
            scope = _default_scope(entry, inventories[entry.source_id])
            scope_sha256 = _sha(_json_bytes(scope))
        try:
            retained, ignored, _ = _scope_parts(scope, inventories[entry.source_id])
        except ValueError as exc:
            code, _, message = str(exc).partition(": ")
            return _blocked([_diagnostic(code or "scope_invalid", message or str(exc), source_id=entry.source_id)], run_id=index.batch_id, inventory_model_loads=inventory_model_loads)
        normalized = dict(scope)
        normalized["retained_sheets"] = retained
        normalized["ignored_sheets"] = ignored
        effective_scopes[entry.source_id] = normalized, scope_sha256
    if supplied_scope_bytes is None:
        scope_document: dict[str, Any] = {
            "schema_version": "analysis.scope.v1", "decision_source": "default_all_sheets",
            "sources": [scope for scope, _ in effective_scopes.values()],
        }
        if len(effective_scopes) == 1:
            scope_document.update(next(iter(effective_scopes.values()))[0])
        supplied_scope_bytes = _json_bytes(scope_document)
    scope_output_sha = _sha(supplied_scope_bytes)

    if dry_run:
        planned = ["manifest.json", "scope.json", "dependency_audit.json", "retained_dependency_cells.json", "step1/HANDOFF.json", "step1/HANDOFF.md", "step2/HANDOFF.json", "step2/HANDOFF.md"]
        planned.extend(f"sources/{entry.source_id}/dependency_graph.json" for entry in usable)
        planned.extend(f"sources/{entry.source_id}/views/*.json" for entry in usable)
        first = usable[0] if usable else None
        return {
            "tool": "step2.prepare", "status": "dry_run",
            "source": {"source_sha256": first.source_sha256, "source_path": first.source_path} if first else None,
            "run_id": first.run_id if first else index.batch_id,
            "inputs": {"index": {"root": "step2_index", "path": index_path.name, "sha256": _sha(index_bytes)},
                       "step1_root": {"root": "step1", "path": "."},
                       "scope_sha256": scope_output_sha},
            "capabilities": {"native_index_validated": True, "scope_validated": True,
                              "implemented_reader_commands": ["step2.prepare"]},
            "artifacts": {}, "planned_files": planned, "diagnostics": [],
            "next_action": {"action": "run_prepare_without_dry_run"},
            "metrics": {"inventory_model_loads": inventory_model_loads, "raw_workbook_parses": 0, "graph_builds": 0, "elapsed_seconds": round(time.perf_counter() - started, 6)},
        }

    index_sha = _sha(index_bytes)
    index_ref = _source_ref(index_path.name, index_sha, "step2_index")
    scope_ref = ReadingFileRef(root="reading", path="scope.json", sha256=scope_output_sha, bytes=len(supplied_scope_bytes))
    input_parts: list[Any] = [
        index_sha, index_ref.path, PREPARE_COMPILER_VERSION, PREPARE_OPTIONS,
        {"step1": str(root), "step2_index": str(index_path.parent)},
        {source_id: scope_sha for source_id, (_, scope_sha) in sorted(effective_scopes.items())},
        {entry.source_id: [(ref.name, ref.path, ref.sha256) for ref in entry.artifacts] for entry in index.entries},
    ]
    input_fingerprint = _sha(_json_bytes(input_parts))
    revision_id = input_fingerprint[:20]
    manifest_path = out / "manifest.json"
    if resume and manifest_path.exists():
        try:
            saved_manifest = Step2ReadingManifest.model_validate_json(manifest_path.read_bytes())
        except (OSError, ValueError) as exc:
            return _blocked([_diagnostic("prepared_manifest_invalid", f"Cannot validate prepared manifest: {exc}")], run_id=index.batch_id, inventory_model_loads=inventory_model_loads)
        saved_index_identity = (saved_manifest.index.root, saved_manifest.index.path, saved_manifest.index.sha256)
        current_index_identity = (index_ref.root, index_ref.path, index_ref.sha256)
        saved_scope_identity = (saved_manifest.scope.root, saved_manifest.scope.path, saved_manifest.scope.sha256, saved_manifest.scope.bytes)
        current_scope_identity = (scope_ref.root, scope_ref.path, scope_ref.sha256, scope_ref.bytes)
        current_roots = {"reading": ".", "step1": str(root), "step2_index": str(index_path.parent)}
        expected_views = {entry.source_id: _expected_view_bindings(entry.source_id, inventories[entry.source_id]) for entry in usable}
        source_bindings_match = saved_manifest.source_count == len(index.entries) and len(saved_manifest.sources) == len(index.entries)
        if source_bindings_match:
            for saved_source, entry in zip(saved_manifest.sources, index.entries):
                effective_scope, effective_scope_sha = effective_scopes.get(entry.source_id, ({}, None))
                expected_artifacts = {
                    Path(ref.name).name: ReadingFileRef(root="step1", path=ref.path, sha256=ref.sha256)
                    for ref in entry.artifacts
                }
                expected_views_for_source = expected_views.get(entry.source_id, [])
                actual_views = [
                    (ref.path, ref.view_id, ref.scope, ref.sheet_name)
                    for ref in saved_source.views
                ]
                source_bindings_match = (
                    saved_source.source_id == entry.source_id
                    and saved_source.source_path == entry.source_path
                    and saved_source.source_sha256 == entry.source_sha256
                    and saved_source.run_id == entry.run_id
                    and saved_source.status == entry.status
                    and saved_source.ready_for_next_step == entry.ready_for_next_step
                    and saved_source.diagnostics == entry.diagnostics
                    and saved_source.scope_sha256 == effective_scope_sha
                    and saved_source.effective_scope == effective_scope
                    and saved_source.retained_sheets == effective_scope.get("retained_sheets", [])
                    and saved_source.excluded_sheets == effective_scope.get("ignored_sheets", [])
                    and saved_source.artifacts == expected_artifacts
                    and actual_views == expected_views_for_source
                    and all(ref.root == "reading" for ref in saved_source.views)
                )
                if entry.source_id in inventories:
                    graph_path = f"sources/{entry.source_id}/dependency_graph.json"
                    graph_ref = saved_source.dependency_graph
                    source_bindings_match = (
                        source_bindings_match and graph_ref is not None and graph_ref.root == "reading"
                        and graph_ref.path == graph_path
                    )
                else:
                    source_bindings_match = source_bindings_match and saved_source.dependency_graph is None
                if source_bindings_match:
                    expected_workbook_view = next((ref for ref in saved_source.views if ref.scope == "workbook"), None)
                    expected_sheet_views = {ref.sheet_name: ref for ref in saved_source.views if ref.scope == "sheet" and ref.sheet_name}
                    expected_region_views = [ref for ref in saved_source.views if ref.scope == "region"]
                    source_bindings_match = (
                        saved_source.lookup.workbook_view == expected_workbook_view
                        and saved_source.lookup.sheet_views == expected_sheet_views
                        and saved_source.lookup.region_views == expected_region_views
                    )
                if source_bindings_match and entry.source_id in inventories:
                    graph_path = out / f"sources/{entry.source_id}/dependency_graph.json"
                    graph_hash = saved_manifest.outputs.get(graph_path.relative_to(out).as_posix())
                    if graph_hash and graph_path.is_file() and _file_sha(graph_path) == graph_hash:
                        try:
                            saved_graph = FormulaGraph.model_validate_json(graph_path.read_bytes())
                        except (OSError, ValueError):
                            source_bindings_match = False
                        else:
                            expected_names = _defined_name_lookup(inventories[entry.source_id], saved_graph)
                            source_bindings_match = saved_source.lookup.defined_names == expected_names
                    audit_path = out / "dependency_audit.json"
                    audit_hash = saved_manifest.outputs.get("dependency_audit.json")
                    if source_bindings_match and audit_hash and audit_path.is_file() and _file_sha(audit_path) == audit_hash:
                        try:
                            saved_audit = Step2DependencyAudit.model_validate_json(audit_path.read_bytes())
                        except (OSError, ValueError):
                            source_bindings_match = False
                        else:
                            audit_source = next((item for item in saved_audit.sources if item.source_id == entry.source_id), None)
                            source_bindings_match = (
                                audit_source is not None
                                and audit_source.source_sha256 == entry.source_sha256
                                and audit_source.run_id == entry.run_id
                                and audit_source.scope_sha256 == effective_scope_sha
                                and audit_source.retained_dependency_range_count == len(audit_source.ranges)
                                and saved_source.allowed_dependency_ranges == audit_source.ranges
                            )
                elif source_bindings_match:
                    source_bindings_match = not saved_source.lookup.defined_names and not saved_source.allowed_dependency_ranges
                if source_bindings_match and entry.source_id in inventories:
                    expected_allowed_paths = [
                        (ref.root, ref.path) for ref in saved_source.views
                        if ref.scope == "workbook"
                        or (ref.scope == "sheet" and ref.sheet_name in effective_scope.get("retained_sheets", []))
                    ]
                    expected_allowed_paths.append(("reading", "retained_dependency_cells.json"))
                    actual_allowed_paths = [(ref.root, ref.path) for ref in saved_source.allowed_read_paths]
                    source_bindings_match = actual_allowed_paths == expected_allowed_paths
                elif source_bindings_match:
                    source_bindings_match = not saved_source.allowed_read_paths
                if not source_bindings_match:
                    break
        if (
            saved_manifest.input_fingerprint == input_fingerprint
            and saved_manifest.compiler_version == PREPARE_COMPILER_VERSION
            and saved_manifest.options == PREPARE_OPTIONS
            and saved_index_identity == current_index_identity
            and saved_scope_identity == current_scope_identity
            and saved_manifest.roots == current_roots
            and source_bindings_match
        ):
            expected_outputs = {
                "scope.json", "dependency_audit.json", "retained_dependency_cells.json",
                "step1/HANDOFF.json", "step1/HANDOFF.md", "step2/HANDOFF.json", "step2/HANDOFF.md",
            }
            for entry in usable:
                expected_outputs.update(binding[0] for binding in expected_views[entry.source_id])
                expected_outputs.add(f"sources/{entry.source_id}/dependency_graph.json")
            missing_entries = sorted(expected_outputs - set(saved_manifest.outputs))
            if missing_entries:
                diagnostics.append(_diagnostic("prepared_output_manifest_incomplete", "Prepared manifest omits required output references.", paths=missing_entries))
            for source in saved_manifest.sources:
                references = [*source.views, *source.allowed_read_paths]
                if source.dependency_graph:
                    references.append(source.dependency_graph)
                for ref in references:
                    expected_hash = saved_manifest.outputs.get(ref.path)
                    if ref.root != "reading" or expected_hash is None:
                        continue
                    if ref.sha256 != expected_hash:
                        diagnostics.append(_diagnostic("prepared_output_reference_mismatch", "Prepared source reference checksum differs from its output entry.", path=ref.path))
                    path = out / ref.path
                    if path.is_file() and (ref.bytes is None or path.stat().st_size != ref.bytes):
                        diagnostics.append(_diagnostic("prepared_output_reference_mismatch", "Prepared source reference byte count differs from its output file.", path=ref.path))
            for relative, expected in saved_manifest.outputs.items():
                path = out / relative
                if not path.is_file():
                    diagnostics.append(_diagnostic("prepared_output_missing", f"Prepared output is missing: {relative}", path=relative))
                elif _file_sha(path) != expected:
                    diagnostics.append(_diagnostic("prepared_output_checksum_mismatch", f"Prepared output checksum differs: {relative}", path=relative))
            if diagnostics:
                return _blocked(diagnostics, source={"source_sha256": usable[0].source_sha256} if usable else None, run_id=index.batch_id, inventory_model_loads=inventory_model_loads)
            output_refs = _result_artifacts(out, saved_manifest.outputs)
            return {
                "tool": "step2.prepare", "status": "reused",
                "source": {"source_sha256": usable[0].source_sha256, "source_path": usable[0].source_path} if usable else None,
                "run_id": usable[0].run_id if usable else index.batch_id,
                "artifacts": output_refs, "planned_files": [], "diagnostics": [],
                "next_action": {"action": "read", "root": "reading", "path": "step2/HANDOFF.md"},
                "metrics": {"inventory_model_loads": inventory_model_loads, "raw_workbook_parses": 0, "graph_builds": 0, "reused_outputs": len(saved_manifest.outputs), "elapsed_seconds": round(time.perf_counter() - started, 6)},
            }

    graph_models: dict[str, FormulaGraph] = {}
    extra_records: dict[str, dict[str, list[ViewRecord]]] = {}
    audit_sources: list[DependencyAuditSource] = []
    snapshot_sources: list[DependencySnapshotSource] = []
    legacy_sidecars: set[str] = set()
    for entry in usable:
        inventory = inventories[entry.source_id]
        graph = RegexFormulaGraphBuilder().build(inventory)
        graph_models[entry.source_id] = graph
        scope, scope_sha256 = effective_scopes[entry.source_id]
        retained, ignored, _ = _scope_parts(scope, inventory)
        audit, snapshot = _dependency_data(entry, inventory, graph, retained, ignored, scope_sha256)
        audit_sources.append(audit)
        snapshot_sources.append(snapshot)
        policy = scope.get("dependency_policy")
        if isinstance(policy, dict):
            legacy_sidecars.update(value for value in (policy.get("snapshot"), policy.get("audit")) if isinstance(value, str) and value)
        extra_records[entry.source_id] = _control_records(entry, inventory, artifact_models)

    view_index = index.model_copy(update={"entries": usable})
    try:
        views = compile_views(None, root, budget=int(PREPARE_OPTIONS["view_token_budget"]), validated_index=view_index,
                              validated_artifacts=artifact_models, graph_overrides=graph_models, extra_records=extra_records)
    except (OSError, ValueError) as exc:
        return _blocked([_diagnostic("view_compile_failed", str(exc))], run_id=index.batch_id, inventory_model_loads=inventory_model_loads)

    outputs: dict[str, bytes] = {"scope.json": supplied_scope_bytes}
    source_models: dict[str, ReadingSource] = {}
    view_objs_by_source: dict[str, list[Any]] = {}
    for view in views:
        source_dir = f"sources/{view.source_id}"
        view_rel = f"{source_dir}/views/{view.view_id}.json"
        data = _json_bytes(view.model_dump(mode="json"))
        outputs[view_rel] = data
        view_objs_by_source.setdefault(view.source_id, []).append((view, view_rel, _sha(data), len(data)))
    for entry in index.entries:
        artifact_refs = {Path(ref.name).name: ReadingFileRef(root="step1", path=ref.path, sha256=ref.sha256) for ref in entry.artifacts}
        usable_entry = entry.source_id in inventories
        scope, scope_hash = effective_scopes.get(entry.source_id, ({}, None))
        source_views = view_objs_by_source.get(entry.source_id, [])
        view_refs = [ReadingFileRef(root="reading", path=rel, sha256=digest, bytes=size, view_id=view.view_id, scope=view.scope, sheet_name=view.sheet_name) for view, rel, digest, size in source_views]
        graph_ref = None
        ranges: list[dict[str, str]] = []
        if usable_entry:
            graph_rel = f"sources/{entry.source_id}/dependency_graph.json"
            graph_bytes = _json_bytes(graph_models[entry.source_id].model_dump(mode="json"))
            outputs[graph_rel] = graph_bytes
            graph_ref = ReadingFileRef(root="reading", path=graph_rel, sha256=_sha(graph_bytes), bytes=len(graph_bytes))
            ranges = next((item.ranges for item in audit_sources if item.source_id == entry.source_id), [])
        defined_names = _defined_name_lookup(inventories[entry.source_id], graph_models[entry.source_id]) if usable_entry else []
        lookup = ReadingLookup(
            workbook_view=next((ref for ref in view_refs if ref.scope == "workbook"), None),
            sheet_views={ref.sheet_name: ref for ref in view_refs if ref.scope == "sheet" and ref.sheet_name},
            region_views=[ref for ref in view_refs if ref.scope == "region"],
            defined_names=defined_names,
        )
        if usable_entry:
            allowed = [ref for ref in view_refs if ref.scope == "workbook" or (ref.scope == "sheet" and ref.sheet_name in scope.get("retained_sheets", []))]
            allowed.append(ReadingFileRef(root="reading", path="retained_dependency_cells.json", sha256="0" * 64))
        else:
            allowed = []
        source_models[entry.source_id] = ReadingSource(
            source_id=entry.source_id, source_path=entry.source_path, source_sha256=entry.source_sha256,
            run_id=entry.run_id, status=entry.status, ready_for_next_step=entry.ready_for_next_step,
            diagnostics=entry.diagnostics, artifacts=artifact_refs, views=view_refs,
            dependency_graph=graph_ref, retained_sheets=scope.get("retained_sheets", []),
            excluded_sheets=scope.get("ignored_sheets", []), allowed_dependency_ranges=ranges,
            allowed_read_paths=allowed, scope_sha256=scope_hash, effective_scope=scope,
            lookup=lookup,
        )

    audit_model = Step2DependencyAudit(ignored_legacy_sidecars=sorted(legacy_sidecars), sources=audit_sources)
    snapshot_model = Step2DependencySnapshot(sources=snapshot_sources)
    audit_bytes = _json_bytes(audit_model.model_dump(mode="json"))
    snapshot_bytes = _json_bytes(snapshot_model.model_dump(mode="json"))
    outputs["dependency_audit.json"] = audit_bytes
    outputs["retained_dependency_cells.json"] = snapshot_bytes
    for source in source_models.values():
        source.allowed_read_paths = [
            ref.model_copy(update={"sha256": _sha(snapshot_bytes), "bytes": len(snapshot_bytes)})
            if ref.path == "retained_dependency_cells.json" else ref
            for ref in source.allowed_read_paths
        ]

    source_rows = [{
        "source_id": item.source_id, "source_path": item.source_path, "source_sha256": item.source_sha256,
        "run_id": item.run_id, "status": item.status, "ready_for_next_step": item.ready_for_next_step,
        "diagnostics": item.diagnostics,
        "artifacts": {name: ref.model_dump(mode="json") for name, ref in item.artifacts.items()},
    } for item in source_models.values()]
    step1_handoff = {
        "schema_version": "step2.prepare.step1_handoff.v1", "artifact_type": "prepared_step1_handoff",
        "source_quality": {"status": index.status, "entry_count": len(index.entries)},
        "reference_integrity": {"status": "pass", "index_sha256": index_sha},
        "input_handoff": {"root": "step1", "path": index.input_handoff_path, "sha256": index.input_handoff_sha256},
        "sources": source_rows,
    }
    audit_ref = {"root": "reading", "path": "dependency_audit.json", "sha256": _sha(audit_bytes), "bytes": len(audit_bytes)}
    snapshot_ref = {"root": "reading", "path": "retained_dependency_cells.json", "sha256": _sha(snapshot_bytes), "bytes": len(snapshot_bytes)}
    step2_handoff = {
        "schema_version": "step2.prepare.handoff.v1", "artifact_type": "prepared_reading_handoff",
        "source_quality": {"status": index.status, "entry_count": len(index.entries), "failed_entries": sum(entry.status not in {"pass", "partial"} for entry in index.entries)},
        "reference_integrity": {"status": "pass", "index_sha256": index_sha, "diagnostics": []},
        "static_readiness": {
            "status": "partial" if any(source.unresolved_static_reference_count or source.diagnostics for source in audit_sources) else "supported_static_references_audited",
            "audit": audit_ref,
            "snapshot": snapshot_ref,
            "sources": [{
                "source_id": source.source_id,
                "source_sha256": source.source_sha256,
                "run_id": source.run_id,
                "scope_sha256": source.scope_sha256,
                "supported_static_inbound_reference_count": source.supported_static_inbound_reference_count,
                "retained_dependency_range_count": source.retained_dependency_range_count,
                "unresolved_static_reference_count": source.unresolved_static_reference_count,
                "stored_cell_count": len(snapshot_sources[i].cells),
            } for i, source in enumerate(audit_sources)],
        },
        "runtime_and_numerical_behavior": {"status": "not_verified", "macros_executed": False, "recalculated": False, "scenario_reconciliation": False},
        "scope_sha256": scope_output_sha,
        "sources": [{"source_id": item.source_id, "source_path": item.source_path, "run_id": item.run_id,
                     "source_sha256": item.source_sha256, "status": item.status,
                     "ready_for_next_step": item.ready_for_next_step, "diagnostics": item.diagnostics,
                     "retained_sheets": item.retained_sheets,
                     "excluded_sheets": item.excluded_sheets,
                     "allowed_dependency_ranges": item.allowed_dependency_ranges,
                     "allowed_read_paths": [ref.model_dump(mode="json") for ref in item.allowed_read_paths]} for item in source_models.values()],
        "artifacts": {"manifest": {"root": "reading", "path": "manifest.json"}, "dependency_audit": audit_ref, "dependency_snapshot": snapshot_ref},
    }
    outputs["step1/HANDOFF.json"] = _json_bytes(step1_handoff)
    outputs["step1/HANDOFF.md"] = _human_step1(step1_handoff).encode("utf-8")
    outputs["step2/HANDOFF.json"] = _json_bytes(step2_handoff)
    outputs["step2/HANDOFF.md"] = _human_step2(step2_handoff).encode("utf-8")
    output_hashes = {name: _sha(data) for name, data in outputs.items()}
    source_manifest = [item for item in source_models.values()]
    manifest = Step2ReadingManifest(
        compiler_version=PREPARE_COMPILER_VERSION, revision_id=revision_id,
        input_fingerprint=input_fingerprint, options=PREPARE_OPTIONS,
        roots={"reading": ".", "step1": str(root), "step2_index": str(index_path.parent)},
        source_count=len(index.entries), index=index_ref, scope=scope_ref,
        sources=source_manifest, outputs=output_hashes,
    )
    manifest_bytes = _json_bytes(manifest.model_dump(mode="json"))
    try:
        out.mkdir(parents=True, exist_ok=True)
        for relative, data in outputs.items():
            path = out / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            temporary = path.with_name(path.name + ".tmp")
            temporary.write_bytes(data)
            temporary.replace(path)
        (out / "manifest.json.tmp").write_bytes(manifest_bytes)
        (out / "manifest.json.tmp").replace(manifest_path)
    except OSError as exc:
        return _blocked([_diagnostic("prepare_write_failed", f"Cannot write reading package: {exc}")], run_id=index.batch_id, inventory_model_loads=inventory_model_loads)
    all_hashes = dict(output_hashes)
    all_hashes["manifest.json"] = _sha(manifest_bytes)
    first = usable[0] if usable else None
    return {
        "tool": "step2.prepare", "status": "prepared",
        "source": {"source_sha256": first.source_sha256, "source_path": first.source_path} if first else None,
        "run_id": first.run_id if first else index.batch_id,
        "artifacts": _result_artifacts(out, all_hashes), "planned_files": [], "diagnostics": [],
        "next_action": {"action": "read", "root": "reading", "path": "step2/HANDOFF.md"},
        "metrics": {
            "inventory_model_loads": inventory_model_loads, "raw_workbook_parses": 0, "graph_builds": len(graph_models),
            "views_written": len(views), "stored_dependency_cells": sum(len(item.cells) for item in snapshot_sources),
            "supported_static_inbound_references": sum(item.supported_static_inbound_reference_count for item in audit_sources),
            "retained_dependency_ranges": sum(item.retained_dependency_range_count for item in audit_sources),
            "elapsed_seconds": round(time.perf_counter() - started, 6),
        },
    }
