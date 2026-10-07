"""Bounded readers for a prepared Step 2 reading package."""

from __future__ import annotations

import base64
import hashlib
import json
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl.utils.cell import get_column_letter, range_boundaries

from excel_to_act.schemas import Step2ReadingManifest, ViewRecord, WorkbookView
from excel_to_act.schemas.step2_prepare import ReadingFileRef, ReadingSource
from excel_to_act.schemas.step2_query import EvidencePacket
from excel_to_act.views import _chunks, _tokens


_KINDS = {"overview", "sheet", "cell", "range", "name", "control", "vba", "feature"}


class QueryFailure(ValueError):
    def __init__(self, status: str, code: str, message: str, **details: Any) -> None:
        super().__init__(message)
        self.status = status
        self.diagnostic = {"code": code, "severity": "error", "message": message, **details}


def _json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _resolved(root: Path, relative: str) -> Path:
    path = (root / Path(*PurePosixPath(relative).parts)).resolve()
    if not path.is_relative_to(root.resolve()):
        raise QueryFailure("integrity_failed", "reference_outside_root", "A referenced file resolves outside its declared root.", path=relative)
    return path


def _root_path(manifest_path: Path, manifest: Step2ReadingManifest, root_name: str) -> Path:
    value = manifest.roots.get(root_name)
    if not isinstance(value, str) or not value:
        raise QueryFailure("integrity_failed", "manifest_root_missing", f"Manifest root {root_name!r} is missing.")
    root = Path(value).expanduser()
    return root.resolve() if root.is_absolute() else (manifest_path.parent / root).resolve()


def _read_ref(ref: ReadingFileRef, roots: dict[str, Path], manifest: Step2ReadingManifest) -> bytes:
    root = roots.get(ref.root)
    if root is None:
        raise QueryFailure("integrity_failed", "reference_root_unknown", "Reference uses an undeclared root.", root=ref.root, path=ref.path)
    try:
        data = _resolved(root, ref.path).read_bytes()
    except OSError as exc:
        raise QueryFailure("unavailable", "evidence_file_unavailable", f"Cannot read referenced evidence: {exc}", root=ref.root, path=ref.path) from exc
    if _sha(data) != ref.sha256 or (ref.bytes is not None and len(data) != ref.bytes):
        raise QueryFailure("integrity_failed", "evidence_hash_mismatch", "Referenced evidence does not match its canonical SHA-256 or byte length.", root=ref.root, path=ref.path)
    if ref.root == "reading" and manifest.outputs.get(ref.path) != ref.sha256:
        raise QueryFailure("integrity_failed", "manifest_output_mismatch", "Reading reference does not match the selected manifest output hash.", path=ref.path)
    return data


def _load_manifest(path: Path) -> tuple[bytes, Step2ReadingManifest, Path, dict[str, Path]]:
    path = path.expanduser().resolve()
    try:
        data = path.read_bytes()
        manifest = Step2ReadingManifest.model_validate_json(data)
    except (OSError, ValueError) as exc:
        raise QueryFailure("integrity_failed", "manifest_invalid", f"Cannot load a valid Step 2 reading manifest: {exc}", path=str(path)) from exc
    if "reading" not in manifest.roots:
        raise QueryFailure("integrity_failed", "manifest_root_missing", "Manifest root 'reading' is missing.", root="reading", path=str(path))
    roots = {name: _root_path(path, manifest, name) for name in manifest.roots}
    return data, manifest, roots["reading"], roots


def _source(manifest: Step2ReadingManifest, source_id: str) -> ReadingSource:
    matches = [item for item in manifest.sources if item.source_id == source_id]
    if len(matches) != 1:
        status = "unavailable" if not matches else "integrity_failed"
        code = "source_not_found" if not matches else "source_id_ambiguous"
        raise QueryFailure(status, code, "The explicit source ID does not identify exactly one manifest source.", source_id=source_id)
    return matches[0]


def _sheet_ref(source: ReadingSource, sheet: str) -> ReadingFileRef | None:
    return next((ref for ref in source.views if ref.scope == "sheet" and ref.sheet_name and ref.sheet_name.casefold() == sheet.casefold()), None)


def _workbook_ref(source: ReadingSource) -> ReadingFileRef | None:
    return source.lookup.workbook_view


def _normalize_address(value: str, *, cell: bool = False) -> tuple[str, tuple[int, int, int, int]]:
    address = value.strip().replace("$", "").upper()
    try:
        bounds = range_boundaries(address)
    except ValueError as exc:
        raise QueryFailure("unsupported", "selector_address_invalid", f"Selector address {value!r} is not an A1 cell/range.", requested=value) from exc
    if any(item is None for item in bounds):
        raise QueryFailure("unsupported", "selector_address_invalid", f"Selector address {value!r} is not an A1 cell/range.", requested=value)
    min_col, min_row, max_col, max_row = bounds
    if min_col < 1 or max_col > 16_384 or min_row < 1 or max_row > 1_048_576:
        raise QueryFailure("unsupported", "selector_address_invalid", f"Selector address {value!r} is outside Excel worksheet bounds.", requested=value)
    if cell and (min_col != max_col or min_row != max_row):
        raise QueryFailure("unsupported", "selector_cell_invalid", "A cell selector must identify exactly one cell.", requested=value)
    normalized = f"{get_column_letter(min_col)}{min_row}"
    if (min_col, min_row) != (max_col, max_row):
        normalized += f":{get_column_letter(max_col)}{max_row}"
    return normalized, bounds  # type: ignore[return-value]


def _parse_sheet_address(target: str, sheet: str | None) -> tuple[str | None, str]:
    if sheet:
        return sheet, target
    if "!" in target:
        raw_sheet, address = target.rsplit("!", 1)
        if raw_sheet.startswith("'") and raw_sheet.endswith("'"):
            raw_sheet = raw_sheet[1:-1].replace("''", "'")
        else:
            raw_sheet = raw_sheet.strip("'")
        return raw_sheet, address
    return None, target


def _is_excluded(source: ReadingSource, sheet: str | None) -> bool:
    return bool(sheet and any(name.casefold() == sheet.casefold() for name in source.excluded_sheets))


def _allowed_ranges(source: ReadingSource, sheet: str) -> list[tuple[str, tuple[int, int, int, int]]]:
    result = []
    for item in source.allowed_dependency_ranges:
        if item.get("sheet_name", "").casefold() != sheet.casefold() or not item.get("address"):
            continue
        try:
            result.append((item["address"], _normalize_address(item["address"])[1]))
        except QueryFailure:
            continue
    return result


def _contains(outer: tuple[int, int, int, int], inner: tuple[int, int, int, int]) -> bool:
    return inner[0] >= outer[0] and inner[1] >= outer[1] and inner[2] <= outer[2] and inner[3] <= outer[3]


def _scope_guard(source: ReadingSource, sheet: str, bounds: tuple[int, int, int, int], selector: dict[str, Any]) -> list[tuple[str, tuple[int, int, int, int]]]:
    if not _is_excluded(source, sheet):
        return []
    allowed = _allowed_ranges(source, sheet)
    if not any(_contains(bounds_allowed, bounds) for _, bounds_allowed in allowed):
        raise QueryFailure(
            "out_of_scope", "selector_out_of_scope", "Requested evidence is outside the approved dependency ranges.",
            requested_selector=selector,
            allowed_targets=[{"sheet": sheet, "range": address} for address, _ in allowed],
        )
    return allowed


def _load_views(refs: list[ReadingFileRef], source: ReadingSource, manifest: Step2ReadingManifest, roots: dict[str, Path]) -> list[tuple[ReadingFileRef, WorkbookView, dict[str, Any], int]]:
    loaded = []
    seen: set[str] = set()
    source_ref_by_path = {ref.path: ref for ref in source.views}
    for ref in refs:
        if ref.path in seen or source_ref_by_path.get(ref.path) != ref:
            raise QueryFailure("integrity_failed", "canonical_view_ref_invalid", "Selected view reference is missing, duplicated, or differs from the manifest source.", path=ref.path)
        seen.add(ref.path)
        data = _read_ref(ref, roots, manifest)
        try:
            raw = json.loads(data)
            view = WorkbookView.model_validate(raw)
        except (ValueError, TypeError) as exc:
            raise QueryFailure("integrity_failed", "canonical_view_invalid", f"Canonical view is invalid: {exc}", path=ref.path) from exc
        if (view.view_id != ref.view_id or view.scope != ref.scope or view.sheet_name != ref.sheet_name
                or view.source_id != source.source_id or view.source_sha256 != source.source_sha256
                or view.source_run_id != source.run_id):
            raise QueryFailure("integrity_failed", "canonical_view_identity_mismatch", "Canonical view identity differs from its manifest reference or source.", path=ref.path)
        loaded.append((ref, view, raw, len(data)))
    return loaded


def _selector(kind: str, target: str | None, sheet: str | None, address: str | None) -> tuple[dict[str, Any], tuple[int, int, int, int] | None]:
    if kind not in _KINDS:
        raise QueryFailure("unsupported", "selector_kind_unsupported", f"Unsupported query kind {kind!r}.", kind=kind)
    normalized_target = target.strip() if target else None
    normalized_sheet = sheet.strip() if sheet else None
    bounds = None
    if kind == "control" and normalized_target and not normalized_sheet and "." in normalized_target:
        normalized_sheet, normalized_target = normalized_target.rsplit(".", 1)
    if kind == "cell":
        normalized_sheet, target_address = _parse_sheet_address(normalized_target or "", normalized_sheet)
        normalized_target, bounds = _normalize_address(target_address, cell=True)
    elif kind == "range":
        normalized_sheet, target_address = _parse_sheet_address(address or normalized_target or "", normalized_sheet)
        normalized_target, bounds = _normalize_address(target_address)
    elif kind in {"sheet", "control", "vba"}:
        if kind == "sheet":
            normalized_sheet = normalized_sheet or normalized_target
            normalized_target = normalized_sheet
        elif not normalized_target:
            raise QueryFailure("unsupported", "selector_target_required", f"Query kind {kind!r} requires --target.", kind=kind)
    elif kind == "name" and not normalized_target:
        raise QueryFailure("unsupported", "selector_target_required", "Query kind 'name' requires --target.", kind=kind)
    elif kind == "feature" and not normalized_target and not normalized_sheet:
        raise QueryFailure("unsupported", "selector_target_required", "Feature query requires --target or --sheet.", kind=kind)
    if kind == "overview":
        normalized_target = normalized_sheet = None
    normalized = {"kind": kind, "target": normalized_target, "sheet": normalized_sheet, "range": normalized_target if kind == "range" else None}
    if kind in {"name", "control", "vba", "feature"} and normalized_target:
        normalized["target"] = normalized_target.casefold()
    if normalized_sheet:
        normalized["sheet"] = normalized_sheet.casefold()
    return normalized, bounds


def _record_bounds(record: ViewRecord) -> tuple[int, int, int, int] | None:
    address = record.source_location.address
    if not address:
        return None
    try:
        return _normalize_address(address)[1]
    except QueryFailure:
        return None


def _in_area(record: ViewRecord, bounds: tuple[int, int, int, int]) -> bool:
    record_bounds = _record_bounds(record)
    if record_bounds is None:
        return False
    if record.record_type == "cell" or record_bounds[0] == record_bounds[2] and record_bounds[1] == record_bounds[3]:
        return _contains(bounds, record_bounds)
    return _contains(bounds, record_bounds)


def _is_excluded_form_control(source: ReadingSource, record: ViewRecord) -> bool:
    return _is_excluded(source, record.source_location.sheet_name) and (
        str(record.source_location.object_type or "").casefold() == "form_control"
        or str(record.facts.get("kind", "")).casefold() == "form_control"
        or record.record_type.casefold() == "form_control"
    )


def _artifact_ref(source: ReadingSource, name: str) -> ReadingFileRef | None:
    return next((ref for key, ref in source.artifacts.items() if Path(key).name.casefold() == name.casefold()), None)


def _control_rows(source: ReadingSource, roots: dict[str, Path], manifest: Step2ReadingManifest) -> list[tuple[dict[str, Any], ReadingFileRef, str, int]]:
    result = []
    for filename, row_key, kind in (
        ("checkbox_bindings.json", "bindings", "checkbox_binding"),
        ("activex_events.json", "controls", "activex_event"),
    ):
        ref = _artifact_ref(source, filename)
        if ref is None:
            continue
        data = _read_ref(ref, roots, manifest)
        try:
            value = json.loads(data)
        except json.JSONDecodeError as exc:
            raise QueryFailure("integrity_failed", "source_artifact_invalid", f"Control artifact is invalid JSON: {exc}", path=ref.path) from exc
        if not isinstance(value, dict) or value.get("workbook_sha256") != source.source_sha256:
            raise QueryFailure("integrity_failed", "control_source_mismatch", "Control artifact belongs to a different workbook source.", path=ref.path)
        rows = value.get(row_key, [])
        if not isinstance(rows, list):
            raise QueryFailure("integrity_failed", "control_artifact_invalid", "Control artifact rows are not a list.", path=ref.path)
        result.extend((row, ref, kind, len(data)) for row in rows if isinstance(row, dict))
    return result


def _matches_control(target: str, row: dict[str, Any]) -> bool:
    wanted = target.casefold()
    return any(str(row.get(key) or "").casefold() == wanted for key in ("control_name", "control_name", "shape_id", "name"))


def _vba_handoff(source: ReadingSource, roots: dict[str, Path], manifest: Step2ReadingManifest) -> tuple[dict[str, Any], ReadingFileRef, int] | None:
    ref = _artifact_ref(source, "vba_handoff.json")
    if ref is None:
        return None
    data = _read_ref(ref, roots, manifest)
    try:
        value = json.loads(data)
    except json.JSONDecodeError as exc:
        raise QueryFailure("integrity_failed", "vba_handoff_invalid", f"VBA handoff is invalid JSON: {exc}", path=ref.path) from exc
    if not isinstance(value, dict) or value.get("workbook_sha256") != source.source_sha256:
        raise QueryFailure("integrity_failed", "vba_source_mismatch", "VBA handoff belongs to a different workbook source.", path=ref.path)
    return value, ref, len(data)


def _module_owner(module: dict[str, Any], source: ReadingSource, control_rows: list[tuple[dict[str, Any], ReadingFileRef, str, int]]) -> str | None:
    if module.get("kind") not in {"Document", "ClassModule"}:
        return None
    name = str(module.get("name", ""))
    known_sheets = {ref.sheet_name.casefold(): ref.sheet_name for ref in source.views if ref.scope == "sheet" and ref.sheet_name}
    if name.casefold() in known_sheets:
        return known_sheets[name.casefold()]
    for row, _, kind, _ in control_rows:
        code_name = row.get("sheet_code_name")
        if kind == "activex_event" and isinstance(code_name, str) and code_name.casefold() == name.casefold():
            sheet = row.get("sheet")
            if isinstance(sheet, str):
                return sheet
    return None


def _expected_selection(
    selector: dict[str, Any], source: ReadingSource, roots: dict[str, Path], manifest: Step2ReadingManifest,
    canonical_by_path: dict[str, dict[str, Any]],
) -> tuple[list[ReadingFileRef], list[tuple[str, str]]]:
    kind, target, sheet = selector["kind"], selector.get("target"), selector.get("sheet")
    refs: list[ReadingFileRef] = []
    if kind == "overview":
        return refs, []
    if kind == "sheet":
        if sheet and not _is_excluded(source, sheet):
            ref = _sheet_ref(source, sheet)
            refs = [ref] if ref else []
    elif kind in {"cell", "range"}:
        if sheet:
            normalized, bounds = _normalize_address(str(target or selector.get("range") or ""), cell=kind == "cell")
            if _is_excluded(source, sheet):
                _scope_guard(source, sheet, bounds, selector)
            ref = _sheet_ref(source, sheet)
            refs = [ref] if ref else []
    elif kind == "name":
        candidates = [item for item in source.lookup.defined_names if item.name.casefold() == str(target or "").casefold()]
        if sheet:
            local = [item for item in candidates if item.scope.casefold() == sheet.casefold()]
            candidates = local or [item for item in candidates if item.scope.casefold() == "workbook"]
        if len(candidates) == 1:
            item = candidates[0]
            name_sheet = item.source_location.sheet_name
            if name_sheet:
                _, bounds = _normalize_address(item.address)
                if _is_excluded(source, name_sheet):
                    _scope_guard(source, name_sheet, bounds, selector)
                ref = _sheet_ref(source, name_sheet)
                refs = [ref] if ref else []
    elif kind == "control":
        rows = _control_rows(source, roots, manifest)
        matches = [item for item in rows if _matches_control(str(target or ""), item[0])]
        if sheet:
            matches = [item for item in matches if str(item[0].get("sheet", "")).casefold() == sheet.casefold()]
        if len(matches) == 1:
            owner = matches[0][0].get("sheet")
            if isinstance(owner, str) and not _is_excluded(source, owner):
                ref = _sheet_ref(source, owner)
                refs = [ref] if ref else []
    elif kind == "vba":
        handoff_data = _vba_handoff(source, roots, manifest)
        modules = [item for item in (handoff_data[0].get("modules", []) if handoff_data else [])
                   if isinstance(item, dict) and str(item.get("name", "")).casefold() == str(target or "").casefold()]
        if len(modules) == 1:
            module = modules[0]
            control_rows = _control_rows(source, roots, manifest) if module.get("kind") in {"Document", "ClassModule"} and source.excluded_sheets else []
            owner = _module_owner(module, source, control_rows)
            unknown_worksheet = module.get("kind") in {"Document", "ClassModule"} and module.get("name") != "ThisWorkbook" and owner is None and source.excluded_sheets
            if not (owner and _is_excluded(source, owner)) and not unknown_worksheet:
                ref = _workbook_ref(source)
                refs = [ref] if ref else []
    elif kind == "feature":
        if sheet:
            if not _is_excluded(source, sheet):
                ref = _sheet_ref(source, sheet)
                refs = [ref] if ref else []
        else:
            ref = _workbook_ref(source)
            refs = [ref] if ref else []

    stream: list[tuple[str, str]] = []
    for ref in refs:
        raw = canonical_by_path.get(ref.path)
        if raw is None:
            raise QueryFailure("integrity_failed", "selector_view_missing", "Packet omits a canonical view required by its selector.", path=ref.path)
        records = raw.get("records", [])
        selected_records: list[dict[str, Any]] = []
        if kind == "sheet":
            selected_records = records
        elif kind == "cell":
            selected_records = [item for item in records if item.get("record_type") == "cell"
                                and str(item.get("source_location", {}).get("address", "")).replace("$", "").upper() == str(target)]
        elif kind == "range":
            _, bounds = _normalize_address(str(target or selector.get("range") or ""))
            for item in records:
                record = ViewRecord.model_validate(item)
                if _in_area(record, bounds) and not _is_excluded_form_control(source, record):
                    selected_records.append(item)
        elif kind == "name":
            candidates = [item for item in source.lookup.defined_names if item.name.casefold() == str(target or "").casefold()]
            if sheet:
                local = [item for item in candidates if item.scope.casefold() == sheet.casefold()]
                candidates = local or [item for item in candidates if item.scope.casefold() == "workbook"]
            if len(candidates) == 1:
                declaration = candidates[0]
                _, name_bounds = _normalize_address(declaration.address)
                for item in records:
                    record = ViewRecord.model_validate(item)
                    if ((record.source_location.sheet_name or "").casefold() == (declaration.source_location.sheet_name or "").casefold()
                            and _in_area(record, name_bounds) and not _is_excluded_form_control(source, record)):
                        selected_records.append(item)
        elif kind == "control":
            control_rows = [item for item in _control_rows(source, roots, manifest) if _matches_control(str(target or ""), item[0])]
            if sheet:
                control_rows = [item for item in control_rows if str(item[0].get("sheet", "")).casefold() == sheet.casefold()]
            if len(control_rows) == 1:
                row = control_rows[0][0]
                selected_records = [item for item in records if item.get("record_type") in {"checkbox_binding", "activex_event"}
                                    and item.get("facts", {}).get("record") == row]
        elif kind == "vba":
            selected_records = [item for item in records if item.get("record_type") == "vba_module"
                                and str(item.get("facts", {}).get("name", "")).casefold() == str(target or "").casefold()]
        elif kind == "feature":
            wanted = str(target or "").casefold()
            selected_records = [item for item in records if item.get("record_type") in {"unsupported", "unresolved_reference"}
                                and (not wanted or wanted in {str(item.get("facts", {}).get("feature_type", "")).casefold(), item.get("record_type"), str(item.get("record_id", "")).casefold()})]
        stream.extend((str(raw.get("view_id")), str(item.get("record_id"))) for item in selected_records)
    return refs, stream


def _finish_packet(
    *, manifest_path: Path, manifest_bytes: bytes, manifest: Step2ReadingManifest, source: ReadingSource,
    selector: dict[str, Any], status: str, selected: list[tuple[ReadingFileRef, WorkbookView, dict[str, Any], int]],
    stream: list[tuple[str, ViewRecord]], budget: int, cursor: str | None, summary: dict[str, Any],
    diagnostics: list[dict[str, Any]], source_file_refs: list[ReadingFileRef], source_bytes: int,
) -> dict[str, Any]:
    canonical_refs = [ref for ref, _, _, _ in selected]
    canonical_hashes = [(ref.path, ref.sha256) for ref in canonical_refs]
    stream_ids = [(view_id, record.record_id) for view_id, record in stream]
    binding_data = {
        "manifest_sha256": _sha(manifest_bytes), "source_id": source.source_id,
        "selector": selector, "canonical_views": canonical_hashes, "stream": stream_ids,
    }
    binding = _sha(json.dumps(binding_data, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"))
    start = 0
    if cursor:
        try:
            raw = base64.urlsafe_b64decode(cursor + "=" * (-len(cursor) % 4))
            token = json.loads(raw)
            start = token["next_index"]
            if token.get("version") != 1 or token.get("binding") != binding or type(start) is not int or not 0 <= start <= len(stream):
                raise ValueError("cursor binding or index is invalid")
        except (ValueError, KeyError, TypeError, json.JSONDecodeError) as exc:
            raise QueryFailure("integrity_failed", "cursor_invalid", f"Cursor is malformed or bound to different evidence: {exc}") from exc

    page: list[tuple[str, ViewRecord]] = []
    used = 0
    index = start
    oversized: dict[str, Any] | None = None
    canonical_ref_by_view = {view.view_id: ref for ref, view, _, _ in selected}
    while index < len(stream):
        view_id, record = stream[index]
        cost = _tokens(record)
        if cost > budget:
            oversized = {"canonical_ref": canonical_ref_by_view[view_id].model_dump(mode="json"),
                         "view_id": view_id, "record_id": record.record_id, "required_estimated_tokens": cost}
            break
        if page and used + cost > budget:
            break
        page.append((view_id, record))
        used += cost
        index += 1
    has_more = index < len(stream)
    next_cursor = None
    if has_more:
        token = {"version": 1, "binding": binding, "next_index": index}
        next_cursor = base64.urlsafe_b64encode(json.dumps(token, sort_keys=True, separators=(",", ":")).encode()).decode().rstrip("=")
    grouped: dict[str, list[ViewRecord]] = {}
    delivered: dict[str, list[str]] = {}
    for view_id, record in page:
        grouped.setdefault(view_id, []).append(record)
        delivered.setdefault(view_id, []).append(record.record_id)
    page_views = []
    for ref, view, _, _ in selected:
        records = grouped.get(view.view_id)
        if not records:
            continue
        page_views.append(view.model_copy(update={
            "records": records,
            "chunks": _chunks(view.view_id, records, budget),
            "opaque_report": [item.record_id for item in records if item.facts.get("opaque") is True],
        }))
    if oversized and not page:
        status = "oversized"
    elif has_more and status in {"ok", "partial"}:
        status = "truncated"
    if oversized and not any(item.get("code") == "record_oversized" for item in diagnostics):
        diagnostics.append({"code": "record_oversized", "severity": "error" if not page else "warning",
                            "message": "The next complete record exceeds this page budget and was left for a larger-budget continuation.", **oversized})
    omitted = [{"view_id": view_id, "record_id": record.record_id} for view_id, record in stream[index:]]
    packet = EvidencePacket(
        status=status,
        manifest={"path": str(manifest_path), "reading_root": str(Path(manifest.roots["reading"])), "sha256": _sha(manifest_bytes), "bytes": len(manifest_bytes)},
        source={
            "source_id": source.source_id, "source_path": source.source_path, "source_sha256": source.source_sha256,
            "run_id": source.run_id, "status": source.status, "revision_id": manifest.revision_id,
            "scope_sha256": source.scope_sha256, "retained_sheets": source.retained_sheets,
            "excluded_sheets": source.excluded_sheets, "allowed_dependency_ranges": source.allowed_dependency_ranges,
        },
        selector=selector,
        canonical_view_refs=canonical_refs,
        source_file_refs=source_file_refs,
        views=page_views,
        delivered_record_ids=delivered,
        summary=summary,
        pagination={
            "budget": budget, "estimated_tokens": used, "delivered_count": len(page),
            "start_index": start,
            "binding_sha256": binding,
            "delivered": [{"view_id": view_id, "record_id": record.record_id} for view_id, record in page],
            "omitted_count": len(omitted), "omitted": omitted[:100], "omitted_ids_truncated": len(omitted) > 100,
            "oversized": oversized, "next_cursor": next_cursor,
        },
        metrics={"selected_views": [{"path": ref.path, "bytes": size} for ref, _, _, size in selected],
                 "selected_view_count": len(selected), "selected_view_bytes": sum(size for _, _, _, size in selected),
                 "source_file_bytes": source_bytes, "raw_workbook_parses": 0, "inventory_deserializations": 0,
                 "graph_builds": 0, "combined_views_loaded": False},
        diagnostics=diagnostics,
    )
    return packet.model_dump(mode="json")


def query(
    manifest_path: Path, source_id: str, kind: str, *, target: str | None = None,
    sheet: str | None = None, address: str | None = None, budget: int = 1200,
    cursor: str | None = None,
) -> dict[str, Any]:
    """Read a bounded, source-bound page from the #30 prepared evidence."""
    if not source_id:
        raise QueryFailure("unsupported", "source_id_required", "--source-id is required, including for a single-source manifest.")
    if budget < 1:
        raise QueryFailure("unsupported", "budget_invalid", "Budget must be a positive integer.")
    selector, bounds = _selector(kind, target, sheet, address)
    manifest_bytes, manifest, reading_root, roots = _load_manifest(manifest_path)
    source = _source(manifest, source_id)
    selected_refs: list[ReadingFileRef] = []
    stream: list[tuple[str, ViewRecord]] = []
    summary: dict[str, Any] = {"label": "derivation", "kind": kind}
    diagnostics: list[dict[str, Any]] = []
    source_file_refs: list[ReadingFileRef] = []
    source_bytes = 0
    status = "ok"
    actual_sheet = selector.get("sheet")
    requested_target = selector.get("target")

    if kind == "overview":
        sheets = []
        for name, ref in source.lookup.sheet_views.items():
            excluded = _is_excluded(source, name)
            sheets.append({"sheet": name, "availability": "excluded" if excluded else "available",
                           "excluded": excluded, "view_id": ref.view_id})
        summary = {"label": "derivation", "source": {"source_id": source.source_id, "source_path": source.source_path,
                    "source_sha256": source.source_sha256, "run_id": source.run_id, "status": source.status,
                    "revision_id": manifest.revision_id}, "sheets": sheets,
                    "sheet_count": len(sheets), "retained_sheets": source.retained_sheets,
                    "excluded_sheets": source.excluded_sheets, "allowed_dependency_ranges": source.allowed_dependency_ranges,
                    "view_count": len(source.views)}
    elif kind == "sheet":
        actual_sheet = next((name for name in source.lookup.sheet_views if name.casefold() == str(actual_sheet).casefold()), None)
        if actual_sheet is None:
            raise QueryFailure("unavailable", "sheet_unavailable", "Selected source has no prepared view for this sheet.", sheet=selector.get("sheet"))
        if _is_excluded(source, actual_sheet):
            status = "excluded"
            summary = {"label": "derivation", "sheet": actual_sheet, "availability": "excluded",
                       "allowed_dependency_ranges": [item for item in source.allowed_dependency_ranges if item.get("sheet_name", "").casefold() == actual_sheet.casefold()]}
        else:
            selected_refs = [_sheet_ref(source, actual_sheet)] if _sheet_ref(source, actual_sheet) else []
            if not selected_refs:
                raise QueryFailure("unavailable", "sheet_view_unavailable", "Prepared sheet view is unavailable.", sheet=actual_sheet)
    elif kind in {"cell", "range"}:
        if not actual_sheet:
            raise QueryFailure("unsupported", "sheet_required", f"{kind} query requires --sheet or a Sheet!A1 target.")
        sheet_entry = next((name for name in source.lookup.sheet_views if name.casefold() == actual_sheet.casefold()), None)
        if sheet_entry is None:
            raise QueryFailure("unavailable", "sheet_unavailable", "Selected source has no prepared view for this sheet.", sheet=actual_sheet)
        actual_sheet = sheet_entry
        if bounds is None:
            raise QueryFailure("unsupported", "selector_address_invalid", "Cell/range selector needs an A1 address.")
        allowed = _scope_guard(source, actual_sheet, bounds, selector)
        excluded = _is_excluded(source, actual_sheet)
        if excluded:
            candidate_addresses = [address for address, allowed_bounds in allowed if _contains(allowed_bounds, bounds)]
            if candidate_addresses:
                ref = _sheet_ref(source, actual_sheet)
                selected_refs = [ref] if ref else []
                if not selected_refs:
                    raise QueryFailure("unavailable", "approved_sheet_view_unavailable", "No prepared canonical view covers the approved target.", sheet=actual_sheet, selector=selector)
        else:
            ref = _sheet_ref(source, actual_sheet)
            if ref:
                selected_refs = [ref]
        if not selected_refs:
            raise QueryFailure("unavailable", "selected_view_unavailable", "No prepared view covers the requested address.", sheet=actual_sheet)
    elif kind == "name":
        candidates = [item for item in source.lookup.defined_names if item.name.casefold() == str(requested_target).casefold()]
        if selector.get("sheet"):
            local = [item for item in candidates if item.scope.casefold() == str(selector["sheet"]).casefold()]
            workbook = [item for item in candidates if item.scope.casefold() == "workbook"]
            candidates = local or workbook
        if len(candidates) > 1:
            status = "needs_selection"
            diagnostics.append({"code": "name_ambiguous", "severity": "error",
                                "message": "The selected name has multiple source declarations; select a scope or target."})
            candidate_rows = []
            for item in candidates:
                owner = item.source_location.sheet_name
                if owner and _is_excluded(source, owner):
                    try:
                        _scope_guard(source, owner, _normalize_address(item.address)[1], selector)
                    except QueryFailure:
                        candidate_rows.append({"name": item.name, "scope": item.scope, "node_id": item.node_id,
                                               "availability": "excluded", "sheet": owner, "requires_scope": True,
                                               "allowed_targets": [{"sheet": owner, "range": address}
                                                                   for address, _ in _allowed_ranges(source, owner)]})
                        continue
                candidate_rows.append(item.model_dump(mode="json"))
            summary = {"label": "derivation", "candidates": candidate_rows, "needs_selection": True}
        elif not candidates:
            raise QueryFailure("not_found", "name_not_found", "No matching defined name exists in the selected source.", target=requested_target)
        else:
            name = candidates[0]
            name_sheet = name.source_location.sheet_name
            if not name_sheet:
                raise QueryFailure("needs_scope_resolution", "name_destination_unknown", "Defined-name destination has no worksheet ownership.", target=requested_target)
            normalized, name_bounds = _normalize_address(name.address)
            _scope_guard(source, name_sheet, name_bounds, selector)
            if _is_excluded(source, name_sheet):
                allowed = [address for address, box in _allowed_ranges(source, name_sheet) if _contains(box, name_bounds)]
                if not allowed:
                    raise QueryFailure("out_of_scope", "name_out_of_scope", "Defined name resolves outside the approved dependency ranges.", requested_selector=selector,
                                       allowed_targets=[{"sheet": name_sheet, "range": address} for address, _ in _allowed_ranges(source, name_sheet)])
                ref = _sheet_ref(source, name_sheet)
                selected_refs = [ref] if ref else []
            else:
                ref = _sheet_ref(source, name_sheet)
                selected_refs = [ref] if ref else []
            if ref is None:
                raise QueryFailure("unavailable", "name_view_unavailable", "Defined-name target has no prepared evidence view.", target=requested_target)
            summary = {"label": "derivation", "declaration": name.model_dump(mode="json"), "resolved_target": {"sheet": name_sheet, "range": normalized}}
    elif kind == "control":
        rows = _control_rows(source, roots, manifest)
        source_bytes += sum({ref.path: size for _, ref, _, size in rows}.values())
        matches = [(row, ref, record_type, size) for row, ref, record_type, size in rows if _matches_control(str(requested_target), row)]
        if selector.get("sheet"):
            matches = [item for item in matches if str(item[0].get("sheet", "")).casefold() == selector["sheet"].casefold()]
        sheets = sorted({str(row.get("sheet")) for row, _, _, _ in matches if row.get("sheet")})
        if len(sheets) > 1 or len(matches) > 1:
            status = "needs_selection"
            diagnostics.append({"code": "control_ambiguous", "severity": "error",
                                "message": "The selected control name has multiple source candidates; select a worksheet."})
            summary = {"label": "derivation", "candidates": [
                {"sheet": row.get("sheet"), "record_type": record_type,
                 "record": {"sheet": row.get("sheet"), "shape_id": row.get("shape_id"),
                            "control_name": row.get("control_name"), "binding_status": row.get("binding_status"),
                            "excluded": _is_excluded(source, row.get("sheet"))}}
                for row, _, record_type, _ in matches
            ], "needs_selection": True}
        elif not matches:
            raise QueryFailure("not_found", "control_not_found", "No matching control exists in the selected source.", target=requested_target)
        elif not sheets:
            raise QueryFailure("needs_scope_resolution", "control_ownership_unknown", "Control ownership is missing from its canonical record.", target=requested_target)
        else:
            row, artifact_ref, record_type, _size = matches[0]
            source_file_refs.append(artifact_ref)
            actual_sheet = sheets[0]
            if _is_excluded(source, actual_sheet):
                status = "excluded"
                summary = {"label": "derivation", "control": {"sheet": actual_sheet, "record_type": record_type,
                           "name": row.get("control_name") or row.get("shape_id"), "binding_status": row.get("binding_status"),
                           "excluded": True}, "available": True}
            else:
                ref = _sheet_ref(source, actual_sheet)
                if not ref:
                    raise QueryFailure("unavailable", "control_sheet_view_unavailable", "Control sheet has no prepared canonical view.", sheet=actual_sheet)
                selected_refs = [ref]
                summary = {"label": "derivation", "control_sheet": actual_sheet, "control_type": record_type}
    elif kind == "vba":
        handoff_data = _vba_handoff(source, roots, manifest)
        if handoff_data is None:
            raise QueryFailure("unavailable", "vba_unavailable", "Source has no canonical VBA handoff artifact.", target=requested_target)
        handoff, handoff_ref, handoff_size = handoff_data
        source_bytes += handoff_size
        modules = [item for item in handoff.get("modules", []) if isinstance(item, dict) and str(item.get("name", "")).casefold() == str(requested_target).casefold()]
        if len(modules) > 1:
            status = "needs_selection"
            diagnostics.append({"code": "vba_ambiguous", "severity": "error",
                                "message": "The selected module name has multiple source candidates."})
            summary = {"label": "derivation", "candidates": modules, "needs_selection": True}
        elif not modules:
            raise QueryFailure("not_found", "vba_module_not_found", "No matching VBA module exists in the selected source.", target=requested_target)
        else:
            module = modules[0]
            control_rows = _control_rows(source, roots, manifest) if module.get("kind") in {"Document", "ClassModule"} and source.excluded_sheets else []
            source_bytes += sum({ref.path: size for _, ref, _, size in control_rows}.values())
            owner = _module_owner(module, source, control_rows)
            if module.get("kind") in {"Document", "ClassModule"} and module.get("name") != "ThisWorkbook" and owner is None and source.excluded_sheets:
                status = "needs_scope_resolution"
                diagnostics.append({"code": "vba_ownership_unknown", "severity": "error",
                                    "message": "Worksheet VBA ownership is unresolved; select or resolve the module's worksheet scope.",
                                    "selector": selector})
                summary = {"label": "derivation", "module": {"name": module.get("name"), "kind": module.get("kind"),
                           "procedures": module.get("procedures", []), "ownership": "unknown"}, "needs_scope_resolution": True}
            elif owner and _is_excluded(source, owner):
                status = "excluded"
                summary = {"label": "derivation", "module": {"name": module.get("name"), "kind": module.get("kind"),
                           "procedures": module.get("procedures", []), "owner_sheet": owner, "excluded": True}, "available": True}
            else:
                target_ref = _workbook_ref(source)
                if not target_ref:
                    raise QueryFailure("unavailable", "vba_view_unavailable", "Prepared workbook view is unavailable.", target=requested_target)
                selected_refs = [target_ref]
                source_file = module.get("source_file")
                expected_path = (PurePosixPath(handoff_ref.path).parent / str(source_file)).as_posix()
                source_ref = next((ref for ref in source.artifacts.values() if ref.root == "step1" and ref.path == expected_path), None)
                if source_ref is None:
                    raise QueryFailure("unavailable", "vba_source_file_missing", "VBA module source file has no canonical manifest reference.", target=requested_target, source_file=source_file)
                source_data = _read_ref(source_ref, roots, manifest)
                if _sha(source_data) != module.get("sha256"):
                    raise QueryFailure("integrity_failed", "vba_source_mismatch", "Canonical VBA source file does not match the selected handoff module SHA-256.", target=requested_target, path=source_ref.path)
                source_bytes += len(source_data)
                source_file_refs.append(source_ref)
                summary = {"label": "derivation", "module": module, "source_file": source_ref.model_dump(mode="json"), "owner_sheet": owner}
    elif kind == "feature":
        if actual_sheet and _is_excluded(source, actual_sheet):
            status = "excluded"
            summary = {"label": "derivation", "sheet": actual_sheet, "availability": "excluded",
                       "allowed_dependency_ranges": [item for item in source.allowed_dependency_ranges if item.get("sheet_name", "").casefold() == actual_sheet.casefold()]}
        elif actual_sheet:
            ref = _sheet_ref(source, actual_sheet)
            if ref:
                selected_refs = [ref]
        else:
            ref = _workbook_ref(source)
            if ref:
                selected_refs = [ref]
    
    selected = _load_views(selected_refs, source, manifest, roots) if selected_refs else []
    if kind == "overview":
        pass
    elif kind == "sheet" and selected:
        view = selected[0][1]
        counts: dict[str, int] = {}
        for record in view.records:
            counts[record.record_type] = counts.get(record.record_type, 0) + 1
        summary = {"label": "derivation", "sheet": actual_sheet, "record_count": len(view.records), "record_types": counts}
        stream.extend((view.view_id, record) for record in view.records)
    elif kind in {"cell", "range"}:
        for _ref, view, _, _ in selected:
            for record in view.records:
                loc_sheet = record.source_location.sheet_name
                if not loc_sheet or loc_sheet.casefold() != str(actual_sheet).casefold():
                    continue
                if kind == "cell":
                    if record.record_type == "cell" and (record.source_location.address or "").replace("$", "").upper() == requested_target:
                        stream.append((view.view_id, record))
                elif _in_area(record, bounds) and not _is_excluded_form_control(source, record):
                    stream.append((view.view_id, record))
        if kind == "cell" and not stream:
            raise QueryFailure("not_found", "cell_missing", "No stored canonical cell record exists at the requested address; missing cells are not synthesized.", sheet=actual_sheet, address=requested_target)
        if kind == "range":
            summary = {"label": "derivation", "sheet": actual_sheet, "range": requested_target,
                       "record_count": len(stream), "missing_cells_not_synthesized": True}
    elif kind == "name" and stream == [] and selected:
        name = summary.get("declaration", {})
        name_sheet = name.get("source_location", {}).get("sheet_name")
        name_bounds = _normalize_address(name.get("address", ""))[1]
        for _ref, view, _, _ in selected:
            for record in view.records:
                if (record.source_location.sheet_name and record.source_location.sheet_name.casefold() == str(name_sheet).casefold()
                        and _in_area(record, name_bounds) and not _is_excluded_form_control(source, record)):
                    stream.append((view.view_id, record))
        summary["record_count"] = len(stream)
    elif kind == "control" and selected:
        row, _, record_type, _size = matches[0]
        view = selected[0][1]
        for record in view.records:
            if record.record_type == record_type and record.facts.get("record") == row:
                stream.append((view.view_id, record))
        if not stream:
            raise QueryFailure("unavailable", "control_record_unavailable", "Canonical view has no complete control record matching its source artifact.", target=requested_target)
        summary["record_count"] = len(stream)
    elif kind == "vba" and selected:
        module = summary.get("module", {})
        view = selected[0][1]
        stream.extend((view.view_id, record) for record in view.records if record.record_type == "vba_module" and record.facts.get("name", "").casefold() == str(requested_target).casefold())
        if not stream:
            raise QueryFailure("unavailable", "vba_record_unavailable", "Canonical workbook view has no complete module record matching the VBA handoff.", target=requested_target)
        summary["record_count"] = len(stream)
    elif kind == "feature" and selected:
        wanted = str(requested_target).casefold() if requested_target else None
        for _ref, view, _, _ in selected:
            for record in view.records:
                feature_type = str(record.facts.get("feature_type", "")).casefold()
                if record.record_type in {"unsupported", "unresolved_reference"} and (wanted is None or wanted in {feature_type, record.record_type, record.record_id.casefold()}):
                    stream.append((view.view_id, record))
        if not stream:
            raise QueryFailure("not_found", "feature_not_found", "No matching prepared feature record exists in the selected evidence view.", target=requested_target, sheet=actual_sheet)
        summary = {"label": "derivation", "feature_count": len(stream), "record_types": sorted({record.record_type for _, record in stream})}

    if kind == "sheet" and not stream:
        # Empty prepared sheet views remain valid evidence, and return their summary.
        pass
    if kind in {"cell", "range", "name", "control", "vba", "feature"} and not stream and status == "ok":
        raise QueryFailure("not_found", "records_not_found", "The selected prepared view contains no matching records.", selector=selector)
    return _finish_packet(
        manifest_path=Path(manifest_path).expanduser().resolve(), manifest_bytes=manifest_bytes, manifest=manifest,
        source=source, selector=selector, status=status, selected=selected, stream=stream,
        budget=budget, cursor=cursor, summary=summary, diagnostics=diagnostics,
        source_file_refs=source_file_refs, source_bytes=source_bytes,
    )


def validate_evidence_packet(packet_value: Any, output: dict[str, Any]) -> dict[str, Any]:
    """Verify canonical files and delivered records before validating claims."""
    from excel_to_act.views import validate_agent_output

    try:
        if not isinstance(packet_value, dict):
            raise QueryFailure("integrity_failed", "packet_invalid", "Evidence packet must be a JSON object.")
        if not isinstance(output, dict):
            raise QueryFailure("integrity_failed", "agent_output_invalid", "Agent output must be a JSON object.")
        if packet_value.get("trace") is not None:
            from excel_to_act.steps.step2.trace import validate_trace_packet
            EvidencePacket.model_validate(packet_value)
            return validate_trace_packet(packet_value, output)
        raw_manifest = packet_value.get("manifest")
        raw_source = packet_value.get("source")
        raw_selector = packet_value.get("selector")
        if not all(isinstance(item, dict) for item in (raw_manifest, raw_source, raw_selector)):
            raise QueryFailure("integrity_failed", "packet_invalid", "Evidence packet needs manifest, source, and selector objects.")
        manifest_path = Path(raw_manifest.get("path", "")).expanduser().resolve()
        manifest_bytes, manifest, _reading_root, roots = _load_manifest(manifest_path)
        if _sha(manifest_bytes) != raw_manifest.get("sha256") or len(manifest_bytes) != raw_manifest.get("bytes"):
            raise QueryFailure("integrity_failed", "manifest_hash_mismatch", "Packet manifest does not match its recorded canonical hash.")
        if raw_manifest.get("reading_root") != manifest.roots.get("reading", "."):
            raise QueryFailure("integrity_failed", "reading_root_mismatch", "Packet reading-root binding differs from the canonical manifest.")
        source = _source(manifest, str(raw_source.get("source_id", "")))
        expected_source = {
            "source_id": source.source_id, "source_path": source.source_path,
            "source_sha256": source.source_sha256, "run_id": source.run_id, "status": source.status,
            "revision_id": manifest.revision_id, "scope_sha256": source.scope_sha256,
            "retained_sheets": source.retained_sheets, "excluded_sheets": source.excluded_sheets,
            "allowed_dependency_ranges": source.allowed_dependency_ranges,
        }
        if _json_bytes(raw_source) != _json_bytes(expected_source):
            raise QueryFailure("integrity_failed", "packet_source_mismatch", "Packet source or scope fields differ from the canonical manifest.")

        raw_refs = packet_value.get("canonical_view_refs")
        raw_source_refs = packet_value.get("source_file_refs", [])
        raw_pages = packet_value.get("views")
        if not isinstance(raw_refs, list) or not isinstance(raw_source_refs, list) or not isinstance(raw_pages, list):
            raise QueryFailure("integrity_failed", "packet_records_invalid", "Packet references and views must be lists.")
        source_view_json = [ref.model_dump(mode="json") for ref in source.views]
        if len(raw_refs) != len({item.get("path") for item in raw_refs if isinstance(item, dict)}):
            raise QueryFailure("integrity_failed", "canonical_view_ref_duplicate", "Packet repeats a canonical view reference.")
        for item in raw_refs:
            if not isinstance(item, dict) or not any(_json_bytes(item) == _json_bytes(ref) for ref in source_view_json):
                raise QueryFailure("integrity_failed", "canonical_view_ref_invalid", "Packet canonical reference differs from the selected source manifest.")
        artifact_json = [ref.model_dump(mode="json") for ref in source.artifacts.values()]
        for item in raw_source_refs:
            if not isinstance(item, dict) or not any(_json_bytes(item) == _json_bytes(ref) for ref in artifact_json):
                raise QueryFailure("integrity_failed", "source_file_ref_invalid", "Packet source-file reference differs from the selected source manifest.")
            _read_ref(ReadingFileRef.model_validate(item), roots, manifest)

        refs = [ReadingFileRef.model_validate(item) for item in raw_refs]
        loaded = _load_views(refs, source, manifest, roots)
        canonical_by_id = {view.view_id: (ref, raw) for ref, view, raw, _ in loaded}
        canonical_by_path = {ref.path: raw for ref, _, raw, _ in loaded}
        slices_raw: dict[str, dict[str, Any]] = {}
        for page in raw_pages:
            if not isinstance(page, dict) or not isinstance(page.get("records"), list):
                raise QueryFailure("integrity_failed", "delivered_view_invalid", "Delivered view slice is not a record-bearing object.")
            view_id = page.get("view_id")
            if view_id in slices_raw or view_id not in canonical_by_id:
                raise QueryFailure("integrity_failed", "delivered_view_invalid", "Delivered slice has an unknown or duplicate view ID.", view_id=view_id)
            _ref, canonical_raw = canonical_by_id[view_id]
            raw_header = {key: value for key, value in canonical_raw.items() if key not in {"records", "chunks", "opaque_report"}}
            slice_header = {key: value for key, value in page.items() if key not in {"records", "chunks", "opaque_report"}}
            if _json_bytes(raw_header) != _json_bytes(slice_header):
                raise QueryFailure("integrity_failed", "delivered_view_header_changed", "Delivered view identity differs from its canonical view.", view_id=view_id)
            canonical_raw_records = canonical_raw.get("records", [])
            canonical_records = {item.get("record_id"): item for item in canonical_raw_records if isinstance(item, dict)}
            seen: set[str] = set()
            for record in page["records"]:
                if not isinstance(record, dict) or not isinstance(record.get("record_id"), str):
                    raise QueryFailure("integrity_failed", "delivered_record_invalid", "Delivered records must retain canonical record IDs.", view_id=view_id)
                record_id = record["record_id"]
                if record_id in seen:
                    raise QueryFailure("integrity_failed", "delivered_record_duplicate", "A delivered slice repeats a record ID within one view.", view_id=view_id, record_id=record_id)
                seen.add(record_id)
                exact = canonical_records.get(record_id)
                if exact is None or _json_bytes(exact) != _json_bytes(record):
                    raise QueryFailure("integrity_failed", "delivered_record_changed", "A delivered full record differs from its canonical record.", view_id=view_id, record_id=record_id)
            record_ids = {item["record_id"] for item in page["records"]}
            if any(not isinstance(chunk, dict) or any(record_id not in record_ids for record_id in chunk.get("record_ids", [])) for chunk in page.get("chunks", [])):
                raise QueryFailure("integrity_failed", "page_chunk_invalid", "Page chunks may refer only to delivered records.", view_id=view_id)
            if any(record_id not in record_ids for record_id in page.get("opaque_report", [])):
                raise QueryFailure("integrity_failed", "page_opaque_report_invalid", "Opaque report may refer only to delivered records.", view_id=view_id)
            slices_raw[view_id] = page
        actual_delivered = {view_id: [item["record_id"] for item in page["records"]] for view_id, page in slices_raw.items()}
        if _json_bytes(actual_delivered) != _json_bytes(packet_value.get("delivered_record_ids")):
            raise QueryFailure("integrity_failed", "delivered_ids_mismatch", "Packet delivered IDs do not match the full records it contains.")
        packet = EvidencePacket.model_validate(packet_value)
        selector = packet.selector.model_dump(mode="json")
        if _json_bytes(raw_selector) != _json_bytes(selector):
            raise QueryFailure("integrity_failed", "selector_invalid", "Packet selector does not preserve its exact normalized fields.")
        pages = [WorkbookView.model_validate(page) for page in slices_raw.values()]
        kind = selector["kind"]
        expected_refs, expected_stream = _expected_selection(selector, source, roots, manifest, canonical_by_path)
        if _json_bytes([ref.model_dump(mode="json") for ref in expected_refs]) != _json_bytes(raw_refs):
            raise QueryFailure("integrity_failed", "selector_reference_mismatch", "Packet canonical references do not match the selector's prepared views.")
        binding_data = {
            "manifest_sha256": _sha(manifest_bytes), "source_id": source.source_id,
            "selector": selector,
            "canonical_views": [(ref.path, ref.sha256) for ref in expected_refs],
            "stream": expected_stream,
        }
        expected_binding = _sha(json.dumps(binding_data, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"))
        if packet.pagination.get("binding_sha256") != expected_binding:
            raise QueryFailure("integrity_failed", "selector_binding_mismatch", "Packet selector or record stream binding is altered.")
        page_pairs = [(page["view_id"], record["record_id"]) for page in raw_pages for record in page["records"]]
        pagination = packet.pagination
        budget = pagination.get("budget")
        start = pagination.get("start_index")
        if type(budget) is not int or budget < 1 or type(start) is not int or not 0 <= start <= len(expected_stream):
            raise QueryFailure("integrity_failed", "pagination_invalid", "Packet page budget or start index is invalid.")
        end = start + len(page_pairs)
        next_cursor = pagination.get("next_cursor")
        if end > len(expected_stream) or page_pairs != expected_stream[start:end]:
            raise QueryFailure("integrity_failed", "pagination_records_mismatch", "Delivered page is not the bound contiguous record prefix.")
        if type(pagination.get("delivered_count")) is not int or pagination.get("delivered_count") != len(page_pairs):
            raise QueryFailure("integrity_failed", "pagination_count_mismatch", "Packet delivered count differs from its complete page records.")
        expected_delivered = [{"view_id": view_id, "record_id": record_id} for view_id, record_id in page_pairs]
        if _json_bytes(pagination.get("delivered")) != _json_bytes(expected_delivered):
            raise QueryFailure("integrity_failed", "pagination_ids_mismatch", "Packet page IDs differ from its delivered records.")
        if type(pagination.get("estimated_tokens")) is not int or pagination["estimated_tokens"] != sum(
            _tokens(record) for page in pages for record in page.records
        ) or pagination["estimated_tokens"] > budget:
            raise QueryFailure("integrity_failed", "pagination_budget_mismatch", "Packet page estimate differs from its complete records or exceeds its budget.")
        omitted = expected_stream[end:]
        expected_omitted = [{"view_id": view_id, "record_id": record_id} for view_id, record_id in omitted[:100]]
        if (pagination.get("omitted_count") != len(omitted)
                or _json_bytes(pagination.get("omitted")) != _json_bytes(expected_omitted)
                or pagination.get("omitted_ids_truncated") is not (len(omitted) > 100)):
            raise QueryFailure("integrity_failed", "pagination_omitted_mismatch", "Packet omitted IDs differ from the remaining canonical record stream.")
        if next_cursor is not None:
            try:
                cursor_data = json.loads(base64.urlsafe_b64decode(str(next_cursor) + "=" * (-len(str(next_cursor)) % 4)))
            except (ValueError, TypeError, json.JSONDecodeError) as exc:
                raise QueryFailure("integrity_failed", "pagination_cursor_invalid", "Packet next cursor is malformed.") from exc
            if not isinstance(cursor_data, dict):
                raise QueryFailure("integrity_failed", "pagination_cursor_invalid", "Packet next cursor must decode to a JSON object.")
            if (cursor_data.get("version") != 1 or cursor_data.get("binding") != expected_binding
                    or type(cursor_data.get("next_index")) is not int or cursor_data["next_index"] != end):
                raise QueryFailure("integrity_failed", "pagination_cursor_invalid", "Packet next cursor does not continue after its delivered prefix.")
        elif end != len(expected_stream):
            raise QueryFailure("integrity_failed", "pagination_cursor_missing", "Packet omits the next cursor while canonical records remain.")
        oversized = pagination.get("oversized")
        if oversized is not None:
            if not isinstance(oversized, dict) or end >= len(expected_stream):
                raise QueryFailure("integrity_failed", "pagination_oversized_invalid", "Oversized metadata has no next canonical record.")
            view_id, record_id = expected_stream[end]
            ref = next((item for item in refs if item.view_id == view_id), None)
            raw_record = next((item for item in canonical_by_id[view_id][1].get("records", []) if item.get("record_id") == record_id), None)
            required = _tokens(ViewRecord.model_validate(raw_record)) if raw_record else None
            if (oversized.get("view_id") != view_id or oversized.get("record_id") != record_id
                    or oversized.get("canonical_ref") != (ref.model_dump(mode="json") if ref else None)
                    or oversized.get("required_estimated_tokens") != required or required is None or required <= budget):
                raise QueryFailure("integrity_failed", "pagination_oversized_invalid", "Oversized metadata differs from the next indivisible canonical record.")
        if output.get("claims") and packet.status not in {"ok", "partial", "truncated"}:
            raise QueryFailure("integrity_failed", "packet_status_not_citable", "Packet status does not permit citation validation.", status=packet.status)
        if kind == "overview" and pages:
            raise QueryFailure("integrity_failed", "selector_record_mismatch", "Overview metadata cannot grant record citations.")
        if kind in {"cell", "range", "sheet"} and selector.get("sheet"):
            requested_sheet = selector["sheet"]
            if _is_excluded(source, requested_sheet) and kind == "sheet" and pages:
                raise QueryFailure("integrity_failed", "excluded_sheet_records_delivered", "Excluded-sheet metadata packets cannot contain worksheet records.")
            selector_bounds = None
            if kind in {"cell", "range"}:
                selector_bounds = _normalize_address(selector.get("target") or selector.get("range") or "", cell=kind == "cell")[1]
                if _is_excluded(source, requested_sheet) and kind in {"cell", "range"}:
                    _scope_guard(source, requested_sheet, selector_bounds, selector)
            for page in pages:
                for record in page.records:
                    if (record.source_location.sheet_name or "").casefold() != requested_sheet.casefold():
                        raise QueryFailure("integrity_failed", "selector_record_mismatch", "Delivered record is outside the packet's selected sheet.", record_id=record.record_id)
                    if kind == "cell" and (record.record_type != "cell" or (record.source_location.address or "").replace("$", "").upper() != selector.get("target")):
                        raise QueryFailure("integrity_failed", "selector_record_mismatch", "Delivered record does not match the selected cell.", record_id=record.record_id)
                    if selector_bounds and kind == "range":
                        if _is_excluded_form_control(source, record):
                            raise QueryFailure("integrity_failed", "excluded_control_record_delivered", "Excluded-sheet form-control facts cannot be delivered by a range query.", record_id=record.record_id)
                        if not _in_area(record, selector_bounds):
                            raise QueryFailure("integrity_failed", "selector_record_mismatch", "Delivered record is outside the selected range.", record_id=record.record_id)
        elif kind == "name":
            candidates = [item for item in source.lookup.defined_names if item.name.casefold() == str(selector.get("target", "")).casefold()]
            if selector.get("sheet"):
                local = [item for item in candidates if item.scope.casefold() == selector["sheet"].casefold()]
                candidates = local or [item for item in candidates if item.scope.casefold() == "workbook"]
            if pages and len(candidates) != 1:
                raise QueryFailure("integrity_failed", "name_selector_ambiguous", "A delivered name page must resolve to exactly one declaration.")
            if pages:
                declaration = candidates[0]
                name_sheet = declaration.source_location.sheet_name
                name_bounds = _normalize_address(declaration.address)[1]
                if _is_excluded(source, name_sheet or ""):
                    _scope_guard(source, name_sheet or "", name_bounds, selector)
                for page in pages:
                    for record in page.records:
                        if _is_excluded_form_control(source, record):
                            raise QueryFailure("integrity_failed", "excluded_control_record_delivered", "Excluded-sheet form-control facts cannot be delivered by a name query.", record_id=record.record_id)
                        if (record.source_location.sheet_name or "").casefold() != (name_sheet or "").casefold() or not _in_area(record, name_bounds):
                            raise QueryFailure("integrity_failed", "name_record_outside_target", "Delivered record is outside the selected defined-name destination.", record_id=record.record_id)
        elif kind in {"control", "vba", "feature"}:
            if kind == "feature" and selector.get("sheet") and _is_excluded(source, selector["sheet"]) and pages:
                raise QueryFailure("integrity_failed", "excluded_feature_records_delivered", "Excluded-sheet feature queries expose metadata only.")
            for page in pages:
                for record in page.records:
                    target_value = str(selector.get("target") or "").casefold()
                    if kind == "control" and (record.record_type not in {"checkbox_binding", "activex_event"} or not _matches_control(str(selector.get("target")), record.facts.get("record", {}))):
                        raise QueryFailure("integrity_failed", "selector_record_mismatch", "Delivered control record does not match the selected control.", record_id=record.record_id)
                    if kind == "control" and _is_excluded(source, record.source_location.sheet_name or ""):
                        raise QueryFailure("integrity_failed", "excluded_control_record_delivered", "Excluded controls expose metadata only.", record_id=record.record_id)
                    if kind == "control" and selector.get("sheet") and (record.source_location.sheet_name or "").casefold() != selector["sheet"].casefold():
                        raise QueryFailure("integrity_failed", "selector_record_mismatch", "Delivered control is outside the selected worksheet.", record_id=record.record_id)
                    if kind == "control":
                        artifact_ref = record.facts.get("artifact_ref", {})
                        if not any(ref.root == artifact_ref.get("root") and ref.path == artifact_ref.get("path") and ref.sha256 == artifact_ref.get("sha256") for ref in packet.source_file_refs):
                            raise QueryFailure("integrity_failed", "control_source_ref_missing", "Delivered control record must retain its canonical source artifact reference.", record_id=record.record_id)
                    if kind == "vba" and (record.record_type != "vba_module" or str(record.facts.get("name", "")).casefold() != target_value):
                        raise QueryFailure("integrity_failed", "selector_record_mismatch", "Delivered module record does not match the selected VBA module.", record_id=record.record_id)
                    if kind == "feature" and (record.record_type not in {"unsupported", "unresolved_reference"} or target_value and target_value not in {record.record_id.casefold(), record.record_type, str(record.facts.get("feature_type", "")).casefold()}):
                        raise QueryFailure("integrity_failed", "selector_record_mismatch", "Delivered feature does not match the selected feature.", record_id=record.record_id)
                    if kind == "feature" and selector.get("sheet") and (record.source_location.sheet_name or "").casefold() != selector["sheet"].casefold():
                        raise QueryFailure("integrity_failed", "selector_record_mismatch", "Delivered feature is outside the selected worksheet.", record_id=record.record_id)
            if kind == "vba" and (pages or packet.source_file_refs):
                handoff_data = _vba_handoff(source, roots, manifest)
                modules = [item for item in (handoff_data[0].get("modules", []) if handoff_data else [])
                           if isinstance(item, dict) and str(item.get("name", "")).casefold() == str(selector.get("target", "")).casefold()]
                if len(modules) != 1 or handoff_data is None:
                    raise QueryFailure("integrity_failed", "vba_module_binding_invalid", "Delivered VBA module has no unique canonical handoff entry.")
                handoff, handoff_ref, _ = handoff_data
                expected_path = (PurePosixPath(handoff_ref.path).parent / str(modules[0].get("source_file", ""))).as_posix()
                source_ref = next((ref for ref in source.artifacts.values() if ref.root == "step1" and ref.path == expected_path), None)
                if source_ref is None or not any(ref.path == source_ref.path and ref.sha256 == source_ref.sha256 for ref in packet.source_file_refs):
                    raise QueryFailure("integrity_failed", "vba_source_ref_missing", "Delivered module packet must include its canonical source-file reference.")
                if source_ref.sha256 != modules[0].get("sha256"):
                    raise QueryFailure("integrity_failed", "vba_source_mismatch", "Canonical VBA source file reference does not match the selected handoff module SHA-256.", path=source_ref.path)
                if any(hashlib.sha256(str(record.facts.get("code", "")).encode("utf-8")).hexdigest() != modules[0].get("sha256")
                       for page in pages for record in page.records):
                    raise QueryFailure("integrity_failed", "vba_source_mismatch", "Delivered module record code differs from the canonical source file hash.")
                control_rows = _control_rows(source, roots, manifest) if modules[0].get("kind") in {"Document", "ClassModule"} and source.excluded_sheets else []
                owner = _module_owner(modules[0], source, control_rows)
                if owner and _is_excluded(source, owner):
                    raise QueryFailure("integrity_failed", "excluded_worksheet_vba_delivered", "Excluded worksheet VBA exposes metadata only.")
                if not owner and modules[0].get("kind") in {"Document", "ClassModule"} and modules[0].get("name") != "ThisWorkbook" and source.excluded_sheets:
                    raise QueryFailure("integrity_failed", "vba_ownership_unknown", "Worksheet VBA ownership needs scope resolution before delivering code.")
        result = validate_agent_output(output, pages)
        return {"valid": True, "provenance_only": True, "claim_validation": result, "diagnostics": []}
    except (QueryFailure, ValueError, TypeError, KeyError) as exc:
        if isinstance(exc, QueryFailure):
            diagnostic = exc.diagnostic
        else:
            diagnostic = {"code": "packet_validation_failed", "severity": "error", "message": str(exc)}
        return {"valid": False, "provenance_only": True, "diagnostics": [diagnostic]}
