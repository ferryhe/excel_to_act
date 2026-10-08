"""Strict source-side OOXML facts and logical-object discovery for Step1."""

from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl.cell.text import Text
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel

from excel_to_act.ingest.data_table import resolve_package_part
from excel_to_act.ingest.ooxml_package import OPAQUE_MARKERS

DOC_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class SourceScanError(ValueError):
    """Required workbook XML could not be read or parsed."""


def object_identity(kind: str, *parts: str) -> str:
    """Stable, unambiguous identity for a workbook logical object."""

    return f"{kind}:{json.dumps(parts, ensure_ascii=False, separators=(',', ':'))}"


def _local(tag: str) -> str:
    return tag.split("}", 1)[-1] if "}" in tag else tag


def _required_xml(zf: zipfile.ZipFile, part: str) -> ET.Element:
    try:
        return ET.fromstring(zf.read(part))
    except (KeyError, ET.ParseError, OSError, zipfile.BadZipFile) as exc:
        raise SourceScanError(f"{part}: {exc}") from exc


def _validate_relationship_targets(root: ET.Element, rels_part: str, base: str, members: set[str]) -> None:
    for rel in root:
        if _local(rel.tag) != "Relationship" or rel.attrib.get("TargetMode") == "External":
            continue
        target = rel.attrib.get("Target", "")
        part = resolve_package_part(base, target) if target else ""
        if not target or part not in members:
            raise SourceScanError(
                f"{rels_part}: relationship {rel.attrib.get('Id', '')!r} "
                f"({rel.attrib.get('Type', '')!r}) has missing internal target "
                f"{target!r} (resolved package part {part!r})"
            )


def _rels(zf: zipfile.ZipFile, part: str, base: str, parsed_parts: set[str]) -> list[dict[str, str]]:
    rels_part = str(PurePosixPath(part).parent / "_rels" / f"{PurePosixPath(part).name}.rels")
    if rels_part not in zf.namelist():
        return []
    parsed_parts.add(rels_part)
    root = _required_xml(zf, rels_part)
    _validate_relationship_targets(root, rels_part, base, set(zf.namelist()))
    result = []
    for rel in root:
        if _local(rel.tag) != "Relationship":
            continue
        target = rel.attrib.get("Target", "")
        result.append(
            {
                "id": rel.attrib.get("Id", ""),
                "type": rel.attrib.get("Type", ""),
                "target": target,
                "part": "" if rel.attrib.get("TargetMode") == "External" else resolve_package_part(base, target),
            }
        )
    return result


def _content_types(zf: zipfile.ZipFile, parsed_parts: set[str]) -> dict[str, str]:
    parsed_parts.add("[Content_Types].xml")
    root = _required_xml(zf, "[Content_Types].xml")
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    for item in root:
        if _local(item.tag) == "Default":
            defaults[item.attrib.get("Extension", "")] = item.attrib.get("ContentType", "")
        elif _local(item.tag) == "Override":
            overrides[item.attrib.get("PartName", "").lstrip("/")] = item.attrib.get("ContentType", "")
    result = dict(overrides)
    for info in zf.infolist():
        if info.filename not in result:
            result[info.filename] = defaults.get(Path(info.filename).suffix.lstrip("."), "")
    return result


def _text(element: ET.Element | None) -> str | None:
    if element is None:
        return None
    return "".join(element.itertext())


def _rich_text_content(element: ET.Element, part: str) -> str:
    try:
        parsed = Text.from_tree(element)
    except Exception as exc:
        raise SourceScanError(f"{part}: openpyxl could not parse rich text: {exc}") from exc
    return parsed.content or ""


def _parse_scalar(raw: str | None, cell_type: str | None) -> Any:
    if raw is None:
        return None
    if cell_type == "b":
        return raw == "1"
    if cell_type in {"str", "e", "d"}:
        return raw
    if raw == "":
        return "" if cell_type in {"s", "inlineStr"} else None
    try:
        if not any(char in raw for char in ".eE"):
            return int(raw)
        return float(raw)
    except ValueError:
        return raw


def _resolve_shared_strings(zf: zipfile.ZipFile, workbook_rels: list[dict[str, str]], parsed_parts: set[str]) -> list[str]:
    part = next((r["part"] for r in workbook_rels if r["type"].endswith("/sharedStrings")), "")
    if not part:
        return []
    parsed_parts.add(part)
    root = _required_xml(zf, part)
    return [_rich_text_content(si, part) for si in root if _local(si.tag) == "si"]


def _shared_string_at(shared_strings: list[str], raw_index: str | None, *, part: str, address: str) -> str:
    try:
        if raw_index is None:
            index = -1
        else:
            decimal = raw_index.strip()
            digits = decimal[1:] if decimal[:1] in {"+", "-"} else decimal
            if not digits or not digits.isdecimal():
                raise ValueError(f"not a decimal integer: {raw_index!r}")
            index = int(decimal)
    except ValueError as exc:
        raise SourceScanError(f"{part}: cell {address} has invalid shared-string index {raw_index!r}") from exc
    if index < 0 or index >= len(shared_strings):
        raise SourceScanError(f"{part}: cell {address} has out-of-range shared-string index {raw_index!r}")
    return shared_strings[index]


def scan_step1_source(path: Path) -> dict[str, Any]:
    """Read raw cell facts and independently enumerate supported logical objects.

    A malformed required XML part is an error. It must never be converted to an
    empty workbook or counted as zero objects.
    """

    source_path = path.expanduser().resolve()
    with zipfile.ZipFile(source_path) as zf:
        file_infos = [info for info in zf.infolist() if not info.is_dir()]
        names = [info.filename for info in file_infos]
        if len(names) != len(set(names)):
            raise SourceScanError("OOXML package contains duplicate part names")
        parsed_parts: set[str] = set()
        content_types = _content_types(zf, parsed_parts)
        root_rels_part = "_rels/.rels"
        if root_rels_part not in names:
            raise SourceScanError(f"package root relationships are missing: {root_rels_part}")
        root_rels = _required_xml(zf, root_rels_part)
        parsed_parts.add(root_rels_part)
        _validate_relationship_targets(root_rels, root_rels_part, "", set(names))
        office_document = next(
            (
                item
                for item in root_rels
                if _local(item.tag) == "Relationship"
                and item.attrib.get("Type", "").endswith("/officeDocument")
                and item.attrib.get("TargetMode") != "External"
            ),
            None,
        )
        if office_document is None or resolve_package_part("", office_document.attrib.get("Target", "")) != "xl/workbook.xml":
            raise SourceScanError(f"{root_rels_part}: missing relationship to xl/workbook.xml")
        # Standard workbook metadata and theme XML are known support parts. They
        # affect presentation/metadata but are not workbook logical objects.
        for support_part in ("docProps/app.xml", "docProps/core.xml", "xl/theme/theme1.xml"):
            if support_part in names:
                _required_xml(zf, support_part)
                parsed_parts.add(support_part)
        parsed_parts.add("xl/workbook.xml")
        workbook = _required_xml(zf, "xl/workbook.xml")
        workbook_rels = _rels(zf, "xl/workbook.xml", "xl", parsed_parts)
        rel_by_id = {rel["id"]: rel for rel in workbook_rels}
        properties = next((item for item in workbook if _local(item.tag) == "workbookPr"), None)
        date_system = "1904" if properties is not None and properties.attrib.get("date1904", "0").lower() in {"1", "true"} else "1900"
        epoch = CALENDAR_MAC_1904 if date_system == "1904" else CALENDAR_WINDOWS_1900
        shared_strings = _resolve_shared_strings(zf, workbook_rels, parsed_parts)
        styles_part = next((r["part"] for r in workbook_rels if r["type"].endswith("/styles")), "")
        style_dates: list[bool] = []
        style_formats: list[str] = []
        if styles_part:
            parsed_parts.add(styles_part)
            styles_root = _required_xml(zf, styles_part)
            custom_formats = {
                int(item.attrib["numFmtId"]): item.attrib.get("formatCode", "")
                for item in styles_root.iter()
                if _local(item.tag) == "numFmt" and item.attrib.get("numFmtId", "").isdigit()
            }
            builtin_date_ids = {14, 15, 16, 17, 18, 19, 20, 21, 22, 27, 28, 29, 30, 31, 32, 33, 34, 35, 36, 45, 46, 47, 50, 51, 52, 53, 54, 55, 56, 57, 58}
            xfs = next((item for item in styles_root if _local(item.tag) == "cellXfs"), None)
            if xfs is not None:
                for xf in xfs:
                    num_fmt = int(xf.attrib.get("numFmtId", "0"))
                    style_dates.append(num_fmt in builtin_date_ids or is_date_format(custom_formats.get(num_fmt, "")))
                    style_formats.append(custom_formats.get(num_fmt, BUILTIN_FORMATS.get(num_fmt, "General")))

        sheets: list[dict[str, Any]] = []
        cells: list[dict[str, Any]] = []
        objects: list[dict[str, Any]] = []
        errors: list[str] = []
        supported_vml_parts: set[str] = set()

        def add_object(kind: str, identity: str, part: str, **details: Any) -> None:
            objects.append(
                {
                    "identity": identity,
                    "kind": kind,
                    "part": part,
                    "source_location": {"workbook_path": str(source_path), "ooxml_part": part, "object_type": kind, "object_id": identity},
                    "details": details,
                }
            )

        sheets_element = next((item for item in workbook.iter() if _local(item.tag) == "sheets"), None)
        for sheet_index, sheet_element in enumerate(list(sheets_element or [])):
            if _local(sheet_element.tag) != "sheet":
                continue
            name = sheet_element.attrib.get("name", "")
            if not name:
                raise SourceScanError("xl/workbook.xml: worksheet is missing its required name")
            sheet_id = sheet_element.attrib.get("sheetId", "")
            try:
                int(sheet_id)
            except ValueError as exc:
                raise SourceScanError(
                    f"xl/workbook.xml: worksheet {name!r} has invalid sheetId {sheet_id!r}"
                ) from exc
            relation = rel_by_id.get(sheet_element.attrib.get(f"{DOC_REL}id", ""), {})
            part = relation.get("part", "")
            if not part or part not in names:
                raise SourceScanError(f"xl/workbook.xml: worksheet {name!r} has no readable relationship target")
            root = _required_xml(zf, part)
            parsed_parts.add(part)
            sheet_id = object_identity("sheet", part, name)
            sheets.append({"name": name, "index": sheet_index, "part": part, "state": sheet_element.attrib.get("state", "visible"), "identity": sheet_id})
            add_object("sheet", sheet_id, part, name=name, index=sheet_index)
            sheet_rels = _rels(zf, part, str(PurePosixPath(part).parent), parsed_parts)
            table_parts = [r["part"] for r in sheet_rels if r["type"].endswith("/table") and r["part"]]

            # Formula masters are expanded by reference for a source-level normalized formula.
            shared_formulas: dict[str, tuple[str, str]] = {}
            for cell in root.iter():
                if _local(cell.tag) != "c":
                    continue
                formula_node = next((child for child in cell if _local(child.tag) == "f"), None)
                if formula_node is not None and formula_node.attrib.get("t") == "shared" and formula_node.text:
                    shared_formulas[formula_node.attrib.get("si", "")] = (cell.attrib.get("r", ""), formula_node.text)

            for cell_node in root.iter():
                if _local(cell_node.tag) != "c":
                    continue
                ref = cell_node.attrib.get("r", "")
                formula_node = next((child for child in cell_node if _local(child.tag) == "f"), None)
                value_node = next((child for child in cell_node if _local(child.tag) == "v"), None)
                inline_node = next((child for child in cell_node if _local(child.tag) == "is"), None)
                if formula_node is None and value_node is None and inline_node is None:
                    continue
                cell_type = cell_node.attrib.get("t")
                raw_value = _text(value_node)
                inline_text = _text(inline_node)
                normalized_inline_text = _rich_text_content(inline_node, part) if inline_node is not None else None
                raw_formula = formula_node.text if formula_node is not None else None
                formula_attrs = dict(formula_node.attrib) if formula_node is not None else {}
                formula = None
                if formula_node is not None:
                    formula_text = raw_formula or ""
                    if formula_attrs.get("t") == "shared" and not formula_text:
                        master = shared_formulas.get(formula_attrs.get("si", ""))
                        if master:
                            try:
                                formula_text = Translator(f"={master[1]}", origin=master[0]).translate_formula(ref)
                            except (ValueError, TypeError):
                                formula_text = ""
                    if formula_attrs.get("t") == "dataTable":
                        formula = f"={formula_text}" if formula_text and not formula_text.startswith("=") else (formula_text or None)
                    else:
                        formula = f"={formula_text}" if formula_text and not formula_text.startswith("=") else (formula_text or None)
                    if formula_attrs.get("t") == "dataTable":
                        add_object(
                            "data_table",
                            object_identity("dataTable", part, ref),
                            part,
                            sheet=name,
                            address=formula_attrs.get("ref") or ref,
                            corner_cell=ref,
                            formula=formula,
                            formula_attributes=formula_attrs,
                        )
                style_index = int(cell_node.attrib.get("s", "0")) if cell_node.attrib.get("s", "0").isdigit() else 0
                numeric_source_value = _parse_scalar(raw_value, cell_type) if cell_type in {None, "n"} else None
                has_numeric_source_value = isinstance(numeric_source_value, int | float) and not isinstance(numeric_source_value, bool)
                parsed_value: Any
                if inline_node is not None:
                    parsed_value = normalized_inline_text or ""
                elif cell_type == "s" and value_node is not None:
                    parsed_value = _shared_string_at(shared_strings, raw_value, part=part, address=ref)
                else:
                    parsed_value = _parse_scalar(raw_value, cell_type)
                if formula is None and isinstance(parsed_value, int | float) and not isinstance(parsed_value, bool) and style_index < len(style_dates) and style_dates[style_index]:
                    try:
                        parsed_value = from_excel(float(raw_value or "0"), epoch=epoch).isoformat()
                    except (ValueError, OverflowError, AttributeError):
                        pass
                cached_value = _parse_scalar(raw_value, cell_type) if formula_node is not None and value_node is not None else None
                if formula_node is not None and value_node is not None and cell_type == "s":
                    cached_value = _shared_string_at(shared_strings, raw_value, part=part, address=ref)
                if formula_node is not None and value_node is not None and cell_type == "inlineStr":
                    cached_value = normalized_inline_text or ""
                if formula_node is not None and isinstance(cached_value, int | float) and not isinstance(cached_value, bool) and style_index < len(style_dates) and style_dates[style_index]:
                    try:
                        cached_value = from_excel(float(raw_value or "0"), epoch=epoch).isoformat()
                    except (ValueError, OverflowError, AttributeError):
                        pass
                cached_available = formula_node is not None and value_node is not None and (raw_value not in {None, ""} or cell_type in {"str", "inlineStr"})
                identity = object_identity("cell", part, ref)
                cells.append(
                    {
                        "identity": identity,
                        "sheet": name,
                        "sheet_part": part,
                        "address": ref,
                        "cell_type": cell_type,
                        "style_id": style_index,
                        "number_format": style_formats[style_index] if style_index < len(style_formats) else "General",
                        "raw_value_text": raw_value,
                        "inline_string_text": inline_text,
                        "raw_formula_text": raw_formula,
                        "raw_formula_attributes": formula_attrs,
                        "formula_present": formula_node is not None,
                        "cached_text": raw_value if formula_node is not None and value_node is not None else None,
                        "cached_text_present": formula_node is not None and value_node is not None,
                        "workbook_date_system": date_system,
                        "date_serial_text": raw_value if has_numeric_source_value and style_index < len(style_dates) and style_dates[style_index] and raw_value is not None else None,
                        "normalized_value": None if formula_node is not None else parsed_value,
                        "normalized_formula": formula,
                        "normalized_cached_value": cached_value,
                        "normalized_cached_value_available": cached_available,
                    }
                )
                add_object("cell", identity, part, sheet=name, address=ref)

            for index, element in enumerate(root.iter()):
                kind = _local(element.tag)
                if kind == "mergeCell":
                    ref = element.attrib.get("ref", "")
                    add_object("merged_range", object_identity(kind, part, ref), part, sheet=name, address=ref)
                elif kind == "dataValidation":
                    ref = element.attrib.get("sqref", "")
                    add_object("data_validation", object_identity(kind, part, str(index), ref), part, sheet=name, address=ref, attributes=dict(element.attrib), formulas=[_text(child) for child in element if _local(child.tag) in {"formula1", "formula2"}])
                elif kind == "conditionalFormatting":
                    ref = element.attrib.get("sqref", "")
                    add_object("conditional_formatting", object_identity(kind, part, str(index), ref), part, sheet=name, address=ref, attributes=dict(element.attrib))
                elif kind == "hyperlink":
                    ref = element.attrib.get("ref", "")
                    add_object("hyperlink", object_identity(kind, part, str(index), ref), part, sheet=name, address=ref, attributes=dict(element.attrib))
                elif kind == "row":
                    row_num = element.attrib.get("r", "")
                    hidden = element.attrib.get("hidden") in {"1", "true"}
                    height = element.attrib.get("ht") if element.attrib.get("customHeight") in {"1", "true"} else None
                    if hidden or height is not None:
                        add_object("row_layout", object_identity(kind, part, row_num), part, sheet=name, address=row_num)
                elif kind == "col" and element.tag == MAIN + "col":
                    raw_min = element.attrib.get("min", "")
                    raw_max = element.attrib.get("max", raw_min)
                    start_index = int(raw_min)
                    end_index = int(raw_max)
                    start = get_column_letter(start_index)
                    end = get_column_letter(end_index)
                    address = f"{start}:{end}" if end_index != start_index else start
                    hidden = element.attrib.get("hidden") in {"1", "true"}
                    width = element.attrib.get("width") if element.attrib.get("customWidth") in {"1", "true"} else None
                    if hidden or width is not None:
                        add_object(
                            "column_layout",
                            object_identity(kind, part, raw_min, raw_max),
                            part,
                            sheet=name,
                            address=address,
                            min=raw_min,
                            max=raw_max,
                            width=width,
                            hidden=hidden,
                        )
                elif kind == "pane" and element.attrib.get("state") in {"frozen", "frozenSplit"}:
                    ref = element.attrib.get("topLeftCell", "")
                    add_object("freeze_panes", object_identity(kind, part, ref), part, sheet=name, address=ref)
                elif kind == "sheetProtection" and any(value in {"1", "true"} for value in element.attrib.values()):
                    add_object("sheet_protection", object_identity(kind, part), part, sheet=name)

            for table_part in table_parts:
                if table_part not in names:
                    raise SourceScanError(f"{part}: table relationship points at missing part {table_part}")
                table_root = _required_xml(zf, table_part)
                parsed_parts.add(table_part)
                name_value = table_root.attrib.get("name", "")
                add_object("table", object_identity("table", table_part), table_part, sheet=name, name=name_value, address=table_root.attrib.get("ref", ""))

            for rel in sheet_rels:
                if not rel["part"] or rel["part"] not in names:
                    continue
                if rel["type"].endswith("/comments"):
                    comments_root = _required_xml(zf, rel["part"])
                    parsed_parts.add(rel["part"])
                    for comment in comments_root.iter():
                        if _local(comment.tag) == "comment":
                            ref = comment.attrib.get("ref", "")
                            add_object("comment", object_identity("comment", rel["part"], ref), rel["part"], sheet=name, address=ref)
                elif "vmlDrawing" in rel["type"] or "vmlDrawing" in rel["target"]:
                    vml_root = _required_xml(zf, rel["part"])
                    parsed_parts.add(rel["part"])
                    parents = {child: parent for parent in vml_root.iter() for child in parent}
                    shapes = [element for element in vml_root.iter() if _local(element.tag) == "shape"]
                    client_data_nodes = [element for element in vml_root.iter() if _local(element.tag) == "ClientData"]
                    vml_object_tags = {"shape", "rect", "roundrect", "oval", "line", "polyline", "curve", "arc", "image", "group"}
                    vml_objects = [
                        element
                        for element in vml_root.iter()
                        if _local(element.tag) in vml_object_tags
                    ]
                    vml_supported = bool(shapes) and len(vml_objects) == len(shapes)
                    clients_in_shapes = 0
                    for shape in shapes:
                        shape_clients = [element for element in shape.iter() if _local(element.tag) == "ClientData"]
                        clients_in_shapes += len(shape_clients)
                        if len(shape_clients) != 1:
                            vml_supported = False
                    if clients_in_shapes != len(client_data_nodes):
                        vml_supported = False
                    for control_index, client_data in enumerate(client_data_nodes):
                        control_type = client_data.attrib.get("ObjectType", "")
                        props = {_local(child.tag): (child.text or "").strip() for child in client_data}
                        has_binding = any(props.get(key) for key in ("FmlaLink", "FmlaRange", "FmlaMacro"))
                        if not control_type or control_type in {"Note", "Pict"} or not has_binding:
                            vml_supported = False
                        if control_type in {"Note", "Pict"} or not has_binding:
                            continue
                        node = client_data
                        control_id = None
                        while node is not None:
                            if _local(node.tag) == "shape" and node.attrib.get("id"):
                                control_id = node.attrib["id"]
                                break
                            node = parents.get(node)
                        control_id = control_id or control_type
                        add_object("form_control", object_identity("form_control", rel["part"], str(control_index)), rel["part"], sheet=name, control_id=control_id, control_type=control_type, linked_cell=props.get("FmlaLink"), list_fill_range=props.get("FmlaRange"), macro=props.get("FmlaMacro"))
                    if vml_supported:
                        supported_vml_parts.add(rel["part"])

        # Workbook and sheet-scoped defined names are distinct declarations.
        defined_names = next((item for item in workbook.iter() if _local(item.tag) == "definedNames"), None)
        for index, item in enumerate(list(defined_names or [])):
            if _local(item.tag) != "definedName":
                continue
            local_sheet = item.attrib.get("localSheetId")
            if local_sheet is None:
                scope = "workbook"
            else:
                if not local_sheet.isascii() or not local_sheet.isdigit() or int(local_sheet) >= len(sheets):
                    raise SourceScanError(f"xl/workbook.xml: defined name localSheetId {local_sheet!r} is outside the worksheet range")
                scope = sheets[int(local_sheet)]["name"]
            name = item.attrib.get("name", "")
            add_object("defined_name", object_identity("defined_name", scope, name, str(index)), "xl/workbook.xml", name=name, scope=scope, formula=item.text or "")

        # Tables are package relationships even when a producer omits the sheet relation.
        seen = {obj["identity"] for obj in objects}
        for name in names:
            if name.startswith("xl/tables/") and name.endswith(".xml"):
                if any(obj["part"] == name and obj["kind"] == "table" for obj in objects):
                    continue
                root = _required_xml(zf, name)
                parsed_parts.add(name)
                identity = object_identity("table", name)
                if identity not in seen:
                    add_object("table", identity, name, name=root.attrib.get("name", ""), address=root.attrib.get("ref", ""))

        part_hashes = {}
        for info in file_infos:
            name = info.filename
            marker = next((description for token, description in OPAQUE_MARKERS.items() if token in name), None)
            parsed = name in parsed_parts
            is_parsed_vml = name.lower().endswith(".vml") and parsed
            if is_parsed_vml and name not in supported_vml_parts:
                opaque_reason = "VML drawing contains unsupported or mixed content"
            elif marker and not (is_parsed_vml and name in supported_vml_parts):
                opaque_reason = marker
            else:
                opaque_reason = "unrecognized package part" if not parsed else None
            part_hashes[name] = {
                "size": info.file_size,
                "sha256": hashlib.sha256(zf.read(name)).hexdigest(),
                "content_type": content_types.get(name, ""),
                "parsed": parsed,
                "opaque": opaque_reason is not None,
                "opaque_reason": opaque_reason,
            }
        return {
            "date_system": date_system,
            "sheets": sheets,
            "cells": cells,
            "objects": objects,
            "parts": part_hashes,
            "parsed_parts": sorted(parsed_parts),
            "scan_errors": errors,
            "shared_strings": len(shared_strings),
        }
