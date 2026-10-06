"""Static checkbox, ActiveX, and VBA handoff facts from OOXML workbooks."""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from excel_to_act.ingest.data_table import resolve_package_part
from excel_to_act.ingest.vba import VbaProject, extract_vba_project
from excel_to_act.schemas import (
    ActiveXControl,
    ActiveXEvents,
    CheckboxBinding,
    CheckboxBindings,
    VbaHandoff,
    VbaSource,
)

_DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_VBA_EXT = {"StdModule": ".bas", "ClassModule": ".cls", "Document": ".cls", "UserForm": ".frm"}
_CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)$")
_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9_.-]+")
ARTIFACT_FILES = {
    "checkbox": "checkbox_bindings.json",
    "activex": "activex_events.json",
    "vba": "vba_handoff.json",
}


def _local(name: str) -> str:
    return name.rsplit("}", 1)[-1].split(":")[-1]


def _attr(element: ET.Element, name: str) -> str | None:
    return next((value for key, value in element.attrib.items() if _local(key) == name), None)


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _sheets(zf: zipfile.ZipFile) -> list[dict[str, object]]:
    workbook = ET.fromstring(zf.read("xl/workbook.xml"))
    relationships = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {
        _attr(rel, "Id"): _attr(rel, "Target") or ""
        for rel in relationships
        if _local(rel.tag) == "Relationship"
    }
    result = []
    for sheet in workbook.iter():
        if _local(sheet.tag) != "sheet":
            continue
        rel_id = sheet.attrib.get(f"{{{_DOC_REL}}}id") or _attr(sheet, "id")
        target = targets.get(rel_id or "")
        if not target:
            continue
        part = resolve_package_part("xl", target)
        root = ET.fromstring(zf.read(part))
        sheet_pr = next((item for item in root if _local(item.tag) == "sheetPr" and _attr(item, "codeName")), None)
        result.append(
            {
                "name": sheet.attrib.get("name", ""),
                "part": part,
                "code_name": _attr(sheet_pr, "codeName") if sheet_pr is not None else None,
                "root": root,
            }
        )
    return result


def _relationships(zf: zipfile.ZipFile, part: str) -> dict[str, dict[str, str]]:
    rels_part = str(PurePosixPath(part).parent / "_rels" / f"{PurePosixPath(part).name}.rels")
    if rels_part not in zf.namelist():
        return {}
    root = ET.fromstring(zf.read(rels_part))
    result = {}
    for rel in root:
        if _local(rel.tag) != "Relationship":
            continue
        target = rel.attrib.get("Target", "")
        result[rel.attrib.get("Id", "")] = {
            "type": rel.attrib.get("Type", ""),
            "part": "" if rel.attrib.get("TargetMode") == "External" else resolve_package_part(str(PurePosixPath(part).parent), target),
        }
    return result


def _normal_shape_id(value: str | None) -> str | None:
    if not value:
        return None
    match = re.search(r"(\d+)$", value)
    return match.group(1) if match else value


def _quote_sheet(value: str) -> str:
    return value if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", value) else "'" + value.replace("'", "''") + "'"


def _parse_binding(raw: str | None, owner: str, known_sheets: dict[str, str]) -> tuple[str, str | None, str | None, str | None, str | None]:
    if raw is None or not raw.strip():
        return "unresolved", None, None, None, "no linked-cell formula is recorded"
    value = raw.strip().removeprefix("=").strip()
    if value.upper() == "#REF!":
        return "invalid", None, None, None, "linked-cell reference is #REF!"
    if any(token in value for token in ("[", "]", "(", ")", ":", ",")):
        return "dynamic", None, None, None, "linked-cell expression is external, dynamic, or multi-cell"
    sheet = owner
    address = value
    if "!" in value:
        raw_sheet, _, address = value.rpartition("!")
        sheet = raw_sheet[1:-1].replace("''", "'") if raw_sheet.startswith("'") and raw_sheet.endswith("'") else raw_sheet
        actual_sheet = known_sheets.get(sheet.casefold())
        if actual_sheet is None:
            return "unresolved", None, sheet, None, f"linked worksheet {sheet!r} is not present"
        sheet = actual_sheet
    cell = _CELL_RE.fullmatch(address.strip())
    if cell:
        try:
            row, column = coordinate_to_tuple(cell.group(1) + cell.group(2))
            if row > 1_048_576 or column > 16_384:
                raise ValueError
        except (ValueError, TypeError):
            return "invalid", None, sheet, None, "linked-cell address is outside Excel worksheet bounds"
        normalized = f"{_quote_sheet(sheet)}!{get_column_letter(column)}{row}"
        return "resolved", normalized, sheet, f"{get_column_letter(column)}{row}", None
    if re.fullmatch(r"[A-Za-z_\\][A-Za-z0-9_.\\]*", address.strip()):
        return "named", None, None, None, f"linked-cell value {address!r} is a defined name; its destination was not resolved"
    return "invalid", None, sheet, None, f"linked-cell reference {raw!r} is not a single A1 cell"


def _same_binding(left: str, right: str, owner: str, known_sheets: dict[str, str]) -> bool:
    left_result = _parse_binding(left, owner, known_sheets)
    right_result = _parse_binding(right, owner, known_sheets)
    if left_result[0] == right_result[0] == "resolved":
        return left_result[1] == right_result[1]
    return left_result[0] == right_result[0] and left.strip().casefold() == right.strip().casefold()


def identify_checkboxes(workbook_path: Path) -> CheckboxBindings:
    workbook_path = workbook_path.expanduser().resolve()
    digest = _hash_file(workbook_path)
    found: dict[tuple[str, str], dict[str, object]] = {}
    diagnostics: list[str] = []
    with zipfile.ZipFile(workbook_path) as zf:
        sheets = _sheets(zf)
        known_sheets = {str(sheet["name"]).casefold(): str(sheet["name"]) for sheet in sheets}
        for sheet in sheets:
            sheet_name, sheet_part = str(sheet["name"]), str(sheet["part"])
            rels = _relationships(zf, sheet_part)
            root = sheet["root"]
            for control in root.iter():
                if _local(control.tag) != "control":
                    continue
                shape_id = _normal_shape_id(control.attrib.get("shapeId"))
                rel_id = control.attrib.get(f"{{{_DOC_REL}}}id") or _attr(control, "id")
                relation = rels.get(rel_id or "", {})
                ctrl_part = relation.get("part", "")
                if not relation.get("type", "").endswith("/ctrlProp") or ctrl_part not in zf.namelist():
                    continue
                try:
                    props = ET.fromstring(zf.read(ctrl_part))
                except ET.ParseError:
                    diagnostics.append(f"{ctrl_part}: malformed ctrlProps XML")
                    continue
                form_props = next((node for node in props.iter() if _local(node.tag) == "formControlPr"), props)
                if (_attr(form_props, "objectType") or "").casefold() != "checkbox":
                    continue
                key = (sheet_part, shape_id or str(control.attrib.get("name", ctrl_part)))
                row = found.setdefault(key, {"sheet": sheet_name, "sheet_part": sheet_part, "shape_id": shape_id, "control_name": control.attrib.get("name"), "legacy_control_id": None, "raw": None, "sources": []})
                raw = _attr(form_props, "fmlaLink")
                if raw:
                    row["raw"] = raw
                row["sources"].append(ctrl_part)

            for rel in rels.values():
                if not (rel["type"].endswith("/vmlDrawing") or "vmlDrawing" in rel["part"]):
                    continue
                part = rel["part"]
                if part not in zf.namelist():
                    continue
                try:
                    vml = ET.fromstring(zf.read(part))
                except ET.ParseError:
                    diagnostics.append(f"{part}: malformed VML XML")
                    continue
                for client in vml.iter():
                    if _local(client.tag) != "ClientData" or (client.attrib.get("ObjectType") or "").casefold() != "checkbox":
                        continue
                    vml_parents = {child: parent for parent in vml.iter() for child in parent}
                    node = client
                    while node is not None and _local(node.tag) != "shape":
                        node = vml_parents.get(node)
                    shape_name = node.attrib.get("id") if node is not None else None
                    shape_id = _normal_shape_id(_attr(node, "spid") if node is not None else None)
                    if not shape_id and shape_name and re.fullmatch(r"_x0000_s\d+", shape_name, re.IGNORECASE):
                        shape_id = shape_name.rsplit("s", 1)[-1]
                    key = (sheet_part, shape_id or f"vml:{shape_name or len(found)}")
                    row = found.setdefault(key, {"sheet": sheet_name, "sheet_part": sheet_part, "shape_id": shape_id, "control_name": shape_name, "legacy_control_id": shape_name, "raw": None, "sources": []})
                    row["legacy_control_id"] = shape_name
                    if not row.get("control_name"):
                        row["control_name"] = shape_name
                    raw = next(((child.text or "").strip() for child in client if _local(child.tag) == "FmlaLink"), None)
                    if raw:
                        existing = row.get("raw")
                        if existing and not _same_binding(str(existing), raw, sheet_name, known_sheets):
                            diagnostics.append(f"{sheet_name} shape {shape_id}: modern and VML linked-cell values differ")
                            row["conflict"] = True
                        else:
                            row["raw"] = raw
                    row["sources"].append(part)

    bindings = []
    for row in found.values():
        raw = row.get("raw")
        status, linked, linked_sheet, address, diagnostic = _parse_binding(raw if isinstance(raw, str) else None, str(row["sheet"]), known_sheets)
        if row.get("conflict"):
            status, linked, linked_sheet, address = "invalid", None, None, None
            diagnostic = "modern and VML linked-cell values differ"
        bindings.append(
            CheckboxBinding(
                sheet=str(row["sheet"]),
                sheet_part=str(row["sheet_part"]),
                shape_id=str(row["shape_id"]) if row.get("shape_id") else None,
                control_name=str(row["control_name"]) if row.get("control_name") else None,
                legacy_control_id=str(row["legacy_control_id"]) if row.get("legacy_control_id") else None,
                linked_cell_raw=str(raw) if raw is not None else None,
                linked_cell=linked,
                linked_sheet=linked_sheet,
                linked_address=address,
                binding_status=status,
                sources=sorted(set(str(source) for source in row["sources"])),
                diagnostic=diagnostic,
            )
        )
    bindings.sort(key=lambda item: (item.sheet_part, f"{int(item.shape_id):012d}" if item.shape_id and item.shape_id.isdigit() else item.shape_id or ""))
    bad = any(item.binding_status != "resolved" for item in bindings)
    return CheckboxBindings(workbook_sha256=digest, status="partial" if bad else "complete" if bindings else "not_applicable", bindings=bindings, diagnostics=diagnostics)


def _active_x_part(zf: zipfile.ZipFile, part: str) -> tuple[str | None, str | None]:
    try:
        root = ET.fromstring(zf.read(part))
        class_id = _attr(root, "classid")
    except (KeyError, ET.ParseError):
        return None, None
    rels = _relationships(zf, part)
    binary = next((rel["part"] for rel in rels.values() if "activeXControlBinary" in rel["type"]), None)
    return class_id, binary


def identify_activex(workbook_path: Path, project: VbaProject | None = None) -> ActiveXEvents:
    workbook_path = workbook_path.expanduser().resolve()
    digest = _hash_file(workbook_path)
    project = project if project is not None else extract_vba_project(workbook_path)
    modules = {module.name.casefold(): module for module in project.modules}
    controls: dict[tuple[str, str], ActiveXControl] = {}
    diagnostics: list[str] = []
    with zipfile.ZipFile(workbook_path) as zf:
        for sheet in _sheets(zf):
            sheet_name, sheet_part = str(sheet["name"]), str(sheet["part"])
            rels = _relationships(zf, sheet_part)
            for control in sheet["root"].iter():
                if _local(control.tag) != "control":
                    continue
                rel_id = control.attrib.get(f"{{{_DOC_REL}}}id") or _attr(control, "id")
                relation = rels.get(rel_id or "", {})
                part = relation.get("part", "")
                if not part.startswith("xl/activeX/") or not part.endswith(".xml") or part not in zf.namelist():
                    continue
                shape_id = _normal_shape_id(control.attrib.get("shapeId"))
                name = control.attrib.get("name", "")
                key = (sheet_part, shape_id or name or part)
                if key in controls:
                    continue
                class_id, binary_part = _active_x_part(zf, part)
                code_name = sheet.get("code_name")
                diagnostic = None
                event_procedures: list[str] = []
                state = "unresolved"
                if project.oletools_missing:
                    state, diagnostic = "unavailable", "oletools is unavailable; VBA event declarations were not extracted"
                elif project.error:
                    state, diagnostic = "unavailable", f"VBA extraction failed: {project.error}"
                elif not code_name:
                    diagnostic = "worksheet codeName is missing; its VBA document module cannot be selected"
                elif code_name.casefold() not in modules:
                    diagnostic = f"VBA document module {code_name!r} was not extracted"
                else:
                    module = modules[code_name.casefold()]
                    prefix = f"{name}_".casefold()
                    event_procedures = [procedure for procedure in module.procedures if procedure.casefold().startswith(prefix)]
                    if event_procedures:
                        state = "resolved"
                    else:
                        diagnostic = f"no event procedure for control {name!r} was declared in module {code_name!r}"
                if not binary_part:
                    diagnostics.append(f"{part}: ActiveX binary relationship is missing")
                controls[key] = ActiveXControl(
                    sheet=sheet_name,
                    sheet_part=sheet_part,
                    sheet_code_name=str(code_name) if code_name else None,
                    shape_id=shape_id or "",
                    control_name=name,
                    class_id=class_id,
                    part=part,
                    binary_part=binary_part,
                    event_procedures=event_procedures,
                    binding_status=state,
                    diagnostic=diagnostic,
                )
                if diagnostic:
                    diagnostics.append(f"{sheet_name} {name}: {diagnostic}")
    rows = sorted(controls.values(), key=lambda item: (item.sheet_part, f"{int(item.shape_id):012d}" if item.shape_id.isdigit() else item.shape_id))
    if not rows:
        status = "not_applicable"
    elif project.oletools_missing or project.error:
        status = "unavailable"
    elif any(item.binding_status != "resolved" for item in rows) or diagnostics:
        status = "partial"
    else:
        status = "complete"
    return ActiveXEvents(workbook_sha256=digest, status=status, controls=rows, diagnostics=diagnostics)


def build_vba_handoff(workbook_path: Path, project: VbaProject | None = None) -> tuple[VbaHandoff, dict[str, bytes]]:
    workbook_path = workbook_path.expanduser().resolve()
    digest = _hash_file(workbook_path)
    project = project if project is not None else extract_vba_project(workbook_path)
    files: dict[str, bytes] = {}
    modules: list[VbaSource] = []
    for index, module in enumerate(project.modules, start=1):
        extension = _VBA_EXT.get(module.kind, ".cls")
        safe_name = _SAFE_NAME_RE.sub("_", module.name).strip("._") or "module"
        source_file = f"vba_sources/{index:03d}-{safe_name}{extension}"
        content = module.code.encode("utf-8")
        files[source_file] = content
        modules.append(VbaSource(name=module.name, kind=module.kind, procedures=module.procedures, source_file=source_file, sha256=hashlib.sha256(content).hexdigest()))
    if not project.available:
        status, diagnostics = "not_applicable", []
    elif project.oletools_missing:
        status, diagnostics = "unavailable", ["oletools is unavailable; VBA source and module inventory were not extracted"]
    elif project.error:
        status, diagnostics = "unavailable", [f"VBA extraction failed: {project.error}"]
    else:
        status, diagnostics = "complete", []
    return VbaHandoff(workbook_sha256=digest, available=project.available, status=status, modules=modules, diagnostics=diagnostics), files


def build_control_artifacts(workbook_path: Path) -> tuple[CheckboxBindings, ActiveXEvents, VbaHandoff, dict[str, bytes]]:
    project = extract_vba_project(workbook_path)
    vba, sources = build_vba_handoff(workbook_path, project)
    return identify_checkboxes(workbook_path), identify_activex(workbook_path, project), vba, sources


def build_single_control_artifact(kind: str, workbook_path: Path):
    if kind == "checkbox":
        return identify_checkboxes(workbook_path), {}
    if kind == "activex":
        return identify_activex(workbook_path), {}
    if kind == "vba":
        return build_vba_handoff(workbook_path)
    raise ValueError(f"unknown control artifact kind: {kind}")


def clear_declared_vba_sources(output_dir: Path, manifest: VbaHandoff | None) -> None:
    root = (output_dir / "vba_sources").resolve()
    if manifest is None:
        return
    for module in manifest.modules:
        relative = PurePosixPath(module.source_file)
        if relative.is_absolute() or relative.parts[:1] != ("vba_sources",) or ".." in relative.parts:
            continue
        path = output_dir.joinpath(*relative.parts)
        try:
            path.resolve().relative_to(root)
        except ValueError:
            continue
        if path.is_file() or path.is_symlink():
            path.unlink()


def enrich_form_controls(inventory: object, checkboxes: CheckboxBindings) -> None:
    """Add modern checkbox links to existing VML inventory records without changing its count."""
    by_vml = {
        (binding.sheet, part, binding.legacy_control_id): binding
        for binding in checkboxes.bindings
        for part in binding.sources
        if part.endswith(".vml")
    }
    for sheet in getattr(inventory, "sheets", []):
        for record in sheet.ranges:
            if record.kind != "form_control":
                continue
            key = (sheet.name, record.source_location.ooxml_part, record.metadata.get("control_id"))
            binding = by_vml.get(key)
            if binding is None:
                continue
            record.metadata.update(
                linked_cell=binding.linked_cell or binding.linked_cell_raw,
                linked_sheet=binding.linked_sheet,
                linked_address=binding.linked_address,
                binding_status=binding.binding_status,
                binding_diagnostic=binding.diagnostic,
                shape_id=binding.shape_id,
                binding_sources=binding.sources,
            )


def controls_handoff_markdown(checkboxes: CheckboxBindings, activex: ActiveXEvents, vba: VbaHandoff) -> str:
    lines = [
        "# Static control and VBA handoff",
        "",
        "This is a source-derived inventory. ActiveX binary streams remain opaque; exported VBA is readable source, not executed or translated behavior.",
        "",
        "## Checkbox bindings",
        "",
        f"**Status:** {checkboxes.status} · **Controls:** {len(checkboxes.bindings)} · **Resolved A1 links:** {sum(row.binding_status == 'resolved' for row in checkboxes.bindings)} · **Invalid source links:** {sum(row.binding_status == 'invalid' for row in checkboxes.bindings)}",
        "",
        "Invalid links such as `#REF!` are source facts and were not guessed or repaired.",
        "",
        "| Sheet | Shape | Control | Linked cell | Status | Sources |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    lines.extend(
        f"| {row.sheet} | {row.shape_id or '—'} | {row.control_name or '—'} | {row.linked_cell or row.linked_cell_raw or '—'} | {row.binding_status} | {', '.join(row.sources)} |"
        for row in checkboxes.bindings
    )
    lines.extend(["", "## ActiveX event declarations", "", f"**Status:** {activex.status} · **Controls:** {len(activex.controls)}", "", "| Sheet | CodeName | Control | Class ID | Binary part | Declared handlers | Status |", "| --- | --- | --- | --- | --- | --- | --- |"])
    lines.extend(
        f"| {row.sheet} | {row.sheet_code_name or '—'} | {row.control_name} | {row.class_id or '—'} | {row.binary_part or '—'} | {', '.join(row.event_procedures) or '—'} | {row.binding_status}{(': ' + row.diagnostic) if row.diagnostic else ''} |"
        for row in activex.controls
    )
    event_sources = [module for module in vba.modules if module.procedures]
    lines.extend(["", "## VBA event-source modules", "", f"**Status:** {vba.status} · **Modules with declared procedures:** {len(event_sources)} · **Total modules:** {len(vba.modules)}", ""])
    if vba.diagnostics:
        lines.extend(f"- {diagnostic}" for diagnostic in vba.diagnostics)
    lines.extend(["| Module | Kind | Procedures | Source | SHA-256 |", "| --- | --- | --- | --- | --- |"])
    lines.extend(f"| {module.name} | {module.kind} | {', '.join(module.procedures)} | [{module.source_file}]({module.source_file}) | `{module.sha256}` |" for module in event_sources)
    lines.extend(["", "## Complete VBA source index", "", "| Module | Kind | Procedures | Source | SHA-256 |", "| --- | --- | --- | --- | --- |"])
    lines.extend(f"| {module.name} | {module.kind} | {', '.join(module.procedures) or '—'} | [{module.source_file}]({module.source_file}) | `{module.sha256}` |" for module in vba.modules)
    lines.extend(["", "## Files", "", "- `checkbox_bindings.json` — typed checkbox records.", "- `activex_events.json` — worksheet controls and declared event handlers.", "- `vba_handoff.json` — module index and source checksums.", ""])
    return "\n".join(lines)


def vba_handoff_markdown(vba: VbaHandoff) -> str:
    lines = ["# VBA source handoff", "", f"**Status:** {vba.status} · **Modules:** {len(vba.modules)}", "", "Exported text is source handoff only; macros were not executed or translated.", ""]
    lines.extend(f"- {module.name} ({module.kind}): {', '.join(module.procedures) or 'no declared procedures'} · [{module.source_file}]({module.source_file}) · SHA-256 `{module.sha256}`" for module in vba.modules)
    lines.extend(["", *[f"- {diagnostic}" for diagnostic in vba.diagnostics], ""])
    return "\n".join(lines)


def evaluate_control_output(kind: str, workbook_path: Path, output_dir: Path) -> dict[str, object]:
    expected, sources = build_single_control_artifact(kind, workbook_path)
    model = {"checkbox": CheckboxBindings, "activex": ActiveXEvents, "vba": VbaHandoff}[kind]
    filename = ARTIFACT_FILES[kind]
    issues = []
    artifact_path = output_dir.expanduser().resolve() / filename
    if not artifact_path.is_file():
        issues.append(f"missing {filename}")
    else:
        try:
            actual = model.model_validate_json(artifact_path.read_bytes())
            if actual.model_dump(mode="json") != expected.model_dump(mode="json"):
                issues.append(f"{filename} differs from the fresh source")
        except (OSError, ValueError) as exc:
            issues.append(f"{filename} is invalid: {exc}")
    if kind == "vba":
        for relative, content in sources.items():
            path = output_dir.expanduser().resolve() / relative
            if not path.is_file() or path.read_bytes() != content:
                issues.append(f"missing or changed {relative}")
        handoff_md = output_dir.expanduser().resolve() / "vba_handoff.md"
        if not handoff_md.is_file() or handoff_md.read_bytes() != vba_handoff_markdown(expected).encode("utf-8"):
            issues.append("missing or changed vba_handoff.md")
    status = "fail" if issues else expected.status
    if status in {"complete"}:
        status = "pass"
    return {"tool": f"step1.{kind}.evaluate", "status": status, "matches_source": not issues, "workbook": str(workbook_path.expanduser().resolve()), "output": str(output_dir.expanduser().resolve()), "artifact": filename, "records": len(getattr(expected, "bindings", getattr(expected, "controls", getattr(expected, "modules", [])))), "issues": issues, "diagnostics": expected.diagnostics}
