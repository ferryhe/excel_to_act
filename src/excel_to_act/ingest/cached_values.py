"""Cached formula results stored inside the workbook.

Excel writes the last calculated result next to every formula. Reading those
results is the only zero-dependency source of "what the workbook actually
evaluates to", which makes it the primary oracle for later reconciliation.

Workbooks that were never recalculated by Excel (for example files written by
openpyxl) contain empty, untyped ``<v/>`` elements, so openpyxl reports
``None``. A formula cell with ``t="str"`` and an empty ``<v/>`` stores the
empty-string result and must remain distinguishable from a missing cache.
"""

from __future__ import annotations

import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from excel_to_act.ingest.data_table import sheet_xml_parts

CellValue = str | int | float | bool | None

# Coordinates are keyed by sheet title because sheet-scoped addresses are the
# only stable identity openpyxl exposes for both the formula and value passes.
CachedValueMap = dict[tuple[str, str], CellValue]


def _safe_value(value: Any) -> CellValue:
    """Coerce a raw openpyxl value into the scalar union allowed by contracts."""

    if value is None or isinstance(value, str | int | float | bool):
        return value
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value)


def _empty_string_formula_caches(workbook_path: Path) -> set[tuple[str, str]]:
    """Find formula cells whose stored string result is present but empty."""

    empty: set[tuple[str, str]] = set()
    parts = sheet_xml_parts(workbook_path)
    if not parts:
        return empty
    with zipfile.ZipFile(workbook_path) as zf:
        for sheet, part in parts.items():
            try:
                root = ET.fromstring(zf.read(part))
            except (KeyError, ET.ParseError):
                continue
            for cell in root.iter():
                if cell.tag.rsplit("}", 1)[-1] != "c" or cell.attrib.get("t") != "str":
                    continue
                formula = next((child for child in cell if child.tag.rsplit("}", 1)[-1] == "f"), None)
                value = next((child for child in cell if child.tag.rsplit("}", 1)[-1] == "v"), None)
                if formula is not None and value is not None and value.text in (None, ""):
                    empty.add((sheet, cell.attrib["r"]))
    return empty


def read_cached_values(workbook_path: Path) -> CachedValueMap:
    """Return ``{(sheet_title, coordinate): cached value}`` for every stored result.

    Cells without a stored result are absent from the map, so callers must use
    ``key in map`` rather than relying on a ``None`` value.
    """

    workbook_path = workbook_path.expanduser().resolve()
    values: CachedValueMap = {}
    empty_string_caches = _empty_string_formula_caches(workbook_path)
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            # In read-only mode the iteration bounds come from the declared
            # <dimension>, which non-Excel writers can omit or understate.
            reset = getattr(ws, "reset_dimensions", None)
            if callable(reset):
                reset()
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    key = (ws.title, cell.coordinate)
                    values[key] = _safe_value(cell.value)
    finally:
        wb.close()
    values.update({key: "" for key in empty_string_caches})
    return values
